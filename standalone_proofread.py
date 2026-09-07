from __future__ import annotations

import argparse
import csv
import hashlib
import json
import tempfile
import zipfile
import os
import re
import shutil
import subprocess
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Mapping
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from check_pronunciation_candidates import (
    VERSION as EXPECTED_RESOLVER_VERSION,
    analyze,
    norm_bopomofo,
    EXPECTED_RESOLVER_SOURCE_FILES,
    DEFAULT_DICT,
    DEFAULT_RULES,
    DEFAULT_REGRESSIONS,
    DEFAULT_CHAR_OVERRIDES,
    DEFAULT_CONTEXT_OVERRIDES,
)
from export_zhuyin_readings import (
    VERSION as ACTUAL_DECODER_VERSION,
    decode, DEFAULT_MAP, DEFAULT_GROUPS, DEFAULT_CFF_MAP, DEFAULT_CFF_CONSENSUS,
    DEFAULT_XREF_OVERRIDES, DEFAULT_TRANSFORMS, DEFAULT_FINGERPRINTS,
    DEFAULT_OUTLINE_SIGNATURES, DEFAULT_MAPPING_CORRECTIONS, DEFAULT_SYMBOL_TEMPLATES,
    DEFAULT_ACTUAL_OVERRIDES, DEFAULT_STRUCTURAL_EXCLUSIONS,
)
from occurrence_ledger import (
    ALL_STATES,
    COMPARABLE_TERMINAL_STATES,
    CONFIRMATION_GATES,
    EXCLUDED_STATES,
    LEDGER_SCHEMA_VERSION,
    NON_TERMINAL_STATES,
    PIPELINE_BLOCKED,
    PROCESSING_FINISHED,
    PROOFREAD_COMPLETE,
    REVIEW_ID_SCHEMA_VERSION,
    SESSION_SCHEMA_VERSION,
    TERMINAL_STATES,
    WORKBOOK_SCHEMA_VERSION,
    DuplicateIdError,
    InvalidTransitionError,
    assert_unique_ids,
    candidate_payload_sha256,
    canonical_bopomofo,
    completion_gate,
    make_review_id as ledger_make_review_id,
    normalize_expected_set,
    reconcile_ledger,
    transition_state,
    validate_occurrence_ledger,
    validate_terminal_state,
    infer_actual_status,
    infer_expected_status,
    derive_comparison_result,
    derive_authoritative_state,
    refresh_derived_state,
)
from runtime_source_validation import (
    SourceValidationError,
    compute_actual_asset_fingerprint,
    compute_expected_asset_fingerprint,
    rewrite_actual_workbook_fingerprint,
    validate_asset_manifest,
    write_pipeline_blocked,
)
from cross_version_compat import (
    fingerprint_compatible,
    metadata_compatibility_warnings,
    review_id_schema_compatible,
    schema_compatible,
)
from runtime_regression_gate import load_mandatory_cases, validate_regression_execution
from actual_review import (
    ACTUAL_REVIEW_SCHEMA_VERSION,
    OCCURRENCE_OVERRIDE_FILE, USER_GLYF_FILE, USER_CFF_FILE, GLYPH_CONFLICT_FILE, GLYPH_PROVENANCE_FILE,
    apply_staged_manual_actual_batch,
    build_actual_review_groups,
    build_actual_group_for_entry,
    apply_verified_actual_group,
    export_actual_review_package,
    import_actual_review_workbook,
    ensure_user_evidence_files,
    actual_workbook_dynamic_dependencies,
    actual_workbook_global_exact_dependencies,
    load_manual_actual_staging,
    stage_manual_actual_group,
)
from global_exact_glyph_library import (
    GlobalExactGlyphRepository,
    global_exact_glyph_evidence_hashes,
)

PROGRAM = "注音校對工具－單機端到端控制器"
VERSION = "5.7.0"
# Tool release number is not a cache/session boundary. Compatibility is decided
# by schema family, durable review IDs, PDF bytes and evidence assets.
COMPATIBLE_SESSION_VERSIONS = None

UNRESOLVED_SHEETS = ["差異候選", "低信心詞境候選", "多音仍待詞條判讀", "現有注音未解碼", "規則衝突"]
COVERAGE_GAP_SHEETS = ["未建立預期音"]
# Historical decision memory may only auto-close a case when the current run has
# independently produced an actual-vs-expected candidate. It cannot rescue
# multi-pronunciation uncertainty, decoder abstention, or rule conflict.
MEMORY_AUTO_SHEETS = {"差異候選", "低信心詞境候選"}
AUTO_ERROR_SHEETS = ["確認錯誤"]
AUTO_CORRECT_SHEETS = ["確認正確回歸"]
SPECIAL_SHEETS = ["專名特殊項目略過"]
DECISIONS = {"補建expected證據", "解決expected證據", "確認現版差異", "確認非校對範圍", "保留待人工"}
GPT_DECISION_BUNDLE_SCHEMA_VERSION = "1.1"
PROJECT_EVIDENCE_DIR = "_專案證據"
PROJECT_ACTUAL_EVIDENCE_DIR = "actual"
NO_ACTUAL_PENDING_MESSAGE = (
    "目前沒有 ACTUAL_DECODE_ERROR／ACTUAL_UNRESOLVED 可匯出；actual 待判定為 0，無須建立 GPT 判定包。"
)


class ManualActualPostApplyError(RuntimeError):
    """Authoritative actual committed, but a later project step did not finish."""

    def __init__(
        self,
        phase: str,
        cause: Exception,
        *,
        batch_result: Mapping[str, Any],
        cleared_event_count: int | None,
    ) -> None:
        self.phase = str(phase)
        self.cause = cause
        self.batch_result = dict(batch_result)
        self.cleared_event_count = cleared_event_count
        if self.phase == "clear_actual_dependent_events":
            recovery = (
                "請先排除人工判定資料庫的錯誤，再清除依賴舊 actual 的 confirmation；"
                "不要重新套用已 acknowledgement 的 staging。"
            )
        else:
            recovery = (
                "可在排除錯誤後使用既有 --refresh-actual／refresh_actual_project() 恢復；"
                "不要重新套用已 acknowledgement 的 staging。"
            )
        super().__init__(
            "actual evidence 已正式套用，但後續 project refresh／驗證未完成；"
            f"phase={self.phase}；{type(cause).__name__}: {cause}。{recovery}"
        )


def project_actual_evidence_root(output_dir: Path) -> Path:
    """Return the project-owned dynamic actual evidence directory.

    v5.5.0 intentionally separates immutable program assets from user/project
    evidence.  A program upgrade must not erase visual actual confirmations,
    and one textbook project must not silently affect another.
    """
    return Path(output_dir).resolve() / PROJECT_EVIDENCE_DIR / PROJECT_ACTUAL_EVIDENCE_DIR


def initialize_project_actual_evidence(output_dir: Path, program_root: Path | None = None) -> Path:
    target = project_actual_evidence_root(output_dir)
    target.mkdir(parents=True, exist_ok=True)
    source_root = Path(program_root).resolve() if program_root else Path(__file__).resolve().parent
    # One-time seed for backward compatibility.  Existing installation evidence
    # is copied, never linked.  Afterwards the project copy is authoritative.
    for name in (OCCURRENCE_OVERRIDE_FILE, USER_GLYF_FILE, USER_CFF_FILE, GLYPH_CONFLICT_FILE, GLYPH_PROVENANCE_FILE):
        dst = target / name
        src = source_root / name
        if not dst.exists() and src.exists():
            shutil.copy2(src, dst)
    ensure_user_evidence_files(target)
    return target

ACTUAL_DECODER_SOURCE_FILES = [
    "export_zhuyin_readings.py",
    "export_pdf_text_diagnostics.py",
    "cff_zhuyin_decoder.py",
    "cff_unseen_family_bootstrap.py",
    "ttf_zhuyin_shape_decoder.py",
    "ttf_symbol_recombination.py",
    "cff_zero_map_batch.py",
    "occurrence_ledger.py",
    "actual_review.py",
    "exact_glyph_identity.py",
    "global_exact_glyph_library.py",
]

DECISION_SEED = Path(__file__).with_name("校對判定記憶.json")
USER_MEMORY = Path(__file__).with_name("使用者判定記憶.json")
REUSABLE_EXPECTED_RULES = Path(__file__).with_name("可重用expected規則.json")
REUSABLE_RULE_SCHEMA_VERSION = "1.0"



def _canonical_rule_doc(doc: Mapping[str, Any] | None = None) -> dict[str, Any]:
    raw = dict(doc or {})
    rules = raw.get("rules") if isinstance(raw.get("rules"), list) else []
    out_rules: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    seen_match_keys: dict[tuple[Any, ...], tuple[str, ...]] = {}
    for index, item in enumerate(rules, 1):
        if not isinstance(item, dict):
            raise ValueError(f"可重用 expected 規則第 {index} 筆不是物件")
        enabled = item.get("enabled", True) is not False
        rule_id = str(item.get("rule_id") or "").strip()
        phrase = str(item.get("phrase") or "").strip()
        target_char = str(item.get("target_char") or "").strip()
        evidence = str(item.get("evidence") or "").strip()
        context_contains = str(item.get("context_contains") or "").strip()
        expected_set = list(normalize_expected_set(item.get("expected_set")))
        try:
            target_index = int(item.get("target_index"))
        except (TypeError, ValueError):
            raise ValueError(f"可重用 expected 規則第 {index} 筆 target_index 無效")
        if not rule_id:
            raise ValueError(f"可重用 expected 規則第 {index} 筆缺少 rule_id")
        if rule_id in seen_ids:
            raise ValueError(f"可重用 expected 規則 rule_id 重複：{rule_id}")
        seen_ids.add(rule_id)
        if not phrase or not target_char or not expected_set or not evidence:
            raise ValueError(f"可重用 expected 規則 {rule_id} 缺少 phrase／target_char／expected_set／evidence")
        if target_index < 0 or target_index >= len(phrase) or phrase[target_index] != target_char:
            raise ValueError(f"可重用 expected 規則 {rule_id} 的 target_index 不對應目標字")
        match_key = (phrase, target_index, target_char, context_contains)
        expected_tuple = tuple(expected_set)
        if match_key in seen_match_keys and seen_match_keys[match_key] != expected_tuple:
            raise ValueError(f"可重用 expected 規則同一條件出現衝突讀音：{match_key}")
        seen_match_keys[match_key] = expected_tuple
        out_rules.append({
            "rule_id": rule_id,
            "enabled": enabled,
            "phrase": phrase,
            "target_char": target_char,
            "target_index": target_index,
            "expected_set": expected_set,
            "evidence": evidence,
            "context_contains": context_contains,
            "scope": str(item.get("scope") or "same_phrase_position").strip(),
            "source": str(item.get("source") or "人工核准規則").strip(),
            "note": str(item.get("note") or "").strip(),
            "approved_at": str(item.get("approved_at") or "").strip(),
        })
    return {"schema_version": REUSABLE_RULE_SCHEMA_VERSION, "rules": out_rules}


def reusable_rules_sha256(doc: Mapping[str, Any]) -> str:
    canonical = _canonical_rule_doc(doc)
    raw = json.dumps(canonical, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def load_reusable_expected_rules(path: Path = REUSABLE_EXPECTED_RULES) -> dict[str, Any]:
    if not path.exists():
        return _canonical_rule_doc({})
    raw = json_load_strict(path)
    if not isinstance(raw, dict):
        raise ValueError("可重用 expected 規則檔根節點必須是物件")
    if str(raw.get("schema_version") or REUSABLE_RULE_SCHEMA_VERSION) != REUSABLE_RULE_SCHEMA_VERSION:
        raise ValueError("可重用 expected 規則 schema 不相容")
    return _canonical_rule_doc(raw)


def _find_phrase_target_index(entry: Mapping[str, Any], phrase: str) -> int:
    source = entry.get("source_record") or {}
    line = str(source.get("所在行") or "")
    try:
        line_pos = int(source.get("行內字元位置"))
    except (TypeError, ValueError):
        line_pos = -1
    target = str(entry.get("char") or "")
    hits: list[int] = []
    start = 0
    while phrase and line:
        pos = line.find(phrase, start)
        if pos < 0:
            break
        rel = line_pos - pos
        if 0 <= rel < len(phrase) and phrase[rel] == target:
            hits.append(rel)
        start = pos + 1
    if len(set(hits)) == 1:
        return hits[0]
    # Safe fallback only when the phrase itself contains the target exactly once.
    indexes = [i for i, ch in enumerate(phrase) if ch == target]
    if len(indexes) == 1:
        return indexes[0]
    raise ValueError("無法唯一確定可重用規則中的目標字位置；請縮短／調整完整詞")


def build_reusable_rule(entry: Mapping[str, Any], *, phrase: str, expected_set: Any, evidence: str, note: str = "", context_contains: str = "") -> dict[str, Any]:
    phrase = str(phrase or "").strip()
    evidence = str(evidence or "").strip()
    context_contains = str(context_contains or "").strip()
    readings = list(normalize_expected_set(expected_set))
    if not phrase or not evidence or not readings:
        raise ValueError("升格規則需要完整詞、expected 與正式來源")
    target_index = _find_phrase_target_index(entry, phrase)
    target_char = str(entry.get("char") or "").strip()
    payload = {
        "phrase": phrase,
        "target_char": target_char,
        "target_index": target_index,
        "expected_set": readings,
        "evidence": evidence,
        "context_contains": context_contains,
        "scope": "same_phrase_position",
        "source": "人工核准可重用 expected 規則",
        "note": str(note or "").strip(),
        "approved_at": datetime.now().isoformat(timespec="seconds"),
        "enabled": True,
    }
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    payload["rule_id"] = "USER-EXPECTED-" + hashlib.sha256(raw).hexdigest()[:20]
    return _canonical_rule_doc({"rules": [payload]})["rules"][0]


def save_reusable_expected_rule(rule: Mapping[str, Any], path: Path = REUSABLE_EXPECTED_RULES) -> dict[str, Any]:
    current = load_reusable_expected_rules(path)
    rules = list(current.get("rules") or [])
    new_rule = dict(rule)
    # Idempotent save for the exact same approved condition/readings.
    for existing in rules:
        key_a = (existing.get("phrase"), existing.get("target_index"), existing.get("target_char"), existing.get("context_contains"), tuple(existing.get("expected_set") or []))
        key_b = (new_rule.get("phrase"), new_rule.get("target_index"), new_rule.get("target_char"), new_rule.get("context_contains"), tuple(new_rule.get("expected_set") or []))
        if key_a == key_b:
            return existing
    rules.append(new_rule)
    doc = _canonical_rule_doc({"rules": rules})
    json_save(path, doc)
    return new_rule


def _matching_reusable_rules(entry: Mapping[str, Any], rules: list[Mapping[str, Any]]) -> list[dict[str, Any]]:
    source = entry.get("source_record") or {}
    line = str(source.get("所在行") or "")
    local = str(source.get("局部詞境") or "")
    target = str(entry.get("char") or "")
    try:
        line_pos = int(source.get("行內字元位置"))
    except (TypeError, ValueError):
        line_pos = -1
    matches: list[dict[str, Any]] = []
    for raw in rules:
        rule = dict(raw)
        if rule.get("enabled", True) is False or str(rule.get("target_char") or "") != target:
            continue
        phrase = str(rule.get("phrase") or "")
        context_contains = str(rule.get("context_contains") or "")
        if context_contains and context_contains not in line and context_contains not in local:
            continue
        try:
            target_index = int(rule.get("target_index"))
        except (TypeError, ValueError):
            continue
        matched = False
        start = 0
        while phrase and line:
            pos = line.find(phrase, start)
            if pos < 0:
                break
            if line_pos >= 0 and pos + target_index == line_pos:
                matched = True
                break
            start = pos + 1
        if not matched and line_pos < 0 and phrase in local:
            matched = True
        if matched:
            matches.append(rule)
    return matches


def _apply_reusable_rules(ledger: list[dict[str, Any]], rules_doc: Mapping[str, Any]) -> list[dict[str, Any]]:
    """Apply approved expected rules without consulting actual as rule input.

    v5.5.0 keeps reusable expected evidence alive even when the actual lane is
    unresolved.  The public workflow state is recomputed from the two lanes only
    after the expected update is complete.
    """
    rules = [dict(rule) for rule in (rules_doc.get("rules") or [])]
    if not rules:
        return [refresh_derived_state(entry) for entry in ledger]
    out: list[dict[str, Any]] = []
    skip_states = {"EXCLUDED_NONINDEPENDENT_LAYER", "EXCLUDED_OUT_OF_SCOPE", "SOURCE_INVALID", "DATA_INTEGRITY_ERROR"}
    for original in ledger:
        entry = refresh_derived_state(dict(original))
        if entry.get("state") in skip_states:
            out.append(entry)
            continue
        hits = _matching_reusable_rules(entry, rules)
        if not hits:
            out.append(entry)
            continue
        hits.sort(key=lambda r: (bool(r.get("context_contains")), len(str(r.get("context_contains") or ""))), reverse=True)
        best_score = (bool(hits[0].get("context_contains")), len(str(hits[0].get("context_contains") or "")))
        best = [r for r in hits if (bool(r.get("context_contains")), len(str(r.get("context_contains") or ""))) == best_score]
        reading_sets = {tuple(normalize_expected_set(r.get("expected_set"))) for r in best}
        if len(reading_sets) != 1:
            entry.update({
                "expected_status": "CONFLICT",
                "expected_set": [],
                "expected_evidence": "人工核准可重用規則彼此衝突",
                "context_evidence": str((entry.get("source_record") or {}).get("所在行") or entry.get("context_evidence") or entry.get("char") or ""),
                "source_view": "規則衝突",
                "note": "｜".join(str(r.get("rule_id") or "") for r in best),
            })
            out.append(refresh_derived_state(entry, preserve_confirmed=False))
            continue
        rule = best[0]
        expected = list(normalize_expected_set(rule.get("expected_set")))
        entry.update({
            "previous_expected_set": list(entry.get("expected_set") or []),
            "previous_expected_evidence": str(entry.get("expected_evidence") or ""),
            "expected_set": expected,
            "expected_evidence": f"{rule.get('source') or '人工核准規則'}｜{rule.get('evidence')}",
            "expected_status": "RESOLVED",
            "context_evidence": f"完整詞：{rule.get('phrase')}；目標位置：{rule.get('target_index')}；{rule.get('context_contains') or ''}".strip("；"),
            "source_view": "人工核准可重用規則",
            "reusable_rule_id": rule.get("rule_id"),
        })
        out.append(refresh_derived_state(entry, preserve_confirmed=False))
    return out


def configure_utf8_stdio() -> None:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


def utf8_env() -> dict[str, str]:
    env = os.environ.copy()
    env["PYTHONUTF8"] = "1"
    env["PYTHONIOENCODING"] = "utf-8"
    return env


def _norm_context(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value)).replace("，", ",").replace("。", ".").replace("？", "?").replace("！", "!")


def semantic_review_key(rec: dict[str, Any]) -> str:
    r = rec.get("row", {})
    try:
        offpage = "offpage" if float(r.get("x0") or 0) < 0 or float(r.get("y0") or 0) < 0 else "visible"
    except Exception:
        offpage = "visible"
    core = [
        rec.get("source_sheet", ""), _norm_context(r.get("字元")),
        _norm_context(r.get("所在行")), _norm_context(r.get("局部詞境")),
        _norm_context(r.get("實際注音")),
        _norm_context(r.get("預期注音") or r.get("字典有效讀音") or r.get("規則候選讀音")),
        _norm_context(r.get("規則ID")), _norm_context(r.get("規則模式")),
        _norm_context(r.get("規則來源")), _norm_context(r.get("預期注音依據")),
        offpage,
    ]
    raw = json.dumps(core, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    return hashlib.sha1(raw).hexdigest()




def semantic_review_key_v1(rec: dict[str, Any]) -> str:
    """Backward-compatible key used by the pre-v5.1.3 built-in memories.

    It is still exact on source sheet, character, line/context, actual, expected
    and visibility. v5.1.3 only permits it on current independently-generated
    actual-vs-expected candidate sheets, so it cannot act as a historical pass.
    """
    r = rec.get("row", {})
    try:
        offpage = "offpage" if float(r.get("x0") or 0) < 0 or float(r.get("y0") or 0) < 0 else "visible"
    except Exception:
        offpage = "visible"
    core = [
        rec.get("source_sheet", ""), _norm_context(r.get("字元")),
        _norm_context(r.get("所在行")), _norm_context(r.get("局部詞境")),
        _norm_context(r.get("實際注音")),
        _norm_context(r.get("預期注音") or r.get("字典有效讀音") or r.get("規則候選讀音")),
        offpage,
    ]
    raw = json.dumps(core, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    return hashlib.sha1(raw).hexdigest()

def load_decision_memory() -> dict[str, Any]:
    # Legacy memories are never loaded into the active v5.2 evidence path.
    return {}


def remember_decision(rec: dict[str, Any], decision: dict[str, Any]) -> None:
    raise ValueError("禁止把人工／GPT 結論寫入跨 occurrence 判定記憶")


def apply_decision_memory(manifest: dict[str, Any], db: dict[str, Any]) -> int:
    # 2.4 and earlier memories are regression/audit material only. They never
    # become v5.2 terminal evidence and therefore are deliberately not applied.
    return 0


def record_is_resolved(rec: dict[str, Any], decision: dict[str, Any] | None) -> bool:
    state = str((decision or {}).get("resulting_state") or rec.get("state") or "")
    return state in TERMINAL_STATES


def expected_gap_group_key(rec: dict[str, Any]) -> str:
    """只合併完全相同的完整行／位置／局部詞境，避免過度泛化。"""
    r = rec.get("row", {})
    core = [
        _norm_context(r.get("字元")),
        _norm_context(r.get("所在行")),
        str(r.get("行內字元位置") if r.get("行內字元位置") is not None else ""),
        _norm_context(r.get("局部詞境")),
        _norm_context(r.get("實際注音")),
        _norm_context(r.get("候選短語")),
    ]
    raw = json.dumps(core, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    return hashlib.sha1(raw).hexdigest()


def derive_gap_decision(actual: Any, expected: Any) -> tuple[str, str, str]:
    """由獨立補建的 expected 與 PDF actual 做純字面比較。

    回傳 (decision, normalized_actual, normalized_expected)。任何一側不是
    單一合法注音時都拒絕匯入，避免把模糊答案升格為完成。
    """
    a = norm_bopomofo(actual)
    e = norm_bopomofo(expected)
    if not a:
        raise ValueError("實際注音不是可比較的單一注音")
    if not e:
        raise ValueError("建議預期注音不是可比較的單一注音")
    return ("PASS" if a == e else "DIFFERENCE_PENDING_CONFIRMATION", a, e)


def actual_workbook_path(actual_dir: Path, pdf: Path) -> Path:
    return actual_dir / f"{pdf.stem}_實際注音解碼.xlsx"


def candidate_workbook_path(cand_dir: Path, pdf: Path) -> Path:
    return cand_dir / f"{pdf.stem}_注音校對候選.xlsx"


def workbook_metadata(path: Path, sheet: str = "v5.2中繼資料") -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"缺少必要工作表：{sheet}")
        ws = wb[sheet]
        out = {}
        for row_number, row in enumerate(ws.iter_rows(values_only=True), 1):
            if not row or row[0] in (None, ""):
                continue
            key = str(row[0])
            if key in out:
                raise ValueError(f"{path.name}/{sheet} metadata key 重複：{key} (row {row_number})")
            out[key] = row[1] if len(row) > 1 else None
        return out
    finally:
        wb.close()


def output_is_reusable(
    output_dir: Path,
    pdf: Path,
    actual: Path,
    actual_fingerprint: Mapping[str, Any] | None = None,
    *,
    baseline_manifest: Mapping[str, Any] | None = None,
) -> bool:
    """Reuse an actual workbook across tool releases when evidence still matches.

    v5.6 deliberately ignores tool-version-only drift. PDF bytes, approved actual
    assets and project-scoped dynamic evidence remain hard reuse boundaries.
    """
    if not actual.exists():
        return False
    manifest = baseline_manifest
    if manifest is None:
        try:
            manifest = json_load_strict(output_dir / "校對工作階段.json")
            validate_manifest_integrity(manifest)
            validate_output_artifact_hashes(manifest, require_candidate=False)
        except (FileNotFoundError, ValueError):
            return False
    if not isinstance(manifest, dict):
        return False
    if not schema_compatible(manifest.get("session_schema_version"), SESSION_SCHEMA_VERSION):
        return False
    if not schema_compatible(manifest.get("ledger_schema_version"), LEDGER_SCHEMA_VERSION):
        return False
    current = dict(actual_fingerprint or {})
    if not current.get("fingerprint"):
        return False
    for info in manifest.get("pdfs", []):
        if info.get("pdf_name") == pdf.name:
            try:
                if info.get("pdf_sha256") != sha256_file(pdf):
                    return False
                metadata = workbook_metadata(actual)
                if not schema_compatible(metadata.get("workbook_schema_version"), WORKBOOK_SCHEMA_VERSION):
                    return False
                if not schema_compatible(metadata.get("ledger_schema_version"), LEDGER_SCHEMA_VERSION):
                    return False
                if metadata.get("pdf_sha256") != info.get("pdf_sha256"):
                    return False
                return fingerprint_compatible(
                    metadata.get("actual_asset_fingerprint"),
                    metadata.get("actual_asset_fingerprint_components"),
                    current,
                    chain="actual",
                )
            except Exception:
                return False
    return False

def load_reuse_baseline_manifest(output_dir: Path) -> dict[str, Any] | None:
    """Load and seal-check the previous session before any incremental writes."""
    try:
        manifest = json_load_strict(Path(output_dir) / "校對工作階段.json")
        validate_manifest_integrity(manifest)
        validate_output_artifact_hashes(manifest)
        return manifest
    except (FileNotFoundError, ValueError):
        return None


def candidate_is_baseline_reusable(
    manifest: Mapping[str, Any] | None,
    pdf: Path,
    candidate: Path,
    current_expected_fingerprint: Mapping[str, Any],
) -> bool:
    if not manifest or not candidate.exists():
        return False
    info = next((row for row in manifest.get("pdfs", []) if row.get("pdf_name") == pdf.name), None)
    if not info:
        return False
    try:
        if info.get("pdf_sha256") != sha256_file(pdf):
            return False
        if info.get("candidate_workbook_sha256") != sha256_file(candidate):
            return False
        summary = summary_dict(candidate)
        if not schema_compatible(summary.get("workbook_schema_version"), WORKBOOK_SCHEMA_VERSION):
            return False
        if not schema_compatible(summary.get("ledger_schema_version"), LEDGER_SCHEMA_VERSION):
            return False
        return fingerprint_compatible(
            summary.get("expected_asset_fingerprint"),
            summary.get("expected_asset_fingerprint_components"),
            current_expected_fingerprint,
            chain="expected",
        )
    except Exception:
        return False

def _export_pending_for_gpt_legacy_disabled(output_dir: Path) -> Path:
    raise ValueError("LEGACY_SCHEMA：v5.1.5 GPT 匯出器已隔離，不得執行")
    manifest = json_load(output_dir / "校對工作階段.json", None)
    if not manifest:
        raise FileNotFoundError("找不到校對工作階段.json，請先執行一次完整校對。")
    db = normalize_db(json_load(output_dir / "人工判定資料庫.json", {}))
    apply_decision_memory(manifest, db)
    json_save(output_dir / "人工判定資料庫.json", db)

    pending: list[dict[str, Any]] = []
    gaps: list[dict[str, Any]] = []
    for rec in manifest.get("records", []):
        d = db["decisions"].get(rec.get("review_id"), {})
        if rec.get("source_sheet") in UNRESOLVED_SHEETS:
            if not record_is_resolved(rec, d):
                pending.append(rec)
        elif rec.get("source_sheet") in COVERAGE_GAP_SHEETS:
            if not record_is_resolved(rec, d):
                gaps.append(rec)

    # 僅合併完全相同「字元＋局部詞境＋actual＋候選短語」的覆蓋缺口。
    # 不跨語境推論；同群組的所有 review_id 仍完整保留，匯入時逐筆回寫。
    gap_groups: dict[str, list[dict[str, Any]]] = {}
    for rec in gaps:
        gap_groups.setdefault(expected_gap_group_key(rec), []).append(rec)

    wb = Workbook()
    guide = wb.active
    guide.title = "使用說明"
    guide.append(["項目", "說明"])
    instructions = [
        ("用途", "本報表同時處理兩種殘餘：①已有 expected 但仍需判定的候選；②實際注音已解碼、但尚未建立 expected 的覆蓋缺口。"),
        ("最高優先規則", "《統一用字手冊113.10.04》的注音規則為最高優先；若與一般字典、語感或其他低優先規則衝突，以公司手冊為準。"),
        ("實際注音原則", "『實際注音』是程式從 PDF 字形解出的觀測值；不得因為語意上應該怎麼念而反過來改寫實際注音。若懷疑 observed 解錯，必須回看原 PDF。"),
        ("『一』變調規則", "依使用者確認規則：『一』在去聲字前標ㄧˊ；在陰平、陽平、上聲字前標ㄧˋ。A一A 同樣適用。"),
        ("工作表一", "『待判定候選』沿用原流程：填『判定』『修正實際注音』『判定理由』『GPT信心』。"),
        ("工作表二", "『待補預期音』只在能從完整詞境安全確定時填『建議預期注音』與『預期音依據』。程式匯入後會自行拿 expected 與 actual 比較，決定教材錯誤或課本正確。"),
        ("待補預期音－不確定", "若無法安全建立 expected，建議預期注音保持空白，處理方式填『需人工確認』或留白；不得猜測。"),
        ("待補預期音－特殊略過", "若回看 PDF 後確認該列不是應校的可見注音物件，可將處理方式填『特殊略過』。"),
        ("解碼錯誤", "只有另外檢視原 PDF 或頁面影像，確定程式讀錯實際注音時才可在『待判定候選』選解碼錯誤；單靠詞境不可以。"),
        ("一般殘餘候選", len(pending)),
        ("未建立預期音座標", len(gaps)),
        ("未建立預期音精確詞境群組", len(gap_groups)),
        ("回寫方式", "GPT 填好後存成 xlsx，再用程式的『匯入 GPT 判定報表』。匯入只回寫目前工作階段，不會把 GPT 補建 expected 自動提升成跨教材永久規則。"),
    ]
    for r in instructions:
        guide.append(r)
    _style_sheet(guide)
    guide.column_dimensions["A"].width = 30
    guide.column_dimensions["B"].width = 105

    ws = wb.create_sheet("待判定候選")
    headers = ["編號","review_id","PDF","課本頁","實體頁碼","字元","所在行","局部詞境","實際注音","預期／候選注音","來源類別","候選短語","上下文","字典讀音原文","字典有效讀音","字典例詞","完整詞語規則","規則ID","規則模式","規則來源","規則備註","穩定注音鍵","font","注音元件ID","x0","y0","x1","y1","判定","修正實際注音","判定理由","GPT信心"]
    ws.append(headers)
    for i, rec in enumerate(pending, 1):
        r = rec.get("row", {})
        ws.append([i,rec["review_id"],rec["pdf_name"],r.get("課本頁"),r.get("實體頁碼"),r.get("字元"),r.get("所在行"),r.get("局部詞境"),r.get("實際注音"),r.get("預期注音") or r.get("字典有效讀音") or r.get("規則候選讀音"),rec.get("source_sheet"),r.get("候選短語"),r.get("上下文"),r.get("字典讀音原文"),r.get("字典有效讀音"),r.get("字典例詞"),r.get("完整詞語規則"),r.get("規則ID"),r.get("規則模式"),r.get("規則來源"),r.get("規則備註"),r.get("穩定注音鍵"),r.get("font"),r.get("注音元件ID"),r.get("x0"),r.get("y0"),r.get("x1"),r.get("y1"),"","","",""])
    _style_sheet(ws)
    dv = DataValidation(type="list", formula1='"教材錯誤,課本正確,特殊略過,解碼錯誤"', allow_blank=True)
    ws.add_data_validation(dv)
    if ws.max_row >= 2:
        dv.add(f"AC2:AC{ws.max_row}")
    dv2 = DataValidation(type="list", formula1='"高,中,低"', allow_blank=True)
    ws.add_data_validation(dv2)
    if ws.max_row >= 2:
        dv2.add(f"AF2:AF{ws.max_row}")
    for col in ["G","H","L","M","P","Q","U","AE"]:
        ws.column_dimensions[col].width = 45
    ws.column_dimensions["B"].width = 42

    wg = wb.create_sheet("待補預期音")
    gap_headers = [
        "編號","group_id","成員數","review_id列表","PDF／頁面清單",
        "代表PDF","課本頁","實體頁碼","字元","所在行","局部詞境","實際注音",
        "來源類別","候選短語","上下文","字典例詞索引提示","穩定注音鍵","font","注音元件ID",
        "建議預期注音","預期音依據","GPT信心","處理方式","備註"
    ]
    wg.append(gap_headers)
    for i, (gid, members) in enumerate(sorted(gap_groups.items(), key=lambda kv: (kv[1][0].get("pdf_name", ""), str(kv[1][0].get("row", {}).get("實體頁碼") or ""), str(kv[1][0].get("row", {}).get("x0") or ""))), 1):
        rec = members[0]
        r = rec.get("row", {})
        pages = []
        for m in members:
            mr = m.get("row", {})
            pages.append(f"{m.get('pdf_name','')}：課本{mr.get('課本頁','')}/實體{mr.get('實體頁碼','')}")
        wg.append([
            i, gid, len(members), "|".join(m["review_id"] for m in members), "；".join(pages),
            rec.get("pdf_name"), r.get("課本頁"), r.get("實體頁碼"), r.get("字元"), r.get("所在行"), r.get("局部詞境"), r.get("實際注音"),
            "未建立預期音", r.get("候選短語"), r.get("上下文"), r.get("字典例詞索引提示"), r.get("穩定注音鍵"), r.get("font"), r.get("注音元件ID"),
            "", "", "", "", ""
        ])
    _style_sheet(wg)
    dv3 = DataValidation(type="list", formula1='"高,中,低"', allow_blank=True)
    wg.add_data_validation(dv3)
    if wg.max_row >= 2:
        dv3.add(f"V2:V{wg.max_row}")
    dv4 = DataValidation(type="list", formula1='"建立預期音,特殊略過,需人工確認"', allow_blank=True)
    wg.add_data_validation(dv4)
    if wg.max_row >= 2:
        dv4.add(f"W2:W{wg.max_row}")
    for col in ["D","E","J","K","N","O","P","U","X"]:
        wg.column_dimensions[col].width = 45
    wg.column_dimensions["B"].width = 42

    out = output_dir / "待判定候選_給GPT.xlsx"
    wb.save(out)
    return out


def _hash_review_payload(payload: Mapping[str, Any]) -> str:
    raw = json.dumps(dict(payload), ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def _review_snapshot(entry: Mapping[str, Any]) -> str:
    """Full snapshot used by actions that depend on both actual and expected.

    A final difference confirmation must be invalidated when *either* evidence
    chain changes, so this intentionally includes both sides and the state.
    """
    payload = {
        "session_schema_version": SESSION_SCHEMA_VERSION,
        "review_id_schema_version": REVIEW_ID_SCHEMA_VERSION,
        "occurrence_id": entry.get("occurrence_id"),
        "review_id": entry.get("review_id"),
        "state": entry.get("state"),
        "actual": entry.get("actual"),
        "expected_set": list(normalize_expected_set(entry.get("expected_set"))),
        "actual_evidence": entry.get("actual_evidence"),
        "expected_evidence": entry.get("expected_evidence"),
        "context_evidence": entry.get("context_evidence"),
    }
    return _hash_review_payload(payload)


def _expected_review_snapshot(entry: Mapping[str, Any]) -> str:
    """Snapshot only of the expected/context chain.

    Expected evidence must remain independent from actual.  Re-decoding actual,
    changing an actual fingerprint, or refreshing actual evidence must therefore
    not make an already completed expected judgment stale.
    """
    payload = {
        "session_schema_version": SESSION_SCHEMA_VERSION,
        "review_id_schema_version": REVIEW_ID_SCHEMA_VERSION,
        "occurrence_id": entry.get("occurrence_id"),
        "review_id": entry.get("review_id"),
        "expected_set": list(normalize_expected_set(entry.get("expected_set"))),
        "expected_evidence": str(entry.get("expected_evidence") or ""),
        "context_evidence": str(entry.get("context_evidence") or ""),
    }
    return _hash_review_payload(payload)


def _expected_review_snapshot_from_export_row(row: Mapping[str, Any]) -> str:
    """Backward-compatible expected snapshot for workbooks exported before hotfix 5."""
    payload = {
        "session_schema_version": SESSION_SCHEMA_VERSION,
        "review_id_schema_version": REVIEW_ID_SCHEMA_VERSION,
        "occurrence_id": str(row.get("occurrence_id") or ""),
        "review_id": str(row.get("review_id") or ""),
        "expected_set": list(normalize_expected_set(row.get("expected_set"))),
        "expected_evidence": str(row.get("expected_evidence") or ""),
        "context_evidence": str(row.get("context_evidence") or ""),
    }
    return _hash_review_payload(payload)


def _expected_target_snapshot(entry: Mapping[str, Any]) -> str:
    """Stable semantic target identity for expected-only review.

    Unlike ``_expected_review_snapshot``, this deliberately excludes both the
    automatic expected resolver output and all actual fields.  It represents
    only *which printed occurrence and context* the reviewer judged.  This lets
    an expected decision survive the common sequence where actual was initially
    undecoded, then becomes available and causes the automatic expected resolver
    to run for the first time.
    """
    source = entry.get("source_record") or entry.get("row") or {}
    payload = {
        "session_schema_version": SESSION_SCHEMA_VERSION,
        "review_id_schema_version": REVIEW_ID_SCHEMA_VERSION,
        "occurrence_id": str(entry.get("occurrence_id") or ""),
        "review_id": str(entry.get("review_id") or ""),
        "pdf_name": str(entry.get("pdf_name") or source.get("pdf_name") or ""),
        "physical_page": str(entry.get("physical_page") or source.get("實體頁碼") or ""),
        "printed_page": str(entry.get("printed_page") or source.get("課本頁") or ""),
        "char": str(entry.get("char") or source.get("字元") or ""),
        "line": str(source.get("所在行") or ""),
        "local_context": str(source.get("局部詞境") or ""),
    }
    return _hash_review_payload(payload)


def _expected_target_snapshot_from_export_row(row: Mapping[str, Any]) -> str:
    payload = {
        "session_schema_version": SESSION_SCHEMA_VERSION,
        "review_id_schema_version": REVIEW_ID_SCHEMA_VERSION,
        "occurrence_id": str(row.get("occurrence_id") or ""),
        "review_id": str(row.get("review_id") or ""),
        "pdf_name": str(row.get("PDF") or ""),
        "physical_page": str(row.get("實體頁碼") or ""),
        "printed_page": str(row.get("課本頁") or ""),
        "char": str(row.get("字元") or ""),
        "line": str(row.get("所在行") or ""),
        "local_context": str(row.get("局部詞境") or ""),
    }
    return _hash_review_payload(payload)


def _expected_target_unchanged(row: Mapping[str, Any], entry: Mapping[str, Any]) -> bool:
    """Fail-closed semantic-target comparison, including old exported workbooks."""
    exported = str(row.get("expected_target_snapshot") or "").strip()
    if not exported:
        exported = _expected_target_snapshot_from_export_row(row)
    return exported == _expected_target_snapshot(entry)


def needs_expected_review(entry: Mapping[str, Any]) -> bool:
    """Return whether this occurrence belongs in the expected/GPT review lane."""
    state = str(entry.get("state") or "")
    if state in {"SOURCE_INVALID", "DATA_INTEGRITY_ERROR", "REGRESSION_BLOCKED", *EXCLUDED_STATES}:
        return False
    expected_status = infer_expected_status(entry)
    if expected_status in {"UNRESOLVED", "AMBIGUOUS", "CONFLICT"}:
        return True
    return state in {"DIFFERENCE_PENDING_CONFIRMATION", "REVIEW_PENDING"}


def export_pending_for_gpt(output_dir: Path) -> Path:
    manifest = json_load_strict(output_dir / "校對工作階段.json")
    validate_manifest_integrity(manifest)
    validate_output_artifact_hashes(manifest)
    if not schema_compatible(manifest.get("session_schema_version"), SESSION_SCHEMA_VERSION):
        raise ValueError("SESSION_SCHEMA_INCOMPATIBLE：工作階段資料布局不同，需有明確轉接器")
    db = load_or_initialize_db(output_dir)
    ledger = materialize_ledger(manifest, db)
    # Expected/GPT workbook is lane-specific.  Pure actual decode failures with
    # an already-resolved expected lane belong only in the actual visual-review
    # package; exporting them here previously created duplicated work and stale
    # cross-lane snapshots.
    pending = [entry for entry in ledger if needs_expected_review(entry)]
    assert_unique_ids(pending, "occurrence_id")
    assert_unique_ids(pending, "review_id")

    wb = Workbook()
    guide = wb.active
    guide.title = "使用說明"
    guide.append(["項目", "說明"])
    instructions = [
        ("用途", "本表用來補充應標讀音證據、解決規則衝突或完成教材錯誤確認；程式仍會自行做 actual／expected 比較。"),
        ("actual 修正", "不得在本表直接改 actual。請用『輸出 actual 待判定給 GPT』產生 actual 視覺判定包，再以『匯入 GPT 判定檔（自動辨識）』匯入；程式會自行維護 occurrence override 並增量重解碼，不需要手動編輯 CSV。"),
        ("補充／解決 expected", "一般未決可選『補建expected證據』；規則衝突或需要以較高優先來源更正 expected 時選『解決expected證據』。都必須填讀音、來源與詞境證據。"),
        ("確認差異", "只有原 state 為 DIFFERENCE_PENDING_CONFIRMATION，且六個 gate 欄全填 Y，才可進 TEXTBOOK_ERROR_CONFIRMED。"),
        ("排除", "action 選『確認非校對範圍』時，排除理由與排除證據都必填；不能只按特殊略過。"),
        ("交易性", "任何一列 duplicate/unknown/stale/schema/證據錯誤，整批拒絕且一列都不寫入。"),
        ("expected／差異待處理數", len(pending)),
    ]
    for item in instructions:
        guide.append(item)

    meta = wb.create_sheet("匯入中繼資料")
    meta.append(["項目", "內容"])
    for item in [
        ("version", VERSION),
        ("session_id", manifest.get("session_id", "")),
        ("session_schema_version", SESSION_SCHEMA_VERSION),
        ("workbook_schema_version", WORKBOOK_SCHEMA_VERSION),
        ("review_id_schema_version", REVIEW_ID_SCHEMA_VERSION),
        ("expected_asset_fingerprint", manifest.get("expected_asset_fingerprint", "")),
        ("exported_review_count", len(pending)),
    ]:
        meta.append(item)

    ws = wb.create_sheet("待判定候選")
    headers = [
        "編號", "occurrence_id", "review_id", "review_snapshot", "expected_snapshot", "expected_target_snapshot", "state", "actual_status", "expected_status", "狀態說明", "PDF", "課本頁", "實體頁碼", "字元",
        "所在行", "局部詞境", "actual", "actual_evidence", "expected_set", "expected_evidence", "context_evidence",
        "source_view", "action", "proposed_expected_set", "proposed_expected_evidence", "proposed_context_evidence",
        "exclusion_reason", "exclusion_evidence", *CONFIRMATION_GATES, "note",
    ]
    ws.append(headers)
    for index, entry in enumerate(pending, 1):
        source = entry.get("source_record") or {}
        row = {
            "編號": index,
            "occurrence_id": entry.get("occurrence_id"),
            "review_id": entry.get("review_id"),
            "review_snapshot": _review_snapshot(entry),
            "expected_snapshot": _expected_review_snapshot(entry),
            "expected_target_snapshot": _expected_target_snapshot(entry),
            "state": entry.get("state"),
            "actual_status": infer_actual_status(entry),
            "expected_status": infer_expected_status(entry),
            "狀態說明": friendly_state(entry.get("state")),
            "PDF": entry.get("pdf_name"),
            "課本頁": entry.get("printed_page"),
            "實體頁碼": entry.get("physical_page"),
            "字元": entry.get("char"),
            "所在行": source.get("所在行", ""),
            "局部詞境": source.get("局部詞境", ""),
            "actual": entry.get("actual"),
            "actual_evidence": entry.get("actual_evidence"),
            "expected_set": "|".join(entry.get("expected_set") or []),
            "expected_evidence": entry.get("expected_evidence"),
            "context_evidence": entry.get("context_evidence"),
            "source_view": entry.get("source_view"),
        }
        ws.append([row.get(header, "") for header in headers])
    action_column = headers.index("action") + 1
    dv = DataValidation(type="list", formula1='"補建expected證據,解決expected證據,確認現版差異,確認非校對範圍,保留待人工"', allow_blank=True)
    ws.add_data_validation(dv)
    if ws.max_row >= 2:
        letter = get_column_letter(action_column)
        dv.add(f"{letter}2:{letter}{ws.max_row}")
    for gate in CONFIRMATION_GATES:
        column = headers.index(gate) + 1
        gate_dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
        ws.add_data_validation(gate_dv)
        if ws.max_row >= 2:
            letter = get_column_letter(column)
            gate_dv.add(f"{letter}2:{letter}{ws.max_row}")
    meta.sheet_state = "hidden"
    for col_name in ("occurrence_id", "review_id", "review_snapshot", "expected_snapshot", "expected_target_snapshot", "state", "source_view"):
        if col_name in headers:
            ws.column_dimensions[get_column_letter(headers.index(col_name) + 1)].hidden = True
    for sheet in wb.worksheets:
        _style_sheet(sheet)
    out = output_dir / "待判定候選_給GPT.xlsx"
    wb.save(out)
    return out


def _import_gpt_decisions_legacy_disabled(output_dir: Path, xlsx: Path) -> tuple[int, int, Path]:
    raise ValueError("LEGACY_SCHEMA：v5.1.5 GPT 匯入器已隔離，不得執行")
    manifest = json_load(output_dir / "校對工作階段.json", None)
    if not manifest:
        raise FileNotFoundError("找不到校對工作階段.json，請先執行一次完整校對。")
    by_id = {r.get("review_id"): r for r in manifest.get("records", [])}
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    if "待判定候選" not in wb.sheetnames and "待補預期音" not in wb.sheetnames:
        raise ValueError("匯入檔缺少『待判定候選』與『待補預期音』工作表。")

    db = normalize_db(json_load(output_dir / "人工判定資料庫.json", {}))
    imported = 0
    skipped = 0

    # 既有候選：沿用 v5.1 流程，由 GPT / 使用者直接填四類判定。
    if "待判定候選" in wb.sheetnames:
        ws = wb["待判定候選"]
        rows = ws.iter_rows(values_only=True)
        try:
            hdr = [str(x or "") for x in next(rows)]
        except StopIteration:
            hdr = []
        ix = {x: i for i, x in enumerate(hdr)}
        if hdr and ("review_id" not in ix or "判定" not in ix):
            raise ValueError("『待判定候選』缺少 review_id 或 判定 欄位。")
        for row in rows:
            rid = str(row[ix["review_id"]] or "").strip()
            decision = str(row[ix["判定"]] or "").strip()
            if not rid or not decision:
                continue
            rec = by_id.get(rid)
            if not rec or rec.get("source_sheet") not in UNRESOLVED_SHEETS or decision not in DECISIONS:
                skipped += 1
                continue
            corr = str(row[ix["修正實際注音"]] or "").strip() if "修正實際注音" in ix else ""
            note = str(row[ix["判定理由"]] or "").strip() if "判定理由" in ix else ""
            if decision == "解碼錯誤" and not corr:
                skipped += 1
                continue
            d = {
                "decision": decision,
                "corrected_actual": corr,
                "note": note,
                "source": "GPT判定報表匯入（使用者匯入）",
                "updated_at": datetime.now().isoformat(timespec="seconds"),
            }
            db["decisions"][rid] = d
            remember_decision(rec, d)
            imported += 1

    # 新流程：未建立 expected 的座標，先補 expected，再由程式比較 actual。
    # GPT 只負責提供 expected 與依據，不直接決定「教材錯誤／課本正確」。
    if "待補預期音" in wb.sheetnames:
        ws = wb["待補預期音"]
        rows = ws.iter_rows(values_only=True)
        try:
            hdr = [str(x or "") for x in next(rows)]
        except StopIteration:
            hdr = []
        ix = {x: i for i, x in enumerate(hdr)}
        if hdr and "review_id列表" not in ix:
            raise ValueError("『待補預期音』缺少 review_id列表 欄位。")
        for row in rows:
            rid_text = str(row[ix["review_id列表"]] or "").strip() if "review_id列表" in ix else ""
            if not rid_text:
                continue
            rids = [x.strip() for x in rid_text.split("|") if x.strip()]
            expected = str(row[ix["建議預期注音"]] or "").strip() if "建議預期注音" in ix else ""
            evidence = str(row[ix["預期音依據"]] or "").strip() if "預期音依據" in ix else ""
            confidence = str(row[ix["GPT信心"]] or "").strip() if "GPT信心" in ix else ""
            mode = str(row[ix["處理方式"]] or "").strip() if "處理方式" in ix else ""
            note = str(row[ix["備註"]] or "").strip() if "備註" in ix else ""
            if mode in {"", "需人工確認"} and not expected:
                continue
            if mode == "需人工確認":
                continue

            for rid in rids:
                rec = by_id.get(rid)
                if not rec or rec.get("source_sheet") not in COVERAGE_GAP_SHEETS:
                    skipped += 1
                    continue
                if mode == "特殊略過":
                    db["decisions"][rid] = {
                        "decision": "特殊略過",
                        "note": note or evidence,
                        "source": "GPT預期音報表匯入（使用者匯入）",
                        "gpt_confidence": confidence,
                        "updated_at": datetime.now().isoformat(timespec="seconds"),
                    }
                    imported += 1
                    continue
                try:
                    decision, actual_norm, expected_norm = derive_gap_decision(rec.get("row", {}).get("實際注音"), expected)
                except ValueError:
                    skipped += 1
                    continue
                db["decisions"][rid] = {
                    "decision": decision,
                    "expected": expected_norm,
                    "expected_evidence": evidence,
                    "gpt_confidence": confidence,
                    "note": note,
                    "source": "GPT補建預期音報表匯入（使用者匯入）",
                    "updated_at": datetime.now().isoformat(timespec="seconds"),
                }
                imported += 1

    json_save(output_dir / "人工判定資料庫.json", db)
    save_pending_json(output_dir, manifest, db)
    report = generate_report(output_dir, manifest, db)
    return imported, skipped, report



def _sha_text(value: Any) -> str:
    return hashlib.sha256(str(value or "").encode("utf-8")).hexdigest()


def detect_gpt_decision_path_kind(path: Path) -> str:
    path = Path(path)
    if path.suffix.lower() == ".zip":
        try:
            with zipfile.ZipFile(path, "r") as zf:
                names = set(zf.namelist())
                if "GPT判定包.json" in names and "actual_occurrence_decisions.csv" in names:
                    return "bundle"
        except zipfile.BadZipFile as exc:
            raise ValueError(f"GPT 判定包不是有效 ZIP：{path.name}") from exc
        raise ValueError("無法辨識 GPT ZIP：缺少 GPT判定包.json 或 actual_occurrence_decisions.csv")
    return detect_gpt_workbook_kind(path)


def _load_gpt_bundle_metadata(zf: zipfile.ZipFile) -> dict[str, Any]:
    try:
        raw = zf.read("GPT判定包.json").decode("utf-8-sig")
        meta = json.loads(raw)
    except Exception as exc:
        raise ValueError("GPT 判定包 manifest 損壞") from exc
    if not isinstance(meta, dict):
        raise ValueError("GPT 判定包 manifest 格式錯誤")
    return meta


def _validate_gpt_bundle_metadata(meta: Mapping[str, Any], manifest: Mapping[str, Any]) -> None:
    bad = {}
    if str(meta.get("bundle_schema_version") or "") != GPT_DECISION_BUNDLE_SCHEMA_VERSION:
        bad["bundle_schema_version"] = (GPT_DECISION_BUNDLE_SCHEMA_VERSION, meta.get("bundle_schema_version"))
    if str(meta.get("session_id") or "") != str(manifest.get("session_id") or ""):
        bad["session_id"] = (manifest.get("session_id"), meta.get("session_id"))
    if not schema_compatible(meta.get("session_schema_version"), SESSION_SCHEMA_VERSION):
        bad["session_schema_version"] = (SESSION_SCHEMA_VERSION, meta.get("session_schema_version"))
    if not schema_compatible(meta.get("workbook_schema_version"), WORKBOOK_SCHEMA_VERSION):
        bad["workbook_schema_version"] = (WORKBOOK_SCHEMA_VERSION, meta.get("workbook_schema_version"))
    if not review_id_schema_compatible(meta.get("review_id_schema_version"), REVIEW_ID_SCHEMA_VERSION):
        bad["review_id_schema_version"] = (REVIEW_ID_SCHEMA_VERSION, meta.get("review_id_schema_version"))
    # meta['version'] is audit-only from v5.6 onward.
    if bad:
        raise ValueError(f"GPT 判定包 session/schema 不相容：{bad}")

def import_actual_occurrence_decisions(output_dir: Path, csv_path: Path, *, package_meta: Mapping[str, Any]) -> tuple[int, int, Path]:
    """Import occurrence-scoped visual actual decisions from a GPT bundle.

    The bundle deliberately carries no expected value for this path.  Current
    occurrence identity and the exported actual/evidence are revalidated before
    any dynamic actual truth file is modified.  Exact-glyph propagation is used
    only when at least two independently reviewed occurrences from the same
    current group are present in the bundle; otherwise the correction remains
    occurrence-specific.
    """
    output_dir = Path(output_dir)
    manifest = json_load_strict(output_dir / "校對工作階段.json")
    validate_manifest_integrity(manifest)
    validate_output_artifact_hashes(manifest)
    _validate_gpt_bundle_metadata(package_meta, manifest)
    db = load_or_initialize_db(output_dir)
    ledger = materialize_ledger(manifest, db)
    by_occ = {str(e.get("occurrence_id") or ""): e for e in ledger}

    required = {
        "occurrence_id", "review_id", "pdf_name", "physical_page", "char",
        "exported_actual", "exported_actual_evidence_sha256", "verified_actual",
        "visual_confirmation", "note",
    }
    with Path(csv_path).open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if not reader.fieldnames or not required.issubset(set(reader.fieldnames)):
            raise ValueError(f"GPT actual occurrence 決策表缺少欄位：{sorted(required - set(reader.fieldnames or []))}")
        rows = list(reader)
    if not rows:
        return 0, 0, generate_report(output_dir, manifest, db)

    staged_by_group: dict[str, dict[str, Any]] = {}
    errors: list[str] = []
    seen_occ: set[str] = set()
    already_applied_ids: set[str] = set()
    for row_no, row in enumerate(rows, 2):
        oid = str(row.get("occurrence_id") or "").strip()
        if not oid:
            errors.append(f"row {row_no}: occurrence_id 空白")
            continue
        if oid in seen_occ:
            errors.append(f"row {row_no}: occurrence_id 重複 {oid}")
            continue
        seen_occ.add(oid)
        entry = by_occ.get(oid)
        if entry is None:
            errors.append(f"row {row_no}: current occurrence 不存在 {oid}")
            continue
        if str(row.get("review_id") or "").strip() != str(entry.get("review_id") or ""):
            errors.append(f"row {row_no}: review_id/occurrence_id mapping 已改變")
            continue
        if str(row.get("pdf_name") or "").strip() != str(entry.get("pdf_name") or ""):
            errors.append(f"row {row_no}: PDF identity 已改變")
            continue
        if str(row.get("physical_page") or "").strip() != str(entry.get("physical_page") or ""):
            errors.append(f"row {row_no}: page identity 已改變")
            continue
        if str(row.get("char") or "").strip() != str(entry.get("char") or ""):
            errors.append(f"row {row_no}: char identity 已改變")
            continue
        exported_actual = canonical_bopomofo(row.get("exported_actual"))
        current_actual = canonical_bopomofo(entry.get("actual"))
        if str(row.get("visual_confirmation") or "").strip().upper() != "Y":
            errors.append(f"row {row_no}: actual 修正必須 visual_confirmation=Y")
            continue
        verified = canonical_bopomofo(row.get("verified_actual"))
        if not verified:
            errors.append(f"row {row_no}: verified_actual 非合法注音")
            continue

        # Crash-safe/idempotent retry: Hotfix 6 could successfully write the
        # dynamic actual truth and rebuild the new manifest, then fail later
        # while replaying an unrelated expected event.  Retrying the same GPT
        # bundle must not demand a fresh visual review merely because current
        # actual now equals the already-verified answer.  Identity is still
        # checked above; this branch does not consult expected.
        if current_actual == verified and current_actual != exported_actual:
            if not entry.get("actual_evidence"):
                errors.append(f"row {row_no}: 已套用 actual 缺少 current actual evidence")
                continue
            already_applied_ids.add(oid)
            continue

        if exported_actual != current_actual:
            errors.append(f"row {row_no}: current actual 已改變，且不等於本判定 verified_actual；請重新覆核")
            continue
        if str(row.get("exported_actual_evidence_sha256") or "").strip() != _sha_text(entry.get("actual_evidence")):
            errors.append(f"row {row_no}: actual evidence 已改變，請重新覆核")
            continue
        group = build_actual_group_for_entry(ledger, entry)
        slot = staged_by_group.setdefault(group["group_id"], {"group": group, "reading": verified, "ids": [], "notes": []})
        if slot["reading"] != verified:
            errors.append(f"row {row_no}: 同一 exact glyph group 出現互斥 actual")
            continue
        slot["ids"].append(oid)
        if str(row.get("note") or "").strip():
            slot["notes"].append(str(row.get("note") or "").strip())
    if errors:
        raise ValueError("GPT actual occurrence 匯入整批拒絕；未寫入任何一列：\n" + "\n".join(errors[:100]))

    root = initialize_project_actual_evidence(output_dir, Path(__file__).resolve().parent)
    dynamic_names = [OCCURRENCE_OVERRIDE_FILE, USER_GLYF_FILE, USER_CFF_FILE, GLYPH_CONFLICT_FILE, GLYPH_PROVENANCE_FILE]
    backups = {root / name: (root / name).read_bytes() if (root / name).exists() else None for name in dynamic_names}
    results = []
    try:
        for slot in staged_by_group.values():
            results.append(apply_verified_actual_group(
                root,
                slot["group"],
                slot["reading"],
                checked_occurrence_ids=slot["ids"],
                source="GPT 判定包 actual 原頁視覺覆核",
                note="；".join(slot["notes"]) or "GPT 判定包：actual 僅由原頁可見注音判定",
            ))
    except Exception:
        for path, data in backups.items():
            if data is None:
                if path.exists():
                    path.unlink()
            else:
                path.write_bytes(data)
        raise

    affected = {oid for item in results for oid in (item.get("affected_occurrence_ids") or item.get("target_occurrence_ids") or []) if oid}
    removed = _clear_actual_dependent_events(output_dir, ledger, affected)
    if already_applied_ids:
        print(
            f"  GPT actual 判定包重試：{len(already_applied_ids)} 筆 current actual 已等於 verified_actual；"
            "視為前次已安全套用。",
            flush=True,
        )
    if results:
        report = refresh_actual_project(output_dir)
    else:
        # A previous bundle attempt may already have completed the expensive
        # actual refresh and then failed in the expected phase.  Re-importing the
        # same bundle should continue directly from expected instead of decoding
        # the same PDFs again.
        report = output_dir / "注音校對_最終報告.xlsx"
        if not report.exists():
            report = generate_report(output_dir, manifest, db)
        if already_applied_ids:
            print("  actual 階段已在前次完成；略過重解碼，直接續跑 expected 階段。", flush=True)
    return len(rows), removed, report


def import_gpt_decision_bundle(output_dir: Path, bundle: Path) -> tuple[int, int, int, int, Path]:
    """One-file round import with expected preflight and lane-ordered commit.

    Expected actions are validated against the immutable occurrence/context target
    before any actual evidence is written.  Actual is then committed and the
    project refreshed, followed by the preflighted expected workbook.
    """
    output_dir = Path(output_dir)
    bundle = Path(bundle)
    manifest = json_load_strict(output_dir / "校對工作階段.json")
    validate_manifest_integrity(manifest)
    validate_output_artifact_hashes(manifest)
    with zipfile.ZipFile(bundle, "r") as zf:
        meta = _load_gpt_bundle_metadata(zf)
        _validate_gpt_bundle_metadata(meta, manifest)
        names = set(zf.namelist())
        expected_name = str(meta.get("expected_workbook") or "").strip()
        if expected_name and expected_name not in names:
            raise ValueError(f"GPT 判定包缺少 expected workbook：{expected_name}")
        with tempfile.TemporaryDirectory(prefix="zhuyin_gpt_bundle_") as td:
            td = Path(td)
            actual_csv = td / "actual_occurrence_decisions.csv"
            actual_csv.write_bytes(zf.read("actual_occurrence_decisions.csv"))
            expected_path = None
            if expected_name:
                expected_path = td / Path(expected_name).name
                expected_path.write_bytes(zf.read(expected_name))
                # Fail before actual mutation if the expected workbook is already
                # invalid for this printed occurrence/context or source version.
                import_gpt_decisions(output_dir, expected_path, dry_run=True)
            actual_n, removed, report = import_actual_occurrence_decisions(output_dir, actual_csv, package_meta=meta)
            expected_n = skipped = 0
            if expected_path is not None:
                expected_n, skipped, report = import_gpt_decisions(output_dir, expected_path)
            return actual_n, removed, expected_n, skipped, report


def current_pending_state_counts(output_dir: Path) -> dict[str, int]:
    manifest = json_load_strict(Path(output_dir) / "校對工作階段.json")
    db = load_or_initialize_db(Path(output_dir))
    return dict(Counter(str(e.get("state") or "") for e in materialize_ledger(manifest, db) if e.get("state") in NON_TERMINAL_STATES))


def detect_gpt_workbook_kind(xlsx: Path) -> str:
    """Fail-closed workbook type detection for GUI handoff imports.

    expected/difference evidence workbooks contain ``待判定候選`` plus
    ``匯入中繼資料``.  Actual visual-review workbooks contain
    ``actual待判定``.  Do not infer from filenames because users often rename
    downloaded GPT results.
    """
    path = Path(xlsx)
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        names = set(wb.sheetnames)
    finally:
        wb.close()
    is_expected = "待判定候選" in names and "匯入中繼資料" in names
    is_actual = "actual待判定" in names
    if is_expected and is_actual:
        raise ValueError("GPT 匯入檔同時含 expected 與 actual 工作表，為避免證據鏈混用，拒絕自動匯入")
    if is_actual:
        return "actual"
    if is_expected:
        return "expected"
    raise ValueError(
        "無法辨識 GPT 匯入檔：expected／差異檔必須含『待判定候選』『匯入中繼資料』；"
        "actual 檔必須含『actual待判定』。"
    )


def plan_gpt_auto_imports(paths: Iterable[str | Path]) -> list[tuple[Path, str]]:
    """Classify and deterministically order one or more GPT result workbooks.

    Actual is applied first because it is the observation chain.  Expected-only
    workbooks are then imported through the action-scoped safe-rebase checks, so
    users may multi-select both files in one GUI operation without manually
    deciding which one goes first.
    """
    seen: set[str] = set()
    items: list[tuple[Path, str]] = []
    for raw in paths:
        path = Path(raw)
        key = str(path.resolve()) if path.exists() else str(path)
        if key in seen:
            continue
        seen.add(key)
        items.append((path, detect_gpt_decision_path_kind(path)))
    order = {"bundle": 0, "actual": 1, "expected": 2}
    return sorted(items, key=lambda item: (order.get(item[1], 99), str(item[0])))


def import_gpt_decisions(output_dir: Path, xlsx: Path, *, dry_run: bool = False) -> tuple[int, int, Path]:
    manifest = json_load_strict(output_dir / "校對工作階段.json")
    validate_manifest_integrity(manifest)
    validate_output_artifact_hashes(manifest)
    if (
        not schema_compatible(manifest.get("session_schema_version"), SESSION_SCHEMA_VERSION)
        or not review_id_schema_compatible(manifest.get("review_id_schema_version"), REVIEW_ID_SCHEMA_VERSION)
    ):
        raise ValueError("SESSION_SCHEMA_INCOMPATIBLE：工作階段／review ID 無法直接沿用")

    metadata = workbook_metadata(xlsx, "匯入中繼資料")
    mismatched = {}
    if str(metadata.get("session_id") or "") != str(manifest.get("session_id") or ""):
        mismatched["session_id"] = (manifest.get("session_id"), metadata.get("session_id"))
    if not schema_compatible(metadata.get("session_schema_version"), SESSION_SCHEMA_VERSION):
        mismatched["session_schema_version"] = (SESSION_SCHEMA_VERSION, metadata.get("session_schema_version"))
    if not schema_compatible(metadata.get("workbook_schema_version"), WORKBOOK_SCHEMA_VERSION):
        mismatched["workbook_schema_version"] = (WORKBOOK_SCHEMA_VERSION, metadata.get("workbook_schema_version"))
    if not review_id_schema_compatible(metadata.get("review_id_schema_version"), REVIEW_ID_SCHEMA_VERSION):
        mismatched["review_id_schema_version"] = (REVIEW_ID_SCHEMA_VERSION, metadata.get("review_id_schema_version"))
    # Tool version is intentionally ignored. The exported expected fingerprint
    # is compared to the session snapshot, so an old workbook from the same
    # project remains importable after an application upgrade.
    workbook_expected_fp = str(metadata.get("expected_asset_fingerprint") or "")
    allowed_expected_fps = {
        str(manifest.get("expected_asset_fingerprint") or ""),
        *{str(x) for x in (manifest.get("compatible_expected_asset_fingerprints") or [])},
    }
    allowed_expected_fps.discard("")
    if (allowed_expected_fps or workbook_expected_fp) and workbook_expected_fp not in allowed_expected_fps:
        mismatched["expected_asset_fingerprint"] = (sorted(allowed_expected_fps), workbook_expected_fp)
    if mismatched:
        raise ValueError(f"匯入 workbook session/evidence 不相容：{mismatched}")

    required_columns = [
        "occurrence_id", "review_id", "review_snapshot", "state", "action", "proposed_expected_set",
        "proposed_expected_evidence", "proposed_context_evidence", "exclusion_reason", "exclusion_evidence",
        *CONFIRMATION_GATES, "note",
    ]
    rows = workbook_rows(xlsx, "待判定候選", required_columns)
    assert_unique_ids(rows, "occurrence_id")
    assert_unique_ids(rows, "review_id")

    db = load_or_initialize_db(output_dir)
    current_ledger = materialize_ledger(manifest, db)
    assert_unique_ids(current_ledger, "occurrence_id")
    assert_unique_ids(current_ledger, "review_id")
    by_review_id = {entry["review_id"]: entry for entry in current_ledger}
    if len(by_review_id) != len(current_ledger):
        raise DuplicateIdError("current ledger duplicate review_id")

    staged_events: dict[str, dict[str, Any]] = {}
    errors: list[str] = []
    safely_rebased = 0
    ignored_noop_rows = 0
    expected_only_actions = {"補建expected證據", "解決expected證據"}
    for row_number, row in enumerate(rows, 2):
        action = str(row.get("action") or "").strip()
        # A workbook can contain many untouched pending rows.  They are not part
        # of the transaction and must not make an otherwise valid import fail
        # merely because the project changed after export.
        if not action or action == "保留待人工":
            ignored_noop_rows += 1
            continue
        if action not in DECISIONS:
            errors.append(f"row {row_number}: 禁止的直接結論或未知 action {action}")
            continue

        occurrence_id = str(row.get("occurrence_id") or "").strip()
        review_id = str(row.get("review_id") or "").strip()
        entry = by_review_id.get(review_id)
        if entry is None:
            errors.append(f"row {row_number}: unknown review_id {review_id}")
            continue
        if occurrence_id != entry.get("occurrence_id"):
            errors.append(f"row {row_number}: review_id/occurrence_id mapping 不一致")
            continue

        full_snapshot_matches = str(row.get("review_snapshot") or "") == _review_snapshot(entry)
        if action in expected_only_actions:
            # v5.5.0: an expected judgment is bound to the printed occurrence
            # and its context, not to mutable actual or automatic-resolver output.
            # The expected asset fingerprint in workbook metadata protects the
            # source/rule version; expected_target_snapshot protects location and
            # context.  This removes the cross-lane stale coupling entirely.
            if not _expected_target_unchanged(row, entry):
                errors.append(
                    f"row {row_number}: stale expected target；PDF／頁碼／字元／詞境定位已改變，請重新覆核"
                )
                continue
            proposed_expected = tuple(normalize_expected_set(row.get("proposed_expected_set")))
            current_expected = tuple(normalize_expected_set(entry.get("expected_set")))
            if action == "補建expected證據" and infer_expected_status(entry) == "RESOLVED" and proposed_expected != current_expected:
                errors.append(
                    f"row {row_number}: expected 已由目前 resolver 建立為不同讀音；請改用『解決expected證據』並提供較高優先依據"
                )
                continue
            safely_rebased += int(not full_snapshot_matches)
            safe_expected_rebase = not full_snapshot_matches
        else:
            safe_expected_rebase = False
            if str(row.get("state") or "") != str(entry.get("state") or ""):
                errors.append(f"row {row_number}: stale state；此 action 依賴目前 actual/expected 狀態，請重新匯出")
                continue
            if not full_snapshot_matches:
                errors.append(f"row {row_number}: stale/modified review snapshot；此 action 需重新覆核目前證據")
                continue

        event = {
            "action": action,
            "expected_set": str(row.get("proposed_expected_set") or "").strip(),
            "expected_evidence": str(row.get("proposed_expected_evidence") or "").strip(),
            "context_evidence": str(row.get("proposed_context_evidence") or "").strip(),
            "exclusion_reason": str(row.get("exclusion_reason") or "").strip(),
            "exclusion_evidence": str(row.get("exclusion_evidence") or "").strip(),
            "note": str(row.get("note") or "").strip(),
            "source": "GPT／人工 expected 證據報表匯入（v5.5 dual-lane transaction）",
            "updated_at": datetime.now().isoformat(timespec="seconds"),
            "_safe_expected_rebase": safe_expected_rebase,
        }
        for gate in CONFIRMATION_GATES:
            value = str(row.get(gate) or "").strip().upper()
            if value not in {"", "Y", "N"}:
                errors.append(f"row {row_number}: gate {gate} 必須是 Y/N")
            event[gate] = value == "Y"
        staged_events[review_id] = event

    if errors:
        raise ValueError("GPT／人工匯入整批拒絕；未寫入任何一列：\n" + "\n".join(errors[:100]))

    staged_db = json.loads(json.dumps(db, ensure_ascii=False))
    staged_db.setdefault("events", {}).update(staged_events)
    # Full materialization validates every transition and terminal invariant
    # before the first persistent write.
    materialize_ledger(manifest, staged_db)
    if dry_run:
        # Bundle preflight: validate expected actions before any actual evidence
        # is committed.  This prevents avoidable half-applied bundles.
        return len(staged_events), ignored_noop_rows, output_dir / "注音校對_最終報告.xlsx"
    json_save(output_dir / "人工判定資料庫.json", staged_db)
    save_pending_json(output_dir, manifest, staged_db)
    report = generate_report(output_dir, manifest, staged_db)
    # skipped counts only explicit no-op/keep rows; safe rebases are successful
    # imports and remain included in len(staged_events).
    return len(staged_events), ignored_noop_rows, report


def unresolved_records(manifest: dict[str, Any], db: dict[str, Any]) -> list[dict[str, Any]]:
    return [entry for entry in materialize_ledger(manifest, db) if entry.get("state") in NON_TERMINAL_STATES]


def unresolved_coverage_gaps(manifest: dict[str, Any], db: dict[str, Any]) -> list[dict[str, Any]]:
    return [entry for entry in materialize_ledger(manifest, db) if entry.get("state") == "EXPECTED_UNRESOLVED"]


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def json_load(path: Path, default):
    if not path.exists():
        return default
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return default


def json_load_strict(path: Path) -> Any:
    if not path.exists():
        raise FileNotFoundError(path)

    def reject_duplicate_keys(pairs):
        out = {}
        for key, value in pairs:
            if key in out:
                raise ValueError(f"JSON 重複 key：{key}")
            out[key] = value
        return out

    try:
        return json.loads(path.read_text(encoding="utf-8"), object_pairs_hook=reject_duplicate_keys)
    except Exception as exc:
        raise ValueError(f"JSON 損壞或含重複 key：{path}") from exc


def json_save(path: Path, data) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)


def manifest_integrity_sha256(manifest: Mapping[str, Any]) -> str:
    payload = {key: value for key, value in dict(manifest).items() if key != "manifest_integrity_sha256"}
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def seal_manifest(manifest: dict[str, Any]) -> dict[str, Any]:
    manifest["manifest_integrity_sha256"] = manifest_integrity_sha256(manifest)
    return manifest


def validate_manifest_integrity(manifest: Mapping[str, Any]) -> None:
    recorded = str(manifest.get("manifest_integrity_sha256") or "")
    computed = manifest_integrity_sha256(manifest)
    if not recorded or recorded != computed:
        raise ValueError("DATA_INTEGRITY_ERROR：校對工作階段 payload hash 缺失或不符")


def validate_output_artifact_hashes(manifest: Mapping[str, Any], *, require_candidate: bool = True) -> None:
    for info in manifest.get("pdfs", []):
        checks = [("actual_workbook", "actual_workbook_sha256")]
        if require_candidate:
            checks.append(("candidate_workbook", "candidate_workbook_sha256"))
        for path_key, hash_key in checks:
            path = Path(str(info.get(path_key) or ""))
            recorded = str(info.get(hash_key) or "")
            if not path.exists() or not recorded or sha256_file(path) != recorded:
                raise ValueError(f"DATA_INTEGRITY_ERROR：輸出資產 hash 不符：{info.get('pdf_name')}/{path_key}")


def resolve_pdfs(input_path: Path) -> list[Path]:
    if input_path.is_file() and input_path.suffix.lower() == ".pdf":
        return [input_path.resolve()]
    if input_path.is_dir():
        return [p.resolve() for p in sorted(input_path.glob("*.pdf"))]
    return []


def summary_dict(workbook: Path) -> dict[str, Any]:
    wb = load_workbook(workbook, read_only=True, data_only=True)
    try:
        if "摘要" not in wb.sheetnames:
            raise ValueError(f"{workbook.name} 缺少必要工作表：摘要")
        ws = wb["摘要"]
        out = {}
        for row_number, row in enumerate(ws.iter_rows(min_row=1, values_only=True), 1):
            if row and row[0] is not None:
                key = str(row[0])
                if key in out:
                    raise ValueError(f"{workbook.name}/摘要 key 重複：{key} (row {row_number})")
                out[key] = row[1] if len(row) > 1 else None
        return out
    finally:
        wb.close()


def _summary_int(summary: Mapping[str, Any], key: str, *, missing: int = -1) -> int:
    """Read an integer summary field without mistaking a legitimate 0 for missing.

    ``value or -1`` is incorrect for zero-occurrence PDFs because numeric zero is
    falsey.  Missing/blank values still map to ``missing`` so the caller's
    integrity comparison fails closed.
    """
    value = summary.get(key)
    if value in (None, ""):
        return int(missing)
    return int(value)


def workbook_rows(path: Path, sheet: str, required_columns: list[str] | None = None) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"{path.name} 缺少必要工作表：{sheet}")
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        try:
            headers = [str(x or "") for x in next(rows)]
        except StopIteration as exc:
            raise ValueError(f"{path.name} 工作表 {sheet} 沒有 header") from exc
        duplicate_headers = [header for header, count in Counter(headers).items() if header and count > 1]
        if duplicate_headers:
            raise ValueError(f"{path.name} 工作表 {sheet} header 重複：{duplicate_headers}")
        if any(not header for header in headers):
            raise ValueError(f"{path.name} 工作表 {sheet} 含空白 header")
        missing = [column for column in (required_columns or []) if column not in headers]
        if missing:
            raise ValueError(f"{path.name} 工作表 {sheet} 缺少欄位：{missing}")
        out = []
        for vals in rows:
            if not any(v not in (None, "") for v in vals):
                continue
            out.append({headers[i]: vals[i] for i in range(min(len(headers), len(vals)))})
        return out
    finally:
        wb.close()


def make_review_id(pdf_hash: str, category: str, row: dict[str, Any]) -> str:
    occurrence_id = str(row.get("occurrence_id") or "").strip()
    if not occurrence_id:
        raise ValueError("v5.5.x review_id 必須由 occurrence_id 建立；不得由分類／actual／expected 重算")
    return ledger_make_review_id(occurrence_id)


AUTHORITATIVE_VIEW_SHEETS = [
    "單音一致", "完整詞語一致", "差異候選", "低信心詞境候選", "多音仍待詞條判讀",
    "專名特殊項目略過", "現有注音未解碼", "未建立預期音", "規則衝突",
    "非獨立技術層", "結構偵測排除",
]


def _parse_ledger_rows(candidate: Path) -> list[dict[str, Any]]:
    rows = workbook_rows(candidate, "Occurrence Ledger", [
        "ledger_schema_version", "review_id_schema_version", "occurrence_id", "review_id", "active_review", "state",
        "pdf_sha256", "actual_status", "expected_set_json", "expected_status", "comparison_result",
        "blocking_state", "exclusion_state", "confirmation_gates_json", "source_record_json",
    ])
    ledger = []
    for row in rows:
        if not schema_compatible(row.get("ledger_schema_version"), LEDGER_SCHEMA_VERSION) or not review_id_schema_compatible(row.get("review_id_schema_version"), REVIEW_ID_SCHEMA_VERSION):
            raise ValueError("candidate ledger schema／review ID 不相容")
        try:
            expected_set = json.loads(str(row.get("expected_set_json") or "[]"))
            gates = json.loads(str(row.get("confirmation_gates_json") or "{}"))
            source_record = json.loads(str(row.get("source_record_json") or "{}"))
        except Exception as exc:
            raise ValueError(f"candidate ledger JSON 欄位損壞：{row.get('occurrence_id')}") from exc
        entry = {
            key: value for key, value in row.items()
            if key not in {"expected_set_json", "confirmation_gates_json", "source_record_json"}
        }
        entry["active_review"] = bool(row.get("active_review"))
        entry["identity_row_fallback"] = str(row.get("identity_row_fallback") or "").upper() in {"TRUE", "Y", "1"}
        entry["expected_set"] = expected_set
        entry["confirmation_gates"] = gates
        entry["source_record"] = source_record
        entry["source_sheet"] = entry.get("source_view") or ""
        entry["row"] = source_record
        ledger.append(entry)
    # A split PDF such as a cutter/die file may legitimately contain no zhuyin
    # occurrences. Empty per-PDF candidate ledgers are valid input to whole-book
    # reconciliation; non-empty ledgers retain the strict validator.
    if ledger:
        validate_occurrence_ledger(ledger)
    return ledger


def _actual_source_ids(actual: Path) -> tuple[set[str], int, int, dict[str, dict[str, Any]]]:
    actual_rows = workbook_rows(actual, "實際注音", [
        "occurrence_id", "review_id", "ledger_schema_version", "workbook_schema_version", "實際注音", "解碼依據",
    ])
    excluded_rows = workbook_rows(actual, "結構偵測排除", [
        "occurrence_id", "review_id", "ledger_schema_version", "workbook_schema_version", "來源", "排除理由",
    ])
    for row in actual_rows + excluded_rows:
        if not schema_compatible(row.get("ledger_schema_version"), LEDGER_SCHEMA_VERSION) or not schema_compatible(row.get("workbook_schema_version"), WORKBOOK_SCHEMA_VERSION):
            raise ValueError("actual occurrence source schema 不相容")
    all_rows = actual_rows + excluded_rows
    assert_unique_ids(all_rows, "occurrence_id")
    assert_unique_ids(all_rows, "review_id")
    source_by_id = {}
    for row in actual_rows:
        source_by_id[str(row["occurrence_id"])] = {
            "kind": "actual",
            "review_id": str(row.get("review_id") or ""),
            "actual": canonical_bopomofo(row.get("實際注音")),
            "actual_evidence": str(row.get("解碼依據") or ""),
        }
    for row in excluded_rows:
        reason = "｜".join(str(row.get(key) or "") for key in ("來源", "排除理由") if row.get(key))
        source_by_id[str(row["occurrence_id"])] = {
            "kind": "structural_exclusion",
            "review_id": str(row.get("review_id") or ""),
            "exclusion_reason": reason,
        }
    return {str(row["occurrence_id"]) for row in all_rows}, len(actual_rows), len(excluded_rows), source_by_id


def _regression_gate_flag(row: Mapping[str, Any]) -> bool:
    raw = str(row.get("計入完成門檻") or "").strip().upper()
    if raw:
        if raw not in {"Y", "N"}:
            raise ValueError(f"歷史回歸計入完成門檻欄位無效：{row.get('案例ID')} / {raw}")
        return raw == "Y"
    # Compatibility for in-memory legacy fixtures and pre-v5.6.2 rows. New
    # candidate workbooks require the explicit field.
    mode = str(row.get("適用性模式") or "hard_gate").strip().lower()
    return mode != "reference_only" and str(row.get("回歸結果") or "") != "NOT_APPLICABLE"


def _candidate_regression_report(candidate: Path) -> tuple[dict[str, Any], dict[str, Any], list[dict[str, Any]], list[dict[str, Any]]]:
    mandatory_cases = load_mandatory_cases(Path(__file__).with_name("mandatory_regression_cases.csv"))
    mandatory_rows = workbook_rows(candidate, "強制回歸測試", ["case_id", "result"])
    mandatory_results = [{str(key): value for key, value in row.items()} for row in mandatory_rows]
    mandatory = validate_regression_execution(mandatory_cases, mandatory_results)

    pdf_rows = workbook_rows(candidate, "回歸測試", [
        "案例ID", "適用性模式", "適用性狀態", "計入完成門檻", "定位命中數", "回歸結果",
    ])
    ids = [str(row.get("案例ID") or "").strip() for row in pdf_rows]
    duplicate_ids = [key for key, count in Counter(ids).items() if key and count > 1]
    gating_rows = [row for row in pdf_rows if _regression_gate_flag(row)]
    reference_rows = [row for row in pdf_rows if str(row.get("適用性模式") or "").strip().lower() == "reference_only"]
    required_ids = {
        str(row.get("案例ID") or "").strip()
        for row in gating_rows
        if str(row.get("案例ID") or "").strip()
    }
    executed = [row for row in gating_rows if str(row.get("回歸結果") or "") in {"PASS", "FAIL"}]
    failed = [row for row in gating_rows if str(row.get("回歸結果") or "") == "FAIL"]
    not_executed = [row for row in gating_rows if str(row.get("回歸結果") or "") == "NOT_EXECUTED"]
    pdf_report = {
        "ok": not duplicate_ids and not failed,
        "required": len(required_ids),
        "executed": len(executed),
        "passed": sum(1 for row in gating_rows if row.get("回歸結果") == "PASS"),
        "failed": len(failed),
        "not_executed": len(not_executed),
        "not_applicable": sum(1 for row in pdf_rows if row.get("回歸結果") == "NOT_APPLICABLE"),
        "duplicate_case_id": len(duplicate_ids),
        "duplicate_case_ids": duplicate_ids,
        "reference_audit_matched": sum(1 for row in reference_rows if row.get("回歸結果") in {"PASS", "FAIL"}),
        "reference_audit_passed": sum(1 for row in reference_rows if row.get("回歸結果") == "PASS"),
        "reference_audit_failed": sum(1 for row in reference_rows if row.get("回歸結果") == "FAIL"),
        "reference_audit_not_found": sum(1 for row in reference_rows if row.get("回歸結果") == "NOT_APPLICABLE"),
    }
    return mandatory, pdf_report, mandatory_results, pdf_rows


def _aggregate_pdf_regression_rows(per_pdf_rows: list[tuple[str, list[dict[str, Any]]]]) -> dict[str, Any]:
    """Aggregate split-PDF historical regressions by unique case ID.

    Each candidate workbook intentionally carries the same applicable historical
    definitions so that a missing occurrence in one split remains auditable. At
    whole-book level a case is required once, not once per split. Exactly one
    split must execute each gating case; the remaining split-local copies may be
    NOT_EXECUTED. Reference-only copies are aggregated as non-gating audit
    results, including matched PASS/FAIL. Historical truth and locator
    definitions must be identical across copies.
    """
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    blank_case_rows: list[dict[str, Any]] = []
    for pdf_name, rows in per_pdf_rows:
        for row in rows:
            item = dict(row)
            item["pdf_name"] = pdf_name
            case_id = str(item.get("案例ID") or "").strip()
            if not case_id:
                blank_case_rows.append(item)
                continue
            grouped[case_id].append(item)

    if blank_case_rows:
        raise ValueError(f"歷史回歸含空白案例ID：{len(blank_case_rows)} 筆")

    signature_fields = (
        "課本頁", "完整詞語", "字元", "目標在詞內序號", "定位上下文", "真值實際注音", "真值預期注音",
        "真值結論", "適用性模式", "來源PDF SHA-256", "來源", "備註",
    )
    definition_conflicts: list[str] = []
    duplicate_execution_case_ids: list[str] = []
    duplicate_reference_case_ids: list[str] = []
    result_rows: list[dict[str, Any]] = []
    for case_id, copies in sorted(grouped.items()):
        signatures = {
            tuple(str(row.get(field) or "") for field in signature_fields)
            for row in copies
        }
        if len(signatures) != 1:
            definition_conflicts.append(case_id)
            continue
        base = copies[0]
        mode = str(base.get("適用性模式") or "hard_gate").strip().lower()
        if mode == "reference_only":
            executed = [row for row in copies if str(row.get("回歸結果") or "") in {"PASS", "FAIL"}]
            gate_flag = "N"
            if len(executed) > 1:
                duplicate_reference_case_ids.append(case_id)
                result = "FAIL"
                state = "REFERENCE_DUPLICATE_MATCH"
                note = "reference-only 在多個 PDF 分檔重現；identity ambiguity：" + ", ".join(row["pdf_name"] for row in executed)
                executed_pdf = "|".join(row["pdf_name"] for row in executed)
            elif len(executed) == 1:
                result = str(executed[0].get("回歸結果") or "")
                state = str(executed[0].get("適用性狀態") or "REFERENCE_MATCHED")
                note = str(executed[0].get("執行說明") or "")
                executed_pdf = executed[0]["pdf_name"]
            else:
                result = "NOT_APPLICABLE"
                state = "REFERENCE_NOT_FOUND"
                note = "reference not found：全冊所有 PDF 分檔均未定位此歷史 occurrence"
                executed_pdf = ""
        else:
            gating_copies = [
                row for row in copies
                if _regression_gate_flag(row)
            ]
            executed = [row for row in gating_copies if str(row.get("回歸結果") or "") in {"PASS", "FAIL"}]
            gate_flag = "Y" if gating_copies else "N"
            if not gating_copies:
                result = "NOT_APPLICABLE"
                state = "SOURCE_NOT_APPLICABLE"
                note = str(copies[0].get("執行說明") or "此歷史案例不適用於現版來源")
                executed_pdf = ""
            elif len(executed) > 1:
                duplicate_execution_case_ids.append(case_id)
                result = "FAIL"
                state = "DUPLICATE_EXECUTION"
                note = "同一歷史回歸案例在多個 PDF 分檔被實際執行：" + ", ".join(row["pdf_name"] for row in executed)
                executed_pdf = "|".join(row["pdf_name"] for row in executed)
            elif len(executed) == 1:
                result = str(executed[0].get("回歸結果") or "")
                state = str(executed[0].get("適用性狀態") or "GATE_EXECUTED")
                note = str(executed[0].get("執行說明") or "")
                executed_pdf = executed[0]["pdf_name"]
            else:
                result = "NOT_EXECUTED"
                state = "GATE_NOT_EXECUTED"
                note = "全冊所有 PDF 分檔均未實際定位／執行此歷史回歸案例"
                executed_pdf = ""
        result_rows.append({
            "案例ID": case_id,
            **{field: base.get(field, "") for field in signature_fields},
            "適用性狀態": state,
            "計入完成門檻": gate_flag,
            "定位命中數": sum(int(row.get("定位命中數") or 0) for row in executed),
            "回歸結果": result,
            "執行PDF": executed_pdf,
            "執行說明": note,
            "分檔副本數": len(copies),
        })

    if definition_conflicts:
        raise ValueError("歷史回歸案例定義在 PDF 分檔間不一致：" + ", ".join(definition_conflicts[:20]))

    gate_rows = [row for row in result_rows if row["計入完成門檻"] == "Y"]
    reference_rows = [row for row in result_rows if str(row.get("適用性模式") or "").lower() == "reference_only"]
    required = len(gate_rows)
    executed_count = sum(1 for row in gate_rows if row["回歸結果"] in {"PASS", "FAIL"})
    passed = sum(1 for row in gate_rows if row["回歸結果"] == "PASS")
    failed = sum(1 for row in gate_rows if row["回歸結果"] == "FAIL")
    not_executed = sum(1 for row in gate_rows if row["回歸結果"] == "NOT_EXECUTED")
    not_applicable = sum(1 for row in result_rows if row["回歸結果"] == "NOT_APPLICABLE")
    duplicate_count = len(duplicate_execution_case_ids)
    return {
        "ok": executed_count == required and failed == 0 and not_executed == 0 and duplicate_count == 0,
        "required": required,
        "executed": executed_count,
        "passed": passed,
        "failed": failed,
        "not_executed": not_executed,
        "not_applicable": not_applicable,
        "duplicate_case_id": duplicate_count,
        "duplicate_execution_case_ids": duplicate_execution_case_ids,
        "reference_duplicate_case_id": len(duplicate_reference_case_ids),
        "duplicate_reference_case_ids": duplicate_reference_case_ids,
        "reference_audit_total": len(reference_rows),
        "reference_audit_matched": sum(1 for row in reference_rows if row["回歸結果"] in {"PASS", "FAIL"}),
        "reference_audit_passed": sum(1 for row in reference_rows if row["回歸結果"] == "PASS"),
        "reference_audit_failed": sum(1 for row in reference_rows if row["回歸結果"] == "FAIL"),
        "reference_audit_not_found": sum(1 for row in reference_rows if row["回歸結果"] == "NOT_APPLICABLE"),
        "required_case_ids": sorted(row["案例ID"] for row in gate_rows),
        "not_applicable_case_ids": sorted(row["案例ID"] for row in result_rows if row["回歸結果"] == "NOT_APPLICABLE"),
        "executed_case_ids": sorted(row["案例ID"] for row in gate_rows if row["回歸結果"] in {"PASS", "FAIL"}),
        "results": result_rows,
    }


def _carry_forward_cross_version_identity(manifest: dict[str, Any], baseline: Mapping[str, Any] | None) -> dict[str, Any]:
    """Keep project session identity and compatible review artifacts across releases."""
    if not baseline:
        return seal_manifest(manifest)
    old_session = str(baseline.get("session_id") or "").strip()
    if old_session:
        manifest["session_id"] = old_session
    manifest["upgraded_from_version"] = str(baseline.get("version") or "")

    compatible = {
        str(x) for x in (baseline.get("compatible_expected_asset_fingerprints") or []) if str(x or "").strip()
    }
    old_fp = str(baseline.get("expected_asset_fingerprint") or "").strip()
    current_payload = {
        "fingerprint": manifest.get("expected_asset_fingerprint"),
        "components": manifest.get("expected_asset_fingerprint_components") or {},
    }
    if old_fp and fingerprint_compatible(
        old_fp, baseline.get("expected_asset_fingerprint_components"), current_payload, chain="expected"
    ):
        compatible.add(old_fp)
    current_fp = str(manifest.get("expected_asset_fingerprint") or "").strip()
    if current_fp:
        compatible.add(current_fp)
    manifest["compatible_expected_asset_fingerprints"] = sorted(compatible)
    return seal_manifest(manifest)


def collect_manifest(
    pdfs: list[Path],
    actual_dir: Path,
    cand_dir: Path,
    *,
    actual_fingerprints: Mapping[str, Mapping[str, Any]],
    source_validation: Mapping[str, Any],
    runtime_root: Path | None = None,
) -> dict[str, Any]:
    root = Path(runtime_root or Path(__file__).resolve().parent).resolve()
    manifest = {
        "version": VERSION,
        "session_schema_version": SESSION_SCHEMA_VERSION,
        "workbook_schema_version": WORKBOOK_SCHEMA_VERSION,
        "ledger_schema_version": LEDGER_SCHEMA_VERSION,
        "review_id_schema_version": REVIEW_ID_SCHEMA_VERSION,
        "session_id": str(uuid4()),
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "source_validation_ok": bool(source_validation.get("ok")),
        "source_validation": dict(source_validation),
        "dynamic_actual_evidence_root": str(project_actual_evidence_root(actual_dir.parent)),
        "pdfs": [],
        "records": [],
        "reconciliation": {},
    }
    expected_chain_fingerprint = compute_expected_asset_fingerprint(
        root,
        source_validation,
        resolver_version=EXPECTED_RESOLVER_VERSION,
        source_files=EXPECTED_RESOLVER_SOURCE_FILES,
    )
    manifest["expected_asset_fingerprint"] = expected_chain_fingerprint["fingerprint"]
    manifest["expected_asset_fingerprint_components"] = expected_chain_fingerprint["components"]
    reusable_rules = load_reusable_expected_rules(root / REUSABLE_EXPECTED_RULES.name)
    manifest["reusable_expected_rules"] = reusable_rules
    manifest["reusable_expected_rules_sha256"] = reusable_rules_sha256(reusable_rules)
    all_actual_ids: set[str] = set()
    worksheet_counts: dict[str, tuple[int, int]] = {}
    mandatory_reports = []
    pdf_regression_reports = []
    pdf_regression_rows_by_pdf: list[tuple[str, list[dict[str, Any]]]] = []
    for pdf in pdfs:
        ph = sha256_file(pdf)
        actual = actual_workbook_path(actual_dir, pdf)
        cand = candidate_workbook_path(cand_dir, pdf)
        if not actual.exists() or not cand.exists():
            raise FileNotFoundError(f"缺少輸出：{pdf.name}")
        actual_metadata = workbook_metadata(actual)
        actual_fingerprint_payload = dict(actual_fingerprints[pdf.name])
        actual_fingerprint = str(actual_fingerprint_payload.get("fingerprint") or "")
        if not fingerprint_compatible(
            actual_metadata.get("actual_asset_fingerprint"),
            actual_metadata.get("actual_asset_fingerprint_components"),
            actual_fingerprint_payload,
            chain="actual",
        ):
            raise ValueError(f"actual evidence assets 不符：{pdf.name}")
        if actual_metadata.get("pdf_sha256") != ph:
            raise ValueError(f"actual PDF hash 不符：{pdf.name}")
        actual_ids, actual_row_count, structural_exclusion_count, actual_source_by_id = _actual_source_ids(actual)
        ledger = _parse_ledger_rows(cand)
        for entry in ledger:
            if entry.get("pdf_sha256") != ph:
                raise ValueError(f"candidate ledger PDF hash 不符：{entry.get('occurrence_id')}")
            entry["pdf"] = str(pdf)
            entry["pdf_name"] = pdf.name
            entry["candidate_workbook"] = str(cand)
            entry["actual_workbook"] = str(actual)
            source = actual_source_by_id.get(str(entry.get("occurrence_id") or ""))
            if source is None:
                raise ValueError(f"candidate occurrence 在 actual source 不存在：{entry.get('occurrence_id')}")
            if entry.get("review_id") != source.get("review_id"):
                raise ValueError(f"actual/candidate review_id mapping 不一致：{entry.get('occurrence_id')}")
            if source.get("kind") == "actual":
                if canonical_bopomofo(entry.get("actual")) != source.get("actual"):
                    raise ValueError(f"actual/candidate 讀音值漂移：{entry.get('occurrence_id')}")
                if str(entry.get("actual_evidence") or "") != source.get("actual_evidence"):
                    raise ValueError(f"actual/candidate 解碼證據漂移：{entry.get('occurrence_id')}")
            else:
                if entry.get("state") != "EXCLUDED_OUT_OF_SCOPE" or entry.get("source_view") != "結構偵測排除":
                    raise ValueError(f"結構排除 occurrence 被改列其他狀態：{entry.get('occurrence_id')}")
                if str(entry.get("exclusion_reason") or "") != source.get("exclusion_reason"):
                    raise ValueError(f"actual/candidate 結構排除證據漂移：{entry.get('occurrence_id')}")
        ledger_ids = {str(entry["occurrence_id"]) for entry in ledger}
        if all_actual_ids & actual_ids:
            raise DuplicateIdError("跨 PDF occurrence_id 碰撞")
        all_actual_ids.update(actual_ids)

        view_ids: set[str] = set()
        expected_views = Counter(str(entry.get("source_view") or "") for entry in ledger)
        for sheet in AUTHORITATIVE_VIEW_SHEETS:
            view_rows = workbook_rows(cand, sheet, ["occurrence_id"])
            observed_ids = [str(row.get("occurrence_id") or "") for row in view_rows]
            if len(observed_ids) != len(set(observed_ids)):
                raise DuplicateIdError(f"{cand.name}/{sheet} occurrence_id 重複")
            overlap = view_ids & set(observed_ids)
            if overlap:
                raise ValueError(f"candidate authoritative views 不互斥：{sheet} {sorted(overlap)[:10]}")
            view_ids.update(observed_ids)
            expected_ids = {str(entry["occurrence_id"]) for entry in ledger if entry.get("source_view") == sheet}
            worksheet_counts[f"{pdf.name}/{sheet}"] = (len(expected_ids), len(observed_ids))
            if set(observed_ids) != expected_ids:
                raise ValueError(f"candidate view 與 ledger 不一致：{pdf.name}/{sheet}")
        if view_ids != ledger_ids:
            raise ValueError(f"candidate views 聯集不等於 ledger：{pdf.name}")

        sm = summary_dict(actual)
        csm = summary_dict(cand)
        if not schema_compatible(csm.get("workbook_schema_version"), WORKBOOK_SCHEMA_VERSION) or not schema_compatible(csm.get("ledger_schema_version"), LEDGER_SCHEMA_VERSION):
            raise ValueError(f"candidate summary schema 不相容：{pdf.name}")
        if not fingerprint_compatible(
            csm.get("expected_asset_fingerprint"),
            csm.get("expected_asset_fingerprint_components"),
            expected_chain_fingerprint,
            chain="expected",
        ):
            raise ValueError(f"candidate expected evidence assets 已失效：{pdf.name}")
        summary_counts = {
            "actual summary total": (actual_row_count, _summary_int(sm, "偵測到注音字形筆數")),
            "candidate mother total": (len(ledger), _summary_int(csm, "PDF偵測母體總數")),
            "candidate in-scope total": (
                sum(1 for entry in ledger if entry.get("state") not in EXCLUDED_STATES),
                _summary_int(csm, "現版注音校對分母"),
            ),
        }
        for label, (expected, observed) in summary_counts.items():
            if expected != observed:
                raise ValueError(f"summary mismatch {pdf.name}/{label}: {expected}!={observed}")

        per_pdf_reconciliation = reconcile_ledger(ledger, actual_ids, worksheet_counts={
            key: value for key, value in worksheet_counts.items() if key.startswith(pdf.name + "/")
        }, summary_counts=summary_counts)
        if not per_pdf_reconciliation.ok:
            raise ValueError(f"集合對帳失敗 {pdf.name}：{'；'.join(per_pdf_reconciliation.errors)}")
        mandatory, pdf_regression, mandatory_rows, pdf_regression_rows = _candidate_regression_report(cand)
        observed_candidate_hash = candidate_payload_sha256(ledger, mandatory_rows, pdf_regression_rows)
        if csm.get("candidate_payload_sha256") != observed_candidate_hash:
            raise ValueError(f"candidate authoritative payload hash 不符：{pdf.name}")
        mandatory_reports.append(mandatory)
        pdf_regression_reports.append(pdf_regression)
        pdf_regression_rows_by_pdf.append((pdf.name, pdf_regression_rows))
        manifest["pdfs"].append({
            "pdf": str(pdf), "pdf_name": pdf.name, "pdf_sha256": ph,
            "actual_workbook": str(actual), "candidate_workbook": str(cand),
            "actual_workbook_sha256": sha256_file(actual),
            "candidate_workbook_sha256": sha256_file(cand),
            "actual_asset_fingerprint": actual_fingerprint,
            "actual_total": actual_row_count,
            "structural_exclusions": structural_exclusion_count,
            "ledger_total": len(ledger),
            "proofread_total": per_pdf_reconciliation.details["in_scope_count"],
            "mandatory_regression": mandatory,
            "pdf_regression": pdf_regression,
        })
        manifest["records"].extend(ledger)

    assert_unique_ids(manifest["records"], "occurrence_id")
    assert_unique_ids(manifest["records"], "review_id")
    # An all-production-artwork batch can legitimately contain zero zhuyin
    # occurrences.  Treat that as a valid empty dataset and let the completion
    # gate report in_scope_occurrence_count_positive=False instead of crashing.
    if manifest["records"]:
        validate_occurrence_ledger(manifest["records"])
    reconciliation = reconcile_ledger(manifest["records"], all_actual_ids, worksheet_counts=worksheet_counts)
    if not reconciliation.ok:
        raise ValueError("全量集合對帳失敗：" + "；".join(reconciliation.errors))
    manifest["reconciliation"] = reconciliation.as_dict()
    manifest["actual_source_ids"] = sorted(all_actual_ids)

    first_mandatory = mandatory_reports[0] if mandatory_reports else {
        "ok": False, "required": 0, "executed": 0, "passed": 0, "failed": 0, "not_executed": 0, "duplicate_case_id": 0
    }
    same_mandatory = all(
        {key: report.get(key) for key in ("ok", "required", "executed", "passed", "failed", "not_executed", "duplicate_case_id")}
        == {key: first_mandatory.get(key) for key in ("ok", "required", "executed", "passed", "failed", "not_executed", "duplicate_case_id")}
        for report in mandatory_reports
    )
    if not same_mandatory:
        raise ValueError("各 PDF mandatory regression 結果不一致")
    # Whole-book historical regressions are unique cases.  The same case may
    # appear as NOT_EXECUTED in all non-owning split PDFs, but it is required
    # only once and must execute exactly once somewhere in the book.
    combined_pdf = _aggregate_pdf_regression_rows(pdf_regression_rows_by_pdf)
    manifest["mandatory_regression"] = first_mandatory
    manifest["pdf_regression"] = combined_pdf
    manifest["regression_gate"] = {
        "ok": bool(first_mandatory.get("ok")) and bool(combined_pdf.get("ok")),
        "required": int(first_mandatory.get("required", 0)) + int(combined_pdf.get("required", 0)),
        "executed": int(first_mandatory.get("executed", 0)) + int(combined_pdf.get("executed", 0)),
        "passed": int(first_mandatory.get("passed", 0)) + int(combined_pdf.get("passed", 0)),
        "failed": int(first_mandatory.get("failed", 0)) + int(combined_pdf.get("failed", 0)),
        "not_executed": int(first_mandatory.get("not_executed", 0)) + int(combined_pdf.get("not_executed", 0)),
        "not_applicable": int(combined_pdf.get("not_applicable", 0)),
        "duplicate_case_id": int(first_mandatory.get("duplicate_case_id", 0)) + int(combined_pdf.get("duplicate_case_id", 0)),
    }
    return seal_manifest(manifest)


def normalize_db(db: dict[str, Any]) -> dict[str, Any]:
    if not db:
        return {
            "version": VERSION,
            "session_schema_version": SESSION_SCHEMA_VERSION,
            "review_id_schema_version": REVIEW_ID_SCHEMA_VERSION,
            "events": {},
        }
    if (
        not schema_compatible(db.get("session_schema_version"), SESSION_SCHEMA_VERSION)
        or not review_id_schema_compatible(db.get("review_id_schema_version"), REVIEW_ID_SCHEMA_VERSION)
        or not isinstance(db.get("events"), dict)
    ):
        raise ValueError("DECISION_SCHEMA_INCOMPATIBLE：資料布局或 review ID 規則不同，需有明確轉接器")
    # Cross-release reuse is intentional. Preserve all durable occurrence events
    # and merely stamp the current tool version for subsequent writes.
    if db.get("version") != VERSION:
        db = dict(db)
        old_version = str(db.get("version") or "")
        db["version"] = VERSION
        db.setdefault("migrated_from_version", old_version)
        db["session_schema_version"] = SESSION_SCHEMA_VERSION
        db["review_id_schema_version"] = REVIEW_ID_SCHEMA_VERSION
    return db

def load_or_initialize_db(output_dir: Path) -> dict[str, Any]:
    path = output_dir / "人工判定資料庫.json"
    if not path.exists():
        return normalize_db({})
    raw = json_load_strict(path)
    if not isinstance(raw, dict):
        raise ValueError("DATA_INTEGRITY_ERROR：現版人工判定資料庫根節點必須是物件")
    appears_current = any(
        raw.get(key) == value
        for key, value in (
            ("version", VERSION),
            ("session_schema_version", SESSION_SCHEMA_VERSION),
            ("review_id_schema_version", REVIEW_ID_SCHEMA_VERSION),
        )
    )
    if appears_current:
        # A current-looking but invalid database is corruption, not legacy.
        # Never quarantine-and-reset it silently.
        return normalize_db(raw)
    try:
        return normalize_db(raw)
    except ValueError:
        orphan_path = output_dir / "legacy_orphan_decisions.json"
        existing = json_load(orphan_path, {"legacy_items": []})
        if not isinstance(existing, dict):
            existing = {"legacy_items": []}
        existing.setdefault("legacy_items", []).append({
            "captured_at": datetime.now().isoformat(timespec="seconds"),
            "source": str(path),
            "reason": "v5.1.5 schema 無法唯一映射；只供稽核，不得取得 v5.5.x terminal status",
            "payload": raw,
        })
        json_save(orphan_path, existing)
        return normalize_db({})


def _apply_review_event(entry: dict[str, Any], event: Mapping[str, Any]) -> dict[str, Any]:
    action = str(event.get("action") or "").strip()
    if action == "保留待人工":
        return dict(entry)
    if action in {"補建expected證據", "解決expected證據"}:
        # v5.5.0: expected events write only the expected lane.  They never
        # consult current actual to decide *what* expected should be, and they
        # never fail merely because the public composite state is ACTUAL_*.
        expected_set = list(normalize_expected_set(event.get("expected_set")))
        evidence = str(event.get("expected_evidence") or "").strip()
        context = str(event.get("context_evidence") or entry.get("context_evidence") or entry.get("char") or "").strip()
        if not expected_set or not evidence or not context:
            raise InvalidTransitionError("expected 必須提供讀音、獨立來源證據與詞語／語境／位置證據")

        current_expected = list(normalize_expected_set(entry.get("expected_set")))
        if (
            action == "補建expected證據"
            and infer_expected_status(entry) == "RESOLVED"
            and current_expected != expected_set
        ):
            # ``補建`` fills an unresolved expected lane; it is not an override.
            # When a newer resolver now supplies a different independent truth,
            # retain the durable event for audit but never overwrite that truth.
            # The current mechanically-derived state remains authoritative and,
            # when it is still a difference, naturally stays fail-closed.
            out = dict(entry)
            out["review_event_replay_status"] = "INVALIDATED_EXPECTED_DRIFT"
            out["review_event_replay_note"] = (
                f"舊補建 expected={expected_set}；現版 resolver expected={current_expected}"
            )
            return out

        state = str(entry.get("state") or "")
        if state in {"EXCLUDED_NONINDEPENDENT_LAYER", "EXCLUDED_OUT_OF_SCOPE"}:
            out = dict(entry)
            out["review_event_replay_status"] = "EXPECTED_EVENT_DORMANT_EXCLUDED"
            return out

        out = dict(entry)
        out.update({
            "previous_expected_set": list(entry.get("expected_set") or []),
            "previous_expected_evidence": str(entry.get("expected_evidence") or ""),
            "expected_set": expected_set,
            "expected_evidence": evidence,
            "expected_status": "RESOLVED",
            "context_evidence": context,
            "expected_resolution_source": str(event.get("source") or "人工現版證據"),
            "expected_resolution_reason": str(event.get("resolution_reason") or "").strip(),
            "note": str(event.get("note") or "").strip(),
        })
        # Integrity/source blocks remain authoritative workflow blockers, but the
        # expected lane is still retained for audit and can be reused after the
        # blocker is repaired. Regression-blocked rows are treated the same way.
        if state in {"SOURCE_INVALID", "DATA_INTEGRITY_ERROR", "REGRESSION_BLOCKED"}:
            out["blocking_state"] = state
        out = refresh_derived_state(out, preserve_confirmed=False)
        out["expected_resolution_deferred"] = infer_actual_status(out) != "RESOLVED"
        out["expected_resolution_deferred_reason"] = (
            out.get("state") if out["expected_resolution_deferred"] else ""
        )
        return out
    if action == "確認現版差異":
        # The decision DB intentionally keeps one authoritative event per review_id.
        # A common two-step GUI path first writes an expected-evidence event and then
        # confirms the resulting mismatch.  The second write replaces the first one,
        # so the confirmation event must carry enough independent expected evidence
        # to reconstruct that intermediate DIFFERENCE state from the frozen manifest.
        working = dict(entry)
        event_expected = list(normalize_expected_set(event.get("expected_set")))
        current_expected = list(normalize_expected_set(working.get("expected_set")))
        if (
            infer_expected_status(working) == "RESOLVED"
            and event_expected
            and current_expected != event_expected
        ):
            # A six-gate conclusion is bound to the expected truth reviewed at
            # that time.  Rebuilding candidates with a different independent
            # expected truth invalidates only the terminal conclusion; it must
            # not resurrect the embedded historical expected evidence.
            working["review_event_replay_status"] = "INVALIDATED_EXPECTED_DRIFT"
            working["review_event_replay_note"] = (
                f"舊確認 expected={event_expected}；現版 resolver expected={current_expected}"
            )
            return working
        if working.get("state") != "DIFFERENCE_PENDING_CONFIRMATION":
            expected_set = event_expected
            evidence = str(event.get("expected_evidence") or "").strip()
            context = str(event.get("context_evidence") or working.get("context_evidence") or working.get("char") or "").strip()
            if expected_set and evidence and context:
                working = _apply_review_event(working, {
                    "action": "解決expected證據",
                    "expected_set": expected_set,
                    "expected_evidence": evidence,
                    "context_evidence": context,
                    "resolution_reason": str(event.get("expected_resolution_reason") or "").strip(),
                    "source": str(event.get("expected_resolution_source") or event.get("source") or "人工 GUI 現版 expected 證據"),
                    "note": str(event.get("expected_resolution_note") or "").strip(),
                })
            if working.get("state") != "DIFFERENCE_PENDING_CONFIRMATION":
                raise InvalidTransitionError(
                    "只有現版機械比較差異可進六閘門確認；若此差異由先前 expected 證據建立，"
                    "確認事件必須一併保留該 expected 讀音、來源與詞境證據"
                )
        gates = {gate: event.get(gate) is True for gate in CONFIRMATION_GATES}
        return transition_state(working, "TEXTBOOK_ERROR_CONFIRMED", {
            "confirmation_gates": gates,
            "note": str(event.get("note") or "").strip(),
        })
    if action == "確認非校對範圍":
        reason = str(event.get("exclusion_reason") or "").strip()
        evidence = str(event.get("exclusion_evidence") or "").strip()
        if not reason or not evidence:
            raise InvalidTransitionError("排除 occurrence 必須提供理由與可稽核證據")
        updates = {
            "exclusion_reason": f"{reason}｜證據：{evidence}",
            "note": str(event.get("note") or "").strip(),
        }
        try:
            return transition_state(entry, "EXCLUDED_OUT_OF_SCOPE", updates)
        except InvalidTransitionError:
            # Scope decisions are independent from actual/expected.  They may be
            # replayed after an actual-only refresh even if the base row has
            # mechanically become PASS.  Never let them override source/data
            # integrity failures or the technical non-independent-layer class.
            if entry.get("state") in {
                "PASS", "EXPECTED_UNRESOLVED", "EXPECTED_AMBIGUOUS", "REVIEW_PENDING",
                "RULE_CONFLICT", "DIFFERENCE_PENDING_CONFIRMATION", "ACTUAL_UNRESOLVED",
                "ACTUAL_DECODE_ERROR", "REGRESSION_BLOCKED", "EXCLUDED_OUT_OF_SCOPE",
            }:
                out = dict(entry)
                out.update(updates)
                out["state"] = "EXCLUDED_OUT_OF_SCOPE"
                out["active_review"] = False
                validate_terminal_state(out)
                return out
            raise
    raise InvalidTransitionError(f"未知或禁止的人工／GPT action：{action}")


def materialize_ledger(manifest: Mapping[str, Any], db: Mapping[str, Any]) -> list[dict[str, Any]]:
    db = normalize_db(dict(db))
    ledger = [dict(entry) for entry in manifest.get("records", [])]
    rules_doc = _canonical_rule_doc(manifest.get("reusable_expected_rules") or {})
    expected_rules_hash = str(manifest.get("reusable_expected_rules_sha256") or reusable_rules_sha256(rules_doc))
    if expected_rules_hash != reusable_rules_sha256(rules_doc):
        raise ValueError("DATA_INTEGRITY_ERROR：工作階段可重用 expected 規則 snapshot hash 不符")
    ledger = _apply_reusable_rules(ledger, rules_doc)
    assert_unique_ids(ledger, "occurrence_id")
    assert_unique_ids(ledger, "review_id")
    by_review_id = {entry["review_id"]: index for index, entry in enumerate(ledger)}
    events = db.get("events", {})
    unknown = sorted(set(events) - set(by_review_id))
    if unknown:
        raise ValueError(f"人工事件含 unknown review_id：{unknown[:20]}")
    for review_id, event in events.items():
        if not isinstance(event, dict):
            raise ValueError(f"review event 格式錯誤：{review_id}")
        index = by_review_id[review_id]
        try:
            ledger[index] = _apply_review_event(ledger[index], event)
        except InvalidTransitionError as exc:
            raise InvalidTransitionError(
                f"review event replay 失敗：review_id={review_id}；"
                f"action={str(event.get('action') or '')}；"
                f"base_state={str(ledger[index].get('state') or '')}；{exc}"
            ) from exc
    if ledger:
        validate_occurrence_ledger(ledger)
    return ledger


def save_pending_json(output_dir: Path, manifest: dict[str, Any], db: dict[str, Any]) -> int:
    ledger = materialize_ledger(manifest, db)
    pending = [dict(entry) for entry in ledger if entry.get("state") in NON_TERMINAL_STATES]
    gaps = [dict(entry) for entry in ledger if infer_expected_status(entry) == "UNRESOLVED"]
    json_save(output_dir / "待人工確認.json", {
        "version": VERSION,
        "session_schema_version": SESSION_SCHEMA_VERSION,
        "ledger_schema_version": LEDGER_SCHEMA_VERSION,
        "pending": pending,
        "expected_gaps": gaps,
    })
    return len(pending)


def _style_sheet(ws, freeze="A2"):
    """Apply workbook styling without allocating a new style object per cell.

    The v5.2 ledger report can exceed one million populated cells across the
    authoritative ledger and mutually-exclusive state views. Reusing immutable
    openpyxl style objects keeps report generation bounded without changing any
    data, state, evidence, or completion result.
    """
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = freeze
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_alignment = Alignment(vertical="top", wrap_text=True)
    for c in ws[1]:
        c.fill = header_fill
        c.font = header_font
        c.alignment = header_alignment
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.value is not None:
                c.alignment = body_alignment
    for col in range(1, ws.max_column + 1):
        max_len = 0
        for r in range(1, min(ws.max_row, 300) + 1):
            v = ws.cell(r, col).value
            if v is not None:
                max_len = max(max_len, min(70, len(str(v))))
        ws.column_dimensions[get_column_letter(col)].width = max(10, min(50, max_len + 2))


def core_record(rec: dict[str, Any], decision: str = "", note: str = "", corrected_actual: str = "", decision_source: str = "", expected_override: str = "") -> list[Any]:
    r = rec["row"]
    return [
        rec["pdf_name"], r.get("課本頁"), r.get("實體頁碼"), r.get("字元"),
        r.get("所在行"), r.get("局部詞境"), r.get("實際注音"),
        expected_override or r.get("預期注音") or r.get("字典有效讀音") or r.get("規則候選讀音"),
        rec["source_sheet"], decision, decision_source, corrected_actual, note,
        r.get("穩定注音鍵"), r.get("font"), r.get("注音元件ID"),
        r.get("x0"), r.get("y0"), r.get("x1"), r.get("y1"), rec["review_id"],
    ]


def _generate_report_legacy_disabled(output_dir: Path, manifest: dict[str, Any], db: dict[str, Any]) -> Path:
    """v5.1.5 renderer retained only for historical diff review; never called."""
    raise ValueError("LEGACY_SCHEMA：v5.1.5 報告器已隔離，不得執行")
    db = normalize_db(db)
    apply_decision_memory(manifest, db)
    json_save(output_dir / "人工判定資料庫.json", db)
    decisions = db["decisions"]

    auto_errors: list[dict[str, Any]] = []
    auto_correct: list[dict[str, Any]] = []
    specials: list[dict[str, Any]] = []
    manual_errors: list[tuple[dict[str, Any], dict[str, Any]]] = []
    manual_correct: list[tuple[dict[str, Any], dict[str, Any]]] = []
    decoder_errors: list[tuple[dict[str, Any], dict[str, Any]]] = []
    pending: list[dict[str, Any]] = []
    coverage_gaps: list[dict[str, Any]] = []
    resolved_expected: list[tuple[dict[str, Any], dict[str, Any]]] = []
    all_rows: list[list[Any]] = []

    for rec in manifest.get("records", []):
        sheet = rec["source_sheet"]
        d = decisions.get(rec["review_id"], {})
        decision = d.get("decision", "") if isinstance(d, dict) else ""
        note = d.get("note", "") if isinstance(d, dict) else ""
        corr = d.get("corrected_actual", "") if isinstance(d, dict) else ""
        dsource = d.get("source", "人工確認") if isinstance(d, dict) and decision else ""
        exp_override = d.get("expected", "") if isinstance(d, dict) else ""

        if sheet in AUTO_ERROR_SHEETS:
            auto_errors.append(rec)
        elif sheet in AUTO_CORRECT_SHEETS:
            auto_correct.append(rec)
        elif sheet in SPECIAL_SHEETS:
            specials.append(rec)
        elif sheet in COVERAGE_GAP_SHEETS:
            if decision == "教材錯誤" and exp_override:
                manual_errors.append((rec, d))
                resolved_expected.append((rec, d))
            elif decision == "課本正確" and exp_override:
                manual_correct.append((rec, d))
                resolved_expected.append((rec, d))
            elif decision == "特殊略過":
                specials.append(rec)
            elif decision == "解碼錯誤":
                decoder_errors.append((rec, d))
                # actual 已被判定不可信，在重新解碼前 expected 仍不能算完整。
                coverage_gaps.append(rec)
            else:
                coverage_gaps.append(rec)
        elif sheet in UNRESOLVED_SHEETS:
            if decision == "教材錯誤":
                manual_errors.append((rec, d))
            elif decision == "課本正確":
                manual_correct.append((rec, d))
            elif decision == "特殊略過":
                specials.append(rec)
            elif decision == "解碼錯誤":
                decoder_errors.append((rec, d))
            else:
                pending.append(rec)

        all_rows.append(core_record(rec, decision, note, corr, dsource, exp_override))

    total = sum(int(x.get("actual_total", 0)) for x in manifest.get("pdfs", []))
    mapped = sum(int(x.get("actual_mapped", 0)) for x in manifest.get("pdfs", []))
    proofread_total = sum(int(x.get("proofread_total", 0)) for x in manifest.get("pdfs", []))
    regression_fail = sum(int(x.get("regression_fail", 0)) for x in manifest.get("pdfs", []))
    operational_finished = (total > 0 and mapped == total and len(pending) == 0 and len(decoder_errors) == 0 and regression_fail == 0)
    finished = operational_finished and len(coverage_gaps) == 0

    wb = Workbook()
    ws = wb.active
    ws.title = "摘要"
    ws.append(["指標", "結果"])
    summary = [
        ("工具", f"{PROGRAM} v{VERSION}"),
        ("PDF檔數", len(manifest.get("pdfs", []))),
        ("實際注音結構總數", total),
        ("成功解碼數", mapped),
        ("實際注音解碼率", mapped / total if total else 0),
        ("現版校對座標總數", proofread_total),
        ("已補建預期音", len(resolved_expected)),
        ("未建立預期音", len(coverage_gaps)),
        ("預期音資料覆蓋率", ((proofread_total - len(coverage_gaps)) / proofread_total) if proofread_total else 0),
        ("回歸測試失敗", regression_fail),
        ("既有確認教材錯誤", len(auto_errors)),
        ("已確認判定教材錯誤", len(manual_errors)),
        ("教材錯誤合計", len(auto_errors) + len(manual_errors)),
        ("既有確認正確回歸", len(auto_correct)),
        ("已確認判定課本正確", len(manual_correct)),
        ("特殊略過", len(specials)),
        ("解碼錯誤待修正", len(decoder_errors)),
        ("尚待人工確認", len(pending)),
        ("候選收斂狀態", "已收斂" if operational_finished else "尚未收斂"),
        ("校對狀態", "完成" if finished else "尚未完整完成"),
        ("完成判定", "只有現版實際注音100%安全解碼、現版預期音覆蓋缺口為0、待人工確認與解碼錯誤均為0、且回歸測試0失敗，才標示全冊校對完成。"),
        ("預期音補建原則", "未建立預期音可由人工或GPT報表補建 expected；程式只做 actual-vs-expected 比較。GPT補建結果只寫入目前工作階段，不自動成為跨教材永久規則。"),
        ("歷史資料限制", "前版相同、前次曾判正確／錯誤、或注音外觀未變，只能作回歸與定位資訊，不能替代現版獨立實際音＋預期音判定。"),
        ("建立時間", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for x in summary:
        ws.append(x)
    for rr in range(2, ws.max_row + 1):
        if ws.cell(rr, 1).value in {"實際注音解碼率", "預期音資料覆蓋率"}:
            ws.cell(rr, 2).number_format = "0.00%"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 95

    gap_by_pdf = Counter(r.get("pdf_name", "") for r in coverage_gaps)
    resolved_by_pdf = Counter(r.get("pdf_name", "") for r, _ in resolved_expected)
    wf = wb.create_sheet("檔案統計")
    wf.append(["PDF","SHA-256","實際注音總數","成功解碼","解碼覆蓋率","現版校對座標","已補建預期音","未建立預期音","預期音覆蓋率","回歸失敗","實際注音工作簿","候選工作簿"])
    for x in manifest.get("pdfs", []):
        t = int(x.get("actual_total", 0))
        m = int(x.get("actual_mapped", 0))
        pt = int(x.get("proofread_total", t))
        gap = int(gap_by_pdf.get(x.get("pdf_name", ""), 0))
        exp_cov = (pt - gap) / pt if pt else 0
        wf.append([x["pdf_name"],x["pdf_sha256"],t,m,m/t if t else 0,pt,int(resolved_by_pdf.get(x.get("pdf_name", ""),0)),gap,exp_cov,int(x.get("regression_fail",0)),x["actual_workbook"],x["candidate_workbook"]])
    for rr in range(2, wf.max_row + 1):
        wf.cell(rr, 5).number_format = "0.00%"
        wf.cell(rr, 9).number_format = "0.00%"

    headers = ["PDF","課本頁","實體頁碼","字元","所在行","局部詞境","實際注音","預期／候選注音","來源類別","判定","判定來源","人工修正實際注音","判定備註","穩定注音鍵","font","注音元件ID","x0","y0","x1","y1","review_id"]
    def add_sheet(name, rows):
        sh = wb.create_sheet(name)
        sh.append(headers)
        for rr in rows:
            sh.append(rr)
        _style_sheet(sh)
        return sh

    err_rows = [core_record(r, "教材錯誤", "既有確認錯誤回歸", decision_source="規則／回歸自動判定") for r in auto_errors]
    err_rows += [core_record(r, "教材錯誤", d.get("note", ""), d.get("corrected_actual", ""), d.get("source", "人工確認"), d.get("expected", "")) for r, d in manual_errors]
    cor_rows = [core_record(r, "課本正確", "既有確認正確回歸", decision_source="規則／回歸自動判定") for r in auto_correct]
    cor_rows += [core_record(r, "課本正確", d.get("note", ""), d.get("corrected_actual", ""), d.get("source", "人工確認"), d.get("expected", "")) for r, d in manual_correct]
    sp_rows = []
    for r in specials:
        d = decisions.get(r["review_id"], {})
        if isinstance(d, dict) and d.get("decision") == "特殊略過":
            sp_rows.append(core_record(r, "特殊略過", d.get("note", ""), d.get("corrected_actual", ""), d.get("source", "人工確認"), d.get("expected", "")))
        else:
            sp_rows.append(core_record(r, "特殊略過", decision_source="規則自動判定"))
    de_rows = [core_record(r, "解碼錯誤", d.get("note", ""), d.get("corrected_actual", ""), d.get("source", "人工確認"), d.get("expected", "")) for r, d in decoder_errors]
    pe_rows = [core_record(r) for r in pending]
    add_sheet("確認教材錯誤", err_rows)
    add_sheet("課本正確", cor_rows)
    add_sheet("特殊略過", sp_rows)
    add_sheet("解碼錯誤", de_rows)
    add_sheet("待人工確認", pe_rows)
    cg_rows = [core_record(r, "", "現版未建立預期音，不得視為通過", decision_source="預期音覆蓋缺口") for r in coverage_gaps]
    add_sheet("未建立預期音", cg_rows)

    wr = wb.create_sheet("已補建預期音")
    wr.append(["PDF","課本頁","實體頁碼","字元","所在行","局部詞境","實際注音","補建預期注音","比較結果","預期音依據","補建來源","GPT信心","review_id"])
    for rec, d in resolved_expected:
        r = rec.get("row", {})
        wr.append([
            rec.get("pdf_name"), r.get("課本頁"), r.get("實體頁碼"), r.get("字元"), r.get("所在行"), r.get("局部詞境"),
            r.get("實際注音"), d.get("expected", ""), d.get("decision", ""), d.get("expected_evidence", ""), d.get("source", ""), d.get("gpt_confidence", ""), rec.get("review_id")
        ])
    _style_sheet(wr)

    add_sheet("所有候選", all_rows)
    wi = wb.create_sheet("執行資訊")
    wi.append(["項目", "內容"])
    wi.append(["程式版本", VERSION])
    wi.append(["資料夾", str(output_dir)])
    wi.append(["人工判定資料庫", str(output_dir / "人工判定資料庫.json")])
    wi.append(["安全原則", "actual 實際注音只由 PDF glyph／人工 occurrence visual 證據決定；expected 規則不得反填 actual；expected 也不得由 actual 選答案。"])
    wi.append(["雙證據鏈", "v5.5.0 對每個 occurrence 分別維護 actual_status 與 expected_status；只有兩邊都 RESOLVED 後才做 MATCH/MISMATCH 機械比較。"])
    wi.append(["專案 actual 證據", str(project_actual_evidence_root(output_dir))])
    wi.append(["glyph truth 衝突隔離", str(project_actual_evidence_root(output_dir) / GLYPH_CONFLICT_FILE)])
    wi.append(["expected補建", "v5.1.5 先以內建教育部《國語辭典簡編本》完整詞條／唯一主音自動補建 expected；只有仍無唯一證據者才交人工或GPT。補建後由程式機械比較 actual 與 expected。"] )
    wi.append(["現版獨立覆核", "前版相同或歷史判定不能直接產生本版『正確』；只有本版 actual 與本版 expected 可獨立重現且一致，才可通過。"])
    wi.append(["判定記憶安全門檻", "只允許在本版已獨立產生 actual-vs-expected 差異候選時，且字元、詞境、actual、expected、規則ID/模式/來源完全一致，才自動沿用既有人工判定；未解碼、多音未決與規則衝突不自動沿用。"])
    wi.append(["單機流程", "整批CFF自舉→實際注音解碼→手冊／專案規則→內建教育部簡編本完整詞條／唯一主音→專案字典與變調→一般殘餘候選／真正覆蓋缺口→人工或GPT→最終報告。"] )
    for sh in [ws, wf, wi]:
        _style_sheet(sh)

    out = output_dir / "注音校對_最終報告.xlsx"
    wb.save(out)
    return out


def _ledger_report_row(entry: Mapping[str, Any]) -> list[Any]:
    source = entry.get("source_record") or {}
    return [
        entry.get("occurrence_id"), entry.get("review_id"), entry.get("state"), entry.get("active_review"),
        infer_actual_status(entry), infer_expected_status(entry),
        entry.get("pdf_name"), entry.get("printed_page"), entry.get("physical_page"), entry.get("char"),
        source.get("所在行", ""), source.get("局部詞境", ""), entry.get("actual"),
        "|".join(entry.get("expected_set") or []), entry.get("comparison_result"),
        entry.get("actual_evidence"), entry.get("expected_evidence"), entry.get("context_evidence"),
        entry.get("source_view"), entry.get("identity_confidence"), entry.get("identity_row_fallback"),
        entry.get("canonical_occurrence_id"), entry.get("exclusion_reason"), entry.get("stable_key"),
        entry.get("font"), entry.get("font_xref"), entry.get("glyph_id"), entry.get("zhuyin_component_id"),
        entry.get("x0"), entry.get("y0"), entry.get("x1"), entry.get("y1"), entry.get("note"),
    ]



STATE_DISPLAY = {
    "PASS": "正確",
    "TEXTBOOK_ERROR_CONFIRMED": "已確認教材錯誤",
    "DIFFERENCE_PENDING_CONFIRMATION": "發現差異，請確認",
    "RULE_CONFLICT": "有兩種規則互相衝突",
    "EXPECTED_AMBIGUOUS": "正確讀音尚未確定",
    "EXPECTED_UNRESOLVED": "尚未找到正確讀音依據",
    "REVIEW_PENDING": "需要人工確認",
    "ACTUAL_DECODE_ERROR": "程式可能讀錯課本注音",
    "ACTUAL_UNRESOLVED": "課本注音尚未安全辨識",
    "REGRESSION_BLOCKED": "內部回歸檢查未通過",
    "SOURCE_INVALID": "校對資料來源檢查失敗",
    "DATA_INTEGRITY_ERROR": "工作資料完整性異常",
    "EXCLUDED_OUT_OF_SCOPE": "此處不需校對",
    "EXCLUDED_NONINDEPENDENT_LAYER": "重複技術層，不重複校對",
}


def friendly_state(state: Any) -> str:
    return STATE_DISPLAY.get(str(state or ""), str(state or ""))


def _entry_phrase(entry: Mapping[str, Any]) -> str:
    source = entry.get("source_record") or {}
    return str(source.get("局部詞境") or entry.get("context_evidence") or entry.get("char") or "").strip()


def _entry_sentence(entry: Mapping[str, Any]) -> str:
    source = entry.get("source_record") or {}
    return str(source.get("所在行") or source.get("局部詞境") or "").strip()


def _friendly_issue(entry: Mapping[str, Any]) -> tuple[str, str]:
    state = str(entry.get("state") or "")
    if state == "DIFFERENCE_PENDING_CONFIRMATION":
        return "課本目前注音與應標注音不同。", "核對原頁與規範；確定有誤再確認教材錯誤，若課本其實正確則補充較高優先的應標讀音依據。"
    if state == "RULE_CONFLICT":
        return "程式找到互相衝突的讀音規則。", "選擇本句正確讀音並提供規範／辭典／公司規定等獨立依據。"
    if state == "EXPECTED_AMBIGUOUS":
        return "目前有多個可能讀音，證據不足以自動決定。", "有正式依據時補充正確讀音；沒有就保留待確認。"
    if state == "EXPECTED_UNRESOLVED":
        return "尚未建立可獨立重現的應標注音。", "查完整詞條或公司規定後補充應標讀音；不能照著課本目前注音倒推。"
    if state == "ACTUAL_DECODE_ERROR":
        return "程式可能把課本現標注音辨識錯誤。", "回頁確認後以指定位置辨識修正處理，再重新解碼。"
    if state == "ACTUAL_UNRESOLVED":
        return "課本現標注音尚未安全辨識。", "先處理課本注音辨識，不要先用應標讀音猜現標。"
    if state == "REVIEW_PENDING":
        return "目前證據不足，需要人工判斷。", "核對原頁、完整詞語與正式來源後再處理。"
    if state in {"REGRESSION_BLOCKED", "SOURCE_INVALID", "DATA_INTEGRITY_ERROR"}:
        return friendly_state(state), "這是程式／資料完整性問題，先修復後再繼續校對。"
    return friendly_state(state), "保留待處理。"

def generate_report(
    output_dir: Path,
    manifest: dict[str, Any],
    db: dict[str, Any],
    *,
    runtime_root: Path | None = None,
) -> Path:
    """Render mutually-exclusive ledger views and the v2.5 completion gate."""
    validate_manifest_integrity(manifest)
    validate_output_artifact_hashes(manifest)
    if (
        not schema_compatible(manifest.get("session_schema_version"), SESSION_SCHEMA_VERSION)
        or not schema_compatible(manifest.get("workbook_schema_version"), WORKBOOK_SCHEMA_VERSION)
        or not schema_compatible(manifest.get("ledger_schema_version"), LEDGER_SCHEMA_VERSION)
        or not review_id_schema_compatible(manifest.get("review_id_schema_version"), REVIEW_ID_SCHEMA_VERSION)
    ):
        raise ValueError("SESSION_SCHEMA_INCOMPATIBLE：session/workbook/review ID 資料布局無法直接沿用")

    root = Path(runtime_root or Path(__file__).resolve().parent).resolve()
    source_validation = validate_asset_manifest(root)
    if not source_validation.get("ok"):
        write_pipeline_blocked(output_dir, source_validation, "SOURCE_INVALID_DURING_REPORT")
        raise SourceValidationError("runtime 核心資料來源已改變或損壞；保留上一份可用最終報告")
    current_expected_fingerprint = compute_expected_asset_fingerprint(
        root,
        source_validation,
        resolver_version=EXPECTED_RESOLVER_VERSION,
        source_files=EXPECTED_RESOLVER_SOURCE_FILES,
    )
    if not fingerprint_compatible(
        manifest.get("expected_asset_fingerprint"),
        manifest.get("expected_asset_fingerprint_components"),
        current_expected_fingerprint,
        chain="expected",
    ):
        blocked_report = dict(source_validation)
        blocked_report["errors"] = list(source_validation.get("errors") or []) + ["candidate expected evidence assets 已失效"]
        write_pipeline_blocked(output_dir, blocked_report, "EXPECTED_CANDIDATE_CACHE_INVALID")
        raise SourceValidationError("expected 資料資產已改變；只需重建受影響候選，未覆寫上一份最終報告")
    compat_warnings = metadata_compatibility_warnings(
        manifest.get("expected_asset_fingerprint_components"), current_expected_fingerprint, chain="expected"
    )

    ledger = materialize_ledger(manifest, db)
    actual_source_ids = manifest.get("actual_source_ids") or []
    reconciliation = reconcile_ledger(ledger, actual_source_ids)
    gate = completion_gate(
        ledger,
        reconciliation,
        manifest.get("regression_gate") or {},
        source_validation_ok=bool(source_validation.get("ok")),
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "摘要"
    ws.append(["指標", "結果"])
    status_text = {
        PROCESSING_FINISHED: "處理結束；校對尚未完成。",
        PROOFREAD_COMPLETE: "全冊注音校對完成。",
        PIPELINE_BLOCKED: "處理被阻擋。",
    }[gate["status"]]
    state_counts = Counter(str(entry.get("state") or "") for entry in ledger)
    summary = [
        ("工具", f"{PROGRAM} v{VERSION}"),
        ("pipeline status", gate["status"]),
        ("狀態文字", status_text),
        ("session_id", manifest.get("session_id", "")),
        ("session_schema_version", SESSION_SCHEMA_VERSION),
        ("workbook_schema_version", WORKBOOK_SCHEMA_VERSION),
        ("ledger_schema_version", LEDGER_SCHEMA_VERSION),
        ("review_id_schema_version", REVIEW_ID_SCHEMA_VERSION),
        ("expected_asset_fingerprint", current_expected_fingerprint["fingerprint"]),
        ("PDF檔數", len(manifest.get("pdfs", []))),
        ("PDF偵測母體總數", len(ledger)),
        ("現版注音校對分母", gate["in_scope_total"]),
        ("actual覆蓋數", gate["actual_covered"]),
        ("actual解碼覆蓋率", gate["actual_coverage"]),
        ("expected覆蓋數", gate["expected_covered"]),
        ("expected覆蓋率", gate["expected_coverage"]),
        ("待人工／非終態", sum(state_counts[state] for state in NON_TERMINAL_STATES)),
        ("未建立預期音", state_counts["EXPECTED_UNRESOLVED"]),
        ("多音待判", state_counts["EXPECTED_AMBIGUOUS"]),
        ("規則衝突", state_counts["RULE_CONFLICT"]),
        ("解碼錯誤", state_counts["ACTUAL_DECODE_ERROR"] + state_counts["ACTUAL_UNRESOLVED"]),
        ("確認通過", state_counts["PASS"]),
        ("現版六閘門確認教材錯誤", state_counts["TEXTBOOK_ERROR_CONFIRMED"]),
        ("非獨立技術層排除", state_counts["EXCLUDED_NONINDEPENDENT_LAYER"]),
        ("真正不在範圍排除", state_counts["EXCLUDED_OUT_OF_SCOPE"]),
        ("mandatory regression required", (manifest.get("mandatory_regression") or {}).get("required", 0)),
        ("mandatory regression executed", (manifest.get("mandatory_regression") or {}).get("executed", 0)),
        ("mandatory regression passed", (manifest.get("mandatory_regression") or {}).get("passed", 0)),
        ("mandatory regression failed", (manifest.get("mandatory_regression") or {}).get("failed", 0)),
        ("mandatory regression not executed", (manifest.get("mandatory_regression") or {}).get("not_executed", 0)),
        ("PDF occurrence regression required", (manifest.get("pdf_regression") or {}).get("required", 0)),
        ("PDF occurrence regression executed", (manifest.get("pdf_regression") or {}).get("executed", 0)),
        ("PDF occurrence regression passed", (manifest.get("pdf_regression") or {}).get("passed", 0)),
        ("PDF occurrence regression failed", (manifest.get("pdf_regression") or {}).get("failed", 0)),
        ("PDF occurrence regression not executed", (manifest.get("pdf_regression") or {}).get("not_executed", 0)),
        ("PDF occurrence regression not applicable", (manifest.get("pdf_regression") or {}).get("not_applicable", 0)),
        ("reference-only audit total", (manifest.get("pdf_regression") or {}).get("reference_audit_total", 0)),
        ("reference-only audit matched", (manifest.get("pdf_regression") or {}).get("reference_audit_matched", 0)),
        ("reference-only audit passed", (manifest.get("pdf_regression") or {}).get("reference_audit_passed", 0)),
        ("reference-only audit failed (non-gating)", (manifest.get("pdf_regression") or {}).get("reference_audit_failed", 0)),
        ("reference-only audit not found", (manifest.get("pdf_regression") or {}).get("reference_audit_not_found", 0)),
        ("集合對帳", "PASS" if reconciliation.ok else "FAIL"),
        ("失敗硬門檻", "｜".join(gate.get("failed_gates") or [])),
        ("完成判定", "只有 actual 100%、expected 100%、所有非終態/錯誤/衝突為 0、mandatory 與適用 PDF regression 全部實際執行且 0 失敗、全量集合對帳成立，才可 PROOFREAD_COMPLETE。"),
        ("建立時間", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for item in summary:
        ws.append(item)
    for row_no in range(2, ws.max_row + 1):
        if ws.cell(row_no, 1).value in {"actual解碼覆蓋率", "expected覆蓋率"}:
            ws.cell(row_no, 2).number_format = "0.00%"

    headers = [
        "occurrence_id", "review_id", "state", "active_review", "actual_status", "expected_status", "PDF", "課本頁", "實體頁碼", "字元",
        "所在行", "局部詞境", "actual", "expected_set", "機械比較", "actual證據", "expected證據", "語境／位置證據",
        "來源view", "identity confidence", "row fallback", "canonical occurrence", "排除理由", "穩定注音鍵",
        "font", "font_xref", "glyph_id", "注音元件ID", "x0", "y0", "x1", "y1", "備註",
    ]
    ws_ledger = wb.create_sheet("Occurrence Ledger")
    ws_ledger.append(headers)
    for entry in ledger:
        ws_ledger.append(_ledger_report_row(entry))

    # Exactly one authoritative state sheet per occurrence. Empty sheets remain
    # present so deletion/rename can be detected by downstream validation.
    for state in sorted(ALL_STATES):
        state_sheet = wb.create_sheet(state[:31])
        state_sheet.append(headers)
        for entry in ledger:
            if entry.get("state") == state:
                state_sheet.append(_ledger_report_row(entry))

    wr = wb.create_sheet("已補建預期音_稽核事件")
    wr.append(["review_id", "occurrence_id", "action", "expected_set", "expected_evidence", "event_source", "updated_at", "note"])
    by_review_id = {entry["review_id"]: entry for entry in ledger}
    for review_id, event in normalize_db(db).get("events", {}).items():
        if event.get("action") != "補建expected證據":
            continue
        entry = by_review_id.get(review_id, {})
        wr.append([
            review_id, entry.get("occurrence_id", ""), event.get("action", ""),
            "|".join(normalize_expected_set(event.get("expected_set"))), event.get("expected_evidence", ""),
            event.get("source", ""), event.get("updated_at", ""), event.get("note", ""),
        ])

    wc = wb.create_sheet("集合對帳")
    wc.append(["項目", "數量／結果", "明細"])
    reconciliation_data = reconciliation.as_dict()
    for key in [
        "ok", "actual_count", "ledger_count", "in_scope_count", "explicitly_excluded_count",
        "actual_only", "ledger_only", "duplicate_occurrence_ids", "duplicate_review_ids",
        "multiple_authoritative_states", "unclassified", "worksheet_count_mismatch", "summary_mismatch", "state_counts", "errors",
    ]:
        value = reconciliation_data.get(key)
        detail = json.dumps(value, ensure_ascii=False, sort_keys=True) if isinstance(value, (dict, list)) else ""
        wc.append([key, value if not isinstance(value, (dict, list)) else len(value), detail])

    wf = wb.create_sheet("檔案統計")
    wf.append(["PDF", "SHA-256", "actual fingerprint", "偵測母體", "校對分母", "actual覆蓋", "expected覆蓋", "mandatory failed", "PDF regression failed", "PDF regression not executed"])
    for info in manifest.get("pdfs", []):
        pdf_entries = [entry for entry in ledger if entry.get("pdf_name") == info.get("pdf_name")]
        in_scope = [entry for entry in pdf_entries if entry.get("state") not in EXCLUDED_STATES]
        actual_covered = sum(1 for entry in in_scope if canonical_bopomofo(entry.get("actual")) and entry.get("actual_evidence"))
        expected_covered = sum(1 for entry in in_scope if normalize_expected_set(entry.get("expected_set")) and entry.get("expected_evidence"))
        wf.append([
            info.get("pdf_name"), info.get("pdf_sha256"), info.get("actual_asset_fingerprint"), len(pdf_entries), len(in_scope),
            actual_covered / len(in_scope) if in_scope else 0, expected_covered / len(in_scope) if in_scope else 0,
            (info.get("mandatory_regression") or {}).get("failed", 0),
            (info.get("pdf_regression") or {}).get("failed", 0), (info.get("pdf_regression") or {}).get("not_executed", 0),
        ])
    for row_no in range(2, wf.max_row + 1):
        wf.cell(row_no, 6).number_format = "0.00%"
        wf.cell(row_no, 7).number_format = "0.00%"

    wa = wb.create_sheet("資料來源稽核")
    wa.append(["name", "chain", "path", "SHA-256", "schema ok", "loaded rows", "validated rows", "errors"])
    for asset in source_validation.get("assets", []):
        schema = asset.get("schema") or {}
        wa.append([
            asset.get("name"), asset.get("chain"), asset.get("path"), asset.get("sha256"),
            schema.get("ok"), schema.get("loaded_row_count"), schema.get("validated_row_count"),
            "｜".join(asset.get("errors") or []),
        ])

    wm = wb.create_sheet("強制回歸摘要")
    wm.append(["指標", "結果"])
    for key in ("ok", "required", "executed", "passed", "failed", "not_executed", "duplicate_case_id"):
        wm.append([key, (manifest.get("mandatory_regression") or {}).get(key)])

    wh = wb.create_sheet("歷史回歸稽核")
    historical_columns = [
        "案例ID", "適用性模式", "適用性狀態", "計入完成門檻", "回歸結果", "定位命中數",
        "執行PDF", "課本頁", "完整詞語", "字元", "目標在詞內序號", "定位上下文",
        "來源PDF SHA-256", "真值實際注音", "真值預期注音", "真值結論", "執行說明", "來源", "備註",
    ]
    wh.append(historical_columns)
    for row in (manifest.get("pdf_regression") or {}).get("results", []):
        wh.append([row.get(column, "") for column in historical_columns])

    wi = wb.create_sheet("執行資訊")
    wi.append(["項目", "內容"])
    wi.append(["pipeline status", gate["status"]])
    wi.append(["處理／完成分離", "PROCESSING_FINISHED 只表示程式跑完；只有所有 2.5 硬門檻成立才是 PROOFREAD_COMPLETE。"])
    wi.append(["actual/expected independence", "actual 僅來自 PDF glyph 解碼與 occurrence override；expected 證據不得參照 actual 反推。"])
    wi.append(["legacy isolation", "v5.1.5 session/workbook/decision 不相容；無法唯一映射者只進 legacy_orphan_decisions。"])

    # v5.3.0: all governance detail is preserved in a separate technical audit
    # workbook. The day-to-day report intentionally contains only information
    # needed by textbook proofreaders.
    for sheet in wb.worksheets:
        _style_sheet(sheet)
    technical_out = output_dir / "注音校對_技術稽核.xlsx"
    wb.save(technical_out)

    user_wb = Workbook()
    us = user_wb.active
    us.title = "校對摘要"
    us.append(["項目", "結果"])
    pending_count = sum(state_counts[state] for state in NON_TERMINAL_STATES)
    decode_problem_count = state_counts["ACTUAL_DECODE_ERROR"] + state_counts["ACTUAL_UNRESOLVED"]
    excluded_count = state_counts["EXCLUDED_NONINDEPENDENT_LAYER"] + state_counts["EXCLUDED_OUT_OF_SCOPE"]
    internal_ok = bool(source_validation.get("ok")) and bool(reconciliation.ok) and bool((manifest.get("regression_gate") or {}).get("ok"))
    if gate["status"] == PROOFREAD_COMPLETE:
        human_status = "全冊注音校對完成"
    elif gate["status"] == PIPELINE_BLOCKED:
        human_status = "處理被阻擋"
    elif pending_count:
        human_status = f"尚有 {pending_count} 筆待處理"
    else:
        human_status = "校對項目已收斂，但尚未達成全冊完成條件"
    user_summary = [
        ("工具版本", VERSION),
        ("目前狀態", human_status),
        ("PDF 檔數", len(manifest.get("pdfs", []))),
        ("可見注音座標", len(ledger)),
        ("已確認正確", state_counts["PASS"]),
        ("已確認教材錯誤", state_counts["TEXTBOOK_ERROR_CONFIRMED"]),
        ("待確認", pending_count),
        ("注音辨識待處理", decode_problem_count),
        ("不列入校對分母", excluded_count),
        ("課本注音辨識覆蓋率", gate["actual_coverage"]),
        ("應標讀音覆蓋率", gate["expected_coverage"]),
        ("內部資料／回歸／集合檢查", "通過" if internal_ok else "未通過（詳見技術稽核報告）"),
        ("完成狀態", "已完成" if gate["status"] == PROOFREAD_COMPLETE else "尚未完成"),
        ("建立時間", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for item in user_summary:
        us.append(item)
    for row_no in range(2, us.max_row + 1):
        if us.cell(row_no, 1).value in {"課本注音辨識覆蓋率", "應標讀音覆蓋率"}:
            us.cell(row_no, 2).number_format = "0.00%"

    ue = user_wb.create_sheet("修正清單")
    ue.append(["PDF", "課本頁", "所在句", "詞語／局部詞境", "字", "現標", "應標", "依據", "確認紀錄"])
    for entry in ledger:
        if entry.get("state") != "TEXTBOOK_ERROR_CONFIRMED":
            continue
        ue.append([
            entry.get("pdf_name", ""), entry.get("printed_page", ""), _entry_sentence(entry), _entry_phrase(entry),
            entry.get("char", ""), entry.get("actual", ""), " | ".join(entry.get("expected_set") or []),
            entry.get("expected_evidence", ""), entry.get("note", ""),
        ])

    up = user_wb.create_sheet("待確認")
    up.append(["PDF", "課本頁", "所在句", "詞語／局部詞境", "字", "目前注音", "應標／候選", "問題", "建議操作"])
    for entry in ledger:
        if entry.get("state") not in NON_TERMINAL_STATES:
            continue
        issue, suggestion = _friendly_issue(entry)
        up.append([
            entry.get("pdf_name", ""), entry.get("printed_page", ""), _entry_sentence(entry), _entry_phrase(entry),
            entry.get("char", ""), entry.get("actual", ""), " | ".join(entry.get("expected_set") or []), issue, suggestion,
        ])

    ur = user_wb.create_sheet("人工與公司規則")
    ur.append(["來源類型", "PDF", "課本頁", "詞語／局部詞境", "字", "應標注音", "依據／來源", "適用範圍", "備註"])

    # Reusable rules are intentionally shown in the human-facing report because
    # they affect future expected evidence. Machine IDs/hashes remain confined to
    # the technical audit and session JSON.
    rules_doc = _canonical_rule_doc(manifest.get("reusable_expected_rules") or {})
    for rule in rules_doc.get("rules", []):
        if rule.get("enabled", True) is False:
            continue
        target_index = int(rule.get("target_index") or 0)
        scope = "相同完整詞＋相同目標位置"
        if rule.get("context_contains"):
            scope += f"；句境需包含：{rule.get('context_contains')}"
        ur.append([
            "可重用規則", "", "", rule.get("phrase", ""), rule.get("target_char", ""),
            " | ".join(rule.get("expected_set") or []), rule.get("evidence", ""),
            f"{scope}（第 {target_index + 1} 字）", rule.get("note", ""),
        ])

    current_by_review = {entry.get("review_id"): entry for entry in ledger}
    for review_id, event in normalize_db(db).get("events", {}).items():
        entry = current_by_review.get(review_id) or {}
        action = str(event.get("action") or "")
        if action == "保留待人工":
            continue
        if action in {"補建expected證據", "解決expected證據"}:
            label = "人工建立／解決應標讀音"
            expected_text = " | ".join(normalize_expected_set(event.get("expected_set")))
            evidence_text = str(event.get("expected_evidence") or event.get("source") or "")
        elif action == "確認現版差異":
            if entry.get("review_event_replay_status") == "INVALIDATED_EXPECTED_DRIFT":
                label = "舊確認已失效（expected 已變更）"
            else:
                label = "確認教材錯誤"
            expected_text = " | ".join(entry.get("expected_set") or [])
            evidence_text = str(entry.get("expected_evidence") or event.get("source") or "")
        elif action == "確認非校對範圍":
            label = "確認不需校對"
            expected_text = ""
            evidence_text = str(event.get("exclusion_evidence") or event.get("source") or "")
        else:
            label = action
            expected_text = ""
            evidence_text = str(event.get("source") or "")
        ur.append([
            "人工判定", entry.get("pdf_name", ""), entry.get("printed_page", ""), _entry_phrase(entry), entry.get("char", ""),
            expected_text, evidence_text, "僅此位置", str(
                entry.get("review_event_replay_note")
                or event.get("note")
                or event.get("resolution_reason")
                or ""
            ),
        ])

    for sheet in user_wb.worksheets:
        _style_sheet(sheet)
    out = output_dir / "注音校對_最終報告.xlsx"
    user_wb.save(out)

    json_save(output_dir / "pipeline_status.json", {
        "version": VERSION,
        "status": gate["status"],
        "status_text": status_text,
        "completion_gate": gate,
        "reconciliation": reconciliation.as_dict(),
        "user_report": str(out),
        "technical_audit_report": str(technical_out),
    })
    return out



def _session_metadata_for_actual(manifest: Mapping[str, Any]) -> dict[str, Any]:
    return {
        "version": VERSION,
        "session_id": manifest.get("session_id"),
        "session_schema_version": SESSION_SCHEMA_VERSION,
        "workbook_schema_version": WORKBOOK_SCHEMA_VERSION,
        "review_id_schema_version": REVIEW_ID_SCHEMA_VERSION,
    }


def _resolve_session_pdfs(output_dir: Path, manifest: Mapping[str, Any]) -> list[Path]:
    """Resolve current PDF bytes even when the project/source folder was moved.

    Stored paths are preferred only when their SHA-256 still matches. Otherwise
    search under the project parent by exact filename and recorded SHA.  Path is
    not part of actual truth; file bytes are.
    """
    output_dir = Path(output_dir).resolve()
    search_roots = [output_dir.parent]
    resolved: list[Path] = []
    for info in manifest.get("pdfs", []):
        name = str(info.get("pdf_name") or Path(str(info.get("pdf") or "")).name).strip()
        recorded_sha = str(info.get("pdf_sha256") or "").strip()
        stored = Path(str(info.get("pdf") or ""))
        candidates: list[Path] = []
        if stored.exists() and stored.is_file():
            candidates.append(stored)
        for base in search_roots:
            direct = base / name
            if name and direct.exists() and direct.is_file():
                candidates.append(direct)
            if name:
                try:
                    candidates.extend(p for p in base.rglob(name) if p.is_file())
                except Exception:
                    pass
        unique: list[Path] = []
        seen = set()
        for candidate in candidates:
            try:
                rp = candidate.resolve()
            except Exception:
                rp = candidate
            key = str(rp)
            if key in seen:
                continue
            seen.add(key)
            if recorded_sha:
                try:
                    if sha256_file(rp) != recorded_sha:
                        continue
                except Exception:
                    continue
            unique.append(rp)
        if not unique:
            raise FileNotFoundError(
                f"找不到現版 PDF：{name}。可移動檔案，但內容必須與原工作階段相同；"
                f"請把 PDF 放在校對專案資料夾的上一層或其子資料夾後再試。"
            )
        resolved.append(unique[0])
    if not resolved:
        raise FileNotFoundError("工作階段沒有可解析的 PDF")
    return resolved


def export_actual_pending_for_gpt(output_dir: Path) -> Path | None:
    output_dir = Path(output_dir)
    manifest = json_load_strict(output_dir / "校對工作階段.json")
    validate_manifest_integrity(manifest)
    validate_output_artifact_hashes(manifest)
    if (
        not schema_compatible(manifest.get("session_schema_version"), SESSION_SCHEMA_VERSION)
        or not review_id_schema_compatible(manifest.get("review_id_schema_version"), REVIEW_ID_SCHEMA_VERSION)
    ):
        raise ValueError("SESSION_SCHEMA_INCOMPATIBLE：此工作階段無法建立 actual 待判定包")
    db = load_or_initialize_db(output_dir)
    ledger = materialize_ledger(manifest, db)
    groups = build_actual_review_groups(ledger)
    if not groups:
        return None
    return export_actual_review_package(
        output_dir,
        ledger,
        version=VERSION,
        session_id=str(manifest.get("session_id") or ""),
        session_schema_version=SESSION_SCHEMA_VERSION,
        workbook_schema_version=WORKBOOK_SCHEMA_VERSION,
        review_id_schema_version=REVIEW_ID_SCHEMA_VERSION,
    )


def manual_actual_staging_summary(output_dir: Path) -> dict[str, Any]:
    """Return a read-only summary of durable, not-yet-authoritative actual intent."""
    staging = load_manual_actual_staging(project_actual_evidence_root(Path(output_dir)))
    groups = list(staging.get("staged_groups") or [])
    return {
        "staged_group_count": len(groups),
        "staged_group_ids": [str(group.get("group_id") or "") for group in groups],
        "staged_member_occurrence_ids": sorted({
            str(occurrence_id)
            for group in groups
            for occurrence_id in (group.get("member_occurrence_ids") or [])
            if str(occurrence_id)
        }),
        "staged_checked_occurrence_ids": sorted({
            str(occurrence_id)
            for group in groups
            for occurrence_id in (group.get("checked_occurrence_ids") or [])
            if str(occurrence_id)
        }),
    }


def stage_manual_actual_correction(
    output_dir: Path,
    review_id: str,
    reading: str,
    checked_occurrence_ids: list[str] | None = None,
    note: str = "",
) -> dict[str, Any]:
    """Rebuild one current live group and durably stage its visual actual decision."""
    output_dir = Path(output_dir)
    manifest = json_load_strict(output_dir / "校對工作階段.json")
    validate_manifest_integrity(manifest)
    validate_output_artifact_hashes(manifest)
    db = load_or_initialize_db(output_dir)
    ledger = materialize_ledger(manifest, db)
    entry = next(
        (row for row in ledger if str(row.get("review_id") or "") == str(review_id)),
        None,
    )
    if entry is None:
        raise ValueError("找不到目前 review_id；請重新開啟人工校對畫面")
    group = build_actual_group_for_entry(ledger, entry)
    checked_ids = (
        [str(entry.get("occurrence_id") or "")]
        if checked_occurrence_ids is None
        else list(checked_occurrence_ids)
    )
    staged = stage_manual_actual_group(
        project_actual_evidence_root(output_dir),
        group,
        reading,
        checked_occurrence_ids=checked_ids,
        source="人工 GUI actual 視覺確認",
        note=note,
    )
    return {
        "staged_group": staged,
        "staging_summary": manual_actual_staging_summary(output_dir),
    }


def _clear_actual_dependent_events(output_dir: Path, ledger: list[dict[str, Any]], occurrence_ids: set[str]) -> int:
    if not occurrence_ids:
        return 0
    db = load_or_initialize_db(output_dir)
    review_ids = {
        str(entry.get("review_id") or "")
        for entry in ledger
        if str(entry.get("occurrence_id") or "") in occurrence_ids
    }
    events = db.setdefault("events", {})
    removed = 0
    for rid in list(review_ids):
        if not rid or rid not in events:
            continue
        event = events.get(rid) or {}
        action = str(event.get("action") or "")
        # Expected evidence and scope decisions are independent from actual and
        # must survive an actual-only correction.  A final mismatch confirmation
        # depends on the old actual, so revoke only the confirmation layer while
        # preserving its embedded expected/context evidence when possible.
        if action == "確認現版差異":
            expected_set = list(normalize_expected_set(event.get("expected_set")))
            expected_evidence = str(event.get("expected_evidence") or "").strip()
            context_evidence = str(event.get("context_evidence") or "").strip()
            if expected_set and expected_evidence and context_evidence:
                old_note = str(event.get("note") or "").strip()
                events[rid] = {
                    "action": "解決expected證據",
                    "expected_set": expected_set,
                    "expected_evidence": expected_evidence,
                    "context_evidence": context_evidence,
                    "resolution_reason": "actual 已更新；撤銷舊差異確認但保留獨立 expected 證據",
                    "source": str(event.get("expected_resolution_source") or event.get("source") or "既有 expected 證據"),
                    "note": (old_note + "；actual 更新後已撤銷舊 human_confirmation").strip("；"),
                    "updated_at": datetime.now().isoformat(timespec="seconds"),
                    "_safe_expected_rebase": True,
                }
            else:
                events.pop(rid, None)
            removed += 1
    json_save(output_dir / "人工判定資料庫.json", db)
    return removed


def refresh_actual_project(output_dir: Path, *, defer_excel_reports: bool = True) -> Path:
    """Incrementally re-decode current project after dynamic actual evidence changes."""
    output_dir = Path(output_dir)
    manifest = json_load_strict(output_dir / "校對工作階段.json")
    validate_manifest_integrity(manifest)
    # Workbooks may now be stale relative to dynamic actual evidence, but must
    # still be intact relative to the last sealed session before we replace them.
    validate_output_artifact_hashes(manifest)
    pdfs = _resolve_session_pdfs(output_dir, manifest)
    return run_pipeline_pdfs(
        pdfs,
        output_dir,
        session_id_override=str(manifest.get("session_id") or ""),
        defer_excel_reports=defer_excel_reports,
    )


def import_actual_gpt_decisions(output_dir: Path, xlsx: Path) -> tuple[int, int, Path]:
    output_dir = Path(output_dir)
    manifest = json_load_strict(output_dir / "校對工作階段.json")
    validate_manifest_integrity(manifest)
    validate_output_artifact_hashes(manifest)
    db = load_or_initialize_db(output_dir)
    ledger = materialize_ledger(manifest, db)
    groups = build_actual_review_groups(ledger)
    result = import_actual_review_workbook(
        initialize_project_actual_evidence(output_dir, Path(__file__).resolve().parent),
        Path(xlsx),
        groups,
        expected_metadata=_session_metadata_for_actual(manifest),
    )
    occurrence_ids = {
        oid
        for item in result.get("results", [])
        for oid in (item.get("affected_occurrence_ids") or item.get("target_occurrence_ids") or [])
        if oid
    }
    removed = _clear_actual_dependent_events(output_dir, ledger, occurrence_ids)
    report = refresh_actual_project(output_dir)
    return int(result.get("imported_groups") or 0), removed, report


def apply_manual_actual_correction(
    output_dir: Path,
    review_id: str,
    reading: str,
    checked_occurrence_ids: list[str] | None = None,
    note: str = "",
) -> tuple[dict[str, Any], Path]:
    """Apply one human visual actual correction and immediately refresh project."""
    output_dir = Path(output_dir)
    manifest = json_load_strict(output_dir / "校對工作階段.json")
    validate_manifest_integrity(manifest)
    validate_output_artifact_hashes(manifest)
    db = load_or_initialize_db(output_dir)
    ledger = materialize_ledger(manifest, db)
    entry = next((row for row in ledger if str(row.get("review_id") or "") == str(review_id)), None)
    if entry is None:
        raise ValueError("找不到目前 review_id；請重新開啟人工校對畫面")
    group = build_actual_group_for_entry(ledger, entry)
    result = apply_verified_actual_group(
        initialize_project_actual_evidence(output_dir, Path(__file__).resolve().parent),
        group,
        reading,
        checked_occurrence_ids=checked_occurrence_ids,
        source="人工 GUI actual 視覺確認",
        note=note,
    )
    affected = set(result.get("affected_occurrence_ids") or result.get("target_occurrence_ids") or [])
    _clear_actual_dependent_events(output_dir, ledger, affected)
    report = refresh_actual_project(output_dir)

    # Hard postcondition: a visual actual correction is not considered applied
    # merely because the refresh command completed.  The rebuilt ledger must
    # show the submitted actual at the same occurrence; otherwise surface a
    # deterministic error instead of silently leaving the same review card.
    refreshed_manifest = json_load_strict(output_dir / "校對工作階段.json")
    refreshed_db = load_or_initialize_db(output_dir)
    refreshed_ledger = materialize_ledger(refreshed_manifest, refreshed_db)
    target_occurrence_id = str(entry.get("occurrence_id") or "")
    refreshed_entry = next(
        (row for row in refreshed_ledger if str(row.get("occurrence_id") or "") == target_occurrence_id),
        None,
    )
    if refreshed_entry is None:
        # Defensive fallback for a future identity-schema migration: coordinates
        # plus PDF/page/character are still occurrence-scoped and auditable.
        def same_location(row):
            if str(row.get("pdf_name") or "") != str(entry.get("pdf_name") or ""):
                return False
            if str(row.get("char") or "") != str(entry.get("char") or ""):
                return False
            if int(row.get("physical_page") or 0) != int(entry.get("physical_page") or 0):
                return False
            try:
                return abs(float(row.get("x0") or 0) - float(entry.get("x0") or 0)) <= 0.8 and abs(float(row.get("y0") or 0) - float(entry.get("y0") or 0)) <= 0.8
            except Exception:
                return False
        refreshed_entry = next((row for row in refreshed_ledger if same_location(row)), None)
    wanted = canonical_bopomofo(reading)
    observed = canonical_bopomofo((refreshed_entry or {}).get("actual"))
    if refreshed_entry is None or not wanted or observed != wanted:
        raise ValueError(
            "actual 視覺修正已寫入，但重新解碼後未套用到目標位置；"
            f"要求={wanted or reading}，重解結果={observed or '未辨識'}。"
            "程式已停止把此筆視為完成，請使用 v5.5.x 的 occurrence 匹配修正後再試。"
        )
    result["post_actual"] = observed
    result["post_state"] = str(refreshed_entry.get("state") or "")
    result["resolved_from_pending"] = result["post_state"] not in NON_TERMINAL_STATES
    return result, report


def _current_live_groups_for_staged_manual_actual(
    ledger: list[dict[str, Any]],
    staged_decisions: list[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    """Rebuild every Phase 2A live group from the current materialized ledger."""
    by_occurrence_id: dict[str, dict[str, Any]] = {}
    for entry in ledger:
        occurrence_id = str(entry.get("occurrence_id") or "")
        if not occurrence_id:
            raise ValueError("manual actual batch current ledger 含空白 occurrence_id")
        if occurrence_id in by_occurrence_id:
            raise DuplicateIdError(f"manual actual batch current ledger occurrence_id 重複：{occurrence_id}")
        by_occurrence_id[occurrence_id] = entry

    live_groups: list[dict[str, Any]] = []
    for staged in sorted(staged_decisions, key=lambda item: str(item.get("group_id") or "")):
        current_member = next(
            (
                by_occurrence_id.get(str(occurrence_id or ""))
                for occurrence_id in (staged.get("member_occurrence_ids") or [])
                if by_occurrence_id.get(str(occurrence_id or "")) is not None
            ),
            None,
        )
        if current_member is None:
            raise ValueError(
                "manual actual batch staged group 在目前 ledger 找不到任何 member："
                f"{staged.get('group_id') or ''}；整批拒絕"
            )
        live_groups.append(build_actual_group_for_entry(ledger, current_member))
    return live_groups


def _manual_actual_batch_postconditions(
    batch_result: Mapping[str, Any],
) -> list[dict[str, str]]:
    """Freeze the directly checked actuals committed by the Phase 2A batch."""
    group_results = batch_result.get("group_results")
    if not isinstance(group_results, list):
        raise ValueError("manual actual batch result 缺少 group_results")
    postconditions: list[dict[str, str]] = []
    seen_occurrence_ids: set[str] = set()
    for index, item in enumerate(group_results):
        if not isinstance(item, Mapping):
            raise ValueError(f"manual actual batch group_results[{index}] 格式錯誤")
        group_id = str(item.get("group_id") or "").strip()
        reading = canonical_bopomofo(item.get("reading"))
        checked_ids = item.get("verified_occurrence_ids")
        if not group_id or not reading or not isinstance(checked_ids, list) or not checked_ids:
            raise ValueError(f"manual actual batch result 缺少直接核對 postcondition：index={index}")
        for occurrence_id_raw in checked_ids:
            occurrence_id = str(occurrence_id_raw or "").strip()
            if not occurrence_id:
                raise ValueError(f"manual actual batch result 含空白 checked occurrence：{group_id}")
            if occurrence_id in seen_occurrence_ids:
                raise ValueError(f"manual actual batch result checked occurrence 重複：{occurrence_id}")
            seen_occurrence_ids.add(occurrence_id)
            postconditions.append({
                "group_id": group_id,
                "occurrence_id": occurrence_id,
                "reading": reading,
            })
    if len(group_results) != int(batch_result.get("applied_group_count") or 0):
        raise ValueError("manual actual batch result applied_group_count 與 group_results 不一致")
    return postconditions


def _verify_manual_actual_batch_postconditions(
    refreshed_ledger: list[dict[str, Any]],
    postconditions: list[Mapping[str, str]],
) -> None:
    by_occurrence_id = {
        str(entry.get("occurrence_id") or ""): entry
        for entry in refreshed_ledger
        if str(entry.get("occurrence_id") or "")
    }
    for condition in postconditions:
        occurrence_id = str(condition.get("occurrence_id") or "")
        entry = by_occurrence_id.get(occurrence_id)
        if entry is None:
            raise ValueError(
                "manual actual batch postcondition 失敗：checked occurrence 已消失："
                f"{occurrence_id}（group={condition.get('group_id') or ''}）"
            )
        wanted = canonical_bopomofo(condition.get("reading"))
        observed = canonical_bopomofo(entry.get("actual"))
        if not wanted or observed != wanted:
            raise ValueError(
                "manual actual batch postcondition 失敗：checked occurrence actual 不符："
                f"{occurrence_id}；要求={wanted or condition.get('reading') or ''}；"
                f"重解結果={observed or '未辨識'}"
            )


def apply_staged_manual_actual_corrections(output_dir: Path) -> dict[str, Any]:
    """Apply the durable GUI queue once, clear once, refresh once, then verify."""
    output_dir = Path(output_dir)
    manifest = json_load_strict(output_dir / "校對工作階段.json")
    validate_manifest_integrity(manifest)
    validate_output_artifact_hashes(manifest)
    db = load_or_initialize_db(output_dir)
    pre_refresh_ledger = materialize_ledger(manifest, db)
    actual_root = project_actual_evidence_root(output_dir)
    staging = load_manual_actual_staging(actual_root)
    frozen_staged_decisions = list(staging.get("staged_groups") or [])
    live_groups = _current_live_groups_for_staged_manual_actual(
        pre_refresh_ledger,
        frozen_staged_decisions,
    )

    batch_result = dict(apply_staged_manual_actual_batch(actual_root, live_groups))
    if int(batch_result.get("applied_group_count") or 0) == 0:
        return {
            **batch_result,
            "cleared_actual_dependent_event_count": 0,
            "refresh_performed": False,
            "refresh_report": "",
            "postcondition_checked_occurrence_ids": [],
        }

    try:
        postconditions = _manual_actual_batch_postconditions(batch_result)
    except Exception as exc:
        raise ManualActualPostApplyError(
            "batch_result_contract",
            exc,
            batch_result=batch_result,
            cleared_event_count=None,
        ) from exc

    affected_occurrence_ids = {
        str(occurrence_id)
        for occurrence_id in (batch_result.get("affected_occurrence_ids") or [])
        if str(occurrence_id)
    }
    try:
        cleared_event_count = _clear_actual_dependent_events(
            output_dir,
            pre_refresh_ledger,
            affected_occurrence_ids,
        )
    except Exception as exc:
        raise ManualActualPostApplyError(
            "clear_actual_dependent_events",
            exc,
            batch_result=batch_result,
            cleared_event_count=None,
        ) from exc

    try:
        report = refresh_actual_project(output_dir)
    except Exception as exc:
        raise ManualActualPostApplyError(
            "refresh_actual_project",
            exc,
            batch_result=batch_result,
            cleared_event_count=cleared_event_count,
        ) from exc

    try:
        refreshed_manifest = json_load_strict(output_dir / "校對工作階段.json")
        validate_manifest_integrity(refreshed_manifest)
        validate_output_artifact_hashes(refreshed_manifest)
        refreshed_db = load_or_initialize_db(output_dir)
        refreshed_ledger = materialize_ledger(refreshed_manifest, refreshed_db)
        _verify_manual_actual_batch_postconditions(refreshed_ledger, postconditions)
    except Exception as exc:
        raise ManualActualPostApplyError(
            "post_refresh_actual_verification",
            exc,
            batch_result=batch_result,
            cleared_event_count=cleared_event_count,
        ) from exc

    return {
        **batch_result,
        "cleared_actual_dependent_event_count": int(cleared_event_count),
        "refresh_performed": True,
        "refresh_report": str(report),
        "postcondition_checked_occurrence_ids": sorted({
            condition["occurrence_id"] for condition in postconditions
        }),
    }


def _write_pipeline_progress(output_dir: Path, phase: str, status_text: str) -> None:
    """Publish lightweight stage progress for the launcher UI.

    This file is not a sealed proofreading session and must never be used as
    proof of completion.  It only lets the GUI reflect [1/3]..[3/3] and fatal
    states before ``校對工作階段.json`` exists.
    """
    json_save(Path(output_dir) / "pipeline_status.json", {
        "version": VERSION,
        "status": "PROCESSING",
        "phase": str(phase),
        "status_text": str(status_text),
        "updated_at": datetime.now().isoformat(timespec="seconds"),
    })


def _has_complete_pre_session_outputs(pdfs: list[Path], actual_dir: Path, cand_dir: Path) -> bool:
    """Whether every PDF already has both stage-1 and stage-2 workbooks.

    This is only a trigger to *attempt* recovery.  ``collect_manifest`` still
    performs the authoritative fingerprint/hash/schema/set reconciliation before
    any unsealed outputs are accepted.
    """
    return bool(pdfs) and all(
        actual_workbook_path(actual_dir, pdf).exists()
        and candidate_workbook_path(cand_dir, pdf).exists()
        for pdf in pdfs
    )


def run_pipeline(input_path: Path, output_dir: Path) -> Path:
    pdfs = resolve_pdfs(input_path)
    if not pdfs:
        raise SystemExit("找不到 PDF。請指定 PDF 檔或含 PDF 的資料夾。")
    return run_pipeline_pdfs(pdfs, output_dir)


def run_pipeline_pdfs(
    pdfs: list[Path],
    output_dir: Path,
    *,
    session_id_override: str | None = None,
    defer_excel_reports: bool = False,
    runtime_root: Path | None = None,
) -> Path:
    pdfs = [Path(pdf).resolve() for pdf in pdfs]
    root=Path(runtime_root or Path(__file__).resolve().parent).resolve()
    source_validation = validate_asset_manifest(root)
    if not source_validation.get("ok"):
        blocked_path = write_pipeline_blocked(output_dir, source_validation, "SOURCE_INVALID")
        raise SourceValidationError(f"核心資料來源驗證失敗；處理被阻擋：{blocked_path}")
    try:
        # One immutable snapshot is the sole global truth view for this whole
        # session.  Fingerprint computation, every PDF decode, and final
        # metadata sealing all receive this same object.
        global_snapshot = GlobalExactGlyphRepository.resolved().load_snapshot()
    except Exception as exc:
        blocked_report = {
            **dict(source_validation),
            "ok": False,
            "errors": [
                *list(source_validation.get("errors") or []),
                f"global exact glyph source validation failed：{type(exc).__name__}: {exc}",
            ],
        }
        blocked_path = write_pipeline_blocked(
            output_dir,
            blocked_report,
            "GLOBAL_EXACT_SOURCE_INVALID",
        )
        raise SourceValidationError(
            f"global exact glyph 資料來源驗證失敗；處理被阻擋：{blocked_path}"
        ) from exc

    output_dir.mkdir(parents=True,exist_ok=True)
    actual_dir=output_dir/"01_實際注音"; cand_dir=output_dir/"02_候選報告"
    actual_dir.mkdir(exist_ok=True); cand_dir.mkdir(exist_ok=True)
    dynamic_actual_root = initialize_project_actual_evidence(output_dir, root)
    # Freeze the previous session before any output file is touched.  Reuse
    # decisions must not be recomputed after the first rewritten workbook,
    # otherwise one changed file would invalidate the old manifest and make all
    # later PDFs look stale.
    try:
        baseline_manifest = load_reuse_baseline_manifest(output_dir)
        expected_chain_fingerprint = compute_expected_asset_fingerprint(
            root,
            source_validation,
            resolver_version=EXPECTED_RESOLVER_VERSION,
            source_files=EXPECTED_RESOLVER_SOURCE_FILES,
        )
        actual_fingerprints: dict[str, dict[str, Any]] = {}
        for pdf in pdfs:
            actual = actual_workbook_path(actual_dir, pdf)
            dependencies = actual_workbook_dynamic_dependencies(actual)
            global_dependencies = actual_workbook_global_exact_dependencies(actual, pdf)
            actual_fingerprints[pdf.name] = compute_actual_asset_fingerprint(
                root,
                pdf,
                source_validation,
                decoder_version=ACTUAL_DECODER_VERSION,
                source_files=ACTUAL_DECODER_SOURCE_FILES,
                global_exact_glyph_evidence_hashes=global_exact_glyph_evidence_hashes(
                    global_snapshot,
                    global_dependencies,
                ),
                dynamic_dependencies=dependencies,
                dynamic_evidence_root=dynamic_actual_root,
            )
    except Exception as exc:
        blocked_report = {
            **dict(source_validation),
            "ok": False,
            "errors": [
                *list(source_validation.get("errors") or []),
                f"actual dependency validation failed：{type(exc).__name__}: {exc}",
            ],
        }
        write_pipeline_blocked(output_dir, blocked_report, "ACTUAL_DEPENDENCY_INVALID")
        raise
    resume_manifest: dict[str, Any] | None = None
    if baseline_manifest is None and _has_complete_pre_session_outputs(pdfs, actual_dir, cand_dir):
        _write_pipeline_progress(
            output_dir,
            "3/3-recover",
            "偵測到前次未封存的 actual／candidate 輸出；先驗證後續接第 3 階段",
        )
        print("[恢復] 偵測到前次未建立工作階段，但 1/3、2/3 輸出齊全；驗證後直接續跑 3/3", flush=True)
        try:
            resume_manifest = collect_manifest(
                pdfs, actual_dir, cand_dir,
                actual_fingerprints=actual_fingerprints,
                source_validation=source_validation,
                runtime_root=root,
            )
            print("  前次輸出驗證通過：略過 actual 重解碼與 candidate 重建。", flush=True)
        except Exception as recovery_exc:
            print(
                "  前次未封存輸出驗證未通過，為避免沿用不一致資料，改跑完整流程："
                f"{type(recovery_exc).__name__}: {recovery_exc}",
                flush=True,
            )
            resume_manifest = None

    batch=root/"cff_zero_map_batch.py"
    try:
        if resume_manifest is None:
            _write_pipeline_progress(output_dir, "1/3", f"[1/3] 整批實際注音解碼：{len(pdfs)} 個 PDF")
            print(f"[1/3] 整批實際注音解碼：{len(pdfs)} 個 PDF",flush=True)
            expected_actual=[]
            redecoded_pdf_names: set[str] = set()
            redecoded_actual_paths: list[Path] = []
            for i,pdf in enumerate(pdfs,1):
                actual=actual_workbook_path(actual_dir,pdf)
                expected_actual.append(actual)
                fingerprint = actual_fingerprints[pdf.name]
                if output_is_reusable(
                    output_dir,
                    pdf,
                    actual,
                    fingerprint,
                    baseline_manifest=baseline_manifest,
                ):
                    print(f"  [{i}/{len(pdfs)}] PDF 與 actual 證據資產相容，reuse：{pdf.name}",flush=True)
                    continue
                print(f"  [{i}/{len(pdfs)}] actual cache 不相容，重新解碼：{pdf.name}",flush=True)
                decode(pdf, actual, root / DEFAULT_MAP.name, root / DEFAULT_GROUPS.name, root / DEFAULT_CFF_MAP.name,
                       root / DEFAULT_XREF_OVERRIDES.name, root / DEFAULT_TRANSFORMS.name, root / DEFAULT_FINGERPRINTS.name,
                       root / DEFAULT_OUTLINE_SIGNATURES.name, root / DEFAULT_MAPPING_CORRECTIONS.name,
                       root / DEFAULT_SYMBOL_TEMPLATES.name, dynamic_actual_root / OCCURRENCE_OVERRIDE_FILE,
                       root / DEFAULT_STRUCTURAL_EXCLUSIONS.name, root / DEFAULT_CFF_CONSENSUS.name,
                       fingerprint, dynamic_actual_root, global_snapshot=global_snapshot)
                # A brand-new/legacy workbook may have required a conservative
                # global-fallback fingerprint before its exact glyph dependencies
                # were known.  Re-seal only the metadata with the final per-PDF
                # dependency fingerprint; pronunciation cells are untouched.
                final_dependencies = actual_workbook_dynamic_dependencies(actual)
                final_global_dependencies = actual_workbook_global_exact_dependencies(actual, pdf)
                if final_global_dependencies is None:
                    raise SourceValidationError(
                        f"新 actual workbook 無法建立完整 global exact dependency roster：{pdf.name}"
                    )
                final_fingerprint = compute_actual_asset_fingerprint(
                    root,
                    pdf,
                    source_validation,
                    decoder_version=ACTUAL_DECODER_VERSION,
                    source_files=ACTUAL_DECODER_SOURCE_FILES,
                    global_exact_glyph_evidence_hashes=global_exact_glyph_evidence_hashes(
                        global_snapshot,
                        final_global_dependencies,
                    ),
                    dynamic_dependencies=final_dependencies,
                    dynamic_evidence_root=dynamic_actual_root,
                )
                if final_fingerprint.get("fingerprint") != fingerprint.get("fingerprint"):
                    rewrite_actual_workbook_fingerprint(actual, final_fingerprint)
                    actual_fingerprints[pdf.name] = final_fingerprint
                redecoded_pdf_names.add(pdf.name)
                redecoded_actual_paths.append(actual)

            # CFF cross-file evidence still reads the whole batch, but v5.5.0 only
            # rewrites workbooks whose actual cache was invalidated.  This preserves
            # byte-stable caches for unrelated PDFs while keeping the same batch
            # evidence pool.
            if redecoded_actual_paths:
                cmd=[
                    sys.executable,"-X","utf8",str(batch),
                    "--existing-workbooks",*map(str,expected_actual),
                    "--write-only-workbooks",*map(str,redecoded_actual_paths),
                    "-o",str(actual_dir),
                ]
                subprocess.run(cmd,cwd=root,check=True,env=utf8_env())
            else:
                print("  actual 全部 reuse；略過 CFF 工作簿改寫", flush=True)

            _write_pipeline_progress(output_dir, "2/3", "[2/3] 詞境／字典校對候選")
            print("[2/3] 詞境／字典校對候選",flush=True)
            current_expected_fp = expected_chain_fingerprint
            for i,pdf in enumerate(pdfs,1):
                actual=actual_workbook_path(actual_dir,pdf)
                cand=candidate_workbook_path(cand_dir,pdf)
                if (
                    pdf.name not in redecoded_pdf_names
                    and candidate_is_baseline_reusable(baseline_manifest, pdf, cand, current_expected_fp)
                ):
                    print(f"  [{i}/{len(pdfs)}] actual 未變且 expected 證據資產相同，reuse：{pdf.name}",flush=True)
                    continue
                print(f"  [{i}/{len(pdfs)}] 重新建立候選：{pdf.name}",flush=True)
                analyze(
                    actual,
                    root / DEFAULT_DICT.name,
                    root / DEFAULT_RULES.name,
                    cand,
                    pdf,
                    root / DEFAULT_REGRESSIONS.name,
                    root / DEFAULT_CHAR_OVERRIDES.name,
                    root / DEFAULT_CONTEXT_OVERRIDES.name,
                    dynamic_actual_root,
                    runtime_root=root,
                    global_snapshot=global_snapshot,
                )

            _write_pipeline_progress(output_dir, "3/3", "[3/3] 建立 occurrence ledger、全量對帳與 completion gate")
            print("[3/3] 建立 occurrence ledger、全量對帳與 completion gate",flush=True)
            manifest=collect_manifest(
                pdfs, actual_dir, cand_dir,
                actual_fingerprints=actual_fingerprints,
                source_validation=source_validation,
                runtime_root=root,
            )
        else:
            manifest = resume_manifest
            _write_pipeline_progress(output_dir, "3/3", "[3/3] 前次輸出已驗證，建立工作階段與 completion gate")
            print("[3/3] 前次輸出已驗證，建立工作階段與 completion gate", flush=True)
        if baseline_manifest is not None:
            manifest = _carry_forward_cross_version_identity(manifest, baseline_manifest)
        if session_id_override:
            manifest["session_id"] = str(session_id_override)
            seal_manifest(manifest)
        json_save(output_dir/"校對工作階段.json",manifest)
        db=load_or_initialize_db(output_dir)
        json_save(output_dir/"人工判定資料庫.json",db)
        pending=save_pending_json(output_dir,manifest,db)
        ledger = materialize_ledger(manifest, db)
        reconciliation = reconcile_ledger(ledger, manifest.get("actual_source_ids") or [])
        gate = completion_gate(ledger, reconciliation, manifest.get("regression_gate") or {}, source_validation_ok=True)
        if defer_excel_reports:
            # Interactive actual correction should refresh the authoritative
            # ledger quickly.  The large Excel technical/user reports and GPT
            # workbook are presentation artifacts and can be regenerated from
            # the sealed manifest on demand via「更新 Excel 報告」.
            report = output_dir / "注音校對_最終報告.xlsx"
            status_text = {
                PROCESSING_FINISHED: "處理結束；校對尚未完成。",
                PROOFREAD_COMPLETE: "全冊注音校對完成。",
                PIPELINE_BLOCKED: "處理被阻擋。",
            }[gate["status"]]
            json_save(output_dir / "pipeline_status.json", {
                "version": VERSION,
                "status": gate["status"],
                "status_text": status_text,
                "completion_gate": gate,
                "reconciliation": reconciliation.as_dict(),
                "user_report": str(report),
                "technical_audit_report": str(output_dir / "注音校對_技術稽核.xlsx"),
                "excel_report_deferred": True,
            })
            gpt_report = None
        else:
            report=generate_report(output_dir,manifest,db,runtime_root=root)
            gpt_report = export_pending_for_gpt(output_dir) if pending else None
        if gate["status"] == PROOFREAD_COMPLETE:
            print(f"全冊注音校對完成。\n報告：{report}",flush=True)
        else:
            print(
                f"處理結束；校對尚未完成。非終態：{pending} 筆；失敗硬門檻：{'、'.join(gate['failed_gates'])}\n報告：{report}",
                flush=True,
            )
        if gpt_report:
            print(f"證據補建／六閘門確認報表：{gpt_report}",flush=True)
        return report
    except Exception as exc:
        blocked_report = dict(source_validation)
        blocked_report.setdefault("errors", []).append(f"runtime fatal error：{type(exc).__name__}: {exc}")
        write_pipeline_blocked(output_dir, blocked_report, "RUNTIME_FATAL_OR_DATA_INTEGRITY_ERROR")
        raise


def _refresh_manifest_from_outputs_legacy_disabled(output_dir: Path, old_manifest: dict[str, Any]) -> dict[str, Any]:
    raise ValueError("LEGACY_SCHEMA：v5.1.5 manifest refresh 已隔離，不得執行")
    """重新讀取既有實際注音／候選工作簿，避免升版後沿用舊的摘要統計。

    不需要重新解析 PDF；沿用工作階段中已記錄的 PDF SHA-256，重新建立
    現版候選紀錄與覆蓋率統計。這也確保「重新產生報告」真的反映目前工作簿。
    """
    manifest = {"version": VERSION, "created_at": datetime.now().isoformat(timespec="seconds"), "pdfs": [], "records": []}
    actual_dir = output_dir / "01_實際注音"
    cand_dir = output_dir / "02_候選報告"
    for old in old_manifest.get("pdfs", []):
        pdf_name = str(old.get("pdf_name") or Path(str(old.get("pdf") or "")).name)
        if not pdf_name:
            continue
        pdf_path = str(old.get("pdf") or pdf_name)
        ph = str(old.get("pdf_sha256") or "")
        stem = Path(pdf_name).stem
        actual = Path(str(old.get("actual_workbook") or ""))
        cand = Path(str(old.get("candidate_workbook") or ""))
        if not actual.exists():
            actual = actual_dir / f"{stem}_實際注音解碼.xlsx"
        if not cand.exists():
            cand = cand_dir / f"{stem}_注音校對候選.xlsx"
        if not actual.exists() or not cand.exists():
            raise FileNotFoundError(f"重新產生報告時缺少輸出：{pdf_name}")
        sm = summary_dict(actual)
        csm = summary_dict(cand)
        info = {
            "pdf": pdf_path, "pdf_name": pdf_name, "pdf_sha256": ph,
            "actual_workbook": str(actual), "candidate_workbook": str(cand),
            "actual_total": int(sm.get("偵測到注音字形筆數") or sm.get("注音結構偵測筆數") or sm.get("偵測到注音結構總筆數") or 0),
            "actual_mapped": int(sm.get("已解碼筆數") or 0),
            "proofread_total": int(csm.get("現版注音座標總數") or sm.get("偵測到注音字形筆數") or 0),
            "expected_uncovered": int(csm.get("未建立預期音") or 0),
            "regression_fail": int(csm.get("回歸測試失敗") or 0),
            "expected_coverage": float(csm.get("預期音資料覆蓋率") or 0),
        }
        manifest["pdfs"].append(info)
        for category in AUTO_ERROR_SHEETS + AUTO_CORRECT_SHEETS + SPECIAL_SHEETS + UNRESOLVED_SHEETS + COVERAGE_GAP_SHEETS:
            for row in workbook_rows(cand, category):
                rid = make_review_id(ph, category, row)
                manifest["records"].append({
                    "review_id": rid,
                    "pdf": pdf_path, "pdf_name": pdf_name, "pdf_sha256": ph,
                    "candidate_workbook": str(cand), "source_sheet": category,
                    "row": row,
                })
    return manifest


def regenerate_report(output_dir: Path) -> Path:
    old_manifest=json_load_strict(output_dir/"校對工作階段.json")
    validate_manifest_integrity(old_manifest)
    validate_output_artifact_hashes(old_manifest)
    if (
        not schema_compatible(old_manifest.get("session_schema_version"), SESSION_SCHEMA_VERSION)
        or not schema_compatible(old_manifest.get("ledger_schema_version"), LEDGER_SCHEMA_VERSION)
        or not review_id_schema_compatible(old_manifest.get("review_id_schema_version"), REVIEW_ID_SCHEMA_VERSION)
    ):
        raise ValueError("SESSION_SCHEMA_INCOMPATIBLE：資料布局或 occurrence identity 規則不同，需有明確轉接器")
    root = Path(__file__).resolve().parent
    source_validation = validate_asset_manifest(root)
    if not source_validation.get("ok"):
        write_pipeline_blocked(output_dir, source_validation, "SOURCE_INVALID_DURING_REGENERATE")
        raise SourceValidationError("資料來源驗證失敗；未覆寫上一份最終報告")
    try:
        global_snapshot = GlobalExactGlyphRepository.resolved().load_snapshot()
    except Exception as exc:
        blocked_report = {
            **dict(source_validation),
            "ok": False,
            "errors": [
                *list(source_validation.get("errors") or []),
                f"global exact glyph source validation failed：{type(exc).__name__}: {exc}",
            ],
        }
        write_pipeline_blocked(
            output_dir,
            blocked_report,
            "GLOBAL_EXACT_SOURCE_INVALID_DURING_REGENERATE",
        )
        raise SourceValidationError("global exact glyph 資料來源驗證失敗；未覆寫最終報告") from exc
    dynamic_actual_root = initialize_project_actual_evidence(output_dir, root)
    pdfs = _resolve_session_pdfs(output_dir, old_manifest)
    fingerprints = {}
    for pdf in pdfs:
        actual_path = actual_workbook_path(output_dir / "01_實際注音", pdf)
        global_dependencies = actual_workbook_global_exact_dependencies(actual_path, pdf)
        fingerprints[pdf.name] = compute_actual_asset_fingerprint(
            root,
            pdf,
            source_validation,
            decoder_version=ACTUAL_DECODER_VERSION,
            source_files=ACTUAL_DECODER_SOURCE_FILES,
            global_exact_glyph_evidence_hashes=global_exact_glyph_evidence_hashes(
                global_snapshot,
                global_dependencies,
            ),
            dynamic_dependencies=actual_workbook_dynamic_dependencies(
                actual_path
            ),
            dynamic_evidence_root=dynamic_actual_root,
        )
    manifest = collect_manifest(
        pdfs,
        output_dir / "01_實際注音",
        output_dir / "02_候選報告",
        actual_fingerprints=fingerprints,
        source_validation=source_validation,
    )
    manifest = _carry_forward_cross_version_identity(manifest, old_manifest)
    db = load_or_initialize_db(output_dir)
    materialize_ledger(manifest, db)
    json_save(output_dir/"校對工作階段.json",manifest)
    save_pending_json(output_dir,manifest,db)
    return generate_report(output_dir,manifest,db)


def repair_project_state(output_dir: Path, *, runtime_root: Path | None = None) -> Path:
    """Rebuild expected candidates and completion state without forcing actual decode.

    This is the supported v5.6.2 upgrade path for an existing sealed project.
    ``run_pipeline_pdfs`` reuses each actual workbook when the exact PDF bytes,
    approved actual assets, and project-owned actual evidence are unchanged. A
    changed expected rule/regression asset rebuilds only candidate workbooks,
    then replays durable manual decisions and regenerates the completion gate.
    """
    output_dir = resolve_existing_project_dir(Path(output_dir))
    old_manifest = json_load_strict(output_dir / "校對工作階段.json")
    validate_manifest_integrity(old_manifest)
    validate_output_artifact_hashes(old_manifest)
    if (
        not schema_compatible(old_manifest.get("session_schema_version"), SESSION_SCHEMA_VERSION)
        or not schema_compatible(old_manifest.get("ledger_schema_version"), LEDGER_SCHEMA_VERSION)
        or not review_id_schema_compatible(old_manifest.get("review_id_schema_version"), REVIEW_ID_SCHEMA_VERSION)
    ):
        raise ValueError("SESSION_SCHEMA_INCOMPATIBLE：現有專案無法安全修復，需明確轉接器")
    # Validate durable decisions before candidate workbooks can be replaced.
    # The generic loader may quarantine unmistakably legacy data for normal UI
    # startup, but repair must never erase/bypass an incompatible decision DB or
    # leave the old manifest pointing at a newly rewritten candidate workbook.
    decision_path = output_dir / "人工判定資料庫.json"
    if decision_path.exists():
        existing_db = normalize_db(json_load_strict(decision_path))
        materialize_ledger(old_manifest, existing_db)
    pdfs = _resolve_session_pdfs(output_dir, old_manifest)
    print(
        "[專案修復] 沿用相同 PDF 與既有人工判定；actual 證據相容時直接 reuse，只重建受影響的 expected 候選與完成門檻。",
        flush=True,
    )
    pipeline_kwargs: dict[str, Any] = {
        "session_id_override": str(old_manifest.get("session_id") or "") or None,
    }
    if runtime_root is not None:
        pipeline_kwargs["runtime_root"] = Path(runtime_root)
    return run_pipeline_pdfs(pdfs, output_dir, **pipeline_kwargs)


def resolve_existing_project_dir(path: Path) -> Path:
    """Resolve a project folder conservatively for management/import operations.

    GUI users can accidentally select either the textbook parent folder or a
    duplicated ``注音校對_輸出/注音校對_輸出`` child.  Never create or
    migrate anything here: only redirect when exactly one adjacent folder
    already contains the sealed session manifest.
    """
    path = Path(path)
    session_name = "校對工作階段.json"
    if (path / session_name).exists():
        return path

    candidates: list[Path] = []
    parent = path.parent
    child = path / "注音校對_輸出"

    # Common accidental duplicate: .../注音校對_輸出/注音校對_輸出
    if path.name == "注音校對_輸出" and (parent / session_name).exists():
        candidates.append(parent)
    # Common selection of the textbook folder instead of its project folder.
    if (child / session_name).exists():
        candidates.append(child)

    # Deduplicate without resolving symlinks/network paths.
    unique: list[Path] = []
    seen: set[str] = set()
    for candidate in candidates:
        key = str(candidate)
        if key not in seen:
            seen.add(key)
            unique.append(candidate)

    return unique[0] if len(unique) == 1 else path


def main():
    ap=argparse.ArgumentParser(description=f"{PROGRAM} v{VERSION}")
    ap.add_argument("input",nargs="?",help="PDF 或含 PDF 的資料夾")
    ap.add_argument("-o","--output-dir")
    ap.add_argument("--report-only",action="store_true",help="只依既有工作階段與判定重新產生最終報告")
    ap.add_argument("--repair-project",action="store_true",help="沿用 actual 與人工判定，重建受影響候選、回歸門檻及報告")
    ap.add_argument("--export-gpt",action="store_true",help="輸出 expected／差異 GPT 證據表")
    ap.add_argument("--import-gpt",help="匯入 expected／差異 GPT 證據表 xlsx")
    ap.add_argument("--import-gpt-auto", nargs="+", help="自動辨識並匯入 GPT xlsx 或單檔 GPT 判定包 zip；多檔時自動先判定包／actual 後 expected")
    ap.add_argument("--export-actual-gpt",action="store_true",help="輸出 actual 字形待判定 GPT 包")
    ap.add_argument("--import-actual-gpt",help="匯入 GPT 已填寫的 actual待判定_給GPT.xlsx，驗證後自動重新解碼")
    ap.add_argument("--refresh-actual",action="store_true",help="依目前 user actual 證據重新解碼現有工作階段")
    args=ap.parse_args()
    if args.report_only or args.repair_project or args.export_gpt or args.import_gpt or args.import_gpt_auto or args.export_actual_gpt or args.import_actual_gpt or args.refresh_actual:
        if not args.output_dir: raise SystemExit("此操作需要 -o 輸出資料夾")
        requested_outdir = Path(args.output_dir)
        outdir = resolve_existing_project_dir(requested_outdir)
        if outdir != requested_outdir:
            print(f"[專案路徑修正] {requested_outdir} -> {outdir}", flush=True)
        if args.repair_project:
            print(repair_project_state(outdir)); return 0
        if args.export_actual_gpt:
            package = export_actual_pending_for_gpt(outdir)
            print(package if package is not None else NO_ACTUAL_PENDING_MESSAGE)
            return 0
        if args.import_gpt_auto:
            planned = plan_gpt_auto_imports(args.import_gpt_auto)
            if not planned:
                raise ValueError("沒有可匯入的 GPT 判定檔")
            if len(planned) > 1:
                print("[GPT 批次匯入] 已自動排序：判定包／actual -> expected；不需手動決定先後。", flush=True)
            for import_path, kind in planned:
                if kind == "bundle":
                    actual_n, removed, expected_n, skipped, report = import_gpt_decision_bundle(outdir, import_path)
                    print(
                        f"自動辨識：{import_path.name} = GPT 單檔判定包。"
                        f"匯入 actual occurrence {actual_n} 筆；撤銷 {removed} 筆依賴舊 actual 的最終確認；"
                        f"匯入 expected／差異 {expected_n} 筆；略過 {skipped} 筆未操作列。\n{report}",
                        flush=True,
                    )
                elif kind == "actual":
                    n, removed, report = import_actual_gpt_decisions(outdir, import_path)
                    print(
                        f"自動辨識：{import_path.name} = actual GPT 判定檔。"
                        f"匯入 actual {n} 組；撤銷 {removed} 筆依賴舊 actual 的最終確認。\n{report}",
                        flush=True,
                    )
                else:
                    n, skipped, report = import_gpt_decisions(outdir, import_path)
                    print(
                        f"自動辨識：{import_path.name} = expected／差異 GPT 證據檔。"
                        f"匯入 {n} 筆，略過 {skipped} 筆未操作列。\n{report}\n"
                        "完成狀態請以本次產生的報告／pipeline_status.json 為準；"
                        "若剛升級規則，請按「修復／更新報告」。",
                        flush=True,
                    )
            return 0
        if args.import_actual_gpt:
            n,removed,report=import_actual_gpt_decisions(outdir,Path(args.import_actual_gpt)); print(f"匯入 actual {n} 組；撤銷 {removed} 筆依賴舊 actual 的人工事件。\n{report}"); return 0
        if args.refresh_actual:
            print(refresh_actual_project(outdir)); return 0
        if args.export_gpt:
            print(export_pending_for_gpt(outdir)); return 0
        if args.import_gpt:
            n,skipped,report=import_gpt_decisions(outdir,Path(args.import_gpt)); print(f"匯入 {n} 筆，略過 {skipped} 筆。\n{report}"); return 0
        print(regenerate_report(outdir)); return 0
    if not args.input: raise SystemExit("請指定 PDF 或 PDF 資料夾。")
    inp=Path(args.input)
    out=Path(args.output_dir) if args.output_dir else inp.parent/(inp.stem+"_注音校對") if inp.is_file() else inp/("注音校對_輸出")
    run_pipeline(inp,out)
    return 0

if __name__=="__main__":
    configure_utf8_stdio()
    raise SystemExit(main())
