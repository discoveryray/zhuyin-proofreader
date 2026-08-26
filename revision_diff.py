from __future__ import annotations

import argparse
import hashlib
import json
import os
import subprocess
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from export_zhuyin_readings import (
    decode, DEFAULT_MAP, DEFAULT_GROUPS, DEFAULT_CFF_MAP, DEFAULT_CFF_CONSENSUS,
    DEFAULT_XREF_OVERRIDES, DEFAULT_TRANSFORMS, DEFAULT_FINGERPRINTS,
    DEFAULT_OUTLINE_SIGNATURES, DEFAULT_MAPPING_CORRECTIONS, DEFAULT_SYMBOL_TEMPLATES,
    DEFAULT_ACTUAL_OVERRIDES, DEFAULT_STRUCTURAL_EXCLUSIONS,
)

PROGRAM = "審次注音差異檢查器"
VERSION = "5.1-handbook-priority"


@dataclass
class PageData:
    side: str
    pdf_name: str
    pdf_path: str
    physical_page: int
    printed_page: str
    rows: list[dict[str, Any]]
    global_index: int = 0

    @property
    def chars(self) -> str:
        return "".join(str(r.get("字元") or "") for r in self.rows)


@dataclass
class PagePair:
    old: PageData
    new: PageData
    method: str
    similarity: float


def configure_utf8_stdio() -> None:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


def utf8_env() -> dict[str, str]:
    env = os.environ.copy()
    env["PYTHONUTF8"] = "1"
    env["PYTHONIOENCODING"] = "utf-8"
    return env


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def json_load(path: Path, default):
    if not path.exists():
        return default
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return default


def json_save(path: Path, data) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)


def resolve_pdfs(input_path: Path) -> list[Path]:
    if input_path.is_file() and input_path.suffix.lower() == ".pdf":
        return [input_path.resolve()]
    if input_path.is_dir():
        return [p.resolve() for p in sorted(input_path.glob("*.pdf"))]
    return []


def actual_path(decoded_dir: Path, pdf: Path) -> Path:
    return decoded_dir / f"{pdf.stem}_實際注音解碼.xlsx"


def _cache_ok(manifest: dict[str, Any], pdf: Path, out: Path) -> bool:
    if not out.exists() or manifest.get("version") != VERSION:
        return False
    try:
        digest = sha256_file(pdf)
    except Exception:
        return False
    return any(x.get("pdf_name") == pdf.name and x.get("sha256") == digest for x in manifest.get("pdfs", []))


def decode_revision(input_path: Path, side_dir: Path, label: str) -> list[tuple[Path, Path]]:
    pdfs = resolve_pdfs(input_path)
    if not pdfs:
        raise FileNotFoundError(f"{label}找不到 PDF：{input_path}")
    decoded = side_dir / "實際注音"
    decoded.mkdir(parents=True, exist_ok=True)
    manifest_path = side_dir / "解碼快取.json"
    old_manifest = json_load(manifest_path, {})
    results: list[tuple[Path, Path]] = []

    print(f"[{label}] 共 {len(pdfs)} 個 PDF", flush=True)
    for i, pdf in enumerate(pdfs, 1):
        out = actual_path(decoded, pdf)
        results.append((pdf, out))
        if _cache_ok(old_manifest, pdf, out):
            print(f"  [{i}/{len(pdfs)}] 未變更，沿用安全解碼：{pdf.name}", flush=True)
            continue
        print(f"  [{i}/{len(pdfs)}] 解碼：{pdf.name}", flush=True)
        decode(
            pdf, out, DEFAULT_MAP, DEFAULT_GROUPS, DEFAULT_CFF_MAP,
            DEFAULT_XREF_OVERRIDES, DEFAULT_TRANSFORMS, DEFAULT_FINGERPRINTS,
            DEFAULT_OUTLINE_SIGNATURES, DEFAULT_MAPPING_CORRECTIONS,
            DEFAULT_SYMBOL_TEMPLATES, DEFAULT_ACTUAL_OVERRIDES,
            DEFAULT_STRUCTURAL_EXCLUSIONS, DEFAULT_CFF_CONSENSUS,
        )

    # 每個審次獨立建立自己的 CFF 工作階段；不讓本審與上一審互相提供自舉證據。
    root = Path(__file__).resolve().parent
    batch = root / "cff_zero_map_batch.py"
    cmd = [sys.executable, "-X", "utf8", str(batch), "--existing-workbooks", *[str(x[1]) for x in results], "-o", str(decoded)]
    subprocess.run(cmd, cwd=root, check=True, env=utf8_env())

    manifest = {
        "version": VERSION,
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "pdfs": [{"pdf_name": p.name, "pdf_path": str(p), "sha256": sha256_file(p), "actual": str(w)} for p, w in results],
    }
    json_save(manifest_path, manifest)
    return results


def _load_actual_rows(workbook: Path) -> list[dict[str, Any]]:
    wb = load_workbook(workbook, read_only=True, data_only=True)
    if "實際注音" not in wb.sheetnames:
        raise ValueError(f"工作簿缺少『實際注音』：{workbook}")
    ws = wb["實際注音"]
    headers = [str(c.value or "") for c in ws[1]]
    ix = {h: i for i, h in enumerate(headers)}
    out = []
    for vals in ws.iter_rows(min_row=2, values_only=True):
        if not vals or not any(v not in (None, "") for v in vals):
            continue
        d = {h: vals[i] if i < len(vals) else None for h, i in ix.items()}
        # 審次差異只比較實際可見版面；版外物件不應造成美編警報。
        try:
            if float(d.get("x0") or 0) < 0 or float(d.get("y0") or 0) < 0:
                continue
        except Exception:
            pass
        out.append(d)
    return out


def build_pages(items: list[tuple[Path, Path]], side: str) -> list[PageData]:
    pages: list[PageData] = []
    g = 0
    for pdf, wb_path in items:
        rows = _load_actual_rows(wb_path)
        by_page: dict[int, list[dict[str, Any]]] = defaultdict(list)
        for r in rows:
            try:
                p = int(r.get("實體頁碼") or 0)
            except Exception:
                p = 0
            by_page[p].append(r)
        for physical in sorted(by_page):
            rs = by_page[physical]
            # export_zhuyin_readings 已輸出閱讀順序；保留該順序，不再依座標重排。
            printed = ""
            for r in rs:
                v = r.get("課本頁")
                if v not in (None, ""):
                    printed = str(v)
                    break
            pages.append(PageData(side, pdf.name, str(pdf), physical, printed, rs, g))
            g += 1
    return pages


def _ratio(a: str, b: str) -> float:
    if not a and not b:
        return 1.0
    if not a or not b:
        return 0.0
    return SequenceMatcher(None, a, b, autojunk=False).ratio()


def pair_pages(old_pages: list[PageData], new_pages: list[PageData]) -> tuple[list[PagePair], list[PageData], list[PageData]]:
    pairs: list[PagePair] = []
    used_old: set[int] = set()
    used_new: set[int] = set()

    old_by_label: dict[str, list[int]] = defaultdict(list)
    new_by_label: dict[str, list[int]] = defaultdict(list)
    for i, p in enumerate(old_pages):
        if p.printed_page:
            old_by_label[p.printed_page].append(i)
    for j, p in enumerate(new_pages):
        if p.printed_page:
            new_by_label[p.printed_page].append(j)

    # 第一層：唯一課本頁碼。為避免改頁後同頁碼硬配錯頁，仍要求最低內容相似度。
    for label in sorted(set(old_by_label) & set(new_by_label)):
        oi = old_by_label[label]
        nj = new_by_label[label]
        if len(oi) == 1 and len(nj) == 1:
            i, j = oi[0], nj[0]
            score = _ratio(old_pages[i].chars, new_pages[j].chars)
            if score >= 0.35:
                pairs.append(PagePair(old_pages[i], new_pages[j], "課本頁碼一致", score))
                used_old.add(i); used_new.add(j)

    # 第二層：剩餘頁以注音中文字序列做高相似度配對；不使用注音本身，避免目標洩漏。
    candidates: list[tuple[float, float, int, int]] = []
    maxn = max(len(old_pages), len(new_pages), 1)
    for i, op in enumerate(old_pages):
        if i in used_old or not op.chars:
            continue
        for j, np in enumerate(new_pages):
            if j in used_new or not np.chars:
                continue
            score = _ratio(op.chars, np.chars)
            if score < 0.72:
                continue
            order_gap = abs(op.global_index - np.global_index) / maxn
            candidates.append((score, -order_gap, i, j))
    for score, neg_gap, i, j in sorted(candidates, reverse=True):
        if i in used_old or j in used_new:
            continue
        pairs.append(PagePair(old_pages[i], new_pages[j], "內容序列高相似度", score))
        used_old.add(i); used_new.add(j)

    pairs.sort(key=lambda p: (p.old.global_index, p.new.global_index))
    return pairs, [p for i, p in enumerate(old_pages) if i not in used_old], [p for j, p in enumerate(new_pages) if j not in used_new]


def _context(rows: list[dict[str, Any]], index: int, radius: int = 7) -> str:
    a = max(0, index - radius); b = min(len(rows), index + radius + 1)
    return "".join(str(r.get("字元") or "") for r in rows[a:b])


def _font_changed(a: dict[str, Any], b: dict[str, Any]) -> str:
    return "是" if str(a.get("font") or "") != str(b.get("font") or "") else "否"


def _base_event(pair: PagePair, old_row: dict[str, Any] | None, new_row: dict[str, Any] | None) -> dict[str, Any]:
    return {
        "上一審PDF": pair.old.pdf_name if old_row is not None else pair.old.pdf_name,
        "本審PDF": pair.new.pdf_name if new_row is not None else pair.new.pdf_name,
        "上一審課本頁": pair.old.printed_page,
        "本審課本頁": pair.new.printed_page,
        "上一審實體頁": pair.old.physical_page,
        "本審實體頁": pair.new.physical_page,
        "頁面配對方式": pair.method,
        "頁面文字相似度": pair.similarity,
    }


def compare_pair(pair: PagePair) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], int]:
    high: list[dict[str, Any]] = []
    changed: list[dict[str, Any]] = []
    adddel: list[dict[str, Any]] = []
    uncertain: list[dict[str, Any]] = []
    text_only = 0
    a = pair.old.rows; b = pair.new.rows
    sa = pair.old.chars; sb = pair.new.chars
    sm = SequenceMatcher(None, sa, sb, autojunk=False)

    def emit_add(row: dict[str, Any], idx: int):
        e = _base_event(pair, None, row)
        e.update({
            "變更": "新增", "字元": row.get("字元"), "注音": row.get("實際注音"),
            "font": row.get("font"), "上下文": _context(b, idx),
            "x0": row.get("x0"), "y0": row.get("y0"), "穩定注音鍵": row.get("穩定注音鍵"),
        }); adddel.append(e)

    def emit_del(row: dict[str, Any], idx: int):
        e = _base_event(pair, row, None)
        e.update({
            "變更": "刪除", "字元": row.get("字元"), "注音": row.get("實際注音"),
            "font": row.get("font"), "上下文": _context(a, idx),
            "x0": row.get("x0"), "y0": row.get("y0"), "穩定注音鍵": row.get("穩定注音鍵"),
        }); adddel.append(e)

    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            for i, j in zip(range(i1, i2), range(j1, j2)):
                old = a[i]; new = b[j]
                ro = str(old.get("實際注音") or ""); rn = str(new.get("實際注音") or "")
                if not ro or not rn:
                    e = _base_event(pair, old, new)
                    e.update({"原因":"同字位置至少一審未安全解碼", "字元":old.get("字元"), "上一審注音":ro, "本審注音":rn,
                              "上一審上下文":_context(a,i), "本審上下文":_context(b,j)})
                    uncertain.append(e)
                    continue
                if ro != rn:
                    e = _base_event(pair, old, new)
                    e.update({
                        "風險": "最高", "字元": old.get("字元"), "上一審注音": ro, "本審注音": rn,
                        "上一審font": old.get("font"), "本審font": new.get("font"), "字型是否變更": _font_changed(old,new),
                        "上一審穩定注音鍵": old.get("穩定注音鍵"), "本審穩定注音鍵": new.get("穩定注音鍵"),
                        "上一審上下文": _context(a, i), "本審上下文": _context(b, j),
                        "上一審x0": old.get("x0"), "上一審y0": old.get("y0"), "本審x0": new.get("x0"), "本審y0": new.get("y0"),
                        "判定說明": "中文字序列未變，但 PDF 實際注音改變；優先檢查是否為排版／注音子字型被誤換。",
                    })
                    high.append(e)
            continue

        if tag == "delete":
            for i in range(i1, i2): emit_del(a[i], i)
            continue
        if tag == "insert":
            for j in range(j1, j2): emit_add(b[j], j)
            continue

        # replace：小範圍等長或近等長改字，可安全做位置配對；大段改寫則以刪除＋新增呈現。
        old_len, new_len = i2 - i1, j2 - j1
        if max(old_len, new_len) <= 6 and abs(old_len - new_len) <= 2:
            paired = min(old_len, new_len)
            for k in range(paired):
                i, j = i1 + k, j1 + k
                old, new = a[i], b[j]
                ro = str(old.get("實際注音") or ""); rn = str(new.get("實際注音") or "")
                if not ro or not rn:
                    e = _base_event(pair, old, new)
                    e.update({"原因":"改字位置其中一審未安全解碼", "上一審字元":old.get("字元"), "本審字元":new.get("字元"),
                              "上一審注音":ro, "本審注音":rn, "上一審上下文":_context(a,i), "本審上下文":_context(b,j)})
                    uncertain.append(e)
                elif ro != rn:
                    e = _base_event(pair, old, new)
                    e.update({
                        "風險":"中", "上一審字元":old.get("字元"), "本審字元":new.get("字元"),
                        "上一審注音":ro, "本審注音":rn,
                        "上一審font":old.get("font"), "本審font":new.get("font"), "字型是否變更":_font_changed(old,new),
                        "上一審上下文":_context(a,i), "本審上下文":_context(b,j),
                        "判定說明":"中文字有修改，實際注音也隨之改變；通常屬內容改稿，但仍應確認改音是否符合修訂意圖。",
                    })
                    changed.append(e)
                else:
                    text_only += 1
            for i in range(i1 + paired, i2): emit_del(a[i], i)
            for j in range(j1 + paired, j2): emit_add(b[j], j)
        else:
            for i in range(i1, i2): emit_del(a[i], i)
            for j in range(j1, j2): emit_add(b[j], j)

    return high, changed, adddel, uncertain, text_only


def compare_pages(old_pages: list[PageData], new_pages: list[PageData]):
    pairs, old_unpaired, new_unpaired = pair_pages(old_pages, new_pages)
    high: list[dict[str, Any]] = []
    changed: list[dict[str, Any]] = []
    adddel: list[dict[str, Any]] = []
    uncertain: list[dict[str, Any]] = []
    text_only = 0
    for pair in pairs:
        h, c, a, u, t = compare_pair(pair)
        high.extend(h); changed.extend(c); adddel.extend(a); uncertain.extend(u); text_only += t

    # 整頁無法配對時，把該頁所有注音字列為新增／刪除；不猜頁面對應。
    for p in old_unpaired:
        dummy = PagePair(p, PageData("new", "", "", 0, "", [], 10**9), "上一審頁面無對應", 0.0)
        for i, r in enumerate(p.rows):
            e = _base_event(dummy, r, None)
            e.update({"變更":"刪除", "字元":r.get("字元"), "注音":r.get("實際注音"), "font":r.get("font"),
                      "上下文":_context(p.rows,i), "x0":r.get("x0"), "y0":r.get("y0"), "穩定注音鍵":r.get("穩定注音鍵")})
            adddel.append(e)
    for p in new_unpaired:
        dummy = PagePair(PageData("old", "", "", 0, "", [], -1), p, "本審新增頁面", 0.0)
        for j, r in enumerate(p.rows):
            e = _base_event(dummy, None, r)
            e.update({"變更":"新增", "字元":r.get("字元"), "注音":r.get("實際注音"), "font":r.get("font"),
                      "上下文":_context(p.rows,j), "x0":r.get("x0"), "y0":r.get("y0"), "穩定注音鍵":r.get("穩定注音鍵")})
            adddel.append(e)
    return pairs, old_unpaired, new_unpaired, high, changed, adddel, uncertain, text_only


def _style_sheet(ws, freeze: str = "A2") -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = freeze
    if ws.max_row >= 1:
        for c in ws[1]:
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.font = Font(color="FFFFFF", bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for col in range(1, ws.max_column + 1):
        max_len = 0
        for r in range(1, min(ws.max_row, 250) + 1):
            v = ws.cell(r, col).value
            if v is not None:
                max_len = max(max_len, min(60, len(str(v))))
        ws.column_dimensions[get_column_letter(col)].width = max(10, min(42, max_len + 2))


def _append_dict_sheet(wb: Workbook, name: str, rows: list[dict[str, Any]], preferred: list[str]) -> None:
    ws = wb.create_sheet(name)
    keys: list[str] = []
    seen = set()
    for k in preferred:
        if k not in seen:
            keys.append(k); seen.add(k)
    for r in rows:
        for k in r:
            if k not in seen:
                keys.append(k); seen.add(k)
    ws.append(keys)
    for r in rows:
        ws.append([r.get(k, "") for k in keys])
    _style_sheet(ws)
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(1, c).value or "")
        if "相似度" in h:
            for rr in range(2, ws.max_row + 1): ws.cell(rr, c).number_format = "0.0%"
    if name == "最高風險_同字異音":
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.fill = PatternFill("solid", fgColor="FCE4D6")


def make_report(output_dir: Path, old_input: Path, new_input: Path, pairs: list[PagePair], old_unpaired: list[PageData], new_unpaired: list[PageData], high, changed, adddel, uncertain, text_only: int) -> Path:
    wb = Workbook()
    ws = wb.active; ws.title = "摘要"
    ws.append(["指標", "結果"])
    summary = [
        ("工具", f"{PROGRAM} v{VERSION}"),
        ("上一審來源", str(old_input)),
        ("本審來源", str(new_input)),
        ("已配對注音頁面", len(pairs)),
        ("上一審未配對頁面", len(old_unpaired)),
        ("本審未配對頁面", len(new_unpaired)),
        ("最高風險：中文字相同、注音不同", len(high)),
        ("中文字修改、注音也改變", len(changed)),
        ("新增／刪除的注音字", len(adddel)),
        ("無法安全比較", len(uncertain)),
        ("文字修改但注音相同（僅統計，不列差異）", text_only),
        ("安全原則", "只比較兩審 PDF 實際解出的注音；不使用字典／語意反推。中文字相同但注音改變列最高風險。任一側未安全解碼即不硬判。"),
        ("建立時間", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for row in summary: ws.append(row)
    ws.column_dimensions["A"].width = 42; ws.column_dimensions["B"].width = 100
    _style_sheet(ws)

    _append_dict_sheet(wb, "最高風險_同字異音", high, [
        "風險","上一審PDF","本審PDF","上一審課本頁","本審課本頁","字元","上一審注音","本審注音",
        "上一審font","本審font","字型是否變更","上一審上下文","本審上下文","頁面配對方式","頁面文字相似度","判定說明"
    ])
    _append_dict_sheet(wb, "文字與注音皆變", changed, [
        "風險","上一審PDF","本審PDF","上一審課本頁","本審課本頁","上一審字元","本審字元","上一審注音","本審注音",
        "上一審font","本審font","字型是否變更","上一審上下文","本審上下文","頁面配對方式","頁面文字相似度","判定說明"
    ])
    _append_dict_sheet(wb, "新增刪除注音字", adddel, [
        "變更","上一審PDF","本審PDF","上一審課本頁","本審課本頁","字元","注音","font","上下文","頁面配對方式","頁面文字相似度","穩定注音鍵"
    ])
    _append_dict_sheet(wb, "無法安全比較", uncertain, [
        "原因","上一審PDF","本審PDF","上一審課本頁","本審課本頁","字元","上一審字元","本審字元","上一審注音","本審注音","上一審上下文","本審上下文","頁面配對方式","頁面文字相似度"
    ])

    page_rows = []
    for p in pairs:
        page_rows.append({
            "上一審PDF":p.old.pdf_name,"上一審實體頁":p.old.physical_page,"上一審課本頁":p.old.printed_page,
            "本審PDF":p.new.pdf_name,"本審實體頁":p.new.physical_page,"本審課本頁":p.new.printed_page,
            "配對方式":p.method,"中文字序列相似度":p.similarity,"上一審注音字數":len(p.old.rows),"本審注音字數":len(p.new.rows),
        })
    for p in old_unpaired:
        page_rows.append({"上一審PDF":p.pdf_name,"上一審實體頁":p.physical_page,"上一審課本頁":p.printed_page,"配對方式":"未配對：本審無安全對應","中文字序列相似度":0,"上一審注音字數":len(p.rows)})
    for p in new_unpaired:
        page_rows.append({"本審PDF":p.pdf_name,"本審實體頁":p.physical_page,"本審課本頁":p.printed_page,"配對方式":"未配對：本審新增頁面","中文字序列相似度":0,"本審注音字數":len(p.rows)})
    _append_dict_sheet(wb, "頁面配對稽核", page_rows, ["上一審PDF","上一審實體頁","上一審課本頁","本審PDF","本審實體頁","本審課本頁","配對方式","中文字序列相似度","上一審注音字數","本審注音字數"])

    wi = wb.create_sheet("執行資訊"); wi.append(["項目","內容"])
    wi.append(["模式","審次注音差異檢查"])
    wi.append(["解碼隔離","上一審與本審各自獨立解碼與 CFF 批次自舉，兩審不得互相提供自舉證據。"])
    wi.append(["頁面配對","優先用唯一課本頁碼＋中文字序列相似度；其餘以高相似中文字序列配對。配對不用注音值。"])
    wi.append(["字元對齊","配對頁面內以注音中文字序列做序列對齊；同字同位置再比較 PDF 實際注音。"])
    wi.append(["最高風險定義","中文字相同且對齊，但實際注音不同。這類最符合美編／注音子字型被非預期切換的風險。"])
    wi.append(["可棄權","任一側注音未安全解碼時不判差異，列入『無法安全比較』。"])
    _style_sheet(wi); wi.column_dimensions["A"].width=28; wi.column_dimensions["B"].width=110

    out = output_dir / "審次注音差異報告.xlsx"
    wb.save(out)
    return out


def run_revision_diff(previous: Path, current: Path, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    print("[1/4] 上一審：獨立解碼", flush=True)
    old_items = decode_revision(previous, output_dir / "01_上一審", "上一審")
    print("[2/4] 本審：獨立解碼", flush=True)
    new_items = decode_revision(current, output_dir / "02_本審", "本審")
    print("[3/4] 頁面與注音字對齊", flush=True)
    old_pages = build_pages(old_items, "old")
    new_pages = build_pages(new_items, "new")
    result = compare_pages(old_pages, new_pages)
    pairs, old_unpaired, new_unpaired, high, changed, adddel, uncertain, text_only = result
    print("[4/4] 產生差異報告", flush=True)
    report = make_report(output_dir, previous, current, pairs, old_unpaired, new_unpaired, high, changed, adddel, uncertain, text_only)
    audit = {
        "version":VERSION,"previous":str(previous),"current":str(current),
        "paired_pages":len(pairs),"old_unpaired_pages":len(old_unpaired),"new_unpaired_pages":len(new_unpaired),
        "same_char_pron_changed":len(high),"text_and_pron_changed":len(changed),"added_deleted":len(adddel),"uncertain":len(uncertain),"text_only":text_only,
        "report":str(report),"created_at":datetime.now().isoformat(timespec="seconds"),
    }
    json_save(output_dir / "審次差異工作階段.json", audit)
    print(f"完成。最高風險 {len(high)} 筆；文字與注音皆變 {len(changed)} 筆；新增／刪除 {len(adddel)} 筆；無法安全比較 {len(uncertain)} 筆。\n報告：{report}", flush=True)
    return report


def main() -> int:
    ap = argparse.ArgumentParser(description=f"{PROGRAM} v{VERSION}")
    ap.add_argument("previous", help="上一審 PDF 或含 PDF 的資料夾")
    ap.add_argument("current", help="本審 PDF 或含 PDF 的資料夾")
    ap.add_argument("-o", "--output-dir", required=True)
    args = ap.parse_args()
    run_revision_diff(Path(args.previous), Path(args.current), Path(args.output_dir))
    return 0


if __name__ == "__main__":
    configure_utf8_stdio()
    raise SystemExit(main())
