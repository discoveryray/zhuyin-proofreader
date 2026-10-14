"""Lossless XLSX writes for rows whose Python strings are contract text.

Use only on text/data writers, not on cells intended to contain formulas.
Native numbers, booleans and None retain openpyxl's existing semantics.
"""

import re

from openpyxl.cell.cell import Cell
from openpyxl.cell.read_only import EMPTY_CELL
from openpyxl.cell.text import Text
from openpyxl.reader.excel import ExcelReader
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.worksheet._reader import WorkSheetParser, INLINE_STRING, FORMULA_TAG
from openpyxl.xml.constants import SHEET_MAIN_NS, SHARED_STRINGS
from openpyxl.xml.functions import iterparse


_STRING_ESCAPE = re.compile(r"_x([0-9A-Fa-f]{4})_")


def _decode_string_fragments(fragments):
    # Decode original ST_Xstring tokens once. In particular _x005F_x000D_
    # produces literal _x000D_; the output is never scanned a second time.
    escaped = "".join(_STRING_ESCAPE.sub(lambda match: chr(int(match[1], 16)), text)
                      for text in fragments)
    try:
        # OOXML escapes represent UTF-16 units, including surrogate pairs.
        return escaped.encode("utf-16-le", "surrogatepass").decode("utf-16-le", "strict")
    except UnicodeError as exc:
        raise ValueError("XLSX text contains an unpaired UTF-16 surrogate") from exc


def _raw_string(node):
    text = Text.from_tree(node)
    fragments = ([text.plain] if text.plain is not None else [])
    fragments.extend(block.t for block in text.formatted if block.t is not None)
    # Tokens cannot be invented by joining two separate rich-text runs.
    return _decode_string_fragments(fragments)


class _TextParser(WorkSheetParser):
    def parse_cell(self, element):
        cell = super().parse_cell(element)
        if element.get("t") == "inlineStr" and element.find(FORMULA_TAG) is None:
            child = element.find(INLINE_STRING)
            if child is not None:
                cell["value"] = _raw_string(child)
        return cell


class _TextWorksheet(ReadOnlyWorksheet):
    def _cells_by_row(self, min_col, min_row, max_col, max_row, values_only=False):
        # Finite openpyxl 3.1.5 row iteration; only parser selection differs.
        # Inherit its cell construction, dimensions and sparse-row semantics.
        filler = None if values_only else EMPTY_CELL
        max_col = max_col or self.max_column
        max_row = max_row or self.max_row
        empty_row = (filler,) * (max_col + 1 - min_col) if max_col is not None else []
        counter, idx = min_row, 1
        with self._get_source() as stream:
            parser = _TextParser(stream, self._shared_strings,
                data_only=self.parent.data_only, epoch=self.parent.epoch,
                date_formats=self.parent._date_formats,
                timedelta_formats=self.parent._timedelta_formats)
            for idx, row in parser.parse():
                if max_row is not None and idx > max_row:
                    break
                for _ in range(counter, idx):
                    counter += 1
                    yield empty_row
                if counter <= idx:
                    yield self._get_row(row, min_col, max_col, values_only)
                    counter += 1
        if max_row is not None and max_row < idx:
            for _ in range(counter, max_row + 1):
                yield empty_row


class _TextReader(ExcelReader):
    def read_strings(self):
        content_type = self.package.find(SHARED_STRINGS)
        if content_type is not None:
            with self.archive.open(content_type.PartName[1:]) as stream:
                self.shared_strings = []
                for _, node in iterparse(stream):
                    if node.tag == "{%s}si" % SHEET_MAIN_NS:
                        self.shared_strings.append(_raw_string(node))
                        node.clear()

    def read_worksheets(self):
        super().read_worksheets()
        for index, original in enumerate(self.wb._sheets):
            if isinstance(original, ReadOnlyWorksheet):
                sheet = _TextWorksheet(self.wb, original.title,
                    original._worksheet_path, self.shared_strings)
                sheet.sheet_state = original.sheet_state
                sheet.defined_names = original.defined_names
                self.wb._sheets[index] = sheet


def load_text_workbook(path, *, data_only=True):
    """Decode raw shared/inline text in one owned read-only workbook handle."""
    import openpyxl
    if openpyxl.__version__ != "3.1.5":
        raise ValueError("XLSX text reader requires the supported openpyxl 3.1.5 contract")
    reader = _TextReader(path, read_only=True, data_only=data_only)
    try:
        reader.read()
        return reader.wb
    except BaseException:
        reader.archive.close()
        raise


class _TransportTextCell(Cell):
    __slots__ = ()

    @property
    def value(self):
        # Public values, sheet.values and carried proof always see the original.
        return Cell._value.__get__(self, type(self))

    @value.setter
    def value(self, value):
        Cell.value.fset(self, value)

    def __getattribute__(self, name):
        value = super().__getattribute__(name)
        if name == "_value" and self.data_type == "s" and isinstance(value, str):
            # The openpyxl serializer alone reads the internal transport value.
            # Escape only the leading underscore of an escape-looking token.
            return _STRING_ESCAPE.sub(lambda match: "_x005F_" + match[0][1:], value)
        return value


def _transport_text_cell(sheet, cell, value):
    if _STRING_ESCAPE.search(value) and not isinstance(cell, _TransportTextCell):
        # Same slots/layout: retain the cell object, style, comments, links and
        # their owners. Only declared text serialization adopts this local type.
        cell.__class__ = _TransportTextCell
    return cell


def _validate_text(sheet, row, column, value, field=None):
    if not isinstance(value, str):
        return
    field = field or (sheet.cell(1, column).value if row > 1 else value)
    location = f"{sheet.title}/{field or get_column_letter(column)} ({get_column_letter(column)}{row})"
    # XLSX strings use XML 1.0. Reject before openpyxl can truncate the value or
    # let the XML serializer fail without naming the affected source field.
    for char in value:
        code = ord(char)
        if not (code in (9, 10, 13) or 0x20 <= code <= 0xD7FF
                or 0xE000 <= code <= 0xFFFD or 0x10000 <= code <= 0x10FFFF):
            raise ValueError(f"XLSX text {location}: XML 1.0 cannot preserve U+{code:04X}")
    units = sum(2 if ord(char) > 0xFFFF else 1 for char in value)
    if units > 32767:
        raise ValueError(f"XLSX text {location}: {units} UTF-16 units exceed the 32767 cell limit")


def write_text_value(sheet, row, column, value, *, field=None):
    """Write a declared text value (or an unchanged native scalar) to a cell."""
    _validate_text(sheet, row, column, value, field)
    cell = sheet.cell(row, column)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"
        cell.number_format = "@"
        cell = _transport_text_cell(sheet, cell, value)
    return cell


def append_text_row(sheet, values):
    """Append contract data; all strings are literal text regardless of content."""
    values = list(values)
    row = sheet._current_row + 1
    # Validate the entire row before append, which otherwise silently truncates
    # long strings. Never change values, whitespace, punctuation or prefixes.
    for column, value in enumerate(values, 1):
        _validate_text(sheet, row, column, value)
    sheet.append(values)
    for column, value in enumerate(values, 1):
        if isinstance(value, str):
            cell = sheet.cell(row, column)
            cell.data_type = "s"
            cell.number_format = "@"
            _transport_text_cell(sheet, cell, value)
