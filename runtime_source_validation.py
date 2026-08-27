from __future__ import annotations

import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence
from unicodedata import normalize

from openpyxl import load_workbook

from occurrence_ledger import LEDGER_SCHEMA_VERSION, canonical_bopomofo
from actual_review import dynamic_actual_hashes
from cross_version_compat import fingerprint_contract_components


ASSET_MANIFEST_SCHEMA_VERSION = "2.6.2"
ACTUAL_FINGERPRINT_SCHEMA_VERSION = "2.9.0"
EXPECTED_FINGERPRINT_SCHEMA_VERSION = "2.7.0"
TOOL_VERSION = "5.6.2"
REQUIRED_ASSET_CHAINS = {
    "moe_concise_dictionary": "expected",
    "project_polyphonic_dictionary": "expected",
    "pronunciation_lexical_rules": "expected",
    "handbook_pronunciation_rules": "expected",
    "handbook_pronunciation_constraints": "expected",
    "pdf_occurrence_regressions": "expected",
    "mandatory_regression_cases": "expected",
    "character_overrides": "expected",
    "source_context_overrides": "expected",
    "zhuyin_component_map": "actual",
    "font_compatibility_groups": "actual",
    "cff_bopomofo_symbol_map": "actual",
    "cff_crossfamily_cid_consensus": "actual",
    "ttf_xref_component_overrides": "actual",
    "ttf_verified_component_transforms": "actual",
    "ttf_verified_glyf_fingerprints": "actual",
    "ttf_verified_outline_signatures": "actual",
    "ttf_verified_mapping_corrections": "actual",
    "ttf_verified_symbol_templates": "actual",
    "structural_detection_exclusions": "actual",
}


class SourceValidationError(RuntimeError):
    pass


def _json_without_duplicate_keys(text: str) -> Any:
    def reject_duplicate_keys(pairs):
        out = {}
        for key, value in pairs:
            if key in out:
                raise ValueError(f"JSON 重複 key：{key}")
            out[key] = value
        return out

    return json.loads(text, object_pairs_hook=reject_duplicate_keys)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        headers = [str(value or "").strip() for value in (reader.fieldnames or [])]
        rows = [dict(row) for row in reader]
    return headers, rows


def _blank_required_cells(rows: Sequence[Mapping[str, Any]], columns: Sequence[str]) -> list[str]:
    errors: list[str] = []
    for index, row in enumerate(rows, 2):
        for column in columns:
            if not normalize("NFKC", str(row.get(column) or "")).strip():
                errors.append(f"row {index} column {column}: blank")
                if len(errors) >= 50:
                    return errors
    return errors


def _blank_composite_keys(rows: Sequence[Mapping[str, Any]], columns: Sequence[str]) -> list[str]:
    if not columns:
        return []
    return [
        f"row {index}: all key columns blank ({', '.join(columns)})"
        for index, row in enumerate(rows, 2)
        if not any(normalize("NFKC", str(row.get(column) or "")).strip() for column in columns)
    ][:50]


def validate_duplicate_keys(rows: Sequence[Mapping[str, Any]], key_columns: Sequence[str], value_columns: Sequence[str] = ()) -> dict[str, Any]:
    groups: dict[tuple[str, ...], list[Mapping[str, Any]]] = defaultdict(list)
    for row in rows:
        key = tuple(normalize("NFKC", str(row.get(column) or "")).strip() for column in key_columns)
        if any(key):
            groups[key].append(row)
    duplicate_keys = [key for key, items in groups.items() if len(items) > 1]
    conflicting_keys = []
    if value_columns:
        for key in duplicate_keys:
            values = {
                tuple(normalize("NFKC", str(row.get(column) or "")).strip() for column in value_columns)
                for row in groups[key]
            }
            if len(values) > 1:
                conflicting_keys.append(key)
    return {
        "ok": not duplicate_keys,
        "duplicate_key_count": len(duplicate_keys),
        "conflicting_duplicate_key_count": len(conflicting_keys),
        "duplicate_keys": [list(key) for key in duplicate_keys[:50]],
        "conflicting_duplicate_keys": [list(key) for key in conflicting_keys[:50]],
    }


def _split_bopomofo_cell(value: Any, mode: str) -> list[str]:
    text = normalize("NFKC", str(value or "")).strip()
    if not text:
        return []
    parts = re.split(r"[;；|]", text)
    out = []
    for part in parts:
        item = part.strip()
        if mode == "syllable_sequence":
            tokens = [token for token in re.split(r"\s+", item) if token]
        else:
            tokens = [item]
        for token in tokens:
            direct = canonical_bopomofo(token)
            if direct:
                out.append(direct)
                continue
            if mode == "annotated_single":
                match = re.fullmatch(r"([ㄅ-ㄩˊˇˋ˙‧・]+)(?:語|讀|限讀|限讀)", token)
                if match and canonical_bopomofo(match.group(1)):
                    out.append(canonical_bopomofo(match.group(1)))
                    continue
            return []
    return out


def _validate_bopomofo_rows(rows: Sequence[Mapping[str, Any]], specs: Sequence[Mapping[str, Any]]) -> list[str]:
    errors: list[str] = []
    for spec in specs:
        column = str(spec.get("column") or "")
        mode = str(spec.get("mode") or "single")
        allow_empty = bool(spec.get("allow_empty", True))
        exempt_values = {normalize("NFKC", str(v or "")).strip() for v in spec.get("exempt_values", [])}
        for index, row in enumerate(rows, 2):
            value = row.get(column)
            text = normalize("NFKC", str(value or "")).strip()
            if not text and allow_empty:
                continue
            if text in exempt_values:
                continue
            if not _split_bopomofo_cell(value, mode):
                errors.append(f"row {index} column {column}: {text[:80]}")
                if len(errors) >= 50:
                    return errors
    return errors


def validate_csv_schema(path: Path, spec: Mapping[str, Any]) -> dict[str, Any]:
    result: dict[str, Any] = {"path": str(path), "kind": "csv", "ok": False, "errors": []}
    try:
        headers, rows = _read_csv(path)
    except Exception as exc:
        result["errors"].append(f"CSV 無法解析：{exc}")
        return result
    required = [str(value) for value in spec.get("required_columns", [])]
    duplicate_headers = sorted(column for column, count in Counter(headers).items() if column and count > 1)
    if duplicate_headers:
        result["errors"].append(f"重複欄位名稱：{duplicate_headers}")
    if any(not column for column in headers):
        result["errors"].append("存在空白欄位名稱")
    malformed_rows = [index for index, row in enumerate(rows, 2) if None in row]
    if malformed_rows:
        result["errors"].append(f"CSV 欄位數超出 header：rows {malformed_rows[:20]}")
    missing = [column for column in required if column not in headers]
    if missing:
        result["errors"].append(f"缺少欄位：{missing}")
    minimum = int(spec.get("min_rows", 0) or 0)
    if len(rows) < minimum:
        result["errors"].append(f"資料列低於安全下限：{len(rows)} < {minimum}")
    expected_rows = spec.get("expected_rows")
    if expected_rows is not None and len(rows) != int(expected_rows):
        result["errors"].append(f"資料列數不符 manifest：{len(rows)} != {expected_rows}")
    duplicate_report = {"ok": True, "duplicate_key_count": 0, "conflicting_duplicate_key_count": 0}
    unique_key = [str(value) for value in spec.get("unique_key", [])]
    required_nonempty = [str(value) for value in spec.get("required_nonempty", [])]
    blank_required = _blank_required_cells(rows, required_nonempty)
    if blank_required:
        result["errors"].append(f"必要鍵值空白：{blank_required[:10]}")
    blank_keys = _blank_composite_keys(rows, unique_key)
    if blank_keys:
        result["errors"].append(f"唯一鍵全空白：{blank_keys[:10]}")
    if spec.get("unique_key"):
        duplicate_report = validate_duplicate_keys(rows, spec["unique_key"], spec.get("duplicate_value_columns", []))
        allow_identical = bool(spec.get("allow_identical_duplicates"))
        duplicate_report["allowed_identical_duplicates"] = bool(
            allow_identical
            and duplicate_report["duplicate_key_count"] > 0
            and duplicate_report["conflicting_duplicate_key_count"] == 0
        )
        if not duplicate_report["ok"] and not duplicate_report["allowed_identical_duplicates"]:
            result["errors"].append(
                f"unique key 重複：{duplicate_report['duplicate_key_count']}；衝突：{duplicate_report['conflicting_duplicate_key_count']}"
            )
    bopomofo_errors = _validate_bopomofo_rows(rows, spec.get("bopomofo_columns", []))
    if bopomofo_errors:
        result["errors"].append(f"注音欄無法 canonicalize：{bopomofo_errors[:10]}")
    result.update({
        "headers": headers,
        "loaded_row_count": len(rows),
        "validated_row_count": len(rows) if not missing and not bopomofo_errors else 0,
        "duplicate_keys": duplicate_report,
    })
    result["ok"] = not result["errors"] and result["validated_row_count"] == result["loaded_row_count"]
    return result


def validate_xlsx_schema(path: Path, spec: Mapping[str, Any]) -> dict[str, Any]:
    result: dict[str, Any] = {"path": str(path), "kind": "xlsx", "ok": False, "errors": []}
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        result["errors"].append(f"Excel 無法解析：{exc}")
        return result
    sheet_name = str(spec.get("sheet") or "")
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            result["errors"].append(f"缺少工作表：{sheet_name}")
            workbook.close()
            return result
        worksheet = workbook[sheet_name]
    else:
        worksheet = workbook.active
        sheet_name = worksheet.title
    iterator = worksheet.iter_rows(values_only=True)
    try:
        headers = [str(value or "").strip() for value in next(iterator)]
    except StopIteration:
        headers = []
    rows = [{headers[index]: values[index] for index in range(min(len(headers), len(values)))} for values in iterator]
    workbook.close()
    required = [str(value) for value in spec.get("required_columns", [])]
    duplicate_headers = sorted(column for column, count in Counter(headers).items() if column and count > 1)
    if duplicate_headers:
        result["errors"].append(f"重複欄位名稱：{duplicate_headers}")
    if any(not column for column in headers):
        result["errors"].append("存在空白欄位名稱")
    missing = [column for column in required if column not in headers]
    if missing:
        result["errors"].append(f"缺少欄位：{missing}")
    minimum = int(spec.get("min_rows", 0) or 0)
    if len(rows) < minimum:
        result["errors"].append(f"資料列低於安全下限：{len(rows)} < {minimum}")
    expected_rows = spec.get("expected_rows")
    if expected_rows is not None and len(rows) != int(expected_rows):
        result["errors"].append(f"資料列數不符 manifest：{len(rows)} != {expected_rows}")
    duplicate_report = {"ok": True, "duplicate_key_count": 0, "conflicting_duplicate_key_count": 0}
    unique_key = [str(value) for value in spec.get("unique_key", [])]
    required_nonempty = [str(value) for value in spec.get("required_nonempty", [])]
    blank_required = _blank_required_cells(rows, required_nonempty)
    if blank_required:
        result["errors"].append(f"必要鍵值空白：{blank_required[:10]}")
    blank_keys = _blank_composite_keys(rows, unique_key)
    if blank_keys:
        result["errors"].append(f"唯一鍵全空白：{blank_keys[:10]}")
    if spec.get("unique_key"):
        duplicate_report = validate_duplicate_keys(rows, spec["unique_key"], spec.get("duplicate_value_columns", []))
        allow_identical = bool(spec.get("allow_identical_duplicates"))
        duplicate_report["allowed_identical_duplicates"] = bool(
            allow_identical
            and duplicate_report["duplicate_key_count"] > 0
            and duplicate_report["conflicting_duplicate_key_count"] == 0
        )
        if not duplicate_report["ok"] and not duplicate_report["allowed_identical_duplicates"]:
            result["errors"].append(
                f"unique key 重複：{duplicate_report['duplicate_key_count']}；衝突：{duplicate_report['conflicting_duplicate_key_count']}"
            )
    bopomofo_errors = _validate_bopomofo_rows(rows, spec.get("bopomofo_columns", []))
    if bopomofo_errors:
        result["errors"].append(f"注音欄無法 canonicalize：{bopomofo_errors[:10]}")
    result.update({
        "sheet": sheet_name,
        "headers": headers,
        "loaded_row_count": len(rows),
        "validated_row_count": len(rows) if not missing and not bopomofo_errors else 0,
        "duplicate_keys": duplicate_report,
    })
    result["ok"] = not result["errors"] and result["validated_row_count"] == result["loaded_row_count"]
    return result


def validate_asset_manifest(root: Path, manifest_path: Path | None = None) -> dict[str, Any]:
    root = Path(root).resolve()
    manifest_path = Path(manifest_path or root / "runtime_asset_manifest.json")
    report: dict[str, Any] = {
        "manifest_schema_version": "",
        "manifest": str(manifest_path),
        "ok": False,
        "assets": [],
        "errors": [],
        "warnings": [],
    }
    try:
        manifest = _json_without_duplicate_keys(manifest_path.read_text(encoding="utf-8"))
    except Exception as exc:
        report["errors"].append(f"asset manifest 缺失或損壞：{exc}")
        return report
    report["manifest_schema_version"] = manifest.get("schema_version") or ""
    report["manifest_tool_version"] = manifest.get("tool_version") or ""
    from cross_version_compat import schema_compatible
    if not schema_compatible(manifest.get("schema_version"), ASSET_MANIFEST_SCHEMA_VERSION):
        report["errors"].append(
            f"asset manifest schema 不相容：{manifest.get('schema_version')} != {ASSET_MANIFEST_SCHEMA_VERSION}"
        )
    elif manifest.get("schema_version") != ASSET_MANIFEST_SCHEMA_VERSION:
        report["warnings"].append(
            f"asset manifest schema patch 版本不同但可相容：{manifest.get('schema_version')} -> {ASSET_MANIFEST_SCHEMA_VERSION}"
        )
    if manifest.get("tool_version") != TOOL_VERSION:
        report["warnings"].append(
            f"asset manifest tool version 不同但不阻擋：{manifest.get('tool_version')} -> {TOOL_VERSION}"
        )
    assets = manifest.get("assets")
    if not isinstance(assets, list) or not assets:
        report["errors"].append("asset manifest 沒有核心資產")
        return report
    names = [str(item.get("name") or "") for item in assets if isinstance(item, dict)]
    if any(not name for name in names):
        report["errors"].append("asset manifest 存在空白 name")
    duplicates = [name for name, count in Counter(names).items() if name and count > 1]
    if duplicates:
        report["errors"].append(f"asset manifest 重複 name：{duplicates}")
    manifest_chains = {
        str(item.get("name") or ""): str(item.get("chain") or "")
        for item in assets if isinstance(item, dict) and item.get("name")
    }
    missing_assets = sorted(set(REQUIRED_ASSET_CHAINS) - set(manifest_chains))
    unexpected_assets = sorted(set(manifest_chains) - set(REQUIRED_ASSET_CHAINS))
    wrong_chains = sorted(
        name for name, chain in manifest_chains.items()
        if name in REQUIRED_ASSET_CHAINS and chain != REQUIRED_ASSET_CHAINS[name]
    )
    report["required_asset_roster_ok"] = not missing_assets and not unexpected_assets and not wrong_chains
    chain_roster_ok = {}
    for chain in ("actual", "expected"):
        required_for_chain = {name for name, required_chain in REQUIRED_ASSET_CHAINS.items() if required_chain == chain}
        observed_for_chain = {name for name, observed_chain in manifest_chains.items() if observed_chain == chain}
        chain_roster_ok[chain] = observed_for_chain == required_for_chain
    report["required_asset_chain_roster_ok"] = chain_roster_ok
    if missing_assets:
        report["errors"].append(f"asset manifest 缺少固定核心資產：{missing_assets}")
    if unexpected_assets:
        report["errors"].append(f"asset manifest 含未核准資產：{unexpected_assets}")
    if wrong_chains:
        report["errors"].append(f"asset manifest actual/expected chain 錯置：{wrong_chains}")
    paths = [str(item.get("path") or "") for item in assets if isinstance(item, dict)]
    duplicate_paths = [path for path, count in Counter(paths).items() if path and count > 1]
    if duplicate_paths:
        report["errors"].append(f"asset manifest 重複 path：{duplicate_paths}")
    for spec in assets:
        if not isinstance(spec, dict):
            report["errors"].append("asset manifest 含非物件項目")
            continue
        item: dict[str, Any] = {
            "name": spec.get("name") or "",
            "chain": spec.get("chain") or "",
            "path": spec.get("path") or "",
            "ok": False,
            "errors": [],
        }
        if item["chain"] not in {"actual", "expected"}:
            item["errors"].append(f"chain 無效：{item['chain']}")
        if not re.fullmatch(r"[0-9a-fA-F]{64}", str(spec.get("sha256") or "")):
            item["errors"].append("manifest SHA-256 格式無效")
        relative = Path(str(spec.get("path") or ""))
        path = (root / relative).resolve()
        try:
            path.relative_to(root)
        except ValueError:
            item["errors"].append("資產路徑逸出工具根目錄")
            report["assets"].append(item)
            continue
        if not path.exists() or not path.is_file():
            item["errors"].append("核心資產不存在")
            report["assets"].append(item)
            continue
        actual_hash = sha256_file(path)
        item["sha256"] = actual_hash
        item["expected_sha256"] = str(spec.get("sha256") or "").lower()
        if actual_hash != item["expected_sha256"]:
            item["errors"].append("SHA-256 不符；runtime 不會自動接受新 hash")
        kind = str(spec.get("kind") or path.suffix.lstrip(".")).lower()
        if kind == "csv":
            schema_report = validate_csv_schema(path, spec)
        elif kind == "xlsx":
            schema_report = validate_xlsx_schema(path, spec)
        elif kind in {"json", "txt"}:
            schema_report = {"ok": True, "loaded_row_count": None, "validated_row_count": None, "errors": []}
            if kind == "json":
                try:
                    _json_without_duplicate_keys(path.read_text(encoding="utf-8"))
                except Exception as exc:
                    schema_report = {"ok": False, "errors": [f"JSON 無法解析：{exc}"]}
        else:
            schema_report = {"ok": False, "errors": [f"不支援的資產格式：{kind}"]}
        item["schema"] = schema_report
        item["errors"].extend(schema_report.get("errors") or [])
        item["ok"] = not item["errors"] and bool(schema_report.get("ok"))
        report["assets"].append(item)
    report["errors"].extend(
        f"{item['name']}：{'；'.join(item['errors'])}" for item in report["assets"] if not item.get("ok")
    )
    report["ok"] = not report["errors"] and all(item.get("ok") for item in report["assets"])
    return report


def _manifest_asset_hashes(validation_report: Mapping[str, Any], chain: str) -> dict[str, str]:
    return {
        str(item.get("name")): str(item.get("sha256"))
        for item in validation_report.get("assets", [])
        if item.get("chain") == chain and item.get("ok")
    }


def source_chain_ok(validation_report: Mapping[str, Any], chain: str) -> bool:
    from cross_version_compat import schema_compatible
    if not schema_compatible(validation_report.get("manifest_schema_version"), ASSET_MANIFEST_SCHEMA_VERSION):
        return False
    # Tool release is audit-only. A manifest produced by an older/newer
    # compatible release remains usable when the validated asset bytes match.
    chain_roster = validation_report.get("required_asset_chain_roster_ok")
    if isinstance(chain_roster, Mapping) and chain_roster.get(chain) is False:
        return False
    items = [item for item in validation_report.get("assets", []) if item.get("chain") == chain]
    return bool(items) and all(item.get("ok") for item in items)


def compute_actual_asset_fingerprint(
    root: Path,
    pdf_path: Path,
    validation_report: Mapping[str, Any],
    *,
    decoder_version: str,
    source_files: Iterable[str | Path],
    dynamic_dependencies: Mapping[str, Any] | None = None,
    dynamic_evidence_root: Path | None = None,
) -> dict[str, Any]:
    if not source_chain_ok(validation_report, "actual"):
        raise SourceValidationError("actual 核心資產尚未通過 fail-closed 驗證，不得建立 actual fingerprint")
    root = Path(root).resolve()
    pdf_path = Path(pdf_path).resolve()
    source_hashes: dict[str, str] = {}
    for source in source_files:
        relative = Path(source)
        path = relative if relative.is_absolute() else root / relative
        if not path.exists():
            raise SourceValidationError(f"actual decoder source 缺失：{path}")
        source_hashes[str(path.resolve().relative_to(root))] = sha256_file(path)
    actual_assets = _manifest_asset_hashes(validation_report, "actual")
    if not actual_assets:
        raise SourceValidationError("asset manifest 沒有已驗證的 actual 核心資產")
    # User-maintained actual evidence is intentionally dynamic.  It is validated
    # independently from the immutable runtime asset manifest, but its hashes
    # are part of the actual fingerprint so every verified correction safely
    # invalidates stale actual workbooks/caches without requiring a manifest edit.
    try:
        dynamic_root = Path(dynamic_evidence_root).resolve() if dynamic_evidence_root else root
        dynamic_hashes = dynamic_actual_hashes(
            dynamic_root,
            pdf_path=pdf_path,
            dependencies=dynamic_dependencies,
        )
    except Exception as exc:
        raise SourceValidationError(f"動態 actual 證據驗證失敗：{exc}") from exc
    components = {
        "fingerprint_schema_version": ACTUAL_FINGERPRINT_SCHEMA_VERSION,
        "pdf_sha256": sha256_file(pdf_path),
        "actual_ledger_schema_version": LEDGER_SCHEMA_VERSION,
        "decoder_version": decoder_version,
        "actual_decoder_source_hashes": dict(sorted(source_hashes.items())),
        "actual_asset_hashes": dict(sorted(actual_assets.items())),
        "dynamic_actual_evidence_hashes": dynamic_hashes,
        "dynamic_dependency_mode": "per_pdf_exact_dependency_v1" if dynamic_dependencies is not None else "global_fallback",
        "cff_batch_algorithm_hash": source_hashes.get("cff_zero_map_batch.py", ""),
        "consensus_hash": actual_assets.get("cff_crossfamily_cid_consensus", ""),
        "reuse_policy": "evidence_assets_v1",
    }
    # v5.6: tool version and implementation source hashes remain auditable but
    # no longer invalidate a cache by themselves.  Reuse is keyed to the
    # explicit fingerprint schema/policy contract, exact PDF bytes, approved
    # actual data assets and the project-scoped dynamic evidence relevant to
    # this PDF.
    reuse_components = fingerprint_contract_components(components, chain="actual")
    raw = json.dumps(reuse_components, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return {"fingerprint": hashlib.sha256(raw).hexdigest(), "components": components, "reuse_components": reuse_components}



def rewrite_actual_workbook_fingerprint(path: Path, fingerprint_payload: Mapping[str, Any]) -> None:
    """Update only the sealed actual fingerprint metadata after dependency scoping.

    Decoding a brand-new PDF cannot know its exact glyph dependency roster until
    the workbook exists.  The caller may therefore decode with a conservative
    global-fallback fingerprint, inspect the resulting exact glyph hashes, then
    replace the two metadata cells with the final per-PDF fingerprint.  No
    pronunciation/evidence row is changed here.
    """
    path = Path(path)
    wb = load_workbook(path)
    try:
        if "v5.2中繼資料" not in wb.sheetnames:
            raise ValueError(f"{path.name} 缺少 v5.2中繼資料")
        ws = wb["v5.2中繼資料"]
        wanted = {
            "actual_asset_fingerprint": str(fingerprint_payload.get("fingerprint") or ""),
            "actual_asset_fingerprint_components": json.dumps(
                fingerprint_payload.get("components") or {}, ensure_ascii=False, sort_keys=True
            ),
        }
        found = set()
        for row in ws.iter_rows(min_row=1, max_col=2):
            key = str(row[0].value or "")
            if key in wanted:
                row[1].value = wanted[key]
                found.add(key)
        missing = set(wanted) - found
        if missing:
            raise ValueError(f"{path.name} fingerprint metadata 缺少欄位：{sorted(missing)}")
        wb.save(path)
    finally:
        wb.close()

def compute_expected_asset_fingerprint(
    root: Path,
    validation_report: Mapping[str, Any],
    *,
    resolver_version: str,
    source_files: Iterable[str | Path],
) -> dict[str, Any]:
    """Fingerprint only the expected evidence chain and its resolver code."""

    if not source_chain_ok(validation_report, "expected"):
        raise SourceValidationError("expected 核心資產尚未通過 fail-closed 驗證，不得建立 expected fingerprint")
    root = Path(root).resolve()
    source_hashes: dict[str, str] = {}
    for source in source_files:
        relative = Path(source)
        path = relative if relative.is_absolute() else root / relative
        if not path.exists():
            raise SourceValidationError(f"expected resolver source 缺失：{path}")
        source_hashes[str(path.resolve().relative_to(root))] = sha256_file(path)
    expected_assets = _manifest_asset_hashes(validation_report, "expected")
    if not expected_assets:
        raise SourceValidationError("asset manifest 沒有已驗證的 expected 核心資產")
    components = {
        "fingerprint_schema_version": EXPECTED_FINGERPRINT_SCHEMA_VERSION,
        "resolver_version": resolver_version,
        "expected_resolver_source_hashes": dict(sorted(source_hashes.items())),
        "expected_asset_hashes": dict(sorted(expected_assets.items())),
        "reuse_policy": "evidence_assets_v1",
    }
    # v5.6: resolver/tool updates are audit information, not automatic cache
    # invalidators.  A candidate cache is rebuilt when its explicit
    # schema/policy contract or expected data assets change; otherwise prior
    # reviewed work remains reusable across releases.
    reuse_components = fingerprint_contract_components(components, chain="expected")
    raw = json.dumps(reuse_components, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return {"fingerprint": hashlib.sha256(raw).hexdigest(), "components": components, "reuse_components": reuse_components}


def write_pipeline_blocked(output_dir: Path, validation_report: Mapping[str, Any], reason: str = "SOURCE_INVALID") -> Path:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "pipeline_blocked.json"
    payload = {
        "status": "PIPELINE_BLOCKED",
        "reason": reason,
        "updated_at": datetime.now().isoformat(timespec="seconds"),
        "source_validation": dict(validation_report),
    }
    temporary = path.with_suffix(".json.tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)
    status_path = output_dir / "pipeline_status.json"
    status_temporary = status_path.with_suffix(".json.tmp")
    status_temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    status_temporary.replace(status_path)
    return path
