"""Read-only, explicitly selected legacy project admission. No decision making.

The plan is rebuilt on every apply. Only CSV learning/conflict readings are
donors; current PDF/workbook/manifest data prove identity and sample locations.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
from dataclasses import asdict
from pathlib import Path

import fitz
from openpyxl import load_workbook

import global_exact_glyph_library as library
from actual_review import (
    USER_GLYF_FILE, USER_GLYF_HEADERS, USER_CFF_FILE, USER_CFF_HEADERS,
    GLYPH_CONFLICT_FILE, GLYPH_CONFLICT_HEADERS, GLYPH_PROVENANCE_FILE,
    GLYPH_PROVENANCE_HEADERS, OCCURRENCE_OVERRIDE_FILE, OVERRIDE_HEADERS,
)
from cross_version_compat import schema_compatible
from occurrence_ledger import LEDGER_SCHEMA_VERSION, SESSION_SCHEMA_VERSION, WORKBOOK_SCHEMA_VERSION, REVIEW_ID_SCHEMA_VERSION
from cff_zhuyin_decoder import CFFZhuyinInspector
from export_pdf_text_diagnostics import TrueTypeGlyphInspector, build_font_lookup, resolve_font, normalize_basefont


class MigrationValidationError(ValueError):
    pass


def _require(condition, message):
    if not condition:
        raise MigrationValidationError(message)


def _sha(raw):
    return hashlib.sha256(raw).hexdigest()


def _json(raw):
    def pairs(items):
        result = {}
        for key, value in items:
            _require(key not in result, "duplicate JSON key: " + key)
            result[key] = value
        return result
    def constant(value):
        raise MigrationValidationError("nonfinite JSON: " + value)
    return json.loads(raw.decode("utf-8-sig"), object_pairs_hook=pairs, parse_constant=constant)


def _token(value, name):
    value = library._strict_text(value, name)
    _require(not any(c in value for c in ("/", "\\", ":")), name + " must be a token")
    return value


def _oid(value):
    _require(isinstance(value, str) and re.fullmatch(r"occ_[0-9a-f]{64}", value), "invalid occurrence_id")
    return value


def _ids(value):
    if value == "":
        return []
    _require(isinstance(value, str), "source examples must be text")
    return sorted({_oid(item) for item in value.split("|")})


def _integer(value, name, minimum=0):
    _require(not isinstance(value, bool), "invalid " + name)
    try:
        number = int(value)
        _require(float(value) == number and number >= minimum, "invalid " + name)
        return number
    except (TypeError, ValueError, OverflowError) as exc:
        raise MigrationValidationError("invalid " + name) from exc


def _number(value):
    _require(not isinstance(value, bool), "invalid bbox")
    result = float(value)
    _require(math.isfinite(result), "invalid bbox")
    return result


def _absolute(value, name):
    path = Path(value)
    _require(path.is_absolute(), name + " must be absolute")
    return path.resolve()


def _filename(value):
    value = _token(value, "PDF filename")
    _require(value not in {".", ".."} and Path(value).name == value, "invalid filename")
    return value


def _read(path, frozen):
    raw = path.read_bytes()
    frozen[path] = _sha(raw)
    return raw


def _csv(path, headers, frozen, *, optional=False):
    if optional and not path.exists():
        frozen[path] = None
        return [], None
    raw = _read(path, frozen)
    reader = csv.DictReader(io.StringIO(raw.decode("utf-8-sig"), newline=""), strict=True)
    _require(reader.fieldnames == headers, path.name + " exact headers required")
    rows = []
    for row in reader:
        _require(None not in row and all(value is not None for value in row.values()),
                 path.name + " malformed CSV row")
        rows.append(row)
    return rows, _sha(raw)


def _identity_key(kind, style, sha):
    library._sha256_text(sha, "glyph_sha256")
    library._strict_text(style, "style_group", allow_empty=True)
    _require(kind in {library.TTF_GLYF_SHA256, library.CFF_GLYPH_SHA256}, "invalid glyph kind")
    _require((kind == library.TTF_GLYF_SHA256 and style == "") or
             (kind == library.CFF_GLYPH_SHA256 and bool(style)), "invalid style contract")
    return kind, style, sha


def _learning(rows, kind):
    result = []
    for row in rows:
        key = _identity_key(kind, row.get("style_group", ""), row["glyph_sha256"])
        reading = library._canonical_reading(row["bopomofo"])
        level = row["verification_level"]
        _require(level in library.LEGACY_VERIFICATION_LEVELS, "unknown learning level")
        _require(re.fullmatch(r"[1-9][0-9]*", row["source_count"]), "invalid source_count")
        count, ids = int(row["source_count"]), _ids(row["source_examples"])
        _require(count == len(ids), "source_count/source_examples mismatch")
        _require(level != "USER_VERIFIED_SINGLE" or count == 1, "single requires one source")
        _require(level != "VERIFIED_EXACT_GLYPH" or count >= 2, "verified requires two sources")
        result.append(dict(key=key, reading=reading, level=level, ids=ids))
    return result


def _conflicts(rows):
    result = []
    for row in rows:
        key = _identity_key(row["kind"], row["style_group"], row["glyph_sha256"])
        _require(row["status"] == "GLYPH_TRUTH_CONFLICT", "invalid project conflict status")
        readings = [library._canonical_reading(value) for value in row["readings"].split("|")]
        _require(len(set(readings)) >= 2, "conflict requires different readings")
        ids = _ids(row["source_occurrence_ids"])
        for reading in sorted(set(readings)):
            result.append(dict(key=key, reading=reading, level=library.QUARANTINED_CONFLICT, ids=ids))
    return result


def _resolve_pdf(info, output, roots, frozen):
    name, digest = _filename(info["pdf_name"]), library._sha256_text(info["pdf_sha256"], "PDF SHA")
    candidates = [output / name, *(root / name for root in roots)]
    stored = Path(str(info.get("pdf") or ""))
    if stored.is_absolute():
        candidates.append(stored)
    for path in dict.fromkeys(candidates):
        if path.is_file():
            raw = path.read_bytes()
            if _sha(raw) == digest:
                frozen[path] = digest
                return raw
    raise MigrationValidationError("PDF missing or SHA mismatch: " + name)


def _artifact(info, field, output, directory, frozen):
    stored = Path(library._strict_text(info[field], field))
    # Fixed project-owned artifact directories, never an old project's files.
    path = output / directory / stored.name
    if stored.is_absolute() and stored.resolve().is_relative_to(output):
        path = stored.resolve()
    _require(path.resolve().is_relative_to(output), "artifact escapes selected project")
    raw = _read(path, frozen)
    _require(_sha(raw) == library._sha256_text(info[field + "_sha256"], field + " SHA"),
             "artifact SHA mismatch: " + field)
    return raw


def _workbook(raw, pdf_sha):
    with io.BytesIO(raw) as stream:
        book = load_workbook(stream, read_only=True, data_only=False)
        try:
            _require("實際注音" in book.sheetnames and "v5.2中繼資料" in book.sheetnames,
                     "actual workbook contract missing")
            meta = {}
            for row in book["v5.2中繼資料"].iter_rows(values_only=True):
                if row and row[0] not in (None, ""):
                    _require(row[0] not in meta, "duplicate workbook metadata")
                    meta[row[0]] = row[1] if len(row) > 1 else None
            _require(meta.get("pdf_sha256") == pdf_sha, "workbook PDF SHA mismatch")
            _require(schema_compatible(meta.get("workbook_schema_version"), WORKBOOK_SCHEMA_VERSION),
                     "workbook schema incompatible")
            result = []
            for sheet in ("實際注音", "結構偵測排除"):
                if sheet not in book.sheetnames:
                    continue
                iterator = book[sheet].iter_rows(values_only=True)
                headers = list(next(iterator, ()))
                _require(len(set(headers)) == len(headers), "duplicate workbook columns")
                required = {"occurrence_id", "review_id", "實體頁碼", "字元", "x0", "y0"}
                _require(required.issubset(headers), "missing actual identity columns")
                for values in iterator:
                    if not any(value is not None for value in values):
                        continue
                    row = dict(zip(headers, values))
                    row["_actual_sheet"] = sheet == "實際注音"
                    result.append(row)
            return result
        finally:
            book.close()


def _current_identity(document, row, key):
    """Reparse exact embedded bytes and prove the locator belongs to this PDF."""
    kind, style, digest = key
    page_index = _integer(row["實體頁碼"], "physical page", 1) - 1
    xref = _integer(row["font_xref"], "font xref", 1)
    root_gid = _integer(row["glyph_id_字形索引"], "root glyph")
    gid = _integer(row["注音元件ID"], "component glyph") if kind == library.TTF_GLYF_SHA256 else root_gid
    page = document[page_index]
    fonts = page.get_fonts(full=True)
    _require(xref in {int(font[0]) for font in fonts}, "font not on current page")
    lookup = build_font_lookup(fonts)
    matches = []
    for span in page.get_texttrace():
        resolved_xref, *_ = resolve_font(lookup, normalize_basefont(span.get("font")))
        if resolved_xref != xref:
            continue
        for codepoint, printed_gid, _origin, bbox in span.get("chars", ()):
            if (chr(codepoint) == row["字元"] and printed_gid == root_gid and
                    abs(float(bbox[0]) - _number(row["x0"])) <= 0.8 and
                    abs(float(bbox[1]) - _number(row["y0"])) <= 0.8):
                matches.append(printed_gid)
    _require(len(matches) == 1, "current occurrence locator unprovable")
    name, extension, _, data = document.extract_font(xref)
    _require(bool(data), "embedded font missing")
    if kind == library.TTF_GLYF_SHA256:
        _require(extension == "ttf", "not current TTF")
        inspector = TrueTypeGlyphInspector(data)
        _require(root_gid == gid or gid in {part["gid"] for part in inspector.components(root_gid)},
                 "component not attached to occurrence")
        identity = inspector.global_exact_identity(gid)
    else:
        _require(extension == "cff", "not current CFF")
        identity = CFFZhuyinInspector(data, normalize_basefont(name), {}).global_exact_identity(gid)
        _require(identity["style_group"] == style, "CFF current style mismatch")
    _require(identity.get("glyph_sha256") == digest, "current glyph SHA mismatch")
    return identity


def _occurrence_index(manifest, pdf_data, workbook_rows, required_keys):
    records = manifest.get("records")
    _require(isinstance(records, list), "sealed occurrence records required")
    by_id, reviews = {}, set()
    for record in records:
        _require(isinstance(record, dict), "invalid occurrence record")
        oid, rid = _oid(record.get("occurrence_id")), _token(record.get("review_id"), "review_id")
        _require(oid not in by_id and rid not in reviews, "duplicate occurrence/review identity")
        by_id[oid] = record
        reviews.add(rid)
    seen, index, proofs = set(), {}, {}
    for name, (pdf_sha, raw) in pdf_data.items():
        with fitz.open(stream=raw, filetype="pdf") as document:
            for row in workbook_rows[name]:
                oid = _oid(row["occurrence_id"])
                _require(oid not in seen and oid in by_id, "workbook occurrence roster mismatch")
                seen.add(oid)
                record = by_id[oid]
                _require(record.get("pdf_name") == name and record.get("pdf_sha256") == pdf_sha and
                         record.get("review_id") == row["review_id"], "manifest/workbook identity mismatch")
                _require(_integer(record.get("physical_page"), "page", 1) ==
                         _integer(row["實體頁碼"], "page", 1) and record.get("char") == row["字元"],
                         "manifest/workbook occurrence mismatch")
                for field in ("x0", "y0", "x1", "y1"):
                    if row.get(field) not in (None, "") or record.get(field) not in (None, ""):
                        _require(abs(_number(row.get(field)) - _number(record.get(field))) <= 0.001,
                                 "manifest/workbook bbox mismatch")
                if not row["_actual_sheet"]:
                    continue
                ttf, cff = row.get("TTF字形SHA256"), row.get("CFF整字字形SHA256")
                _require(not (ttf and cff), "ambiguous actual exact identity")
                if not ttf and not cff:
                    continue
                key = _identity_key(library.TTF_GLYF_SHA256 if ttf else library.CFF_GLYPH_SHA256,
                                    "" if ttf else row.get("CFF樣式群組"), ttf or cff)
                if key not in required_keys:
                    continue
                source = record.get("source_record")
                _require(isinstance(source, dict), "sealed actual source missing")
                for field in ("TTF字形SHA256", "CFF整字字形SHA256", "CFF樣式群組"):
                    _require((source.get(field) or "") == (row.get(field) or ""),
                             "sealed/current actual source mismatch: " + field)
                for field in ("font_xref", "glyph_id_字形索引") + (("注音元件ID",) if ttf else ()):
                    _require(_integer(source.get(field), field) == _integer(row.get(field), field),
                             "sealed/current actual source mismatch: " + field)
                proof = _current_identity(document, row, key)
                if key in proofs:
                    _require(proofs[key] == proof, "inconsistent exact identity proof")
                proofs[key] = proof
                index[oid] = {"key": key, "occurrence_id": oid, "review_id": row["review_id"],
                              "pdf_name": name, "pdf_sha256": pdf_sha,
                              "physical_page": _integer(row["實體頁碼"], "page", 1),
                              "character": row["字元"],
                              "bbox": [row.get(field) for field in ("x0", "y0", "x1", "y1")]}
    _require(seen == set(by_id), "sealed/current occurrence roster mismatch")
    _require(required_keys.issubset(proofs), "current source cannot prove legacy exact identity")
    return index, proofs, set(by_id)


def build_migration_plan(project_output, *, pdf_roots=()):
    output = _absolute(project_output, "project-output")
    roots = [_absolute(root, "pdf-root") for root in pdf_roots]
    frozen = {}
    manifest_raw = _read(output / "校對工作階段.json", frozen)
    manifest = _json(manifest_raw)
    _require(isinstance(manifest, dict), "manifest must be object")
    seal = library._sha256_text(manifest.get("manifest_integrity_sha256"), "manifest seal")
    _require(seal == library._canonical_sha256({key: value for key, value in manifest.items()
                                               if key != "manifest_integrity_sha256"}), "manifest seal mismatch")
    _require(schema_compatible(manifest.get("session_schema_version"), SESSION_SCHEMA_VERSION),
             "session schema incompatible")
    _require(schema_compatible(manifest.get("ledger_schema_version"), LEDGER_SCHEMA_VERSION),
             "ledger schema incompatible")
    _require(schema_compatible(manifest.get("workbook_schema_version"), WORKBOOK_SCHEMA_VERSION),
             "manifest workbook schema incompatible")
    _require(manifest.get("review_id_schema_version") == REVIEW_ID_SCHEMA_VERSION, "review schema incompatible")
    project_id = _token(manifest.get("session_id"), "session_id")
    evidence_root = output / "_專案證據" / "actual"
    # Pending Phase 4 transactions cannot masquerade as a stable legacy source.
    for name in ("global_exact_glyph_project_transaction.json",):
        _require(not (evidence_root / name).exists(), "pending project transaction blocks migration")
        frozen[evidence_root / name] = None
    batches = []
    for name, headers, kind in ((USER_GLYF_FILE, USER_GLYF_HEADERS, library.TTF_GLYF_SHA256),
                                (USER_CFF_FILE, USER_CFF_HEADERS, library.CFF_GLYPH_SHA256),
                                (GLYPH_CONFLICT_FILE, GLYPH_CONFLICT_HEADERS, None)):
        rows, digest = _csv(evidence_root / name, headers, frozen)
        rows = _learning(rows, kind) if kind else _conflicts(rows)
        seen = set()
        for row in rows:
            key = (row["key"], row["reading"], row["level"])
            _require(key not in seen, "duplicate canonical legacy row")
            seen.add(key)
        batches.append((name, digest, rows))
    # These files are audit-only; validate headers/reading syntax without ever
    # seeding missing files or running the historical override repair helper.
    for name, headers, reading_field in ((OCCURRENCE_OVERRIDE_FILE, OVERRIDE_HEADERS, "actual_reading"),
                                         (GLYPH_PROVENANCE_FILE, GLYPH_PROVENANCE_HEADERS, "bopomofo")):
        rows, _ = _csv(evidence_root / name, headers, frozen, optional=True)
        for row in rows:
            if row[reading_field]:
                library._canonical_reading(row[reading_field])
    required_keys = {row["key"] for _, _, rows in batches for row in rows}
    pdf_data, workbook_rows = {}, {}
    roster = manifest.get("pdfs")
    _require(isinstance(roster, list) and bool(roster), "PDF roster required")
    for info in roster:
        _require(isinstance(info, dict), "invalid PDF roster item")
        name = _filename(info.get("pdf_name"))
        _require(name not in pdf_data, "duplicate PDF name")
        raw = _resolve_pdf(info, output, roots, frozen)
        pdf_data[name] = (info["pdf_sha256"], raw)
        actual = _artifact(info, "actual_workbook", output, "01_實際注音", frozen)
        _artifact(info, "candidate_workbook", output, "02_候選報告", frozen)  # bytes/hash only
        workbook_rows[name] = _workbook(actual, info["pdf_sha256"])
    index, proofs, current_ids = _occurrence_index(manifest, pdf_data, workbook_rows, required_keys)
    conflict_readings = {(row["key"], row["reading"]) for row in batches[-1][2]}
    retained_samples = {}
    # Even a row omitted from delivery remains source evidence to validate.
    # Preserve its usable samples on the matching authoritative conflict import.
    for name, _, rows in batches:
        for row in rows:
            key = row["key"]
            _require(all(oid not in current_ids or (oid in index and index[oid]["key"] == key)
                         for oid in row["ids"]), "source example points to a different current identity")
            if row["level"] == library.QUARANTINED_CONFLICT and name != GLYPH_CONFLICT_FILE:
                _require((key, row["reading"]) in conflict_readings,
                         "quarantined learning retained reading missing from conflict authority")
                retained_samples.setdefault((key, row["reading"]), set()).update(row["ids"])
    intents, targets, skipped = [], [], []
    for name, digest, rows in batches:
        for row in rows:
            proof, key = proofs[row["key"]], row["key"]
            if proof["eligibility"] == library.NON_GLOBAL_ELIGIBLE:
                skipped.append({"source_file": name, "kind": key[0], "style_group": key[1],
                                "glyph_sha256": key[2], "reading": row["reading"],
                                "status": library.NON_GLOBAL_ELIGIBLE, "reason": proof.get("reason", "")})
                continue
            if row["level"] == library.QUARANTINED_CONFLICT and name != GLYPH_CONFLICT_FILE:
                continue
            sample_ids = set(row["ids"])
            if name == GLYPH_CONFLICT_FILE:
                sample_ids.update(retained_samples.get((key, row["reading"]), ()))
            mapped = [index[oid] for oid in sorted(sample_ids) if oid in index]
            status = library.MIGRATION_CONFLICT if name == GLYPH_CONFLICT_FILE else (
                library.MIGRATION_CANDIDATE if mapped else library.MIGRATION_INSUFFICIENT)
            exact = {"kind": key[0], "style_group": key[1], "glyph_sha256": key[2],
                     "identity_eligibility": proof["eligibility"]}
            item = library.make_migration_intent(
                exact, reading=row["reading"], legacy_project_id=project_id,
                legacy_project_sha256=_sha(manifest_raw), legacy_evidence_sha256=digest,
                old_verification_level=row["level"], status=status)
            intents.append(item)
            for occurrence in mapped:
                targets.append({key: value for key, value in occurrence.items() if key != "key"} | {
                    "import_id": item["payload"]["import_id"], "glyph_id": item["payload"]["glyph_id"],
                    "identity": exact, "source_project_id": project_id, "legacy_status": status})
    # Freeze/recheck all bytes after parsing. No path enters the import payload.
    for path, digest in frozen.items():
        _require((not path.exists()) if digest is None else (path.is_file() and _sha(path.read_bytes()) == digest),
                 "source changed during validation: " + path.name)
    return {"migration_import_contract_version": library.GLOBAL_MIGRATION_IMPORT_CONTRACT_VERSION,
            "source_project_id": project_id, "intents": library.canonical_migration_batch(intents),
            "reconfirmation_targets": sorted(targets, key=lambda item: (item["import_id"], item["occurrence_id"])),
            "skipped": sorted(skipped, key=library._canonical_json)}


def migrate_project(project_output, *, pdf_roots=(), global_library_root=None, apply=False):
    """Default-safe service. Applying always validates a fresh source snapshot."""
    plan = build_migration_plan(project_output, pdf_roots=pdf_roots)
    repository = library.GlobalExactGlyphRepository.resolved(global_library_root)
    _require(not repository.path.resolve().is_relative_to(_absolute(project_output, "project-output")),
             "Global library must be outside the read-only source project")
    if apply:
        receipts = library.deliver_global_migration(repository, plan["intents"])
        return {**plan, "mode": "APPLIED", "receipts": [asdict(receipt) for receipt in receipts]}
    snapshot = repository.load_snapshot()  # Strict, read-only; absence stays absent.
    return {**plan, "mode": "DRY_RUN", "global_store_status": snapshot.store_status}
