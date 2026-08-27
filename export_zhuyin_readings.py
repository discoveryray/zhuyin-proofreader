from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
import sys

import fitz
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from export_pdf_text_diagnostics import (
    PROGRAM as DIAG_PROGRAM,
    VERSION as DIAG_VERSION,
    ZH_FONT_RE,
    TrueTypeGlyphInspector,
    build_font_lookup,
    get_page_fonts,
    guess_main_and_zhuyin_component,
    resolve_font,
    normalize_basefont,
    safe_char,
    safe_round,
)

PROGRAM = "PDF 實際注音解碼器"
VERSION = "5.6.2"
DEFAULT_MAP = Path(__file__).with_name("zhuyin_component_map.csv")
DEFAULT_GROUPS = Path(__file__).with_name("font_compatibility_groups.csv")
DEFAULT_CFF_MAP = Path(__file__).with_name("cff_bopomofo_symbol_map.csv")
DEFAULT_CFF_CONSENSUS = Path(__file__).with_name("cff_crossfamily_cid_consensus.csv")
DEFAULT_XREF_OVERRIDES = Path(__file__).with_name("ttf_xref_component_overrides.csv")
DEFAULT_TRANSFORMS = Path(__file__).with_name("ttf_verified_component_transforms.csv")
DEFAULT_FINGERPRINTS = Path(__file__).with_name("ttf_verified_glyf_fingerprints.csv")
DEFAULT_OUTLINE_SIGNATURES = Path(__file__).with_name("ttf_verified_outline_signatures.csv")
DEFAULT_MAPPING_CORRECTIONS = Path(__file__).with_name("ttf_verified_mapping_corrections.csv")
DEFAULT_SYMBOL_TEMPLATES = Path(__file__).with_name("ttf_verified_symbol_templates.csv")
DEFAULT_ACTUAL_OVERRIDES = Path(__file__).with_name("manual_actual_occurrence_overrides.csv")
DEFAULT_STRUCTURAL_EXCLUSIONS = Path(__file__).with_name("structural_detection_exclusions.csv")
ACTUAL_DECODER_SOURCE_FILES = [
    "export_zhuyin_readings.py", "export_pdf_text_diagnostics.py", "cff_zhuyin_decoder.py",
    "cff_unseen_family_bootstrap.py", "ttf_zhuyin_shape_decoder.py", "ttf_symbol_recombination.py",
    "cff_zero_map_batch.py", "occurrence_ledger.py", "actual_review.py",
]

from cff_zhuyin_decoder import CFFZhuyinInspector, cff_style_group, load_cff_symbol_map
from actual_review import (
    USER_GLYF_FILE, USER_CFF_FILE, OCCURRENCE_OVERRIDE_FILE,
    load_user_verified_glyf, load_user_verified_cff, load_glyph_truth_quarantine, ensure_user_evidence_files,
    actual_workbook_dynamic_dependencies,
)
from cff_unseen_family_bootstrap import (
    cff_variant, cff_unseen_family_key, load_crossfamily_consensus,
    bootstrap_unseen_families, decode_with_bootstrap,
)
from ttf_zhuyin_shape_decoder import ShapeEvent, cross_validate_font, train_model, contour_group_signature
from ttf_symbol_recombination import ExactSymbolRecombinationModel, evidence_text as symbol_evidence_text
from occurrence_ledger import LEDGER_SCHEMA_VERSION, WORKBOOK_SCHEMA_VERSION, prepare_occurrence_rows
from runtime_source_validation import compute_actual_asset_fingerprint, rewrite_actual_workbook_fingerprint, source_chain_ok, validate_asset_manifest




def infer_printed_page_number(page, physical_page_no: int) -> str:
    """Return the printed textbook page number when it can be read safely.

    Prefer an explicit PDF page label when present.  Many textbook PDFs do not
    define page labels, so fall back to a numeric footer near the bottom outer
    edge of the page.  If no trustworthy printed number is found, return an
    empty string instead of inventing one.
    """
    try:
        label = (page.get_label() or "").strip()
    except Exception:
        label = ""
    if label and label != str(physical_page_no):
        return label

    try:
        words = page.get_text("words") or []
    except Exception:
        words = []
    width = float(page.rect.width or 0)
    height = float(page.rect.height or 0)
    candidates = []
    for w in words:
        if len(w) < 5:
            continue
        x0, y0, x1, y1, text = w[:5]
        text = str(text or "").strip()
        if not re.fullmatch(r"\d{1,3}", text):
            continue
        # Printed page numbers in the textbook sit at the bottom outer margin.
        if height and float(y0) < height * 0.88:
            continue
        outer = (float(x0) <= width * 0.22) or (float(x1) >= width * 0.78)
        if not outer:
            continue
        candidates.append((float(y0), -min(float(x0), max(0.0, width-float(x1))), text))
    if not candidates:
        return ""
    candidates.sort(reverse=True)
    return candidates[0][2]

def unique_output_path(pdf_path: Path) -> Path:
    base = pdf_path.with_name(f"{pdf_path.stem}_實際注音解碼.xlsx")
    if not base.exists():
        return base
    i = 2
    while True:
        p = pdf_path.with_name(f"{pdf_path.stem}_實際注音解碼_{i}.xlsx")
        if not p.exists():
            return p
        i += 1


def style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    fill = PatternFill("solid", fgColor="D9EAF7")
    for c in ws[1]:
        c.fill = fill
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        mx = 0
        for c in ws[letter][: min(ws.max_row, 600)]:
            mx = max(mx, len("" if c.value is None else str(c.value)))
        ws.column_dimensions[letter].width = min(max(mx + 2, 10), 48)


def load_map(path: Path):
    mapping = {}
    if not path.exists():
        return mapping
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            font = (row.get("font") or "").strip()
            try:
                gid = int(row.get("zhuyin_component_id") or "")
            except Exception:
                continue
            bop = (row.get("bopomofo") or "").strip()
            if not font or not bop:
                continue
            mapping[(font, gid)] = {
                "bopomofo": bop,
                "verification": (row.get("verification") or "").strip(),
                "notes": (row.get("notes") or "").strip(),
                "source_font": font,
            }
    return mapping


def load_xref_overrides(path: Path, pdf_path: Path):
    """Load PDF-scoped TTF overrides, optionally restricted to one printed page.

    Global entries are keyed by ``(None, xref, font, component_id)``.  A row
    with ``page`` set is keyed by that printed page and is looked up first.
    Page scope is intentionally supported for human-confirmed corrections so a
    local truth cannot contaminate other occurrences that reuse the same
    embedded font xref/component identifier.
    """
    mapping = {}
    if not path.exists():
        return mapping
    stem = pdf_path.stem
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            scope = (row.get("pdf_contains") or "").strip()
            if scope and scope not in stem:
                continue
            try:
                xref = int(row.get("font_xref") or "")
                gid = int(row.get("zhuyin_component_id") or "")
            except Exception:
                continue
            page_scope = str(row.get("page") or "").strip() or None
            font = (row.get("font") or "").strip()
            bop = (row.get("bopomofo") or "").strip()
            if not font or not bop:
                continue
            mapping[(page_scope, xref, font, gid)] = {
                "bopomofo": bop,
                "verification": (row.get("verification") or "").strip(),
                "notes": (row.get("notes") or "").strip(),
                "source_font": f"{font}@xref{xref}" + (f"@p{page_scope}" if page_scope else ""),
                "page_scope": page_scope or "",
            }
    return mapping


def load_actual_occurrence_overrides(path: Path | None, pdf_path: Path):
    """Load user-verified actual-pronunciation corrections at one exact occurrence.

    These rows are deliberately occurrence-scoped and must never be promoted to
    a global font/component mapping.  They are applied only after all automatic
    glyph decoders have run, preserving the original automatic value for audit.
    """
    out=[]
    if not path or not path.exists():
        return out
    stem=pdf_path.stem
    with path.open("r",encoding="utf-8-sig",newline="") as f:
        for row in csv.DictReader(f):
            scope=(row.get("pdf_contains") or "").strip()
            if scope and scope not in stem:
                continue
            excludes=(row.get("pdf_excludes") or "").strip()
            if excludes:
                parts=[x.strip() for x in re.split(r"[|;,]", excludes) if x.strip()]
                if any(x in stem for x in parts):
                    continue
            try:
                row["x0"]=float(row.get("x0") or 0)
                row["y0"]=float(row.get("y0") or 0)
            except Exception:
                continue
            out.append(row)
    return out


def _page_token(value) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    try:
        number = float(text)
        if number.is_integer():
            return str(int(number))
    except Exception:
        pass
    return text


def match_actual_occurrence_override(row_values: dict, overrides: list[dict], tol: float = 0.8):
    """Match one visual actual correction to the exact printed occurrence.

    v5.5.0 accepts either printed-page or physical-page numbering in legacy/user
    override rows.  Stable-key drift after a decoder upgrade is tolerated only
    when page+character+coordinates match at a tighter threshold; this keeps the
    correction occurrence-scoped while avoiding a stale page/key from silently
    preventing a verified visual correction from taking effect.
    """
    printed_page = _page_token(row_values.get("課本頁"))
    physical_page = _page_token(row_values.get("實體頁碼"))
    target_char = str(row_values.get("字元") or "").strip()
    current_key = str(row_values.get("穩定注音鍵") or "").strip()
    for ov in overrides:
        ov_page = _page_token(ov.get("page"))
        if ov_page and ov_page not in {printed_page, physical_page}:
            continue
        if target_char != str(ov.get("target_char") or "").strip():
            continue
        try:
            dx = abs(float(row_values.get("x0") or 0)-float(ov.get("x0") or 0))
            dy = abs(float(row_values.get("y0") or 0)-float(ov.get("y0") or 0))
        except Exception:
            continue
        if dx > tol or dy > tol:
            continue
        sk=(ov.get("stable_key") or "").strip()
        if sk and current_key != sk:
            # Decoder versions can legitimately rebuild a stable key while the
            # physical glyph remains at the same coordinate.  Only accept that
            # drift when the coordinate match is essentially exact.
            if dx > 0.15 or dy > 0.15:
                continue
        if (ov.get("actual_reading") or "").strip():
            return ov
    return None


def load_structural_exclusions(path: Path | None, pdf_path: Path):
    """Load source-scoped false-positive structural detections.

    Exclusions require a direct page/glyph audit showing that the shifted
    component is part of the Han glyph and that no visible Bopomofo exists.
    """
    out=[]
    if not path or not path.exists():
        return out
    stem=pdf_path.stem
    with path.open("r",encoding="utf-8-sig",newline="") as f:
        for row in csv.DictReader(f):
            scope=(row.get("pdf_contains") or "").strip()
            if scope and scope not in stem:
                continue
            try:
                for k in ("physical_page","font_xref","glyph_id","zhuyin_component_id"):
                    row[k]=int(row.get(k) or 0)
                row["x0"]=float(row.get("x0") or 0); row["y0"]=float(row.get("y0") or 0)
            except Exception:
                continue
            out.append(row)
    return out


def match_structural_exclusion(page_no, printed_page, char, font, xref, glyph_id, zh_id, x0, y0, exclusions, tol=0.8):
    for ex in exclusions:
        if int(page_no) != int(ex.get("physical_page") or -1): continue
        if str(printed_page or "").strip() != str(ex.get("page") or "").strip(): continue
        if str(char or "").strip() != str(ex.get("target_char") or "").strip(): continue
        if str(font or "").strip() != str(ex.get("font") or "").strip(): continue
        if int(xref or 0) != int(ex.get("font_xref") or -1): continue
        if int(glyph_id or 0) != int(ex.get("glyph_id") or -1): continue
        if int(zh_id or 0) != int(ex.get("zhuyin_component_id") or -1): continue
        if abs(float(x0 or 0)-float(ex.get("x0") or 0)) > tol: continue
        if abs(float(y0 or 0)-float(ex.get("y0") or 0)) > tol: continue
        return ex
    return None


def load_verified_transforms(path: Path):
    """Load blind-test verified cross-font component mappings.

    Each row was established from glyph-outline evidence only.  Runtime use
    is fail-closed: the referenced source font/component must still carry the
    same exact-map reading *and* the current target component raw ``glyf``
    SHA-256 must match the stored blind-test fingerprint.  This keeps the
    observation chain independent from Chinese character semantics and blocks
    same-name/same-GID font-subset collisions.
    """
    out = {}
    if not path.exists():
        return out
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if (row.get("confidence") or "").strip().lower() != "safe":
                continue
            target_font = (row.get("target_font") or "").strip()
            source_font = (row.get("source_font") or "").strip()
            try:
                target_gid = int(row.get("target_component_id") or "")
                source_gid = int(row.get("source_component_id") or "")
            except Exception:
                continue
            bop = (row.get("actual_zhuyin") or "").strip()
            if not target_font or not source_font or not bop:
                continue
            target_hashes = tuple(
                x.strip().lower()
                for x in (row.get("target_glyph_sha256") or "").replace(";", " ").split()
                if x.strip()
            )
            # v3.2 transform evidence is intentionally fail-closed: rows without
            # an embedded target-glyph fingerprint are audit-only and cannot decode.
            if not target_hashes:
                continue
            out[(target_font, target_gid)] = {
                "bopomofo": bop,
                "source_font": source_font,
                "source_gid": source_gid,
                "target_hashes": target_hashes,
                "offset": (row.get("component_id_offset") or "").strip(),
                "verification": (row.get("validation_method") or "").strip(),
                "distance": (row.get("visual_chamfer_distance") or "").strip(),
                "threshold": (row.get("distance_threshold") or "").strip(),
                "notes": f"blind-test transform; source={source_font}#{source_gid}; target_sha256={','.join(target_hashes)}; distance={row.get('visual_chamfer_distance','')}; threshold={row.get('distance_threshold','')}",
            }
    return out


def lookup_verified_transform(font: str, gid: int, transform_map, exact_map, target_glyph_sha256: str = ""):
    """Return a verified transform only when both provenance gates pass.

    Gate 1: the referenced source font/component still has the same manually
    verified Bopomofo in the ordinary exact map.
    Gate 2: the current PDF's target component has the exact ``glyf`` SHA-256
    fingerprint observed during blind-test outline validation.  This prevents a
    same-name/same-component-ID but different embedded font subset from silently
    inheriting a reading.  Failure of either gate returns unknown.
    """
    rec = transform_map.get((font, gid))
    if not rec:
        return None
    current_hash = str(target_glyph_sha256 or "").strip().lower()
    allowed_hashes = set(rec.get("target_hashes") or ())
    if not current_hash or current_hash not in allowed_hashes:
        return None
    source = exact_map.get((rec["source_font"], rec["source_gid"]))
    if not source or source.get("bopomofo") != rec.get("bopomofo"):
        return None
    merged = dict(rec)
    merged["source_font"] = f"{rec['source_font']}#{rec['source_gid']}"
    merged["target_glyph_sha256"] = current_hash
    return merged


def load_compatibility_groups(path: Path):
    groups = []
    if not path.exists():
        return groups
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            group_id = (row.get("group_id") or "").strip()
            pattern = (row.get("font_regex") or "").strip()
            if not group_id or not pattern:
                continue
            try:
                rx = re.compile(pattern, re.IGNORECASE)
            except re.error:
                continue
            groups.append({
                "group_id": group_id,
                "regex": rx,
                "notes": (row.get("notes") or "").strip(),
            })
    return groups


def font_group(font: str, groups):
    for g in groups:
        if g["regex"].search(font or ""):
            return g["group_id"]
    return ""



def load_verified_glyf_fingerprints(path: Path):
    """Load directly visually verified raw-glyf fingerprints.

    These records are full Bopomofo annotation glyphs whose isolated outlines
    were manually read from the glyph itself.  Runtime use remains exact-only:
    a target raw ``glyf`` SHA-256 must equal a stored fingerprint byte-for-byte.
    """
    out = {}
    if not path.exists():
        return out
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            sha = (row.get("glyph_sha256") or "").strip().lower()
            bop = (row.get("bopomofo") or "").strip()
            if not sha or not bop:
                continue
            try:
                gid = int(row.get("source_gid") or "")
            except Exception:
                gid = -1
            out[sha] = {
                "bopomofo": bop,
                "source_font": (row.get("source_font") or "").strip(),
                "source_gid": gid,
                "verification": (row.get("verification") or "").strip(),
                "notes": (row.get("notes") or "").strip(),
            }
    return out



def load_verified_outline_signatures(path: Path):
    """Load conflict-free exact normalized TTF outline signatures."""
    out = {}
    if not path.exists():
        return out
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            sig = (row.get("outline_signature") or "").strip().lower()
            bop = (row.get("bopomofo") or "").strip()
            if not sig or not bop:
                continue
            out[sig] = {
                "bopomofo": bop,
                "source_count": (row.get("source_count") or "").strip(),
                "source_examples": (row.get("source_examples") or "").strip(),
                "verification": (row.get("verification") or "").strip(),
                "notes": (row.get("notes") or "").strip(),
            }
    return out



def load_verified_mapping_corrections(path: Path):
    """Load SHA-bound corrections for previously mislabeled TTF glyph keys.

    A correction is never selected by Chinese text or lexical expectation.  It
    applies only when font, component ID *and* raw glyf SHA-256 all match the
    isolated-glyph evidence captured during the v3.8 self-audit.
    """
    out = {}
    if not path.exists():
        return out
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            font = (row.get("font") or "").strip()
            sha = (row.get("glyph_sha256") or "").strip().lower()
            try:
                gid = int(row.get("component_id") or "")
            except Exception:
                continue
            new_reading = (row.get("new_reading") or "").strip()
            if not font or not sha or not new_reading:
                continue
            out[(font, gid, sha)] = {
                "bopomofo": new_reading,
                "old_reading": (row.get("old_reading") or "").strip(),
                "verification": (row.get("verification") or "").strip(),
                "notes": (row.get("notes") or "").strip(),
                "evidence": (row.get("evidence") or "").strip(),
                "source_book": (row.get("source_book") or "").strip(),
                "source_font": f"SHA-CORRECTION:{font}#{gid}",
            }
    return out


def lookup_verified_mapping_correction(font: str, gid: int, glyph_sha256: str, corrections):
    sha = str(glyph_sha256 or "").strip().lower()
    if not sha:
        return None
    return corrections.get((str(font or ""), int(gid), sha))

def is_han_char(c: str) -> bool:
    return bool(c) and any(("\u3400" <= ch <= "\u9fff") or ("\uf900" <= ch <= "\ufaff") for ch in str(c))

def build_group_map(exact_map, groups):
    """Build safe group-level lookup.

    A group fallback is allowed only when every non-empty exact mapping already
    present for the same (group, component id) agrees on the Bopomofo value.
    Conflicting IDs are deliberately excluded instead of guessed.
    """
    candidates = defaultdict(list)
    for (font, gid), rec in exact_map.items():
        gid_group = font_group(font, groups)
        if gid_group:
            candidates[(gid_group, gid)].append(rec)

    group_map = {}
    conflicts = {}
    for key, recs in candidates.items():
        readings = {r["bopomofo"] for r in recs if r.get("bopomofo")}
        if len(readings) == 1:
            # Prefer a record with the strongest/most specific note, while keeping sources.
            chosen = max(recs, key=lambda r: len(r.get("verification", "")) + len(r.get("notes", "")))
            merged = dict(chosen)
            merged["source_fonts"] = sorted({r.get("source_font", "") for r in recs if r.get("source_font")})
            group_map[key] = merged
        elif len(readings) > 1:
            conflicts[key] = sorted(readings)
    return group_map, conflicts


def lookup_reading(font: str, gid: int, exact_map, group_map, groups):
    exact = exact_map.get((font, gid))
    if exact:
        return exact, "精確字型對照", font_group(font, groups), font

    gid_group = font_group(font, groups)
    if gid_group:
        rec = group_map.get((gid_group, gid))
        if rec:
            srcs = rec.get("source_fonts") or [rec.get("source_font", "")]
            src = " / ".join(x for x in srcs if x)
            return rec, "相容字型群組對照", gid_group, src

    return None, "待建立對照", gid_group, ""


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _attach_occurrence_metadata(ws, pdf_path: Path, pdf_sha256: str) -> list[dict]:
    headers = [str(cell.value or "") for cell in ws[1]]
    metadata_headers = [
        "occurrence_id", "review_id", "identity_confidence",
        "identity_row_fallback", "identity_collision_base", "source_row_number",
        "ledger_schema_version", "workbook_schema_version",
    ]
    for name in metadata_headers:
        if name not in headers:
            ws.cell(1, len(headers) + 1, name)
            headers.append(name)
    rows: list[dict] = []
    for excel_row in range(2, ws.max_row + 1):
        record = {headers[index]: ws.cell(excel_row, index + 1).value for index in range(len(headers))}
        record.update({
            "source_row_number": excel_row - 1,
            "pdf_sha256": pdf_sha256,
            "pdf": str(pdf_path),
            "pdf_name": pdf_path.name,
        })
        rows.append(record)
    if rows:
        prepare_occurrence_rows(pdf_sha256, rows)
    header_index = {name: index + 1 for index, name in enumerate(headers)}
    for excel_row, record in enumerate(rows, 2):
        values = {
            "occurrence_id": record.get("occurrence_id", ""),
            "review_id": record.get("review_id", ""),
            "identity_confidence": record.get("identity_confidence", ""),
            "identity_row_fallback": "Y" if record.get("identity_row_fallback") else "N",
            "identity_collision_base": record.get("identity_collision_base", ""),
            "source_row_number": record.get("source_row_number", ""),
            "ledger_schema_version": LEDGER_SCHEMA_VERSION,
            "workbook_schema_version": WORKBOOK_SCHEMA_VERSION,
        }
        for name, value in values.items():
            ws.cell(excel_row, header_index[name], value)
    return rows


def decode(pdf_path: Path, output_path: Path, map_path: Path, groups_path: Path, cff_map_path: Path | None = None, xref_overrides_path: Path | None = None, transforms_path: Path | None = None, fingerprints_path: Path | None = None, outline_signatures_path: Path | None = None, mapping_corrections_path: Path | None = None, symbol_templates_path: Path | None = None, actual_overrides_path: Path | None = None, structural_exclusions_path: Path | None = None, cff_consensus_path: Path | None = None, actual_asset_fingerprint: dict | None = None, dynamic_evidence_root: Path | None = None):
    """Decode actual printed Bopomofo from two independent PDF font structures.

    * TrueType composite fonts: existing component-ID mapping.
    * CFF Type0 fonts: read the Bopomofo shapes physically embedded at the right
      side of each whole Han glyph.  No Han-character semantics are used to fill
      an unknown CFF symbol.
    """
    exact_map = load_map(map_path)
    groups = load_compatibility_groups(groups_path)
    group_map, group_conflicts = build_group_map(exact_map, groups)
    cff_map_path = Path(cff_map_path) if cff_map_path else DEFAULT_CFF_MAP
    cff_symbol_map = load_cff_symbol_map(cff_map_path)
    cff_consensus_path = Path(cff_consensus_path) if cff_consensus_path else DEFAULT_CFF_CONSENSUS
    cff_variant_ref, cff_gid_ref = load_crossfamily_consensus(cff_consensus_path)
    xref_overrides_path = Path(xref_overrides_path) if xref_overrides_path else DEFAULT_XREF_OVERRIDES
    xref_overrides = load_xref_overrides(xref_overrides_path, pdf_path)
    transforms_path = Path(transforms_path) if transforms_path else DEFAULT_TRANSFORMS
    verified_transforms = load_verified_transforms(transforms_path)
    fingerprints_path = Path(fingerprints_path) if fingerprints_path else DEFAULT_FINGERPRINTS
    verified_fingerprints = load_verified_glyf_fingerprints(fingerprints_path)
    # User-confirmed exact glyph truths are a separate dynamic evidence chain.
    # They are accepted only after promotion to VERIFIED_EXACT_GLYPH and are
    # merged exact-only; conflicts with built-in verified truths fail closed.
    root = Path(__file__).resolve().parent
    # v5.5.0: user-maintained actual evidence belongs to the proofreading
    # project, not to the program installation directory.  Keeping this lane
    # project-scoped prevents upgrades/hotfix folder changes from silently
    # losing or cross-contaminating visual actual decisions.
    dynamic_root = Path(dynamic_evidence_root).resolve() if dynamic_evidence_root else root
    dynamic_root.mkdir(parents=True, exist_ok=True)
    ensure_user_evidence_files(dynamic_root)
    quarantine = load_glyph_truth_quarantine(dynamic_root)
    quarantined_ttf = set(quarantine.get("ttf") or set())
    quarantined_cff = set(quarantine.get("cff") or set())
    # A contradictory exact SHA is not a decoder failure.  It is an explicit
    # reason to stop global SHA reuse while occurrence-local visual overrides
    # remain valid.  Both built-in and user reusable donors are suppressed for
    # quarantined keys so the conflict cannot silently propagate.
    for sha in list(verified_fingerprints):
        if sha in quarantined_ttf:
            verified_fingerprints.pop(sha, None)
    user_verified_fingerprints = load_user_verified_glyf(dynamic_root / USER_GLYF_FILE)
    for sha, urec in user_verified_fingerprints.items():
        if sha in quarantined_ttf:
            continue
        existing = verified_fingerprints.get(sha)
        if existing and str(existing.get("bopomofo") or "") != str(urec.get("bopomofo") or ""):
            # Defensive fail-closed fallback for legacy evidence that has not yet
            # been migrated into the explicit conflict registry.  Do not choose
            # either reading as a reusable donor.
            verified_fingerprints.pop(sha, None)
            quarantined_ttf.add(sha)
            continue
        verified_fingerprints[sha] = urec
    user_verified_cff = {
        key: rec for key, rec in load_user_verified_cff(dynamic_root / USER_CFF_FILE).items()
        if key not in quarantined_cff
    }
    outline_signatures_path = Path(outline_signatures_path) if outline_signatures_path else DEFAULT_OUTLINE_SIGNATURES
    verified_outline_signatures = load_verified_outline_signatures(outline_signatures_path)
    mapping_corrections_path = Path(mapping_corrections_path) if mapping_corrections_path else DEFAULT_MAPPING_CORRECTIONS
    verified_mapping_corrections = load_verified_mapping_corrections(mapping_corrections_path)
    symbol_templates_path = Path(symbol_templates_path) if symbol_templates_path else DEFAULT_SYMBOL_TEMPLATES
    symbol_model = ExactSymbolRecombinationModel.load_csv(symbol_templates_path)
    actual_overrides_path = (
        Path(actual_overrides_path) if actual_overrides_path
        else (dynamic_root / OCCURRENCE_OVERRIDE_FILE if dynamic_evidence_root else DEFAULT_ACTUAL_OVERRIDES)
    )
    actual_overrides = load_actual_occurrence_overrides(actual_overrides_path, pdf_path)
    structural_exclusions_path = Path(structural_exclusions_path) if structural_exclusions_path else DEFAULT_STRUCTURAL_EXCLUSIONS
    structural_exclusions = load_structural_exclusions(structural_exclusions_path, pdf_path)

    doc = fitz.open(pdf_path)
    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "摘要"
    ws = wb.create_sheet("實際注音")
    ws_keys = wb.create_sheet("注音鍵彙整")
    ws_missing = wb.create_sheet("未解碼注音鍵")
    ws_groups = wb.create_sheet("字型群組")
    ws_shape = wb.create_sheet("TTF字形自驗稽核")
    ws_bridge = wb.create_sheet("TTF跨字型指紋橋接")
    ws_symbol = wb.create_sheet("TTF符號重組稽核")
    ws_struct = wb.create_sheet("結構注音偵測")
    ws_actual_override = wb.create_sheet("人工實際注音覆寫")
    ws_struct_excl = wb.create_sheet("結構偵測排除")
    ws_cff_bootstrap = wb.create_sheet("CFF零對照自舉稽核")
    ws_runtime = wb.create_sheet("v5.2中繼資料")

    ws.append([
        "實體頁碼", "課本頁", "頁面標籤", "字元", "Unicode", "實際注音", "解碼狀態",
        "解碼依據", "字型相容群組", "對照來源font", "穩定注音鍵", "群組注音鍵",
        "font", "font_xref", "glyph_id_字形索引", "主字元件ID", "注音元件ID",
        "x0", "y0", "x1", "y1", "size", "判定方式",
        "字形架構", "CFF樣式群組", "CFF符號簽名", "CFF符號序列", "CFF聲調", "CFF聲調證據",
        "TTF字形SHA256", "TTF字形自驗預測", "TTF字形自驗狀態", "TTF字形模型CV", "TTF字形證據",
        "TTF正規化輪廓簽名", "注音結構偵測來源",
        "自動解碼原值", "人工實際注音覆寫", "人工覆寫來源", "人工覆寫備註",
        "CFF聲調簽名", "CFF輕聲簽名", "CFF完整注音簽名", "CFF整字字形SHA256"
    ])
    ws_keys.append([
        "穩定注音鍵", "群組注音鍵", "font", "字型相容群組", "注音元件ID", "實際注音",
        "解碼依據", "對照來源font", "出現次數", "字元例", "頁碼", "驗證方式", "備註",
        "字形架構", "CFF符號簽名", "CFF符號序列", "CFF聲調證據"
    ])
    ws_missing.append([
        "穩定注音鍵", "群組注音鍵", "font", "字型相容群組", "注音元件ID",
        "出現次數", "字元例", "頁碼", "下一步", "字形架構", "CFF未知符號簽名"
    ])
    ws_shape.append([
        "font", "注音元件ID", "出現次數", "字元例", "頁碼", "原解碼", "字形預測", "處置",
        "OOF預測", "模型可用", "模型已知鍵", "OOF預測鍵", "OOF正確鍵", "OOF精確率", "OOF覆蓋率",
        "字形SHA256例", "字形證據", "安全說明"
    ])
    ws_bridge.append([
        "目標font", "目標注音元件ID", "出現次數", "字元例", "頁碼", "橋接注音",
        "目標glyf SHA256", "安全來源鍵", "來源font", "來源安全讀音", "處置", "安全說明"
    ])
    ws_symbol.append([
        "font", "注音元件ID", "出現次數", "字元例", "頁碼", "進入重組前讀音",
        "exact符號重組預測", "處置", "所有xref變體一致", "最小來源鍵支持", "最小來源字型支持",
        "字形證據", "安全說明"
    ])
    ws_struct.append(["font", "偵測架構", "舊版名稱規則命中", "出現次數", "已解碼", "未解碼", "說明"])
    ws_actual_override.append(["課本頁", "字元", "穩定注音鍵", "font", "font_xref", "x0", "y0", "自動解碼原值", "人工覆寫後實際注音", "來源", "備註"])
    ws_struct_excl.append(["實體頁碼", "課本頁", "字元", "font", "font_xref", "glyph_id", "候選注音元件ID", "x0", "y0", "來源", "排除理由"])
    ws_cff_bootstrap.append(["未見家族鍵", "CFF符號簽名", "暫存符號", "處置", "目標glyph_id支持數", "最高安全雙重來源家族支持", "證據筆數", "來源已知樣式", "目標font", "安全說明"])
    ws_runtime.append(["項目", "內容"])
    ws_groups.append(["群組ID", "font_regex", "說明"])
    for g in groups:
        ws_groups.append([g["group_id"], g["regex"].pattern, g["notes"]])
    ws_groups.append(["CFF:BIAOKAI_W5", r"^DFBiaoKai(?:ZhuIn|PoIn[123])-W5$", "CFF 整字右側注音輪廓"])
    ws_groups.append(["CFF:YUAN_W3W5", r"^DFYuan(?:ZhuIn|PoIn[123])-W[35]$", "CFF 整字右側注音輪廓"])
    ws_groups.append(["CFF:YUAN_W7", r"^DFYuan(?:ZhuIn|PoIn[123])-W7$", "CFF 整字右側注音輪廓"])
    ws_groups.append(["CFF:HEI_W5", r"^DFHei(?:ZhuIn|PoIn[123])-W5$", "CFF 整字右側注音輪廓"])
    ws_groups.append(["CFF:HEI_W7", r"^DFHei(?:ZhuIn|PoIn[123])-W7$", "CFF 整字右側注音輪廓"])
    ws_groups.append(["CFF:KAICHU_MD", r"^DFKaiChuIn-Md-(?:BPMF|PoIn[123])-BF$", "CFF 整字右側注音輪廓；v4.3 康軒 exact signature family"])

    tt_cache = {}
    cff_cache = {}
    tt_component_cache = {}
    tt_shape_cache = {}
    cff_decode_cache = {}
    key_count = Counter()
    key_chars = defaultdict(list)
    key_pages = defaultdict(set)
    key_meta = {}
    shape_rows = []
    unseen_cff_rows = []

    total = mapped = 0
    tt_total = tt_mapped = exact_mapped = group_mapped = transform_mapped = 0
    outline_mapped = 0
    cff_total = cff_mapped = 0
    structural_counts = Counter()
    structural_mapped = Counter()

    def tt_components_for(xref, gid):
        if not isinstance(xref, int):
            return []
        ck = (xref, int(gid))
        if ck in tt_component_cache:
            return tt_component_cache[ck]
        if xref not in tt_cache:
            try:
                _name, _ext, _type, data = doc.extract_font(xref)
                tt_cache[xref] = TrueTypeGlyphInspector(data) if str(_ext).lower() == "ttf" else None
            except Exception:
                tt_cache[xref] = None
        ins = tt_cache.get(xref)
        val = ins.components(gid) if ins else []
        tt_component_cache[ck] = val
        return val

    def tt_shape_for(xref, gid):
        if not isinstance(xref, int):
            return [], ""
        ck = (xref, int(gid))
        if ck in tt_shape_cache:
            return tt_shape_cache[ck]
        if xref not in tt_cache:
            try:
                _name, _ext, _type, data = doc.extract_font(xref)
                tt_cache[xref] = TrueTypeGlyphInspector(data) if str(_ext).lower() == "ttf" else None
            except Exception:
                tt_cache[xref] = None
        ins = tt_cache.get(xref)
        if ins is None:
            val = ([], "")
        else:
            try:
                # First pass fingerprints only.  Contour parsing is deferred
                # until after page traversal and deduplicated by (font, gid,
                # raw-glyf SHA256).  Large books may repeat the same annotation
                # component thousands of times across embedded font subsets;
                # parsing every occurrence would be unnecessarily expensive.
                val = ([], ins.glyph_sha256(gid))
            except Exception:
                val = ([], "")
        tt_shape_cache[ck] = val
        return val

    def tt_outline_signature_for(xref, gid):
        if not isinstance(xref, int):
            return ""
        ck = ("outline", xref, int(gid))
        if ck in tt_shape_cache:
            return tt_shape_cache[ck]
        if xref not in tt_cache:
            try:
                _name, _ext, _type, data = doc.extract_font(xref)
                tt_cache[xref] = TrueTypeGlyphInspector(data) if str(_ext).lower() == "ttf" else None
            except Exception:
                tt_cache[xref] = None
        ins = tt_cache.get(xref)
        sig = ""
        if ins is not None:
            try:
                contours = ins.simple_contours(gid)
                if contours:
                    sig = contour_group_signature(contours)
            except Exception:
                sig = ""
        tt_shape_cache[ck] = sig
        return sig

    def cff_inspector_for(xref, font):
        if not isinstance(xref, int):
            return None
        key = (xref, font)
        if key not in cff_cache:
            try:
                _name, ext, _type, data = doc.extract_font(xref)
                if str(ext).lower() != "cid":
                    cff_cache[key] = None
                else:
                    cff_cache[key] = CFFZhuyinInspector(data, font, cff_symbol_map, user_verified_cff)
            except Exception:
                cff_cache[key] = None
        return cff_cache.get(key)

    def remember_key(key, page_no, char, meta):
        key_count[key] += 1
        key_pages[key].add(page_no)
        if char and char not in key_chars[key] and len(key_chars[key]) < 20:
            key_chars[key].append(char)
        key_meta[key] = meta

    for page_no, page in enumerate(doc, start=1):
        printed_page = infer_printed_page_number(page, page_no)
        label = page.get_label() or str(page_no)
        raw_fonts = get_page_fonts(page)
        font_lookup = build_font_lookup(raw_fonts)
        ext_lookup = defaultdict(list)
        for rec in raw_fonts:
            vals = list(rec) + [None] * 7
            xref0, ext0, ftype0, basefont0 = vals[:4]
            ext_lookup[normalize_basefont(basefont0)].append((xref0, str(ext0 or "").lower(), ftype0))
        try:
            traces = page.get_texttrace() or []
        except Exception:
            traces = []

        # Some textbook headings contain two glyph layers at exactly the same
        # coordinates (for example a white underlay plus the visible coloured
        # text).  Older versions counted both and could create a false mismatch
        # when the hidden underlay used a different Bopomofo component.  Keep
        # only the topmost / visibly coloured trace for an identical
        # (character, bbox) position.
        visible_choice = {}
        def _trace_rank(sp):
            color = sp.get("color")
            nonwhite = 1
            if isinstance(color, (tuple, list)) and len(color) >= 3:
                try:
                    rgb = [float(x) for x in color[:3]]
                    nonwhite = 0 if min(rgb) >= 0.965 else 1
                except Exception:
                    pass
            opacity = float(sp.get("opacity") or 0.0)
            seqno = int(sp.get("seqno") or 0)
            return (nonwhite, opacity, seqno)

        for span_idx, sp in enumerate(traces):
            # v3.7 selects the visible layer across every font.  Edited pages
            # can place a visible nonstandard annotation font above a hidden
            # ZhuIn/PoIn underlay.
            font0 = str(sp.get("font") or "")
            rank = _trace_rank(sp)
            for ch0 in sp.get("chars", ()):
                if len(ch0) < 4:
                    continue
                u0, _g0, _o0, bb0 = ch0[:4]
                try:
                    key0 = (int(u0),) + tuple(round(float(v), 0) for v in bb0[:4])
                except Exception:
                    continue
                prev = visible_choice.get(key0)
                if prev is None or rank > prev[0]:
                    visible_choice[key0] = (rank, span_idx)

        for span_idx, span in enumerate(traces):
            # Font naming is no longer an admission gate.  TTF annotation
            # membership is established by the shifted composite component.
            font = str(span.get("font") or "")
            size = safe_round(span.get("size"))
            xref, _basefont, _resource, _encoding = resolve_font(font_lookup, font)
            font_meta = ext_lookup.get(font, [])
            if not isinstance(xref, int) and len(font_meta) == 1:
                xref = font_meta[0][0]
            ext = font_meta[0][1] if len(font_meta) == 1 else ""
            cff_group = cff_style_group(font)
            cff_ins = cff_inspector_for(xref, font) if ext == "cid" else None

            for ch in span.get("chars", ()):
                if len(ch) < 4:
                    continue
                ucs, glyph_id, origin, bbox = ch[:4]
                try:
                    pos_key = (int(ucs),) + tuple(round(float(v), 0) for v in bbox[:4])
                except Exception:
                    pos_key = None
                if pos_key is not None:
                    chosen = visible_choice.get(pos_key)
                    if chosen is not None and chosen[1] != span_idx:
                        continue
                c = safe_char(ucs)
                if not is_han_char(c):
                    continue

                # CFF whole-glyph path.  Punctuation/Latin/no-annotation glyphs are
                # deliberately excluded from the denominator.
                if cff_ins is not None:
                    dkey = (xref, int(glyph_id))
                    dec = cff_decode_cache.get(dkey)
                    if dec is None:
                        dec = cff_ins.decode(glyph_id)
                        cff_decode_cache[dkey] = dec
                    if not dec.get("detected"):
                        continue
                    cff_detection = "known-CFF-family" if cff_group else "unseen-CFF-structure"
                    structural_counts[(font, cff_detection)] += 1
                    total += 1
                    cff_total += 1
                    bop = dec.get("reading", "")
                    if bop:
                        mapped += 1
                        cff_mapped += 1
                        structural_mapped[(font, cff_detection)] += 1
                    status = "已解碼" if bop else "待建立對照"
                    stable_key = f"CFF:{font}#{int(glyph_id)}"
                    cff_group_key = cff_group or f"UNSEEN:{font}"
                    group_key = f"CFF:{cff_group_key}#{int(glyph_id)}"
                    decode_basis = "CFF整字注音輪廓" if bop else ("CFF符號待建立對照" if cff_group else "未知CFF家族：結構已偵測，符號尚未驗證")
                    sigs = ";".join(dec.get("signatures", []))
                    symbols = "".join(x or "?" for x in dec.get("symbols", []))
                    tone = dec.get("tone", "")
                    tone_evidence = dec.get("tone_evidence", "")
                    tone_signature = dec.get("tone_signature", "")
                    neutral_signature = dec.get("neutral_signature", "")
                    full_signature = dec.get("full_signature", "")
                    cff_glyph_sha256 = dec.get("glyph_sha256", "")
                    unknown = ";".join(dec.get("unknown_signatures", []))
                    notes = f"CFF符號簽名={sigs}; 符號序列={symbols}; {tone_evidence}"
                    ws.append([
                        page_no, printed_page, label, c, f"U+{int(ucs):04X}" if isinstance(ucs, int) else "",
                        bop, status, decode_basis, f"CFF:{cff_group_key}", font, stable_key, group_key,
                        font, xref, glyph_id, "", "",
                        safe_round(bbox[0]), safe_round(bbox[1]), safe_round(bbox[2]), safe_round(bbox[3]),
                        size, dec.get("method", ""),
                        "CFF整字注音", cff_group, sigs, symbols, tone, tone_evidence,
                        "", "", "", "", "", "", cff_detection,
                        "", "", "", "", tone_signature, neutral_signature, full_signature, cff_glyph_sha256,
                    ])
                    if not cff_group:
                        unseen_cff_rows.append({
                            "row": ws._current_row,
                            "font": font,
                            "family_key": cff_unseen_family_key(font),
                            "variant": cff_variant(font),
                            "glyph_id": int(glyph_id),
                            "signatures": tuple(dec.get("signatures", [])),
                            "tone": tone,
                            "tone_evidence": tone_evidence,
                            "stable_key": stable_key,
                            "group_key": group_key,
                            "char": c,
                            "page_no": page_no,
                        })
                    rec = {
                        "bopomofo": bop,
                        "verification": "CFF右側注音符號字形直接辨識；符號表逐一視覺核驗",
                        "notes": notes,
                    } if bop else None
                    remember_key(stable_key, page_no, c, {
                        "font": font, "zh_id": "", "rec": rec, "decode_basis": decode_basis,
                        "gid_group": f"CFF:{cff_group_key}", "source_font": font, "group_key": group_key,
                        "architecture": "CFF整字注音", "cff_signatures": sigs, "cff_symbols": symbols,
                        "tone_evidence": tone_evidence, "unknown": unknown,
                    })
                    continue

                # Existing TrueType composite component path.
                comps = tt_components_for(xref, glyph_id)
                main_id, zh_id, method = guess_main_and_zhuyin_component(font, comps)
                if zh_id is None:
                    continue
                zh_id = int(zh_id)
                struct_ex = match_structural_exclusion(
                    page_no, printed_page, c, font, xref, glyph_id, zh_id,
                    safe_round(bbox[0]), safe_round(bbox[1]), structural_exclusions
                )
                if struct_ex:
                    ws_struct_excl.append([
                        page_no, printed_page, c, font, xref, glyph_id, zh_id,
                        safe_round(bbox[0]), safe_round(bbox[1]),
                        struct_ex.get("source", ""), struct_ex.get("note", ""),
                    ])
                    continue
                total += 1
                tt_total += 1
                outline_sig = ""
                shape_contours, target_glyph_hash = [], ""
                if isinstance(xref, int):
                    _unused_contours, target_glyph_hash = tt_shape_for(xref, zh_id)
                xref_rec = None
                if isinstance(xref, int):
                    page_key = str(printed_page or "").strip() or None
                    xref_rec = xref_overrides.get((page_key, int(xref), font, zh_id))
                    if xref_rec is None:
                        xref_rec = xref_overrides.get((None, int(xref), font, zh_id))
                correction_rec = None if xref_rec else lookup_verified_mapping_correction(
                    font, zh_id, target_glyph_hash, verified_mapping_corrections
                )
                if xref_rec:
                    stable_key = f"{font}@xref{xref}#{zh_id}"
                    rec = xref_rec
                    decode_basis = "PDF字型子集精確覆寫"
                    gid_group = font_group(font, groups)
                    source_font = xref_rec.get("source_font", f"{font}@xref{xref}")
                elif correction_rec:
                    stable_key = f"{font}#{zh_id}"
                    rec = correction_rec
                    decode_basis = "SHA綁定字形真值修正"
                    gid_group = "GLYF-CORRECTION"
                    source_font = correction_rec.get("source_font", f"SHA-CORRECTION:{font}#{zh_id}")
                else:
                    stable_key = f"{font}#{zh_id}"
                    rec, decode_basis, gid_group, source_font = lookup_reading(
                        font, zh_id, exact_map, group_map, groups
                    )
                    if rec is None:
                        # Preserve v3.2 performance: the SHA-256 transform gate
                        # is evaluated only when a transform lookup is actually
                        # needed.  v3.3 shape modelling is deferred until after
                        # page traversal and uses one representative per key.
                        # target_glyph_hash was already computed once above so
                        # SHA-bound corrections and transforms share the same
                        # exact raw-glyf evidence.
                        transform_rec = lookup_verified_transform(
                            font, zh_id, verified_transforms, exact_map, target_glyph_hash
                        )
                        if transform_rec:
                            rec = transform_rec
                            decode_basis = "已驗證跨字型輪廓轉換（SHA-256字形閘門）"
                            gid_group = f"TRANSFORM:{font}"
                            source_font = transform_rec.get("source_font", "")
                    if rec is None:
                        outline_sig = tt_outline_signature_for(xref, zh_id)
                        outline_rec = verified_outline_signatures.get(outline_sig) if outline_sig else None
                        if outline_rec:
                            rec = {
                                "bopomofo": outline_rec.get("bopomofo", ""),
                                "verification": outline_rec.get("verification", ""),
                                "notes": outline_rec.get("notes", ""),
                            }
                            decode_basis = "TTF跨字型exact正規化輪廓簽名"
                            gid_group = "OUTLINE-SIGNATURE"
                            source_font = outline_rec.get("source_examples", "")
                group_key = f"{gid_group}#{zh_id}" if gid_group else ""
                bop = rec["bopomofo"] if rec else ""
                status = "已解碼" if bop else "待建立對照"
                if bop:
                    mapped += 1
                    tt_mapped += 1
                    if decode_basis in ("精確字型對照", "PDF字型子集精確覆寫", "SHA綁定字形真值修正"):
                        exact_mapped += 1
                    elif decode_basis == "相容字型群組對照":
                        group_mapped += 1
                    elif decode_basis.startswith("已驗證跨字型輪廓轉換"):
                        transform_mapped += 1
                    elif decode_basis.startswith("TTF跨字型exact正規化輪廓簽名"):
                        outline_mapped += 1
                if not outline_sig:
                    outline_sig = tt_outline_signature_for(xref, zh_id)
                structural_source = "font-name+shifted-component" if ZH_FONT_RE.search(font) else "shifted-component-only"
                structural_counts[(font, structural_source)] += 1
                if bop:
                    structural_mapped[(font, structural_source)] += 1
                ws.append([
                    page_no, printed_page, label, c, f"U+{int(ucs):04X}" if isinstance(ucs, int) else "",
                    bop, status, decode_basis, gid_group, source_font, stable_key, group_key,
                    font, xref, glyph_id, main_id, zh_id,
                    safe_round(bbox[0]), safe_round(bbox[1]), safe_round(bbox[2]), safe_round(bbox[3]),
                    size, method,
                    "TrueType複合元件", "", "", "", "", "",
                    target_glyph_hash, "", "", "", "", outline_sig, structural_source,
                    "", "", "", "", "", "", "",
                ])
                shape_rows.append({
                    # ``ws.max_row`` scans the worksheet cell map and becomes
                    # O(n^2) when called once per glyph.  ``_current_row`` is
                    # updated by ``append`` and is constant-time here.
                    "row": ws._current_row, "font": font, "xref": xref, "zh_id": zh_id,
                    "stable_key": stable_key, "glyph_sha256": target_glyph_hash,
                    "contours": shape_contours, "initial_bopomofo": bop,
                    "initial_basis": decode_basis, "char": c, "page_no": page_no,
                })
                remember_key(stable_key, page_no, c, {
                    "font": font, "zh_id": zh_id, "rec": rec, "decode_basis": decode_basis,
                    "gid_group": gid_group, "source_font": source_font, "group_key": group_key,
                    "architecture": "TrueType複合元件", "cff_signatures": "", "cff_symbols": "",
                    "tone_evidence": "", "unknown": "",
                })

    # ------------------------------------------------------------------
    # v4.6 單檔層：沿用嚴格未見 CFF 家族零對照自舉；跨檔批次由 cff_zero_map_batch.py 再收斂。
    #
    # 這一層只對 cff_style_group() 完全未知的家族生效。目標家族沒有任何
    # 持久 mapping，也不讀中文字、詞義或句境。已知家族只提供 CID body
    # 的雙重共識錨點；目標家族再以自身 exact CFF 符號簽名做暫存自舉。
    # 任一衝突、證據不足、簽名缺漏或聲調落在安全區外都直接棄權。
    # ------------------------------------------------------------------
    cff_bootstrap_mapped_rows = 0
    cff_bootstrap_family_maps = {}
    cff_bootstrap_stats = {}
    cff_bootstrap_audit_conflicts = 0
    if unseen_cff_rows and cff_variant_ref and cff_gid_ref:
        cff_bootstrap_family_maps, cff_bootstrap_audit, cff_bootstrap_stats = bootstrap_unseen_families(
            unseen_cff_rows, cff_variant_ref, cff_gid_ref, min_source_styles=2
        )
        for a in cff_bootstrap_audit:
            if str(a.get("status", "")).startswith("棄權：同一簽名"):
                cff_bootstrap_audit_conflicts += 1
            ws_cff_bootstrap.append([
                a.get("family_key", ""), a.get("signature", ""), a.get("symbol", ""),
                a.get("status", ""), a.get("target_gid_count", 0), a.get("max_source_support", 0),
                a.get("evidence_count", 0), a.get("source_styles", ""), a.get("fonts", ""),
                "只使用已知家族CID body雙重共識與目標家族exact signature；不寫回持久mapping；不使用中文字/句境/字典；衝突或證據不足即棄權。",
            ])

        key_bootstrap_results = defaultdict(list)
        for item in unseen_cff_rows:
            temp_map = cff_bootstrap_family_maps.get(item["family_key"], {})
            reading = decode_with_bootstrap(item, temp_map)
            key_bootstrap_results[item["stable_key"]].append(reading)
            if not reading:
                continue
            row_no = item["row"]
            symbols = "".join(temp_map[s] for s in item["signatures"])
            auto_group = f"AUTOBOOT:{item['family_key']}"
            ws.cell(row_no, 6).value = reading
            ws.cell(row_no, 7).value = "已解碼"
            ws.cell(row_no, 8).value = "CFF未見家族零對照自舉（暫存exact簽名）"
            ws.cell(row_no, 9).value = f"CFF:{auto_group}"
            ws.cell(row_no, 12).value = f"CFF:{auto_group}#{item['glyph_id']}"
            ws.cell(row_no, 23).value = "跨已知家族CID雙重共識→目標家族exact簽名暫存→exact解碼"
            ws.cell(row_no, 25).value = auto_group
            ws.cell(row_no, 27).value = symbols
            cff_bootstrap_mapped_rows += 1

        # 注音鍵彙整只有在同一穩定鍵的所有出現位置都得到同一讀音時才升格。
        # 只要有一個子集變體棄權或出現衝突，該鍵仍保留未解碼，避免把局部
        # 成功誤寫成全鍵真值。
        for stable_key, readings in key_bootstrap_results.items():
            if not readings or any(not x for x in readings) or len(set(readings)) != 1:
                continue
            if stable_key not in key_meta:
                continue
            reading = readings[0]
            m = key_meta[stable_key]
            m["rec"] = {
                "bopomofo": reading,
                "verification": "CFF未見家族零對照自舉：CID雙重共識＋目標exact簽名暫存",
                "notes": "未新增target mapping；暫存簽名只在本次PDF內有效；安全門檻為雙重reference共識且來源家族支持>=2，或同一target簽名由>=2個不同glyph_id互證。",
            }
            m["decode_basis"] = "CFF未見家族零對照自舉（暫存exact簽名）"
            m["gid_group"] = f"CFF:AUTOBOOT:{cff_unseen_family_key(m['font'])}"
            m["group_key"] = f"{m['gid_group']}#{stable_key.rsplit('#', 1)[-1]}"
            m["cff_symbols"] = ""
            m["unknown"] = ""

    # ------------------------------------------------------------------
    # v3.3 TrueType glyph self-audit + automatic Bopomofo decomposition.
    #
    # The training labels below come only from the already-established glyph
    # mapping layers.  Unknown glyphs are decoded from exact outline signatures
    # of reusable Bopomofo symbols and tone marks.  Chinese characters, sentence
    # context, and dictionary expectations are never passed to the shape model.
    # A font may decode unknowns only after deterministic key-level OOF
    # validation reaches the strict readiness threshold.
    # ------------------------------------------------------------------
    # Build one representative simple outline per (font, component ID) for
    # training / OOF.  This keeps full-book runtime near v3.2.  For *unknown*
    # keys, every distinct embedded xref is parsed on demand below before an
    # automatic reading is accepted, so subset collisions still fail closed.
    shape_group_rows = defaultdict(list)
    for item in shape_rows:
        shape_group_rows[(item["font"], item["zh_id"])].append(item)

    shape_representatives = {}
    for key, items in shape_group_rows.items():
        representative = None
        seen_xrefs = set()
        for item in items:
            xref0 = item.get("xref")
            if not isinstance(xref0, int) or xref0 in seen_xrefs:
                continue
            seen_xrefs.add(xref0)
            ins = tt_cache.get(xref0)
            if ins is None:
                continue
            try:
                contours = ins.simple_contours(item["zh_id"])
                glyph_hash = ins.glyph_sha256(item["zh_id"])
            except Exception:
                contours, glyph_hash = [], ""
            if contours:
                representative = {
                    "contours": contours, "glyph_sha256": glyph_hash, "xref": xref0
                }
                break
        shape_representatives[key] = representative or {
            "contours": [], "glyph_sha256": "", "xref": None
        }

    shape_events_by_font = defaultdict(dict)
    for (font, zh_id), items in shape_group_rows.items():
        readings = {x.get("initial_bopomofo", "") for x in items if x.get("initial_bopomofo", "")}
        rep = shape_representatives.get((font, zh_id)) or {}
        contours = rep.get("contours") or []
        if len(readings) != 1 or not contours:
            continue
        reading = next(iter(readings))
        shape_events_by_font[font][(font, zh_id)] = ShapeEvent(
            key=(font, zh_id), contours=contours, reading=reading
        )

    shape_validation = {}
    shape_models = {}
    for font, event_map in shape_events_by_font.items():
        events = list(event_map.values())
        validation = cross_validate_font(font, events)
        shape_validation[font] = validation
        if validation.ready:
            shape_models[font] = train_model(events)

    shape_auto_keys = set()
    shape_auto_occurrences = 0
    shape_conflict_keys = set()
    shape_conflict_occurrences = 0
    shape_decode_cache = {}
    # Keys decoded by the exact same-font shape model may later donate their
    # raw-glyf fingerprint to another font.  This is deliberately separate
    # from OOF-held-out static keys: the *font model* must be OOF-ready and
    # every embedded-xref variant of the source key must decode identically.
    # Runtime bridge matching remains exact SHA-256 only.
    shape_auto_donor_hashes = {}

    def _shape_cv_text(validation):
        if validation is None:
            return "無模型"
        if validation.predicted:
            return (
                f"OOF {validation.correct}/{validation.predicted}="
                f"{validation.precision:.2%}; coverage={validation.coverage:.2%}; "
                f"known_keys={validation.unique_keys}; ready={'Y' if validation.ready else 'N'}"
            )
        return f"known_keys={validation.unique_keys}; ready=N"

    def _shape_evidence(decoded):
        if not decoded or not decoded.get("reading"):
            return decoded.get("reason", "") if decoded else ""
        symbol_text = ",".join(
            f"{x.get('symbol')}[n={x.get('support')},p={x.get('purity', 0):.3f}]"
            for x in decoded.get("symbols", [])
        )
        tone_support = decoded.get("tone_support")
        tone_text = f"tone={decoded.get('tone')}"
        if tone_support is not None:
            tone_text += f"[n={tone_support},p={decoded.get('tone_purity', 0):.3f}]"
        return f"symbols={symbol_text}; {tone_text}"

    for (font, zh_id), items in sorted(shape_group_rows.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        validation = shape_validation.get(font)
        model = shape_models.get(font)
        cv_text = _shape_cv_text(validation)
        initial_readings = {x.get("initial_bopomofo", "") for x in items if x.get("initial_bopomofo", "")}
        initial_text = next(iter(initial_readings)) if len(initial_readings) == 1 else (" / ".join(sorted(initial_readings)) if initial_readings else "")
        oof_prediction = validation.oof_predictions.get((font, zh_id), "") if validation else ""

        decoded_variants = []
        variant_hashes = []
        xref_hash_map = {}
        variant_complete = False
        if model:
            if initial_readings:
                # Existing mapping audit uses the representative outline.  The
                # independent evidence is the key-level OOF prediction; the full
                # model prediction here is only supplementary.
                rep = shape_representatives.get((font, zh_id)) or {}
                contours = rep.get("contours") or []
                if contours:
                    decoded_variants = [model.decode(contours)]
                    if rep.get("glyph_sha256"):
                        variant_hashes.append(rep.get("glyph_sha256"))
                    variant_complete = True
            else:
                # Unknown-key auto-decode is stricter: parse every distinct
                # embedded xref carrying this font/component and require every
                # variant to decode to the same pronunciation.
                seen_xrefs = set()
                variant_complete = True
                for item in items:
                    xref0 = item.get("xref")
                    if not isinstance(xref0, int) or xref0 in seen_xrefs:
                        continue
                    seen_xrefs.add(xref0)
                    ins = tt_cache.get(xref0)
                    if ins is None:
                        variant_complete = False
                        break
                    try:
                        glyph_hash = ins.glyph_sha256(zh_id)
                        contours = ins.simple_contours(zh_id)
                    except Exception:
                        glyph_hash, contours = "", []
                    if not contours:
                        variant_complete = False
                        break
                    xref_hash_map[xref0] = glyph_hash
                    variant_hashes.append(glyph_hash)
                    cache_key = (font, zh_id, glyph_hash or f"xref{xref0}")
                    if cache_key not in shape_decode_cache:
                        shape_decode_cache[cache_key] = model.decode(contours)
                    decoded_variants.append(shape_decode_cache[cache_key])
                if not seen_xrefs:
                    variant_complete = False

        predicted_readings = {d.get("reading", "") for d in decoded_variants if d.get("reading", "")}
        all_variants_decoded = variant_complete and decoded_variants and all(d.get("reading", "") for d in decoded_variants)
        shape_prediction = next(iter(predicted_readings)) if all_variants_decoded and len(predicted_readings) == 1 else ""
        evidence = " | ".join(sorted({_shape_evidence(d) for d in decoded_variants if d}))

        action = ""
        self_state = ""
        final_prediction = oof_prediction or shape_prediction

        # Independent OOF disagreement means the static mapping can no longer be
        # trusted blindly.  Fail closed: remove the reading and surface a conflict
        # instead of choosing either side automatically.
        if initial_readings and validation and validation.ready and oof_prediction and len(initial_readings) == 1 and oof_prediction != initial_text:
            action = "既有對照與OOF字形自驗衝突；保守停用"
            self_state = "衝突-停用"
            shape_conflict_keys.add((font, zh_id))
            shape_conflict_occurrences += len(items)
            for item in items:
                row_no = item["row"]
                ws.cell(row_no, 6).value = ""
                ws.cell(row_no, 7).value = "字形自驗衝突"
                ws.cell(row_no, 8).value = "既有對照與TTF字形OOF自驗衝突（保守停用）"
                ws.cell(row_no, 31).value = oof_prediction
                ws.cell(row_no, 32).value = self_state
                ws.cell(row_no, 33).value = cv_text
                ws.cell(row_no, 34).value = evidence
                sk = item["stable_key"]
                if sk in key_meta:
                    key_meta[sk]["rec"] = None
                    key_meta[sk]["decode_basis"] = "既有對照與TTF字形OOF自驗衝突（保守停用）"
                    key_meta[sk]["source_font"] = ""

        elif not initial_readings and validation and validation.ready and shape_prediction:
            action = "以TTF字形分解新增解碼"
            self_state = "自動泛化-已解碼"
            shape_auto_keys.add((font, zh_id))
            shape_auto_occurrences += len(items)
            rec = {
                "bopomofo": shape_prediction,
                "verification": (
                    "TTF glyf 直接輪廓分解；同字型 Bopomofo 符號/聲調模板；"
                    f"{cv_text}"
                ),
                "notes": evidence,
            }
            safe_hashes = {h for h in variant_hashes if h}
            if safe_hashes:
                shape_auto_donor_hashes[(font, zh_id)] = {
                    "reading": shape_prediction,
                    "hashes": safe_hashes,
                    "cv": cv_text,
                }
            for item in items:
                row_no = item["row"]
                ws.cell(row_no, 6).value = shape_prediction
                ws.cell(row_no, 7).value = "已解碼"
                ws.cell(row_no, 8).value = "TTF字形分解自動泛化（OOF驗證模型）"
                ws.cell(row_no, 9).value = f"SHAPE:{font}"
                ws.cell(row_no, 10).value = "同字型已驗證Bopomofo輪廓模板"
                ws.cell(row_no, 12).value = f"SHAPE:{font}#{zh_id}"
                if isinstance(item.get("xref"), int) and xref_hash_map.get(item.get("xref")):
                    ws.cell(row_no, 30).value = xref_hash_map.get(item.get("xref"))
                ws.cell(row_no, 31).value = shape_prediction
                ws.cell(row_no, 32).value = self_state
                ws.cell(row_no, 33).value = cv_text
                ws.cell(row_no, 34).value = evidence
                sk = item["stable_key"]
                if sk in key_meta:
                    key_meta[sk]["rec"] = rec
                    key_meta[sk]["decode_basis"] = "TTF字形分解自動泛化（OOF驗證模型）"
                    key_meta[sk]["gid_group"] = f"SHAPE:{font}"
                    key_meta[sk]["source_font"] = "同字型已驗證Bopomofo輪廓模板"
                    key_meta[sk]["group_key"] = f"SHAPE:{font}#{zh_id}"

        else:
            if initial_readings:
                if validation and validation.ready and oof_prediction == initial_text:
                    action = "既有對照通過OOF字形自驗"
                    self_state = "OOF一致"
                elif validation and validation.ready and shape_prediction == initial_text:
                    action = "既有對照與全模型字形一致（非獨立OOF）"
                    self_state = "全模型一致"
                elif validation and validation.ready:
                    action = "既有對照保留；字形模型未能獨立判定"
                    self_state = "模型未覆蓋"
                else:
                    action = "既有對照保留；字形模型未達啟用門檻"
                    self_state = "模型未啟用"
            else:
                action = "仍未知；字形證據不足或模型未達啟用門檻"
                self_state = "仍未知"
            for item in items:
                row_no = item["row"]
                ws.cell(row_no, 31).value = final_prediction
                ws.cell(row_no, 32).value = self_state
                ws.cell(row_no, 33).value = cv_text
                ws.cell(row_no, 34).value = evidence

        chars = "、".join(dict.fromkeys(x.get("char", "") for x in items if x.get("char")))
        pages = ",".join(str(x) for x in sorted({x.get("page_no") for x in items if x.get("page_no")}))
        hashes = ",".join(dict.fromkeys(x for x in variant_hashes if x))
        if not hashes:
            rep_hash = (shape_representatives.get((font, zh_id)) or {}).get("glyph_sha256", "")
            hashes = rep_hash or ""
        ws_shape.append([
            font, zh_id, len(items), chars, pages, initial_text, shape_prediction, action,
            oof_prediction, "Y" if validation and validation.ready else "N",
            validation.unique_keys if validation else 0,
            validation.predicted if validation else 0,
            validation.correct if validation else 0,
            validation.precision if validation else 0,
            validation.coverage if validation else 0,
            hashes[:500], evidence[:1500],
            "未知解碼只採 exact outline signature；模型須通過 key-level OOF；所有字形變體須一致；衝突一律停用。",
        ])

    # ------------------------------------------------------------------
    # v3.5 cross-font exact raw-glyf fingerprint bridge.
    #
    # This layer never compares Chinese text, lexical context, or dictionary
    # expectations.  A source raw ``glyf`` SHA-256 is eligible only from one
    # of three conservative evidence classes:
    #   (1) a static key independently confirmed by key-level OOF;
    #   (2) a previously unknown key decoded by an OOF-ready same-font exact
    #       shape model, with every embedded-xref source variant agreeing; or
    #   (3) an isolated-glyph fingerprint directly visually verified and stored
    #       in the explicit fingerprint table.
    # The target remains fail-closed: every embedded-xref variant for the target
    # component must resolve uniquely to one pronunciation.  Runtime matching
    # itself is exact SHA-256 equality; no nearest-neighbour or fuzzy threshold
    # is used.
    # ------------------------------------------------------------------
    fp_sources = defaultdict(lambda: defaultdict(set))
    source_key_oof = {}
    for (src_font, src_gid), items in shape_group_rows.items():
        validation = shape_validation.get(src_font)
        if not validation or not validation.ready:
            continue
        src_readings = {x.get("initial_bopomofo", "") for x in items if x.get("initial_bopomofo", "")}
        if len(src_readings) != 1:
            continue
        src_reading = next(iter(src_readings))
        oof = validation.oof_predictions.get((src_font, src_gid), "")
        if not oof or oof != src_reading:
            continue
        source_key_oof[(src_font, src_gid)] = oof
        seen_src_xrefs = set()
        for item in items:
            xref0 = item.get("xref")
            if not isinstance(xref0, int) or xref0 in seen_src_xrefs:
                continue
            seen_src_xrefs.add(xref0)
            ins = tt_cache.get(xref0)
            if ins is None:
                continue
            try:
                sha = ins.glyph_sha256(src_gid)
            except Exception:
                sha = ""
            if sha and sha not in quarantined_ttf:
                fp_sources[sha][src_reading].add((src_font, src_gid))

    # Same-font shape-auto keys are a second safe source class.  They are not
    # themselves held out during OOF; instead the font model has passed strict
    # key-level OOF and the source key was accepted only after all of its xref
    # variants exact-decomposed to one reading.  Their raw hashes can therefore
    # act as exact bridge donors without introducing fuzzy matching.
    for (src_font, src_gid), info in shape_auto_donor_hashes.items():
        src_reading = info.get("reading", "")
        if not src_reading:
            continue
        for sha in info.get("hashes", set()):
            if sha and sha not in quarantined_ttf:
                fp_sources[sha][src_reading].add((f"SHAPE-AUTO:{src_font}", src_gid))

    # Directly verified full-glyph fingerprints are a third, independent safe
    # source for signatures that the OOF decomposition model could not cover.
    for sha, rec in verified_fingerprints.items():
        if sha in quarantined_ttf:
            continue
        sf = rec.get("source_font", "MANUAL-GLYF") or "MANUAL-GLYF"
        sg = rec.get("source_gid", -1)
        fp_sources[sha][rec["bopomofo"]].add((sf, sg))

    fp_conflicts = {sha: by_reading for sha, by_reading in fp_sources.items() if len(by_reading) > 1}
    fp_safe = {sha: next(iter(by_reading)) for sha, by_reading in fp_sources.items() if len(by_reading) == 1 and sha not in quarantined_ttf}
    bridge_keys = set()
    bridge_occurrences = 0

    for (font, zh_id), items in sorted(shape_group_rows.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        # Only bridge keys that are still unresolved after v3.3 self-audit and
        # same-font shape auto-decomposition.
        current_readings = {str(ws.cell(item["row"], 6).value or "") for item in items}
        current_readings.discard("")
        if current_readings:
            continue

        target_hashes = []
        target_sources = set()
        target_readings = set()
        complete = True
        seen_xrefs = set()
        for item in items:
            xref0 = item.get("xref")
            if not isinstance(xref0, int) or xref0 in seen_xrefs:
                continue
            seen_xrefs.add(xref0)
            ins = tt_cache.get(xref0)
            if ins is None:
                complete = False
                break
            try:
                sha = ins.glyph_sha256(zh_id)
            except Exception:
                sha = ""
            if not sha or sha in quarantined_ttf or sha not in fp_safe:
                complete = False
                break
            by_reading = fp_sources.get(sha, {})
            if len(by_reading) != 1:
                complete = False
                break
            reading = next(iter(by_reading))
            target_hashes.append(sha)
            target_readings.add(reading)
            target_sources.update(by_reading[reading])
        if not seen_xrefs:
            complete = False
        if not complete or len(target_readings) != 1:
            continue

        bridge_reading = next(iter(target_readings))
        bridge_keys.add((font, zh_id))
        bridge_occurrences += len(items)
        source_text = "; ".join(f"{sf}#{sg}" for sf, sg in sorted(target_sources))
        hash_text = ",".join(dict.fromkeys(target_hashes))
        rec = {
            "bopomofo": bridge_reading,
            "verification": "TTF raw glyf SHA-256 exact match to a conservative safe glyph source (OOF static / shape-auto exact / isolated manual)",
            "notes": f"source={source_text}; sha256={hash_text}",
        }
        for item in items:
            row_no = item["row"]
            ws.cell(row_no, 6).value = bridge_reading
            ws.cell(row_no, 7).value = "已解碼"
            ws.cell(row_no, 8).value = "TTF跨字型glyf指紋橋接（安全字形來源）"
            ws.cell(row_no, 9).value = "GLYF-SHA256"
            ws.cell(row_no, 10).value = source_text
            ws.cell(row_no, 12).value = f"GLYF-SHA256#{zh_id}"
            ws.cell(row_no, 31).value = bridge_reading
            ws.cell(row_no, 32).value = "跨字型指紋橋接"
            ws.cell(row_no, 33).value = "exact SHA-256 + conservative safe glyph source"
            ws.cell(row_no, 34).value = f"sha256={hash_text}; source={source_text}"
            sk = item["stable_key"]
            if sk in key_meta:
                key_meta[sk]["rec"] = rec
                key_meta[sk]["decode_basis"] = "TTF跨字型glyf指紋橋接（安全字形來源）"
                key_meta[sk]["gid_group"] = "GLYF-SHA256"
                key_meta[sk]["source_font"] = source_text
                key_meta[sk]["group_key"] = f"GLYF-SHA256#{zh_id}"

        chars = "、".join(dict.fromkeys(x.get("char", "") for x in items if x.get("char")))
        pages = ",".join(str(x) for x in sorted({x.get("page_no") for x in items if x.get("page_no")}))
        ws_bridge.append([
            font, zh_id, len(items), chars, pages, bridge_reading, hash_text[:1000],
            source_text, ",".join(sorted({sf for sf, _ in target_sources})), bridge_reading,
            "已橋接",
            "目標 raw glyf SHA-256 與安全來源完全一致；來源限 OOF 靜態鍵、OOF-ready shape-auto exact 鍵或 isolated 人工驗證指紋；目標所有 xref 變體須一致；禁止中文字/句境/字典反填。",
        ])

    # ------------------------------------------------------------------
    # v3.9 exact Bopomofo-symbol recombination, including nested composite annotations.
    #
    # The template table contains two separately admitted evidence classes:
    # (1) conflict-free automatic templates supported by >=2 verified full-glyph
    # keys, and (2) exact isolated Bopomofo / tone shapes that were directly
    # visually verified without Han text or lexical context. Nested Bpmf*
    # annotations are decomposed into child glyphs and use the same exact table.
    # Runtime never receives Han text or lexical expectations; all xref variants
    # for one target key must exact-recombine to the same reading.
    # ------------------------------------------------------------------
    symbol_recomb_keys = set()
    symbol_recomb_occurrences = 0
    symbol_audit_conflict_keys = set()
    symbol_audit_conflict_occurrences = 0
    symbol_decode_cache = {}

    def _symbol_decode_for(xref0, gid0):
        if not isinstance(xref0, int):
            return {"reading": "", "reason": "no_xref"}
        ins = tt_cache.get(xref0)
        if ins is None:
            return {"reading": "", "reason": "no_ttf_inspector"}
        try:
            sha = ins.glyph_sha256(gid0)
        except Exception:
            sha = ""
        ck = (xref0, int(gid0), sha)
        if ck in symbol_decode_cache:
            return symbol_decode_cache[ck]
        try:
            contours = ins.simple_contours(gid0)
        except Exception:
            contours = []
        decoded = symbol_model.decode(contours)
        if not decoded.get("reading") and not contours:
            # v3.9: some Bpmf* fonts store the annotation component itself as
            # a composite of standalone Bopomofo children. Decode that nested
            # structure by exact child signatures and relative component geometry.
            decoded = symbol_model.decode_nested_composite(ins, gid0)
        symbol_decode_cache[ck] = decoded
        return decoded

    for (font, zh_id), items in sorted(shape_group_rows.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        current_readings = {str(ws.cell(item["row"], 6).value or "") for item in items}
        current_readings.discard("")
        decoded_variants = []
        seen_xrefs = set()
        complete = True
        for item in items:
            xref0 = item.get("xref")
            if not isinstance(xref0, int) or xref0 in seen_xrefs:
                continue
            seen_xrefs.add(xref0)
            decoded = _symbol_decode_for(xref0, zh_id)
            if not decoded.get("reading"):
                complete = False
                decoded_variants.append(decoded)
                break
            decoded_variants.append(decoded)
        if not seen_xrefs:
            complete = False
        predictions = {d.get("reading", "") for d in decoded_variants if d.get("reading")}
        prediction = next(iter(predictions)) if complete and len(predictions) == 1 else ""
        min_support_keys = min(
            (d.get("min_support_keys", 0) for d in decoded_variants if d.get("reading")),
            default=0,
        )
        min_support_fonts = min(
            (d.get("min_support_fonts", 0) for d in decoded_variants if d.get("reading")),
            default=0,
        )
        evidence = " || ".join(symbol_evidence_text(d) for d in decoded_variants if d.get("reading"))
        action = "未命中"

        if current_readings:
            if len(current_readings) == 1 and prediction:
                current = next(iter(current_readings))
                if prediction == current:
                    action = "既有解碼一致（自驗PASS）"
                else:
                    action = "既有解碼與符號重組衝突（只報告，不自動覆寫）"
                    symbol_audit_conflict_keys.add((font, zh_id))
                    symbol_audit_conflict_occurrences += len(items)
                    for item in items:
                        row_no = item["row"]
                        ws.cell(row_no, 32).value = "符號重組稽核衝突"
                        ws.cell(row_no, 33).value = "exact symbol recombination disagrees with current decode"
                        ws.cell(row_no, 34).value = f"current={current}; recombined={prediction}; {evidence}"[:3000]
            elif len(current_readings) > 1:
                action = "既有同鍵多讀音，略過自驗"
        elif prediction:
            action = "新增解碼"
            symbol_recomb_keys.add((font, zh_id))
            symbol_recomb_occurrences += len(items)
            source_text = f"STATIC-SYMBOL-TEMPLATES:{symbol_templates_path.name}"
            rec = {
                "bopomofo": prediction,
                "verification": "exact Bopomofo symbol recombination; template admitted by support>=2 verified keys or explicit direct isolated-symbol visual verification; all target xref variants agree",
                "notes": evidence[:2500],
            }
            for item in items:
                row_no = item["row"]
                ws.cell(row_no, 6).value = prediction
                ws.cell(row_no, 7).value = "已解碼"
                ws.cell(row_no, 8).value = "TTF exact注音符號重組"
                ws.cell(row_no, 9).value = "SYMBOL-RECOMB"
                ws.cell(row_no, 10).value = source_text
                ws.cell(row_no, 12).value = f"SYMBOL-RECOMB#{zh_id}"
                ws.cell(row_no, 31).value = prediction
                ws.cell(row_no, 32).value = "符號重組新增解碼"
                ws.cell(row_no, 33).value = f"exact signatures; min source keys={min_support_keys}"
                ws.cell(row_no, 34).value = evidence[:3000]
                sk = item["stable_key"]
                if sk in key_meta:
                    key_meta[sk]["rec"] = rec
                    key_meta[sk]["decode_basis"] = "TTF exact注音符號重組"
                    key_meta[sk]["gid_group"] = "SYMBOL-RECOMB"
                    key_meta[sk]["source_font"] = source_text
                    key_meta[sk]["group_key"] = f"SYMBOL-RECOMB#{zh_id}"

        chars = "、".join(dict.fromkeys(x.get("char", "") for x in items if x.get("char")))
        pages = ",".join(str(x) for x in sorted({x.get("page_no") for x in items if x.get("page_no")}))
        current_text = "/".join(sorted(current_readings))
        ws_symbol.append([
            font, zh_id, len(items), chars, pages, current_text, prediction, action,
            "Y" if prediction else ("N" if seen_xrefs else ""),
            min_support_keys, min_support_fonts, evidence[:4000],
            "執行期只接受預先驗證 static exact symbol/tone/neutral signatures；每個模板須有至少2個驗證來源鍵，或為明確標記的 isolated-symbol 人工視覺驗證；所有xref變體須一致；禁止中文字/句境/字典反填；不啟用fuzzy/nearest-neighbour。",
        ])

    # ------------------------------------------------------------------
    # v4.0 user-verified occurrence-level actual-pronunciation corrections.
    # These are intentionally applied after every automatic decoder and are
    # never written back into font/component maps.  The automatic value remains
    # visible in audit columns and a dedicated sheet.
    # ------------------------------------------------------------------
    manual_actual_override_count = 0
    actual_headers = {str(c.value): i for i, c in enumerate(ws[1], start=1)}
    for row_no in range(2, ws.max_row + 1):
        rv = {name: ws.cell(row_no, col).value for name, col in actual_headers.items()}
        ov = match_actual_occurrence_override(rv, actual_overrides)
        if not ov:
            continue
        old_bop = str(rv.get("實際注音") or "")
        new_bop = (ov.get("actual_reading") or "").strip()
        if not new_bop:
            continue
        ws.cell(row_no, actual_headers["自動解碼原值"]).value = old_bop
        ws.cell(row_no, actual_headers["人工實際注音覆寫"]).value = new_bop
        ws.cell(row_no, actual_headers["人工覆寫來源"]).value = ov.get("source", "")
        ws.cell(row_no, actual_headers["人工覆寫備註"]).value = ov.get("note", "")
        ws.cell(row_no, actual_headers["實際注音"]).value = new_bop
        ws.cell(row_no, actual_headers["解碼狀態"]).value = "已解碼"
        ws.cell(row_no, actual_headers["解碼依據"]).value = "使用者原頁人工覆核（出現位置限定）"
        manual_actual_override_count += 1
        ws_actual_override.append([
            rv.get("課本頁"), rv.get("字元"), rv.get("穩定注音鍵"), rv.get("font"), rv.get("font_xref"),
            rv.get("x0"), rv.get("y0"), old_bop, new_bop, ov.get("source", ""), ov.get("note", ""),
        ])

    # Recompute final counters from the post-audit worksheet so the summary is
    # guaranteed to describe the actual delivered rows rather than the first-pass
    # provisional decode.
    mapped = tt_mapped = exact_mapped = group_mapped = transform_mapped = cff_mapped = 0
    shape_mapped = bridge_mapped = outline_mapped = correction_mapped = symbol_recomb_mapped = 0
    for data_row in ws.iter_rows(min_row=2, values_only=True):
        bop = data_row[5]
        basis = str(data_row[7] or "")
        architecture = str(data_row[23] or "")
        if not bop:
            continue
        mapped += 1
        if architecture == "CFF整字注音":
            cff_mapped += 1
        elif architecture == "TrueType複合元件":
            tt_mapped += 1
            if basis in ("精確字型對照", "PDF字型子集精確覆寫"):
                exact_mapped += 1
            elif basis == "SHA綁定字形真值修正":
                correction_mapped += 1
            elif basis == "相容字型群組對照":
                group_mapped += 1
            elif basis.startswith("已驗證跨字型輪廓轉換"):
                transform_mapped += 1
            elif basis.startswith("TTF跨字型exact正規化輪廓簽名"):
                outline_mapped += 1
            elif basis.startswith("TTF字形分解自動泛化"):
                shape_mapped += 1
            elif basis.startswith("TTF跨字型glyf指紋橋接"):
                bridge_mapped += 1
            elif basis.startswith("TTF exact注音符號重組"):
                symbol_recomb_mapped += 1


    unseen_cff_total = 0
    cff_bootstrap_final_mapped = 0
    cff_bootstrap_abstained = 0
    known_cff_direct_mapped = 0
    for data_row in ws.iter_rows(min_row=2, values_only=True):
        architecture = str(data_row[23] or "")
        if architecture != "CFF整字注音":
            continue
        detection = str(data_row[35] or "")
        basis = str(data_row[7] or "")
        bop = str(data_row[5] or "")
        if detection == "unseen-CFF-structure":
            unseen_cff_total += 1
            if bop and basis.startswith("CFF未見家族零對照自舉"):
                cff_bootstrap_final_mapped += 1
            elif not bop:
                cff_bootstrap_abstained += 1
        elif bop:
            known_cff_direct_mapped += 1


    # v3.8: structural decoded counts must reflect post-pass shape/bridge/symbol
    # additions, not just first-pass mappings.
    structural_mapped = Counter()
    for data_row in ws.iter_rows(min_row=2, values_only=True):
        if data_row[5]:
            structural_mapped[(str(data_row[12] or ""), str(data_row[35] or ""))] += 1

    for (sfont, detection), occ in sorted(structural_counts.items(), key=lambda kv: (-kv[1], kv[0])):
        decn = structural_mapped.get((sfont, detection), 0)
        ws_struct.append([
            sfont, detection, "Y" if ZH_FONT_RE.search(sfont) else "N", occ, decn, occ - decn,
            "TTF 由 shifted component 結構辨識；字型名稱只供稽核。" if detection != "known-CFF-family" else "既有 CFF family 輪廓偵測。",
        ])

    for key, count in sorted(key_count.items(), key=lambda x: (-x[1], x[0])):
        m = key_meta[key]
        rec = m["rec"]
        bop = rec.get("bopomofo", "") if rec else ""
        verification = rec.get("verification", "") if rec else ""
        notes = rec.get("notes", "") if rec else ""
        chars = "、".join(key_chars[key])
        pages = ",".join(str(p) for p in sorted(key_pages[key]))
        ws_keys.append([
            key, m["group_key"], m["font"], m["gid_group"], m["zh_id"], bop,
            m["decode_basis"], m["source_font"], count, chars, pages, verification, notes,
            m["architecture"], m["cff_signatures"], m["cff_symbols"], m["tone_evidence"]
        ])
        if not bop:
            if m["architecture"] == "CFF整字注音":
                next_step = "直接檢視未知 CFF 注音符號字形，新增『符號形狀→ㄅㄆㄇㄈ』對照；禁止用中文字預期讀音反填。"
            else:
                next_step = "渲染此注音元件字形後加入對照表；若同群組已有相同ID但發生衝突，先人工確認。"
            ws_missing.append([
                key, m["group_key"], m["font"], m["gid_group"], m["zh_id"], count, chars, pages,
                next_step, m["architecture"], m["unknown"]
            ])

    coverage = (mapped / total) if total else 0
    summary = [
        ("工具", f"{PROGRAM} v{VERSION}"),
        ("來源 PDF", str(pdf_path)),
        ("匯出時間", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("PDF 頁數", doc.page_count),
        ("偵測到注音字形筆數", total),
        ("其中：TrueType複合注音", tt_total),
        ("其中：CFF整字注音", cff_total),
        ("已解碼筆數", mapped),
        ("其中：TTF精確字型對照（含PDF子集覆寫）", exact_mapped),
        ("其中：TTF SHA綁定字形真值修正", correction_mapped),
        ("其中：TTF相容字型群組", group_mapped),
        ("其中：TTF已驗證跨字型轉換", transform_mapped),
        ("其中：TTF exact正規化輪廓簽名橋接", outline_mapped),
        ("其中：TTF字形分解自動泛化", shape_mapped),
        ("其中：TTF跨字型glyf指紋橋接", bridge_mapped),
        ("其中：TTF exact注音符號重組", symbol_recomb_mapped),
        ("其中：CFF輪廓解碼總筆數", cff_mapped),
        ("其中：已知CFF家族exact直接解碼", known_cff_direct_mapped),
        ("其中：未見CFF家族零對照自舉", cff_bootstrap_final_mapped),
        ("未見CFF家族結構偵測筆數", unseen_cff_total),
        ("未見CFF家族自舉後棄權筆數", cff_bootstrap_abstained),
        ("未見CFF暫存家族數", len(cff_bootstrap_family_maps)),
        ("未見CFF暫存exact簽名數", sum(len(x) for x in cff_bootstrap_family_maps.values())),
        ("未見CFF簽名標籤衝突數", cff_bootstrap_audit_conflicts),
        ("TTF字形分解新增鍵數", len(shape_auto_keys)),
        ("TTF字形分解新增筆數", shape_auto_occurrences),
        ("TTF字形自驗衝突鍵數", len(shape_conflict_keys)),
        ("TTF字形自驗衝突筆數", shape_conflict_occurrences),
        ("TTF字形模型啟用數", len(shape_models)),
        ("TTF跨字型glyf指紋橋接新增鍵數", len(bridge_keys)),
        ("TTF跨字型glyf指紋橋接新增筆數", bridge_occurrences),
        ("TTF glyf指紋來源衝突數", len(fp_conflicts)),
        ("TTF shape-auto安全指紋來源鍵數", len(shape_auto_donor_hashes)),
        ("TTF shape-auto安全指紋來源SHA數", sum(len(x.get("hashes", set())) for x in shape_auto_donor_hashes.values())),
        ("TTF符號重組新增鍵數", len(symbol_recomb_keys)),
        ("TTF符號重組新增筆數", symbol_recomb_occurrences),
        ("TTF符號重組既有解碼衝突鍵數", len(symbol_audit_conflict_keys)),
        ("TTF符號重組既有解碼衝突筆數", symbol_audit_conflict_occurrences),
        ("TTF已驗證符號模板數", len(symbol_model.templates)),
        ("使用者原頁人工實際注音覆寫筆數", manual_actual_override_count),
        ("結構偵測人工排除筆數", max(0, ws_struct_excl.max_row - 1)),
        ("TTF exact glyf 指紋總數（內建＋使用者驗證）", len(verified_fingerprints)),
        ("使用者已升格 exact TTF glyph 數", len(user_verified_fingerprints)),
        ("TTF exact glyph truth 衝突隔離數", len(quarantined_ttf)),
        ("使用者已升格 exact CFF 整字字形 SHA-256 數", len(user_verified_cff)),
        ("TTF無衝突正規化輪廓簽名數", len(verified_outline_signatures)),
        ("舊版字型名稱規則未涵蓋之結構注音筆數", sum(v for (f,d),v in structural_counts.items() if d == "shifted-component-only")),
        ("待建立對照筆數", total - mapped),
        ("目前解碼覆蓋率", coverage),
        ("使用TTF對照表", str(map_path)),
        ("使用字型群組設定", str(groups_path)),
        ("使用CFF符號對照表", str(cff_map_path)),
        ("使用CFF跨家族CID共識表", str(cff_consensus_path)),
        ("使用TTF PDF字型子集覆寫", str(xref_overrides_path)),
        ("使用TTF已驗證跨字型轉換", str(transforms_path)),
        ("使用TTF直接驗證glyf指紋", str(fingerprints_path)),
        ("使用TTF exact正規化輪廓簽名", str(outline_signatures_path)),
        ("使用TTF SHA綁定字形真值修正", str(mapping_corrections_path)),
        ("使用TTF exact注音符號模板", str(symbol_templates_path)),
        ("使用人工實際注音出現位置覆寫", str(actual_overrides_path)),
        ("使用結構偵測排除清單", str(structural_exclusions_path)),
        ("TTF精確對照已知鍵數", len(exact_map)),
        ("TTF已驗證跨字型轉換鍵數", len(verified_transforms)),
        ("TTF安全群組對照鍵數", len(group_map)),
        ("TTF群組內衝突鍵數", len(group_conflicts)),
        ("CFF已驗證符號形狀鍵數", len(cff_symbol_map)),
        ("CFF原則", "已知家族仍以人工驗證的exact符號簽名解碼；未見家族不新增持久mapping，先以variant+CID與CID-only兩路已知家族body共識互相核對，再在本次PDF內建立暫存exact簽名。任一衝突或證據不足即棄權。"),
        ("單檔未見CFF自舉安全門檻", "每個target簽名必須無標籤衝突，且符合：同一提議由至少2個不同target glyph_id互證，或兩路CID共識的較弱來源仍至少有2個已知樣式家族支持；最終glyph每個符號都要exact命中。"),
        ("聲調安全門檻", "二聲corr 0.85~0.98；四聲corr -0.98~-0.85；三聲corr -0.20~0.60；輕聲只接受完整注音stack上方的小圓點；落在安全區外一律棄權。"),
        ("既有CFF修正", "KAICHU_MD signature 7a7543f3... 由ㄇ修正為ㄖ；依原始CFF輪廓重新目視覆核，並由2個獨立已知樣式家族同CID共識交叉確認。"),
        ("TTF字形自驗原則", "將 simple glyf 輪廓分解成注音符號與聲調；只用 exact outline signature；每個 font 必須通過 key-level OOF 後才可替未知鍵新增解碼；同一鍵所有字形變體必須一致。"),
        ("TTF跨字型指紋橋接原則", "只接受 raw glyf SHA-256 完全相同；來源限獨立OOF一致靜態鍵、OOF-ready模型對所有xref一致解出的shape-auto鍵、或isolated人工驗證指紋；目標所有xref變體須一致；任一指紋對應多個讀音即停用。"),
        ("CFF v3.4補強原則", "跨冊幾何距離只作離線候選排序；新增符號須再直接檢視字形確認；執行期仍只以已驗證 exact CFF signature 解碼，不啟用模糊最近鄰。"),
        ("v4.0結構偵測原則", "TTF 以 shifted component 結構偵測注音；已直接視覺證明為漢字部件、且原頁沒有可見注音的來源限定假陽性可列入排除清單，排除需保留 glyph/座標稽核證據。"),
        ("TTF exact輪廓簽名原則", "只接受 translation-invariant exact normalized outline signature；來源 signature 若曾對應多個讀音即排除；禁止 fuzzy/nearest-neighbour。"),
        ("v4.2符號重組原則", "將注音元件依垂直幾何拆成獨立注音符號群與聲調；執行期只接受 exact signature。模板來源分為：support>=2 的無衝突自驗模板，或直接將 isolated 注音子符號單獨渲染後人工確認的 exact 模板；nested composite 也只依 exact child signature 與相對幾何重組。禁止 fuzzy/nearest-neighbour、中文字、詞境或字典反填。"),
        ("v4.2符號重組驗證", "v3.8 基礎模板先通過靜態已知鍵與跨冊 holdout 自驗；v3.9 再加入 isolated 直接字形模板，且每次新增後以單音一致、完整詞境差異、跨冊回歸作負向稽核。研究過程曾發現並撤回多個誤標模板；最終另修正 DFYuan signature 4a098a...：ㄎ→ㄅ，修正由 isolated 字形與標準注音字形對照決定。"),
        ("安全原則", "實際注音只由 PDF 可見字形證據決定；自動解碼維持 exact/fail-closed。使用者人工確認的辨識錯誤只能以來源+頁碼+字元+穩定鍵+座標限定的 occurrence override 修正，保留自動原值且禁止回灌全域字型鍵；禁止用中文字、詞義或字典預期讀音反填。"),
        ("診斷核心", f"TTF 沿用 {DIAG_PROGRAM} v{DIAG_VERSION}；CFF 使用整字輪廓解析。"),
    ]
    for k, v in summary:
        ws_sum.append([k, v])
    ws_sum.column_dimensions["A"].width = 32
    ws_sum.column_dimensions["B"].width = 105
    ws_sum.sheet_view.showGridLines = False
    ws_sum["A1"].font = Font(bold=True)
    for row in ws_sum.iter_rows(min_row=1, max_col=2):
        if row[0].value == "目前解碼覆蓋率":
            row[1].number_format = "0.0%"
            break

    # v5.2.0 metadata-only post-pass. This does not read expected evidence and
    # does not alter any decoded pronunciation cell.
    pdf_sha256 = _sha256_file(pdf_path)
    actual_source_rows = _attach_occurrence_metadata(ws, pdf_path, pdf_sha256)
    exclusion_source_rows = _attach_occurrence_metadata(ws_struct_excl, pdf_path, pdf_sha256)
    fingerprint_payload = dict(actual_asset_fingerprint or {})
    for item in [
        ("工具版本", VERSION),
        ("workbook_schema_version", WORKBOOK_SCHEMA_VERSION),
        ("ledger_schema_version", LEDGER_SCHEMA_VERSION),
        ("pdf_sha256", pdf_sha256),
        ("actual_asset_fingerprint", fingerprint_payload.get("fingerprint", "")),
        ("actual_asset_fingerprint_components", json.dumps(fingerprint_payload.get("components", {}), ensure_ascii=False, sort_keys=True)),
        ("actual_source_row_count", len(actual_source_rows)),
        ("structural_exclusion_source_row_count", len(exclusion_source_rows)),
    ]:
        ws_runtime.append(item)

    style_sheet(ws)
    style_sheet(ws_keys)
    style_sheet(ws_missing)
    style_sheet(ws_groups)
    style_sheet(ws_shape)
    style_sheet(ws_bridge)
    style_sheet(ws_symbol)
    style_sheet(ws_struct)
    style_sheet(ws_actual_override)
    style_sheet(ws_struct_excl)
    style_sheet(ws_cff_bootstrap)
    style_sheet(ws_runtime)
    for row in ws_shape.iter_rows(min_row=2):
        row[13].number_format = "0.00%"
        row[14].number_format = "0.00%"
    wb.save(output_path)
    doc.close()
    return {
        "total": total,
        "mapped": mapped,
        "exact_mapped": exact_mapped,
        "group_mapped": group_mapped,
        "transform_mapped": transform_mapped,
        "outline_mapped": outline_mapped,
        "shape_mapped": shape_mapped,
        "shape_auto_keys": len(shape_auto_keys),
        "shape_auto_occurrences": shape_auto_occurrences,
        "shape_conflict_keys": len(shape_conflict_keys),
        "shape_conflict_occurrences": shape_conflict_occurrences,
        "shape_models": len(shape_models),
        "bridge_mapped": bridge_mapped,
        "bridge_keys": len(bridge_keys),
        "bridge_occurrences": bridge_occurrences,
        "correction_mapped": correction_mapped,
        "symbol_recomb_mapped": symbol_recomb_mapped,
        "symbol_recomb_keys": len(symbol_recomb_keys),
        "symbol_recomb_occurrences": symbol_recomb_occurrences,
        "symbol_audit_conflict_keys": len(symbol_audit_conflict_keys),
        "symbol_audit_conflict_occurrences": symbol_audit_conflict_occurrences,
        "fingerprint_conflicts": len(fp_conflicts),
        "cff_mapped": cff_mapped,
        "known_cff_direct_mapped": known_cff_direct_mapped,
        "cff_bootstrap_mapped": cff_bootstrap_final_mapped,
        "unseen_cff_total": unseen_cff_total,
        "cff_bootstrap_abstained": cff_bootstrap_abstained,
        "cff_bootstrap_families": len(cff_bootstrap_family_maps),
        "cff_bootstrap_signatures": sum(len(x) for x in cff_bootstrap_family_maps.values()),
        "cff_bootstrap_conflicts": cff_bootstrap_audit_conflicts,
        "tt_total": tt_total,
        "cff_total": cff_total,
        "coverage": coverage,
        "output": output_path,
        "group_conflicts": len(group_conflicts),
        "manual_actual_overrides": manual_actual_override_count,
        "structural_exclusions": max(0, ws_struct_excl.max_row - 1),
    }

def choose_pdf_gui():
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk(); root.withdraw(); root.attributes("-topmost", True)
        fn = filedialog.askopenfilename(title="選擇要解碼實際注音的 PDF", filetypes=[("PDF", "*.pdf")])
        root.destroy()
        return Path(fn) if fn else None
    except Exception:
        return None


def notify(title, message, error=False):
    try:
        import tkinter as tk
        from tkinter import messagebox
        root = tk.Tk(); root.withdraw(); root.attributes("-topmost", True)
        (messagebox.showerror if error else messagebox.showinfo)(title, message)
        root.destroy()
    except Exception:
        pass


def main():
    ap = argparse.ArgumentParser(description=f"{PROGRAM} v{VERSION}")
    ap.add_argument("pdf", nargs="?")
    ap.add_argument("-o", "--output")
    ap.add_argument("--map", default=str(DEFAULT_MAP))
    ap.add_argument("--groups", default=str(DEFAULT_GROUPS))
    ap.add_argument("--cff-map", default=str(DEFAULT_CFF_MAP))
    ap.add_argument("--cff-consensus", default=str(DEFAULT_CFF_CONSENSUS))
    ap.add_argument("--xref-overrides", default=str(DEFAULT_XREF_OVERRIDES))
    ap.add_argument("--transforms", default=str(DEFAULT_TRANSFORMS))
    ap.add_argument("--fingerprints", default=str(DEFAULT_FINGERPRINTS))
    ap.add_argument("--outline-signatures", default=str(DEFAULT_OUTLINE_SIGNATURES))
    ap.add_argument("--mapping-corrections", default=str(DEFAULT_MAPPING_CORRECTIONS))
    ap.add_argument("--symbol-templates", default=str(DEFAULT_SYMBOL_TEMPLATES))
    ap.add_argument("--actual-overrides", default=str(DEFAULT_ACTUAL_OVERRIDES))
    ap.add_argument("--structural-exclusions", default=str(DEFAULT_STRUCTURAL_EXCLUSIONS))
    args = ap.parse_args()
    pdf = Path(args.pdf) if args.pdf else choose_pdf_gui()
    if not pdf:
        return 0
    if not pdf.exists() or pdf.suffix.lower() != ".pdf":
        msg = f"找不到有效 PDF：{pdf}"; print(msg, file=sys.stderr); notify(PROGRAM, msg, True); return 2
    out = Path(args.output) if args.output else unique_output_path(pdf)
    try:
        root = Path(__file__).resolve().parent
        requested_assets = {
            "map": (Path(args.map).resolve(), DEFAULT_MAP.resolve()),
            "groups": (Path(args.groups).resolve(), DEFAULT_GROUPS.resolve()),
            "cff_map": (Path(args.cff_map).resolve(), DEFAULT_CFF_MAP.resolve()),
            "cff_consensus": (Path(args.cff_consensus).resolve(), DEFAULT_CFF_CONSENSUS.resolve()),
            "xref_overrides": (Path(args.xref_overrides).resolve(), DEFAULT_XREF_OVERRIDES.resolve()),
            "transforms": (Path(args.transforms).resolve(), DEFAULT_TRANSFORMS.resolve()),
            "fingerprints": (Path(args.fingerprints).resolve(), DEFAULT_FINGERPRINTS.resolve()),
            "outline_signatures": (Path(args.outline_signatures).resolve(), DEFAULT_OUTLINE_SIGNATURES.resolve()),
            "mapping_corrections": (Path(args.mapping_corrections).resolve(), DEFAULT_MAPPING_CORRECTIONS.resolve()),
            "symbol_templates": (Path(args.symbol_templates).resolve(), DEFAULT_SYMBOL_TEMPLATES.resolve()),
            "actual_overrides": (Path(args.actual_overrides).resolve(), DEFAULT_ACTUAL_OVERRIDES.resolve()),
            "structural_exclusions": (Path(args.structural_exclusions).resolve(), DEFAULT_STRUCTURAL_EXCLUSIONS.resolve()),
        }
        unmanifested = [name for name, (actual_path, required_path) in requested_assets.items() if actual_path != required_path]
        if unmanifested:
            raise ValueError(f"v5.2.0 不接受未列入 runtime manifest 的 actual 資產：{unmanifested}")
        source_validation = validate_asset_manifest(root)
        if not source_chain_ok(source_validation, "actual"):
            actual_errors = [
                f"{item.get('name')}：{'；'.join(item.get('errors') or [])}"
                for item in source_validation.get("assets", [])
                if item.get("chain") == "actual" and not item.get("ok")
            ]
            raise ValueError("actual 核心資料來源驗證失敗：" + "；".join(actual_errors))
        fingerprint = compute_actual_asset_fingerprint(
            root, pdf, source_validation, decoder_version=VERSION, source_files=ACTUAL_DECODER_SOURCE_FILES
        )
        r = decode(pdf, out, Path(args.map), Path(args.groups), Path(args.cff_map), Path(args.xref_overrides), Path(args.transforms), Path(args.fingerprints), Path(args.outline_signatures), Path(args.mapping_corrections), Path(args.symbol_templates), Path(args.actual_overrides), Path(args.structural_exclusions), Path(args.cff_consensus), fingerprint)
        final_fingerprint = compute_actual_asset_fingerprint(
            root,
            pdf,
            source_validation,
            decoder_version=VERSION,
            source_files=ACTUAL_DECODER_SOURCE_FILES,
            dynamic_dependencies=actual_workbook_dynamic_dependencies(out),
        )
        if final_fingerprint.get("fingerprint") != fingerprint.get("fingerprint"):
            rewrite_actual_workbook_fingerprint(out, final_fingerprint)
    except Exception as e:
        msg = f"解碼失敗：{e}"; print(msg, file=sys.stderr); notify(PROGRAM, msg, True); return 1
    msg = (
        f"actual 解碼處理結束（不等同全冊校對完成）。\n\n偵測注音：{r['total']} 筆\n已解碼：{r['mapped']} 筆\n"
        f"TTF精確字型：{r['exact_mapped']} 筆\nTTF群組回退：{r['group_mapped']} 筆\n"
        f"TTF跨字型轉換：{r['transform_mapped']} 筆\n"
        f"TTF字形分解自動泛化：{r['shape_mapped']} 筆（{r['shape_auto_keys']} 鍵）\n"
        f"TTF跨字型glyf指紋橋接：{r['bridge_mapped']} 筆（{r['bridge_keys']} 鍵）\n"
        f"TTF exact注音符號重組：{r['symbol_recomb_mapped']} 筆（{r['symbol_recomb_keys']} 鍵）\n"
        f"TTF符號重組稽核衝突：{r['symbol_audit_conflict_occurrences']} 筆（{r['symbol_audit_conflict_keys']} 鍵）\n"
        f"TTF字形自驗衝突：{r['shape_conflict_occurrences']} 筆（{r['shape_conflict_keys']} 鍵）\n"
        f"CFF輪廓解碼總數：{r['cff_mapped']} 筆\n"
        f"CFF未見家族零對照自舉：{r['cff_bootstrap_mapped']} / {r['unseen_cff_total']} 筆；棄權 {r['cff_bootstrap_abstained']} 筆\n"
        f"人工出現位置覆寫：{r['manual_actual_overrides']} 筆\n"
        f"結構偵測排除：{r['structural_exclusions']} 筆\n"
        f"覆蓋率：{r['coverage']:.1%}\n群組衝突鍵：{r['group_conflicts']}\n\n輸出：\n{r['output']}"
    )
    print(msg); notify(PROGRAM, msg); return 0


if __name__ == "__main__":
    raise SystemExit(main())
