from __future__ import annotations

import contextlib
import csv
import io
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import fitz
from openpyxl import load_workbook

import actual_review as actual
import cff_zero_map_batch as batch
import check_pronunciation_candidates as candidates
import export_zhuyin_readings as decoder
import global_exact_glyph_library as library
import runtime_source_validation as validation
from cross_version_compat import fingerprint_compatible
from tests.test_global_exact_glyph_read_reuse_v580 import _cff_inspector, _identity, _record, _snapshot


ROOT = Path(__file__).resolve().parents[1]
IDENTITY_COLUMNS = ("字型相容群組", "群組注音鍵", "CFF樣式群組", "CFF整字字形SHA256", "穩定注音鍵")


class _Page:
    rect = SimpleNamespace(width=600, height=800)

    def __init__(self, font, gid):
        self.font, self.gid = font, gid

    def get_label(self):
        return "1"

    def get_fonts(self, full=True):
        return [(10, "cid", "Type0", self.font, "F1", "Identity-H", 0)]

    def get_texttrace(self):
        return [{"font": self.font, "size": 12, "opacity": 1, "seqno": 1,
                 "chars": ((ord("字"), self.gid, (10, 10), (10, 10, 20, 20)),)}]

    def get_text(self, kind):
        return [] if kind == "words" else {}


class _Document:
    page_count = 1

    def __init__(self, font, gid):
        self.page = _Page(font, gid)

    def __iter__(self):
        return iter((self.page,))

    def extract_font(self, xref):
        return (self.page.font, "cid", "Type0", b"isolated font surface")

    def close(self):
        pass


def _rows(path, sheet="實際注音"):
    with contextlib.closing(load_workbook(path, read_only=True, data_only=True)) as wb:
        iterator = wb[sheet].iter_rows(values_only=True)
        headers = next(iterator)
        return [dict(zip(headers, row)) for row in iterator]


def _metadata(path):
    with contextlib.closing(load_workbook(path, read_only=True, data_only=True)) as wb:
        return dict(wb["v5.2中繼資料"].iter_rows(max_col=2, values_only=True))


class CFFBatchFingerprintTests(unittest.TestCase):
    """Real decoder bootstrap, batch, dependency/hash and analyzer boundaries.

    PDF traversal/font loading alone use a bounded synthetic surface; the real
    CFF inspector resolves a complete recording. The oracle readings are literal
    actual-only CID reference labels, independent of the expected resolver.
    """

    @classmethod
    def setUpClass(cls):
        cls.sources = validation.validate_asset_manifest(ROOT)
        if not cls.sources.get("ok"):
            raise AssertionError(cls.sources.get("errors"))

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.dynamic = self.root / "actual-evidence"
        actual.ensure_user_evidence_files(self.dynamic)
        self.snapshot = _snapshot()
        self.consensus = self.root / "consensus.csv"
        with self.consensus.open("w", encoding="utf-8", newline="") as fh:
            writer = csv.DictWriter(fh, fieldnames=("reference_mode", "variant", "glyph_id", "body", "style_support_count", "style_groups", "occurrence_count", "conflict"))
            writer.writeheader()
            for gid, body in ((1, "ㄅ"), (2, "ㄆ")):
                for mode in ("variant_gid", "gid"):
                    writer.writerow({"reference_mode": mode, "variant": "Z", "glyph_id": gid,
                                     "body": body, "style_support_count": 2, "style_groups": "REF_A|REF_B", "occurrence_count": 2})

    def decode(self, name="target", font="FixtureUnseen", gid=1):
        pdf, workbook = self.root / f"{name}.pdf", self.root / f"{name}.xlsx"
        with fitz.open() as document:
            document.new_page()
            document.save(pdf)
        inspector, _sha = _cff_inspector(self.snapshot, style_group="")
        # Patch acquisition only; inspector.decode and single-file bootstrap
        # execute unchanged, including complete-recording SHA and symbol gates.
        with patch.object(decoder.fitz, "open", return_value=_Document(font, gid)), \
                patch.object(decoder, "CFFZhuyinInspector", return_value=inspector), \
                contextlib.redirect_stdout(io.StringIO()):
            decoder.decode(pdf, workbook, decoder.DEFAULT_MAP, decoder.DEFAULT_GROUPS,
                           cff_consensus_path=self.consensus, dynamic_evidence_root=self.dynamic,
                           global_snapshot=self.snapshot)
        fingerprint = self.fingerprint(workbook, pdf)
        validation.rewrite_actual_workbook_fingerprint(workbook, fingerprint)
        return pdf, workbook, fingerprint

    def fingerprint(self, workbook, pdf, snapshot=None):
        return validation.compute_actual_asset_fingerprint(
            ROOT, pdf, self.sources, decoder_version=decoder.VERSION,
            source_files=decoder.ACTUAL_DECODER_SOURCE_FILES,
            dynamic_dependencies=actual.actual_workbook_dynamic_dependencies(workbook),
            dynamic_evidence_root=self.dynamic,
            global_exact_glyph_evidence_hashes=library.global_exact_glyph_evidence_hashes(
                snapshot or self.snapshot, actual.actual_workbook_global_exact_dependencies(workbook, pdf)),
        )

    def run_batch(self, workbooks, write_only=None):
        command = ["cff_zero_map_batch.py", "--existing-workbooks", *map(str, workbooks),
                   "--cff-consensus", str(self.consensus), "-o", str(self.root)]
        if write_only:
            command += ["--write-only-workbooks", *map(str, write_only)]
        log = io.StringIO()
        with patch.object(sys, "argv", command), contextlib.redirect_stdout(log):
            self.assertEqual(batch.main(), 0)
        return log.getvalue()

    def analyze(self, pdf, workbook, snapshot=None):
        with contextlib.redirect_stdout(io.StringIO()):
            return candidates.analyze(
                workbook, candidates.DEFAULT_DICT, candidates.DEFAULT_RULES,
                workbook.with_name(workbook.stem + "-candidates.xlsx"), pdf,
                dynamic_evidence_root=self.dynamic, runtime_root=ROOT,
                global_snapshot=snapshot or self.snapshot,
            )

    def assert_sealed(self, pdf, workbook, fingerprint, before):
        rows = _rows(workbook)
        for key in IDENTITY_COLUMNS:
            self.assertEqual(rows[0][key], before[key], key)
        current = self.fingerprint(workbook, pdf)
        metadata = _metadata(workbook)
        self.assertEqual(current["fingerprint"], fingerprint["fingerprint"])
        self.assertEqual(metadata["actual_asset_fingerprint"], fingerprint["fingerprint"])
        self.assertEqual(json.loads(metadata["actual_asset_fingerprint_components"]), fingerprint["components"])
        self.assertTrue(fingerprint_compatible(metadata["actual_asset_fingerprint"],
                                              metadata["actual_asset_fingerprint_components"], current, chain="actual"))

    def test_single_decode_seal_real_batch_then_candidate_accepts_unchanged_reading(self):
        pdf, workbook, fingerprint = self.decode()
        before = _rows(workbook)[0]
        self.assertEqual(before["實際注音"], "ㄅ")
        self.assertEqual(before["CFF樣式群組"], "AUTOBOOT:FIXTUREUNSEEN")
        log = self.run_batch([workbook])
        self.assertIn("批次新增 0；撤回 0", log)
        analyzed = self.analyze(pdf, workbook)
        self.assert_sealed(pdf, workbook, fingerprint, before)
        self.assertEqual(_rows(workbook)[0]["實際注音"], "ㄅ")
        self.assertIn("跨檔批次", _rows(workbook)[0]["解碼依據"])
        self.assertEqual(analyzed["ledger"][0]["actual"], "ㄅ")
        self.assertTrue(analyzed["reconciliation"]["ok"])

    def test_cross_file_addition_preserves_unknown_identity_and_unrelated_bytes(self):
        pdf, workbook, fingerprint = self.decode(gid=3)
        donor_pdf, donor, donor_fingerprint = self.decode("donor", "OtherUnseen", gid=1)
        before = _rows(workbook)[0]
        self.assertFalse(before["實際注音"])
        self.assertFalse(before["CFF樣式群組"])
        donor_bytes = donor.read_bytes()
        log = self.run_batch([workbook, donor], [workbook])
        self.assertIn("批次新增 1；撤回 0", log)
        self.assert_sealed(pdf, workbook, fingerprint, before)
        self.assertEqual(_rows(workbook)[0]["實際注音"], "ㄅ")
        self.assertEqual(actual.actual_workbook_dynamic_dependencies(workbook)["cff_glyph_keys"], [])
        self.assertEqual(actual.actual_workbook_global_exact_dependencies(workbook, pdf), ())
        self.assertEqual(donor.read_bytes(), donor_bytes)
        self.assertEqual(self.fingerprint(donor, donor_pdf)["fingerprint"], donor_fingerprint["fingerprint"])
        self.analyze(pdf, workbook)

    def test_conflicting_batch_retracts_reading_without_changing_dependency_or_seal(self):
        pdf, workbook, fingerprint = self.decode()
        _donor_pdf, donor, _ = self.decode("conflict", gid=2)
        before = _rows(workbook)[0]
        log = self.run_batch([workbook, donor])
        self.assertIn("批次新增 0；撤回 2", log)
        self.assert_sealed(pdf, workbook, fingerprint, before)
        after = _rows(workbook)[0]
        self.assertFalse(after["實際注音"])
        self.assertEqual(after["CFF符號序列"], "?")
        self.assertIn("安全門檻未通過", after["解碼依據"])
        self.assertFalse(_rows(workbook, "注音鍵彙整")[0]["實際注音"])
        self.assertTrue(any("互斥" in str(row["處置"]) for row in _rows(workbook, "CFF批次自舉稽核")))
        self.analyze(pdf, workbook)

    def test_pdf_applicable_global_change_and_unprovable_roster_still_rejected(self):
        pdf, workbook, _ = self.decode()
        self.run_batch([workbook])
        row = _rows(workbook)[0]
        identity = _identity(row["CFF整字字形SHA256"], kind="CFF_GLYPH_SHA256", style_group=row["CFF樣式群組"])
        changed = _snapshot(trusted=(_record(identity, reading="ㄆ"),))
        with self.assertRaisesRegex(ValueError, "evidence assets 已變更"):
            self.analyze(pdf, workbook, changed)
        broken = self.root / "missing-roster.xlsx"
        shutil.copyfile(workbook, broken)
        with contextlib.closing(load_workbook(broken)) as wb:
            ws = wb["實際注音"]
            h = batch._headers(ws)
            ws.cell(2, h["CFF整字字形SHA256"]).value = ""
            wb.save(broken)
        self.assertIsNone(actual.actual_workbook_global_exact_dependencies(broken, pdf))
        with self.assertRaisesRegex(ValueError, "evidence assets 已變更"):
            self.analyze(pdf, broken)
        pdf.write_bytes(pdf.read_bytes() + b"\n% changed PDF bytes\n")
        with self.assertRaisesRegex(ValueError, "PDF SHA-256 與現版 PDF 不一致"):
            self.analyze(pdf, workbook)

    def test_manual_occurrence_override_survives_conflicting_batch(self):
        (self.dynamic / actual.OCCURRENCE_OVERRIDE_FILE).write_text(
            ",".join(actual.OVERRIDE_HEADERS) + "\ntarget,,1,字,,10,10,ㄈ,direct visual fixture,occurrence only\n",
            encoding="utf-8")
        pdf, workbook, fingerprint = self.decode()
        _donor_pdf, donor, _ = self.decode("conflict", gid=2)
        before = _rows(workbook)[0]
        self.assertEqual(before["實際注音"], "ㄈ")
        self.assertEqual(before["自動解碼原值"], "ㄅ")
        self.run_batch([workbook, donor])
        self.assertEqual(_rows(workbook)[0], before)
        self.assert_sealed(pdf, workbook, fingerprint, before)
        self.analyze(pdf, workbook)

    def test_insufficient_batch_abstains_without_creating_reusable_identity(self):
        pdf, workbook, fingerprint = self.decode(gid=3)
        before = _rows(workbook)[0]
        self.assertFalse(before["實際注音"])
        log = self.run_batch([workbook])
        self.assertIn("批次新增 0；撤回 0", log)
        self.assert_sealed(pdf, workbook, fingerprint, before)
        self.assertFalse(_rows(workbook)[0]["實際注音"])
        self.assertFalse(_rows(workbook)[0]["CFF樣式群組"])
        analyzed = self.analyze(pdf, workbook)
        self.assertFalse(analyzed["ledger"][0]["actual"])
        self.assertEqual(analyzed["ledger"][0]["expected_status"], "RESOLVED")

    def test_non_unseen_rows_remain_unchanged_and_candidate_compatible(self):
        pdf, workbook, fingerprint = self.decode()
        with contextlib.closing(load_workbook(workbook)) as wb:
            ws = wb["實際注音"]
            ws.cell(2, batch._headers(ws)["注音結構偵測來源"]).value = "known-CFF-family"
            wb.save(workbook)
        before = _rows(workbook)[0]
        self.run_batch([workbook])
        self.assertEqual(_rows(workbook)[0], before)
        self.assert_sealed(pdf, workbook, fingerprint, before)
        self.analyze(pdf, workbook)


if __name__ == "__main__":
    unittest.main()
