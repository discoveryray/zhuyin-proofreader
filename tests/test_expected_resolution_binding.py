from __future__ import annotations

import copy
import os
import tempfile
import tkinter as tk
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

import occurrence_ledger as ol
import review_gui as gui
import standalone_proofread as sp
from review_save_service import ReviewSaveService
from tests.test_staged_review_navigation_v580 import create_staged_navigation_fixture
from tests.review_save_test_support import wait_for_save


def captured_confirmation(row, db):
    saved = []
    app = SimpleNamespace(db=db, save_event=lambda entry, event: saved.append(event))
    gui.ReviewApp._confirm_difference_entry(app, row)
    return saved[0]


class _ResolutionFixture:
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.output = Path(self.tmp.name) / "isolated-project"
        self.localapp = patch.dict(os.environ, {"LOCALAPPDATA": str(Path(self.tmp.name) / "LocalAppData")})
        self.localapp.start()
        self.addCleanup(self.localapp.stop)
        self.manifest = create_staged_navigation_fixture(self.output)
        self.base = next(row for row in self.manifest["records"] if row["state"] == "DIFFERENCE_PENDING_CONFIRMATION")

    def import_resolution(self):
        workbook = sp.export_pending_for_gpt(self.output)
        wb = load_workbook(workbook)
        sheet = wb["待判定候選"]
        columns = {c.value: c.column for c in sheet[1]}
        number = next(n for n in range(2, sheet.max_row + 1)
                      if sheet.cell(n, columns["review_id"]).value == self.base["review_id"])
        for key, value in {"action": "解決expected證據",
                           "proposed_expected_set": "|".join(self.base["expected_set"]),
                           "proposed_expected_evidence": "isolated imported independent dictionary proof",
                           "proposed_context_evidence": "isolated imported occurrence-specific context",
                           "note": "isolated original resolution note"}.items():
            sheet.cell(number, columns[key], value)
        wb.save(workbook)
        wb.close()
        sp.import_gpt_decisions(self.output, workbook)
        return workbook, sp.load_or_initialize_db(self.output)

    def bound_confirmation(self):
        source = {"action": "解決expected證據", "expected_set": list(self.base["expected_set"]),
                  "expected_evidence": "isolated independent resolution proof",
                  "context_evidence": "isolated resolution context", "source": "isolated explicit expected source",
                  "resolution_reason": "", "note": "original expected note"}
        expected = sp.bind_expected_resolution(self.base, source, self.manifest)
        reviewed = sp._apply_review_event(self.base, expected, expected_manifest=self.manifest)
        db = {**sp.normalize_db({}), "events": {self.base["review_id"]: expected}}
        return expected, captured_confirmation(reviewed, db)

    def legacy_resolution(self, *, with_fingerprint):
        if not with_fingerprint:
            for key in ("expected_asset_fingerprint", "expected_asset_fingerprint_components",
                        "compatible_expected_asset_fingerprints"):
                self.manifest.pop(key, None)
            sp.seal_manifest(self.manifest)
            sp.json_save(self.output / "校對工作階段.json", self.manifest)
        source = {"action": "解決expected證據", "expected_set": list(self.base["expected_set"]),
                  "expected_evidence": "independent historical expected source",
                  "context_evidence": "historical reviewed context",
                  "source": "GPT／人工 expected 證據報表匯入（v5.5 dual-lane transaction）",
                  "resolution_reason": "", "note": "preserved historical expected note",
                  "updated_at": "2026-09-29T00:00:00"}
        db = sp.normalize_db({})
        db["events"][self.base["review_id"]] = copy.deepcopy(source)
        sp.json_save(self.output / "人工判定資料庫.json", db)
        reviewed = next(row for row in sp.materialize_ledger(self.manifest, db)
                        if row["review_id"] == self.base["review_id"])
        return db, source, reviewed

class ExpectedResolutionReplayTests(_ResolutionFixture, unittest.TestCase):
    def test_unbound_legacy_resolution_invalid_confirmation_never_replaces_durable_source(self):
        fingerprinted = copy.deepcopy(self.manifest)
        for with_fingerprint in (True, False):
            with self.subTest(with_fingerprint=with_fingerprint):
                self.manifest = copy.deepcopy(fingerprinted)
                sp.json_save(self.output / "校對工作階段.json", self.manifest)
                db, source, reviewed = self.legacy_resolution(with_fingerprint=with_fingerprint)
                self.assertEqual(reviewed["state"], "DIFFERENCE_PENDING_CONFIRMATION")
                self.assertEqual(reviewed["expected_evidence"], source["expected_evidence"])
                self.assertEqual(reviewed["context_evidence"], source["context_evidence"])
                self.assertEqual(reviewed["expected_set"], self.base["expected_set"])
                event = captured_confirmation(reviewed, db)  # The real GUI generator; no Tk claim.
                self.assertNotIn("expected_resolution_binding", event)
                before = (self.output / "人工判定資料庫.json").read_bytes()
                with patch.object(sp, "json_save", wraps=sp.json_save) as writer:
                    with self.assertRaisesRegex(ValueError, "確認.*失效"):
                        ReviewSaveService(self.output).save_event(self.base["review_id"], event,
                            expected_manifest=self.manifest, expected_db=db)
                writer.assert_not_called()
                self.assertEqual((self.output / "人工判定資料庫.json").read_bytes(), before)
                durable = sp.load_or_initialize_db(self.output)
                self.assertEqual(durable, db)
                self.assertEqual(durable["events"][self.base["review_id"]], source)
                reopened = next(row for row in sp.materialize_ledger(self.manifest, durable)
                                if row["review_id"] == self.base["review_id"])
                self.assertEqual(reopened, reviewed)

    def test_direct_baseline_gui_confirmation_without_binding_saves_and_reopens(self):
        fingerprinted = copy.deepcopy(self.manifest)
        for with_fingerprint in (True, False):
            with self.subTest(with_fingerprint=with_fingerprint):
                self.manifest = copy.deepcopy(fingerprinted)
                if not with_fingerprint:
                    for key in ("expected_asset_fingerprint", "expected_asset_fingerprint_components"):
                        self.manifest.pop(key, None)
                    sp.seal_manifest(self.manifest)
                sp.json_save(self.output / "校對工作階段.json", self.manifest)
                db = sp.normalize_db({})
                sp.json_save(self.output / "人工判定資料庫.json", db)
                event = captured_confirmation(self.base, db)
                self.assertNotIn("expected_resolution_binding", event)
                result = ReviewSaveService(self.output).save_event(self.base["review_id"], event,
                    expected_manifest=self.manifest, expected_db=db)
                self.assertEqual(result.resolved_entry["state"], "TEXTBOOK_ERROR_CONFIRMED")
                durable = sp.load_or_initialize_db(self.output)
                reopened = next(row for row in sp.materialize_ledger(self.manifest, durable)
                                if row["review_id"] == self.base["review_id"])
                self.assertEqual(reopened, result.resolved_entry)
                self.assertEqual(reopened["expected_evidence"], self.base["expected_evidence"])
                self.assertNotIn("expected_resolution_binding", durable["events"][self.base["review_id"]])

    def test_invalid_gui_actual_target_or_expected_evidence_snapshot_is_not_written(self):
        db = sp.load_or_initialize_db(self.output)
        original = captured_confirmation(self.base, db)
        before = (self.output / "人工判定資料庫.json").read_bytes()
        for section, key, value in (("actual_snapshot", "actual", "ㄎㄢˊ"),
                                    ("target", "review_id", "different-target"),
                                    ("expected_snapshot", "expected_evidence", "different source proof")):
            with self.subTest(section=section):
                event = copy.deepcopy(original)
                event["gui_confirmation"][section][key] = value
                with self.assertRaisesRegex(ValueError, "確認.*失效"):
                    ReviewSaveService(self.output).save_event(self.base["review_id"], event,
                        expected_manifest=self.manifest, expected_db=db)
                self.assertEqual((self.output / "人工判定資料庫.json").read_bytes(), before)
                self.assertEqual(sp.load_or_initialize_db(self.output), db)

    def test_same_reading_baseline_source_context_or_fingerprint_drift_is_not_overwritten(self):
        _, event = self.bound_confirmation()
        for update in ({"expected_evidence": "new independent base proof"},
                       {"context_evidence": "new independent base context"},
                       {"expected_set": ["ㄎㄢˊ"]}, {"printed_page": "different page"}):
            with self.subTest(update=update):
                changed = {**self.base, **update}
                replay = sp._apply_review_event(changed, event, expected_manifest=self.manifest)
                self.assertEqual(replay["review_event_replay_status"], "INVALIDATED_EXPECTED_BINDING_DRIFT")
                self.assertEqual(replay["expected_evidence"], changed["expected_evidence"])
                self.assertEqual(replay["expected_set"], changed["expected_set"])
                self.assertNotIn("gui_confirmation", replay)
        for field in ("expected_asset_fingerprint", "session_id"):
            changed_manifest = copy.deepcopy(self.manifest)
            changed_manifest[field] = "d" * 64 if field == "expected_asset_fingerprint" else "other-session"
            changed_manifest = sp.seal_manifest(changed_manifest)
            replay = sp._apply_review_event(self.base, event, expected_manifest=changed_manifest)
            self.assertEqual(replay["review_event_replay_status"], "INVALIDATED_EXPECTED_BINDING_DRIFT")
            self.assertEqual(replay["expected_evidence"], self.base["expected_evidence"])

    def test_actual_drift_revokes_confirmation_but_retains_independent_resolution(self):
        _, event = self.bound_confirmation()
        for update in ({"actual": "ㄎㄢˊ"}, {"actual_evidence": "new actual proof"}):
            replay = sp._apply_review_event({**self.base, **update}, event, expected_manifest=self.manifest)
            self.assertNotEqual(replay["state"], "TEXTBOOK_ERROR_CONFIRMED")
            self.assertNotIn("gui_confirmation", replay)
            self.assertEqual(replay["expected_evidence"], event["expected_evidence"])
        db = {**sp.normalize_db({}), "events": {self.base["review_id"]: event}}
        sp.json_save(self.output / "人工判定資料庫.json", db)
        self.assertEqual(sp._clear_actual_dependent_events(self.output, [self.base], {self.base["occurrence_id"]}), 1)
        retained = sp.load_or_initialize_db(self.output)["events"][self.base["review_id"]]
        self.assertEqual(retained["expected_resolution_binding"], event["expected_resolution_binding"])
        self.assertEqual(sp.expected_resolution_payload(retained), event["expected_resolution_binding"]["resolution"])
        replay = sp.materialize_ledger(self.manifest, sp.load_or_initialize_db(self.output))
        self.assertEqual(next(r for r in replay if r["review_id"] == self.base["review_id"])["expected_evidence"], event["expected_evidence"])

    def test_missing_unknown_tampered_or_noncanonical_binding_is_rejected(self):
        _, event = self.bound_confirmation()
        cases = [None, {}, {**event["expected_resolution_binding"], "version": 2},
                 {**event["expected_resolution_binding"], "version": True},
                 {**event["expected_resolution_binding"], "extra": "unknown"},
                 {**event["expected_resolution_binding"], "integrity_sha256": "0" * 64}]
        for binding in cases:
            with self.subTest(binding=binding), self.assertRaises(ol.InvalidTransitionError):
                sp._apply_review_event(self.base, {**event, "expected_resolution_binding": binding}, expected_manifest=self.manifest)
        bad = copy.deepcopy(event)
        bad["expected_evidence"] = "tampered saved source"
        with self.assertRaises(ol.InvalidTransitionError):
            sp._apply_review_event(self.base, bad, expected_manifest=self.manifest)
        bad = copy.deepcopy(event)
        binding = bad["expected_resolution_binding"]
        binding["resolution"]["expected_set"] = ["not bopomofo", *self.base["expected_set"]]
        binding["integrity_sha256"] = sp._hash_review_payload({k:v for k,v in binding.items() if k != "integrity_sha256"})
        with self.assertRaises(ol.InvalidTransitionError):
            sp._apply_review_event(self.base, bad, expected_manifest=self.manifest)
        with self.assertRaises(ol.InvalidTransitionError):
            sp._apply_review_event(self.base, event)

    def test_binding_contains_no_actual_and_confirmation_cannot_mint_or_drop_it(self):
        expected,event = self.bound_confirmation()
        self.assertEqual(sp.bind_expected_resolution({**self.base, "actual": "ㄎㄢˊ", "actual_evidence": "different"},
                         expected["expected_resolution_binding"]["resolution"], self.manifest)["expected_resolution_binding"],
                         expected["expected_resolution_binding"])
        db = {**sp.normalize_db({}), "events": {self.base["review_id"]: expected}}
        sp.json_save(self.output / "人工判定資料庫.json", db)
        before = (self.output / "人工判定資料庫.json").read_bytes()
        without = {k:v for k,v in event.items() if k != "expected_resolution_binding"}
        with self.assertRaisesRegex(ValueError, "binding"):
            ReviewSaveService(self.output).save_event(self.base["review_id"], without,
                                                     expected_manifest=self.manifest, expected_db=db)
        self.assertEqual((self.output / "人工判定資料庫.json").read_bytes(), before)

    def test_legacy_adapter_requires_exact_workbook_and_preserves_original_event(self):
        workbook,db = self.import_resolution()
        reviewed = next(r for r in sp.materialize_ledger(self.manifest, db) if r["review_id"] == self.base["review_id"])
        event = captured_confirmation(reviewed, db)
        event.pop("expected_resolution_binding")
        db["events"][self.base["review_id"]] = event
        sp.json_save(self.output / "人工判定資料庫.json", db)
        arguments = dict(review_ids=[self.base["review_id"]],
                         expected_manifest_sha256=sp.sha256_file(self.output / "校對工作階段.json"),
                         expected_db_sha256=sp.sha256_file(self.output / "人工判定資料庫.json"),
                         expected_workbook_sha256=sp.sha256_file(workbook))
        before = (self.output / "人工判定資料庫.json").read_bytes()
        with self.assertRaisesRegex(ValueError, "SHA"):
            sp.bind_legacy_gui_expected_resolutions(self.output, workbook,
                **{**arguments, "expected_workbook_sha256": "0" * 64})
        self.assertEqual((self.output / "人工判定資料庫.json").read_bytes(), before)
        result = sp.bind_legacy_gui_expected_resolutions(self.output, workbook, **arguments, dry_run=True)
        self.assertTrue(result["dry_run"])
        self.assertEqual((self.output / "人工判定資料庫.json").read_bytes(), before)
        manifest_path = self.output / "校對工作階段.json"
        original_manifest_bytes = manifest_path.read_bytes()
        changed = copy.deepcopy(self.manifest)
        next(r for r in changed["records"] if r["review_id"] == self.base["review_id"])["expected_evidence"] = "new base proof, same reading"
        sp.json_save(manifest_path, sp.seal_manifest(changed))
        with self.assertRaisesRegex(ValueError, "baseline"):
            sp.bind_legacy_gui_expected_resolutions(self.output, workbook,
                **{**arguments, "expected_manifest_sha256": sp.sha256_file(manifest_path)})
        self.assertEqual((self.output / "人工判定資料庫.json").read_bytes(), before)
        manifest_path.write_bytes(original_manifest_bytes)
        sp.bind_legacy_gui_expected_resolutions(self.output, workbook, **arguments)
        durable = sp.load_or_initialize_db(self.output)
        bound = durable["events"][self.base["review_id"]]
        self.assertEqual({k:v for k,v in bound.items() if k != "expected_resolution_binding"}, event)
        self.assertEqual(sp.materialize_ledger(self.manifest, durable)[4]["state"], "TEXTBOOK_ERROR_CONFIRMED")

    def test_import_gui_confirm_save_reopen_repair_pending_export(self):
        _, db = self.import_resolution()
        reviewed = next(row for row in sp.materialize_ledger(self.manifest, db)
                        if row["review_id"] == self.base["review_id"])
        self.assertEqual(reviewed["expected_set"], self.base["expected_set"])
        self.assertNotEqual(reviewed["expected_evidence"], self.base["expected_evidence"])
        confirmation = captured_confirmation(reviewed, db)
        saved = ReviewSaveService(self.output).save_event(self.base["review_id"], confirmation,
                                                        expected_manifest=self.manifest, expected_db=db)
        self.assertEqual(saved.resolved_entry["state"], "TEXTBOOK_ERROR_CONFIRMED")
        self.assertEqual(saved.resolved_entry["actual"], self.base["actual"])
        durable = sp.load_or_initialize_db(self.output)
        self.assertEqual(durable["events"][self.base["review_id"]], saved.db["events"][self.base["review_id"]])

        def unchanged_pipeline(pdfs, output, **kwargs):
            # The real repair preflight and all replay/report/export validators run;
            # only decoding is substituted for this explicitly synthetic PDF.
            self.assertEqual(kwargs["session_id_override"], self.manifest["session_id"])
            sp.save_pending_json(output, self.manifest, durable)
            return sp.generate_report(output, self.manifest, durable)

        with patch.object(sp, "run_pipeline_pdfs", side_effect=unchanged_pipeline):
            sp.repair_project_state(self.output)
        reopened = sp.materialize_ledger(self.manifest, sp.load_or_initialize_db(self.output))
        confirmed = next(row for row in reopened if row["review_id"] == self.base["review_id"])
        self.assertEqual(confirmed["state"], "TEXTBOOK_ERROR_CONFIRMED")
        self.assertEqual(confirmed["gui_confirmation"]["confirmed_at"], confirmation["gui_confirmation"]["confirmed_at"])
        exported = sp.export_pending_for_gpt(self.output)
        wb = load_workbook(exported, read_only=True)
        try:
            rows = list(wb["待判定候選"].values)
            index = rows[0].index("review_id")
            self.assertNotIn(self.base["review_id"], [row[index] for row in rows[1:]])
            self.assertGreater(len(rows), 1)
        finally:
            wb.close()
        gate = sp.json_load_strict(self.output / "pipeline_status.json")
        self.assertFalse(gate["completion_gate"]["complete"])


class ExpectedResolutionMaterializationTests(_ResolutionFixture, unittest.TestCase):
    def bound_batch(self, count):
        baseline, _ = sp.prepare_review_ledger(self.manifest, {})
        db = sp.normalize_db({})
        expected = copy.deepcopy(baseline)
        for index, row in enumerate(baseline[:count]):
            event = sp.bind_expected_resolution(row, {
                "action": "解決expected證據", "expected_set": ["ㄎㄢ"],
                "expected_evidence": f"independent dictionary proof {index}",
                "context_evidence": f"independent occurrence context {index}",
                "source": "isolated expected evidence", "resolution_reason": "", "note": "",
            }, self.manifest)
            reviewed = sp._apply_review_event(row, event, expected_manifest=self.manifest)
            db["events"][row["review_id"]] = event
            if reviewed["state"] == "DIFFERENCE_PENDING_CONFIRMATION":
                event = captured_confirmation(reviewed, db)
                db["events"][row["review_id"]] = event
                reviewed = sp._apply_review_event(row, event, expected_manifest=self.manifest)
            expected[index] = reviewed
        return baseline, db, expected

    def test_whole_manifest_hash_is_once_for_one_or_many_bound_events(self):
        for count in (1, len(self.manifest["records"])):
            baseline, db, expected = self.bound_batch(count)
            original = copy.deepcopy((self.manifest, db))
            for operation in (sp.prepare_review_ledger, sp.materialize_ledger):
                with self.subTest(count=count, operation=operation.__name__), \
                        patch.object(sp, "manifest_integrity_sha256", wraps=sp.manifest_integrity_sha256) as hash_manifest:
                    result = operation(self.manifest, db)
                    actual = result[1] if operation is sp.prepare_review_ledger else result
                    self.assertEqual(actual, expected)
                    self.assertEqual(hash_manifest.call_count, 1)
                self.assertEqual(actual[0]["state"], "TEXTBOOK_ERROR_CONFIRMED")
                self.assertEqual(actual[0]["expected_set"], ["ㄎㄢ"])
                self.assertEqual(actual[0]["expected_evidence"], "independent dictionary proof 0")
                self.assertEqual([row["actual"] for row in actual], [row["actual"] for row in baseline])
                if operation is sp.prepare_review_ledger:
                    self.assertEqual(result[0], baseline)
                self.assertEqual((self.manifest, db), original)

    def test_validation_and_replay_use_one_detached_nested_snapshot(self):
        self.manifest["reusable_expected_rules"] = sp._canonical_rule_doc({})
        sp.seal_manifest(self.manifest)
        baseline, db, expected = self.bound_batch(2)
        original = copy.deepcopy(self.manifest)
        validate = sp.validate_manifest_integrity

        def mutate_source_after_validation(snapshot):
            validate(snapshot)
            self.manifest["records"][0]["source_record"]["所在行"] = "changed source context"
            self.manifest["records"][1]["expected_set"].append("ㄎㄢˊ")
            self.manifest["expected_asset_fingerprint_components"]["expected_resolver_semantics_epoch"] = "future"
            self.manifest["reusable_expected_rules"]["rules"].append({"invalid": "new rule"})

        with patch.object(sp, "validate_manifest_integrity", side_effect=mutate_source_after_validation) as validator:
            result = sp.prepare_review_ledger(self.manifest, db)
        self.assertEqual(validator.call_count, 1)
        self.assertEqual(result, (baseline, expected))
        self.assertNotEqual(self.manifest, original)
        with self.assertRaisesRegex(ValueError, "payload hash"):
            sp.materialize_ledger(self.manifest, db)

    def test_reloaded_changed_manifest_is_revalidated_and_stale_event_invalidated(self):
        _, db, expected = self.bound_batch(2)
        self.assertEqual(sp.materialize_ledger(self.manifest, db), expected)
        changed = copy.deepcopy(self.manifest)
        changed["records"][0]["source_record"]["所在行"] = "new sealed source context"
        sp.json_save(self.output / "校對工作階段.json", sp.seal_manifest(changed))
        reloaded = sp.json_load_strict(self.output / "校對工作階段.json")
        with patch.object(sp, "manifest_integrity_sha256", wraps=sp.manifest_integrity_sha256) as hash_manifest:
            baseline, actual = sp.prepare_review_ledger(reloaded, db)
        self.assertEqual(hash_manifest.call_count, 1)
        self.assertEqual(actual[0]["review_event_replay_status"], "INVALIDATED_EXPECTED_BINDING_DRIFT")
        self.assertEqual(actual[0]["expected_evidence"], baseline[0]["expected_evidence"])
        self.assertNotIn("gui_confirmation", actual[0])
        self.assertEqual(actual[1:], expected[1:])

    def test_later_event_integrity_source_and_gui_checks_are_not_skipped(self):
        _, db, expected = self.bound_batch(2)
        review_id = self.manifest["records"][1]["review_id"]
        for mutation in ("binding", "source", "legacy"):
            changed = copy.deepcopy(db)
            event = changed["events"][review_id]
            binding = event["expected_resolution_binding"]
            if mutation == "binding":
                binding["integrity_sha256"] = "0" * 64
            elif mutation == "source":
                event["expected_evidence"] = "tampered source payload"
            else:
                binding["legacy_evidence"] = {"adapter": "unknown"}
                binding["integrity_sha256"] = sp._hash_review_payload({
                    k: v for k, v in binding.items() if k != "integrity_sha256"})
            with self.subTest(mutation=mutation), self.assertRaisesRegex(ol.InvalidTransitionError, review_id):
                sp.materialize_ledger(self.manifest, changed)
        changed = copy.deepcopy(db)
        changed["events"][review_id]["gui_confirmation"]["actual_snapshot"]["actual"] = "ㄎㄢˊ"
        actual = sp.materialize_ledger(self.manifest, changed)
        self.assertEqual(actual[0], expected[0])
        self.assertEqual(actual[1]["review_event_replay_status"], "INVALIDATED_CONFIRMATION_SNAPSHOT")
        self.assertEqual(actual[1]["expected_evidence"], "independent dictionary proof 1")
        self.assertNotIn("gui_confirmation", actual[1])

    def test_standalone_validation_never_reuses_previous_manifest_trust(self):
        baseline, db, expected = self.bound_batch(1)
        event = next(iter(db["events"].values()))
        self.assertEqual(sp.materialize_ledger(self.manifest, db), expected)
        for operation in (
            lambda: sp._validate_expected_resolution_binding(baseline[0], event, self.manifest),
            lambda: sp._apply_review_event(baseline[0], event, expected_manifest=self.manifest),
        ):
            with patch.object(sp, "manifest_integrity_sha256", wraps=sp.manifest_integrity_sha256) as hash_manifest:
                operation()
                self.assertEqual(hash_manifest.call_count, 1)
        self.manifest["records"][0]["source_record"]["所在行"] = "unsealed edit"
        for operation in (
            lambda: sp._validate_expected_resolution_binding(baseline[0], event, self.manifest),
            lambda: sp._apply_review_event(baseline[0], event, expected_manifest=self.manifest),
            lambda: sp.materialize_ledger(self.manifest, db),
        ):
            with self.assertRaisesRegex(ValueError, "payload hash"):
                operation()


class ExpectedResolutionGuiTests(_ResolutionFixture, unittest.TestCase):
    def test_real_tk_invalid_unbound_confirmation_reports_failure_and_keeps_current_row(self):
        db, source, reviewed = self.legacy_resolution(with_fingerprint=False)
        root = tk.Tk()
        self.addCleanup(root.destroy)
        root.withdraw()
        window = tk.Toplevel(root)
        app = gui.ReviewApp(window, self.output)
        target = next(row for row in app.records if row["review_id"] == self.base["review_id"])
        app.index = app.records.index(target)
        app.show()
        before_index, before_row = app.index, copy.deepcopy(app.current())
        before = (self.output / "人工判定資料庫.json").read_bytes()
        with patch.object(gui.messagebox, "showerror") as error, \
                patch.object(gui.messagebox, "showinfo") as success:
            app.primary.invoke()
            wait_for_save(app)
        error.assert_called_once()
        self.assertIn("失效", error.call_args.args[1])
        success.assert_not_called()
        self.assertFalse(app._last_event_saved)
        self.assertEqual(app.index, before_index)
        self.assertEqual(app.current(), before_row)
        self.assertEqual(app.db, db)
        self.assertEqual((self.output / "人工判定資料庫.json").read_bytes(), before)
        window.destroy()
        reopened_window = tk.Toplevel(root)
        reopened = gui.ReviewApp(reopened_window, self.output)
        reopened_row = next(row for row in reopened.records if row["review_id"] == self.base["review_id"])
        self.assertEqual(reopened_row, reviewed)
        self.assertEqual(reopened.db["events"][self.base["review_id"]], source)
        reopened_window.destroy()

    def test_real_tk_imported_expected_confirm_saves_reopens_and_advances(self):
        _, imported = self.import_resolution()
        root = tk.Tk()  # Required Windows capability; do not skip missing Tk.
        self.addCleanup(root.destroy)
        window = tk.Toplevel(root)
        app = gui.ReviewApp(window, self.output)
        root.withdraw()
        target = next(row for row in app.records if row["review_id"] == self.base["review_id"])
        self.assertEqual(target["expected_set"], self.base["expected_set"])
        self.assertNotEqual(target["expected_evidence"], self.base["expected_evidence"])
        app.index = app.records.index(target)
        app.show()
        self.assertEqual(app.primary.cget("text"), "確認教材錯誤")
        with patch.object(gui.messagebox, "showinfo"), patch.object(gui.messagebox, "showerror") as error:
            app.primary.invoke()
            wait_for_save(app)
        error.assert_not_called()
        self.assertTrue(app._last_event_saved)
        self.assertNotEqual(app.current()["review_id"], self.base["review_id"])
        durable = sp.load_or_initialize_db(self.output)
        event = durable["events"][self.base["review_id"]]
        self.assertEqual(event["expected_resolution_binding"], imported["events"][self.base["review_id"]]["expected_resolution_binding"])
        ledger = sp.materialize_ledger(self.manifest, durable)
        self.assertEqual(next(r for r in ledger if r["review_id"] == self.base["review_id"])["state"], "TEXTBOOK_ERROR_CONFIRMED")
        window.destroy()
        reopened_window = tk.Toplevel(root)
        reopened = gui.ReviewApp(reopened_window, self.output)
        self.assertEqual(reopened.db["events"][self.base["review_id"]], event)
        self.assertNotIn(self.base["review_id"], [row["review_id"] for row in reopened.records])
        reopened_window.destroy()
