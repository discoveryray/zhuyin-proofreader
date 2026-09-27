"""Round 24: exact historical imports and target PDF bytes at final writes."""
import copy
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import load_workbook
import pdf_portability as p
import standalone_proofread as sp
from review_save_service import ReviewSaveService
from tests.test_pdf_portability import pdf, project, _filled_expected_excel
from tests.test_pdf_portability_prewrite import pair, state


def edited_excel(first, output, *, evidence=None, title=None):
    book = load_workbook(first)
    if evidence is not None:
        sheet = book['待判定候選']; headers = [cell.value for cell in sheet[1]]
        sheet.cell(2, headers.index('proposed_expected_evidence') + 1, evidence)
    if title is not None: book.properties.title = title
    book.save(output); book.close()
    return output


def adjudicate(b, rid, evidence='LOCAL ADJUDICATION'):
    ReviewSaveService(b).save_event(rid, {'action': '解決expected證據',
        'expected_set': ['ㄐㄩㄝˊ'], 'expected_evidence': evidence, 'context_evidence': '角色/角@1'})


def excel_pair(a, b, tmp_path):
    sp.json_save(a / '人工判定資料庫.json', sp.normalize_db({}))
    first = _filled_expected_excel(a, tmp_path / 'first.xlsx')
    second = edited_excel(first, tmp_path / 'second.xlsx', evidence='OTHER ORIGINAL RULE')
    sp.import_gpt_decisions(b, first)
    with pytest.raises(ValueError, match='不同判定'): sp.import_gpt_decisions(b, second)
    return first, second


@pytest.mark.parametrize('lane', ['project', 'expected_first', 'expected_second'])
def test_exact_historical_reimport_preserves_local_adjudication(pair, tmp_path, lane):
    a, b, bp, bm = pair; rid = bm['records'][0]['review_id']
    if lane == 'project':
        db = sp.json_load_strict(b / '人工判定資料庫.json')
        db['events'][rid] = {'action': '確認非校對範圍', 'exclusion_reason': 'B decision', 'exclusion_evidence': 'page 1'}
        sp.json_save(b / '人工判定資料庫.json', db)
        assert p.import_project_decisions(a, b)['conflicts']
        operation = lambda: p.import_project_decisions(a, b)
    else:
        first, second = excel_pair(a, b, tmp_path)
        operation = lambda: sp.import_gpt_decisions(b, first if lane == 'expected_first' else second)
    adjudicate(b, rid)
    before = state(b); event = copy.deepcopy(sp.json_load_strict(b / '人工判定資料庫.json')['events'][rid])
    operation()
    assert state(b) == before
    assert sp.json_load_strict(b / '人工判定資料庫.json')['events'][rid] == event
    assert len(sp.json_load_strict(b / p.CONFLICT_FILE)['conflicts']) == 1
    assert p.validate_conflict_state(b, bm, sp.json_load_strict(b / '人工判定資料庫.json')) == []


@pytest.mark.parametrize('lane', ['project', 'expected'])
@pytest.mark.parametrize('changed_source', ['same', 'new'])
def test_real_source_change_after_adjudication_is_still_conflict(pair, tmp_path, lane, changed_source):
    a, b, bp, bm = pair; rid = bm['records'][0]['review_id']
    if lane == 'project':
        adjudicate(b, rid, 'INITIAL B')
        assert p.import_project_decisions(a, b)['conflicts']
    else:
        first, _ = excel_pair(a, b, tmp_path)
    adjudicate(b, rid)
    original_receipt = sp.json_load_strict(b / p.CONFLICT_FILE)['conflicts']
    source = a
    if changed_source == 'new':
        cp = tmp_path / 'c.pdf'; pdf(cp, title='C')
        source = tmp_path / 'c'; project(source, cp, session='C', event_positions=(0,) if lane == 'project' else ())
    if lane == 'project':
        source_db = sp.json_load_strict(source / '人工判定資料庫.json')
        next(iter(source_db['events'].values()))['expected_evidence'] = 'NEW SOURCE RULE'
        sp.json_save(source / '人工判定資料庫.json', source_db)
        assert p.import_project_decisions(source, b)['conflicts']
    else:
        new = (edited_excel(first, tmp_path / 'third.xlsx', evidence='NEW SOURCE RULE')
               if changed_source == 'same' else _filled_expected_excel(source, tmp_path / 'third.xlsx'))
        with pytest.raises(ValueError, match='不同判定'): sp.import_gpt_decisions(b, new)
    current = sp.json_load_strict(b / p.CONFLICT_FILE)['conflicts']
    assert len(current) == 2 and current[0] == original_receipt[0]
    assert rid not in sp.json_load_strict(b / '人工判定資料庫.json')['events']


@pytest.mark.parametrize('lane', ['project', 'expected'])
def test_retained_duplicate_source_replay_preserves_adjudication(pair, tmp_path, lane):
    a, b, bp, bm = pair; rid = bm['records'][0]['review_id']
    if lane == 'project':
        p.import_project_decisions(a, b)
        cp = tmp_path / 'c.pdf'; pdf(cp, title='C')
        c = tmp_path / 'c'; project(c, cp, session='C', event_positions=(0,))
        assert p.import_project_decisions(c, b)['provenance_updates'] == 1
        replay = lambda: p.import_project_decisions(c, b)
        ad = sp.json_load_strict(a / '人工判定資料庫.json')
        next(iter(ad['events'].values()))['expected_evidence'] = 'DIFFERENT RULE'
        sp.json_save(a / '人工判定資料庫.json', ad)
        assert p.import_project_decisions(a, b)['conflicts']
    else:
        sp.json_save(a / '人工判定資料庫.json', sp.normalize_db({}))
        first = _filled_expected_excel(a, tmp_path / 'first.xlsx')
        duplicate = edited_excel(first, tmp_path / 'duplicate.xlsx', title='independent file copy')
        different = edited_excel(first, tmp_path / 'different.xlsx', evidence='DIFFERENT RULE')
        sp.import_gpt_decisions(b, first)
        sp.import_gpt_decisions(b, duplicate)
        replay = lambda: sp.import_gpt_decisions(b, duplicate)
        with pytest.raises(ValueError, match='不同判定'): sp.import_gpt_decisions(b, different)
    receipt = sp.json_load_strict(b / p.CONFLICT_FILE)
    assert len(receipt['conflicts'][0]['target_event']['portability_duplicate_sources']) == 1
    adjudicate(b, rid)
    before = state(b)
    replay()
    assert state(b) == before


def test_new_matching_source_keeps_provenance_then_replays_from_history(pair, tmp_path):
    a, b, bp, bm = pair; rid = bm['records'][0]['review_id']
    adjudicate(b, rid, 'INITIAL B')
    p.import_project_decisions(a, b)
    adjudicate(b, rid)
    cp = tmp_path / 'c.pdf'; pdf(cp, title='C')
    c = tmp_path / 'c'; cm, cd = project(c, cp, session='C')
    local = sp.json_load_strict(b / '人工判定資料庫.json')['events'][rid]
    cd['events'][cm['records'][0]['review_id']] = p._event_payload(local)
    sp.json_save(c / '人工判定資料庫.json', cd)
    assert p.import_project_decisions(c, b)['provenance_updates'] == 1
    assert sp.json_load_strict(b / '人工判定資料庫.json')['events'][rid]['portability_conflict_resolution'] == local['portability_conflict_resolution']
    ad = sp.json_load_strict(a / '人工判定資料庫.json')
    next(iter(ad['events'].values()))['expected_evidence'] = 'CHANGED AGAIN'
    sp.json_save(a / '人工判定資料庫.json', ad)
    assert p.import_project_decisions(a, b)['conflicts']
    adjudicate(b, rid, 'SECOND LOCAL')
    before = state(b)
    p.import_project_decisions(c, b)
    assert state(b) == before


@pytest.mark.parametrize('moment', ['after_plan', 'after_receipt', 'foreign_after_receipt'])
def test_cross_expected_binds_pdf_at_each_final_write(pair, tmp_path, monkeypatch, moment):
    a, b, bp, bm = pair
    sp.json_save(a / '人工判定資料庫.json', sp.normalize_db({}))
    x = _filled_expected_excel(a, tmp_path / 'expected.xlsx')
    if moment != 'after_plan': adjudicate(b, bm['records'][0]['review_id'], 'B original')
    report = b / '注音校對_最終報告.xlsx'; report.write_bytes(b'ORIGINAL REPORT')
    before = state(b)
    changed = tmp_path / 'changed.pdf'; pdf(changed, title='Changed', text='different whole textbook')
    original_plan, original_save = p._new_presentation_marker, sp.json_save
    foreign_marker = b'{"owner":"FOREIGN"}'; foreign_receipt = b'{"owner":"FOREIGN RECEIPT"}'
    injected = False
    def swap():
        nonlocal injected
        injected = True; bp.write_bytes(changed.read_bytes())
        if moment == 'foreign_after_receipt':
            (b / p.INCOMPLETE_FILE).write_bytes(foreign_marker)
            (b / p.CONFLICT_FILE).write_bytes(foreign_receipt)
    def plan(*args, **kwargs):
        result = original_plan(*args, **kwargs)
        if moment == 'after_plan': swap()
        return result
    def save(path, value, **kwargs):
        result = original_save(path, value, **kwargs)
        if Path(path) == b / p.CONFLICT_FILE and moment != 'after_plan': swap()
        return result
    monkeypatch.setattr(p, '_new_presentation_marker', plan)
    monkeypatch.setattr(sp, 'json_save', save)
    with patch.object(p, '_pdf_content_snapshot', wraps=p._pdf_content_snapshot) as snapshots:
        with pytest.raises(ValueError): sp.import_gpt_decisions(b, x)
    assert injected and p._sha(bp) != bm['pdfs'][0]['pdf_sha256']
    assert len([call for call in snapshots.call_args_list if Path(call.args[0]) == bp]) == 1
    if moment == 'foreign_after_receipt':
        before[p.CONFLICT_FILE] = foreign_receipt
        assert (b / p.INCOMPLETE_FILE).read_bytes() == foreign_marker
    else:
        assert not (b / p.INCOMPLETE_FILE).exists()
    assert state(b) == before
