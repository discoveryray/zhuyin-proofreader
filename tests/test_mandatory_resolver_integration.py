from __future__ import annotations

import hashlib
import sys
import tempfile
import unittest
from pathlib import Path

import fitz
from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from check_pronunciation_candidates import (  # noqa: E402
    DEFAULT_DICT,
    DEFAULT_RULES,
    analyze,
)
from export_zhuyin_readings import (  # noqa: E402
    ACTUAL_DECODER_SOURCE_FILES,
    VERSION as DECODER_VERSION,
)
from occurrence_ledger import (  # noqa: E402
    LEDGER_SCHEMA_VERSION,
    WORKBOOK_SCHEMA_VERSION,
    prepare_occurrence_rows,
)
from runtime_source_validation import (  # noqa: E402
    compute_actual_asset_fingerprint,
    validate_asset_manifest,
)
from standalone_proofread import (  # noqa: E402
    actual_workbook_path,
    candidate_workbook_path,
    collect_manifest,
    validate_manifest_integrity,
    validate_output_artifact_hashes,
)


ACTUAL_COLUMNS = [
    "實體頁碼", "課本頁", "頁面標籤", "字元", "實際注音", "解碼狀態", "解碼依據",
    "穩定注音鍵", "群組注音鍵", "font", "font_xref", "注音元件ID",
    "glyph_id_字形索引", "x0", "y0", "x1", "y1",
    "occurrence_id", "review_id", "identity_confidence", "identity_row_fallback",
    "identity_collision_base", "source_row_number", "ledger_schema_version", "workbook_schema_version",
]

EXCLUDED_COLUMNS = [
    "實體頁碼", "課本頁", "字元", "font", "font_xref", "glyph_id", "候選注音元件ID",
    "x0", "y0", "來源", "排除理由", "occurrence_id", "review_id", "identity_confidence",
    "identity_row_fallback", "source_row_number", "ledger_schema_version", "workbook_schema_version",
]


class MandatoryResolverIntegrationTests(unittest.TestCase):
    def test_all_mandatory_cases_reach_the_frozen_resolver(self):
        source_report = validate_asset_manifest(ROOT)
        self.assertTrue(source_report["ok"], source_report.get("errors"))

        with tempfile.TemporaryDirectory() as directory:
            temp_root = Path(directory)
            # Use the 3上 filename scope that previously loaded the over-broad
            # V38 「轉動」 exception and made M25-ZHUAN-DONG fail.
            pdf_path = temp_root / "07-115國小健體3上課本-L06-0225.pdf"
            document = fitz.open()
            document.new_page()
            document.save(pdf_path)
            document.close()

            pdf_hash = hashlib.sha256(pdf_path.read_bytes()).hexdigest()
            fingerprint = compute_actual_asset_fingerprint(
                ROOT,
                pdf_path,
                source_report,
                decoder_version=DECODER_VERSION,
                source_files=ACTUAL_DECODER_SOURCE_FILES,
            )

            source = {
                "實體頁碼": 1,
                "課本頁": 1,
                "頁面標籤": "1",
                "字元": "角",
                "實際注音": "ㄐㄩㄝˊ",
                "解碼狀態": "已映射",
                "解碼依據": "synthetic actual evidence",
                "穩定注音鍵": "fixture-key",
                "群組注音鍵": "fixture-group",
                "font": "fixture-font",
                "font_xref": 1,
                "注音元件ID": "fixture-component",
                "glyph_id_字形索引": 1,
                "x0": 10,
                "y0": 10,
                "x1": 20,
                "y1": 20,
                "source_row_number": 2,
                "ledger_schema_version": LEDGER_SCHEMA_VERSION,
                "workbook_schema_version": WORKBOOK_SCHEMA_VERSION,
            }
            prepare_occurrence_rows(pdf_hash, [source])

            actual_dir = temp_root / "01_實際注音"
            candidate_dir = temp_root / "02_候選報告"
            actual_dir.mkdir()
            candidate_dir.mkdir()
            actual_path = actual_workbook_path(actual_dir, pdf_path)
            workbook = Workbook()
            actual_sheet = workbook.active
            actual_sheet.title = "實際注音"
            actual_sheet.append(ACTUAL_COLUMNS)
            actual_sheet.append([source.get(column, "") for column in ACTUAL_COLUMNS])
            excluded_sheet = workbook.create_sheet("結構偵測排除")
            excluded_sheet.append(EXCLUDED_COLUMNS)
            metadata_sheet = workbook.create_sheet("v5.2中繼資料")
            metadata_sheet.append(["項目", "內容"])
            for key, value in (
                ("workbook_schema_version", WORKBOOK_SCHEMA_VERSION),
                ("ledger_schema_version", LEDGER_SCHEMA_VERSION),
                ("pdf_sha256", pdf_hash),
                ("actual_asset_fingerprint", fingerprint["fingerprint"]),
            ):
                metadata_sheet.append([key, value])
            summary = workbook.create_sheet("摘要")
            summary.append(["項目", "內容"])
            summary.append(["偵測到注音字形筆數", 1])
            workbook.save(actual_path)

            output_path = candidate_workbook_path(candidate_dir, pdf_path)
            result = analyze(actual_path, DEFAULT_DICT, DEFAULT_RULES, output_path, pdf_path=pdf_path)
            report = result["mandatory_regression"]

            self.assertGreater(report["required"], 0)
            self.assertEqual(report["required"], 36)
            self.assertEqual(report["executed"], report["required"])
            self.assertEqual(report["not_executed"], 0)
            self.assertEqual(report["duplicate_case_id"], 0)
            self.assertEqual(report["executed_case_ids"], report["required_case_ids"])
            self.assertEqual(report["failed"], 0)
            self.assertTrue(report["ok"])

            manifest = collect_manifest(
                [pdf_path],
                actual_dir,
                candidate_dir,
                actual_fingerprints={pdf_path.name: fingerprint},
                source_validation=source_report,
            )
            validate_manifest_integrity(manifest)
            validate_output_artifact_hashes(manifest)
            self.assertEqual(len(manifest["records"]), 1)
            self.assertEqual(manifest["records"][0]["actual"], "ㄐㄩㄝˊ")

            tampered_actual = load_workbook(actual_path)
            actual_sheet = tampered_actual["實際注音"]
            actual_headers = {str(cell.value or ""): cell.column for cell in actual_sheet[1]}
            actual_sheet.cell(2, actual_headers["實際注音"], "ㄐㄧㄠˇ")
            tampered_actual.save(actual_path)
            tampered_actual.close()
            with self.assertRaises(ValueError):
                validate_output_artifact_hashes(manifest)
            with self.assertRaises(ValueError):
                collect_manifest(
                    [pdf_path],
                    actual_dir,
                    candidate_dir,
                    actual_fingerprints={pdf_path.name: fingerprint},
                    source_validation=source_report,
                )
            restored_actual = load_workbook(actual_path)
            restored_actual["實際注音"].cell(2, actual_headers["實際注音"], "ㄐㄩㄝˊ")
            restored_actual.save(actual_path)
            restored_actual.close()

            tampered = load_workbook(output_path)
            ledger_sheet = tampered["Occurrence Ledger"]
            ledger_headers = {str(cell.value or ""): cell.column for cell in ledger_sheet[1]}
            ledger_sheet.cell(2, ledger_headers["expected_evidence"], "tampered evidence")
            tampered.save(output_path)
            tampered.close()
            with self.assertRaises(ValueError):
                collect_manifest(
                    [pdf_path],
                    actual_dir,
                    candidate_dir,
                    actual_fingerprints={pdf_path.name: fingerprint},
                    source_validation=source_report,
                )


if __name__ == "__main__":
    unittest.main()
