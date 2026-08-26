from __future__ import annotations

import csv
import tempfile
from pathlib import Path

import fitz

from actual_review import (
    GLYPH_CONFLICT_FILE,
    GLYPH_PROVENANCE_FILE,
    OCCURRENCE_OVERRIDE_FILE,
    USER_GLYF_FILE,
    apply_verified_actual_group,
    build_actual_review_groups,
    dynamic_actual_hashes,
    ensure_user_evidence_files,
    load_glyph_truth_quarantine,
    load_user_verified_glyf,
)

SHA = "5adfc5e3" + "0" * 56


def make_pdf(path: Path):
    doc = fitz.open()
    page = doc.new_page(width=300, height=300)
    page.insert_text((60, 100), "ABCD", fontsize=24)
    doc.save(path)
    doc.close()


def entry(oid: str, rid: str, pdf: Path, x: float, *, state="ACTUAL_DECODE_ERROR"):
    return {
        "occurrence_id": oid,
        "review_id": rid,
        "state": state,
        "pdf": str(pdf),
        "pdf_name": pdf.name,
        "physical_page": 1,
        "printed_page": "1",
        "char": "樂",
        "stable_key": f"key-{oid}",
        "x0": x,
        "y0": 70,
        "x1": x + 20,
        "y1": 110,
        "actual": "",
        "actual_evidence": "decoder unresolved",
        "source_record": {
            "TTF字形SHA256": SHA,
            "穩定注音鍵": f"key-{oid}",
        },
    }


def read_csv(path: Path):
    with path.open(encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def test_conflicting_global_truth_does_not_block_occurrence_local_correction():
    with tempfile.TemporaryDirectory() as td:
        root = Path(td)
        ensure_user_evidence_files(root)
        pdf = root / "book.pdf"
        make_pdf(pdf)
        members = [
            entry("o1", "r1", pdf, 40),
            entry("o2", "r2", pdf, 80),
            entry("o3", "r3", pdf, 120),
            entry("o4", "r4", pdf, 160),
        ]
        group = build_actual_review_groups(members)[0]

        first = apply_verified_actual_group(
            root, group, "ㄩㄝˋ", checked_occurrence_ids=["o1", "o2"],
            source="old visual review", note="old truth",
        )
        assert first["learning_level"] == "VERIFIED_EXACT_GLYPH"
        assert first["propagated_to_group"] is True
        assert len(read_csv(root / OCCURRENCE_OVERRIDE_FILE)) == 4

        second = apply_verified_actual_group(
            root, group, "ㄌㄜˋ", checked_occurrence_ids=["o1", "o2"],
            source="new direct visual review", note="new evidence",
        )
        assert second["learning_level"] == "GLYPH_TRUTH_CONFLICT"
        assert second["glyph_truth_conflict"] is True
        assert second["propagated_to_group"] is False
        assert set(second["reopened_occurrence_ids"]) == {"o3", "o4"}
        assert set(second["affected_occurrence_ids"]) == {"o1", "o2", "o3", "o4"}

        # Only the newly/directly checked positions remain occurrence-pinned.
        overrides = read_csv(root / OCCURRENCE_OVERRIDE_FILE)
        assert len(overrides) == 2
        assert {row["actual_reading"] for row in overrides} == {"ㄌㄜˋ"}

        # Old global truth is retained for audit but no longer loadable as reusable truth.
        learning = read_csv(root / USER_GLYF_FILE)
        assert learning[0]["verification_level"] == "QUARANTINED_CONFLICT"
        assert load_user_verified_glyf(root / USER_GLYF_FILE) == {}

        conflicts = read_csv(root / GLYPH_CONFLICT_FILE)
        assert len(conflicts) == 1
        assert conflicts[0]["status"] == "GLYPH_TRUTH_CONFLICT"
        assert set(conflicts[0]["readings"].split("|")) == {"ㄩㄝˋ", "ㄌㄜˋ"}
        assert SHA in load_glyph_truth_quarantine(root)["ttf"]
        assert read_csv(root / GLYPH_PROVENANCE_FILE)


def test_conflict_registry_is_part_of_per_pdf_dynamic_fingerprint_scope():
    with tempfile.TemporaryDirectory() as td:
        root = Path(td)
        ensure_user_evidence_files(root)
        pdf = root / "book.pdf"
        make_pdf(pdf)
        deps_hit = {"ttf_glyph_sha256": [SHA], "cff_glyph_keys": []}
        deps_miss = {"ttf_glyph_sha256": ["f" * 64], "cff_glyph_keys": []}
        before_hit = dynamic_actual_hashes(root, pdf_path=pdf, dependencies=deps_hit)
        before_miss = dynamic_actual_hashes(root, pdf_path=pdf, dependencies=deps_miss)

        (root / GLYPH_CONFLICT_FILE).write_text(
            "kind,style_group,glyph_sha256,status,readings,source_occurrence_ids,updated_at,notes\n"
            f"TTF_GLYF_SHA256,,{SHA},GLYPH_TRUTH_CONFLICT,ㄩㄝˋ|ㄌㄜˋ,o1|o2,2026-08-24T10:00:00,test\n",
            encoding="utf-8-sig",
        )
        after_hit = dynamic_actual_hashes(root, pdf_path=pdf, dependencies=deps_hit)
        after_miss = dynamic_actual_hashes(root, pdf_path=pdf, dependencies=deps_miss)
        assert before_hit[GLYPH_CONFLICT_FILE] != after_hit[GLYPH_CONFLICT_FILE]
        assert before_miss[GLYPH_CONFLICT_FILE] == after_miss[GLYPH_CONFLICT_FILE]


def test_gpt_import_conflict_is_not_batch_fatal_and_local_override_commits():
    from openpyxl import load_workbook
    from actual_review import export_actual_review_package, import_actual_review_workbook

    with tempfile.TemporaryDirectory() as td:
        root = Path(td)
        ensure_user_evidence_files(root)
        pdf = root / "book.pdf"
        make_pdf(pdf)
        members = [entry("o1", "r1", pdf, 40), entry("o2", "r2", pdf, 80), entry("o3", "r3", pdf, 120)]
        group = build_actual_review_groups(members)[0]
        # Establish the old reusable truth first, including propagated overrides.
        apply_verified_actual_group(root, group, "ㄩㄝˋ", checked_occurrence_ids=["o1", "o2"], source="old", note="old")

        meta = {
            "version": "5.5.1",
            "session_id": "sess",
            "session_schema_version": "2.6.0",
            "workbook_schema_version": "2.6.0",
            "review_id_schema_version": "2.5.0",
        }
        export_actual_review_package(root, members, **meta)
        xlsx = root / "actual待判定_GPT包" / "actual待判定_給GPT.xlsx"
        wb = load_workbook(xlsx)
        ws = wb["actual待判定"]
        headers = [c.value for c in ws[1]]
        ix = {h: i + 1 for i, h in enumerate(headers)}
        ws.cell(2, ix["decision"], "VERIFIED")
        ws.cell(2, ix["actual_reading"], "ㄌㄜˋ")
        ws.cell(2, ix["confidence"], "高")
        ws.cell(2, ix["sample_a_checked"], "Y")
        ws.cell(2, ix["sample_b_checked"], "Y")
        wb.save(xlsx)
        wb.close()

        result = import_actual_review_workbook(root, xlsx, [group], expected_metadata=meta)
        assert result["imported_groups"] == 1
        item = result["results"][0]
        assert item["learning_level"] == "GLYPH_TRUTH_CONFLICT"
        assert set(item["reopened_occurrence_ids"]) == {"o3"}
        overrides = read_csv(root / OCCURRENCE_OVERRIDE_FILE)
        assert len(overrides) == 2
        assert {row["actual_reading"] for row in overrides} == {"ㄌㄜˋ"}
