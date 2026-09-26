from __future__ import annotations

import copy
from contextlib import closing
import hashlib
import json
import tempfile
import threading
import tkinter as tk
import tkinter.font as tkfont
import unittest
from pathlib import Path
from unittest.mock import patch

import fitz
from openpyxl import Workbook, load_workbook

import occurrence_ledger as ol
import review_gui as gui
import standalone_proofread as sp
from tests.review_save_test_support import wait_for_save

ROOT = Path(__file__).resolve().parents[1]
REGRESSION = {"ok": True, "required": 1, "executed": 1, "passed": 1, "failed": 0,
              "not_executed": 0, "duplicate_case_id": 0}


def make_entry(n=1, *, actual="ㄎㄢˋ", expected=(), blocked="", excluded=False, pdf="fixture.pdf", pdf_sha="f" * 64):
    source = {"pdf_sha256": pdf_sha, "pdf": str(pdf), "pdf_name": Path(pdf).name,
              "實體頁碼": 1, "課本頁": "1", "字元": "看", "穩定注音鍵": f"fixture-{n}",
              "font": "fixture", "font_xref": 1, "glyph_id_字形索引": n,
              "x0": 72.0, "y0": float(72 + n * 45), "x1": 96.0, "y1": float(96 + n * 45),
              "source_row_number": n, "所在行": f"看一看窗外的風景（隔離操作樣本 {n}）", "局部詞境": "看一看"}
    ol.prepare_occurrence_rows(pdf_sha, [source])
    classified = {"actual": actual, "actual_evidence": "isolated fixture reading" if actual else "",
                  "actual_status": "RESOLVED" if actual else "UNRESOLVED",
                  "expected_set": list(expected), "expected_status": "RESOLVED" if expected else "UNRESOLVED",
                  "expected_evidence": "isolated expected fixture" if expected else "", "context_evidence": "看一看",
                  "state": blocked or ("EXCLUDED_OUT_OF_SCOPE" if excluded else "REVIEW_PENDING"),
                  "exclusion_reason": "isolated out-of-scope fixture" if excluded else ""}
    classified["state"] = ol.derive_authoritative_state(classified)
    return ol.build_occurrence_ledger([source], {source["occurrence_id"]: classified})[0]


def make_manifest(entries):
    validation = sp.validate_asset_manifest(ROOT)
    fp = sp.compute_expected_asset_fingerprint(ROOT, validation, resolver_version=sp.EXPECTED_RESOLVER_VERSION,
                                              source_files=sp.EXPECTED_RESOLVER_SOURCE_FILES)
    return sp.seal_manifest({
        "version": sp.VERSION, "session_schema_version": ol.SESSION_SCHEMA_VERSION,
        "workbook_schema_version": ol.WORKBOOK_SCHEMA_VERSION, "ledger_schema_version": ol.LEDGER_SCHEMA_VERSION,
        "review_id_schema_version": ol.REVIEW_ID_SCHEMA_VERSION, "session_id": "isolated-manual-usability",
        "records": entries, "actual_source_ids": [e["occurrence_id"] for e in entries], "pdfs": [],
        "expected_asset_fingerprint": fp["fingerprint"], "expected_asset_fingerprint_components": fp["components"],
        "regression_gate": REGRESSION, "mandatory_regression": REGRESSION,
        "pdf_regression": {"failed": 0, "not_executed": 0},
    })


def create_gui_fixture(output_dir: Path, *, mixed=True):
    """Isolated synthetic PDF/session for tests and actual Windows GUI evidence.

    The source reading and regression rows are explicit fixture inputs, not a
    claim that a real textbook or decoder acceptance run passed.
    """
    output_dir.mkdir(parents=True, exist_ok=True)
    pdf = output_dir / "隔離操作樣本.pdf"
    with fitz.open() as doc:
        page = doc.new_page(width=595, height=842)
        page.insert_text((40, 45), "人工確認操作隔離樣本（非使用者教材）", fontname="china-t", fontsize=16)
        for n in range(1, 8):
            page.insert_text((72, 91 + n * 45), f"看一看窗外的風景  樣本 {n}   ㄎㄢˋ", fontname="china-t", fontsize=15)
        doc.save(pdf)
    sha = hashlib.sha256(pdf.read_bytes()).hexdigest()
    entries = [make_entry(1, pdf=pdf, pdf_sha=sha), make_entry(2, pdf=pdf, pdf_sha=sha)]
    if mixed:
        entries.extend([
            make_entry(3, actual="", pdf=pdf, pdf_sha=sha),
            make_entry(4, actual="", expected=["ㄎㄢˋ"], pdf=pdf, pdf_sha=sha),
            make_entry(5, expected=["ㄎㄢ"], pdf=pdf, pdf_sha=sha),
            make_entry(6, blocked="DATA_INTEGRITY_ERROR", pdf=pdf, pdf_sha=sha),
            make_entry(7, excluded=True, pdf=pdf, pdf_sha=sha),
        ])
    manifest = make_manifest(entries)
    sp.json_save(output_dir / "校對工作階段.json", manifest)
    sp.json_save(output_dir / "人工判定資料庫.json", sp.normalize_db({}))
    return manifest


class ManualExpectedContractTests(unittest.TestCase):
    def test_no_operation_no_event_and_shortcut_freezes_one_occurrence(self):
        first, second = make_entry(), make_entry(2)
        manifest = make_manifest([first, second])
        self.assertEqual(sp.materialize_ledger(manifest, sp.normalize_db({})), [first, second])
        event = sp.build_manual_expected_event(first, operation="CONFIRM_CURRENT_AS_EXPECTED")
        resolved = sp.materialize_ledger(manifest, sp.normalize_db({**sp.normalize_db({}), "events": {first["review_id"]: event}}))
        self.assertEqual(resolved[0]["state"], "PASS")
        self.assertEqual(resolved[1], second)
        self.assertEqual(event["expected_evidence"], "")
        self.assertTrue(ol.valid_manual_expected_decision(resolved[0]))
        drifted = dict(first, actual="ㄎㄢ", actual_evidence="new fixture actual")
        reviewed = sp._apply_review_event(drifted, event)
        self.assertEqual(reviewed["expected_set"], ["ㄎㄢˋ"])
        self.assertEqual(reviewed["state"], "DIFFERENCE_PENDING_CONFIRMATION")
        self.assertEqual(reviewed["manual_expected_decision"], event["manual_expected_decision"])

    def test_shortcut_rejects_unresolved_invalid_or_missing_actual_proof(self):
        for updates in ({"actual_status": "UNRESOLVED"}, {"actual_status": "DECODE_ERROR"},
                        {"actual": "bad"}, {"actual": ""}, {"actual_evidence": "  "},
                        {"state": "SOURCE_INVALID"}):
            with self.subTest(updates=updates), self.assertRaises(ol.InvalidTransitionError):
                sp.build_manual_expected_event(dict(make_entry(), **updates), operation="CONFIRM_CURRENT_AS_EXPECTED")

    def test_blank_rationale_reopen_replay_report_and_completion(self):
        for rationale in ("", " \n\t "):
            with self.subTest(rationale=rationale), tempfile.TemporaryDirectory() as td:
                output = Path(td)
                row = make_entry()
                manifest = make_manifest([row])
                sp.json_save(output / "校對工作階段.json", manifest)
                event = sp.build_manual_expected_event(row, operation="ENTER_EXPECTED", expected_set="ㄎㄢˋ", rationale=rationale)
                sp.json_save(output / "人工判定資料庫.json", sp.normalize_db({**sp.normalize_db({}), "events": {row["review_id"]: event}}))
                db = sp.load_or_initialize_db(output)
                ledger = sp.materialize_ledger(sp.json_load_strict(output / "校對工作階段.json"), db)
                ol.validate_occurrence_ledger(ledger)
                gate = ol.completion_gate(ledger, {"ok": True}, REGRESSION)
                self.assertEqual(gate["expected_covered"], 1)
                self.assertEqual(gate["status"], ol.PROOFREAD_COMPLETE)
                path = sp.generate_report(output, sp.json_load_strict(output / "校對工作階段.json"), sp.load_or_initialize_db(output))
                status = sp.json_load_strict(output / "pipeline_status.json")
                self.assertEqual(status["completion_gate"]["expected_covered"], 1)
                with closing(load_workbook(path, read_only=True)) as wb:
                    rows = list(wb["人工與公司規則"].values)
                    self.assertIsNone(rows[1][6])
                    self.assertEqual(rows[1][9], "人工輸入應標注音")
                    self.assertEqual(rows[1][10], event["manual_expected_decision"]["decided_at"])
                with closing(load_workbook(output / "注音校對_技術稽核.xlsx", read_only=True)) as wb:
                    headers, values = list(wb["Occurrence Ledger"].values)
                    record = json.loads(values[headers.index("本筆人工應標判定紀錄")])
                    self.assertEqual(record, event["manual_expected_decision"])

    def test_legacy_requires_evidence_and_malformed_new_record_never_upgrades(self):
        row = make_entry()
        legacy = {"action": "解決expected證據", "expected_set": ["ㄎㄢˋ"], "expected_evidence": "真實舊來源", "context_evidence": "看一看"}
        self.assertEqual(sp._apply_review_event(row, legacy)["state"], "PASS")
        for source in ("人工 GUI 現版 expected 證據", "external", ""):
            with self.subTest(source=source), self.assertRaises(ol.InvalidTransitionError):
                sp._apply_review_event(row, dict(legacy, expected_evidence="", source=source))
        good = sp.build_manual_expected_event(row, operation="ENTER_EXPECTED", expected_set="ㄎㄢˋ")
        corruptions = [lambda r: r.pop("decided_at"), lambda r: r.update(version=2),
                       lambda r: r.update(version=True), lambda r: r.update(operation="EXTERNAL_IMPORT"),
                       lambda r: r["target"].update(occurrence_id="foreign"),
                       lambda r: r["target"].update(line="different context"),
                       lambda r: r.update(expected_set=["ㄎㄢ"]),
                       lambda r: r.update(context_evidence="another context"),
                       lambda r: r.update(decided_at="2026-01-01"), lambda r: r.update(extra="unknown")]
        for mutate in corruptions:
            bad = copy.deepcopy(good)
            mutate(bad["manual_expected_decision"])
            bad["expected_evidence"] = "cannot disguise invalid record as legacy"
            with self.subTest(record=bad), self.assertRaises(ol.InvalidTransitionError):
                sp._apply_review_event(row, bad)

    def test_illegal_expected_mixed_with_valid_is_rejected(self):
        for reading in ("", "ㄎㄢˋ|invalid", ["ㄎㄢˋ", "bad"]):
            with self.subTest(reading=reading), self.assertRaises(ol.InvalidTransitionError):
                sp.build_manual_expected_event(make_entry(), operation="ENTER_EXPECTED", expected_set=reading)

    def test_numeric_workbook_roundtrip_preserves_target_and_changed_position_rejects(self):
        row = make_entry()
        event = sp.build_manual_expected_event(row, operation="ENTER_EXPECTED", expected_set="ㄎㄢˋ")
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "source.xlsx"
            wb = Workbook()
            wb.active.append(list(row["source_record"]))
            wb.active.append(list(row["source_record"].values()))
            wb.save(path)
            with closing(load_workbook(path, read_only=True)) as loaded:
                header, values = list(loaded.active.values)
                source = dict(zip(header, values))
            self.assertIsInstance(source["x0"], int)
            source["pdf_sha256"] = row["pdf_sha256"]
            ol.prepare_occurrence_rows(row["pdf_sha256"], [source])
            self.assertEqual(source["occurrence_id"], row["occurrence_id"])
            refreshed = copy.deepcopy(row)
            refreshed.update({k: source[k] for k in ("x0", "y0", "x1", "y1")})
            refreshed["source_record"] = source
            manifest = make_manifest([refreshed])
            sp.json_save(Path(td) / "manifest.json", manifest)
            replay = sp.materialize_ledger(sp.json_load_strict(Path(td) / "manifest.json"), sp.normalize_db({**sp.normalize_db({}), "events": {row["review_id"]: event}}))
            self.assertEqual(replay[0]["state"], "PASS")
            with self.assertRaises(ol.InvalidTransitionError):
                sp._apply_review_event(dict(refreshed, x0=73), event)

    def test_confirmation_then_actual_refresh_preserves_expected_and_revokes_comparison(self):
        row = make_entry()
        event = sp.build_manual_expected_event(row, operation="ENTER_EXPECTED", expected_set="ㄎㄢ")
        reviewed = sp._apply_review_event(row, event)
        confirmed_event = {**event, "action": "確認現版差異", **{gate: True for gate in ol.CONFIRMATION_GATES},
                           "confirmation_actual_snapshot": sp.actual_confirmation_snapshot(reviewed)}
        self.assertEqual(sp._apply_review_event(row, confirmed_event)["state"], "TEXTBOOK_ERROR_CONFIRMED")
        # Even another mismatching actual reading invalidates all old six gates.
        drifted = dict(row, actual="ㄎㄢˊ", actual_evidence="new source")
        reopened = sp._apply_review_event(drifted, confirmed_event)
        self.assertEqual(reopened["state"], "DIFFERENCE_PENDING_CONFIRMATION")
        self.assertEqual(reopened["review_event_replay_status"], "INVALIDATED_ACTUAL_DRIFT")
        with tempfile.TemporaryDirectory() as td:
            output = Path(td)
            sp.json_save(output / "人工判定資料庫.json", sp.normalize_db({**sp.normalize_db({}), "events": {row["review_id"]: confirmed_event}}))
            self.assertEqual(sp._clear_actual_dependent_events(output, [reviewed], {row["occurrence_id"]}), 1)
            saved = sp.load_or_initialize_db(output)["events"][row["review_id"]]
            self.assertEqual(saved["manual_expected_decision"], event["manual_expected_decision"])
            self.assertEqual(saved["expected_evidence"], "")
            self.assertNotIn("confirmation_actual_snapshot", saved)
            updated = sp._apply_review_event(dict(row, actual="ㄎㄢ"), saved)
            self.assertEqual(updated["state"], "PASS")

    def test_blocks_exclusion_and_missing_evidence_remain_fail_closed(self):
        event = sp.build_manual_expected_event(make_entry(), operation="ENTER_EXPECTED", expected_set="ㄎㄢˋ")
        for state in ol.HARD_BLOCKING_STATES:
            row = dict(make_entry(), state=state)
            self.assertEqual(sp._apply_review_event(row, event)["state"], state)
        excluded = dict(make_entry(), state="EXCLUDED_OUT_OF_SCOPE", exclusion_reason="fixture exclusion")
        self.assertEqual(sp._apply_review_event(excluded, event)["state"], "EXCLUDED_OUT_OF_SCOPE")
        invalid = dict(make_entry(), state="PASS", actual_status="RESOLVED", expected_status="RESOLVED", expected_set=["ㄎㄢˋ"])
        with self.assertRaises(ol.LedgerError):
            ol.validate_occurrence_ledger([invalid])
        self.assertFalse(ol.completion_gate([invalid], {"ok": True}, REGRESSION)["hard_gates"]["expected_coverage_100"])


class ManualReviewGuiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # This task requires a real Windows/Tk GUI. Missing capability is a test
        # error requiring the authorized desktop test environment, never a skip.
        cls.root = tk.Tk()
        cls.root.withdraw()

    @classmethod
    def tearDownClass(cls):
        cls.root.destroy()

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.output = Path(self.temp.name)
        self.manifest = create_gui_fixture(self.output)
        self.window = tk.Toplevel(self.root)
        self.app = gui.ReviewApp(self.window, self.output)
        self.window.update()

    def tearDown(self):
        self.window.destroy()
        self.temp.cleanup()

    def test_one_real_button_invoke_saves_without_dialog_and_duplicate_cannot_apply_next(self):
        app = self.app
        first = app.current()
        self.assertEqual(app.db["events"], {})
        with patch.object(gui, "ExpectedDialog") as dialog, patch.object(gui.messagebox, "askyesno") as confirm, patch.object(gui, "save_reusable_expected_rule") as reusable:
            app.primary.invoke()
            self.window.update()
            app.primary.invoke()
            wait_for_save(app)
            app.primary.invoke()
            self.window.update()
        self.assertEqual(len(app.db["events"]), 1)
        self.assertIn(first["review_id"], app.db["events"])
        self.assertNotEqual(app.current()["review_id"], first["review_id"])
        self.assertEqual(gui.review_lane(app.current()), "expected")
        dialog.assert_not_called()
        confirm.assert_not_called()
        reusable.assert_not_called()
        # The save shares the actual transaction lock, but creates no actual truth.
        actual_root = sp.project_actual_evidence_root(self.output)
        self.assertEqual({p.name for p in actual_root.iterdir()}, {".global_exact_glyph_delivery.lock"})

    def test_last_pending_confirmation_can_be_undone_after_it_leaves_queue(self):
        app = self.app
        first = app.current()
        other = next(row for row in app.records if row["review_id"] != first["review_id"])
        # Leave only this item in the visible queue; durable events still cover
        # the full manifest and the other occurrence remains untouched.
        app.deferred_items = {(gui.review_lane(row), row["occurrence_id"])
                              for row in app.records if row["review_id"] != first["review_id"]}
        app.reload_records()
        self.assertEqual(len(app.records), 1)
        app.primary.invoke()
        wait_for_save(app)
        self.assertEqual(app.records, [])
        self.assertIn(first["review_id"], app.db["events"])
        app._review_action_cooldown_until = 0
        with patch.object(gui.messagebox, "askyesno", return_value=True):
            app.undo_expected_button.invoke()
            wait_for_save(app)
        self.assertEqual(app.current()["review_id"], first["review_id"])
        self.assertNotIn(first["review_id"], app.db["events"])
        self.assertEqual(ol.infer_expected_status(app.current()), "UNRESOLVED")
        self.assertNotIn("manual_expected_decision", app.current())
        self.assertNotIn(other["review_id"], app.db["events"])
        self.assertIn(first["review_id"], [row["review_id"] for row in app.records])

    def test_correction_undo_restores_previous_manual_decision(self):
        app = self.app
        first = app.current()
        with patch.object(gui, "ExpectedDialog") as dialog, patch.object(gui.messagebox, "showwarning"):
            dialog.return_value.result = {"expected_set": "ㄎㄢ", "expected_evidence": "", "resolution_reason": ""}
            app.resolve_expected()
            wait_for_save(app)
        prior = copy.deepcopy(app.db["events"][first["review_id"]])
        prior.pop("undo_previous_event")
        app._review_action_cooldown_until = 0
        app._shortcut_cooldown_until = 0
        app.focused_entry = next(row for row in app.records if row["review_id"] == first["review_id"])
        app.show()
        self.assertEqual(app.primary.cget("text"), "確認目前注音就是應標注音")
        app.primary.invoke()
        wait_for_save(app)
        self.assertNotIn(first["review_id"], [row["review_id"] for row in app.records])
        self.assertEqual(app.db["events"][first["review_id"]]["undo_previous_event"], prior)
        app._review_action_cooldown_until = 0
        with patch.object(gui.messagebox, "askyesno", return_value=True):
            app.undo_last_expected()
            wait_for_save(app)
        self.assertEqual(app.current()["review_id"], first["review_id"])
        self.assertEqual(app.current()["expected_set"], ["ㄎㄢ"])
        self.assertEqual(app.current()["state"], "DIFFERENCE_PENDING_CONFIRMATION")
        self.assertEqual(app.db["events"][first["review_id"]], prior)

    def test_confirmed_lookup_reopens_with_filters_and_corrects_one_item(self):
        app = self.app
        first = app.current()
        app.primary.invoke()
        wait_for_save(app)
        reopened_window = tk.Toplevel(self.root)
        try:
            reopened = gui.ReviewApp(reopened_window, self.output)
            reopened_window.update()
            ledger = sp.materialize_ledger(reopened.manifest, reopened.db)
            confirmed = [row for row in ledger if row["review_id"] == first["review_id"]]
            dialog = gui.ConfirmedItemsDialog(reopened_window, confirmed, wait=False)
            dialog.pdf_filter.set("隔離操作樣本")
            dialog.page_filter.set("1")
            dialog.text_filter.set("窗外")
            dialog.refresh()
            self.assertEqual(dialog.table.get_children(), (first["review_id"],))
            dialog.text_filter.set("不存在")
            dialog.refresh()
            self.assertEqual(dialog.table.get_children(), ())
            dialog.destroy()
            with patch.object(gui, "ConfirmedItemsDialog") as choose:
                choose.return_value.result = first["review_id"]
                reopened.open_confirmed_items()
            self.assertEqual(reopened.current()["review_id"], first["review_id"])
            self.assertEqual(reopened.primary.cget("text"), "確認目前注音就是應標注音")
            with patch.object(gui, "ExpectedDialog") as correction, patch.object(gui.messagebox, "showwarning"):
                correction.return_value.result = {"expected_set": "ㄎㄢ", "expected_evidence": "", "resolution_reason": ""}
                reopened.resolve_expected()
                wait_for_save(reopened)
            self.assertEqual(reopened.current()["expected_set"], ["ㄎㄢ"])
            self.assertIn(first["review_id"], [row["review_id"] for row in reopened.records])
            self.assertEqual(sp.materialize_ledger(reopened.manifest, reopened.db)[1]["expected_status"], "UNRESOLVED")
            persisted = sp.load_or_initialize_db(self.output)
            self.assertEqual(persisted["events"][first["review_id"]]["expected_set"], ["ㄎㄢ"])
            reopened._review_action_cooldown_until = 0
            with patch.object(gui.messagebox, "askyesno", return_value=True):
                reopened.undo_last_expected()
                wait_for_save(reopened)
            self.assertEqual(reopened.current()["review_id"], first["review_id"])
            self.assertEqual(reopened.current()["expected_set"], ["ㄎㄢˋ"])
            self.assertEqual(reopened.current()["state"], "PASS")
            self.assertEqual(sp.load_or_initialize_db(self.output)["events"][first["review_id"]]["expected_set"], ["ㄎㄢˋ"])
            reopened.more_menu.invoke(reopened.more_menu.index("返回待辦"))
            self.assertIsNone(reopened.focused_entry)
            self.assertNotEqual(reopened.current()["review_id"], first["review_id"])
        finally:
            reopened_window.destroy()

    def test_failed_undo_keeps_saved_decision_and_does_not_show_success(self):
        app = self.app
        first = app.current()
        app.primary.invoke()
        wait_for_save(app)
        before = (self.output / "人工判定資料庫.json").read_bytes()
        old_current = app.current()
        app._review_action_cooldown_until = 0
        with patch.object(gui.messagebox, "askyesno", return_value=True), \
                patch.object(sp, "json_save", side_effect=OSError("isolated undo failure")), \
                patch.object(gui.messagebox, "showerror") as error, \
                patch.object(gui.messagebox, "showinfo") as success:
            app.undo_last_expected()
            wait_for_save(app)
        error.assert_called_once()
        success.assert_not_called()
        self.assertEqual(app.current(), old_current)
        self.assertEqual((self.output / "人工判定資料庫.json").read_bytes(), before)
        self.assertIn(first["review_id"], app.db["events"])

    def test_optional_blank_dialog_actual_unresolved_then_actual_group(self):
        app = self.app
        both = next(e for e in app.records if not e["actual"] and ol.infer_expected_status(e) == "UNRESOLVED")
        app.index = app.records.index(both)
        app.show()
        self.assertEqual(gui.review_lane(app.current()), "expected")
        self.assertEqual(app.primary.cget("text"), "輸入其他應標注音")
        dialog = gui.ExpectedDialog(self.window, both, "輸入其他應標注音", wait=False)
        dialog.expected.set("ㄎㄢˋ")
        dialog.evidence.set(" \t ")
        dialog.submit()
        with patch.object(gui, "ExpectedDialog") as cls:
            cls.return_value.result = dialog.result
            app.resolve_expected()
            wait_for_save(app)
        self.assertEqual(gui.review_lane(app.current()), "expected")
        new_both = next(e for e in app.records if e["review_id"] == both["review_id"])
        self.assertEqual(gui.review_lane(new_both), "actual")
        self.assertTrue(ol.valid_manual_expected_decision(new_both))
        app.deferred_items = {(gui.review_lane(e), e["occurrence_id"]) for e in app.records if gui.review_lane(e) == "expected"}
        app.reload_records()
        self.assertEqual(gui.review_lane(app.current()), "actual")

    def test_staged_actual_does_not_hide_expected_or_integrity_and_defer_is_finite(self):
        app = self.app
        both = next(e for e in app.records if not e["actual"] and ol.infer_expected_status(e) == "UNRESOLVED")
        actual_only = next(e for e in app.records if gui.review_lane(e) == "actual")
        sp.stage_manual_actual_correction(self.output, both["review_id"], "ㄎㄢˋ")
        sp.stage_manual_actual_correction(self.output, actual_only["review_id"], "ㄎㄢˋ")
        app.reload_records()
        self.assertIn(both["review_id"], [e["review_id"] for e in app.records])
        self.assertNotIn(actual_only["review_id"], [e["review_id"] for e in app.records])
        self.assertIn("DIFFERENCE_PENDING_CONFIRMATION", [e["state"] for e in app.records])
        self.assertIn("DATA_INTEGRITY_ERROR", [e["state"] for e in app.records])
        self.assertNotIn("EXCLUDED_OUT_OF_SCOPE", [e["state"] for e in app.records])
        app.reload_records()
        pending = len(app.records)
        for _ in range(pending):
            app.defer_current()
        self.assertEqual(app.records, [])
        self.assertEqual(len(app.deferred_items), pending)
        self.assertEqual(app.db["events"], {})
        app.defer_current()
        self.assertEqual(len(app.deferred_items), pending)
        self.assertIn("未視為完成", app._empty_actionable_state()[1])
        app.revisit_deferred()
        self.assertEqual(len(app.records), pending)
        self.assertEqual(gui.review_lane(app.current()), "expected")

    def test_failed_save_has_no_db_or_navigation_change_and_no_success(self):
        app = self.app
        first = app.current()
        before = (self.output / "人工判定資料庫.json").read_bytes()
        for failing in ("prepare_review_ledger", "json_save"):
            with self.subTest(failing=failing), patch.object(sp, failing, side_effect=OSError("isolated save failure")), patch.object(gui.messagebox, "showerror") as error, patch.object(gui.messagebox, "showinfo") as success:
                app.primary.invoke()
                wait_for_save(app)
                error.assert_called_once()
                success.assert_not_called()
            self.assertEqual(app.current(), first)
            self.assertEqual(app.db["events"], {})
            self.assertEqual((self.output / "人工判定資料庫.json").read_bytes(), before)

    def test_pdf_read_failure_disables_shortcut(self):
        app = self.app
        Path(app.current()["pdf"]).unlink()
        app.show()
        self.assertEqual(app.primary.cget("state"), "disabled")
        app.confirm_current_expected(app.current()["review_id"])
        self.assertEqual(app.db["events"], {})

    def test_ctrl_enter_dialog_repetition_and_busy_action_save_only_one_item(self):
        app = self.app
        first_id = app.current()["review_id"]
        original_dialog = gui.ExpectedDialog

        def keyboard_dialog(*args, **kwargs):
            dialog = original_dialog(*args, **kwargs, wait=False)
            dialog.expected.set("ㄎㄢˋ")
            dialog.update()
            dialog.focus_force()
            # Queue genuine Tk key events. Destroying the submitted dialog must
            # prevent the second event from becoming another item's decision.
            dialog.event_generate("<Control-Return>", when="tail")
            dialog.event_generate("<Control-Return>", when="tail")
            dialog.wait_window()
            return dialog

        with patch.object(gui, "ExpectedDialog", side_effect=keyboard_dialog) as dialog:
            app.resolve_expected()
            app.resolve_expected()
            wait_for_save(app)
        dialog.assert_called_once()
        self.assertEqual(set(app.db["events"]), {first_id})

    def test_saved_event_survives_native_show_failure_and_reopens(self):
        app = self.app
        first_id = app.current()["review_id"]
        with patch.object(app, "show", side_effect=RuntimeError("native publication failure")), \
                patch.object(gui.messagebox, "showwarning") as warning:
            app.primary.invoke()
            wait_for_save(app)
        warning.assert_called_once()
        self.assertIn("已保存", warning.call_args.args[0])
        self.assertTrue(app._last_event_saved)
        self.assertFalse(app._last_event_refreshed)
        self.assertEqual(app.primary.cget("state"), "disabled")
        reopened_window = tk.Toplevel(self.root)
        try:
            reopened = gui.ReviewApp(reopened_window, self.output)
            reopened_window.update()
            self.assertIn(first_id, reopened.db["events"])
            self.assertNotIn(first_id, [row["review_id"] for row in reopened.records])
        finally:
            reopened_window.destroy()

    def test_next_preview_failure_after_save_does_not_enable_fast_confirmation(self):
        app = self.app
        first_id = app.current()["review_id"]
        with patch.object(app.image, "load", return_value=False):
            app.primary.invoke()
            wait_for_save(app)
        self.assertTrue(app._last_event_saved)
        self.assertIsNone(app._rendered_review_id)
        self.assertEqual(app.primary.cget("state"), "disabled")
        app._shortcut_cooldown_until = app._review_action_cooldown_until = 0
        app.confirm_current_expected(app.current()["review_id"])
        self.assertEqual(set(app.db["events"]), {first_id})

    def test_preview_failure_during_failed_save_stays_disabled(self):
        app = self.app
        first = app.current()
        entered, release = threading.Event(), threading.Event()

        def failed_save(*_args, **_kwargs):
            entered.set()
            if not release.wait(5):
                raise RuntimeError("test did not release worker")
            raise OSError("isolated failed write")

        with patch.object(app._save_service, "save_event", side_effect=failed_save), \
                patch.object(gui.messagebox, "showerror") as error:
            app.primary.invoke()
            self.assertTrue(entered.wait(5))
            try:
                app._preview_failed()
            finally:
                release.set()
            wait_for_save(app)
        error.assert_called_once()
        self.assertEqual(app.current(), first)
        self.assertEqual(app.db["events"], {})
        self.assertIsNone(app._rendered_review_id)
        self.assertEqual(app.primary.cget("state"), "disabled")

    def test_shortcut_full_label_and_staged_batch_footer_remain_visible(self):
        app = self.app
        self.window.update()
        label_width = tkfont.Font(font=app.primary.cget("font")).measure(app.primary.cget("text"))
        self.assertGreaterEqual(app.primary.winfo_width(), label_width + 4)
        both = next(e for e in app.records if not e["actual"] and ol.infer_expected_status(e) == "UNRESOLVED")
        sp.stage_manual_actual_correction(self.output, both["review_id"], "ㄎㄢˋ")
        app.reload_records()
        app.index = next(n for n, e in enumerate(app.records) if e["review_id"] == both["review_id"])
        app.show()
        self.window.update()
        self.assertIn("已暫存", app.summary_text.cget("text"))
        bottom = self.window.winfo_rooty() + self.window.winfo_height()
        for button in (app.primary, app.later, app.apply_actual_button):
            self.assertTrue(button.winfo_viewable())
            self.assertGreater(button.winfo_height(), 1)
            self.assertLessEqual(button.winfo_rooty() + button.winfo_height(), bottom)

    def test_reusable_failure_after_saved_item_does_not_undo_judgment(self):
        app = self.app
        first_id = app.current()["review_id"]
        app.primary.invoke()
        wait_for_save(app)
        before = (self.output / "人工判定資料庫.json").read_bytes()
        with patch.object(gui.simpledialog, "askstring", side_effect=["看一看", ""]), patch.object(gui.messagebox, "showerror") as error, patch.object(gui, "save_reusable_expected_rule") as reusable:
            app.create_reusable_expected_rule()
            error.assert_called_once()
            reusable.assert_not_called()
        self.assertIn(first_id, app.db["events"])
        self.assertEqual((self.output / "人工判定資料庫.json").read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
