"""Operation-bound source reuse; no Tk, live DB, or inferred lane evidence."""
import copy
import hashlib
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest
from openpyxl import Workbook

import actual_review as ar
import pdf_portability as pp
import standalone_proofread as sp
from excel_text import append_text_row
from tests.test_import_transaction import make_manifest


def workbook(path, *, note='=literal', profile=False):
    wb = Workbook(); wb.remove(wb.active)
    meta = wb.create_sheet('匯入中繼資料')
    append_text_row(meta, ['session_id', 'source-session'])
    append_text_row(meta, ['actual_review_schema_version', '1.2' if profile else '1.1'])
    sheet = wb.create_sheet('actual待判定')
    headers = ['group_id', 'group_snapshot', 'decision', 'actual_reading', 'confidence',
               'sample_a_checked', 'sample_b_checked', 'note', 'sample_a_occurrence_id',
               'sample_b_occurrence_id', 'sample_a_image', 'sample_b_image']
    if profile:
        headers += ['sample_a_availability', 'sample_b_availability', 'source_availability_reason']
    append_text_row(sheet, headers)
    if not profile:
        append_text_row(sheet, ['g', 'snapshot', 'UNRESOLVED', '', '', 'N', 'N', note,
                               'occ-a', '', 'sample.png', ''])
    else:
        ar._write_actual_sample_profile(wb, {'schema': ar.ACTUAL_SAMPLE_PROFILE_SCHEMA, 'groups': []})
    wb.save(path); wb.close()
    return path


def test_snapshot_reuses_finite_modes_and_isolates_parsed_views(tmp_path, monkeypatch):
    path = workbook(tmp_path / 'source.xlsx', profile=True)
    opened = []; native = openpyxl.load_workbook
    def load(*args, **kwargs):
        value = native(*args, **kwargs); opened.append(value); return value
    monkeypatch.setattr(openpyxl, 'load_workbook', load)
    with pp._workbook_snapshot(path) as (snapshot, digest):
        for _ in range(3):
            meta = sp.workbook_metadata(snapshot, '匯入中繼資料'); meta['session_id'] = 'caller mutation'
            _, rows = ar._load_sheet_rows(snapshot, 'actual待判定'); rows.append({'fake': True})
            profile = ar.load_actual_sample_profile(snapshot); profile['groups'].append({'fake': True})
            assert sp.workbook_metadata(snapshot, '匯入中繼資料')['session_id'] == 'source-session'
            assert ar._load_sheet_rows(snapshot, 'actual待判定')[1] == []
            assert ar.load_actual_sample_profile(snapshot)['groups'] == []
            with pp._read_excel_workbook(snapshot, data_only=False) as formulas:
                assert formulas.read_only
                with pytest.raises(TypeError): formulas['actual待判定']['A2'] = 'write'
        assert len(opened) == 2
        assert all(wb._archive.fp is not None for wb in opened)
        pp._verify_excel_import_snapshot(snapshot, path, digest)
    assert all(wb._archive.fp is None for wb in opened)
    assert pp._excel_snapshot_reads.get() is None


def test_snapshot_exception_and_nested_scope_restore_and_close(tmp_path):
    path = workbook(tmp_path / 'source.xlsx')
    with pp._workbook_snapshot(path) as (outer, _):
        with pp._read_excel_workbook(outer) as parent:
            with pytest.raises(RuntimeError, match='abort'):
                with pp._workbook_snapshot(path) as (inner, _):
                    with pp._read_excel_workbook(inner) as child:
                        assert parent is not child
                        raise RuntimeError('abort')
            assert child._archive.fp is None
            assert parent._archive.fp is not None
            with pp._read_excel_workbook(outer) as resumed:
                assert resumed is parent
    assert parent._archive.fp is None
    assert pp._excel_snapshot_reads.get() is None
    assert not outer.exists() and not inner.exists()


def test_original_change_stays_private_until_final_sha_guard(tmp_path):
    path = workbook(tmp_path / 'source.xlsx')
    with pp._workbook_snapshot(path) as (snapshot, digest):
        assert sp.workbook_rows(snapshot, 'actual待判定')[0]['note'] == '=literal'
        workbook(path, note='changed')
        assert sp.workbook_rows(snapshot, 'actual待判定')[0]['note'] == '=literal'
        with pytest.raises(ValueError, match='匯入期間變動'):
            pp._verify_excel_import_snapshot(snapshot, path, digest)
    assert sp.workbook_rows(path, 'actual待判定')[0]['note'] == 'changed'
    with pp._workbook_snapshot(path) as (snapshot, _):
        assert sp.workbook_rows(snapshot, 'actual待判定')[0]['note'] == 'changed'


def test_private_snapshot_drift_rejects_cached_result_even_same_stat(tmp_path):
    path = workbook(tmp_path / 'source.xlsx')
    with pp._workbook_snapshot(path) as (snapshot, _):
        sp.workbook_rows(snapshot, 'actual待判定')
        before = snapshot.stat(); raw = bytearray(snapshot.read_bytes()); raw[-1] ^= 1
        snapshot.write_bytes(raw)
        import os
        os.utime(snapshot, ns=(before.st_atime_ns, before.st_mtime_ns))
        assert snapshot.stat().st_size == before.st_size
        with pytest.raises(ValueError, match='私有快照'):
            sp.workbook_rows(snapshot, 'actual待判定')


def test_uncached_live_reads_observe_changed_bytes_and_required_columns(tmp_path):
    path = workbook(tmp_path / 'source.xlsx')
    assert sp.workbook_rows(path, 'actual待判定')[0]['note'] == '=literal'
    workbook(path, note='live change')
    assert sp.workbook_rows(path, 'actual待判定')[0]['note'] == 'live change'
    with pp._workbook_snapshot(path) as (snapshot, _):
        sp.workbook_rows(snapshot, 'actual待判定', ['group_id'])
        with pytest.raises(ValueError, match='缺少欄位'):
            sp.workbook_rows(snapshot, 'actual待判定', ['absent'])


def test_proof_modes_and_decoded_views_do_not_share_mutable_or_formula_state(tmp_path):
    path = workbook(tmp_path / 'source.xlsx')
    wb = openpyxl.load_workbook(path); ws = wb.create_sheet(pp.EXCEL_PROOF_SHEET)
    ws.append(['formula', '=1+2']); wb.save(path); wb.close()
    with pp._workbook_snapshot(path) as (snapshot, _):
        with pp._read_excel_workbook(snapshot, data_only=False) as formulas:
            assert pp._excel_proof_rows(formulas)[0][1] == '=1+2'
            rows = pp._excel_proof_rows(formulas); rows.append(('caller', 'mutation'))
            assert len(pp._excel_proof_rows(formulas)) == 1
        with pp._read_excel_workbook(snapshot, data_only=True) as values:
            assert pp._excel_proof_rows(values)[0][1] is None
    with pytest.raises(ValueError, match='unknown Excel snapshot parser'):
        pp._memoized_excel_read(lambda path: {})


def test_actual_cheap_invalid_values_reject_before_profile_or_transaction(tmp_path, monkeypatch):
    path = workbook(tmp_path / 'source.xlsx')
    wb = openpyxl.load_workbook(path); wb['actual待判定']['C2'] = 'UNKNOWN'; wb.save(path); wb.close()
    def expensive(*args, **kwargs): raise AssertionError('expensive verification reached')
    monkeypatch.setattr(ar, 'validate_actual_sample_profile', expensive)
    monkeypatch.setattr(ar, 'apply_direct_visual_actual_batch', expensive)
    with pytest.raises(ValueError, match='decision 必須'):
        ar.import_actual_review_workbook(tmp_path / 'evidence', path, [], expected_metadata={})
    assert not (tmp_path / 'evidence').exists()


def test_expected_cheap_invalid_action_rejects_before_manifest_validation(tmp_path, monkeypatch):
    monkeypatch.setenv('LOCALAPPDATA', str(tmp_path / 'local'))
    sp.json_save(tmp_path / '校對工作階段.json', make_manifest())
    sp.json_save(tmp_path / '人工判定資料庫.json', sp.normalize_db({}))
    path = sp.export_pending_for_gpt(tmp_path)
    wb = openpyxl.load_workbook(path); ws = wb['待判定候選']; cols = {c.value: c.column for c in ws[1]}
    ws.cell(2, cols['action'], 'UNKNOWN'); wb.save(path); wb.close()
    before = (tmp_path / '人工判定資料庫.json').read_bytes()
    def expensive(*args, **kwargs): raise AssertionError('manifest verification reached')
    monkeypatch.setattr(sp, 'validate_manifest_integrity', expensive)
    with pytest.raises(ValueError, match='未知 action'):
        sp.import_gpt_decisions(tmp_path, path)
    assert (tmp_path / '人工判定資料庫.json').read_bytes() == before


def test_actual_missing_empty_header_is_rejected_without_profile(tmp_path):
    path = workbook(tmp_path / 'source.xlsx', profile=True)
    wb = openpyxl.load_workbook(path); ws = wb['actual待判定']; ws.delete_cols(8); wb.save(path); wb.close()
    with pytest.raises(ValueError, match='缺少欄位'):
        ar._screen_actual_import_rows(path)


def test_cached_carried_proof_still_compares_fresh_source_manifest(tmp_path, monkeypatch):
    import base64
    import json
    import zlib
    path = workbook(tmp_path / 'source.xlsx', profile=True)
    wb = openpyxl.load_workbook(path)
    source = {'session_id': 'source-session', 'seal': 'original independent seal'}
    payload = {'version': 1, 'kind': 'actual', 'manifest': source,
               'actual_profile_metadata': ar._actual_profile_metadata(wb),
               'bound_rows': pp._excel_bound_rows(wb, 'actual')}
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode('utf-8')
    sheet = wb.create_sheet(pp.EXCEL_PROOF_SHEET)
    append_text_row(sheet, ['version', 1]); append_text_row(sheet, ['sha256', hashlib.sha256(raw).hexdigest()])
    append_text_row(sheet, ['chunk_count', 1])
    append_text_row(sheet, [1, base64.b64encode(zlib.compress(raw)).decode('ascii')])
    wb.save(path); wb.close()
    calls = []; decode = pp._decode_excel_content_proof
    def observed(*args, **kwargs):
        calls.append(1); return decode(*args, **kwargs)
    monkeypatch.setattr(pp, '_decode_excel_content_proof', observed)
    with pp._workbook_snapshot(path) as (snapshot, _):
        with pp._read_excel_workbook(snapshot, data_only=False) as read:
            pp.validate_actual_profile_content_proof(read, source_manifest=source)
            pp.validate_actual_profile_content_proof(read, source_manifest=copy.deepcopy(source))
            with pytest.raises(ValueError, match='來源封印不符'):
                pp.validate_actual_profile_content_proof(read, source_manifest={**source, 'seal': 'changed'})
    assert len(calls) == 1


def test_formal_script_entry_help_uses_finite_main_alias_without_window():
    import subprocess
    import sys
    root = Path(__file__).resolve().parents[1]
    result = subprocess.run([sys.executable, str(root / 'standalone_proofread.py'), '--help'],
                            cwd=root, capture_output=True, text=True, encoding='utf-8', timeout=30)
    assert result.returncode == 0, result.stderr
    assert '--output-dir' in result.stdout
    assert 'unknown Excel snapshot parser' not in result.stderr
