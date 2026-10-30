"""Small independent actual fixtures; real sources, SQLite and save/recovery.

Only the ordinary decoder output is synthetic in project integration cases.
No expected answer supplies an actual reading, and no Tk is initialized.
"""
import copy
import hashlib
import json
from pathlib import Path

import fitz
import pytest
from openpyxl import Workbook

import actual_review as ar
import global_exact_glyph_library as lib
import global_glyph_promotion as promotion
import glyph_learning_service as learning
import occurrence_ledger as ol
import standalone_proofread as sp
from glyph_learning_images import LearningImageStore
from review_save_service import ReviewSaveService
from tests.actual_source_test_support import register_fixture_glyph, write_fixture_pdf
from tests.test_actual_visual_write_guards import rows_at, context_for
from tests.test_project_repair_integration_v562 import ACTUAL_COLUMNS, EXCLUDED_COLUMNS


def _hash(value):
    return hashlib.sha256(value.encode()).hexdigest()


def _isolated(tmp_path, monkeypatch):
    # Keep content-addressed filenames below Windows' ordinary MAX_PATH without
    # changing the product's paths or relaxing any filesystem check.
    root = tmp_path.parent / _hash(tmp_path.name)[:6]
    root.mkdir()
    monkeypatch.setenv("LOCALAPPDATA", str(root / "u"))
    return root


def _intent(label, *, identity=None, reading="ㄅ"):
    return lib.make_direct_evidence_intent(identity or {"kind": lib.TTF_GLYF_SHA256,
        "style_group": "", "glyph_sha256": "a" * 64,
        "identity_eligibility": lib.GLOBAL_ELIGIBLE_SIMPLE_GLYF_V1}, reading=reading,
        source_project_id="project-" + label, source_pdf_sha256=_hash("pdf-" + label),
        source_font_program_sha256=_hash("font-" + label), source_occurrence_id="occ_" + _hash(label),
        source_review_id="rev_" + _hash(label), decision_snapshot_sha256=_hash("decision-" + label))


def _decode(pdf, output, *args, global_snapshot=None):
    """Fixed fixture reading plus the actual override, with genuine source IDs."""
    from export_pdf_text_diagnostics import TrueTypeGlyphInspector, build_font_lookup, resolve_font, normalize_basefont
    from export_zhuyin_readings import load_actual_occurrence_overrides, match_actual_occurrence_override
    fingerprint, evidence_root = args[-2:]
    digest = sp.sha256_file(pdf)
    overrides = load_actual_occurrence_overrides(evidence_root / ar.OCCURRENCE_OVERRIDE_FILE, Path(pdf))
    rows = []
    with fitz.open(pdf) as document:
        page = document[0]
        lookup = build_font_lookup(page.get_fonts(full=True))
        for span in page.get_texttrace():
            xref, *_ = resolve_font(lookup, normalize_basefont(span["font"]))
            inspector = TrueTypeGlyphInspector(document.extract_font(xref)[3])
            for codepoint, gid, origin, box in span["chars"]:
                row = {"實體頁碼": 1, "課本頁": 1, "頁面標籤": "1", "字元": chr(codepoint),
                    "實際注音": "ㄅ", "解碼狀態": "已映射", "解碼依據": "isolated explicit actual fixture",
                    "穩定注音鍵": f"fixture-{gid}", "群組注音鍵": f"fixture-{gid}",
                    "字形架構": "TrueType複合元件", "font": span["font"], "font_xref": xref,
                    "注音元件ID": gid, "glyph_id_字形索引": gid,
                    "TTF字形SHA256": inspector.glyph_sha256(gid),
                    **dict(zip(("x0", "y0", "x1", "y1"), box)),
                    "source_row_number": len(rows) + 2, "ledger_schema_version": ol.LEDGER_SCHEMA_VERSION,
                    "workbook_schema_version": ol.WORKBOOK_SCHEMA_VERSION}
                override = match_actual_occurrence_override(row, overrides)
                if override:
                    row["實際注音"] = override["actual_reading"]
                    row["解碼依據"] = "isolated direct actual override"
                rows.append(row)
    ol.prepare_occurrence_rows(digest, rows)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "實際注音"
    columns = list(dict.fromkeys([*ACTUAL_COLUMNS, "字形架構", "TTF字形SHA256", "CFF樣式群組",
        "CFF整字字形SHA256", "注音結構偵測來源", "CFF符號簽名", "CFF聲調"]))
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(column, "") for column in columns])
    workbook.create_sheet("結構偵測排除").append(EXCLUDED_COLUMNS)
    metadata = workbook.create_sheet("v5.2中繼資料")
    for row in [("項目", "內容"), ("workbook_schema_version", ol.WORKBOOK_SCHEMA_VERSION),
        ("ledger_schema_version", ol.LEDGER_SCHEMA_VERSION), ("pdf_sha256", digest),
        ("actual_decoder_semantics_epoch", fingerprint["components"]["actual_decoder_semantics_epoch"]),
        ("actual_asset_fingerprint", fingerprint["fingerprint"]),
        ("actual_asset_fingerprint_components", json.dumps(fingerprint["components"], ensure_ascii=False, sort_keys=True))]:
        metadata.append(row)
    summary = workbook.create_sheet("摘要")
    summary.append(["項目", "內容"])
    summary.append(["偵測到注音字形筆數", len(rows)])
    workbook.save(output)
    workbook.close()


def _project(root, monkeypatch, *, name="one", extra_character=False):
    monkeypatch.setattr(sp, "decode", _decode)
    pdf = root / (name + ".pdf")
    sha = register_fixture_glyph("daily-learning-source")
    specs = [{"glyph_sha256": sha, "char": "看", "x0": 40, "y0": 40},
             {"glyph_sha256": sha, "char": "瞧" if extra_character else "看", "x0": 40, "y0": 90},
             {"glyph_sha256": sha, "char": "看", "x0": 40, "y0": 140}]
    write_fixture_pdf(pdf, specs)
    output = root / (name + "-project")
    sp.run_pipeline_pdfs([pdf], output, defer_excel_reports=True)
    return output


def _plan(output, *, reading="ㄅ", checked_count=1, checked_offset=0):
    snapshot = ReviewSaveService(output).load_resume_snapshot()
    entry = snapshot.ledger[0]
    group = ar.build_actual_group_for_entry(snapshot.ledger, entry)
    selected = group["members"][checked_offset:checked_offset + checked_count]
    checked = [row["occurrence_id"] for row in selected]
    store = LearningImageStore()
    images = []
    from review_operation_worker import preview_png
    for row in selected:
        value = preview_png(row, padding=(24, 24), output_dir=output, display_width=400)
        assert value["target"] is not None
        images.append(store.save_evidence(row["occurrence_id"], value["png"]))
    return learning.prepare_learning_confirmation(output, snapshot, group["group_id"], reading,
        checked, image_evidence=images, image_store_root=store.root)


def test_complete_roster_includes_matching_actual_and_groups_exact_identity(tmp_path):
    rows = rows_at(tmp_path / "source")
    for row in rows:
        row.update(actual="ㄅ", expected_set=["ㄅ"], state="PASS")
    other = rows_at(tmp_path / "other", count=1)[0]
    # A genuinely different complete glyf, same printed character.
    from tests.actual_source_test_support import source_entry
    other = source_entry("different-glyph", occurrence_id="other-position", char=rows[0]["char"],
                         actual="ㄅ", state="PASS")
    all_rows = rows + [other]
    context = context_for(all_rows)
    context["pdfs"][other["pdf_name"]] = {"pdf": other["pdf"], "pdf_sha256": other["pdf_sha256"]}
    # Distinct names bind the two real source documents without aliasing.
    other_pdf = Path(other["pdf"])
    other["pdf_name"] = other_pdf.name
    context = {"session_id": "candidate-test", "pdfs": {}}
    for index, row in enumerate(all_rows):
        if row is other:
            target = tmp_path / "different.pdf"
            target.write_bytes(other_pdf.read_bytes())
            row.update(pdf=str(target), pdf_name=target.name)
        context["pdfs"][row["pdf_name"]] = {"pdf": row["pdf"], "pdf_sha256": row["pdf_sha256"]}
    repo = lib.GlobalExactGlyphRepository(tmp_path / "absent.sqlite3")
    result = learning.collect_learning_candidates(all_rows, context, repository=repo)
    assert sorted(group["occurrence_count"] for group in result["groups"]) == [1, 2]
    assert sum(len(group["candidate_members"]) for group in result["groups"]) == 3
    assert not repo.path.exists()
    changed = copy.deepcopy(all_rows)
    for row in changed:
        row["expected_set"], row["expected_evidence"] = ["ㄆ"], "unrelated expected edit"
    changed_result = learning.collect_learning_candidates(changed, context, repository=repo)
    roster = lambda value: [[row["occurrence_id"] for row in group["candidate_members"]]
                            for group in value["groups"]]
    assert roster(changed_result) == roster(result) and not repo.path.exists()


def test_source_descriptors_are_not_direct_evidence_and_copies_do_not_add_source(tmp_path):
    rows = rows_at(tmp_path / "source")
    for row in rows:
        row.update(actual="ㄅ", state="PASS")
    group = ar.build_actual_group_for_entry(rows, rows[0])
    context = context_for(rows)
    sources, admissions = promotion.visual_source_descriptors(group, [rows[0]["occurrence_id"]], context)
    assert "reading" not in sources[0] and "evidence_class" not in sources[0]
    rejected, rejected_admissions = promotion.direct_visual_intents(group, "not-bopomofo",
        [rows[0]["occurrence_id"]], context)
    assert rejected == [] and rejected_admissions[0]["status"] == "NON_GLOBAL_ELIGIBLE"
    repo = lib.GlobalExactGlyphRepository(tmp_path / "global.sqlite3")
    intents, _ = promotion.direct_visual_intents(group, "ㄅ", [rows[0]["occurrence_id"]], context)
    lib.deliver_global_direct_evidence(repo, intents)
    copied = copy.deepcopy(rows)
    for row in copied:
        row["occurrence_id"] += "-copy"
        row["review_id"] += "-copy"
    copy_context = {**context, "session_id": "renamed-project"}
    result = learning.collect_learning_candidates(copied, copy_context, repository=repo)
    assert result["groups"] == []
    assert all(item["reason"] == "CORRELATED_SOURCE_ALREADY_CONFIRMED" for item in result["suppressed"])
    assert repo.load_inspection_snapshot().records[0].direct_source_count == 1


def test_learning_wait_does_not_change_textbook_completion_or_clear_actual_work(tmp_path):
    rows = rows_at(tmp_path / "source")
    # Actual is a separately authored decoder fixture. Expected is an explicit
    # independent fixture assertion, never copied from row['actual'].
    for row in rows:
        row.update(actual="ㄅ", actual_status="RESOLVED", actual_evidence="isolated native glyph fixture",
            expected_set=["ㄅ"], expected_status="RESOLVED", expected_evidence="independent fixed expected fixture",
            state="PASS")
    original = copy.deepcopy(rows)
    regression = {"ok": True, "required": 1, "executed": 1, "passed": 1,
                  "failed": 0, "not_executed": 0, "duplicate_case_id": 0}
    reconciliation = ol.reconcile_ledger(rows, {row["occurrence_id"] for row in rows})
    before = ol.completion_gate(rows, reconciliation, regression)
    assert before["complete"] and before["status"] == ol.PROOFREAD_COMPLETE
    context = context_for(rows)
    repo = lib.GlobalExactGlyphRepository(tmp_path / "global.sqlite3")
    candidates = learning.collect_learning_candidates(rows, context, repository=repo)
    assert candidates["groups"] and rows == original
    group = ar.build_actual_group_for_entry(rows, rows[0])
    intents, _ = promotion.direct_visual_intents(group, "ㄅ", [rows[0]["occurrence_id"]], context)
    lib.deliver_global_direct_evidence(repo, intents)
    record = repo.load_inspection_snapshot().records[0]
    assert record.direct_source_count == 1 and not record.global_reuse_allowed
    assert repo.load_approval_preview(record.glyph_id).request is None
    learning.collect_learning_candidates(rows, context, repository=repo)
    assert rows == original and ol.completion_gate(rows, reconciliation, regression) == before
    pending = copy.deepcopy(rows)
    pending[0].update(actual="", actual_status="DECODE_ERROR", state="ACTUAL_DECODE_ERROR")
    retained = copy.deepcopy(pending)
    waiting = learning.collect_learning_candidates(pending, context, repository=repo)
    assert pending[0]["occurrence_id"] in waiting["groups"][0]["pending_occurrence_ids"]
    assert pending == retained  # leaving/defering the queue creates no decision
    assert not ol.completion_gate(pending, ol.reconcile_ledger(pending,
        {row["occurrence_id"] for row in pending}), regression)["complete"]


def test_projected_approval_is_read_only_and_matches_real_delivery(tmp_path):
    repo = lib.GlobalExactGlyphRepository(tmp_path / "global.sqlite3")
    first, second = _intent("one"), _intent("two")
    gid = first["payload"]["evidence"]["glyph_id"]
    assert repo.project_direct_evidence_approval([first], gid).request is None
    assert not repo.path.exists()
    lib.deliver_global_direct_evidence(repo, [first])
    before = repo.path.read_bytes()
    generation = repo.load_snapshot().generation
    projected = repo.project_direct_evidence_approval([second], gid)
    assert projected.request is not None and projected.record is None
    assert projected.request.database_path == str(repo.path)
    assert repo.path.read_bytes() == before and repo.load_snapshot().generation == generation
    lib.deliver_global_direct_evidence(repo, [second])
    assert repo.load_approval_preview(gid).request == projected.request
    lib.approve_global_exact_glyph(repo, **projected.request.service_arguments())
    assert repo.lookup_approval_outcome(projected.request).current_reuse_allowed


def test_changed_evidence_cannot_be_approved_by_old_click(tmp_path):
    repo = lib.GlobalExactGlyphRepository(tmp_path / "global.sqlite3")
    first, second, third = (_intent(label) for label in ("one", "two", "three"))
    gid = first["payload"]["evidence"]["glyph_id"]
    lib.deliver_global_direct_evidence(repo, [first])
    request = repo.project_direct_evidence_approval([second], gid).request
    lib.deliver_global_direct_evidence(repo, [second, third])
    with pytest.raises(lib.GlobalLibraryValidationError, match="stale"):
        lib.approve_global_exact_glyph(repo, **request.service_arguments())
    assert not repo.load_inspection_snapshot().records[0].global_reuse_allowed


def test_machine_difference_is_candidate_until_visual_conflict(tmp_path):
    rows = rows_at(tmp_path / "source", count=1)
    rows[0].update(actual="ㄆ", state="PASS")
    group = ar.build_actual_group_for_entry(rows, rows[0])
    sources, _ = promotion.visual_source_descriptors(group, [rows[0]["occurrence_id"]], context_for(rows))
    identity = sources[0]["identity"]
    first, second = _intent("one", identity=identity), _intent("two", identity=identity)
    repo = lib.GlobalExactGlyphRepository(tmp_path / "global.sqlite3")
    lib.deliver_global_direct_evidence(repo, [first, second])
    gid = first["payload"]["evidence"]["glyph_id"]
    request = repo.load_approval_preview(gid).request
    lib.approve_global_exact_glyph(repo, **request.service_arguments())
    assert len(learning.collect_learning_candidates(rows, context_for(rows), repository=repo)["groups"]) == 1
    assert repo.load_inspection_snapshot().records[0].stored_state == lib.VERIFIED_GLOBAL
    actual_intents, _ = promotion.direct_visual_intents(group, "ㄆ", [rows[0]["occurrence_id"]], context_for(rows))
    lib.deliver_global_direct_evidence(repo, actual_intents)
    record = repo.load_inspection_snapshot().records[0]
    assert record.stored_state == lib.QUARANTINED_CONFLICT and not record.global_reuse_allowed
    assert record.readings == ("ㄅ", "ㄆ")


def test_first_save_reopen_second_independent_source_and_retry(tmp_path, monkeypatch):
    tmp_path = _isolated(tmp_path, monkeypatch)
    first = _project(tmp_path, monkeypatch)
    plan = _plan(first)
    assert plan["approval"] is None
    with pytest.raises(ValueError, match="actual final user click"):
        learning.confirm_learning(plan)
    result = learning.confirm_learning(plan, explicit_action=True)
    assert result["local_status"] == "SAVED" and result["status"] == "LOCAL_SAVED_WAITING_SOURCE"
    reopened = learning.list_learning_confirmations(first)
    assert reopened[0]["operation_id"] == plan["operation_id"]
    repo = lib.GlobalExactGlyphRepository.resolved()
    assert repo.load_inspection_snapshot().records[0].direct_source_count == 1
    again = learning.retry_learning_confirmation(first, plan["operation_id"])
    assert again["local_status"] == "SAVED"
    assert repo.load_inspection_snapshot().records[0].direct_source_count == 1
    second = _project(tmp_path, monkeypatch, name="two", extra_character=True)
    other = _plan(second)
    assert other["approval"] is not None
    activated = learning.confirm_learning(other, explicit_action=True)
    assert activated["reusable"] and activated["global_status"] == "APPROVAL_COMMITTED"
    assert repo.load_inspection_snapshot().records[0].direct_source_count == 2
    snapshot = ReviewSaveService(second).load_resume_snapshot()
    assert learning.collect_learning_candidates(snapshot.ledger,
        sp._global_direct_source_context(second, snapshot.manifest), repository=repo,
        evidence_root=sp.project_actual_evidence_root(second))["groups"] == []


def test_source_image_and_project_drift_refuse_before_local_truth(tmp_path, monkeypatch):
    tmp_path = _isolated(tmp_path, monkeypatch)
    output = _project(tmp_path, monkeypatch)
    plan = _plan(output)
    image = Path(plan["image_store_root"]) / plan["image_evidence"][0]["relative_path"]
    raw = image.read_bytes()
    image.write_bytes(raw + b"changed")
    with pytest.raises(ValueError):
        learning.confirm_learning(plan, explicit_action=True)
    image.write_bytes(raw)
    db = output / "人工判定資料庫.json"
    db.write_bytes(db.read_bytes() + b"\n")
    with pytest.raises(ValueError, match="LEARNING_STALE"):
        learning.confirm_learning(plan, explicit_action=True)
    assert learning.list_learning_confirmations(output) == []
    assert not lib.GlobalExactGlyphRepository.resolved().path.exists()


def test_interruption_after_local_commit_retries_without_recount(tmp_path, monkeypatch):
    tmp_path = _isolated(tmp_path, monkeypatch)
    output = _project(tmp_path, monkeypatch)
    plan = _plan(output)
    deliver = promotion.deliver_pending_promotion_outbox
    monkeypatch.setattr(promotion, "deliver_pending_promotion_outbox",
        lambda *args, **kwargs: (_ for _ in ()).throw(OSError("isolated interruption after local commit")))
    result = learning.confirm_learning(plan, explicit_action=True)
    assert result["local_status"] == "SAVED" and result["status"] == "REFRESH_PENDING"
    assert learning.list_learning_confirmations(output)[0]["local_status"] == "SAVED"
    monkeypatch.setattr(promotion, "deliver_pending_promotion_outbox", deliver)
    recovered = learning.retry_learning_confirmation(output, plan["operation_id"])
    assert recovered["status"] == "LOCAL_SAVED_WAITING_SOURCE"
    assert lib.GlobalExactGlyphRepository.resolved().load_inspection_snapshot().records[0].direct_source_count == 1


def test_only_viewed_position_gets_direct_provenance(tmp_path, monkeypatch):
    tmp_path = _isolated(tmp_path, monkeypatch)
    output = _project(tmp_path, monkeypatch)
    plan = _plan(output, checked_count=1, checked_offset=1)
    viewed = plan["group"]["members"][1]
    assert plan["group"]["members"][0]["occurrence_id"] not in plan["checked_occurrence_ids"]
    result = learning.confirm_learning(plan, explicit_action=True)
    assert result["local_status"] == "SAVED"
    assert result["review_id"] == viewed["review_id"]
    assert learning.list_learning_confirmations(output)[0]["review_id"] == viewed["review_id"]
    snapshot = ReviewSaveService(output).load_resume_snapshot()
    checked = set(plan["checked_occurrence_ids"])
    root = sp.project_actual_evidence_root(output)
    assert [row["occurrence_id"] for row in snapshot.ledger if ar.has_direct_occurrence_actual(root, row, "ㄅ")] == list(checked)
    before = lib.GlobalExactGlyphRepository.resolved().load_inspection_snapshot().records[0].direct_source_count
    # Changing an independent expected decision does not add actual evidence.
    row = snapshot.ledger[0]
    ReviewSaveService(output).save_event(row["review_id"],
        sp.build_manual_expected_event(row, operation="ENTER_EXPECTED", expected_set="ㄆ"))
    assert lib.GlobalExactGlyphRepository.resolved().load_inspection_snapshot().records[0].direct_source_count == before


def test_overpaint_is_never_a_projected_global_donor(tmp_path):
    rows = rows_at(tmp_path / "covered", painting="replacement")
    rows[0].update(actual="ㄅ", expected_set=["ㄅ"], state="PASS")
    group = ar.build_actual_group_for_entry(rows, rows[0])
    intents, admissions = promotion.direct_visual_intents(group, "ㄅ",
        [rows[0]["occurrence_id"]], context_for(rows))
    assert intents == [] and admissions[0]["status"] == "NON_GLOBAL_ELIGIBLE"
    root = tmp_path / "actual"
    with pytest.raises(ValueError, match="禁止確認或寫入"):
        ar.apply_verified_actual_group(root, group, "ㄅ",
            checked_occurrence_ids=[rows[0]["occurrence_id"]], source="isolated visual action")
    assert not root.exists()


def test_project_lock_and_interruption_before_commit_keep_original_action(tmp_path, monkeypatch):
    import os
    tmp_path = _isolated(tmp_path, monkeypatch)
    output = _project(tmp_path, monkeypatch)
    plan = _plan(output)
    root = sp.project_actual_evidence_root(output)
    with (root / ".global_exact_glyph_delivery.lock").open("r+b") as handle:
        if os.name == "nt":
            import msvcrt
            msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
        else:
            import fcntl
            fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        try:
            with pytest.raises(OSError):
                learning.confirm_learning(plan, explicit_action=True)
        finally:
            handle.seek(0)
            if os.name == "nt":
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
    assert learning.list_learning_confirmations(output) == []
    apply = ar.apply_direct_visual_actual_batch
    monkeypatch.setattr(ar, "apply_direct_visual_actual_batch",
        lambda *args, **kwargs: (_ for _ in ()).throw(OSError("isolated before local commit")))
    with pytest.raises(OSError, match="before local commit"):
        learning.confirm_learning(plan, explicit_action=True)
    assert learning.list_learning_confirmations(output)[0]["status"] == "LOCAL_PENDING"
    monkeypatch.setattr(ar, "apply_direct_visual_actual_batch", apply)
    result = learning.retry_learning_confirmation(output, plan["operation_id"])
    assert result["local_status"] == "SAVED" and result["status"] == "LOCAL_SAVED_WAITING_SOURCE"


def test_stale_combined_approval_requires_fresh_visual_action_without_recount(tmp_path, monkeypatch):
    tmp_path = _isolated(tmp_path, monkeypatch)
    first = _project(tmp_path, monkeypatch)
    learning.confirm_learning(_plan(first), explicit_action=True)
    second = _project(tmp_path, monkeypatch, name="two", extra_character=True)
    old = _plan(second)
    repo = lib.GlobalExactGlyphRepository.resolved()
    identity = old["intents"][0]["payload"]["identity"]
    lib.deliver_global_direct_evidence(repo, [_intent("independent-concurrent", identity=identity)])
    result = learning.confirm_learning(old, explicit_action=True)
    assert result["local_status"] == "SAVED" and not result["reusable"]
    assert result["global_status"] == "STALE_APPROVAL_REQUIRES_NEW_ACTION"
    snapshot = ReviewSaveService(second).load_resume_snapshot()
    candidates = learning.collect_learning_candidates(snapshot.ledger,
        sp._global_direct_source_context(second, snapshot.manifest), repository=repo,
        evidence_root=sp.project_actual_evidence_root(second))["groups"]
    assert len(candidates) == 1 and candidates[0]["requires_fresh_approval_action"]
    fresh = _plan(second)
    assert fresh["operation_id"] != old["operation_id"]
    assert fresh["approval"]["quorum_evidence_ids"] != old["approval"]["quorum_evidence_ids"]
    saved = learning.confirm_learning(fresh, explicit_action=True)
    assert saved["reusable"]
    assert repo.load_inspection_snapshot().records[0].direct_source_count == 3
    assert repo.lookup_approval_outcome(learning._approval_request(old["approval"])).status == "NOT_OBSERVED"
    old_view = next(row for row in learning.list_learning_confirmations(second)
                    if row["operation_id"] == old["operation_id"])
    assert old_view["status"] == "SUPERSEDED_BY_CURRENT_TRUTH" and not old_view["reusable"]


def test_visual_conflict_removes_unseen_propagated_override_and_keeps_expected(tmp_path, monkeypatch):
    tmp_path = _isolated(tmp_path, monkeypatch)
    output = _project(tmp_path, monkeypatch)
    snapshot = ReviewSaveService(output).load_resume_snapshot()
    first, second, unseen = snapshot.ledger
    sp.apply_manual_actual_correction(output, first["review_id"], "ㄅ",
        [first["occurrence_id"], second["occurrence_id"]])
    snapshot = ReviewSaveService(output).load_resume_snapshot()
    root = sp.project_actual_evidence_root(output)
    assert learning.collect_learning_candidates(snapshot.ledger,
        sp._global_direct_source_context(output, snapshot.manifest), evidence_root=root)["groups"] == []
    event = sp.build_manual_expected_event(snapshot.ledger[0], operation="ENTER_EXPECTED", expected_set="ㄆ")
    ReviewSaveService(output).save_event(first["review_id"], event)
    original_expected = sp.json_load_strict(output / "人工判定資料庫.json")["events"][first["review_id"]]
    plan = _plan(output, reading="ㄆ")
    assert len(plan["group"]["members"]) == 3 and len(plan["checked_occurrence_ids"]) == 1
    result = learning.confirm_learning(plan, explicit_action=True)
    assert result["global_status"] == "QUARANTINED_CONFLICT" and not result["reusable"]
    overrides = ar._read_csv(root / ar.OCCURRENCE_OVERRIDE_FILE, ar.OVERRIDE_HEADERS)
    assert not any(float(row["y0"]) == pytest.approx(unseen["y0"], abs=.01) for row in overrides)
    assert sp.json_load_strict(output / "人工判定資料庫.json")["events"][first["review_id"]] == original_expected
    assert lib.GlobalExactGlyphRepository.resolved().load_inspection_snapshot().records[0].stored_state == lib.QUARANTINED_CONFLICT


def test_moved_completed_and_pending_decisions_keep_original_evidence_and_count(tmp_path, monkeypatch):
    root = _isolated(tmp_path, monkeypatch)
    for stage in ("completed", "before_local", "after_local"):
        # Each scenario owns its Global store, original PDF and selected project.
        area = root / stage
        area.mkdir()
        monkeypatch.setenv("LOCALAPPDATA", str(area / "u"))
        output = _project(area, monkeypatch)
        plan = _plan(output)
        with monkeypatch.context() as failure:
            if stage == "before_local":
                failure.setattr(ar, "apply_direct_visual_actual_batch",
                    lambda *a, **k: (_ for _ in ()).throw(OSError("before local")))
                with pytest.raises(OSError, match="before local"):
                    learning.confirm_learning(plan, explicit_action=True)
            else:
                if stage == "after_local":
                    failure.setattr(promotion, "deliver_pending_promotion_outbox",
                        lambda *a, **k: (_ for _ in ()).throw(OSError("after local")))
                learning.confirm_learning(plan, explicit_action=True)
        decision, _ = learning._paths(output, plan["operation_id"])
        raw = decision.read_bytes()
        moved_area = root / (stage + "-moved")
        area.rename(moved_area)
        # User data is outside the relocated project in production. Keep the
        # isolated Global/image location unchanged while moving the PDF+project.
        area.mkdir()
        (moved_area / "u").rename(area / "u")
        selected = moved_area / "renamed-project"
        (moved_area / output.name).rename(selected)
        view = learning.list_learning_confirmations(selected)
        assert len(view) == 1 and view[0]["status"] != "UNRECOVERABLE", view
        result = learning.retry_learning_confirmation(selected, plan["operation_id"])
        assert result["local_status"] == "SAVED" and result["status"] == "LOCAL_SAVED_WAITING_SOURCE"
        learning.retry_learning_confirmation(selected, plan["operation_id"])
        assert lib.GlobalExactGlyphRepository.resolved().load_inspection_snapshot().records[0].direct_source_count == 1
        assert learning._paths(selected, plan["operation_id"])[0].read_bytes() == raw
        assert not output.exists() and not (area / "one.pdf").exists()


def test_foreign_or_broken_learning_history_does_not_hide_legal_candidates(tmp_path, monkeypatch):
    root = _isolated(tmp_path, monkeypatch)
    first = _project(root, monkeypatch)
    plan = _plan(first)
    with monkeypatch.context() as interrupted:
        interrupted.setattr(ar, "apply_direct_visual_actual_batch",
            lambda *a, **k: (_ for _ in ()).throw(OSError("before local")))
        with pytest.raises(OSError, match="before local"):
            learning.confirm_learning(plan, explicit_action=True)
    original = learning._paths(first, plan["operation_id"])[0].read_bytes()
    other = _project(root, monkeypatch, name="other")
    copied, _ = learning._paths(other, plan["operation_id"])
    copied.parent.mkdir(parents=True)
    copied.write_bytes(original)
    with pytest.raises(ValueError, match="selected session"):
        learning.retry_learning_confirmation(other, plan["operation_id"])
    assert learning.list_learning_confirmations(other)[0]["status"] == "UNRECOVERABLE"
    # A malformed independent history entry cannot block all candidate viewing.
    broken = copied.with_name("f" * 64 + ".json")
    broken.write_text('{"broken": true}', encoding="utf-8")
    from review_operation_worker import ReviewOperationState
    snapshot = ReviewSaveService(other).load_resume_snapshot()
    result = ReviewOperationState(other).dispatch("load_learning", dict(
        expected_manifest=snapshot.manifest, expected_db=snapshot.db), lambda _: None, lambda: False)
    assert len(result["pending"]) == 2 and all(row["status"] == "UNRECOVERABLE" for row in result["pending"])
    # The readable source can still be explicitly opened, independent of history.
    assert result["snapshot"].ledger and result["candidates"]["groups"]
    image = Path(plan["image_store_root"]) / plan["image_evidence"][0]["relative_path"]
    image.unlink()
    assert learning.list_learning_confirmations(first)[0]["status"] == "UNRECOVERABLE"
    with pytest.raises(OSError):
        learning.retry_learning_confirmation(first, plan["operation_id"])
    assert copied.read_bytes() == original


def test_moved_decision_rejects_changed_workbook_and_source_without_old_writes(tmp_path, monkeypatch):
    root = _isolated(tmp_path, monkeypatch)
    output = _project(root, monkeypatch)
    plan = _plan(output)
    learning.confirm_learning(plan, explicit_action=True)
    selected = root / "renamed"
    output.rename(selected)
    from sealed_workbook_paths import descriptor
    manifest = sp.json_load_strict(selected / "校對工作階段.json")
    workbook = selected / descriptor(manifest["pdfs"][0], "actual")
    pdf = root / "one.pdf"
    for path in (workbook, pdf):
        original = path.read_bytes()
        path.write_bytes(original + b"changed")
        assert learning.list_learning_confirmations(selected)[0]["status"] == "UNRECOVERABLE"
        with pytest.raises(ValueError):
            learning.retry_learning_confirmation(selected, plan["operation_id"])
        assert not output.exists()
        path.write_bytes(original)
    assert learning.retry_learning_confirmation(selected, plan["operation_id"])["local_status"] == "SAVED"
    assert lib.GlobalExactGlyphRepository.resolved().load_inspection_snapshot().records[0].direct_source_count == 1


def test_moved_owned_journal_uses_authorized_binding_without_new_vote(tmp_path, monkeypatch):
    from tests.test_glyph_learning_gui import surface_request
    from review_operation_worker import ReviewOperationState
    root = _isolated(tmp_path, monkeypatch)
    output = _project(root, monkeypatch)
    _snapshot, request = surface_request(output)
    ReviewOperationState(output).dispatch("save_learning_surface", request, lambda _: None, lambda: False)
    plan = _plan(output, checked_offset=1)
    with monkeypatch.context() as interrupted:
        interrupted.setattr(promotion, "deliver_pending_promotion_outbox",
            lambda *a, **k: (_ for _ in ()).throw(OSError("after owned local commit")))
        saved = learning.confirm_learning(plan, explicit_action=True)
    assert saved["local_status"] == "SAVED"
    journal_path = sp.project_actual_evidence_root(output) / promotion.PROJECT_TRANSACTION_FILE
    journal_raw = journal_path.read_bytes()
    assert promotion._validated_project_journal(journal_raw)["schema_version"] == "3.0"
    selected = root / "owned-moved"
    output.rename(selected)
    decision_raw = learning._paths(selected, plan['operation_id'])[0].read_bytes()
    assert learning.list_learning_confirmations(selected)[0]["status"] != "UNRECOVERABLE"
    result = learning.retry_learning_confirmation(selected, plan["operation_id"])
    assert result['local_status'] == 'SAVED' and result['refresh_status'] == 'SUCCESS'
    learning.retry_learning_confirmation(selected, plan["operation_id"])
    assert not output.exists()
    assert not (sp.project_actual_evidence_root(selected) / promotion.PROJECT_TRANSACTION_FILE).exists()
    assert learning._paths(selected, plan['operation_id'])[0].read_bytes() == decision_raw
    assert lib.GlobalExactGlyphRepository.resolved().load_inspection_snapshot().records[0].direct_source_count == 1


def test_learning_retry_one_entry_rolls_back_prepared_then_preserves_fixed_decision(tmp_path, monkeypatch):
    from tests.test_owned_journal_relocation import _profile, _files, _relocate
    root, area, output, before = _profile(tmp_path, monkeypatch)
    plan = _plan(output, checked_offset=1)
    actual = sp.project_actual_evidence_root(output)
    atomic = promotion._atomic_bytes
    crashed = []
    def stop_after_actual_write(path, raw, **kwargs):
        atomic(path, raw, **kwargs)
        if Path(path).name == ar.OCCURRENCE_OVERRIDE_FILE and not crashed:
            journal = promotion._validated_project_journal((actual / promotion.PROJECT_TRANSACTION_FILE).read_bytes())
            assert journal['schema_version'] == '3.0' and journal['state'] == 'PREPARED'
            crashed.append(_files(output))
            raise OSError('isolated learning PREPARED process loss')
    with monkeypatch.context() as fault:
        fault.setattr(promotion, '_atomic_bytes', stop_after_actual_write)
        with pytest.raises(OSError, match='PREPARED process loss'):
            learning.confirm_learning(plan, explicit_action=True)
    assert crashed and not (actual / promotion.PROJECT_TRANSACTION_FILE).exists()
    for name, raw in crashed[0].items():
        path = output / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(raw)
    decision_raw = learning._paths(output, plan['operation_id'])[0].read_bytes()
    selected, old = _relocate(root, area, output)
    result = learning.retry_learning_confirmation(selected, plan['operation_id'])
    assert result['operation_id'] == plan['operation_id'] and result['refresh_status'] == 'SUCCESS'
    assert result['local_status'] == 'SAVED'
    learning.retry_learning_confirmation(selected, plan['operation_id'])
    assert learning._paths(selected, plan['operation_id'])[0].read_bytes() == decision_raw
    assert lib.GlobalExactGlyphRepository.resolved().load_inspection_snapshot().records[0].direct_source_count == 1
    assert not lib.GlobalExactGlyphRepository.resolved().load_inspection_snapshot().records[0].global_reuse_allowed
    assert ReviewSaveService(selected).load_resume_snapshot().db == before.db
    assert _files(area) == old
    assert not (sp.project_actual_evidence_root(selected) / promotion.PROJECT_TRANSACTION_FILE).exists()


def test_moved_learning_partial_publication_retains_workbook_gate(tmp_path, monkeypatch):
    from tests.test_glyph_learning_gui import surface_request
    from review_operation_worker import ReviewOperationState
    root = _isolated(tmp_path, monkeypatch)
    output = _project(root, monkeypatch)
    _snapshot, request = surface_request(output)
    ReviewOperationState(output).dispatch('save_learning_surface', request, lambda _: None, lambda: False)
    plan = _plan(output, checked_offset=1)
    move, resume = sp._surface_phase_move, sp._resume_actual_refresh_publication
    moved = []
    def stop_after_move(source, target, phase):
        move(source, target, phase)
        if not moved:
            moved.append(True)
            raise OSError('isolated learning publication prefix interruption')
    def no_hot_resume(*args, **kwargs):
        if moved:
            raise OSError('isolated stopped process')
        return resume(*args, **kwargs)
    with monkeypatch.context() as fault:
        fault.setattr(sp, '_surface_phase_move', stop_after_move)
        fault.setattr(sp, '_resume_actual_refresh_publication', no_hot_resume)
        result = learning.confirm_learning(plan, explicit_action=True)
    assert result['refresh_status'] == 'PENDING' and moved
    assert sp.json_load_strict(output / sp.ACTUAL_REFRESH_PUBLICATION)['phase'] == 'PREPARED'
    decision_raw = learning._paths(output, plan['operation_id'])[0].read_bytes()
    selected = root / 'partial-moved'
    output.rename(selected)
    assert learning.list_learning_confirmations(selected)[0]['status'] != 'UNRECOVERABLE'
    manifest = sp.json_load_strict(selected / '校對工作階段.json')
    workbook = selected / sp.descriptor(manifest['pdfs'][0], 'actual')
    original = workbook.read_bytes()
    workbook.write_bytes(original + b'foreign workbook bytes')
    assert learning.list_learning_confirmations(selected)[0]['status'] == 'UNRECOVERABLE'
    with pytest.raises(ValueError):
        learning.retry_learning_confirmation(selected, plan['operation_id'])
    workbook.write_bytes(original)
    result = learning.retry_learning_confirmation(selected, plan['operation_id'])
    assert result['local_status'] == 'SAVED' and result['refresh_status'] == 'SUCCESS'
    learning.retry_learning_confirmation(selected, plan['operation_id'])
    assert learning._paths(selected, plan['operation_id'])[0].read_bytes() == decision_raw
    assert not output.exists()
    assert lib.GlobalExactGlyphRepository.resolved().load_inspection_snapshot().records[0].direct_source_count == 1


def test_valid_surface_confirmation_suppresses_only_applied_position(tmp_path, monkeypatch):
    import manual_surface_actual as surface
    from tests.test_actual_source_visual_safety import make_visual_source
    from review_operation_worker import ReviewOperationState
    root = _isolated(tmp_path, monkeypatch)
    monkeypatch.setattr(sp, "decode", _decode)
    rows = make_visual_source(root / "src", count=2)
    pdf = root / "src" / "painted.pdf"
    with fitz.open(rows[0]["pdf"]) as document:
        document[0].draw_rect(fitz.Rect(41, 45, 44, 48), fill=(1, 1, 1), color=None)
        document.save(pdf)
    output = root / "project"
    sp.run_pipeline_pdfs([pdf], output, defer_excel_reports=True)
    snapshot = ReviewSaveService(output).load_resume_snapshot()
    target = next(row for row in snapshot.ledger if row["char"] == "A" and row["y0"] < 60)
    group = ar.build_actual_group_for_entry(snapshot.ledger, target)
    state = ReviewOperationState(output)
    selected = state.dispatch("prepare_learning_target", dict(group_id=group["group_id"],
        review_id=target["review_id"], expected_manifest=snapshot.manifest, expected_db=snapshot.db),
        lambda _: None, lambda: False)
    assert selected["surface"] is not None
    state.dispatch("save_learning_surface", dict(expected_manifest=snapshot.manifest, expected_db=snapshot.db,
        review_id=target["review_id"], reading="ㄆ", proof=selected["surface"]["proof"], actor="isolated fixture",
        whole_unit_readable=True, target_character_matches=True, reading_from_original=True, enable_format=True),
        lambda _: None, lambda: False)
    reopened = ReviewSaveService(output).load_resume_snapshot()
    context = sp._global_direct_source_context(output, reopened.manifest)
    evidence = sp.project_actual_evidence_root(output)
    result = learning.collect_learning_candidates(reopened.ledger, context, evidence_root=evidence)
    candidates = {row["occurrence_id"] for group in result["groups"] for row in group["candidate_members"]}
    assert target["occurrence_id"] not in candidates
    assert candidates == {row["occurrence_id"] for row in reopened.ledger} - {target["occurrence_id"]}
    assert any(row["reason"] == "SURFACE_POSITION_ALREADY_CONFIRMED" for row in result["suppressed"])
    assert reopened.db == snapshot.db and not lib.GlobalExactGlyphRepository.resolved().path.exists()
    # The ordinary actual correction route remains available for this position.
    edit = state.dispatch("prepare_actual", dict(review_id=target["review_id"],
        expected_manifest=reopened.manifest, expected_db=reopened.db), lambda _: None, lambda: False)
    assert edit["surface"] is not None
    with monkeypatch.context() as damaged:
        damaged.setattr(surface, "verify_sources", lambda *a, **k: (_ for _ in ()).throw(ValueError("source lost")))
        unverified = learning.collect_learning_candidates(reopened.ledger, context, evidence_root=evidence)
    assert target["occurrence_id"] in {row["occurrence_id"] for group in unverified["groups"] for row in group["candidate_members"]}
    assert any(row["status"] == "SURFACE_CONFIRMATION_UNVERIFIED" for row in unverified["diagnostics"])
    changed = copy.deepcopy(reopened.ledger)
    next(row for row in changed if row["occurrence_id"] == target["occurrence_id"])["manual_surface_actual"]["receipt_sha256"] = "0" * 64
    invalid = learning.collect_learning_candidates(changed, context, evidence_root=evidence)
    assert target["occurrence_id"] in {row["occurrence_id"] for group in invalid["groups"] for row in group["candidate_members"]}
    original_pdf = pdf.read_bytes()
    pdf.write_bytes(original_pdf + b"changed sealed source")
    invalid_source = learning.collect_learning_candidates(reopened.ledger, context, evidence_root=evidence)
    assert target["occurrence_id"] in {row["occurrence_id"] for group in invalid_source["groups"] for row in group["candidate_members"]}
    assert any(row["status"] == "SURFACE_CONFIRMATION_UNVERIFIED" for row in invalid_source["diagnostics"])
    pdf.write_bytes(original_pdf)
