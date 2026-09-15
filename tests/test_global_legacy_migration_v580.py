from __future__ import annotations

import ast
import copy
import csv
import hashlib
import io
import json
import multiprocessing
import os
from pathlib import Path
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from contextlib import closing, redirect_stdout, redirect_stderr

import fitz
from fontTools.fontBuilder import FontBuilder
from fontTools.pens.ttGlyphPen import TTGlyphPen
from fontTools.pens.t2CharStringPen import T2CharStringPen
from openpyxl import Workbook, load_workbook

import actual_review as review
import global_exact_glyph_library as lib
import global_glyph_promotion as promotion
import legacy_global_migration as migration
import migrate_legacy_global_candidates as cli
from occurrence_ledger import LEDGER_SCHEMA_VERSION, SESSION_SCHEMA_VERSION, WORKBOOK_SCHEMA_VERSION, REVIEW_ID_SCHEMA_VERSION
from test_exact_glyph_identity_v580 import synthetic_sfnt, simple_glyph, composite_glyph
from test_global_promotion_v580 import intent, approve, sql_rows, FakeDocument


ROOT = Path(__file__).resolve().parents[1]


def sha(raw):
    return hashlib.sha256(raw if isinstance(raw, bytes) else raw.encode()).hexdigest()


def write_csv(path, headers, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows([{key: row.get(key, "") for key in headers} for row in rows])


def tree_bytes(root):
    return {str(path.relative_to(root)): path.read_bytes() for path in root.rglob("*") if path.is_file()}


def fixture_path_alias(path):
    """Same isolated directory via a guaranteed alias, plus 8.3 when available.

    The parent segment always exercises canonical path comparisons, including
    on filesystems without short names. GetShortPathNameW is read-only and does
    not enable 8.3 names, create links, or change any filesystem policy.
    """
    canonical = Path(path).resolve()
    alias = canonical
    if os.name == "nt":
        import ctypes
        from ctypes import wintypes
        get_short_path = ctypes.WinDLL("kernel32", use_last_error=True).GetShortPathNameW
        get_short_path.argtypes = (wintypes.LPCWSTR, wintypes.LPWSTR, wintypes.DWORD)
        get_short_path.restype = wintypes.DWORD
        size = get_short_path(str(canonical), None, 0)
        if size:
            buffer = ctypes.create_unicode_buffer(size)
            length = get_short_path(str(canonical), buffer, size)
            if 0 < length < size:
                alias = Path(buffer.value)
    return alias / ".." / alias.name


def seal(project):
    project.manifest["manifest_integrity_sha256"] = lib._canonical_sha256({
        key: value for key, value in project.manifest.items() if key != "manifest_integrity_sha256"})
    (project.root / "校對工作階段.json").write_text(
        json.dumps(project.manifest, ensure_ascii=False, indent=2), encoding="utf-8")


def embed_cff(document, page, raw, name, count):
    """A real Type1C stream with a Unicode map and deterministic glyph encoding."""
    def obj(value):
        xref = document.get_new_xref()
        document.update_object(xref, value)
        return xref
    stream = obj("<< /Subtype /Type1C >>")
    document.update_stream(stream, raw)
    descriptor = obj(f"<< /Type /FontDescriptor /FontName /{name} /Flags 4 /FontBBox [0 0 600 800] "
                     f"/ItalicAngle 0 /Ascent 800 /Descent -200 /CapHeight 800 /StemV 80 /FontFile3 {stream} 0 R >>")
    unicode_map = obj("<< >>")
    document.update_stream(unicode_map,
        b"/CIDInit /ProcSet findresource begin 12 dict begin begincmap "
        b"/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def "
        b"/CMapName /MigrationTest def /CMapType 2 def "
        b"1 begincodespacerange <00> <FF> endcodespacerange "
        b"1 beginbfchar <41> <5B57> endbfchar endcmap "
        b"CMapName currentdict /CMap defineresource pop end end")
    font = obj(f"<< /Type /Font /Subtype /Type1 /BaseFont /{name} /FontDescriptor {descriptor} 0 R "
               f"/FirstChar 65 /LastChar 65 /Widths [600] /Encoding << /Type /Encoding "
               f"/Differences [65 /cid00001] >> /ToUnicode {unicode_map} 0 R >>")
    document.xref_set_key(page.xref, "Resources", f"<< /Font << /F5 {font} 0 R >> >>")
    contents = obj("<< >>")
    document.update_stream(contents, "\n".join(
        f"BT /F5 16 Tf {60 + i * 30} 760 Td <41> Tj ET" for i in range(count)).encode())
    document.xref_set_key(page.xref, "Contents", f"{contents} 0 R")
    return font


class Project:
    """Real embedded-font PDF, hashed workbook and sealed actual-only roster."""
    def __init__(self, root, *, session="legacy-session", reading="ㄅ", count=1, cff=False,
                 style_name="DFBiaoKaiZhuIn-W5", record=None):
        self.root = Path(root)
        self.root.mkdir(parents=True)
        self.evidence = self.root / "_專案證據" / "actual"
        self.pdf = self.root / "book.pdf"
        builder = FontBuilder(1000, isTTF=not cff)
        glyph_name = "cid00001" if cff else "testGlyph"
        builder.setupGlyphOrder([".notdef", glyph_name])
        builder.setupCharacterMap({ord("字"): glyph_name})
        if cff:
            pen = T2CharStringPen(600, None)
            pen.moveTo((0, 0)); pen.lineTo((200, 0)); pen.lineTo((200, 300)); pen.closePath()
            empty = T2CharStringPen(600, None)
            builder.setupCFF(style_name, {}, {".notdef": empty.getCharString(), glyph_name: pen.getCharString()}, {})
        else:
            pen = TTGlyphPen(None)
            pen.moveTo((0, 0)); pen.lineTo((200, 0)); pen.lineTo((200, 300)); pen.closePath()
            builder.setupGlyf({".notdef": TTGlyphPen(None).glyph(), "testGlyph": pen.glyph()})
        builder.setupHorizontalMetrics({".notdef": (600, 0), glyph_name: (600, 0)})
        builder.setupHorizontalHeader(ascent=800, descent=-200)
        builder.setupNameTable({"familyName": style_name if cff else "PhaseFiveTest", "styleName": "Regular"})
        builder.setupOS2(sTypoAscender=800, sTypoDescender=-200, usWinAscent=800, usWinDescent=200)
        builder.setupPost()
        data = io.BytesIO(); builder.save(data)
        with fitz.open() as doc:
            page = doc.new_page()
            if cff:
                xref = embed_cff(doc, page, builder.font["CFF "].compile(builder.font), style_name, count)
            else:
                xref = page.insert_font(fontname="MigrationFont", fontbuffer=data.getvalue())
                for i in range(max(count, 1)):
                    page.insert_text((60 + i * 30, 80), "字", fontname="MigrationFont", fontsize=16)
            doc.save(self.pdf)
        with fitz.open(self.pdf) as doc:
            characters = [char for span in doc[0].get_texttrace() for char in span["chars"]]
            font_name, extension, _, font = doc.extract_font(xref)
        self.cff_data = builder.font["CFF "].compile(builder.font) if cff else None
        self.style_name = style_name
        self.identity = (promotion.CFFZhuyinInspector(self.cff_data, style_name, {}).global_exact_identity(1)
                         if cff else promotion.TrueTypeGlyphInspector(font).global_exact_identity(1))
        if record is not None:
            # For structural rejection tests only: exact locator/PDF APIs remain
            # controlled by FakeDocument, while the real Phase 1 parser runs.
            self.identity = promotion.TrueTypeGlyphInspector(synthetic_sfnt([record], font_name_marker=b"test")).global_exact_identity(0)
        self.exact = {"kind": lib.CFF_GLYPH_SHA256 if cff else lib.TTF_GLYF_SHA256,
                      "style_group": self.identity.get("style_group", ""),
                      "glyph_sha256": self.identity["glyph_sha256"],
                      "identity_eligibility": self.identity["eligibility"]}
        self.rows, records = [], []
        for i, char in enumerate(characters):
            oid, rid = "occ_" + sha(session + str(i)), "review-" + sha(session + str(i))
            row = {"occurrence_id": oid, "review_id": rid, "實體頁碼": 1, "字元": "字",
                   "font_xref": xref, "glyph_id_字形索引": char[1], "注音元件ID": char[1],
                   "TTF字形SHA256": "" if cff else self.exact["glyph_sha256"],
                   "CFF整字字形SHA256": self.exact["glyph_sha256"] if cff else "",
                   "CFF樣式群組": self.exact["style_group"],
                   **dict(zip(("x0", "y0", "x1", "y1"), char[3]))}
            if record is not None:
                row.update(font_xref=7, glyph_id_字形索引=0, 注音元件ID=0, x0=10, y0=10, x1=20, y1=30)
            self.rows.append(row)
            records.append({"occurrence_id": oid, "review_id": rid, "pdf_name": self.pdf.name,
                            "pdf_sha256": sha(self.pdf.read_bytes()), "physical_page": 1, "char": "字",
                            **{key: row[key] for key in ("x0", "y0", "x1", "y1")}, "source_record": dict(row),
                            "expected_set": ["ㄆ"], "expected_evidence": "must never be a donor"})
        self.actual = self.root / "01_實際注音" / "actual.xlsx"
        self.actual.parent.mkdir()
        book = Workbook(); sheet = book.active; sheet.title = "實際注音"
        sheet.append(list(self.rows[0]))
        for row in self.rows:
            sheet.append(list(row.values()))
        meta = book.create_sheet("v5.2中繼資料")
        meta.append(["pdf_sha256", sha(self.pdf.read_bytes())])
        meta.append(["workbook_schema_version", WORKBOOK_SCHEMA_VERSION])
        book.save(self.actual); book.close()
        candidate = self.root / "02_候選報告" / "candidate.xlsx"
        candidate.parent.mkdir()
        book = Workbook(); book.save(candidate); book.close()
        self.manifest = {"session_id": session, "session_schema_version": SESSION_SCHEMA_VERSION,
                         "ledger_schema_version": LEDGER_SCHEMA_VERSION, "workbook_schema_version": WORKBOOK_SCHEMA_VERSION,
                         "review_id_schema_version": REVIEW_ID_SCHEMA_VERSION,
                         "pdfs": [{"pdf": str(self.pdf), "pdf_name": self.pdf.name,
                                   "pdf_sha256": sha(self.pdf.read_bytes()), "actual_workbook": str(self.actual),
                                   "actual_workbook_sha256": sha(self.actual.read_bytes()),
                                   "candidate_workbook": str(candidate), "candidate_workbook_sha256": sha(candidate.read_bytes())}],
                         "records": records, "reusable_expected_rules": {"poison": "never interpreted"}}
        seal(self)
        self.learning = dict(glyph_sha256=self.exact["glyph_sha256"], bopomofo=reading,
                             verification_level="USER_VERIFIED_SINGLE" if count == 1 else "VERIFIED_EXACT_GLYPH",
                             source_count=str(count), source_examples="|".join(row["occurrence_id"] for row in self.rows),
                             style_group=self.exact["style_group"])
        self.learning_file = review.USER_CFF_FILE if cff else review.USER_GLYF_FILE
        self.learning_headers = review.USER_CFF_HEADERS if cff else review.USER_GLYF_HEADERS
        for name, headers in ((review.USER_GLYF_FILE, review.USER_GLYF_HEADERS),
                              (review.USER_CFF_FILE, review.USER_CFF_HEADERS),
                              (review.GLYPH_CONFLICT_FILE, review.GLYPH_CONFLICT_HEADERS),
                              (review.GLYPH_PROVENANCE_FILE, review.GLYPH_PROVENANCE_HEADERS),
                              (review.OCCURRENCE_OVERRIDE_FILE, review.OVERRIDE_HEADERS)):
            write_csv(self.evidence / name, headers, [])
        self.save_learning()

    def save_learning(self, rows=None):
        write_csv(self.evidence / self.learning_file, self.learning_headers,
                  [self.learning] if rows is None else rows)

    def conflict(self, *, ids="", readings="ㄅ|ㄆ"):
        write_csv(self.evidence / review.GLYPH_CONFLICT_FILE, review.GLYPH_CONFLICT_HEADERS, [{
            "kind": self.exact["kind"], "style_group": self.exact["style_group"],
            "glyph_sha256": self.exact["glyph_sha256"], "status": "GLYPH_TRUTH_CONFLICT",
            "readings": readings, "source_occurrence_ids": ids}])


def concurrent_import(root, items, ready, output):
    try:
        ready.wait(30)
        receipts = lib.deliver_global_migration(lib.GlobalExactGlyphRepository.resolved(root, busy_timeout_ms=10000), items)
        output.put(("ok", [r.receipt_digest for r in receipts]))
    except Exception as exc:
        output.put(("error", type(exc).__name__, str(exc)))


class MigrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="phase5-")
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.project = Project(self.base / "project")
        self.repo = lib.GlobalExactGlyphRepository.resolved(self.base / "global")

    def plan(self, project=None):
        return migration.build_migration_plan((project or self.project).root)

    def apply(self, project=None):
        return migration.migrate_project((project or self.project).root, global_library_root=self.repo.path.parent, apply=True)

    def assert_blocked(self):
        before = tree_bytes(self.project.root)
        with self.assertRaises((ValueError, lib.GlobalLibraryError, KeyError, OSError)):
            self.apply()
        self.assertFalse(self.repo.path.exists())
        self.assertEqual(tree_bytes(self.project.root), before)

    def test_dry_run_is_read_only_absent_and_no_historical_override_repair(self):
        write_csv(self.project.evidence / review.OCCURRENCE_OVERRIDE_FILE, review.OVERRIDE_HEADERS,
                  [dict(review.KNOWN_BAD_LEGACY_OVERRIDE, source="2026-08-20 GPT 原頁回標")])
        before = tree_bytes(self.project.root)
        with patch.object(review, "ensure_user_evidence_files", side_effect=AssertionError), \
             patch.object(review, "_migrate_v540_known_bad_override", side_effect=AssertionError), \
             patch.object(Path, "rglob", side_effect=AssertionError), \
             patch.object(lib.GlobalExactGlyphRepository, "initialize", side_effect=AssertionError):
            result = migration.migrate_project(self.project.root, global_library_root=self.repo.path.parent)
        self.assertEqual(result["mode"], "DRY_RUN")
        self.assertEqual(result["global_store_status"], lib.ABSENT)
        self.assertFalse(self.repo.path.parent.exists())
        self.assertEqual(tree_bytes(self.project.root), before)

    def test_apply_preserves_source_bytes_and_cannot_put_global_inside_project(self):
        before = tree_bytes(self.project.root)
        self.apply()
        self.assertEqual(tree_bytes(self.project.root), before)
        with self.assertRaises(ValueError):
            migration.migrate_project(self.project.root, global_library_root=self.project.root / "global", apply=True)
        self.assertEqual(tree_bytes(self.project.root), before)

    def test_partial_reconfirmation_mapping_and_deduped_source_examples(self):
        self.project.learning.update(verification_level="VERIFIED_EXACT_GLYPH", source_count="2",
            source_examples=self.project.rows[0]["occurrence_id"] + "|occ_" + "a" * 64)
        self.project.save_learning()
        plan = self.plan()
        self.assertEqual(plan["intents"][0]["payload"]["status"], lib.MIGRATION_CANDIDATE)
        self.assertEqual(len(plan["reconfirmation_targets"]), 1)
        self.project.learning["source_examples"] += "|" + self.project.rows[0]["occurrence_id"]
        self.project.save_learning()
        self.assertEqual(len(self.plan()["reconfirmation_targets"]), 1)

    def test_current_source_example_cannot_point_to_another_identity(self):
        self.project = Project(self.base / "two", count=2)
        # Second current row is a non-exact occurrence; its old sample ID must
        # not be quietly treated as a missing historical sample.
        book = load_workbook(self.project.actual)
        sheet = book["實際注音"]
        headers = [cell.value for cell in sheet[1]]
        sheet.cell(3, headers.index("TTF字形SHA256") + 1).value = None
        book.save(self.project.actual); book.close()
        self.project.manifest["records"][1]["source_record"]["TTF字形SHA256"] = ""
        self.project.manifest["pdfs"][0]["actual_workbook_sha256"] = sha(self.project.actual.read_bytes())
        seal(self.project)
        self.assert_blocked()

    def test_missing_numeric_sealed_locator_does_not_equal_zero(self):
        self.project.manifest["records"][0]["source_record"].pop("font_xref")
        seal(self.project); self.assert_blocked()

    def test_source_changes_during_validation_block_before_global_mutation(self):
        original = migration._occurrence_index
        def change(*args):
            result = original(*args)
            self.project.learning["notes"] = "concurrent edit"
            self.project.save_learning()
            return result
        with patch.object(migration, "_occurrence_index", side_effect=change), self.assertRaises(ValueError):
            self.apply()
        self.assertFalse(self.repo.path.exists())

    def test_pending_project_transaction_blocks_readonly(self):
        (self.project.evidence / promotion.PROJECT_TRANSACTION_FILE).write_text("{}")
        self.assert_blocked()

    def test_single_and_verified_are_nonreusable_with_no_source_or_approval(self):
        for count in (1, 2, 3):
            with self.subTest(count=count):
                project = Project(self.base / f"p{count}", session=f"session-{count}", count=count)
                with patch.object(lib, "approve_global_exact_glyph", side_effect=AssertionError):
                    self.apply(project)
        self.assertEqual(sql_rows(self.repo, "source_evidence"), [])
        self.assertEqual(sql_rows(self.repo, "promotion_approval"), [])
        self.assertEqual(len(sql_rows(self.repo, "migration_candidate")), 3)
        for glyph in sql_rows(self.repo, "glyph_truth"):
            self.assertEqual((glyph["state"], glyph["direct_source_count"], glyph["independent_source_count"], glyph["active_reading"]),
                             (lib.CANDIDATE, 0, 0, None))
        self.assertFalse(self.repo.load_snapshot().trusted_identities)

    def test_learning_validation_is_whole_project_before_write(self):
        original = dict(self.project.learning)
        cases = [dict(source_count="2"), dict(source_count="0"), dict(source_count="01"),
                 dict(source_examples="bad-id"), dict(verification_level="VERIFIED_EXACT_GLYPH"),
                 dict(verification_level="FUTURE"), dict(bopomofo="not bopomofo"), dict(bopomofo=" ㄅ"),
                 dict(glyph_sha256="A" * 64), dict(glyph_sha256="a" * 63)]
        for changes in cases:
            with self.subTest(changes=changes):
                self.project.learning = original | changes
                self.project.save_learning()
                self.assert_blocked()

    def test_duplicate_canonical_row_rejects_even_different_notes(self):
        self.project.save_learning([self.project.learning, self.project.learning | {"notes": "different"}])
        self.assert_blocked()

    def test_missing_csv_and_bad_headers_do_not_seed_files(self):
        path = self.project.evidence / review.USER_GLYF_FILE
        raw = path.read_bytes()
        path.unlink(); self.assert_blocked()
        path.write_bytes(raw.replace(b"glyph_sha256", b"wrong_header")); self.assert_blocked()

    def test_bad_manifest_schema_seal_duplicate_keys_and_roster(self):
        for field, value in (("session_schema_version", "99.0"), ("session_id", " ../path"),
                             ("review_id_schema_version", "future"), ("pdfs", [])):
            old = copy.deepcopy(self.project.manifest)
            self.project.manifest[field] = value; seal(self.project)
            self.assert_blocked()
            self.project.manifest = old; seal(self.project)
        path = self.project.root / "校對工作階段.json"
        path.write_bytes(b'{"session_id":"a","session_id":"b"}'); self.assert_blocked()

    def test_pdf_missing_mismatch_and_no_parent_scan(self):
        self.project.pdf.rename(self.base / "book.pdf")
        self.assert_blocked()
        result = migration.build_migration_plan(self.project.root, pdf_roots=[self.base])
        self.assertEqual(len(result["intents"]), 1)
        (self.base / "book.pdf").write_bytes(b"wrong PDF")
        with self.assertRaises(ValueError):
            migration.build_migration_plan(self.project.root, pdf_roots=[self.base])
        with self.assertRaises(ValueError):
            migration.build_migration_plan("relative-project")

    def test_changed_workbook_and_stale_source_block(self):
        self.project.actual.write_bytes(b"broken")
        self.assert_blocked()
        self.project.manifest["pdfs"][0]["actual_workbook_sha256"] = sha(b"broken")
        seal(self.project)
        with self.assertRaises(Exception):
            self.apply()
        self.assertFalse(self.repo.path.exists())

    def test_unprovable_exact_identity_and_stale_sealed_locator_block(self):
        self.project.learning["glyph_sha256"] = "f" * 64
        self.project.save_learning(); self.assert_blocked()
        self.project.learning["glyph_sha256"] = self.project.exact["glyph_sha256"]
        self.project.save_learning()
        self.project.manifest["records"][0]["x0"] += 1
        seal(self.project); self.assert_blocked()

    def test_structurally_ineligible_skip_not_global_identity(self):
        for i, record in enumerate((composite_glyph(), b"\0" * 10, b"\0", simple_glyph()[:-3])):
            with self.subTest(i=i):
                project = Project(self.base / f"unsupported{i}", record=record)
                font = synthetic_sfnt([record], font_name_marker=b"test")
                with patch.object(migration.fitz, "open", return_value=FakeDocument(font)):
                    result = self.apply(project)
                self.assertEqual(result["intents"], [])
                self.assertEqual(result["skipped"][0]["status"], lib.NON_GLOBAL_ELIGIBLE)
                self.assertFalse(self.repo.path.exists())

    def test_cff_complete_sha_and_style_are_separate(self):
        identities = []
        for index, name in enumerate(("DFBiaoKaiZhuIn-W5", "DFHeiZhuIn-W5")):
            project = Project(self.base / f"cff{index}", cff=True, style_name=name, session=f"cff-{index}")
            result = self.apply(project)
            identities.append(result["intents"][0]["payload"]["identity"])
        self.assertEqual(identities[0]["glyph_sha256"], identities[1]["glyph_sha256"])
        self.assertNotEqual(identities[0]["style_group"], identities[1]["style_group"])
        self.assertEqual(len(sql_rows(self.repo, "glyph_truth")), 2)

    def test_cff_invalid_and_stale_style_fail(self):
        self.project = Project(self.base / "cff", cff=True)
        for style in ("", " BIAOKAI_W5", "HEI_W5"):
            self.project.learning["style_group"] = style
            self.project.save_learning(); self.assert_blocked()

    def test_replay_and_changed_source_bytes_append_not_overwrite(self):
        first = self.apply(); before = self.repo.load_snapshot()
        rows = sql_rows(self.repo, "provenance_event")
        second = self.apply()
        self.assertTrue(second["receipts"][0]["already_processed"])
        self.assertEqual(first["receipts"][0]["receipt_digest"], second["receipts"][0]["receipt_digest"])
        self.assertEqual(self.repo.load_snapshot().generation, before.generation)
        self.assertEqual(sql_rows(self.repo, "provenance_event"), rows)
        self.project.learning["notes"] = "audit change"
        self.project.save_learning(); third = self.apply()
        self.assertNotEqual(first["intents"][0]["intent_id"], third["intents"][0]["intent_id"])
        self.assertEqual(len(sql_rows(self.repo, "migration_candidate")), 2)
        self.assertEqual(self.repo.load_snapshot().generation, before.generation + 1)

    def test_copy_relocation_keeps_import_identity_and_project_identity(self):
        first = self.plan()
        alias = fixture_path_alias(self.base)
        self.assertNotEqual(str(alias), str(self.base.resolve()))
        self.assertEqual(alias.resolve(), self.base.resolve())
        moved = alias / "copied"
        shutil.copytree(self.project.root, moved)
        self.project.pdf.unlink()  # Original stored path no longer resolves.
        copied = migration.build_migration_plan(moved)
        # Relocation changes presentation-only source paths, never the logical
        # plan, canonical import payloads or deterministic import identity.
        audit_paths = {"resolved_pdfs", "source_input_hashes"}
        self.assertEqual({key: value for key, value in first.items() if key not in audit_paths},
                         {key: value for key, value in copied.items() if key not in audit_paths})
        self.assertEqual(copied["resolved_pdfs"]["book.pdf"], str((moved / "book.pdf").resolve()))
        self.assertTrue(all(Path(path).is_relative_to(moved.resolve()) for path in copied["source_input_hashes"]))

    def test_insufficient_and_reconfirmation_do_not_create_decisions(self):
        plan = self.plan()
        target = plan["reconfirmation_targets"][0]
        self.assertEqual(target["occurrence_id"], self.project.rows[0]["occurrence_id"])
        self.assertNotIn("reading", target)
        self.project.learning["source_examples"] = "occ_" + "a" * 64
        self.project.save_learning()
        result = self.apply()
        self.assertEqual(result["reconfirmation_targets"], [])
        self.assertEqual(result["intents"][0]["payload"]["status"], lib.MIGRATION_INSUFFICIENT)
        self.assertEqual(sql_rows(self.repo, "source_evidence"), [])

    def test_fresh_phase4_confirmation_is_only_new_direct_evidence(self):
        result = self.apply()
        row = self.project.rows[0]
        member = self.project.manifest["records"][0] | {"state": "ACTUAL_UNRESOLVED", "pdf": str(self.project.pdf)}
        group = review.build_actual_group_for_entry([member], member)
        context = {"session_id": self.project.manifest["session_id"], "pdfs": {"book.pdf": {
            "pdf": str(self.project.pdf), "pdf_sha256": sha(self.project.pdf.read_bytes())}}}
        items, admissions = promotion.direct_visual_intents(group, "ㄅ", [row["occurrence_id"]], context)
        self.assertEqual(len(items), 1, admissions)
        lib.deliver_global_direct_evidence(self.repo, items)
        evidence = sql_rows(self.repo, "source_evidence")
        self.assertEqual(len(evidence), 1)
        self.assertEqual((evidence[0]["evidence_class"], evidence[0]["counts_toward_global_quorum"]),
                         (lib.DIRECT_VISUAL_ACTUAL, 1))
        self.assertEqual(sql_rows(self.repo, "glyph_truth")[0]["state"], lib.CANDIDATE)
        self.assertEqual(len(sql_rows(self.repo, "migration_candidate")), 1)

    def test_bidirectional_legacy_direct_and_legacy_legacy_conflicts(self):
        for order in ("legacy-direct", "direct-legacy", "legacy-legacy"):
            with self.subTest(order=order):
                self.repo = lib.GlobalExactGlyphRepository.resolved(self.base / order)
                direct = intent(identity=self.project.exact, reading="ㄆ")
                if order == "direct-legacy":
                    lib.deliver_global_direct_evidence(self.repo, [direct])
                self.apply()
                if order == "legacy-direct":
                    lib.deliver_global_direct_evidence(self.repo, [direct])
                if order == "legacy-legacy":
                    other = Project(self.base / "other", session="other", reading="ㄆ")
                    self.apply(other)
                self.assertEqual(sql_rows(self.repo, "glyph_truth")[0]["state"], lib.QUARANTINED_CONFLICT)
                self.assertEqual(self.repo.load_snapshot().quarantined_identities[0].conflicting_readings, ("ㄅ", "ㄆ"))
                self.assertTrue(all(row["status"] == lib.MIGRATION_CONFLICT for row in sql_rows(self.repo, "migration_candidate")))

    def test_verified_matching_audit_keeps_digest_revision_approval_conflict_revokes(self):
        lib.deliver_global_direct_evidence(self.repo, [intent(1, identity=self.project.exact), intent(2, identity=self.project.exact)])
        approve(self.repo)
        before = sql_rows(self.repo, "glyph_truth")[0]
        exact = lib.canonical_global_exact_identity(**self.project.exact)
        snapshot = self.repo.load_snapshot()
        digest = lib.canonical_logical_subset_digest(snapshot, [exact]).sha256
        self.apply()
        self.assertEqual(sql_rows(self.repo, "glyph_truth")[0], before)
        self.assertEqual(lib.canonical_logical_subset_digest(self.repo.load_snapshot(), [exact]).sha256, digest)
        self.project.learning["bopomofo"] = "ㄆ"; self.project.save_learning(); self.apply()
        glyph = sql_rows(self.repo, "glyph_truth")[0]
        self.assertEqual((glyph["state"], glyph["active_reading"]), (lib.QUARANTINED_CONFLICT, None))
        self.assertEqual({r["status"] for r in sql_rows(self.repo, "promotion_approval")}, {lib.APPROVED, lib.REVOKED})
        self.assertNotEqual(lib.canonical_logical_subset_digest(self.repo.load_snapshot(), [exact]).sha256, digest)

    def test_project_conflict_all_readings_no_sources_and_never_auto_recovers(self):
        self.project.learning["verification_level"] = lib.QUARANTINED_CONFLICT
        self.project.save_learning()
        self.project.conflict()
        self.apply()
        self.assertEqual({r["reading"] for r in sql_rows(self.repo, "migration_candidate")}, {"ㄅ", "ㄆ"})
        self.assertEqual(sql_rows(self.repo, "source_evidence"), [])
        self.assertEqual(len(sql_rows(self.repo, "glyph_conflict")), 1)
        self.project.conflict(readings="ㄅ|ㄇ")
        self.apply()
        lib.deliver_global_direct_evidence(self.repo, [intent(1, identity=self.project.exact, reading="ㄅ")])
        self.assertEqual(self.repo.load_snapshot().quarantined_identities[0].conflicting_readings, ("ㄅ", "ㄆ", "ㄇ"))
        self.assertEqual(sql_rows(self.repo, "glyph_truth")[0]["direct_source_count"], 1)

    def test_invalid_project_conflicts_block_whole_project(self):
        for readings in ("ㄅ", "ㄅ|ㄅ", "ㄅ|invalid"):
            self.project.conflict(readings=readings)
            self.assert_blocked()

    def test_duplicate_conflict_and_bad_kind_style_sha_status_are_rejected(self):
        self.project.conflict()
        path = self.project.evidence / review.GLYPH_CONFLICT_FILE
        with path.open(encoding="utf-8-sig", newline="") as handle:
            row = next(csv.DictReader(handle))
        for changes in ({"kind": "FUTURE"}, {"style_group": "invalid-ttf-style"},
                        {"glyph_sha256": "A" * 64}, {"status": "RESOLVED"}):
            write_csv(path, review.GLYPH_CONFLICT_HEADERS, [row | changes])
            self.assert_blocked()
        write_csv(path, review.GLYPH_CONFLICT_HEADERS, [row, row | {"notes": "duplicate"}])
        self.assert_blocked()

    def test_audit_readings_and_workbook_values_are_not_donors(self):
        write_csv(self.project.evidence / review.OCCURRENCE_OVERRIDE_FILE, review.OVERRIDE_HEADERS,
                  [{"actual_reading": "ㄆ", "source": "propagated"}])
        write_csv(self.project.evidence / review.GLYPH_PROVENANCE_FILE, review.GLYPH_PROVENANCE_HEADERS,
                  [{"bopomofo": "ㄇ", "event_id": "historical", "source": "audit"}])
        plan = self.plan()
        self.assertEqual([item["payload"]["reading"] for item in plan["intents"]], ["ㄅ"])
        self.project.save_learning([])
        self.assertEqual(self.plan()["intents"], [])

    def test_pure_legacy_never_changes_independent_quorum(self):
        before = lib.independent_direct_quorum([intent()["payload"]["evidence"] | {
            "evidence_id": "a" * 64, "counts_toward_global_quorum": 1}])
        self.apply()
        self.assertEqual(before, ())
        self.assertEqual(lib.independent_direct_quorum(sql_rows(self.repo, "source_evidence")), ())

    def test_cli_default_dry_run_explicit_apply_and_invalid_project(self):
        args = ["--project-output", str(self.project.root), "--global-library-root", str(self.repo.path.parent)]
        with redirect_stdout(io.StringIO()) as out:
            self.assertEqual(cli.main(args), 0)
        self.assertEqual(json.loads(out.getvalue())["mode"], "DRY_RUN")
        self.assertFalse(self.repo.path.exists())
        with redirect_stdout(io.StringIO()):
            self.assertEqual(cli.main(args + ["--apply"]), 0)
        with redirect_stderr(io.StringIO()):
            self.assertEqual(cli.main(["--project-output", "relative"]), 1)

    def test_payload_unknown_fields_and_bad_last_intent_fail_before_initialize(self):
        good = self.plan()["intents"][0]
        for field in ("expected", "expected_set", "dictionary", "absolute_path", "semantic_context"):
            bad = copy.deepcopy(good); bad["payload"][field] = "poison"
            with self.assertRaises(lib.GlobalLibraryError):
                lib.deliver_global_migration(self.repo, [good, bad])
            self.assertFalse(self.repo.path.exists())

    def test_global_failure_rolls_back_whole_batch(self):
        self.project.conflict()
        plan = self.plan()
        self.repo.initialize()
        before = {name: sql_rows(self.repo, name) for name in lib._REQUIRED_COLUMNS}
        original = lib._ingest_migration
        def fail(tx, item):
            original(tx, item)
            raise sqlite3.OperationalError("injected failure after mutation")
        with patch.object(lib, "_ingest_migration", side_effect=fail), self.assertRaises(lib.GlobalLibraryError):
            lib.deliver_global_migration(self.repo, plan["intents"])
        self.assertEqual({name: sql_rows(self.repo, name) for name in lib._REQUIRED_COLUMNS}, before)

    def test_late_conflict_and_receipt_failure_roll_back_imports_generation_approval(self):
        self.project.conflict()
        plan = self.plan()
        for stage in ("conflict", "receipt"):
            self.repo = lib.GlobalExactGlyphRepository.resolved(self.base / stage)
            lib.deliver_global_direct_evidence(self.repo, [intent(1, identity=self.project.exact), intent(2, identity=self.project.exact)])
            approve(self.repo)
            before = {name: sql_rows(self.repo, name) for name in lib._REQUIRED_COLUMNS}
            if stage == "conflict":
                original = lib._migration_conflict_gate
                def failure(*args, **kwargs):
                    original(*args, **kwargs)
                    raise sqlite3.OperationalError("after conflict")
                patcher = patch.object(lib, "_migration_conflict_gate", side_effect=failure)
            else:
                original = lib.GlobalWriteTransaction.record_processed_intent
                def failure(*args, **kwargs):
                    original(*args, **kwargs)
                    raise sqlite3.OperationalError("after receipt")
                patcher = patch.object(lib.GlobalWriteTransaction, "record_processed_intent", failure)
            with patcher, self.assertRaises(lib.GlobalLibraryError):
                lib.deliver_global_migration(self.repo, plan["intents"])
            self.assertEqual({name: sql_rows(self.repo, name) for name in lib._REQUIRED_COLUMNS}, before)

    def test_partial_conflict_and_multiple_projects_reject_before_initialize(self):
        self.project.conflict()
        batch = self.plan()["intents"]
        conflict = next(item for item in batch if item["payload"]["old_verification_level"] == lib.QUARANTINED_CONFLICT)
        other = Project(self.base / "other", session="other")
        for values in ([conflict], batch + self.plan(other)["intents"]):
            with self.assertRaises(lib.GlobalLibraryError):
                lib.deliver_global_migration(self.repo, values)
            self.assertFalse(self.repo.path.exists())

    def test_strict_db_reader_rejects_legacy_conflict_gate_bypass(self):
        self.apply()
        other = Project(self.base / "other", session="other", reading="ㄆ")
        p = self.plan(other)["intents"][0]["payload"]
        columns = lib._REQUIRED_COLUMNS["migration_candidate"]
        row = {key: p[key] for key in columns if key != "imported_at"} | {"imported_at": lib._utc_now()}
        with closing(sqlite3.connect(self.repo.path)) as connection:
            connection.execute("INSERT INTO migration_candidate (" + ",".join(columns) + ") VALUES (" +
                               ",".join("?" for _ in columns) + ")", tuple(row[key] for key in columns))
            connection.commit()
        with self.assertRaises(lib.GlobalLibraryValidationError):
            self.repo.load_snapshot()

    def test_all_replay_checks_before_first_mutation(self):
        self.apply()
        old = self.plan()["intents"][0]
        changed = lib.make_migration_intent(**{key: old["payload"][key] for key in (
            "identity", "reading", "legacy_project_id", "legacy_project_sha256", "legacy_evidence_sha256",
            "old_verification_level")}, status=lib.MIGRATION_INSUFFICIENT)
        fresh = lib.make_migration_intent(**{key: old["payload"][key] for key in (
            "identity", "reading", "legacy_project_id", "legacy_project_sha256", "old_verification_level", "status")},
            legacy_evidence_sha256="f" * 64)
        with patch.object(lib, "_ingest_migration", side_effect=AssertionError), self.assertRaises(lib.GlobalLibraryIntentConflictError):
            lib.deliver_global_migration(self.repo, [fresh, changed])
        self.assertEqual(len(sql_rows(self.repo, "migration_candidate")), 1)

    def test_old_draft_exact_id_is_readable_without_rewrite(self):
        payload = self.plan()["intents"][0]["payload"]
        self.repo.initialize()
        with self.repo.write_transaction() as tx:
            tx.insert_candidate_identity(self.project.exact)
        columns = lib._REQUIRED_COLUMNS["migration_candidate"]
        row = {key: payload[key] for key in columns if key != "imported_at"} | {"imported_at": lib._utc_now()}
        args = {key: row[key] for key in ("glyph_id", "reading", "legacy_project_id", "legacy_project_sha256",
                                        "legacy_evidence_sha256", "old_verification_level")}
        old_id = lib._legacy_v0_migration_import_id(**args)
        self.assertNotEqual(old_id, row["import_id"])
        row["import_id"] = old_id
        with closing(sqlite3.connect(self.repo.path)) as conn:
            conn.execute("INSERT INTO migration_candidate (" + ",".join(columns) + ") VALUES (" +
                         ",".join("?" for _ in columns) + ")", tuple(row[key] for key in columns))
            conn.commit()
        self.assertEqual(self.repo.load_snapshot().store_status, lib.VALID)
        self.assertEqual(sql_rows(self.repo, "migration_candidate")[0]["import_id"], old_id)
        with closing(sqlite3.connect(self.repo.path)) as conn:
            conn.execute("UPDATE migration_candidate SET import_id=?", ("a" * 64,)); conn.commit()
        with self.assertRaises(lib.GlobalLibraryError):
            self.repo.load_snapshot()

    def test_current_migration_requires_linked_import_receipt_and_provenance(self):
        for table in ("migration_candidate", "processed_intent", "provenance_event"):
            self.repo = lib.GlobalExactGlyphRepository.resolved(self.base / table)
            self.apply()
            with closing(sqlite3.connect(self.repo.path)) as connection:
                connection.execute("DELETE FROM " + table)
                connection.commit()
            with self.assertRaises(lib.GlobalLibraryValidationError):
                self.repo.load_snapshot()
            with self.assertRaises(lib.GlobalLibraryValidationError):
                self.apply()

    def test_deterministic_row_key_and_current_contract(self):
        p = self.plan()["intents"][0]["payload"]
        self.assertEqual(p["legacy_row_key"], lib._canonical_sha256({key: p[key] for key in (
            "glyph_id", "reading", "old_verification_level")}))
        args = {key: p[key] for key in ("glyph_id", "reading", "legacy_project_id", "legacy_project_sha256",
                                      "legacy_evidence_sha256", "old_verification_level")}
        with patch.object(lib, "GLOBAL_LIBRARY_SCHEMA_VERSION", "future"):
            self.assertEqual(lib.compute_migration_import_id(**args), p["import_id"])
        with patch.object(lib, "compute_legacy_row_key", return_value="f" * 64):
            self.assertNotEqual(lib.compute_migration_import_id(**args), p["import_id"])

    def test_corrupt_and_unknown_schema_db_never_recreated(self):
        self.repo.path.parent.mkdir()
        self.repo.path.write_bytes(b"corrupt sqlite")
        with self.assertRaises(lib.GlobalLibraryError):
            self.apply()
        self.assertEqual(self.repo.path.read_bytes(), b"corrupt sqlite")
        self.repo = lib.GlobalExactGlyphRepository.resolved(self.base / "future")
        self.repo.initialize()
        with closing(sqlite3.connect(self.repo.path)) as conn:
            conn.execute("PRAGMA user_version=999"); conn.commit()
        with self.assertRaises(lib.GlobalLibraryError):
            self.apply()

    def test_concurrent_same_and_contradictory_imports(self):
        first = self.plan()["intents"]
        other = Project(self.base / "other", session="other", reading="ㄆ")
        second = self.plan(other)["intents"]
        ctx = multiprocessing.get_context("spawn")
        for name, batches in (("same", [first, first]), ("different", [first, second])):
            root = self.base / name
            barrier, queue = ctx.Barrier(2), ctx.Queue()
            processes = [ctx.Process(target=concurrent_import, args=(str(root), batch, barrier, queue)) for batch in batches]
            try:
                for process in processes: process.start()
                results = [queue.get(timeout=45) for _ in processes]
                for process in processes:
                    process.join(45); self.assertEqual(process.exitcode, 0)
                self.assertTrue(all(result[0] == "ok" for result in results), results)
                repo = lib.GlobalExactGlyphRepository.resolved(root)
                if name == "same":
                    self.assertEqual(results[0][1], results[1][1])
                    self.assertEqual(len(sql_rows(repo, "migration_candidate")), 1)
                    self.assertEqual(repo.load_snapshot().generation, 1)
                else:
                    self.assertEqual(repo.load_snapshot().quarantined_identities[0].conflicting_readings, ("ㄅ", "ㄆ"))
            finally:
                for process in processes:
                    if process.is_alive(): process.terminate(); process.join()
                queue.close(); queue.join_thread()


class MigrationArchitectureTests(unittest.TestCase):
    PHASE5_BASELINE = "d2d558808bd2902857fad30f824d8bdc352beb73"

    @staticmethod
    def _standalone_fingerprint_contract(tree):
        """Freeze Phase 5 cache boundaries, not unrelated manual/report code.

        The producer/comparator modules remain byte-pinned below. Here retain
        their wiring, dependency declaration and the two fail-closed reuse
        guards. For larger application functions, compare only contract calls
        and their arguments, allowing other human-review/report statements.
        """
        declarations = {"ACTUAL_DECODER_SOURCE_FILES", "COMPATIBLE_SESSION_VERSIONS"}
        guards = {"output_is_reusable", "candidate_is_baseline_reusable"}
        calls = {
            "compute_actual_asset_fingerprint", "compute_expected_asset_fingerprint",
            "rewrite_actual_workbook_fingerprint", "fingerprint_compatible",
            "schema_compatible", "review_id_schema_compatible", "validate_asset_manifest",
            "actual_workbook_dynamic_dependencies", "actual_workbook_global_exact_dependencies",
            "global_exact_glyph_evidence_hashes",
        }
        imports = calls | {
            "EXPECTED_RESOLVER_SOURCE_FILES", "LEDGER_SCHEMA_VERSION", "SESSION_SCHEMA_VERSION",
            "WORKBOOK_SCHEMA_VERSION", "REVIEW_ID_SCHEMA_VERSION", "GlobalExactGlyphRepository",
        }
        return {
            "declarations": [ast.dump(node) for node in tree.body if isinstance(node, ast.Assign)
                             and any(isinstance(target, ast.Name) and target.id in declarations for target in node.targets)],
            "imports": [(node.module, alias.name, alias.asname) for node in tree.body
                        if isinstance(node, ast.ImportFrom) for alias in node.names if alias.name in imports],
            "reuse_guards": {node.name: ast.dump(node) for node in tree.body
                             if isinstance(node, ast.FunctionDef) and node.name in guards},
            "contract_calls": {node.name: selected for node in tree.body if isinstance(node, ast.FunctionDef)
                               if (selected := [ast.dump(call) for call in ast.walk(node)
                                                if isinstance(call, ast.Call) and isinstance(call.func, ast.Name)
                                                and call.func.id in calls])},
        }

    @classmethod
    def _historical_standalone(cls):
        return ast.parse(subprocess.check_output(
            ["git", "show", cls.PHASE5_BASELINE + ":standalone_proofread.py"], cwd=ROOT,
        ).decode("utf-8"))

    def _assert_contract_versions(self):
        import cross_version_compat as compatibility
        self.assertEqual((lib.GLOBAL_LIBRARY_SCHEMA_VERSION, lib.GLOBAL_LIBRARY_SQLITE_USER_VERSION), ("1.0", 1))
        self.assertEqual(lib.GLOBAL_MIGRATION_IMPORT_CONTRACT_VERSION, "1.0")
        self.assertEqual(lib.GLOBAL_EXACT_DEPENDENCY_CONTRACT_VERSION, "1.0")
        self.assertEqual(compatibility.ACTUAL_DECODER_SEMANTICS_EPOCH, "1")
        self.assertEqual(compatibility.EXPECTED_RESOLVER_SEMANTICS_EPOCH, "1")

    def test_import_chain_has_no_expected_resolver(self):
        code = "import legacy_global_migration, sys; assert 'standalone_proofread' not in sys.modules; assert 'check_pronunciation_candidates' not in sys.modules"
        result = subprocess.run([sys.executable, "-B", "-c", code], cwd=ROOT, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        source = (ROOT / "legacy_global_migration.py").read_text(encoding="utf-8")
        tree = ast.parse(source)
        forbidden = {"rglob", "walk", "materialize_ledger", "validate_dynamic_actual_evidence", "ensure_user_evidence_files",
                     "_resolve_session_pdfs", "direct_visual_intents", "approve_global_exact_glyph"}
        for node in ast.walk(tree):
            if isinstance(node, ast.Call):
                self.assertNotIn(getattr(node.func, "attr", getattr(node.func, "id", "")), forbidden)
            if isinstance(node, ast.Constant) and isinstance(node.value, str):
                self.assertNotIn(node.value, {"expected_set", "expected_evidence", "reusable_expected_rules", "dictionary"})

    def test_schema_and_fingerprint_boundaries_unchanged(self):
        import runtime_source_validation as runtime
        self._assert_contract_versions()
        self.assertTrue(runtime.validate_asset_manifest(ROOT)["ok"])
        for name in ("runtime_asset_manifest.json", "runtime_source_validation.py", "cross_version_compat.py",
                     "VERSION.txt"):
            committed = subprocess.check_output(["git", "show", self.PHASE5_BASELINE + ":" + name], cwd=ROOT)
            self.assertEqual((ROOT / name).read_bytes().replace(b"\r\n", b"\n"), committed.replace(b"\r\n", b"\n"))
        current = ast.parse((ROOT / "standalone_proofread.py").read_text(encoding="utf-8"))
        self.assertEqual(self._standalone_fingerprint_contract(current),
                         self._standalone_fingerprint_contract(self._historical_standalone()))

    def test_unrelated_manual_and_report_changes_do_not_change_contract(self):
        historical = self._historical_standalone()
        changed = copy.deepcopy(historical)
        for node in changed.body:
            if isinstance(node, ast.FunctionDef) and node.name in {"_apply_review_event", "generate_report"}:
                node.body.append(ast.Pass())
        changed.body.extend(ast.parse("def explicit_local_manual_judgment():\n    return 'human operation'\n").body)
        self.assertNotEqual(ast.dump(changed), ast.dump(historical))
        self.assertEqual(self._standalone_fingerprint_contract(changed),
                         self._standalone_fingerprint_contract(historical))

    def test_changed_dependencies_chain_and_fail_open_guard_are_detected(self):
        historical = self._historical_standalone()
        expected = self._standalone_fingerprint_contract(historical)
        for mutation in ("source_dependency", "actual_chain", "per_pdf_dependency", "fail_open_guard", "import_origin"):
            changed = copy.deepcopy(historical)
            if mutation == "source_dependency":
                declaration = next(node for node in changed.body if isinstance(node, ast.Assign)
                                   and any(isinstance(target, ast.Name) and target.id == "ACTUAL_DECODER_SOURCE_FILES" for target in node.targets))
                declaration.value.elts.pop()
            elif mutation == "actual_chain":
                call = next(node for node in ast.walk(changed) if isinstance(node, ast.Call)
                            and isinstance(node.func, ast.Name) and node.func.id == "fingerprint_compatible"
                            and any(kw.arg == "chain" and isinstance(kw.value, ast.Constant) and kw.value.value == "actual" for kw in node.keywords))
                next(kw for kw in call.keywords if kw.arg == "chain").value = ast.Constant(value="expected")
            elif mutation == "per_pdf_dependency":
                call = next(node for node in ast.walk(changed) if isinstance(node, ast.Call)
                            and isinstance(node.func, ast.Name) and node.func.id == "compute_actual_asset_fingerprint")
                next(kw for kw in call.keywords if kw.arg == "global_exact_glyph_evidence_hashes").value = ast.Dict(keys=[], values=[])
            elif mutation == "fail_open_guard":
                guard = next(node for node in changed.body if isinstance(node, ast.FunctionDef) and node.name == "output_is_reusable")
                guard.body = [ast.Return(value=ast.Constant(value=True))]
            else:
                imported = next(node for node in changed.body if isinstance(node, ast.ImportFrom)
                                and any(alias.name == "compute_actual_asset_fingerprint" for alias in node.names))
                imported.module = "unverified_fingerprint_provider"
            with self.subTest(mutation=mutation), self.assertRaises(AssertionError):
                self.assertEqual(self._standalone_fingerprint_contract(changed), expected)

    def test_changed_schema_epoch_and_dependency_versions_are_detected(self):
        import cross_version_compat as compatibility
        cases = [(lib, "GLOBAL_LIBRARY_SCHEMA_VERSION", "2.0"),
                 (lib, "GLOBAL_LIBRARY_SQLITE_USER_VERSION", 2),
                 (lib, "GLOBAL_MIGRATION_IMPORT_CONTRACT_VERSION", "2.0"),
                 (lib, "GLOBAL_EXACT_DEPENDENCY_CONTRACT_VERSION", "2.0"),
                 (compatibility, "ACTUAL_DECODER_SEMANTICS_EPOCH", "2"),
                 (compatibility, "EXPECTED_RESOLVER_SEMANTICS_EPOCH", "2")]
        for module, field, value in cases:
            with self.subTest(field=field), patch.object(module, field, value), self.assertRaises(AssertionError):
                self._assert_contract_versions()
