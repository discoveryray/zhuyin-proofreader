"""Targeted actual-only admission tests; no production data or PDF writes."""
from __future__ import annotations

import csv
import json
import hashlib
import base64
import io
import shutil
from contextlib import closing
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import Workbook, load_workbook

import cross_version_compat as compat
import export_zhuyin_readings as decoder
import runtime_source_validation as source
import actual_review as actual
from ttf_symbol_recombination import ExactSymbolRecombinationModel
from ttf_zhuyin_shape_decoder import contour_group_signature
from ttf_zhuyin_shape_decoder import ValidationResult
from tests.test_global_exact_glyph_read_reuse_v580 import (
    SHA_A, SHA_B, _DecodeDocument, _DecodeTrueTypeInspector, _absent_snapshot,
    _identity, _read_actual_and_audit, _record, _snapshot,
)


def csv_file(path, fields, rows):
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    return path


def check_four_historical_proposals_have_no_truth_or_training_qualification(font, gid, old, tmp_path):
    # Frozen historical assertions keep exercising rejection even when the
    # shipped candidate table is corrected or the ambiguous value is cleared.
    source_font = "YenRoundZuinnEG-Bd" if font == "YenRoundZuinnEG-Md" else font
    path = csv_file(tmp_path / f"legacy-{gid}.csv", ["font", "zhuyin_component_id", "bopomofo", "verification"], [
        {"font": source_font, "zhuyin_component_id": gid, "bopomofo": old, "verification": "historical v1.7 proposal"},
    ])
    exact = decoder.load_map(path)
    groups = decoder.load_compatibility_groups(decoder.DEFAULT_GROUPS)
    grouped, _ = decoder.build_group_map(exact, groups)
    rec, _, _, _ = decoder.lookup_reading(font, gid, exact, grouped, groups)
    assert rec["bopomofo"] == old
    assert rec["candidate_only"]
    assert not rec.get("evidence_roots")
    assert decoder.training_reading([{"initial_bopomofo": old, **rec}]) == ""


def check_all_ten_real_corrections_keep_font_gid_and_exact_hash_scope(index):
    rows = list(csv.DictReader(decoder.DEFAULT_MAPPING_CORRECTIONS.open(encoding="utf-8-sig")))
    row = rows[index]
    corrections = decoder.load_verified_mapping_corrections(decoder.DEFAULT_MAPPING_CORRECTIONS)
    font, gid, sha = row["font"], int(row["component_id"]), row["glyph_sha256"]
    accepted = decoder.lookup_verified_mapping_correction(font, gid, sha, corrections)
    assert accepted["bopomofo"] == row["new_reading"]
    assert accepted["evidence_roots"]
    assert accepted["donation_eligible"] is False
    assert not decoder.lookup_verified_mapping_correction(font, gid, "0" * 64, corrections)
    assert not decoder.lookup_verified_mapping_correction(font + "-other", gid, sha, corrections)
    assert not decoder.lookup_verified_mapping_correction(font, gid + 1, sha, corrections)


def check_calibrated_graphics_keep_original_claims_and_scoped_roles(tmp_path):
    rows = list(csv.DictReader(decoder.DEFAULT_MAPPING_CORRECTIONS.open(encoding="utf-8-sig")))
    # Frozen from baseline 02ac433's parsed original ten assertions; no Git
    # history/network is needed by the test or by runtime.
    canonical = json.dumps(rows[:10], ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    assert hashlib.sha256(canonical).hexdigest() == "07d39c7c514511238bf63682db34d3d2c4bca4e1276157d3a5c1054c264002d9"
    root = Path(decoder.__file__).resolve().parent
    proof_path = root / "docs/evidence/map-glyph-review-20261019/observations.json"
    if not proof_path.exists():
        assert len(rows) == 10  # Released pre-calibration contract only.
        return
    proof = json.loads(proof_path.read_text(encoding="utf-8"))
    assert len(rows) == 10 + proof["new_scope_rows"]
    corrections = decoder.load_verified_mapping_corrections(decoder.DEFAULT_MAPPING_CORRECTIONS)
    for observed in proof["observations"]:
        font, gid, sha = observed["font"], observed["component_id"], observed["glyph_sha256"]
        rec = decoder.lookup_verified_mapping_correction(font, gid, sha, corrections,
                                                        target_identity_result=observed["record_contract"])
        assert rec["bopomofo"] == observed["reading"]
        root_key = tuple(observed.get("source_truth_root_key") or (font, gid, sha))
        assert rec["evidence_roots"] == corrections[root_key]["evidence_roots"]
        assert rec["is_alias"] == (root_key != (font, gid, sha))
        if rec["is_alias"]:
            assert rec["binding_refs"] and rec["source_key"] == root_key
        assert rec["donation_eligible"] is False
        assert observed["human_confirmation"] is False
        assert observed["global_approval"] is False
        assert observed["new_global_independent_source"] is False
        assert not decoder.lookup_verified_mapping_correction(font + "-unbound", gid, sha, corrections)
        assert not decoder.lookup_verified_mapping_correction(font, gid, "0" * 64, corrections)
        for field in ("original_crop", "isolated_svg"):
            assert hashlib.sha256((root / observed[field]).read_bytes()).hexdigest() == observed[field + "_sha256"]


def check_scoped_aliases_inherit_only_current_direct_truth_and_require_target_binding(tmp_path):
    fields = ["font", "component_id", "glyph_sha256", "old_reading", "new_reading",
              "source_book", "evidence", "verification", "notes"]
    direct = {"font": "Source", "component_id": 1, "glyph_sha256": SHA_A,
              "new_reading": "ㄅ", "evidence": "original isolated glyph",
              "verification": "isolated glyph visual + SHA256 exact gate"}
    independent = {**direct, "font": "Independent", "component_id": 2}
    alias = {**direct, "font": "Target", "component_id": 3,
             "verification": "SCOPED_CORRECTION_ALIAS_V1",
             "evidence": json.dumps({"kind": "SCOPED_CORRECTION_ALIAS_V1",
                                     "source_key": ["Source", 1, SHA_A], "target_proof": "original-target-binding"})}
    path = tmp_path / "corrections.csv"
    source_key, target_key = ("Source", 1, SHA_A), ("Target", 3, SHA_A)
    def load(rows):
        return decoder.load_verified_mapping_corrections(csv_file(path, fields, rows))
    admitted = load([direct, alias, independent])
    rec = admitted[target_key]
    assert rec["is_alias"] and rec["source_key"] == source_key
    assert rec["evidence_roots"] == admitted[source_key]["evidence_roots"]
    assert rec["binding_refs"][0] != rec["evidence_roots"][0]
    assert rec["binding_refs"][1] == "original-target-binding"
    assert rec["donation_eligible"] is False
    identity = {"eligibility": decoder.GLOBAL_ELIGIBLE_SIMPLE_GLYF_V1, "glyph_sha256": SHA_A}
    assert decoder.lookup_verified_mapping_correction(*target_key, admitted, target_identity_result=identity) == rec
    assert decoder.lookup_verified_mapping_correction(*source_key, admitted)  # Direct contract remains independent of Global.
    for invalid_identity in (None, {}, {**identity, "glyph_sha256": SHA_B},
                             {**identity, "eligibility": "NON_GLOBAL_ELIGIBLE"}):
        assert decoder.lookup_verified_mapping_correction(*target_key, admitted, target_identity_result=invalid_identity) is None
    # A surviving independent root does not silently replace the specified root.
    for sources in ([], [{**direct, "verification": "REVOKED"}], [{**direct, "new_reading": "ㄆ"}],
                    [{**direct, "glyph_sha256": SHA_B}], [direct, {**direct, "new_reading": "ㄆ"}]):
        after = load([*sources, alias, independent])
        assert target_key not in after
        assert ("Independent", 2, SHA_A) in after
    def bound_row(font, gid, source, *, reading="ㄅ", proof="target"):
        return {**alias, "font": font, "component_id": gid, "new_reading": reading,
                "evidence": json.dumps({"kind": "SCOPED_CORRECTION_ALIAS_V1", "source_key": source, "target_proof": proof})}
    chain = bound_row("Chain", 4, ["Target", 3, SHA_A])
    assert ("Chain", 4, SHA_A) not in load([direct, alias, chain])
    assert not load([bound_row("Target", 3, ["Chain", 4, SHA_A]), chain])
    assert not load([bound_row("Target", 3, ["Target", 3, SHA_A])])
    for invalid in (bound_row("Target", 3, ["Source", True, SHA_A]),
                    bound_row("Target", 3, ["Source", 1, SHA_A], proof=""),
                    {**alias, "evidence": "not json"}, {**alias, "verification": "UNKNOWN_ROLE"}):
        assert target_key not in load([direct, invalid])
    other = {**direct, "font": "Other", "component_id": 8, "new_reading": "ㄆ"}
    conflicting_alias = bound_row("Target", 3, ["Other", 8, SHA_A], reading="ㄆ")
    assert target_key not in load([direct, other, alias, conflicting_alias])
    # A local independently supported direct claim retains its own contract.
    target_direct = {**direct, "font": "Target", "component_id": 3}
    assert load([direct, other, target_direct, conflicting_alias])[target_key]["is_alias"] is False


def check_revised_graphic_claims_do_not_revive_withdrawn_truth(tmp_path):
    root = Path(decoder.__file__).resolve().parent
    relative = "docs/evidence/map-glyph-review-20261019/corrective2/claim-revisions.json"
    proof_path = root / relative
    evidence = json.loads(proof_path.read_text(encoding="utf-8"))
    assert evidence["schema"] == "finite-original-graphic-claim-revisions/1"
    assert evidence["source_head"] == "3c638b3b7d3492dc5d0dc0539aa8f5a5e8f78ece"
    assert evidence["changed_direct_roots"] == len(evidence["revisions"]) == 9
    assert evidence["changed_dependent_aliases"] == 3
    assert evidence["independent_sources_added"] == 0
    frozen = json.loads((root / evidence["final_decisions_ref"]).read_text(encoding="utf-8"))
    assert hashlib.sha256((root / evidence["final_decisions_ref"]).read_bytes()).hexdigest() == evidence["final_decisions_sha256"]
    assert len(frozen["records"]) == 899
    for name, digest in frozen["source_frozen_files"].items():
        assert hashlib.sha256((proof_path.parent / name).read_bytes()).hexdigest() == digest
    assert "/docs/evidence/map-glyph-review-20261019/corrective2/** -text" in (root / ".gitattributes").read_text(encoding="utf-8")
    rows = list(csv.DictReader(decoder.DEFAULT_MAPPING_CORRECTIONS.open(encoding="utf-8-sig")))
    key = lambda r: (r["font"], int(r["component_id"]), r["glyph_sha256"])
    current = decoder.load_verified_mapping_corrections(decoder.DEFAULT_MAPPING_CORRECTIONS)
    bykey = {key(r): r for r in rows}
    root_keys = {tuple(r["root_key"]) for r in evidence["revisions"]}
    targets = {tuple(k) for r in evidence["revisions"] for k in r["current_target_keys"]}
    assert len(targets) == 12
    fingerprints = decoder.load_verified_glyf_fingerprints(decoder.DEFAULT_FINGERPRINTS)
    outlines = decoder.load_verified_outline_signatures(decoder.DEFAULT_OUTLINE_SIGNATURES)
    symbols = ExactSymbolRecombinationModel.load_csv(decoder.DEFAULT_SYMBOL_TEMPLATES)
    legacy_rows = []
    from exact_glyph_identity import classify_ttf_glyf_record
    for revision in evidence["revisions"]:
        source_key = tuple(revision["root_key"])
        reading, withdrawn = revision["current_reading"], revision["withdrawn_reading"]
        assert reading != withdrawn
        assert current[source_key]["bopomofo"] == reading
        assert current[source_key]["donation_eligible"] is False
        for prior in revision["prior_rows"]:
            old = prior["row"]
            assert old["new_reading"] == withdrawn
            assert decoder._row_claim_ref(decoder.DEFAULT_MAPPING_CORRECTIONS, old,
                    asset_sha256=evidence["prior_asset_hashes"]["ttf_verified_mapping_corrections.csv"]) == prior["claim_ref"]
            assert all(bykey[key(old)][column] == value for column, value in old.items() if column != "new_reading")
        for target_key in revision["current_target_keys"]:
            rec = current[tuple(target_key)]
            assert rec["bopomofo"] == reading
            assert rec["evidence_roots"] == current[source_key]["evidence_roots"]
            assert revision["withdrawn_truth_root_receipt"] not in rec["evidence_roots"]
        for graphic in revision["graphic_evidence"].values():
            assert hashlib.sha256((root / graphic["path"]).read_bytes()).hexdigest() == graphic["sha256"]
        # Exact raw source bytes permit independent entrance regression in CI
        # without the private original PDF or a guessed replacement geometry.
        record = revision["source_glyph_record"]
        raw = base64.b64decode(record["raw_glyf"], validate=True)
        assert hashlib.sha256(raw).hexdigest() == source_key[2] == record["glyph_sha256"]
        assert classify_ttf_glyf_record(raw)["eligibility"] == decoder.GLOBAL_ELIGIBLE_SIMPLE_GLYF_V1
        inspector = decoder.TrueTypeGlyphInspector.__new__(decoder.TrueTypeGlyphInspector)
        inspector.data, inspector.num_glyphs, inspector.glyf_off, inspector.loca = raw, 1, 0, [0, len(raw)]
        contours = inspector.simple_contours(0)
        signature = contour_group_signature(contours)
        assert signature == record["exact_outline_signature"]
        for result in (fingerprints.get(source_key[2], {}).get("bopomofo", ""),
                       outlines.get(signature, {}).get("bopomofo", ""),
                       symbols.decode(contours).get("reading", "")):
            assert result in ("", reading)
        assert all(rec["bopomofo"] != withdrawn for k, rec in current.items() if k[2] == source_key[2])
        legacy_rows.append({"font": source_key[0], "zhuyin_component_id": source_key[1], "bopomofo": withdrawn})
    # A label copied into a legacy proposal/transform never re-acquires roots.
    proposals = decoder.load_map(csv_file(tmp_path / "legacy.csv", list(legacy_rows[0]), legacy_rows))
    for revision in evidence["revisions"]:
        font, gid, sha = revision["root_key"]
        rec = proposals[(font, gid)]
        assert rec["candidate_only"] and not rec.get("evidence_roots")
        assert not decoder.training_reading([{"initial_bopomofo": rec["bopomofo"], **rec}])
        transform = {(font + "-target", gid): {"source_font": font, "source_gid": gid,
                     "target_hashes": (sha,), "bopomofo": revision["withdrawn_reading"]}}
        assert decoder.lookup_verified_transform(font + "-target", gid, transform, proposals, sha) is None
    # Withdrawn or missing exact source assertions remove only their explicit
    # dependants; unaffected independent claims retain their own qualification.
    for missing in (False, True):
        changed = [({**r, "verification": "REVOKED_CORRECTIVE2_EXACT_CLAIM"} if key(r) in root_keys else r)
                   for r in rows if not missing or key(r) not in root_keys]
        after = decoder.load_verified_mapping_corrections(csv_file(tmp_path / f"withdrawn-{missing}.csv", list(rows[0]), changed))
        assert not targets.intersection(after)
        assert set(after) == set(current) - targets
        assert all(rec["bopomofo"] == current[k]["bopomofo"] for k, rec in after.items())
    stale_values = {tuple(k): r["withdrawn_reading"] for r in evidence["revisions"]
                    for k in r["current_target_keys"] if tuple(k) != tuple(r["root_key"])}
    stale = [{**r, "new_reading": stale_values[key(r)]} if key(r) in stale_values else r for r in rows]
    after = decoder.load_verified_mapping_corrections(csv_file(tmp_path / "stale-alias.csv", list(rows[0]), stale))
    assert set(after) == set(current) - set(stale_values)
    assert root_keys.issubset(after)
    # Reconstruct the exact prior direct/alias asset bytes from the preserved
    # claims. The low-level parser can inspect this historical file, but the
    # current official manifest must not admit it as a current runtime asset.
    prior_rows = {key(p["row"]): p["row"] for r in evidence["revisions"] for p in r["prior_rows"]}
    prior = [prior_rows.get(key(r), r) for r in rows]
    buffer = io.StringIO(newline="")
    current_raw = decoder.DEFAULT_MAPPING_CORRECTIONS.read_bytes()
    writer = csv.DictWriter(buffer, fieldnames=list(rows[0]),
                            lineterminator="\r\n" if b"\r\n" in current_raw else "\n")
    writer.writeheader()
    writer.writerows(prior)
    prior_raw = (b"\xef\xbb\xbf" if current_raw.startswith(b"\xef\xbb\xbf") else b"") + buffer.getvalue().encode("utf-8")
    old_asset_sha = evidence["prior_asset_hashes"]["ttf_verified_mapping_corrections.csv"]
    assert hashlib.sha256(prior_raw).hexdigest() == old_asset_sha
    runtime = tmp_path / "runtime"
    runtime.mkdir()
    manifest_path = root / "runtime_asset_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    shutil.copyfile(manifest_path, runtime / manifest_path.name)
    for asset in manifest["assets"]:
        destination = runtime / asset["path"]
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(root / asset["path"], destination)
    assert source.validate_asset_manifest(runtime)["ok"]
    (runtime / decoder.DEFAULT_MAPPING_CORRECTIONS.name).write_bytes(prior_raw)
    report = source.validate_asset_manifest(runtime)
    assert not source.source_chain_ok(report, "actual")
    assert source.source_chain_ok(report, "expected")
    assert any(a["name"] == "ttf_verified_mapping_corrections" and not a["ok"] for a in report["assets"])
    # A prior epoch2 receipt cannot be rebound to corrected static evidence.
    # Refusal preserves the entire workbook, including unrelated/manual data.
    from tests.test_semantics_epoch_contract import actual_components, compatible
    old_components = actual_components(actual_asset_hashes={"ttf_verified_mapping_corrections": old_asset_sha})
    new_components = actual_components(actual_asset_hashes={"ttf_verified_mapping_corrections": hashlib.sha256(decoder.DEFAULT_MAPPING_CORRECTIONS.read_bytes()).hexdigest()})
    assert not compatible(old_components, new_components, chain="actual")
    workbook = tmp_path / "stale-direct-receipt.xlsx"
    wb = Workbook()
    meta = wb.active
    meta.title = "v5.2中繼資料"
    meta.append(["actual_decoder_semantics_epoch", "2"])
    meta.append(["actual_asset_fingerprint_components", json.dumps(old_components)])
    wb.create_sheet("historical_manual_expected").append(["unchanged evidence"])
    wb.save(workbook)
    before = workbook.read_bytes()
    with pytest.raises(ValueError, match="actual contract changed"):
        source.rewrite_actual_workbook_fingerprint(workbook, {"fingerprint": "current", "components": new_components})
    assert workbook.read_bytes() == before
    # Official single-file CLI refuses custom static corrections before PDF
    # acquisition, model training or any Global repository access can occur.
    pdf = tmp_path / "source.pdf"
    pdf.write_bytes(b"path admission fixture, never opened")
    custom = runtime / decoder.DEFAULT_MAPPING_CORRECTIONS.name
    with patch("sys.argv", ["decoder", str(pdf), "--mapping-corrections", str(custom)]), \
            patch.object(decoder, "notify"), \
            patch.object(decoder, "decode", side_effect=AssertionError("custom old asset reached decoder")), \
            patch.object(decoder.GlobalExactGlyphRepository, "resolved", side_effect=AssertionError("custom old asset reached Global repository")):
        assert decoder.main() == 1
    _check_current_correction_acquisition_boundary(tmp_path, evidence, current, prior_raw, runtime)


def _check_current_correction_acquisition_boundary(tmp_path, evidence, current, prior_raw, runtime):
    copy_path = tmp_path / "renamed-current-copy.csv"
    copy_path.write_bytes(decoder.DEFAULT_MAPPING_CORRECTIONS.read_bytes())
    current_sha = hashlib.sha256(copy_path.read_bytes()).hexdigest()
    official = decoder.load_verified_mapping_corrections(decoder.DEFAULT_MAPPING_CORRECTIONS,
                                                         expected_asset_sha256=current_sha)
    copied = decoder.load_verified_mapping_corrections(copy_path, expected_asset_sha256=current_sha)
    assert all(copied[k]["evidence_roots"] == official[k]["evidence_roots"] for k in official)
    old = tmp_path / "renamed-stale-copy.csv"
    old.write_bytes(prior_raw)
    historical = decoder.load_verified_mapping_corrections(old)  # Inspection is allowed.
    for revision in evidence["revisions"]:
        assert historical[tuple(revision["root_key"])]["bopomofo"] == revision["withdrawn_reading"]
    with pytest.raises(ValueError, match="current correction asset binding mismatch"):
        decoder.load_verified_mapping_corrections(old, expected_asset_sha256=current_sha)
    # A valid old manifest next to a selected old copy is not the current
    # deployed module's source authority. No PDF/model/output is touched.
    manifest_path = runtime / "runtime_asset_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    for asset in manifest["assets"]:
        if asset["name"] == "ttf_verified_mapping_corrections":
            asset["sha256"] = hashlib.sha256(prior_raw).hexdigest()
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    assert source.validate_asset_manifest(runtime)["ok"]
    for index, payload in enumerate((None, {}, {"components": {"actual_decoder_semantics_epoch": "2",
                               "actual_asset_hashes": {"ttf_verified_mapping_corrections": current_sha}},
                               "fingerprint": "claimed-current"})):
        output = tmp_path / f"stale-{index}.xlsx"
        with patch.object(decoder.fitz, "open", side_effect=AssertionError("stale source reached PDF")), \
                patch.object(decoder, "cross_validate_font", side_effect=AssertionError("stale source reached OOF")), \
                patch.object(decoder, "train_model", side_effect=AssertionError("stale source reached training")):
            with pytest.raises(ValueError, match="current correction asset binding mismatch"):
                decoder.decode(tmp_path / "source.pdf", output, decoder.DEFAULT_MAP, decoder.DEFAULT_GROUPS,
                               mapping_corrections_path=(old if index != 1 else runtime / decoder.DEFAULT_MAPPING_CORRECTIONS.name),
                               actual_asset_fingerprint=payload, dynamic_evidence_root=tmp_path / f"not-created-{index}",
                               global_snapshot=_absent_snapshot())
        assert not output.exists()
        assert not (tmp_path / f"not-created-{index}").exists()
    # Simulate replacing the selected default path after validation, without
    # ever modifying an installed source file. The loader checks its own read.
    read_bytes = Path.read_bytes
    def replaced_after_validation(path):
        return prior_raw if path.resolve() == decoder.DEFAULT_MAPPING_CORRECTIONS.resolve() else read_bytes(path)
    with patch.object(Path, "read_bytes", replaced_after_validation), \
            patch.object(decoder.fitz, "open", side_effect=AssertionError("changed source reached PDF")):
        with pytest.raises(ValueError, match="current correction asset binding mismatch"):
            decoder.decode(tmp_path / "source.pdf", tmp_path / "replaced.xlsx", decoder.DEFAULT_MAP,
                           decoder.DEFAULT_GROUPS, global_snapshot=_absent_snapshot())
    # Missing optional fingerprint remains supported. A bytematch copy and the
    # deployed file each feed only the corrected real glyph label into OOF.
    revision = evidence["revisions"][0]
    font, gid, sha = revision["root_key"]
    raw = base64.b64decode(revision["source_glyph_record"]["raw_glyf"])
    source_inspector = decoder.TrueTypeGlyphInspector.__new__(decoder.TrueTypeGlyphInspector)
    source_inspector.data, source_inspector.num_glyphs, source_inspector.glyf_off, source_inspector.loca = raw, 1, 0, [0, len(raw)]
    contours = source_inspector.simple_contours(0)
    class Inspector(_DecodeTrueTypeInspector):
        def components(self, glyph_id):
            return [{"gid": 1, "x": 0, "y": 0, "flags": 0}, {"gid": gid, "x": 1024, "y": 0, "flags": 0}]
        def glyph_sha256(self, glyph_id):
            return sha
        def global_exact_identity(self, glyph_id):
            return {"eligibility": decoder.GLOBAL_ELIGIBLE_SIMPLE_GLYF_V1, "glyph_sha256": sha}
        def simple_contours(self, glyph_id):
            return contours
    validate = decoder.cross_validate_font
    for index, path in enumerate((decoder.DEFAULT_MAPPING_CORRECTIONS, copy_path)):
        isolated = tmp_path / f"accepted-{index}"
        isolated.mkdir()
        document = _DecodeDocument()
        document.page.get_fonts = lambda full=True: [(10, "ttf", "TrueType", font, "F1", "Identity-H", 0)]
        document.page.get_texttrace = lambda: [{"font": font, "size": 12, "opacity": 1, "seqno": 1,
                                             "chars": ((ord("字"), 5, (10, 10), (10, 10, 20, 20)),)}]
        training = []
        def observe(current_font, events):
            training.extend(events)
            return validate(current_font, events)
        _, (actual, _) = run_ttf_fixture(isolated, document=document, inspector=Inspector,
                                         corrections_path=path, observe_training=observe)
        assert actual["實際注音"] == revision["current_reading"]
        assert actual["解碼依據"] == "SHA綁定字形真值修正"
        assert len(training) == 1 and training[0].reading == revision["current_reading"]
        assert current[tuple(revision["root_key"])]["donation_eligible"] is False
    # The acquisition guard reads only manifest metadata and the correction
    # body. Real corrupted expected XLSX bytes must not block independent actual.
    root = Path(decoder.DEFAULT_MAPPING_CORRECTIONS).parent
    current_manifest = json.loads((root / "runtime_asset_manifest.json").read_text(encoding="utf-8"))
    manifest_path.write_text(json.dumps(current_manifest), encoding="utf-8")
    (runtime / decoder.DEFAULT_MAPPING_CORRECTIONS.name).write_bytes(copy_path.read_bytes())
    expected_paths = {(runtime / a["path"]).resolve() for a in current_manifest["assets"] if a["chain"] == "expected"}
    expected_xlsx = next(runtime / a["path"] for a in current_manifest["assets"] if a["chain"] == "expected" and a["kind"] == "xlsx")
    expected_xlsx.write_bytes(b"corrupted XLSX body: not a ZIP")
    open_path = Path.open
    def forbid_expected_body(path, *args, **kwargs):
        assert path.resolve() not in expected_paths, "actual correction guard read expected asset body"
        return open_path(path, *args, **kwargs)
    isolated = tmp_path / "actual-independent-of-expected"
    isolated.mkdir()
    with patch.object(decoder, "__file__", str(runtime / "export_zhuyin_readings.py")), \
            patch.object(Path, "open", forbid_expected_body):
        _, (actual, _) = run_ttf_fixture(isolated, document=document, inspector=Inspector,
                                         corrections_path=copy_path, observe_training=observe)
    assert actual["實際注音"] == revision["current_reading"]
    # Structural metadata and the one selected asset remain strict. These
    # are acquisition failures, never fallbacks to an old/self-issued digest.
    variants = []
    for change in ({"schema_version": "999.0"},):
        variants.append({**current_manifest, **change})
    for field, value in (("chain", "expected"), ("path", "../outside.csv"),
                         ("kind", "json"), ("sha256", "0" * 64)):
        variant = json.loads(json.dumps(current_manifest))
        next(a for a in variant["assets"] if a["name"] == "ttf_verified_mapping_corrections")[field] = value
        variants.append(variant)
    duplicate = json.loads(json.dumps(current_manifest))
    duplicate["assets"].append(duplicate["assets"][0])
    variants.extend((duplicate, [], {**current_manifest, "assets": [None]}))
    for variant in variants:
        manifest_path.write_text(json.dumps(variant), encoding="utf-8")
        with pytest.raises(ValueError, match="current correction asset manifest"):
            decoder._current_mapping_correction_asset_sha(runtime)
    manifest_path.write_text('{"schema_version":"1.0","schema_version":"1.0"}', encoding="utf-8")
    with pytest.raises(ValueError):
        decoder._current_mapping_correction_asset_sha(runtime)
    manifest_path.write_text(json.dumps(current_manifest), encoding="utf-8")
    (runtime / decoder.DEFAULT_MAPPING_CORRECTIONS.name).write_bytes(prior_raw)
    with pytest.raises(ValueError, match="current correction asset manifest binding"):
        decoder._current_mapping_correction_asset_sha(runtime)


def check_same_name_gid_other_variant_and_local_override_cannot_train():
    admitted = {"initial_bopomofo": "ㄅ", "training_eligible": True, "evidence_roots": ("direct:sha-A",)}
    assert decoder.training_reading([admitted]) == "ㄅ"
    assert not decoder.training_reading([admitted, {"initial_bopomofo": "", "training_eligible": False}])
    assert not decoder.training_reading([{**admitted, "training_eligible": False}])
    assert not decoder.training_reading([{**admitted, "evidence_roots": ()}])


def check_transform_target_binding_cannot_launder_unbound_source():
    transforms = {("Target", 7): {"source_font": "Source", "source_gid": 3,
                                 "target_hashes": (SHA_A,), "bopomofo": "ㄅ"}}
    source_map = {("Source", 3): {"bopomofo": "ㄅ", "candidate_only": True}}
    assert decoder.lookup_verified_transform("Target", 7, transforms, source_map, SHA_A) is None


def check_table_agreement_and_long_notes_are_not_truth_and_duplicate_conflict_is_closed(tmp_path):
    path = csv_file(tmp_path / "legacy.csv", ["font", "zhuyin_component_id", "bopomofo", "notes"], [
        {"font": "F", "zhuyin_component_id": 7, "bopomofo": "ㄅ", "notes": "verified " * 100},
        {"font": "F", "zhuyin_component_id": 7, "bopomofo": "ㄆ"},
        {"font": "F", "zhuyin_component_id": 7, "bopomofo": "ㄅ"},
    ])
    assert not decoder.load_map(path)


def check_unbound_automatic_outline_and_templates_are_excluded_before_model_loading(tmp_path):
    sig = "a" * 40
    path = csv_file(tmp_path / "outline.csv", ["outline_signature", "bopomofo", "verification"], [
        {"outline_signature": sig, "bopomofo": "ㄅ", "verification": "v3.7 decoded labels; OOF PASS"},
    ])
    assert not decoder.load_verified_outline_signatures(path)
    # A once-admissible direct outline assertion can be specifically withdrawn;
    # its preserved historical verification is audit data, not current truth.
    direct_role = "v4.9 一上解碼漏失修正：原 PDF 可見注音字形直接覆核後凍結 exact normalized outline signature"
    fields = ["outline_signature", "bopomofo", "verification", "legacy_verification", "revocation_evidence"]
    claim = {"outline_signature": sig, "bopomofo": "ㄏㄢˋ", "verification": direct_role,
             "legacy_verification": direct_role}
    csv_file(path, fields, [claim])
    assert decoder.load_verified_outline_signatures(path)[sig]["bopomofo"] == "ㄏㄢˋ"
    csv_file(path, fields, [{**claim, "verification": "REVOKED_TARGET_COUNTEREXAMPLE_V1",
                             "revocation_evidence": "specific original-page glyph counterexample"}])
    assert not decoder.load_verified_outline_signatures(path)
    path = csv_file(tmp_path / "templates.csv", ["template_type", "outline_signature", "label", "support_keys", "support_fonts", "validation"], [
        {"template_type": "symbol", "outline_signature": sig, "label": "ㄅ", "support_keys": 99,
         "support_fonts": 99, "validation": "static full-data audit + key LOO + cross-book holdout; support>=2"},
    ])
    assert not ExactSymbolRecombinationModel.load_csv(path).templates
    # A directly observed isolated symbol is a distinct surviving truth root.
    with path.open("a", encoding="utf-8") as fh:
        fh.write(f"symbol,{sig},ㄅ,1,1,direct_isolated_visual_v39\n")
    templates = ExactSymbolRecombinationModel.load_csv(path).templates
    assert len(templates) == 1 and templates[0].evidence_ref


def run_ttf_fixture(tmp_path, *, snapshot=None, static=None, symbol_model=None, document=None,
                    inspector=_DecodeTrueTypeInspector, observe_training=None, allow_training=False, fingerprint=None,
                    corrections_path=None):
    pdf = tmp_path / "sample.pdf"
    pdf.write_bytes(b"isolated acquisition fixture")
    out = tmp_path / "actual.xlsx"
    mapping = csv_file(tmp_path / "map.csv", ["font", "zhuyin_component_id", "bopomofo"], [
        {"font": "FixtureFont", "zhuyin_component_id": 7, "bopomofo": "ㄇ"},
    ])
    patches = [
        patch.object(decoder.fitz, "open", return_value=document or _DecodeDocument()),
        patch.object(decoder, "TrueTypeGlyphInspector", inspector),
        patch.object(decoder, "load_verified_glyf_fingerprints", return_value=static or {}),
        patch.object(decoder, "cross_validate_font", side_effect=observe_training or AssertionError("unbound label entered training")),
        patch.object(decoder, "train_model", wraps=decoder.train_model) if allow_training else
        patch.object(decoder, "train_model", side_effect=AssertionError("unbound label entered model")),
    ]
    from contextlib import ExitStack
    with ExitStack() as stack:
        for p in patches:
            stack.enter_context(p)
        if symbol_model:
            stack.enter_context(patch.object(decoder.ExactSymbolRecombinationModel, "load_csv", return_value=symbol_model))
        decoder.decode(pdf, out, mapping, tmp_path / "groups.csv", dynamic_evidence_root=tmp_path / "dynamic",
                       global_snapshot=snapshot or _absent_snapshot(), actual_asset_fingerprint=fingerprint,
                       mapping_corrections_path=corrections_path)
    return out, _read_actual_and_audit(out)


def check_same_glyph_copies_do_not_inflate_shape_training(tmp_path):
    # Three acquired keys include two copies of exactly the same raw glyph.
    # Exercise the actual training-input construction, retaining real OOF rules.
    document = _DecodeDocument()
    document.page.get_texttrace = lambda: [{
        "font": "FixtureFont", "size": 12, "opacity": 1, "seqno": 1,
        "chars": tuple((ord("字"), gid, (10, 10), (10 + offset, 10, 20 + offset, 20))
                       for offset, gid in enumerate((5, 6, 8))),
    }]
    class Inspector(_DecodeTrueTypeInspector):
        def components(self, glyph_id):
            return [{"gid": 1, "x": 0, "y": 0, "flags": 0},
                    {"gid": {5: 7, 6: 8, 8: 9}[glyph_id], "x": 1024, "y": 0, "flags": 0}]
        def glyph_sha256(self, glyph_id):
            return SHA_A if glyph_id in (7, 8) else SHA_B
        def global_exact_identity(self, glyph_id):
            return {"eligibility": decoder.GLOBAL_ELIGIBLE_SIMPLE_GLYF_V1,
                    "glyph_sha256": self.glyph_sha256(glyph_id)}
        def simple_contours(self, glyph_id):
            extent = 100 if glyph_id in (7, 8) else 200
            return [[(0, 0, True), (extent, 0, True), (extent, extent, True), (0, extent, True)]]
    observed = []
    validate = decoder.cross_validate_font
    def observe(font, events):
        observed.append((font, events))
        return validate(font, events)
    static = {sha: {"bopomofo": reading, "evidence_roots": (f"direct:{sha}",)}
              for sha, reading in ((SHA_A, "ㄅ"), (SHA_B, "ㄆ"))}
    run_ttf_fixture(tmp_path, document=document, inspector=Inspector, static=static, observe_training=observe)
    assert len(observed) == 1
    assert {event.key for event in observed[0][1]} == {("FixtureFont", 7), ("FixtureFont", 9)}


def check_mixed_variants_fill_only_independently_decoded_pending_rows(tmp_path):
    def square(extent=100):
        return [[(0, 0, True), (extent, 0, True), (extent, extent, True), (0, extent, True)]]
    def acquisition(training_font=None, *, unmatched=False, extra_pending=False):
        specs = [("FixtureFont", 10, [(5, 7, SHA_A, square())]),
                 ("FixtureFont", 11, [(5, 7, SHA_B, square(200 if unmatched else 100))])]
        if training_font:
            specs.append((training_font, 12, [(6, 9, "c" * 64, square()), (8, 10, "d" * 64, square())]))
        if training_font == "SourceFont":
            specs.append((training_font, 13, [(5, 7, SHA_B, square())]))
        if extra_pending:
            specs.append(("FixtureFont", 14, [(5, 7, "e" * 64, square(300))]))
        class Document(_DecodeDocument):
            page_count = len(specs)
            def __iter__(self):
                for font, xref, entries in specs:
                    page = type(self.page)()
                    page.get_fonts = lambda full=True, f=font, x=xref: [(x, "ttf", "TrueType", f, "F1", "Identity-H", 0)]
                    page.get_texttrace = lambda f=font, es=entries: [{
                        "font": f, "size": 12, "opacity": 1, "seqno": 1,
                        "chars": tuple((ord("字"), gid, (10, 10), (10 + i * 20, 10, 20 + i * 20, 20))
                                       for i, (gid, _, _, _) in enumerate(es)),
                    }]
                    yield page
            def extract_font(self, xref):
                return ("fixture", "ttf", "TrueType", str(xref).encode())
        by_xref = {xref: entries for _, xref, entries in specs}
        class Inspector(_DecodeTrueTypeInspector):
            def __init__(self, data):
                self.entries = by_xref[int(data)]
            def components(self, glyph_id):
                zh_id = next(zh for gid, zh, _, _ in self.entries if gid == glyph_id)
                return [{"gid": 1, "x": 0, "y": 0, "flags": 0}, {"gid": zh_id, "x": 1024, "y": 0, "flags": 0}]
            def glyph_sha256(self, glyph_id):
                return next(sha for _, zh, sha, _ in self.entries if zh == glyph_id)
            def simple_contours(self, glyph_id):
                return next(contours for _, zh, _, contours in self.entries if zh == glyph_id)
            def global_exact_identity(self, glyph_id):
                return {"eligibility": decoder.GLOBAL_ELIGIBLE_SIMPLE_GLYF_V1, "glyph_sha256": self.glyph_sha256(glyph_id)}
        return Document(), Inspector
    # Inject only a previously validated ready-state boundary to isolate the
    # orchestration under test. Symbol/shape matching and model training remain
    # real. This small fixture is not evidence that it passes production OOF.
    def ready(font, events):
        return ValidationResult(font, unique_keys=len(events), ready=True,
                                oof_predictions={event.key: event.reading for event in events})
    for method, expected_basis in (("symbol", "TTF exact注音符號重組"),
                                   ("shape", "TTF字形分解自動泛化（OOF驗證模型）"),
                                   ("bridge", "TTF跨字型glyf指紋橋接（安全字形來源）")):
        for positive in (True, False):
            isolated = tmp_path / f"{method}-{positive}"
            isolated.mkdir()
            training_font = "FixtureFont" if method == "shape" else ("SourceFont" if method == "bridge" else None)
            document, inspector = acquisition(training_font, unmatched=not positive and method != "bridge",
                                               extra_pending=not positive and method == "bridge")
            first_reading = "ㄅ" if method == "bridge" else "ㄆ"
            static = {SHA_A: {"bopomofo": first_reading, "evidence_roots": ("direct:A",)}}
            if training_font:
                static.update({sha: {"bopomofo": "ㄅ", "evidence_roots": (f"direct:{sha}",)} for sha in ("c" * 64, "d" * 64)})
            symbol_path = csv_file(isolated / "symbols.csv", ["template_type", "outline_signature", "label", "validation"],
                                   [{"template_type": "symbol", "outline_signature": contour_group_signature(square()), "label": "ㄅ",
                                     "validation": "direct_isolated_visual_v39"}] if method == "symbol" else [])
            output, _ = run_ttf_fixture(isolated, document=document, inspector=inspector, static=static,
                                         symbol_model=ExactSymbolRecombinationModel.load_csv(symbol_path),
                                         observe_training=ready if training_font else None, allow_training=bool(training_font))
            with closing(load_workbook(output, read_only=True)) as wb:
                values = list(wb["實際注音"].values)
                rows = [dict(zip(values[0], row)) for row in values[1:]]
                assert rows[0]["實際注音"] == first_reading
                assert rows[0]["解碼依據"] == decoder.PROJECT_STATIC_TTF_SOURCE_LABEL
                assert rows[1]["實際注音"] == ("ㄅ" if positive else None)
                if positive:
                    assert rows[1]["解碼依據"] == expected_basis
                else:
                    assert all(not row["實際注音"] for row in rows if row["font_xref"] == 14)


def check_derived_oof_disagreement_does_not_revoke_qualified_direct_truth(tmp_path):
    contours = [[(0, 0, True), (100, 0, True), (100, 100, True), (0, 100, True)]]
    def disagree(font, events):
        return ValidationResult(font, unique_keys=len(events), ready=True,
                                oof_predictions={event.key: "ㄆ" for event in events})
    with patch.object(_DecodeTrueTypeInspector, "simple_contours", return_value=contours):
        output, (decoded, _) = run_ttf_fixture(tmp_path, static={SHA_A: {
            "bopomofo": "ㄅ", "evidence_roots": ("independently-admitted-direct:A",),
        }}, observe_training=disagree, allow_training=True)
    assert decoded["實際注音"] == "ㄅ"
    assert decoded["解碼依據"] == decoder.PROJECT_STATIC_TTF_SOURCE_LABEL
    with closing(load_workbook(output, read_only=True)) as wb:
        assert "真值保留" in decoded["TTF字形自驗狀態"]
        assert wb["TTF跨字型指紋橋接"].max_row == 1
        roots = list(wb["TTF證據准入"].values)
        assert "independently-admitted-direct:A" in roots[1][7]


def check_unbound_candidate_never_enters_actual_training_or_fingerprint_bridge(tmp_path):
    out, (actual, _) = run_ttf_fixture(tmp_path)
    assert not actual["實際注音"]
    with closing(load_workbook(out, read_only=True)) as wb:
        assert wb["舊對照候選稽核"].max_row == 2
        assert wb["TTF跨字型指紋橋接"].max_row == 1


def check_changed_project_source_is_rejected_before_training_and_output(tmp_path):
    def changed_after_capture(path, *, snapshot=None):
        path.write_bytes(path.read_bytes() + b"\r\n")
        return {}
    with patch.object(decoder, "load_user_verified_glyf", side_effect=changed_after_capture):
        with pytest.raises(ValueError, match="snapshot source changed"):
            run_ttf_fixture(tmp_path)
    assert not (tmp_path / "actual.xlsx").exists()


def check_project_readers_consume_one_validated_snapshot(tmp_path):
    actual.ensure_user_evidence_files(tmp_path)
    csv_file(tmp_path / actual.USER_GLYF_FILE, actual.USER_GLYF_HEADERS, [{
        "glyph_sha256": SHA_A, "bopomofo": "ㄅ", "verification_level": "VERIFIED_EXACT_GLYPH",
        "source_count": 2, "notes": "existing local truth",
    }])
    snapshot = actual.load_dynamic_actual_snapshot(tmp_path, read_only=True)
    with patch.object(actual, "_read_csv", side_effect=AssertionError("live rows read after capture")):
        assert actual.load_user_verified_glyf(tmp_path / actual.USER_GLYF_FILE, snapshot=snapshot)[SHA_A]["bopomofo"] == "ㄅ"
        assert actual.load_user_verified_cff(tmp_path / actual.USER_CFF_FILE, snapshot=snapshot) == {}
        assert actual.load_glyph_truth_quarantine(tmp_path, snapshot=snapshot) == {"ttf": set(), "cff": set()}
        assert decoder.load_actual_occurrence_overrides(tmp_path / actual.OCCURRENCE_OVERRIDE_FILE, tmp_path / "source.pdf", snapshot=snapshot) == []


def check_static_only_exact_can_replace_candidate_before_training(tmp_path):
    _, (actual, _) = run_ttf_fixture(tmp_path, static={SHA_A: {
        "bopomofo": "ㄅ", "source_font": "direct isolated", "source_gid": 7,
        "evidence_roots": ("direct-isolated-glyph:sha-A",),
    }})
    assert actual["實際注音"] == "ㄅ"
    assert actual["解碼依據"] == "Project/static exact TTF glyf SHA-256"


def check_independent_symbol_fallback_can_succeed_without_clearing_global_quarantine(tmp_path):
    snapshot = _snapshot(quarantined=(_record(_identity(SHA_A), conflict=("ㄅ", "ㄆ")),))
    # Acquiring geometry is isolated; execution still crosses the real initial,
    # training/donor and all-xref fallback paths. No Global state is modified.
    contours = [[(0, 0, True), (100, 0, True), (100, 100, True), (0, 100, True)]]
    path = csv_file(tmp_path / "symbols.csv", ["template_type", "outline_signature", "label", "support_keys", "support_fonts", "validation"], [
        {"template_type": "symbol", "outline_signature": contour_group_signature(contours), "label": "ㄈ",
         "support_keys": 1, "support_fonts": 1, "validation": "direct_isolated_visual_v39"},
    ])
    model = ExactSymbolRecombinationModel.load_csv(path)
    with patch.object(_DecodeTrueTypeInspector, "simple_contours", return_value=contours):
        output, (actual, audit) = run_ttf_fixture(tmp_path, snapshot=snapshot, symbol_model=model)
    assert actual["實際注音"] == "ㄈ"
    assert actual["解碼依據"] == "TTF exact注音符號重組"
    assert audit[0]["處置"] == decoder.GLOBAL_EXACT_RUNTIME_CONFLICT_LABEL
    assert snapshot.quarantined_identities
    # The produced audit must describe the actual direct-role admission gate,
    # retaining support counts as metadata rather than a truth certificate.
    with closing(load_workbook(output, read_only=True)) as wb:
        symbol_rows = list(wb["TTF符號重組稽核"].values)
        record = dict(zip(symbol_rows[0], symbol_rows[1]))
        assert record["最小來源鍵支持"] == 1
        assert "來源支持數僅供稽核" in record["安全說明"]
        assert "direct_isolated_visual_" in record["安全說明"]
        key_rows = list(wb["注音鍵彙整"].values)
        assert "direct visual truth role only" in key_rows[1][11]
        summary = dict(wb["摘要"].values)
        assert "來源支持數僅供稽核" in summary["v4.2符號重組原則"]
        assert "support>=2" not in summary["v4.2符號重組原則"]


def check_old_actual_cannot_be_resealed_and_fresh_epoch_two_can(tmp_path):
    path = tmp_path / "old.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "v5.2中繼資料"
    ws.append(["actual_asset_fingerprint", "old"])
    ws.append(["actual_asset_fingerprint_components", json.dumps({"actual_decoder_semantics_epoch": "1"})])
    wb.save(path)
    payload = {"fingerprint": "new", "components": {"actual_decoder_semantics_epoch": "2"}}
    before = path.read_bytes()
    with pytest.raises(ValueError, match="fresh decoding|missing fields"):
        source.rewrite_actual_workbook_fingerprint(path, payload)
    assert path.read_bytes() == before
    ws.append(["actual_decoder_semantics_epoch", "2"])
    ws["B2"] = json.dumps({"actual_decoder_semantics_epoch": "2"})
    wb.save(path)
    with pytest.raises(ValueError, match="missing fields"):
        source.rewrite_actual_workbook_fingerprint(path, payload)
    # Even a current-epoch workbook cannot wash revoked static claims by
    # replacing its sealed asset SHA. Only dependency scoping may differ.
    from tests.test_semantics_epoch_contract import actual_components
    current = actual_components()
    ws["B2"] = json.dumps(current)
    wb.save(path)
    before = path.read_bytes()
    changed = {**current, "actual_asset_hashes": {"actual-map": "e" * 64}}
    with pytest.raises(ValueError, match="actual contract changed"):
        source.rewrite_actual_workbook_fingerprint(path, {"fingerprint": "changed", "components": changed})
    assert path.read_bytes() == before
    scoped = {**current, "dynamic_actual_evidence_hashes": {"per-pdf": "f" * 64}}
    source.rewrite_actual_workbook_fingerprint(path, {"fingerprint": "scoped", "components": scoped})
    with closing(load_workbook(path, read_only=True)) as saved:
        assert saved["v5.2中繼資料"]["B1"].value == "scoped"
    # A real optional-fingerprint decode emits epoch2/{} and actual rows. It
    # must not acquire a new PDF/static-evidence contract after the fact.
    unsealed = tmp_path / "unsealed"
    unsealed.mkdir()
    static = {SHA_A: {"bopomofo": "ㄅ", "evidence_roots": ("direct:original",)}}
    unsealed_path, (decoded, _) = run_ttf_fixture(unsealed, static=static)
    assert decoded["實際注音"] == "ㄅ"
    before = unsealed_path.read_bytes()
    with pytest.raises(ValueError, match="missing fields"):
        source.rewrite_actual_workbook_fingerprint(unsealed_path, {"fingerprint": "fabricated", "components": changed})
    assert unsealed_path.read_bytes() == before
    # The supported producer path records a complete provisional contract at
    # decode time; dependency scoping then leaves original actual rows intact.
    fresh = tmp_path / "fresh"
    fresh.mkdir()
    from tests.test_semantics_epoch_contract import fingerprint as digest
    provisional = {"components": current, "fingerprint": digest(current, chain="actual")}
    fresh_path, (decoded, _) = run_ttf_fixture(fresh, static=static, fingerprint=provisional)
    source.rewrite_actual_workbook_fingerprint(fresh_path, {"components": scoped, "fingerprint": digest(scoped, chain="actual")})
    assert _read_actual_and_audit(fresh_path)[0]["實際注音"] == decoded["實際注音"] == "ㄅ"
    assert compat.EXPECTED_RESOLVER_SEMANTICS_EPOCH == "1"


def run_admission_checks(tmp_path):
    cases = [
        ("DFPBiaoKai-W5-ZhuIn-BFW", 14598, "ㄨㄟˊ"),
        ("DFPBiaoKai-W5-ZhuIn-BFW", 14674, "ㄉㄢˇ"),
        ("DFPBiaoKai-W5-ZhuIn-BFW", 15284, "ㄕㄡˊ"),
        ("YenRoundZuinnEG-Md", 14670, "ㄌㄨˋ"),
    ]
    for font, gid, old in cases:
        check_four_historical_proposals_have_no_truth_or_training_qualification(font, gid, old, tmp_path)
    for index in range(10):
        check_all_ten_real_corrections_keep_font_gid_and_exact_hash_scope(index)
    check_same_name_gid_other_variant_and_local_override_cannot_train()
    check_transform_target_binding_cannot_launder_unbound_source()
    checks = [
        check_table_agreement_and_long_notes_are_not_truth_and_duplicate_conflict_is_closed,
        check_unbound_automatic_outline_and_templates_are_excluded_before_model_loading,
        check_unbound_candidate_never_enters_actual_training_or_fingerprint_bridge,
        check_changed_project_source_is_rejected_before_training_and_output,
        check_project_readers_consume_one_validated_snapshot,
        check_calibrated_graphics_keep_original_claims_and_scoped_roles,
        check_scoped_aliases_inherit_only_current_direct_truth_and_require_target_binding,
        check_revised_graphic_claims_do_not_revive_withdrawn_truth,
        check_same_glyph_copies_do_not_inflate_shape_training,
        check_mixed_variants_fill_only_independently_decoded_pending_rows,
        check_derived_oof_disagreement_does_not_revoke_qualified_direct_truth,
        check_static_only_exact_can_replace_candidate_before_training,
        check_independent_symbol_fallback_can_succeed_without_clearing_global_quarantine,
        check_old_actual_cannot_be_resealed_and_fresh_epoch_two_can,
    ]
    for index, check in enumerate(checks):
        isolated = tmp_path / str(index)
        isolated.mkdir()
        check(isolated)
