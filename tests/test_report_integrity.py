from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from check_pronunciation_candidates import (  # noqa: E402
    EXPECTED_RESOLVER_SOURCE_FILES,
    VERSION as EXPECTED_RESOLVER_VERSION,
)
from occurrence_ledger import (  # noqa: E402
    ALL_STATES,
    LEDGER_SCHEMA_VERSION,
    PROCESSING_FINISHED,
    REVIEW_ID_SCHEMA_VERSION,
    SESSION_SCHEMA_VERSION,
    WORKBOOK_SCHEMA_VERSION,
    build_occurrence_ledger,
    prepare_occurrence_rows,
)
from runtime_source_validation import (  # noqa: E402
    compute_expected_asset_fingerprint,
    validate_asset_manifest,
)
from standalone_proofread import VERSION, generate_report, normalize_db, reusable_rules_sha256, seal_manifest  # noqa: E402


def make_source(row_number: int, char: str, actual: str):
    return {
        "pdf_sha256": "b" * 64,
        "pdf": "book.pdf",
        "pdf_name": "book.pdf",
        "實體頁碼": 1,
        "課本頁": 1,
        "字元": char,
        "實際注音": actual,
        "解碼依據": "glyph evidence",
        "穩定注音鍵": f"F#{row_number}",
        "font": "F",
        "font_xref": 1,
        "glyph_id_字形索引": row_number,
        "注音元件ID": row_number,
        "x0": row_number * 10,
        "y0": 20,
        "x1": row_number * 10 + 5,
        "y1": 30,
        "source_row_number": row_number,
        "所在行": "角色與測試",
        "局部詞境": "角色",
    }


class ReportIntegrityTests(unittest.TestCase):
    def test_user_report_is_simple_and_technical_audit_keeps_full_state_views(self):
        first = make_source(1, "角", "ㄐㄩㄝˊ")
        second = make_source(2, "龘", "ㄉㄚˊ")
        sources = [first, second]
        prepare_occurrence_rows("b" * 64, sources)
        ledger = build_occurrence_ledger(sources, {
            first["occurrence_id"]: {
                "state": "PASS",
                "actual": "ㄐㄩㄝˊ",
                "actual_evidence": "glyph evidence",
                "expected_set": ["ㄐㄩㄝˊ"],
                "expected_evidence": "現版規則",
                "context_evidence": "角色/角@0",
                "comparison_result": "MATCH",
                "source_view": "單音一致",
                "source_record": first,
            },
            second["occurrence_id"]: {
                "state": "EXPECTED_UNRESOLVED",
                "actual": "ㄉㄚˊ",
                "actual_evidence": "glyph evidence",
                "source_view": "未建立預期音",
                "source_record": second,
            },
        })
        source_report = validate_asset_manifest(ROOT)
        expected_fingerprint = compute_expected_asset_fingerprint(
            ROOT,
            source_report,
            resolver_version=EXPECTED_RESOLVER_VERSION,
            source_files=EXPECTED_RESOLVER_SOURCE_FILES,
        )
        regression = {
            "ok": False,
            "required": 36,
            "executed": 36,
            "passed": 28,
            "failed": 8,
            "not_executed": 0,
            "duplicate_case_id": 0,
        }
        reusable_rules = {
            "schema_version": "1.0",
            "rules": [{
                "rule_id": "USER-EXPECTED-REPORTTEST",
                "enabled": True,
                "phrase": "角色",
                "target_char": "角",
                "target_index": 0,
                "expected_set": ["ㄐㄩㄝˊ"],
                "evidence": "公司現行規定：角色",
                "context_contains": "",
                "scope": "same_phrase_position",
                "source": "人工核准可重用 expected 規則",
                "note": "報告顯示測試",
                "approved_at": "2026-08-19T00:00:00",
            }],
        }
        manifest = seal_manifest({
            "version": VERSION,
            "session_schema_version": SESSION_SCHEMA_VERSION,
            "workbook_schema_version": WORKBOOK_SCHEMA_VERSION,
            "ledger_schema_version": LEDGER_SCHEMA_VERSION,
            "review_id_schema_version": REVIEW_ID_SCHEMA_VERSION,
            "session_id": "report-test",
            "records": ledger,
            "actual_source_ids": [row["occurrence_id"] for row in sources],
            "expected_asset_fingerprint": expected_fingerprint["fingerprint"],
            "regression_gate": regression,
            "mandatory_regression": regression,
            "pdf_regression": {"failed": 0, "not_executed": 0},
            "pdfs": [],
            "reusable_expected_rules": reusable_rules,
            "reusable_expected_rules_sha256": reusable_rules_sha256(reusable_rules),
        })

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            report_path = generate_report(output, manifest, normalize_db({}))
            status = json.loads((output / "pipeline_status.json").read_text(encoding="utf-8"))
            self.assertEqual(status["status"], PROCESSING_FINISHED)
            self.assertTrue((output / "注音校對_技術稽核.xlsx").exists())

            user_workbook = load_workbook(report_path, read_only=True, data_only=True)
            try:
                self.assertEqual(user_workbook.sheetnames, ["校對摘要", "修正清單", "待確認", "人工與公司規則"])
                pending_sheet = user_workbook["待確認"]
                rows = list(pending_sheet.iter_rows(values_only=True))
                self.assertEqual(len(rows), 2)  # header + one unresolved item
                self.assertIn("尚未建立可獨立重現", str(rows[1]))
                self.assertNotIn("occurrence_id", [str(v or "") for v in rows[0]])
                rules_sheet = user_workbook["人工與公司規則"]
                rule_headers = [str(v or "") for v in next(rules_sheet.iter_rows(values_only=True))]
                self.assertEqual(rule_headers, ["來源類型", "PDF", "課本頁", "詞語／局部詞境", "字", "應標注音", "依據／來源", "適用範圍", "備註"])
                self.assertNotIn("review_id", rule_headers)
                rule_rows = list(rules_sheet.iter_rows(min_row=2, values_only=True))
                self.assertTrue(any(row[0] == "可重用規則" and row[3] == "角色" and row[5] == "ㄐㄩㄝˊ" for row in rule_rows))
            finally:
                user_workbook.close()

            technical = load_workbook(output / "注音校對_技術稽核.xlsx", read_only=True, data_only=True)
            try:
                state_memberships = {}
                for state in ALL_STATES:
                    sheet = technical[state[:31]]
                    rows = sheet.iter_rows(values_only=True)
                    headers = [str(value or "") for value in next(rows)]
                    occurrence_index = headers.index("occurrence_id")
                    for row in rows:
                        occurrence_id = str(row[occurrence_index] or "")
                        if occurrence_id:
                            state_memberships.setdefault(occurrence_id, []).append(state)
                self.assertEqual(set(state_memberships), {row["occurrence_id"] for row in sources})
                self.assertTrue(all(len(states) == 1 for states in state_memberships.values()))
                self.assertIn("集合對帳", technical.sheetnames)
                self.assertIn("Occurrence Ledger", technical.sheetnames)
            finally:
                technical.close()


if __name__ == "__main__":
    unittest.main()
