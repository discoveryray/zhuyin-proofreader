from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
import sys
sys.path.insert(0, str(ROOT))

from occurrence_ledger import (  # noqa: E402
    LEDGER_SCHEMA_VERSION,
    REVIEW_ID_SCHEMA_VERSION,
    SESSION_SCHEMA_VERSION,
    WORKBOOK_SCHEMA_VERSION,
    build_occurrence_ledger,
    prepare_occurrence_rows,
)
from standalone_proofread import (  # noqa: E402
    VERSION,
    export_pending_for_gpt,
    import_gpt_decisions,
    json_load,
    json_save,
    load_or_initialize_db,
    normalize_db,
    seal_manifest,
    workbook_rows,
)


def make_manifest():
    source = {
        "pdf_sha256": "a" * 64, "pdf": "book.pdf", "pdf_name": "book.pdf",
        "實體頁碼": 1, "課本頁": 1, "字元": "角", "實際注音": "ㄐㄩㄝˊ",
        "解碼依據": "glyph evidence", "穩定注音鍵": "F#1", "font": "F", "font_xref": 1,
        "glyph_id_字形索引": 2, "注音元件ID": 1, "x0": 1, "y0": 2, "x1": 3, "y1": 4,
        "source_row_number": 1,
    }
    prepare_occurrence_rows("a" * 64, [source])
    ledger = build_occurrence_ledger([source], {
        source["occurrence_id"]: {
            "state": "EXPECTED_UNRESOLVED", "actual": "ㄐㄩㄝˊ", "actual_evidence": "glyph evidence",
            "context_evidence": "角色/角@1", "source_view": "未建立預期音", "source_record": source,
        }
    })
    return seal_manifest({
        "version": VERSION,
        "session_schema_version": SESSION_SCHEMA_VERSION,
        "workbook_schema_version": WORKBOOK_SCHEMA_VERSION,
        "ledger_schema_version": LEDGER_SCHEMA_VERSION,
        "review_id_schema_version": REVIEW_ID_SCHEMA_VERSION,
        "session_id": "session-test",
        "records": ledger,
        "actual_source_ids": [source["occurrence_id"]],
        "regression_gate": {"ok": False, "required": 1, "executed": 1, "passed": 0, "failed": 1, "not_executed": 0, "duplicate_case_id": 0},
    })


class ImportTransactionTests(unittest.TestCase):
    def test_successful_expected_evidence_import_is_atomic(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            json_save(output / "校對工作階段.json", make_manifest())
            json_save(output / "人工判定資料庫.json", normalize_db({}))
            xlsx = export_pending_for_gpt(output)
            wb = load_workbook(xlsx)
            ws = wb["待判定候選"]
            headers = {str(cell.value or ""): cell.column for cell in ws[1]}
            ws.cell(2, headers["action"], "補建expected證據")
            ws.cell(2, headers["proposed_expected_set"], "ㄐㄩㄝˊ")
            ws.cell(2, headers["proposed_expected_evidence"], "現版手冊")
            ws.cell(2, headers["proposed_context_evidence"], "角色/角@1")
            wb.save(xlsx)
            fake_report = output / "fake.xlsx"
            with patch("standalone_proofread.generate_report", return_value=fake_report):
                imported, skipped, report = import_gpt_decisions(output, xlsx)
            self.assertEqual((imported, skipped, report), (1, 0, fake_report))
            db = json_load(output / "人工判定資料庫.json", {})
            self.assertEqual(len(db.get("events", {})), 1)

    def test_duplicate_row_rejects_entire_batch(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            json_save(output / "校對工作階段.json", make_manifest())
            json_save(output / "人工判定資料庫.json", normalize_db({}))
            xlsx = export_pending_for_gpt(output)
            wb = load_workbook(xlsx)
            ws = wb["待判定候選"]
            values = [cell.value for cell in ws[2]]
            ws.append(values)
            headers = {str(cell.value or ""): cell.column for cell in ws[1]}
            ws.cell(2, headers["action"], "補建expected證據")
            ws.cell(3, headers["action"], "補建expected證據")
            wb.save(xlsx)
            with self.assertRaises(Exception):
                import_gpt_decisions(output, xlsx)
            db = json_load(output / "人工判定資料庫.json", {})
            self.assertEqual(db.get("events"), {})

    def test_duplicate_excel_header_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "duplicate-header.xlsx"
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "待判定候選"
            ws.append(["review_id", "review_id"])
            ws.append(["a", "b"])
            wb.save(path)
            with self.assertRaises(ValueError):
                workbook_rows(path, "待判定候選", ["review_id"])

    def test_legacy_schema_rejected(self):
        with self.assertRaises(ValueError):
            normalize_db({"version": 1, "decisions": {"old": {"decision": "課本正確"}}})

    def test_tampered_session_manifest_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            manifest = make_manifest()
            manifest["regression_gate"]["failed"] = 0
            manifest["regression_gate"]["ok"] = True
            json_save(output / "校對工作階段.json", manifest)
            with self.assertRaises(ValueError):
                export_pending_for_gpt(output)

    def test_current_database_corruption_is_not_silently_reset(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            path = output / "人工判定資料庫.json"
            original = '{"version":"5.2.0","session_schema_version":"2.5.0","events":'
            path.write_text(original, encoding="utf-8")
            with self.assertRaises(ValueError):
                load_or_initialize_db(output)
            self.assertEqual(path.read_text(encoding="utf-8"), original)

    def test_current_database_duplicate_json_key_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            path = output / "人工判定資料庫.json"
            path.write_text(
                '{"version":"5.2.0","session_schema_version":"2.5.0",'
                '"review_id_schema_version":"2.5.0","events":{},"events":{}}',
                encoding="utf-8",
            )
            with self.assertRaises(ValueError):
                load_or_initialize_db(output)

    def test_legacy_database_is_quarantined_without_migration(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            json_save(output / "人工判定資料庫.json", {"version": "5.1.5", "decisions": {"old": {"decision": "課本正確"}}})
            db = load_or_initialize_db(output)
            self.assertEqual(db.get("events"), {})
            orphan = json_load(output / "legacy_orphan_decisions.json", {})
            self.assertEqual(orphan["legacy_items"][0]["payload"]["version"], "5.1.5")


if __name__ == "__main__":
    unittest.main()
