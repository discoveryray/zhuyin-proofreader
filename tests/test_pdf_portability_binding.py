"""Consumed-source, locked-write and marker-owner regressions (synthetic projects)."""
import json
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import load_workbook
import actual_review as ar
import global_glyph_promotion as gp
import pdf_portability as p
import standalone_proofread as sp
from review_save_service import ReviewSaveService
from tests.test_pdf_portability import pdf, project, _filled_expected_excel, add_override, reseal_actual_dynamic
from tests.test_pdf_portability_integration import _pdfs, _synthetic_unresolved_decode


def read(path):
    return path.read_bytes() if path.exists() else None


@pytest.fixture
def actual_projects(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "localappdata"))
    first, second = _pdfs(tmp_path)
    with ExitStack() as stack:
        stack.enter_context(patch.object(sp, "decode", side_effect=_synthetic_unresolved_decode))
        stack.enter_context(patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()))
        stack.enter_context(patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()))
        source, target = tmp_path / "source", tmp_path / "target"
        sp.run_pipeline_pdfs([first], source, defer_excel_reports=True)
        sp.run_pipeline_pdfs([second], target, defer_excel_reports=True)
        yield source, target, second
    assert not list((tmp_path / "localappdata").rglob("*.sqlite*"))


def filled_actual(output, path, *, proof=True, reading="ㄐㄩㄝˊ"):
    sp.export_actual_pending_for_gpt(output)
    w = load_workbook(output / "actual待判定_GPT包" / "actual待判定_給GPT.xlsx")
    if not proof:
        del w[p.EXCEL_PROOF_SHEET]
    sheet = w["actual待判定"]
    headers = [c.value for c in sheet[1]]
    for key, value in {"decision": "VERIFIED", "actual_reading": reading,
                       "confidence": "高", "sample_a_checked": "Y"}.items():
        sheet.cell(2, headers.index(key) + 1, value)
    w.save(path)
    w.close()
    return path


def watched(output):
    return [output / n for n in ("人工判定資料庫.json", "校對工作階段.json",
            "pipeline_status.json", "待人工確認.json", "注音校對_最終報告.xlsx",
            p.CONFLICT_FILE, p.ACTUAL_EXCEL_CONFLICT_FILE, p.INCOMPLETE_FILE)] + [
            sp.project_actual_evidence_root(output) / n for n in
            (*gp.PROJECT_FILES, gp.PROJECT_TRANSACTION_FILE)]


def test_attachment_expected_uses_one_workbook(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "localappdata"))
    a_pdf = tmp_path / "a.pdf"
    pdf(a_pdf, title="a")
    source = tmp_path / "a"
    project(source, a_pdf, session="A")
    x = _filled_expected_excel(source, tmp_path / "old.xlsx", action="補建expected證據")
    w = load_workbook(x)
    del w[p.EXCEL_PROOF_SHEET]
    w.save(x)
    sheet = w["待判定候選"]
    headers = [c.value for c in sheet[1]]
    sheet.cell(2, headers.index("proposed_expected_evidence") + 1, "VALIDATED_B")
    replacement = tmp_path / "replacement.xlsx"
    w.save(replacement)
    w.close()
    with patch.object(p, "_publish_portable_outputs", return_value=source / "report.xlsx"):
        sp.import_gpt_decisions(source, replacement)
    original = load_workbook
    def swap(path, *args, **kwargs):
        result = original(path, *args, **kwargs)
        if Path(path).name == x.name:
            x.write_bytes(replacement.read_bytes())
        return result
    with patch("openpyxl.load_workbook", side_effect=swap):
        with pytest.raises(ValueError):
            p.attach_filled_excel_proof(source, x)
    assert not list(tmp_path.glob("old_跨電腦_*.xlsx"))


def test_attachment_actual_uses_one_workbook(actual_projects, tmp_path):
    source, _, _ = actual_projects
    x = filled_actual(source, tmp_path / "old.xlsx", proof=False, reading="ㄐㄩㄝˇ")
    good = filled_actual(source, tmp_path / "good.xlsx", proof=False)
    assert sp.import_actual_gpt_decisions(source, good)[0] == 1
    original = load_workbook
    def swap(path, *args, **kwargs):
        result = original(path, *args, **kwargs)
        if Path(path).name == x.name:
            x.write_bytes(good.read_bytes())
        return result
    with patch("openpyxl.load_workbook", side_effect=swap):
        with pytest.raises(ValueError):
            p.attach_filled_excel_proof(source, x)
    assert not list(tmp_path.glob("old_跨電腦_*.xlsx"))


@pytest.mark.parametrize("restore", [False, True])
def test_source_actual_workbook_swap_after_seal_is_rejected(tmp_path, monkeypatch, restore):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "localappdata"))
    a_pdf, b_pdf = tmp_path / "a.pdf", tmp_path / "b.pdf"
    pdf(a_pdf, title="a")
    pdf(b_pdf, title="b")
    a, b = tmp_path / "a", tmp_path / "b"
    am, _ = project(a, a_pdf, session="A", event_positions=(0,))
    bm, _ = project(b, b_pdf, session="B")
    add_override(a, am["records"][0], reading="ㄐㄩㄝˊ")
    reseal_actual_dynamic(a, a_pdf)
    add_override(b, bm["records"][0], reading="ㄐㄩㄝˇ")
    reseal_actual_dynamic(b, b_pdf)
    workbook_path = a / "01_實際注音" / "actual.xlsx"
    sealed = workbook_path.read_bytes()
    source_seal = (a / "校對工作階段.json").read_bytes()
    before = {path: read(path) for path in watched(b)}
    original = sp.workbook_metadata
    injected = False
    def swap(path, *args, **kwargs):
        nonlocal injected
        if Path(path).name == workbook_path.name and not injected:
            injected = True
            add_override(a, am["records"][0], reading="ㄐㄩㄝˇ")
            dynamic = ar.dynamic_actual_hashes(sp.project_actual_evidence_root(a), pdf_path=a_pdf,
                                              dependencies=None, read_only=True)
            w = load_workbook(workbook_path)
            for row in w["v5.2中繼資料"].iter_rows(min_row=2, max_col=2):
                if row[0].value == "actual_asset_fingerprint_components":
                    row[1].value = json.dumps({"dynamic_actual_evidence_hashes": dynamic})
            w.save(workbook_path)
            w.close()
        result = original(path, *args, **kwargs)
        if restore:
            workbook_path.write_bytes(sealed)
        return result
    with patch.object(sp, "workbook_metadata", side_effect=swap), \
         patch.object(p, "_publish_portable_outputs", return_value=b / "report.xlsx"):
        with pytest.raises(ValueError):
            p.import_project_decisions(a, b)
    assert injected
    assert (a / "校對工作階段.json").read_bytes() == source_seal
    assert {path: read(path) for path in watched(b)} == before


@pytest.mark.parametrize("auto", [False, True])
def test_same_actual_rechecks_late_marker(actual_projects, tmp_path, auto):
    _, target, _ = actual_projects
    x = filled_actual(target, tmp_path / "same.xlsx")
    marker = target / p.INCOMPLETE_FILE
    foreign = {"status": "PRESENTATION_PENDING", "owner": "foreign"}
    before = {path: read(path) for path in watched(target) if path != marker}
    original = ar._load_sheet_rows
    injected = False
    def inject(*args, **kwargs):
        nonlocal injected
        result = original(*args, **kwargs)
        if not injected:
            injected = True
            sp.json_save(marker, foreign)
        return result
    with patch.object(ar, "_load_sheet_rows", side_effect=inject):
        with pytest.raises(ValueError):
            if auto:
                with patch("sys.argv", ["standalone_proofread", "-o", str(target), "--import-gpt-auto", str(x)]):
                    sp.main()
            else:
                sp.import_actual_gpt_decisions(target, x)
    assert injected and sp.json_load_strict(marker) == foreign
    assert {path: read(path) for path in before} == before


def test_actual_recovery_preserves_foreign_marker(actual_projects, tmp_path):
    source, target, _ = actual_projects
    x = filled_actual(source, tmp_path / "filled.xlsx")
    marker = target / p.INCOMPLETE_FILE
    foreign = {"status": "PRESENTATION_PENDING", "owner": "foreign after refresh"}
    original = sp.recover_committed_actual_project
    def inject(*args, **kwargs):
        result = original(*args, **kwargs)
        sp.json_save(marker, foreign)
        return result
    with patch.object(sp, "recover_committed_actual_project", side_effect=inject):
        with pytest.raises(ValueError):
            sp.import_actual_gpt_decisions(target, x)
    assert sp.json_load_strict(marker) == foreign
    assert gp.committed_project_recovery(sp.project_actual_evidence_root(target)) is not None


@pytest.mark.parametrize("action", ["補建expected證據", "解決expected證據"])
def test_resolved_adjudication_cannot_be_overwritten_by_old_expected(tmp_path, monkeypatch, action):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "localappdata"))
    a_pdf, b_pdf = tmp_path / "a.pdf", tmp_path / "b.pdf"
    pdf(a_pdf, title="a")
    pdf(b_pdf, title="b")
    a, b = tmp_path / "a", tmp_path / "b"
    project(a, a_pdf, session="A", event_positions=(0,))
    bm, bd = project(b, b_pdf, session="B")
    old = _filled_expected_excel(b, tmp_path / "old.xlsx", action=action)
    rid = bm["records"][0]["review_id"]
    bd["events"][rid] = {"action": "保留待人工"}
    sp.json_save(b / "人工判定資料庫.json", bd)
    with patch.object(p, "_publish_portable_outputs", return_value=b / "report.xlsx"):
        p.import_project_decisions(a, b)
    ReviewSaveService(b).save_event(rid, {"action": "解決expected證據", "expected_set": ["ㄐㄩㄝˊ"],
                    "expected_evidence": "LOCAL_ADJUDICATION", "context_evidence": "角色/角@1"})
    before = {path: read(path) for path in watched(b)}
    with pytest.raises(ValueError):
        sp.import_gpt_decisions(b, old)
    assert {path: read(path) for path in before} == before

    # An unrelated pending row remains importable without replacing the verdict.
    local_event = sp.json_load_strict(b / "人工判定資料庫.json")["events"][rid]
    receipt = read(b / p.CONFLICT_FILE)
    unrelated = _filled_expected_excel(b, tmp_path / "unrelated.xlsx", action=action)
    with patch.object(p, "_publish_portable_outputs", return_value=b / "report.xlsx"):
        assert sp.import_gpt_decisions(b, unrelated)[0] == 1
    assert sp.json_load_strict(b / "人工判定資料庫.json")["events"][rid] == local_event
    assert read(b / p.CONFLICT_FILE) == receipt
    assert not (b / p.INCOMPLETE_FILE).exists()


def test_actual_pdf_change_after_mapping_prevents_any_commit(actual_projects, tmp_path):
    source, target, target_pdf = actual_projects
    x = filled_actual(source, tmp_path / "filled.xlsx")
    changed = tmp_path / "changed.pdf"
    pdf(changed, title="changed", text="different full page")
    before = {path: read(path) for path in watched(target)}
    original = p.load_excel_content_proof
    def swap(*args, **kwargs):
        result = original(*args, **kwargs)
        target_pdf.write_bytes(changed.read_bytes())
        return result
    with patch.object(p, "load_excel_content_proof", side_effect=swap):
        with pytest.raises((ValueError, FileNotFoundError)):
            sp.import_actual_gpt_decisions(target, x)
    assert {path: read(path) for path in before} == before

@pytest.mark.parametrize("window", ["creation", "transition", "acknowledgement", "precommit_failure", "rollback", "post_ack_resume"])
def test_actual_marker_ownership_at_each_boundary(actual_projects, tmp_path, window):
    source, target, _ = actual_projects
    x = filled_actual(source, tmp_path / "filled.xlsx")
    marker = target / p.INCOMPLETE_FILE
    root = sp.project_actual_evidence_root(target)
    foreign = {"status": "PRESENTATION_PENDING", "owner": "foreign " + window}
    original_save = sp.json_save
    if window in {"creation", "transition"}:
        before = {path: read(path) for path in watched(target) if path != marker}
        wanted = "ACTUAL_EXCEL_REFRESH_PENDING" if window == "creation" else "ACTUAL_EXCEL_RECOVERED"
        def inject(path, payload, **kwargs):
            if Path(path) == marker and payload.get("status") == wanted:
                original_save(marker, foreign)
            return original_save(path, payload, **kwargs)
        with patch.object(sp, "json_save", side_effect=inject):
            with pytest.raises(ValueError):
                sp.import_actual_gpt_decisions(target, x)
        if window == "creation":
            assert {path: read(path) for path in before} == before
        else:
            assert gp.committed_project_recovery(root) is not None
    elif window == "precommit_failure":
        before = {path: read(path) for path in watched(target) if path != marker}
        def inject(*args, **kwargs):
            original_save(marker, foreign)
            raise OSError("precommit write failure")
        with patch.object(ar, "_write_csv", side_effect=inject):
            with pytest.raises((ValueError, OSError)):
                sp.import_actual_gpt_decisions(target, x)
        assert {path: read(path) for path in before} == before
    elif window == "acknowledgement":
        original_ack = gp.acknowledge_project_refresh
        def inject(*args, **kwargs):
            result = original_ack(*args, **kwargs)
            original_save(marker, foreign)
            return result
        with patch.object(gp, "acknowledge_project_refresh", side_effect=inject):
            with pytest.raises(ValueError):
                sp.import_actual_gpt_decisions(target, x)
        assert gp.committed_project_recovery(root) is not None
    else:
        if window == "rollback":
            with patch.object(ar, "_write_csv", side_effect=SystemExit("power loss")):
                with pytest.raises(SystemExit):
                    sp.import_actual_gpt_decisions(target, x)
        else:
            original_unlink = Path.unlink
            def crash(path, *args, **kwargs):
                if path == marker:
                    raise SystemExit("power loss after ack")
                return original_unlink(path, *args, **kwargs)
            with patch.object(Path, "unlink", new=crash):
                with pytest.raises(SystemExit):
                    sp.import_actual_gpt_decisions(target, x)
            assert gp.committed_project_recovery(root) is None
        original_snapshot = p._actual_excel_file_snapshot
        def inject(*args, **kwargs):
            result = original_snapshot(*args, **kwargs)
            original_save(marker, foreign)
            return result
        with patch.object(p, "_actual_excel_file_snapshot", side_effect=inject):
            with pytest.raises(ValueError):
                p.resume_actual_excel_project(target)
    assert sp.json_load_strict(marker) == foreign


@pytest.mark.parametrize("auto", [False, True])
def test_same_actual_rechecks_late_unresolved_conflict(actual_projects, tmp_path, auto):
    _, target, _ = actual_projects
    x = filled_actual(target, tmp_path / "same.xlsx")
    manifest = sp.json_load_strict(target / "校對工作階段.json")
    row = manifest["records"][0]
    receipt = target / p.CONFLICT_FILE
    conflict = {"version": 1, "conflicts": [{"target_review_id": row["review_id"],
                "source_event": {"action": "保留待人工", "note": "source"},
                "target_event": {"action": "保留待人工", "note": "target"},
                "source_pdf_sha256": "a" * 64, "target_pdf_sha256": row["pdf_sha256"]}]}
    before = {path: read(path) for path in watched(target) if path != receipt}
    original = ar._load_sheet_rows
    def inject(*args, **kwargs):
        result = original(*args, **kwargs)
        sp.json_save(receipt, conflict)
        return result
    with patch.object(ar, "_load_sheet_rows", side_effect=inject):
        with pytest.raises(ValueError, match="衝突|裁決"):
            if auto:
                with patch("sys.argv", ["standalone_proofread", "-o", str(target), "--import-gpt-auto", str(x)]):
                    sp.main()
            else:
                sp.import_actual_gpt_decisions(target, x)
    assert sp.json_load_strict(receipt) == conflict
    assert {path: read(path) for path in before} == before


@pytest.mark.parametrize("lane", ["conflict", "no_changes"])
def test_actual_pdf_change_blocks_receipt_and_no_change_paths(actual_projects, tmp_path, lane):
    source, target, target_pdf = actual_projects
    x = filled_actual(source, tmp_path / "filled.xlsx")
    other = filled_actual(source, tmp_path / "other.xlsx", reading="ㄐㄩㄝˇ")
    assert sp.import_actual_gpt_decisions(target, x)[0] == 1
    if lane == "conflict":
        x = other
    changed = tmp_path / "changed.pdf"
    pdf(changed, title="changed", text="different full page")
    before = {path: read(path) for path in watched(target)}
    original = p.load_excel_content_proof
    def swap(*args, **kwargs):
        result = original(*args, **kwargs)
        target_pdf.write_bytes(changed.read_bytes())
        return result
    with patch.object(p, "load_excel_content_proof", side_effect=swap):
        with pytest.raises((ValueError, FileNotFoundError)):
            sp.import_actual_gpt_decisions(target, x)
    assert {path: read(path) for path in before} == before

@pytest.mark.parametrize("auto", [False, True])
def test_bundle_actual_late_marker_prevents_commit(actual_projects, tmp_path, auto):
    import csv
    import io
    import zipfile
    source, target, _ = actual_projects
    manifest = sp.json_load_strict(target / "校對工作階段.json")
    entry = sp.materialize_ledger(manifest, sp.load_or_initialize_db(target))[0]
    meta = {key: manifest[key] for key in ("version", "session_id", "session_schema_version",
                                         "workbook_schema_version", "review_id_schema_version")}
    meta.update(bundle_schema_version=sp.GPT_DECISION_BUNDLE_SCHEMA_VERSION, expected_workbook="")
    row = {key: entry[key] for key in ("occurrence_id", "review_id", "pdf_name", "physical_page", "char")}
    row.update(exported_actual=entry.get("actual", ""),
               exported_actual_evidence_sha256=sp._sha_text(entry.get("actual_evidence")),
               verified_actual="ㄐㄩㄝˊ", visual_confirmation="Y", note="synthetic visual")
    stream = io.StringIO()
    writer = csv.DictWriter(stream, fieldnames=list(row))
    writer.writeheader()
    writer.writerow(row)
    bundle = tmp_path / "filled.zip"
    with zipfile.ZipFile(bundle, "w") as archive:
        archive.writestr("GPT判定包.json", json.dumps(meta))
        archive.writestr("actual_occurrence_decisions.csv", stream.getvalue())
    marker = target / p.INCOMPLETE_FILE
    foreign = {"status": "PRESENTATION_PENDING", "owner": "foreign bundle"}
    before = {path: read(path) for path in watched(target) if path != marker}
    original = sp.apply_direct_visual_actual_batch
    def inject(*args, **kwargs):
        sp.json_save(marker, foreign)
        return original(*args, **kwargs)
    with patch.object(sp, "apply_direct_visual_actual_batch", side_effect=inject):
        with pytest.raises(ValueError):
            if auto:
                with patch("sys.argv", ["standalone_proofread", "-o", str(target), "--import-gpt-auto", str(bundle)]):
                    sp.main()
            else:
                sp.import_gpt_decision_bundle(target, bundle)
    assert sp.json_load_strict(marker) == foreign
    assert {path: read(path) for path in before} == before

@pytest.mark.parametrize("kind", ["actual", "expected"])
def test_same_session_excel_source_change_is_rejected_before_write(actual_projects, tmp_path, kind):
    from tests.test_pdf_portability_presentation import expected_workbook
    _, target, _ = actual_projects
    if kind == "actual":
        x = filled_actual(target, tmp_path / "same.xlsx")
        sheet_name, field = "actual待判定", "actual_reading"
        original = ar._load_sheet_rows
    else:
        x = expected_workbook(tmp_path, target)
        sheet_name, field = "待判定候選", "exclusion_reason"
        original = sp.workbook_rows
    before = {path: read(path) for path in watched(target)}
    def edit_after_read(path, sheet, *args):
        result = original(path, sheet, *args)
        if sheet == sheet_name:
            workbook = load_workbook(x)
            visible = workbook[sheet]
            columns = [c.value for c in visible[1]]
            visible.cell(2, columns.index(field) + 1, "ㄐㄩㄝˇ" if kind == "actual" else "changed externally")
            workbook.save(x)
            workbook.close()
        return result
    reader = patch.object(ar, "_load_sheet_rows", side_effect=edit_after_read) if kind == "actual" else \
             patch.object(sp, "workbook_rows", side_effect=edit_after_read)
    with reader:
        with pytest.raises(ValueError, match="Excel.*變動"):
            (sp.import_actual_gpt_decisions if kind == "actual" else sp.import_gpt_decisions)(target, x)
    assert {path: read(path) for path in before} == before

def test_same_actual_preserves_existing_nested_pdf_relocation(actual_projects, tmp_path):
    _, target, target_pdf = actual_projects
    x = filled_actual(target, tmp_path / "same.xlsx")
    moved = tmp_path / "moved"
    moved.mkdir()
    target_pdf.rename(moved / target_pdf.name)
    assert sp.import_actual_gpt_decisions(target, x)[0] == 1
    assert gp.committed_project_recovery(sp.project_actual_evidence_root(target)) is None
    assert not (target / p.INCOMPLETE_FILE).exists()
