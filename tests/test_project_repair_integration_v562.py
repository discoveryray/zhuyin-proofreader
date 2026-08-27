from __future__ import annotations

import csv
import hashlib
import json
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import fitz
from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import standalone_proofread as proof  # noqa: E402
from actual_review import (  # noqa: E402
    USER_GLYF_FILE,
    USER_GLYF_HEADERS,
    actual_workbook_dynamic_dependencies,
    dynamic_actual_hashes,
)
from check_pronunciation_candidates import (  # noqa: E402
    DEFAULT_CHAR_OVERRIDES,
    DEFAULT_CONTEXT_OVERRIDES,
    DEFAULT_DICT,
    DEFAULT_REGRESSIONS,
    DEFAULT_RULES,
    EXPECTED_RESOLVER_SOURCE_FILES,
    analyze,
)
from occurrence_ledger import (  # noqa: E402
    CONFIRMATION_GATES,
    LEDGER_SCHEMA_VERSION,
    PROOFREAD_COMPLETE,
    REVIEW_ID_SCHEMA_VERSION,
    SESSION_SCHEMA_VERSION,
    WORKBOOK_SCHEMA_VERSION,
    prepare_occurrence_rows,
)
from runtime_source_validation import (  # noqa: E402
    compute_actual_asset_fingerprint,
    compute_expected_asset_fingerprint,
    sha256_file,
    validate_asset_manifest,
)


PDF_NAME = "repair-v562-fixture.pdf"
SESSION_ID = "session-v561-compatible-repair-fixture"
RELEVANT_GLYPH_SHA = "a" * 64
UNRELATED_GLYPH_SHA = "b" * 64

ACTUAL_COLUMNS = [
    "實體頁碼", "課本頁", "頁面標籤", "字元", "實際注音", "解碼狀態", "解碼依據",
    "穩定注音鍵", "群組注音鍵", "font", "font_xref", "注音元件ID",
    "glyph_id_字形索引", "x0", "y0", "x1", "y1", "TTF字形SHA256",
    "occurrence_id", "review_id", "identity_confidence", "identity_row_fallback",
    "identity_collision_base", "source_row_number", "ledger_schema_version", "workbook_schema_version",
]

EXCLUDED_COLUMNS = [
    "實體頁碼", "課本頁", "字元", "font", "font_xref", "glyph_id", "候選注音元件ID",
    "x0", "y0", "來源", "排除理由", "occurrence_id", "review_id", "identity_confidence",
    "identity_row_fallback", "source_row_number", "ledger_schema_version", "workbook_schema_version",
]

APPROVED_REUSABLE_RULES = {
    ("額頭", 1, "頭"): (
        "USER-EXPECTED-8e23e6643db8887ccabf", "˙ㄊㄡ", "簡編本", "2026-08-25T13:49:46"
    ),
    ("OK繃", 2, "繃"): (
        "USER-EXPECTED-04c50c1c3bf9f7d609b1", "ㄅㄥ", "簡編本", "2026-08-25T13:54:44"
    ),
    ("一會兒", 2, "兒"): (
        "USER-EXPECTED-6b8ac667a7d5fdb257db", "ㄦ", "簡編本", "2026-08-25T14:20:35"
    ),
}


def _write_csv(path: Path, fieldnames: list[str], rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows([{field: row.get(field, "") for field in fieldnames} for row in rows])


def _refresh_fixture_manifest(runtime_root: Path) -> None:
    path = runtime_root / "runtime_asset_manifest.json"
    manifest = json.loads(path.read_text(encoding="utf-8"))
    for asset in manifest["assets"]:
        asset["sha256"] = sha256_file(runtime_root / asset["path"])
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _actual_payload_sha256(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        payload = {
            sheet: [list(row) for row in workbook[sheet].iter_rows(values_only=True)]
            for sheet in ("實際注音", "結構偵測排除")
        }
    finally:
        workbook.close()
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def _summary(path: Path, sheet_name: str = "摘要") -> dict[str, object]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        return {
            str(row[0]): row[1] if len(row) > 1 else None
            for row in sheet.iter_rows(values_only=True)
            if row and row[0] not in (None, "")
        }
    finally:
        workbook.close()


class ProjectRepairIntegrationV562Tests(unittest.TestCase):
    """A sealed v5.6.1-compatible project repaired by the real v5.6.2 pipeline."""

    @classmethod
    def setUpClass(cls):
        cls._temporary = tempfile.TemporaryDirectory()
        cls.fixture_root = Path(cls._temporary.name)
        cls.runtime_root = cls.fixture_root / "runtime-current"
        cls.runtime_root.mkdir()
        cls.pdf_path = cls.fixture_root / PDF_NAME
        cls._create_pdf(cls.pdf_path)
        cls.pdf_sha256 = sha256_file(cls.pdf_path)
        cls._copy_runtime_fixture(cls.runtime_root)
        cls._install_regression_fixture(cls.runtime_root)
        _refresh_fixture_manifest(cls.runtime_root)
        validation = validate_asset_manifest(cls.runtime_root)
        if not validation.get("ok"):
            raise AssertionError(validation.get("errors"))
        cls.old_expected_fingerprint = compute_expected_asset_fingerprint(
            cls.runtime_root,
            validation,
            resolver_version=proof.EXPECTED_RESOLVER_VERSION,
            source_files=EXPECTED_RESOLVER_SOURCE_FILES,
        )["fingerprint"]

        cls.baseline_output = cls.fixture_root / "baseline-project"
        cls._create_v561_compatible_project(cls.baseline_output, validation)
        cls._change_expected_truth(cls.runtime_root)
        _refresh_fixture_manifest(cls.runtime_root)
        validation = validate_asset_manifest(cls.runtime_root)
        if not validation.get("ok"):
            raise AssertionError(validation.get("errors"))
        cls.current_expected_fingerprint = compute_expected_asset_fingerprint(
            cls.runtime_root,
            validation,
            resolver_version=proof.EXPECTED_RESOLVER_VERSION,
            source_files=EXPECTED_RESOLVER_SOURCE_FILES,
        )["fingerprint"]
        if cls.current_expected_fingerprint == cls.old_expected_fingerprint:
            raise AssertionError("expected fixture fingerprint did not change")

    @classmethod
    def tearDownClass(cls):
        cls._temporary.cleanup()

    @classmethod
    def _copy_runtime_fixture(cls, destination: Path) -> None:
        manifest = json.loads((ROOT / "runtime_asset_manifest.json").read_text(encoding="utf-8"))
        names = {
            "runtime_asset_manifest.json",
            "可重用expected規則.json",
            *(str(asset["path"]) for asset in manifest["assets"]),
            *EXPECTED_RESOLVER_SOURCE_FILES,
            *proof.ACTUAL_DECODER_SOURCE_FILES,
        }
        for name in sorted(names):
            source = ROOT / name
            if not source.exists():
                raise AssertionError(f"fixture source missing: {source}")
            target = destination / name
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, target)

    @classmethod
    def _create_pdf(cls, path: Path) -> None:
        lines = [
            "角色",
            "第一個一級棒",
            "第二個一級棒",
            "想一想",
            "為什麼",
            "下載",
            "龘",
            "額頭",
            "OK繃",
            "一會兒",
        ]
        document = fitz.open()
        page = document.new_page()
        for index, line in enumerate(lines):
            page.insert_text((72, 72 + index * 24), line, fontname="china-t", fontsize=14)
        document.save(path)
        document.close()

    @classmethod
    def _install_regression_fixture(cls, runtime_root: Path) -> None:
        path = runtime_root / DEFAULT_REGRESSIONS.name
        with path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            fieldnames = list(reader.fieldnames or [])
            rows = list(reader)
        replacements = [
            {
                "case_id": "RPR-HARD-SHA", "pdf_contains": "historical-renamed-book", "page": "1",
                "phrase": "角色", "target_char": "角", "target_occurrence_index": "0",
                "locator_context": "角色", "actual_reading": "ㄐㄩㄝˊ", "expected_reading": "ㄐㄩㄝˊ",
                "status": "確認正確", "source": "repair integration fixture", "note": "exact bytes hard gate",
                "applicability_mode": "hard_gate", "source_pdf_sha256": cls.pdf_sha256,
            },
            {
                "case_id": "RPR-PORTABLE", "pdf_contains": "repair-v562-fixture", "page": "999",
                "phrase": "一級棒", "target_char": "一", "target_occurrence_index": "0",
                "locator_context": "第一個一級棒", "actual_reading": "ㄧˋ", "expected_reading": "ㄧˋ",
                "status": "確認正確", "source": "repair integration fixture", "note": "unique portable locator",
                "applicability_mode": "portable_gate", "source_pdf_sha256": "",
            },
            {
                "case_id": "RPR-REF-PASS", "pdf_contains": "repair-v562-fixture", "page": "1",
                "phrase": "想一想", "target_char": "一", "target_occurrence_index": "0",
                "locator_context": "想一想", "actual_reading": "ㄧˋ", "expected_reading": "ㄧˋ",
                "status": "確認正確", "source": "repair integration fixture", "note": "reference pass",
                "applicability_mode": "reference_only", "source_pdf_sha256": "",
            },
            {
                "case_id": "RPR-REF-FAIL", "pdf_contains": "repair-v562-fixture", "page": "1",
                "phrase": "為什麼", "target_char": "為", "target_occurrence_index": "0",
                "locator_context": "為什麼", "actual_reading": "ㄨㄟˊ", "expected_reading": "ㄨㄟˋ",
                "status": "確認錯誤", "source": "repair integration fixture", "note": "new expected truth invalidates reference",
                "applicability_mode": "reference_only", "source_pdf_sha256": "",
            },
            {
                "case_id": "RPR-REF-NOT-FOUND", "pdf_contains": "repair-v562-fixture", "page": "1",
                "phrase": "不存在", "target_char": "不", "target_occurrence_index": "0",
                "locator_context": "不存在", "actual_reading": "ㄅㄨˋ", "expected_reading": "ㄅㄨˋ",
                "status": "確認正確", "source": "repair integration fixture", "note": "reference not found",
                "applicability_mode": "reference_only", "source_pdf_sha256": "",
            },
            {
                "case_id": "RPR-REF-DIFFERENCE", "pdf_contains": "repair-v562-fixture", "page": "1",
                "phrase": "下載", "target_char": "載", "target_occurrence_index": "0",
                "locator_context": "下載", "actual_reading": "ㄗㄞˇ", "expected_reading": "ㄗㄞˋ",
                "status": "確認錯誤", "source": "repair integration fixture", "note": "reference difference",
                "applicability_mode": "reference_only", "source_pdf_sha256": "",
            },
        ]
        if len(rows) != 166:
            raise AssertionError(f"unexpected production regression row count: {len(rows)}")
        rows[: len(replacements)] = replacements
        _write_csv(path, fieldnames, rows)

    @classmethod
    def _pdf_target_rows(cls) -> list[dict[str, object]]:
        targets = [
            ("角色", "角", 0, "ㄐㄩㄝˊ"),
            ("第一個一級棒", "一", 3, "ㄧˋ"),
            ("第二個一級棒", "一", 3, "ㄧˋ"),
            ("想一想", "一", 1, "ㄧˋ"),
            ("為什麼", "為", 0, "ㄨㄟˊ"),
            ("下載", "載", 1, "ㄗㄞˇ"),
            ("龘", "龘", 0, "ㄉㄚˊ"),
            ("額頭", "頭", 1, "˙ㄊㄡ"),
            ("OK繃", "繃", 2, "ㄅㄥ"),
            ("一會兒", "兒", 2, "ㄦ"),
        ]
        document = fitz.open(cls.pdf_path)
        try:
            raw = document[0].get_text("rawdict")
            line_chars: dict[str, list[dict[str, object]]] = {}
            for block in raw.get("blocks", []):
                for line in block.get("lines", []):
                    chars = [char for span in line.get("spans", []) for char in span.get("chars", [])]
                    line_chars["".join(str(char.get("c") or "") for char in chars)] = chars
        finally:
            document.close()
        rows: list[dict[str, object]] = []
        for index, (line, target, offset, reading) in enumerate(targets, 1):
            chars = line_chars[line]
            bbox = chars[offset]["bbox"]
            rows.append({
                "實體頁碼": 1,
                "課本頁": 1,
                "頁面標籤": "1",
                "字元": target,
                "實際注音": reading,
                "解碼狀態": "已映射",
                "解碼依據": "synthetic independent glyph evidence",
                "穩定注音鍵": f"fixture-key-{index}",
                "群組注音鍵": f"fixture-group-{index}",
                "font": "fixture-cjk-font",
                "font_xref": 1,
                "注音元件ID": f"fixture-component-{index}",
                "glyph_id_字形索引": index,
                "x0": float(bbox[0]),
                "y0": float(bbox[1]),
                "x1": float(bbox[2]),
                "y1": float(bbox[3]),
                "TTF字形SHA256": RELEVANT_GLYPH_SHA if index == 1 else "",
                "source_row_number": index + 1,
                "ledger_schema_version": LEDGER_SCHEMA_VERSION,
                "workbook_schema_version": WORKBOOK_SCHEMA_VERSION,
            })
        prepare_occurrence_rows(cls.pdf_sha256, rows)
        return rows

    @classmethod
    def _write_actual_workbook(
        cls,
        path: Path,
        rows: list[dict[str, object]],
        fingerprint: dict[str, object],
    ) -> None:
        workbook = Workbook()
        actual = workbook.active
        actual.title = "實際注音"
        actual.append(ACTUAL_COLUMNS)
        for row in rows:
            actual.append([row.get(column, "") for column in ACTUAL_COLUMNS])
        excluded = workbook.create_sheet("結構偵測排除")
        excluded.append(EXCLUDED_COLUMNS)
        metadata = workbook.create_sheet("v5.2中繼資料")
        metadata.append(["項目", "內容"])
        for key, value in (
            ("workbook_schema_version", WORKBOOK_SCHEMA_VERSION),
            ("ledger_schema_version", LEDGER_SCHEMA_VERSION),
            ("pdf_sha256", cls.pdf_sha256),
            ("actual_asset_fingerprint", fingerprint["fingerprint"]),
            (
                "actual_asset_fingerprint_components",
                json.dumps(fingerprint["components"], ensure_ascii=False, sort_keys=True),
            ),
        ):
            metadata.append([key, value])
        summary = workbook.create_sheet("摘要")
        summary.append(["項目", "內容"])
        summary.append(["偵測到注音字形筆數", len(rows)])
        workbook.save(path)
        workbook.close()

    @classmethod
    def _create_v561_compatible_project(cls, output: Path, validation: dict[str, object]) -> None:
        actual_dir = output / "01_實際注音"
        candidate_dir = output / "02_候選報告"
        actual_dir.mkdir(parents=True)
        candidate_dir.mkdir(parents=True)
        dynamic_root = proof.initialize_project_actual_evidence(output, cls.runtime_root)
        dependencies = {"ttf_glyph_sha256": [RELEVANT_GLYPH_SHA], "cff_glyph_keys": []}
        fingerprint = compute_actual_asset_fingerprint(
            cls.runtime_root,
            cls.pdf_path,
            validation,
            decoder_version=proof.ACTUAL_DECODER_VERSION,
            source_files=proof.ACTUAL_DECODER_SOURCE_FILES,
            dynamic_dependencies=dependencies,
            dynamic_evidence_root=dynamic_root,
        )
        rows = cls._pdf_target_rows()
        actual_path = proof.actual_workbook_path(actual_dir, cls.pdf_path)
        cls._write_actual_workbook(actual_path, rows, fingerprint)
        candidate_path = proof.candidate_workbook_path(candidate_dir, cls.pdf_path)
        analyze(
            actual_path,
            cls.runtime_root / DEFAULT_DICT.name,
            cls.runtime_root / DEFAULT_RULES.name,
            candidate_path,
            cls.pdf_path,
            cls.runtime_root / DEFAULT_REGRESSIONS.name,
            cls.runtime_root / DEFAULT_CHAR_OVERRIDES.name,
            cls.runtime_root / DEFAULT_CONTEXT_OVERRIDES.name,
            dynamic_root,
            runtime_root=cls.runtime_root,
        )
        manifest = proof.collect_manifest(
            [cls.pdf_path],
            actual_dir,
            candidate_dir,
            actual_fingerprints={cls.pdf_path.name: fingerprint},
            source_validation=validation,
            runtime_root=cls.runtime_root,
        )
        manifest["version"] = "5.6.1"
        manifest["session_id"] = SESSION_ID
        proof.seal_manifest(manifest)

        records = manifest["records"]

        def find(phrase: str, char: str) -> dict[str, object]:
            matches = [
                row for row in records
                if row.get("char") == char and phrase in str((row.get("source_record") or {}).get("所在行") or "")
            ]
            if len(matches) != 1:
                raise AssertionError(f"fixture occurrence lookup failed: {phrase}/{char}: {len(matches)}")
            return matches[0]

        db = proof.normalize_db({})
        dragon = find("龘", "龘")
        db["events"][dragon["review_id"]] = {
            "action": "補建expected證據",
            "expected_set": ["ㄉㄚˊ"],
            "expected_evidence": "教育部異體字字典：龘",
            "context_evidence": "完整詞：龘；目標位置：0",
            "source": "v5.6.1 durable expected fixture",
            "note": "repair replay expected event",
        }
        for phrase, char in (("下載", "載"), ("為什麼", "為")):
            entry = find(phrase, char)
            event = {
                "action": "確認現版差異",
                "expected_set": list(entry.get("expected_set") or []),
                "expected_evidence": str(entry.get("expected_evidence") or "fixture independent rule"),
                "context_evidence": str(entry.get("context_evidence") or phrase),
                "source": "v5.6.1 durable human confirmation fixture",
                "note": f"confirmed before upgrade: {phrase}",
            }
            event.update({gate: True for gate in CONFIRMATION_GATES})
            db["events"][entry["review_id"]] = event

        materialized = proof.materialize_ledger(manifest, db)
        by_line = {
            str((row.get("source_record") or {}).get("所在行") or ""): row
            for row in materialized
        }
        if by_line["下載"]["state"] != "TEXTBOOK_ERROR_CONFIRMED":
            raise AssertionError(by_line["下載"])
        if by_line["為什麼"]["state"] != "TEXTBOOK_ERROR_CONFIRMED":
            raise AssertionError(by_line["為什麼"])
        if by_line["龘"]["state"] != "PASS":
            raise AssertionError(by_line["龘"])

        proof.json_save(output / "校對工作階段.json", manifest)
        proof.json_save(output / "人工判定資料庫.json", db)
        proof.save_pending_json(output, manifest, db)
        proof.generate_report(output, manifest, db, runtime_root=cls.runtime_root)

    @classmethod
    def _change_expected_truth(cls, runtime_root: Path) -> None:
        path = runtime_root / DEFAULT_RULES.name
        with path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            fieldnames = list(reader.fieldnames or [])
            rows = list(reader)
        match = [row for row in rows if row.get("rule_id") == "V13-EX-049"]
        if len(match) != 1 or match[0].get("expected_reading") != "ㄨㄟˋ":
            raise AssertionError("fixture expected seed rule missing")
        match[0]["expected_reading"] = "ㄨㄟˊ"
        match[0]["source"] = "repair integration fixture: independently revised expected rule"
        _write_csv(path, fieldnames, rows)

    def _clone_case(self, name: str) -> tuple[Path, Path]:
        case_root = self.fixture_root / f"case-{name}"
        if case_root.exists():
            shutil.rmtree(case_root)
        case_root.mkdir()
        pdf = case_root / PDF_NAME
        shutil.copy2(self.pdf_path, pdf)
        output = case_root / "project"
        shutil.copytree(self.baseline_output, output)
        manifest_path = output / "校對工作階段.json"
        manifest = proof.json_load_strict(manifest_path)
        for info in manifest["pdfs"]:
            info["pdf"] = str(pdf)
            info["actual_workbook"] = str(output / "01_實際注音" / Path(info["actual_workbook"]).name)
            info["candidate_workbook"] = str(output / "02_候選報告" / Path(info["candidate_workbook"]).name)
        for record in manifest["records"]:
            record["pdf"] = str(pdf)
            record["actual_workbook"] = str(output / "01_實際注音" / Path(record["actual_workbook"]).name)
            record["candidate_workbook"] = str(output / "02_候選報告" / Path(record["candidate_workbook"]).name)
        proof.seal_manifest(manifest)
        proof.json_save(manifest_path, manifest)
        return output, pdf

    def _runtime_variant(self, name: str, case_id: str, **changes: str) -> Path:
        destination = self.fixture_root / f"runtime-{name}"
        if destination.exists():
            shutil.rmtree(destination)
        shutil.copytree(self.runtime_root, destination, copy_function=os.link)
        regression_path = destination / DEFAULT_REGRESSIONS.name
        with regression_path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            fieldnames = list(reader.fieldnames or [])
            rows = list(reader)
        selected = [row for row in rows if row.get("case_id") == case_id]
        self.assertEqual(len(selected), 1)
        selected[0].update(changes)
        regression_path.unlink()
        _write_csv(regression_path, fieldnames, rows)
        manifest_path = destination / "runtime_asset_manifest.json"
        manifest_text = manifest_path.read_text(encoding="utf-8")
        manifest_path.unlink()
        manifest_path.write_text(manifest_text, encoding="utf-8")
        _refresh_fixture_manifest(destination)
        validation = validate_asset_manifest(destination)
        self.assertTrue(validation["ok"], validation.get("errors"))
        return destination

    @staticmethod
    def _record_key(record: dict[str, object]) -> tuple[str, str, int]:
        source = record.get("source_record") or {}
        return (
            str(source.get("所在行") or ""),
            str(record.get("char") or ""),
            int(source.get("行內字元位置") or 0),
        )

    @staticmethod
    def _append_glyph_evidence(output: Path, glyph_sha: str) -> None:
        path = proof.project_actual_evidence_root(output) / USER_GLYF_FILE
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
        rows.append({
            "glyph_sha256": glyph_sha,
            "bopomofo": "ㄐㄩㄝˊ",
            "verification_level": "VERIFIED_EXACT_GLYPH",
            "source_count": "1",
            "source_examples": "repair integration fixture",
            "updated_at": "2026-08-26T00:00:00",
            "notes": "test-only dynamic actual evidence",
        })
        _write_csv(path, USER_GLYF_HEADERS, rows)

    def test_repair_reuses_actual_rebuilds_expected_replays_events_and_recomputes_completion(self):
        output, pdf = self._clone_case("success")
        manifest_before = proof.json_load_strict(output / "校對工作階段.json")
        db_before = proof.json_load_strict(output / "人工判定資料庫.json")
        info_before = manifest_before["pdfs"][0]
        actual_path = Path(info_before["actual_workbook"])
        candidate_path = Path(info_before["candidate_workbook"])
        dependencies = actual_workbook_dynamic_dependencies(actual_path)
        evidence_hashes_before = dynamic_actual_hashes(
            proof.project_actual_evidence_root(output), pdf_path=pdf, dependencies=dependencies
        )
        before = {
            "session_id": manifest_before["session_id"],
            "pdf_sha256": sha256_file(pdf),
            "actual_workbook_sha256": sha256_file(actual_path),
            "actual_authoritative_payload_sha256": _actual_payload_sha256(actual_path),
            "candidate_workbook_sha256": sha256_file(candidate_path),
            "actual_asset_fingerprint": info_before["actual_asset_fingerprint"],
            "identities": {
                self._record_key(record): (record["occurrence_id"], record["review_id"])
                for record in manifest_before["records"]
            },
            "project_actual_evidence_hashes": evidence_hashes_before,
        }
        proof.json_save(output / "pipeline_status.json", {
            "status": "PROCESSING_FINISHED", "completion_gate": {"failed_gates": ["stale pre-repair"]}
        })
        proof.json_save(output / "待人工確認.json", {"pending": [{"stale": True}], "expected_gaps": []})

        with patch.object(proof, "decode", side_effect=AssertionError("actual decoder must not run")) as decoder:
            report = proof.repair_project_state(output, runtime_root=self.runtime_root)
        self.assertEqual(decoder.call_count, 0)
        self.assertEqual(report, output / "注音校對_最終報告.xlsx")

        manifest_after = proof.json_load_strict(output / "校對工作階段.json")
        proof.validate_manifest_integrity(manifest_after)
        proof.validate_output_artifact_hashes(manifest_after)
        info_after = manifest_after["pdfs"][0]
        after = {
            "session_id": manifest_after["session_id"],
            "pdf_sha256": sha256_file(pdf),
            "actual_workbook_sha256": sha256_file(actual_path),
            "actual_authoritative_payload_sha256": _actual_payload_sha256(actual_path),
            "candidate_workbook_sha256": sha256_file(candidate_path),
            "actual_asset_fingerprint": info_after["actual_asset_fingerprint"],
            "identities": {
                self._record_key(record): (record["occurrence_id"], record["review_id"])
                for record in manifest_after["records"]
            },
            "project_actual_evidence_hashes": dynamic_actual_hashes(
                proof.project_actual_evidence_root(output), pdf_path=pdf, dependencies=dependencies
            ),
        }
        for key in (
            "session_id", "pdf_sha256", "actual_workbook_sha256", "actual_authoritative_payload_sha256",
            "actual_asset_fingerprint", "identities", "project_actual_evidence_hashes",
        ):
            self.assertEqual(after[key], before[key], key)
        self.assertNotEqual(after["candidate_workbook_sha256"], before["candidate_workbook_sha256"])
        self.assertEqual(manifest_after["expected_asset_fingerprint"], self.current_expected_fingerprint)
        self.assertNotEqual(manifest_after["expected_asset_fingerprint"], self.old_expected_fingerprint)

        db_after = proof.json_load_strict(output / "人工判定資料庫.json")
        self.assertEqual(db_after["events"], db_before["events"])
        ledger = proof.materialize_ledger(manifest_after, db_after)
        by_line = {
            str((row.get("source_record") or {}).get("所在行") or ""): row
            for row in ledger
        }
        self.assertEqual(by_line["龘"]["state"], "PASS")
        self.assertEqual(by_line["龘"]["expected_set"], ["ㄉㄚˊ"])
        self.assertEqual(by_line["下載"]["state"], "TEXTBOOK_ERROR_CONFIRMED")
        self.assertEqual(by_line["為什麼"]["state"], "PASS")
        self.assertEqual(by_line["為什麼"]["expected_set"], ["ㄨㄟˊ"])
        self.assertEqual(by_line["為什麼"]["review_event_replay_status"], "INVALIDATED_EXPECTED_DRIFT")

        rules = {
            (rule["phrase"], rule["target_index"], rule["target_char"]): rule
            for rule in manifest_after["reusable_expected_rules"]["rules"]
        }
        self.assertEqual(set(rules), set(APPROVED_REUSABLE_RULES))
        for key, (rule_id, expected, evidence, approved_at) in APPROVED_REUSABLE_RULES.items():
            self.assertEqual(rules[key]["rule_id"], rule_id)
            self.assertEqual(rules[key]["expected_set"], [expected])
            self.assertEqual(rules[key]["evidence"], evidence)
            self.assertEqual(rules[key]["approved_at"], approved_at)
        self.assertEqual(by_line["額頭"]["expected_set"], ["˙ㄊㄡ"])
        self.assertEqual(by_line["OK繃"]["expected_set"], ["ㄅㄥ"])
        self.assertEqual(by_line["一會兒"]["expected_set"], ["ㄦ"])

        mandatory = manifest_after["mandatory_regression"]
        self.assertEqual((mandatory["required"], mandatory["executed"], mandatory["failed"]), (36, 36, 0))
        regression = manifest_after["pdf_regression"]
        results = {row["案例ID"]: row for row in regression["results"]}
        self.assertEqual(results["RPR-HARD-SHA"]["回歸結果"], "PASS")
        self.assertEqual(results["RPR-PORTABLE"]["回歸結果"], "PASS")
        self.assertEqual(results["RPR-REF-PASS"]["回歸結果"], "PASS")
        self.assertEqual(results["RPR-REF-FAIL"]["回歸結果"], "FAIL")
        self.assertEqual(results["RPR-REF-NOT-FOUND"]["回歸結果"], "NOT_APPLICABLE")
        self.assertEqual(results["RPR-REF-DIFFERENCE"]["回歸結果"], "PASS")
        self.assertEqual((regression["required"], regression["executed"], regression["failed"]), (2, 2, 0))
        self.assertEqual(regression["reference_audit_total"], 4)
        self.assertEqual(regression["reference_audit_failed"], 1)
        self.assertEqual(regression["reference_audit_not_found"], 1)

        status = proof.json_load_strict(output / "pipeline_status.json")
        pending = proof.json_load_strict(output / "待人工確認.json")
        self.assertEqual(status["status"], PROOFREAD_COMPLETE)
        self.assertEqual(status["completion_gate"]["failed_gates"], [])
        self.assertEqual(pending["pending"], [])
        self.assertEqual(pending["expected_gaps"], [])
        self.assertTrue((output / "注音校對_技術稽核.xlsx").exists())
        self.assertEqual(_summary(output / "注音校對_技術稽核.xlsx")["pipeline status"], PROOFREAD_COMPLETE)
        self.assertEqual(
            _summary(output / "注音校對_最終報告.xlsx", "校對摘要")["目前狀態"],
            "全冊注音校對完成",
        )
        def reportable(snapshot: dict[str, object]) -> dict[str, object]:
            result = dict(snapshot)
            result["identities"] = {
                " | ".join(map(str, key)): list(value)
                for key, value in snapshot["identities"].items()
            }
            return result

        self.__class__.last_success_evidence = {
            "before": reportable(before),
            "after": reportable(after),
            "decoder_calls": decoder.call_count,
        }

    def test_portable_zero_match_is_not_executed_and_repair_fails_closed(self):
        output, _ = self._clone_case("portable-zero")
        runtime = self._runtime_variant(
            "portable-zero", "RPR-PORTABLE", phrase="不存在", target_char="不",
            target_occurrence_index="0", locator_context="不存在", actual_reading="ㄅㄨˋ", expected_reading="ㄅㄨˋ",
        )
        with patch.object(proof, "decode", side_effect=AssertionError("actual decoder must not run")) as decoder:
            proof.repair_project_state(output, runtime_root=runtime)
        self.assertEqual(decoder.call_count, 0)
        manifest = proof.json_load_strict(output / "校對工作階段.json")
        row = next(item for item in manifest["pdf_regression"]["results"] if item["案例ID"] == "RPR-PORTABLE")
        self.assertEqual(row["回歸結果"], "NOT_EXECUTED")
        self.assertEqual(row["適用性狀態"], "GATE_NOT_EXECUTED")
        self.assertEqual(proof.json_load_strict(output / "pipeline_status.json")["status"], "PROCESSING_FINISHED")

    def test_portable_ambiguous_match_fails_without_selecting_an_occurrence(self):
        output, _ = self._clone_case("portable-ambiguous")
        runtime = self._runtime_variant("portable-ambiguous", "RPR-PORTABLE", locator_context="")
        with patch.object(proof, "decode", side_effect=AssertionError("actual decoder must not run")) as decoder:
            proof.repair_project_state(output, runtime_root=runtime)
        self.assertEqual(decoder.call_count, 0)
        manifest = proof.json_load_strict(output / "校對工作階段.json")
        row = next(item for item in manifest["pdf_regression"]["results"] if item["案例ID"] == "RPR-PORTABLE")
        self.assertEqual(row["回歸結果"], "FAIL")
        self.assertEqual(row["適用性狀態"], "IDENTITY_AMBIGUITY")
        self.assertEqual(row["定位命中數"], 2)
        self.assertEqual(proof.json_load_strict(output / "pipeline_status.json")["status"], "PROCESSING_FINISHED")

    def test_exact_sha_hard_gate_not_found_is_blocking(self):
        output, _ = self._clone_case("hard-not-found")
        runtime = self._runtime_variant(
            "hard-not-found", "RPR-HARD-SHA", phrase="不存在", target_char="不",
            target_occurrence_index="0", locator_context="不存在", actual_reading="ㄅㄨˋ", expected_reading="ㄅㄨˋ",
        )
        with patch.object(proof, "decode", side_effect=AssertionError("actual decoder must not run")) as decoder:
            proof.repair_project_state(output, runtime_root=runtime)
        self.assertEqual(decoder.call_count, 0)
        manifest = proof.json_load_strict(output / "校對工作階段.json")
        row = next(item for item in manifest["pdf_regression"]["results"] if item["案例ID"] == "RPR-HARD-SHA")
        self.assertEqual(row["回歸結果"], "NOT_EXECUTED")
        self.assertEqual(row["適用性狀態"], "GATE_NOT_EXECUTED")
        self.assertEqual(proof.json_load_strict(output / "pipeline_status.json")["status"], "PROCESSING_FINISHED")

    def test_pdf_bytes_change_rejects_reuse_before_pipeline(self):
        output, pdf = self._clone_case("pdf-changed")
        pdf.write_bytes(pdf.read_bytes() + b"changed")
        with patch.object(proof, "decode", side_effect=AssertionError("resolver must reject before decode")) as decoder:
            with self.assertRaises(FileNotFoundError):
                proof.repair_project_state(output, runtime_root=self.runtime_root)
        self.assertEqual(decoder.call_count, 0)

    def test_relevant_actual_evidence_invalidates_only_actual_and_calls_decoder(self):
        output, _ = self._clone_case("relevant-actual")
        manifest = proof.json_load_strict(output / "校對工作階段.json")
        info = manifest["pdfs"][0]
        actual_path = Path(info["actual_workbook"])
        candidate_path = Path(info["candidate_workbook"])
        actual_sha = sha256_file(actual_path)
        candidate_sha = sha256_file(candidate_path)
        self._append_glyph_evidence(output, RELEVANT_GLYPH_SHA)
        with patch.object(proof, "decode", side_effect=RuntimeError("DECODER_SENTINEL")) as decoder:
            with self.assertRaisesRegex(RuntimeError, "DECODER_SENTINEL"):
                proof.repair_project_state(output, runtime_root=self.runtime_root)
        self.assertEqual(decoder.call_count, 1)
        self.assertEqual(sha256_file(actual_path), actual_sha)
        self.assertEqual(sha256_file(candidate_path), candidate_sha)

    def test_unrelated_actual_evidence_keeps_per_pdf_actual_reusable(self):
        output, _ = self._clone_case("unrelated-actual")
        manifest = proof.json_load_strict(output / "校對工作階段.json")
        actual_path = Path(manifest["pdfs"][0]["actual_workbook"])
        actual_sha = sha256_file(actual_path)
        self._append_glyph_evidence(output, UNRELATED_GLYPH_SHA)
        with patch.object(proof, "decode", side_effect=AssertionError("unrelated evidence must not decode")) as decoder:
            proof.repair_project_state(output, runtime_root=self.runtime_root)
        self.assertEqual(decoder.call_count, 0)
        self.assertEqual(sha256_file(actual_path), actual_sha)

    def test_tampered_actual_workbook_is_rejected_by_artifact_hash(self):
        output, _ = self._clone_case("tampered-actual")
        manifest = proof.json_load_strict(output / "校對工作階段.json")
        actual_path = Path(manifest["pdfs"][0]["actual_workbook"])
        workbook = load_workbook(actual_path)
        workbook["實際注音"].cell(2, 5, "ㄐㄧㄠˇ")
        workbook.save(actual_path)
        workbook.close()
        with patch.object(proof, "decode", side_effect=AssertionError("integrity must reject before decode")) as decoder:
            with self.assertRaisesRegex(ValueError, "輸出資產 hash 不符"):
                proof.repair_project_state(output, runtime_root=self.runtime_root)
        self.assertEqual(decoder.call_count, 0)

    def test_incompatible_session_and_review_identity_schemas_are_rejected(self):
        for label, field, value in (
            ("session-schema", "session_schema_version", "99.0"),
            ("review-schema", "review_id_schema_version", "99.0"),
        ):
            with self.subTest(field=field):
                output, _ = self._clone_case(label)
                manifest_path = output / "校對工作階段.json"
                manifest = proof.json_load_strict(manifest_path)
                manifest[field] = value
                proof.seal_manifest(manifest)
                proof.json_save(manifest_path, manifest)
                with patch.object(proof, "decode", side_effect=AssertionError("schema must reject before decode")) as decoder:
                    with self.assertRaisesRegex(ValueError, "SESSION_SCHEMA_INCOMPATIBLE"):
                        proof.repair_project_state(output, runtime_root=self.runtime_root)
                self.assertEqual(decoder.call_count, 0)

        output, _ = self._clone_case("decision-review-schema")
        manifest = proof.json_load_strict(output / "校對工作階段.json")
        candidate_path = Path(manifest["pdfs"][0]["candidate_workbook"])
        candidate_sha = sha256_file(candidate_path)
        db_path = output / "人工判定資料庫.json"
        db = proof.json_load_strict(db_path)
        db["review_id_schema_version"] = "99.0"
        proof.json_save(db_path, db)
        with patch.object(proof, "decode", side_effect=AssertionError("decision schema must reject before decode")) as decoder:
            with self.assertRaisesRegex(ValueError, "DECISION_SCHEMA_INCOMPATIBLE"):
                proof.repair_project_state(output, runtime_root=self.runtime_root)
        self.assertEqual(decoder.call_count, 0)
        self.assertEqual(sha256_file(candidate_path), candidate_sha)


if __name__ == "__main__":
    unittest.main(verbosity=2)
