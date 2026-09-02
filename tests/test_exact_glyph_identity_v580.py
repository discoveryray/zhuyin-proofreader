from __future__ import annotations

import ast
import copy
import hashlib
import json
import struct
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from actual_review import (
    ACTUAL_REVIEW_SCHEMA_VERSION,
    LEGACY_MANUAL_ACTUAL_STAGING_SCHEMA_VERSION,
    MANUAL_ACTUAL_STAGING_SCHEMA_VERSION,
    apply_staged_manual_actual_batch,
    apply_verified_actual_group,
    actual_group_identity,
    build_actual_group_for_entry,
    build_actual_review_groups,
    import_actual_review_workbook,
    load_manual_actual_staging,
    manual_actual_staging_path,
    stage_manual_actual_group,
)
from exact_glyph_identity import (
    CANONICAL_EXACT_IDENTITY,
    CFF_GLYPH_SHA256,
    GLOBAL_ELIGIBLE_SIMPLE_GLYF_V1,
    NON_GLOBAL_ELIGIBLE,
    TTF_GLYF_SHA256,
    UNSUPPORTED_EXACT_IDENTITY,
    actual_group_id,
    canonical_actual_group_identity,
    canonical_cff_exact_identity,
    classify_ttf_glyf_record,
    legacy_v1_actual_group_id,
)
from export_pdf_text_diagnostics import TrueTypeGlyphInspector


ROOT = Path(__file__).resolve().parents[1]
SHA_A = "a" * 64
SHA_B = "b" * 64


def payload_sha256(payload: dict) -> str:
    raw = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def simple_glyph(x: int = 0) -> bytes:
    if not 0 <= x <= 255:
        raise ValueError("synthetic fixture x must fit a positive short vector")
    header = struct.pack(">hhhhh", 1, 0, 0, x, 0)
    end_points = struct.pack(">H", 0)
    instruction_length = struct.pack(">H", 0)
    if x:
        # ON_CURVE | X_SHORT | X_POSITIVE | Y_SAME, then one x byte.
        body = bytes((0x33, x))
    else:
        # ON_CURVE | X_SAME | Y_SAME.  The final zero is legal alignment padding.
        body = bytes((0x31, 0x00))
    return header + end_points + instruction_length + body


def composite_glyph(component_gid: int = 1) -> bytes:
    header = struct.pack(">hhhhh", -1, 0, 0, 0, 0)
    # ARGS_ARE_XY_VALUES with byte-sized zero offsets; no more components.
    return header + struct.pack(">HHbb", 0x0002, component_gid, 0, 0)


def synthetic_sfnt(glyph_records: list[bytes], *, font_name_marker: bytes) -> bytes:
    head = bytearray(54)
    struct.pack_into(">h", head, 50, 1)  # long loca
    maxp = struct.pack(">IH", 0x00010000, len(glyph_records))
    offsets = [0]
    glyf = bytearray()
    for record in glyph_records:
        glyf.extend(record)
        offsets.append(len(glyf))
    loca = struct.pack(f">{len(offsets)}I", *offsets)
    tables = [
        (b"head", bytes(head)),
        (b"maxp", maxp),
        (b"loca", loca),
        (b"glyf", bytes(glyf)),
        (b"name", bytes(font_name_marker)),
    ]

    directory_size = 12 + 16 * len(tables)
    cursor = directory_size
    directory = bytearray()
    table_data = bytearray()
    for tag, data in tables:
        padding = (-cursor) % 4
        if padding:
            table_data.extend(b"\0" * padding)
            cursor += padding
        directory.extend(struct.pack(">4sIII", tag, 0, cursor, len(data)))
        table_data.extend(data)
        cursor += len(data)
    header = struct.pack(">IHHHH", 0x00010000, len(tables), 0, 0, 0)
    return header + bytes(directory) + bytes(table_data)


def ledger_entry(
    occurrence_id: str,
    *,
    ttf_sha: str = "",
    cff_sha: str = "",
    style_group: str = "",
    state: str = "ACTUAL_DECODE_ERROR",
    x0: float = 10.0,
    full_signature: str = "audit-only-signature",
) -> dict:
    return {
        "occurrence_id": occurrence_id,
        "review_id": "review-" + occurrence_id,
        "state": state,
        "pdf_name": "book.pdf",
        "physical_page": 1,
        "printed_page": "1",
        "char": "字",
        "stable_key": "stable-" + occurrence_id,
        "x0": x0,
        "y0": 20.0,
        "actual": "",
        "actual_evidence": "decoder unresolved",
        "source_record": {
            "TTF字形SHA256": ttf_sha,
            "CFF整字字形SHA256": cff_sha,
            "CFF樣式群組": style_group,
            "CFF完整注音簽名": full_signature,
            "穩定注音鍵": "stable-" + occurrence_id,
        },
    }


def stage_group(root: Path, group: dict, reading: str = "ㄅ") -> dict:
    return stage_manual_actual_group(
        root,
        group,
        reading,
        checked_occurrence_ids=[group["members"][0]["occurrence_id"]],
        source="direct visual actual test",
    )


def legacy_single_group_snapshot(group: dict, legacy_group_id: str) -> str:
    return payload_sha256({
        "group_id": legacy_group_id,
        "kind": group["kind"],
        "exact_key": group["exact_key"],
        "members": [
            {
                "occurrence_id": member.get("occurrence_id"),
                "review_id": member.get("review_id"),
                "state": member.get("state"),
                "actual": member.get("actual"),
            }
            for member in group["members"]
        ],
    })


class TTFExactIdentityTests(unittest.TestCase):
    def test_legal_simple_glyf_is_globally_eligible_and_hashes_complete_raw_record(self):
        raw = simple_glyph(23)
        result = classify_ttf_glyf_record(raw)
        self.assertEqual(result["status"], CANONICAL_EXACT_IDENTITY)
        self.assertEqual(result["kind"], TTF_GLYF_SHA256)
        self.assertEqual(result["style_group"], "")
        self.assertEqual(result["eligibility"], GLOBAL_ELIGIBLE_SIMPLE_GLYF_V1)
        self.assertEqual(result["glyph_sha256"], hashlib.sha256(raw).hexdigest())

    def test_same_simple_record_ignores_font_name_and_glyph_index(self):
        shared = simple_glyph(17)
        font_a = synthetic_sfnt([shared], font_name_marker=b"Family A")
        font_b = synthetic_sfnt([simple_glyph(9), shared], font_name_marker=b"Unrelated Family B")
        first = TrueTypeGlyphInspector(font_a).global_exact_identity(0)
        second = TrueTypeGlyphInspector(font_b).global_exact_identity(1)
        self.assertEqual(first, second)
        self.assertEqual(first["eligibility"], GLOBAL_ELIGIBLE_SIMPLE_GLYF_V1)
        self.assertEqual(first["glyph_sha256"], hashlib.sha256(shared).hexdigest())

    def test_equal_composite_raw_sha_with_different_component_outlines_is_never_global_eligible(self):
        shared_composite = composite_glyph(1)
        font_a = TrueTypeGlyphInspector(
            synthetic_sfnt([shared_composite, simple_glyph(11)], font_name_marker=b"A")
        )
        font_b = TrueTypeGlyphInspector(
            synthetic_sfnt([shared_composite, simple_glyph(29)], font_name_marker=b"B")
        )
        self.assertEqual(font_a.glyph_bytes(0), font_b.glyph_bytes(0))
        self.assertEqual(font_a.glyph_sha256(0), font_b.glyph_sha256(0))
        self.assertNotEqual(font_a.glyph_sha256(1), font_b.glyph_sha256(1))
        for result in (font_a.global_exact_identity(0), font_b.global_exact_identity(0)):
            self.assertEqual(result["status"], UNSUPPORTED_EXACT_IDENTITY)
            self.assertEqual(result["eligibility"], NON_GLOBAL_ELIGIBLE)
            self.assertEqual(result["reason"], "COMPOSITE_GLYF")
            self.assertEqual(result["glyph_sha256"], hashlib.sha256(shared_composite).hexdigest())

    def test_truncated_invalid_range_and_illegal_simple_structure_fail_closed(self):
        truncated = classify_ttf_glyf_record(b"\0" * 9)
        illegal = classify_ttf_glyf_record(
            struct.pack(">hhhhhHHB", 1, 0, 0, 0, 0, 0, 0, 0xB1)
        )
        font_bytes = bytearray(
            synthetic_sfnt([simple_glyph(7)], font_name_marker=b"bad-loca")
        )
        probe = TrueTypeGlyphInspector(bytes(font_bytes))
        loca_offset, _ = probe.tables["loca"]
        struct.pack_into(">I", font_bytes, loca_offset + 4, probe.glyf_length + 100)
        invalid_range = TrueTypeGlyphInspector(bytes(font_bytes)).global_exact_identity(0)
        out_of_range = probe.global_exact_identity(99)

        for result in (truncated, illegal, invalid_range, out_of_range):
            self.assertEqual(result["status"], UNSUPPORTED_EXACT_IDENTITY)
            self.assertEqual(result["eligibility"], NON_GLOBAL_ELIGIBLE)
        self.assertEqual(illegal["reason"], "INVALID_SIMPLE_RESERVED_FLAG")
        self.assertEqual(invalid_range["reason"], "INVALID_GLYF_RECORD_RANGE")

    def test_project_local_raw_sha_group_and_promotion_remain_available_for_composite(self):
        raw_sha = hashlib.sha256(composite_glyph()).hexdigest()
        members = [
            ledger_entry("composite-a", ttf_sha=raw_sha, x0=10.0),
            ledger_entry("composite-b", ttf_sha=raw_sha, x0=30.0),
        ]
        group = build_actual_group_for_entry(members, members[0])
        with tempfile.TemporaryDirectory() as directory:
            result = apply_verified_actual_group(
                Path(directory),
                group,
                "ㄅ",
                checked_occurrence_ids=["composite-a", "composite-b"],
                source="direct visual project-local test",
            )
        self.assertEqual(group["kind"], TTF_GLYF_SHA256)
        self.assertEqual(group["occurrence_count"], 2)
        self.assertEqual(result["learning_level"], "VERIFIED_EXACT_GLYPH")


class CFFExactIdentityTests(unittest.TestCase):
    def test_cff_identity_is_style_plus_complete_sha_only(self):
        first = canonical_cff_exact_identity("KAICHU_MD", SHA_A)
        same = canonical_cff_exact_identity("KAICHU_MD", SHA_A.upper())
        other_style = canonical_cff_exact_identity("MING_STD", SHA_A)
        self.assertEqual(first, same)
        self.assertNotEqual(first, other_style)
        self.assertEqual(first["glyph_sha256"], SHA_A)

    def test_same_audit_signature_with_different_complete_sha_is_different_identity(self):
        first = ledger_entry("cff-a", cff_sha=SHA_A, style_group="KAICHU_MD")
        second = ledger_entry("cff-b", cff_sha=SHA_B, style_group="KAICHU_MD")
        self.assertEqual(
            first["source_record"]["CFF完整注音簽名"],
            second["source_record"]["CFF完整注音簽名"],
        )
        self.assertNotEqual(actual_group_identity(first), actual_group_identity(second))
        self.assertEqual(len(build_actual_review_groups([first, second])), 2)

    def test_same_cff_sha_different_style_creates_distinct_groups_ids_and_peers(self):
        first = ledger_entry("cff-style-a", cff_sha=SHA_A, style_group="KAICHU_MD", x0=10.0)
        second = ledger_entry("cff-style-b", cff_sha=SHA_A, style_group="MING_STD", x0=30.0)
        groups = build_actual_review_groups([first, second])
        self.assertEqual(len(groups), 2)
        self.assertEqual(len({group["group_id"] for group in groups}), 2)
        self.assertEqual({group["occurrence_count"] for group in groups}, {1})
        self.assertEqual(
            {group["style_group"] for group in groups},
            {"KAICHU_MD", "MING_STD"},
        )

        rebuilt_first = build_actual_group_for_entry([first, second], first)
        rebuilt_second = build_actual_group_for_entry([first, second], second)
        self.assertEqual(rebuilt_first["occurrence_count"], 1)
        self.assertEqual(rebuilt_second["occurrence_count"], 1)
        self.assertEqual(
            {rebuilt_first["group_id"], rebuilt_second["group_id"]},
            {group["group_id"] for group in groups},
        )

    def test_cff_group_id_and_snapshot_are_explicitly_style_bound(self):
        base = ledger_entry("same-member", cff_sha=SHA_A, style_group="KAICHU_MD")
        other = copy.deepcopy(base)
        other["source_record"]["CFF樣式群組"] = "MING_STD"
        first = build_actual_group_for_entry([base], base)
        second = build_actual_group_for_entry([other], other)
        self.assertNotEqual(first["group_id"], second["group_id"])
        self.assertNotEqual(first["group_snapshot"], second["group_snapshot"])
        self.assertEqual(first["group_id"], actual_group_id(actual_group_identity(base)))

    def test_cff_live_revalidation_rejects_member_from_another_style_before_mutation(self):
        member = ledger_entry("live-cff", cff_sha=SHA_A, style_group="KAICHU_MD")
        group = build_actual_group_for_entry([member], member)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            stage_group(root, group)
            staging_before = manual_actual_staging_path(root).read_bytes()
            live = copy.deepcopy(group)
            live["members"][0]["source_record"]["CFF樣式群組"] = "MING_STD"
            with self.assertRaisesRegex(ValueError, "member 不屬於 canonical exact identity"):
                apply_staged_manual_actual_batch(root, [live])
            self.assertEqual(manual_actual_staging_path(root).read_bytes(), staging_before)
            self.assertEqual(list(root.glob("*.csv")), [])


class ExactIdentitySchemaTests(unittest.TestCase):
    def test_ttf_group_id_and_single_group_snapshot_remain_v1_compatible(self):
        members = [
            ledger_entry("ttf-a", ttf_sha=SHA_A, x0=10.0),
            ledger_entry("ttf-b", ttf_sha=SHA_A, x0=30.0),
        ]
        group = build_actual_group_for_entry(members, members[0])
        legacy_id = "agr_" + payload_sha256({
            "kind": TTF_GLYF_SHA256,
            "exact_key": SHA_A,
        })[:24]
        self.assertEqual(group["group_id"], legacy_id)
        self.assertEqual(group["group_snapshot"], legacy_single_group_snapshot(group, legacy_id))

    def test_new_cff_staging_is_schema_1_1_and_style_bound(self):
        member = ledger_entry("new-cff", cff_sha=SHA_A, style_group="KAICHU_MD")
        group = build_actual_group_for_entry([member], member)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            staged = stage_group(root, group)
            raw = json.loads(manual_actual_staging_path(root).read_text(encoding="utf-8"))
        identity = canonical_actual_group_identity(CFF_GLYPH_SHA256, SHA_A, "KAICHU_MD")
        self.assertEqual(MANUAL_ACTUAL_STAGING_SCHEMA_VERSION, "1.1")
        self.assertEqual(raw["schema_version"], "1.1")
        self.assertEqual(staged["group_id"], actual_group_id(identity))
        self.assertNotEqual(staged["group_id"], legacy_v1_actual_group_id(identity))

    def test_legacy_v1_empty_and_ttf_staging_use_read_only_explicit_adapter(self):
        member = ledger_entry("legacy-ttf", ttf_sha=SHA_A)
        group = build_actual_group_for_entry([member], member)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = manual_actual_staging_path(root)
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(
                json.dumps({
                    "schema_version": LEGACY_MANUAL_ACTUAL_STAGING_SCHEMA_VERSION,
                    "staged_groups": [],
                }),
                encoding="utf-8",
            )
            empty_before = path.read_bytes()
            self.assertEqual(load_manual_actual_staging(root), {
                "schema_version": MANUAL_ACTUAL_STAGING_SCHEMA_VERSION,
                "staged_groups": [],
            })
            self.assertEqual(path.read_bytes(), empty_before)

            path.unlink()
            staged = stage_group(root, group)
            raw = json.loads(path.read_text(encoding="utf-8"))
            raw["schema_version"] = LEGACY_MANUAL_ACTUAL_STAGING_SCHEMA_VERSION
            path.write_text(json.dumps(raw, ensure_ascii=False, sort_keys=True), encoding="utf-8")
            ttf_before = path.read_bytes()
            loaded = load_manual_actual_staging(root)
            self.assertEqual(loaded["schema_version"], MANUAL_ACTUAL_STAGING_SCHEMA_VERSION)
            self.assertEqual(loaded["staged_groups"], [staged])
            self.assertEqual(path.read_bytes(), ttf_before)

    def test_legacy_v1_pending_cff_fails_closed_without_rewrite(self):
        member = ledger_entry("legacy-cff", cff_sha=SHA_A, style_group="KAICHU_MD")
        group = build_actual_group_for_entry([member], member)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            stage_group(root, group)
            path = manual_actual_staging_path(root)
            raw = json.loads(path.read_text(encoding="utf-8"))
            identity = canonical_actual_group_identity(CFF_GLYPH_SHA256, SHA_A, "KAICHU_MD")
            legacy_id = legacy_v1_actual_group_id(identity)
            raw["schema_version"] = LEGACY_MANUAL_ACTUAL_STAGING_SCHEMA_VERSION
            raw["staged_groups"][0]["group_id"] = legacy_id
            raw["staged_groups"][0]["group_snapshot"] = legacy_single_group_snapshot(group, legacy_id)
            path.write_text(json.dumps(raw, ensure_ascii=False, sort_keys=True), encoding="utf-8")
            before = path.read_bytes()
            with self.assertRaisesRegex(ValueError, "v1.0 含 pending CFF decision"):
                load_manual_actual_staging(root)
            self.assertEqual(path.read_bytes(), before)

    def test_unknown_staging_schema_and_stale_snapshot_fail_closed(self):
        member = ledger_entry("stale-ttf", ttf_sha=SHA_A)
        group = build_actual_group_for_entry([member], member)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = manual_actual_staging_path(root)
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(
                json.dumps({"schema_version": "99.0", "staged_groups": []}),
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "schema 不相容"):
                load_manual_actual_staging(root)

            path.unlink()
            stage_group(root, group)
            before = path.read_bytes()
            live = copy.deepcopy(group)
            live["group_snapshot"] = "f" * 64
            with self.assertRaisesRegex(ValueError, "group_snapshot"):
                apply_staged_manual_actual_batch(root, [live])
            self.assertEqual(path.read_bytes(), before)
            self.assertEqual(list(root.glob("*.csv")), [])

    def test_old_actual_review_workbook_schema_is_rejected_before_evidence_write(self):
        member = ledger_entry("review-cff", cff_sha=SHA_A, style_group="KAICHU_MD")
        group = build_actual_group_for_entry([member], member)
        metadata = {
            "version": "5.7.0",
            "session_id": "session",
            "session_schema_version": "2.6.2",
            "workbook_schema_version": "2.6.2",
            "review_id_schema_version": "2.5.0",
        }
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            xlsx = root / "legacy-actual-review.xlsx"
            workbook = Workbook()
            meta = workbook.active
            meta.title = "匯入中繼資料"
            meta.append(["項目", "內容"])
            for key, value in metadata.items():
                meta.append([key, value])
            meta.append(["actual_review_schema_version", "1.0"])
            rows = workbook.create_sheet("actual待判定")
            rows.append([
                "group_id", "group_snapshot", "decision", "actual_reading", "confidence",
                "sample_a_checked", "sample_b_checked", "note",
            ])
            workbook.save(xlsx)
            workbook.close()

            with self.assertRaisesRegex(ValueError, "session/schema 不相容"):
                import_actual_review_workbook(
                    root,
                    xlsx,
                    [group],
                    expected_metadata=metadata,
                )
            self.assertEqual(ACTUAL_REVIEW_SCHEMA_VERSION, "1.1")
            self.assertEqual(list(root.glob("*.csv")), [])

    def test_identity_module_is_actual_only_and_has_no_sqlite_or_expected_imports(self):
        path = ROOT / "exact_glyph_identity.py"
        tree = ast.parse(path.read_text(encoding="utf-8"))
        imported: set[str] = set()
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                imported.update(alias.name for alias in node.names)
            elif isinstance(node, ast.ImportFrom) and node.module:
                imported.add(node.module)
        forbidden = {
            "sqlite3",
            "concise_expected_resolver",
            "pronunciation_rule_engine",
            "check_pronunciation_candidates",
        }
        self.assertFalse(imported & forbidden)

    def test_identity_helper_is_present_in_both_actual_source_audit_rosters(self):
        def assigned_source_list(path: Path) -> list[str]:
            tree = ast.parse(path.read_text(encoding="utf-8"))
            assignment = next(
                node
                for node in tree.body
                if isinstance(node, (ast.Assign, ast.AnnAssign))
                and any(
                    isinstance(target, ast.Name) and target.id == "ACTUAL_DECODER_SOURCE_FILES"
                    for target in (node.targets if isinstance(node, ast.Assign) else [node.target])
                )
            )
            return ast.literal_eval(assignment.value)

        exported = assigned_source_list(ROOT / "export_zhuyin_readings.py")
        pipeline = assigned_source_list(ROOT / "standalone_proofread.py")
        self.assertEqual(exported, pipeline)
        self.assertIn("exact_glyph_identity.py", exported)


if __name__ == "__main__":
    unittest.main()
