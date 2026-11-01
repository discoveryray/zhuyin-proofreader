"""PA-06 pure evidence reuse with genuine, isolated nonwindow sources."""
from contextlib import contextmanager
from datetime import datetime, timezone
from dataclasses import asdict
import copy
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import types
from unittest.mock import patch

import pytest
from openpyxl import load_workbook
import actual_review as ar
import glyph_learning_impact as impact
import global_exact_glyph_library as lib
import standalone_proofread as sp
import review_gui as gui
from tests.test_glyph_learning_impact import fixture_project
from tests.test_global_promotion_v580 import intent

BASELINE = "4482f90553ba3ebd85dbb12adf1d6bfa4fc58a9e"


def evidence(name, data):
    if root := os.environ.get("PA06_EVIDENCE_DIR"):
        (Path(root) / (name + '.json')).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')




@contextmanager
def counts(ar_module, *, baseline=False):
    targets = [(ar_module, 'actual_workbook_global_exact_dependencies' if baseline else '_global_exact_dependency_reader'),
               (sp, 'workbook_metadata'), (ar_module.fitz, 'open'),
               (ar_module.TrueTypeGlyphInspector, 'global_exact_identity'), (lib, '_validate_connection')]
    values = dict(roster=0, metadata=0, pdf=0, font=0, validator=0)
    from contextlib import ExitStack
    with ExitStack() as stack:
        for key, (module, name) in zip(values, targets):
            original = getattr(module, name)
            def observed(*args, _key=key, _original=original, **kwargs):
                values[_key] += 1
                return _original(*args, **kwargs)
            stack.enter_context(patch.object(module, name, observed))
        yield values


def repo_for(tmp_path):
    repo = lib.GlobalExactGlyphRepository(tmp_path / 'global.sqlite3')
    repo.initialize()
    return repo


def compare_project(expected, actual, before, after):
    assert set(expected) == set(actual)
    for project in (expected, actual):
        timestamp = project['registered_at']
        assert timestamp.endswith('Z')
        when = datetime.fromisoformat(timestamp.replace('Z', '+00:00'))
        assert before.replace(microsecond=0) <= when <= after
    assert {k: v for k, v in expected.items() if k != 'registered_at'} == {
        k: v for k, v in actual.items() if k != 'registered_at'}


def test_real_parser_counts_and_next_operation(tmp_path):
    repo = repo_for(tmp_path)
    root = tmp_path / 'project'
    manifest, identities = fixture_project(root, source_snapshot=repo.load_snapshot())
    before = datetime.now(timezone.utc)
    with counts(ar) as observed:
        with impact._dependency_parse_operation(root) as owner:
            first = impact.recheck_project_dependencies(root, repository=repo, manifest=manifest)
            assert first['status'] == 'CURRENT'
            assert [row['dependencies'][0]['glyph_sha256'] for row in first['pdfs']] == [i.glyph_sha256 for i in identities]
            expected = copy.deepcopy(first)
            first['pdfs'][0]['dependencies'].clear()
            project = impact.register_usage(root, repository=repo, registry_path=tmp_path/'usage.json')
            assert project['pdfs'] == expected['pdfs']
            assert len(owner.parsed) == 2
            for dependencies, pdf_sha, components_raw in owner.parsed.values():
                assert isinstance(dependencies, tuple) and all(type(i) is lib.GlobalExactGlyphIdentity for i in dependencies)
                assert isinstance(pdf_sha, str) and len(pdf_sha) == 64
                assert isinstance(components_raw, str) and isinstance(json.loads(components_raw), dict)
            with pytest.raises(TypeError):
                owner.parsed['foreign'] = ()
        assert owner.closed and not owner.parsed and owner.key is None
    after = datetime.now(timezone.utc)
    compare_project(project, json.loads((tmp_path/'usage.json').read_text(encoding='utf-8'))['projects'][0], before, after)
    assert observed == dict(roster=2, metadata=2, pdf=2, font=2, validator=2)
    with counts(ar) as next_counts, impact._dependency_parse_operation(root):
        assert impact.recheck_project_dependencies(root, repository=repo) == expected
        impact.register_usage(root, repository=repo, registry_path=tmp_path/'next.json')
    assert next_counts == observed



def seal(root, manifest):
    manifest = sp.seal_manifest(manifest)
    (root/'校對工作階段.json').write_text(json.dumps(manifest, ensure_ascii=False), encoding='utf-8')
    return manifest


@pytest.mark.parametrize('change', ['format', 'new_seal', 'bad_seal', 'pdf_same_stat', 'xlsx_same_stat', 'missing', 'duplicate', 'replacement', 'relocate'])
def test_second_source_boundary_retains_original_full_decision(tmp_path, change):
    repo=repo_for(tmp_path);root=tmp_path/'project';manifest,_=fixture_project(root,source_snapshot=repo.load_snapshot(),glyphs=('one',))
    with impact._dependency_parse_operation(root) as owner:
        impact.recheck_project_dependencies(root,repository=repo)
        info=manifest['pdfs'][0];pdf=Path(info['pdf']);workbook=root/info['actual_workbook']
        if change=='format':
            (root/'校對工作階段.json').write_text(json.dumps(manifest, indent=2),encoding='utf-8')
        elif change=='new_seal': manifest['session_id']+='-new';seal(root,manifest)
        elif change=='bad_seal': manifest['session_id']+='-bad';(root/'校對工作階段.json').write_text(json.dumps(manifest),encoding='utf-8')
        elif change in {'pdf_same_stat','xlsx_same_stat'}:
            path=pdf if change=='pdf_same_stat' else workbook;stat=path.stat();data=bytearray(path.read_bytes());data[len(data)//2]^=1;path.write_bytes(data);os.utime(path,ns=(stat.st_atime_ns,stat.st_mtime_ns));assert path.stat().st_size==stat.st_size
        elif change=='missing': workbook.unlink()
        elif change=='duplicate': (workbook.parent/'duplicate.xlsx').write_bytes(workbook.read_bytes())
        elif change=='replacement':
            stat=workbook.stat();data=workbook.read_bytes();workbook.unlink();workbook.write_bytes(data);os.utime(workbook,ns=(stat.st_atime_ns,stat.st_mtime_ns))
        elif change=='relocate':
            moved=workbook.with_name('renamed.xlsx');workbook.rename(moved)
        if change in {'bad_seal', 'pdf_same_stat', 'xlsx_same_stat', 'missing', 'duplicate', 'relocate'}:
            with pytest.raises(ValueError):
                impact.recheck_project_dependencies(root, repository=repo)
            error = True
        else:
            actual = impact.recheck_project_dependencies(root, repository=repo)
            assert actual['status'] == 'CURRENT'
            assert actual['session_id'] == manifest['session_id']
            error = None
        if change in {'format','new_seal','relocate'} and error is None: assert owner.closed and not owner.parsed


def test_captured_pdf_and_xlsx_are_same_bytes_under_transient_live_change(tmp_path,monkeypatch):
    repo=repo_for(tmp_path);root=tmp_path/'project';manifest,_=fixture_project(root,source_snapshot=repo.load_snapshot(),glyphs=('capture',))
    workbook=root/manifest['pdfs'][0]['actual_workbook'];pdf=Path(manifest['pdfs'][0]['pdf']);original=ar._global_exact_dependency_reader
    seen=[]
    def changed(captured, open_pdf, *, source_label):
        pdf_data=pdf.read_bytes();xlsx_data=workbook.read_bytes()
        pdf.write_bytes(b'not a PDF');workbook.write_bytes(b'not an XLSX')
        try:
            assert Path(captured)!=workbook and Path(captured).read_bytes()==xlsx_data
            result=original(captured,open_pdf,source_label=source_label);seen.append(Path(captured));return result
        finally: pdf.write_bytes(pdf_data);workbook.write_bytes(xlsx_data)
    monkeypatch.setattr(ar,'_global_exact_dependency_reader',changed)
    with impact._dependency_parse_operation(root):
        assert impact.recheck_project_dependencies(root,repository=repo)['status']=='CURRENT'
        assert impact.recheck_project_dependencies(root,repository=repo)['status']=='CURRENT'
    assert len(seen)==1 and not seen[0].exists()


@pytest.mark.parametrize('related',[True,False])
def test_fresh_global_legal_writer_changes_second_status(tmp_path,related):
    repo=repo_for(tmp_path);root=tmp_path/'project';manifest,identities=fixture_project(root,source_snapshot=repo.load_snapshot(),glyphs=('global',))
    identity=identities[0] if related else lib.canonical_global_exact_identity(lib.TTF_GLYF_SHA256,'','a'*64,lib.GLOBAL_ELIGIBLE_SIMPLE_GLYF_V1)
    with impact._dependency_parse_operation(root),counts(ar) as observed:
        assert impact.recheck_project_dependencies(root,repository=repo)['status']=='CURRENT'
        # Legal transactional writer, deterministic between two reads; no sleeps or main-file hash.
        lib.deliver_global_direct_evidence(repo,[intent(1,'ㄅ',identity=asdict(identity)),intent(2,'ㄆ',identity=asdict(identity))])
        project=impact.register_usage(root,repository=repo,registry_path=tmp_path/'usage.json')
    assert project['pdfs'][0]['status']==('STALE' if related else 'CURRENT')
    assert observed['roster']==observed['metadata']==observed['pdf']==1
    # Writer validates too; the two read-only rechecks still each run the complete validator.
    assert observed['validator']>=2


def test_owner_foreign_closed_wrong_type_and_root_are_rejected(tmp_path):
    root=tmp_path/'root';root.mkdir();other=tmp_path/'other';other.mkdir()
    for invalid in (object(),impact._DependencyParseOperation(root)):
        token=impact._dependency_operation.set(invalid)
        try:
            with pytest.raises(ValueError,match='owner'): impact._active_dependency_operation(root)
        finally:impact._dependency_operation.reset(token)
    with impact._dependency_parse_operation(root) as owner:
        with pytest.raises(ValueError,match='owner'):impact._active_dependency_operation(other)
        with pytest.raises(ValueError,match='跨操作'):
            with impact._dependency_parse_operation(root):pass
    token=impact._dependency_operation.set(owner)
    try:
        with pytest.raises(ValueError,match='owner'):impact._active_dependency_operation(root)
    finally:impact._dependency_operation.reset(token)


def test_parse_exception_closes_document_workbook_snapshot_and_owner(tmp_path,monkeypatch):
    repo=repo_for(tmp_path);root=tmp_path/'project';manifest,_=fixture_project(root,source_snapshot=repo.load_snapshot(),glyphs=('close',))
    reader=ar.load_workbook;opened=[]
    pdf_reader=ar.fitz.open;documents=[]
    def track_pdf(*args,**kwargs):
        document=pdf_reader(*args,**kwargs);documents.append(document);return document
    monkeypatch.setattr(ar.fitz,"open",track_pdf)
    def track(path,**kwargs):
        result=reader(path,**kwargs);original=result.close
        def close():opened.append(str(path));return original()
        result.close=close;return result
    monkeypatch.setattr(ar,'load_workbook',track)
    monkeypatch.setattr(ar.TrueTypeGlyphInspector,'global_exact_identity',lambda *a: (_ for _ in ()).throw(RuntimeError('parse failure')))
    with pytest.raises(RuntimeError,match='parse failure'):
        with impact._dependency_parse_operation(root) as owner:impact.recheck_project_dependencies(root,repository=repo)
    assert owner.closed and not owner.parsed and impact._dependency_operation.get() is None
    assert documents and all(document.is_closed for document in documents)
    assert opened and all(not Path(path).exists() for path in opened)


def test_old_api_preserves_source_label_none_and_empty_tuple(tmp_path):
    from openpyxl import Workbook
    path=tmp_path/'actual.xlsx';pdf=tmp_path/'empty.pdf'
    doc=ar.fitz.open();doc.new_page();doc.save(pdf);doc.close()
    book=Workbook();book.active.title='實際注音';book.save(path);book.close()
    assert ar.actual_workbook_global_exact_dependencies(path,pdf) is None
    book=load_workbook(path);sheet=book.active;sheet.append(['字形架構','font_xref','注音元件ID','TTF字形SHA256','CFF樣式群組','CFF整字字形SHA256']);book.save(path);book.close()
    assert ar.actual_workbook_global_exact_dependencies(path,pdf)==()
    path.write_bytes(b'corrupt')
    for invoke in (lambda:ar.actual_workbook_global_exact_dependencies(path,pdf),lambda:ar._global_exact_dependency_reader(path,lambda:ar.fitz.open(pdf),source_label=path)):
        with pytest.raises(ValueError,match=str(path).replace('\\','\\\\')) as caught:invoke()
        assert caught.value.__cause__ is not None

@pytest.mark.parametrize('stage',['first_stale','first_unprovable','second_unprovable','registration_failure'])
def test_loader_branches_close_owner_and_keep_observation_errors(tmp_path,monkeypatch,stage):
    from types import SimpleNamespace
    import pdf_portability as portability
    root=tmp_path/'loader';root.mkdir()
    snapshot=SimpleNamespace(manifest={},db={},ledger=[])
    service=SimpleNamespace(load_resume_snapshot=lambda:snapshot,evidence_anchor=object())
    statuses=iter(['STALE' if stage=='first_stale' else 'UNPROVABLE' if stage=='first_unprovable' else 'CURRENT'])
    owners=[];events=[]
    def checked(*args,**kwargs):
        owner=impact._dependency_operation.get();owners.append(owner)
        assert owner is not None and not owner.closed
        return {'status':next(statuses)}
    def refreshed(*args):
        assert impact._dependency_operation.get() is None
        assert owners[0].closed and not owners[0].parsed
        events.append('refresh')
    def registered(*args):
        if stage=='first_stale':assert impact._dependency_operation.get() is None
        else:assert impact._dependency_operation.get() is owners[0]
        events.append('register')
        if stage=='second_unprovable':raise ValueError('依賴無法核實')
        if stage=='registration_failure':raise OSError('registration failed')
    monkeypatch.setattr(impact,'recheck_project_dependencies',checked)
    monkeypatch.setattr(impact,'refresh_stale_project_dependencies',refreshed)
    monkeypatch.setattr(impact,'register_usage',registered)
    monkeypatch.setattr(gui,'ReviewSaveService',lambda *a,**k:service)
    monkeypatch.setattr(portability,'validate_conflict_state',lambda *a:[])
    monkeypatch.setattr(portability,'actual_excel_conflict_state',lambda *a:[])
    monkeypatch.setattr(portability,'expected_conflict_review_view',lambda *a:[])
    monkeypatch.setattr(gui,'prepare_review_queue',lambda *a,**k:{'items':[]})
    if stage=='first_unprovable':
        with pytest.raises(ValueError,match='無法核實'):gui.load_review_project(root,lambda _:None,service=service)
        assert events==[]
    else:
        result=gui.load_review_project(root,lambda _:None,service=service)
        assert events==(['refresh','register'] if stage=='first_stale' else ['register'])
        assert result['learning_registration_error']==('依賴無法核實' if stage=='second_unprovable' else 'registration failed' if stage=='registration_failure' else None)
    assert owners[0].closed and not owners[0].parsed and impact._dependency_operation.get() is None


@pytest.mark.parametrize('invalid',['duplicate_metadata','unknown_contract','wrong_pdf','missing_columns','duplicate_headers'])
def test_captured_parser_and_metadata_rejections_preserve_contract(tmp_path,invalid):
    repo=repo_for(tmp_path);root=tmp_path/'project';manifest,_=fixture_project(root,source_snapshot=repo.load_snapshot(),glyphs=('invalid',))
    info=manifest['pdfs'][0];path=root/info['actual_workbook'];book=load_workbook(path)
    if invalid=='duplicate_metadata':book['v5.2中繼資料'].append(['pdf_sha256',info['pdf_sha256']])
    elif invalid=='unknown_contract':book['v5.2中繼資料']['B2']='{}'
    elif invalid=='wrong_pdf':book['v5.2中繼資料']['B1']='f'*64
    elif invalid=='missing_columns':book['實際注音'].delete_cols(1)
    elif invalid=='duplicate_headers':book['實際注音']['B1']='字形架構'
    book.save(path);book.close();info['actual_workbook_sha256']=hashlib.sha256(path.read_bytes()).hexdigest();seal(root,manifest)
    with impact._dependency_parse_operation(root) as owner:
        if invalid == 'duplicate_headers':
            with pytest.raises(ValueError, match='欄位重複'):
                impact.recheck_project_dependencies(root, repository=repo)
        else:
            actual = impact.recheck_project_dependencies(root, repository=repo)
            assert actual['status'] == 'UNPROVABLE'
    assert owner.closed and not owner.parsed

@pytest.mark.parametrize('second',['absent','corrupt','unprovable'])
def test_second_real_global_or_metadata_state_is_not_cached(tmp_path,second):
    repo=repo_for(tmp_path);root=tmp_path/'project';manifest,_=fixture_project(root,source_snapshot=repo.load_snapshot(),glyphs=('second',))
    with impact._dependency_parse_operation(root):
        assert impact.recheck_project_dependencies(root,repository=repo)['status']=='CURRENT'
        if second=='absent':
            repo.path.unlink()
            assert impact.register_usage(root,repository=repo,registry_path=tmp_path/'usage.json')['pdfs'][0]['status']=='CURRENT'
        elif second=='corrupt':
            repo.path.write_bytes(b'corrupt sqlite')
            with pytest.raises(lib.GlobalLibraryError):impact.register_usage(root,repository=repo,registry_path=tmp_path/'usage.json')
        else:
            info=manifest['pdfs'][0];path=root/info['actual_workbook'];book=load_workbook(path);book['v5.2中繼資料']['B2']='{}';book.save(path);book.close()
            info['actual_workbook_sha256']=hashlib.sha256(path.read_bytes()).hexdigest();seal(root,manifest)
            with pytest.raises(ValueError,match='無法核實'):impact.register_usage(root,repository=repo,registry_path=tmp_path/'usage.json')
            assert not (tmp_path/'usage.json').exists()


def test_usage_publish_cleanup_does_not_add_postpublication_source_guard(tmp_path,monkeypatch):
    repo=repo_for(tmp_path);root=tmp_path/'project';manifest,_=fixture_project(root,source_snapshot=repo.load_snapshot(),glyphs=('published',))
    write=impact._atomic_write
    def published(path,raw):
        write(path,raw)
        source=Path(manifest['pdfs'][0]['pdf']);source.write_bytes(source.read_bytes()+b'changed after publication')
    monkeypatch.setattr(impact,'_atomic_write',published)
    with impact._dependency_parse_operation(root) as owner:
        impact.recheck_project_dependencies(root,repository=repo)
        result=impact.register_usage(root,repository=repo,registry_path=tmp_path/'usage.json')
    assert owner.closed and not owner.parsed
    assert json.loads((tmp_path/'usage.json').read_text(encoding='utf-8'))['projects'][0]==result


def test_legal_new_workbook_seal_misses_without_extra_rejection(tmp_path):
    repo=repo_for(tmp_path);root=tmp_path/'project';manifest,_=fixture_project(root,source_snapshot=repo.load_snapshot(),glyphs=('new-seal',))
    with impact._dependency_parse_operation(root) as owner,counts(ar) as observed:
        impact.recheck_project_dependencies(root,repository=repo)
        info=manifest['pdfs'][0];path=root/info['actual_workbook'];book=load_workbook(path);book['v5.2中繼資料'].append(['audit-only','new合法metadata']);book.save(path);book.close()
        info['actual_workbook_sha256']=hashlib.sha256(path.read_bytes()).hexdigest();seal(root,manifest)
        actual=impact.register_usage(root,repository=repo,registry_path=tmp_path/'usage.json')
        assert actual['pdfs'][0]['status']=='CURRENT'
        assert owner.closed and not owner.parsed
    assert observed['roster']==observed['metadata']==2 and observed['validator']==2
    assert actual['pdfs'][0]['actual_workbook_sha256'] == info['actual_workbook_sha256']
