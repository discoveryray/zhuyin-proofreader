"""Real tiny sealed PDF/XLSX import/save transactions; no Tk or producer shim."""
import copy
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

import pdf_portability as p
import standalone_proofread as sp
from review_gui import expected_conflict_summary, prepare_review_queue, review_lane
from review_save_service import ReviewSaveService
from tests.test_pdf_portability import pdf, project
from tests.test_expected_business_import import state


@pytest.fixture
def recheck_batch(tmp_path, monkeypatch, request):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "user"))
    monkeypatch.setattr(sp, "USER_MEMORY", tmp_path / "memory.json")
    monkeypatch.setattr(sp, "REUSABLE_EXPECTED_RULES", tmp_path / "rules.json")
    source, target = tmp_path / "source", tmp_path / "target"
    a, b = tmp_path / "a.pdf", tmp_path / "b.pdf"
    pdf(a, title="source recheck")
    pdf(b, title="target recheck")
    boxes = [(10 + n * 30, 20, 25 + n * 30, 45) for n in range(4)]
    sm, sd = project(source, a, session="source", boxes=boxes)
    tm, td = project(target, b, session="target", boxes=boxes)
    assets = sp.validate_asset_manifest(Path(sp.__file__).parent)
    assert assets["ok"]
    fp = sp.compute_expected_asset_fingerprint(Path(sp.__file__).parent, assets,
        resolver_version=sp.EXPECTED_RESOLVER_VERSION, source_files=sp.EXPECTED_RESOLVER_SOURCE_FILES)
    for root, manifest in ((source, sm), (target, tm)):
        manifest["expected_asset_fingerprint"] = fp["fingerprint"]
        manifest["expected_asset_fingerprint_components"] = fp["components"]
        for entry in manifest["records"]:
            entry["source_record"].update(所在行="角色原文", 局部詞境="角色")
        candidate = Path(manifest["pdfs"][0]["candidate_workbook"])
        wb = Workbook()
        wb.active.append(["synthetic candidate"])
        wb.save(candidate)
        wb.close()
        manifest["pdfs"][0]["candidate_workbook_sha256"] = sp.sha256_file(candidate)
    for n in (1, 2, 3):
        entry = sm["records"][n]
        entry.update(expected_set=["ㄐㄩㄝˋ"], expected_evidence="independent dictionary",
                     expected_status="RESOLVED", context_evidence="full original context")
        entry.update(sp.refresh_derived_state(entry, preserve_confirmed=False))
    tm["records"][1].update({key: copy.deepcopy(sm["records"][1][key]) for key in
                            ("expected_set", "expected_evidence", "expected_status", "context_evidence")})
    fault = getattr(request, "param", "actual")
    key, value = {"actual": ("actual", "ㄐㄩㄝˋ"),
                  "actual_evidence": ("actual_evidence", "new glyph evidence"),
                  "expected_set": ("expected_set", ["ㄐㄩㄝˊ"]),
                  "expected_evidence": ("expected_evidence", "new independent dictionary")}[fault]
    tm["records"][1][key] = value
    tm["records"][1].update(sp.refresh_derived_state(tm["records"][1], preserve_confirmed=False))
    rid = tm["records"][2]["review_id"]
    event = {"action": "補建expected證據", "expected_set": "ㄐㄩㄝˋ",
             "expected_evidence": "independent dictionary", "context_evidence": "full original context"}
    td["events"][rid] = {**event, "portability_source": {"session_id": "previous validated import"}}
    local = sp.build_manual_expected_event(tm["records"][3], operation="ENTER_EXPECTED",
                                          expected_set="ㄐㄩㄝˊ", rationale="completed local decision")
    td["events"][tm["records"][3]["review_id"]] = local
    for root, manifest, db in ((source, sm, sd), (target, tm, td)):
        sp.seal_manifest(manifest)
        sp.json_save(root / "校對工作階段.json", manifest)
        sp.json_save(root / "人工判定資料庫.json", db)
    exported = sp.export_pending_for_gpt(source)
    filled = tmp_path / "input.xlsx"
    wb = load_workbook(exported)
    sheet = wb["待判定候選"]
    headers = [cell.value for cell in sheet[1]]
    for n in range(4):
        values = {"action": "補建expected證據" if n == 0 else "確認現版差異" if n == 1 else "確認非校對範圍",
                  "proposed_expected_set": "ㄐㄩㄝˋ", "proposed_expected_evidence": "independent dictionary",
                  "proposed_context_evidence": "full original context",
                  "exclusion_reason": "source structural exclusion", "exclusion_evidence": "source page geometry"}
        if n == 1:
            values.update({gate: "Y" for gate in sp.CONFIRMATION_GATES})
        for key, value in values.items():
            sheet.cell(n + 2, headers.index(key) + 1, value)
    wb.save(filled)
    wb.close()
    return source, target, sm, tm, td, filled


@pytest.mark.parametrize("recheck_batch", ["actual", "actual_evidence", "expected_set", "expected_evidence"], indirect=True)
def test_mixed_valid_and_stale_dependencies_preserve_truth(recheck_batch):
    source, target, sm, tm, td, filled = recheck_batch
    source_before, target_before = state(source), sp.materialize_ledger(tm, td)
    result = sp.import_gpt_decisions(target, filled)
    assert result.status["imported"] == 1
    assert result.status["dependency_rechecks"] == 2
    assert result.status["preserved_local_decisions"] == 1
    db = sp.json_load_strict(target / "人工判定資料庫.json")
    assert sp.materialize_ledger(tm, db)[1:] == target_before[1:]
    assert all(db["events"][key] == value for key, value in td["events"].items())
    receipt = sp.json_load_strict(target / p.CONFLICT_FILE)
    assert len(receipt["conflicts"]) == 3
    record = receipt["conflicts"][1]
    assert record["source_event"]["action"] == "確認非校對範圍"
    assert record["dependency_recheck"]["disposition"] == "RECHECK_LOCAL_REQUIRED"
    # Same source/current semantics do not excuse a different raw baseline.
    assert all(record["dependency_recheck"]["source_snapshot"][key] ==
               record["dependency_recheck"]["target_current_snapshot"][key] for key in p._RECHECK_KEYS)
    assert all(d["comparison"] == "baseline" for d in record["dependency_recheck"]["differences"])
    view = p.expected_conflict_review_view(target, tm, db, sp.materialize_ledger(tm, db))
    assert review_lane(view[2]) == "expected"
    text = expected_conflict_summary(view[2])
    assert "確認非校對範圍" in text and "目標原始基線" in text and "expected_evidence" in text
    queue = prepare_review_queue(tm, view, [], 0, set(), set())
    assert view[2]["review_id"] in {e["review_id"] for e in queue["records"]}
    assert not view[3].get("expected_business_conflict_pending")
    before_receipt = (target / p.CONFLICT_FILE).read_bytes()
    retry = sp.import_gpt_decisions(target, filled)
    assert retry.status["duplicates"] == 4 and retry.status["new_conflicts"] == 0
    assert (target / p.CONFLICT_FILE).read_bytes() == before_receipt
    assert state(source) == source_before


@pytest.mark.parametrize("choice", ["exclude", "expected"])
def test_local_recheck_save_reload_retry_does_not_revive(recheck_batch, choice):
    _, target, _, tm, _, filled = recheck_batch
    sp.import_gpt_decisions(target, filled)
    service = ReviewSaveService(target)
    snapshot = service.load_resume_snapshot()
    view = p.expected_conflict_review_view(target, tm, snapshot.db, snapshot.ledger)
    row = view[2]
    event = ({"action": "確認非校對範圍", "exclusion_reason": "fresh local scope review",
              "exclusion_evidence": "current page geometry", "source": "人工 GUI 排除證據"}
             if choice == "exclude" else sp.build_manual_expected_event(
                 row, operation="ENTER_EXPECTED", expected_set="ㄐㄩㄝˊ", rationale="fresh local dictionary"))
    saved = service.save_event(row["review_id"], event)
    assert row["review_id"] not in p.validate_conflict_state(target, tm, saved.db)
    receipt_before = (target / p.CONFLICT_FILE).read_bytes()
    db_before = sp.json_load_strict(target / "人工判定資料庫.json")
    assert service.load_resume_snapshot().db == db_before
    sp.import_gpt_decisions(target, filled)
    assert sp.json_load_strict(target / "人工判定資料庫.json") == db_before
    assert (target / p.CONFLICT_FILE).read_bytes() == receipt_before
    assert row["review_id"] not in p.validate_conflict_state(target, tm, db_before)


def test_invalid_stale_source_still_aborts_before_any_write(recheck_batch):
    _, target, _, _, _, filled = recheck_batch
    wb = load_workbook(filled)
    sheet = wb["待判定候選"]
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(3, headers.index(sp.CONFIRMATION_GATES[0]) + 1, "N")
    wb.save(filled)
    wb.close()
    before = state(target)
    with pytest.raises(ValueError, match="TEXTBOOK_ERROR_CONFIRMED 六閘門未全數通過"):
        sp.import_gpt_decisions(target, filled)
    assert state(target) == before


def test_unfinished_local_event_is_still_pending_recheck(recheck_batch):
    _, target, _, tm, td, filled = recheck_batch
    rid = tm["records"][3]["review_id"]
    td["events"][rid] = sp.build_manual_expected_event(
        tm["records"][3], operation="ENTER_EXPECTED", expected_set="ㄐㄩㄝˋ",
        rationale="valid but not completed local difference")
    sp.json_save(target / "人工判定資料庫.json", td)
    before = sp.materialize_ledger(tm, td)
    assert before[3]["state"] == "DIFFERENCE_PENDING_CONFIRMATION"
    result = sp.import_gpt_decisions(target, filled)
    assert result.status["dependency_rechecks"] == 3
    assert result.status["preserved_local_decisions"] == 0
    db = sp.json_load_strict(target / "人工判定資料庫.json")
    assert db["events"][rid] == td["events"][rid]
    assert rid in p.validate_conflict_state(target, tm, db)


def test_recheck_receipt_write_failure_rolls_back(recheck_batch, monkeypatch):
    _, target, _, _, _, filled = recheck_batch
    before = state(target)
    original = sp.json_save
    def fail(path, *args, **kwargs):
        if Path(path) == target / p.CONFLICT_FILE:
            raise OSError("injected receipt failure")
        return original(path, *args, **kwargs)
    monkeypatch.setattr(sp, "json_save", fail)
    with pytest.raises(OSError, match="injected receipt"):
        sp.import_gpt_decisions(target, filled)
    assert state(target) == before


@pytest.mark.parametrize("fault", ["differences", "source_action", "snapshot_identity", "future_contract"])
def test_recheck_receipt_structural_corruption_rejected(recheck_batch, fault):
    _, target, _, tm, _, filled = recheck_batch
    sp.import_gpt_decisions(target, filled)
    receipt = sp.json_load_strict(target / p.CONFLICT_FILE)
    entry = receipt["conflicts"][0]
    if fault == "differences":
        entry["dependency_recheck"]["differences"] = []
    elif fault == "source_action":
        entry["source_event"]["action"] = "補建expected證據"
    elif fault == "snapshot_identity":
        entry["dependency_recheck"]["source_snapshot"]["occurrence_id"] = "wrong"
    else:
        entry["contract"] = "review-dependency-recheck/future"
    receipt = p._sealed_expected_receipt(receipt["conflicts"], receipt["supplements"])
    with pytest.raises(ValueError):
        p._validated_expected_receipt(receipt, tm)
