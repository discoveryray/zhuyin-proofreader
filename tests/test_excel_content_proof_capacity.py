"""Small sealed fixtures exercise bounded proof bytes, never Tk or live data.

Boundary cases lower only the shared byte limit, avoiding 256 MiB allocations.
The independent production-limit assertion binds the adopted capacity.
"""

import base64
from datetime import datetime
import hashlib
import json
import zlib

from openpyxl import Workbook
import pytest

import actual_review as ar
import pdf_portability as portable
import standalone_proofread as sp
from tests.test_pdf_portability import pdf, project


class FixedDatetime(datetime):
    @classmethod
    def now(cls):
        return cls(2026, 10, 10, 12, 0, 0)


@pytest.fixture
def source(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    monkeypatch.setattr(portable, "datetime", FixedDatetime)
    source_pdf = tmp_path / "source.pdf"
    pdf(source_pdf, title="capacity fixture")
    root = tmp_path / "source"
    manifest, db = project(root, source_pdf, session="capacity-source")
    return root, manifest, db


def workbook_for(kind):
    workbook = Workbook()
    workbook.active.title = "actual待判定" if kind == "actual" else "待判定候選"
    headers = (["sample_a_occurrence_id", "sample_b_occurrence_id",
                "sample_a_image", "sample_b_image"] if kind == "actual" else
               ["occurrence_id", "review_id"])
    workbook.active.append(headers)
    metadata = workbook.create_sheet("匯入中繼資料")
    metadata.append(["項目", "內容"])
    metadata.append(["session_id", "capacity-source"])
    metadata.append(["actual_review_schema_version", "1.1"])
    return workbook


def proof_bytes(workbook):
    rows = list(workbook[portable.EXCEL_PROOF_SHEET].values)
    packed = base64.b64decode("".join(row[1] for row in rows[3:]), validate=True)
    return zlib.decompress(packed)


def replace_proof(workbook, raw, *, packed=None, digest=None):
    del workbook[portable.EXCEL_PROOF_SHEET]
    sheet = workbook.create_sheet(portable.EXCEL_PROOF_SHEET)
    sheet.append(["version", 1])
    sheet.append(["sha256", digest or hashlib.sha256(raw).hexdigest()])
    encoded = base64.b64encode(zlib.compress(raw) if packed is None else packed).decode("ascii")
    chunks = [encoded[i:i + 30000] for i in range(0, len(encoded), 30000)]
    sheet.append(["chunk_count", len(chunks)])
    for i, chunk in enumerate(chunks, 1):
        sheet.append([i, chunk])
    sheet.sheet_state = "veryHidden"


def read_proof(workbook, route, source, tmp_path):
    root, manifest, _ = source
    if route == "profile":
        return portable.validate_actual_profile_content_proof(workbook, source_manifest=manifest)
    path = tmp_path / f"{route}.xlsx"
    workbook.save(path)
    return portable.load_excel_content_proof(
        path, root, manifest, kind=route, workbook_session="capacity-source")


@pytest.fixture(params=["actual", "expected", "profile"])
def carried(request, source):
    route = request.param
    kind = "actual" if route == "profile" else route
    workbook = workbook_for(kind)
    root, manifest, db = source
    portable.write_excel_content_proof(workbook, root, manifest, db, kind=kind)
    yield workbook, route, kind, proof_bytes(workbook)
    workbook.close()


def test_adopted_capacity_is_256_mib():
    assert portable.EXCEL_CONTENT_PROOF_MAX_BYTES == 268435456


@pytest.mark.parametrize("remaining", [0, 1])
def test_all_readers_accept_shared_inclusive_boundary(carried, source, tmp_path, monkeypatch, remaining):
    workbook, route, _, raw = carried
    monkeypatch.setattr(portable, "EXCEL_CONTENT_PROOF_MAX_BYTES", len(raw) + remaining)
    result = read_proof(workbook, route, source, tmp_path)
    if route != "profile":
        assert result[0] == source[1]
        assert result[6]["kind"] == route


def test_all_readers_reject_one_byte_over_with_capacity_reason(carried, source, tmp_path, monkeypatch):
    workbook, route, _, raw = carried
    monkeypatch.setattr(portable, "EXCEL_CONTENT_PROOF_MAX_BYTES", len(raw) - 1)
    with pytest.raises(ValueError, match="解壓容量超限") as error:
        read_proof(workbook, route, source, tmp_path)
    assert "無法解碼" not in str(error.value)


@pytest.mark.parametrize("damage", ["truncated", "trailing", "second-stream", "corrupt", "json", "sha", "noncanonical"])
def test_all_readers_preserve_stream_json_sha_guards(carried, source, tmp_path, damage):
    workbook, route, _, raw = carried
    packed = zlib.compress(raw)
    digest = None
    if damage == "truncated":
        packed = packed[:-1]
    elif damage == "trailing":
        packed += b"trailing"
    elif damage == "second-stream":
        packed += zlib.compress(b"{}")
    elif damage == "corrupt":
        packed = b"invalid zlib"
    elif damage == "json":
        raw = b"{invalid json"
        packed = zlib.compress(raw)
    elif damage == "sha":
        digest = "0" * 64
    else:
        raw = json.dumps(json.loads(raw), ensure_ascii=False, indent=2).encode("utf-8")
        packed = zlib.compress(raw)
    replace_proof(workbook, raw, packed=packed, digest=digest)
    expected = ("串流不完整或含尾隨" if damage in {"truncated", "trailing", "second-stream"} else
                "壓縮資料損壞" if damage == "corrupt" else
                "JSON 格式損壞" if damage == "json" else "不符")
    with pytest.raises(ValueError, match=expected) as error:
        read_proof(workbook, route, source, tmp_path)
    assert "容量超限" not in str(error.value)


@pytest.mark.parametrize("kind", ["actual", "expected"])
def test_writer_and_readers_share_same_exact_boundary(source, tmp_path, monkeypatch, kind):
    root, manifest, db = source
    workbook = workbook_for(kind)
    portable.write_excel_content_proof(workbook, root, manifest, db, kind=kind)
    raw = proof_bytes(workbook)
    del workbook[portable.EXCEL_PROOF_SHEET]
    monkeypatch.setattr(portable, "EXCEL_CONTENT_PROOF_MAX_BYTES", len(raw))
    portable.write_excel_content_proof(workbook, root, manifest, db, kind=kind)
    assert proof_bytes(workbook) == raw
    read_proof(workbook, kind, source, tmp_path)
    if kind == "actual":
        read_proof(workbook, "profile", source, tmp_path)
    del workbook[portable.EXCEL_PROOF_SHEET]
    monkeypatch.setattr(portable, "EXCEL_CONTENT_PROOF_MAX_BYTES", len(raw) - 1)
    before = [(sheet.title, list(sheet.values)) for sheet in workbook]
    with pytest.raises(ValueError, match="解壓容量超限"):
        portable.write_excel_content_proof(workbook, root, manifest, db, kind=kind)
    assert [(sheet.title, list(sheet.values)) for sheet in workbook] == before
    workbook.close()


@pytest.mark.parametrize("kind", ["actual", "expected"])
@pytest.mark.parametrize("existing", [False, True])
def test_export_over_limit_never_publishes_success_workbook(source, monkeypatch, kind, existing):
    root, manifest, db = source
    expected = root / "待判定候選_給GPT.xlsx"
    actual = root / "actual待判定_GPT包" / "actual待判定_給GPT.xlsx"
    archive = root / "actual待判定_GPT包.zip"
    paths = [expected] if kind == "expected" else [actual, archive]
    if existing:
        for path in paths:
            path.parent.mkdir(exist_ok=True)
            path.write_bytes(b"previous completed artifact")
    before = {path: path.read_bytes() if path.exists() else None for path in paths}
    monkeypatch.setattr(portable, "EXCEL_CONTENT_PROOF_MAX_BYTES", 1)
    with pytest.raises(ValueError, match="解壓容量超限"):
        if kind == "expected":
            sp.export_pending_for_gpt(root)
        else:
            ar.export_actual_review_package(root, [], version=sp.VERSION,
                session_id=manifest["session_id"], session_schema_version=manifest["session_schema_version"],
                workbook_schema_version=manifest["workbook_schema_version"],
                review_id_schema_version=manifest["review_id_schema_version"], portable_source=(manifest, db))
    assert {path: path.read_bytes() if path.exists() else None for path in paths} == before
    assert not list(root.glob(".ag-*"))
