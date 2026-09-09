"""Independent-review regressions, including the actual pinned baseline reader."""
from __future__ import annotations

from contextlib import closing, nullcontext, redirect_stderr, redirect_stdout
import io
import json
import multiprocessing
from pathlib import Path
import sqlite3
import subprocess
import sys
import tempfile
import types
import unittest
from unittest.mock import patch

import fitz
from openpyxl import load_workbook

import actual_review as review
import global_exact_glyph_library as lib
import legacy_global_migration as migration
import migrate_legacy_global_candidates as cli
from test_global_legacy_migration_v580 import Project, ROOT, seal, sha, tree_bytes
from test_global_promotion_v580 import FakeDocument, approve, intent, sql_rows
from test_exact_glyph_identity_v580 import composite_glyph, synthetic_sfnt


BASELINE = "d2d558808bd2902857fad30f824d8bdc352beb73"
NOW = "2026-09-09T00:00:00Z"


def store_evidence_bytes(repo):
    # SQLite may create/remove SHM and an empty WAL on a read-only open. SHM is
    # the transient lock/index, not durable evidence. Compare the main DB, every
    # nonempty WAL byte and any other file; logical tables are also checked below.
    return {name: raw for name, raw in tree_bytes(repo.path.parent).items()
            if name != repo.path.name + "-shm" and not (name == repo.path.name + "-wal" and not raw)}


def approval_request(repo, exact):
    glyph_id = lib.compute_global_glyph_id(**exact)
    glyph = next(row for row in sql_rows(repo, "glyph_truth") if row["glyph_id"] == glyph_id)
    evidence = [row for row in sql_rows(repo, "source_evidence") if row["glyph_id"] == glyph_id]
    ids = sorted(row["evidence_id"] for row in evidence)
    reading = evidence[0]["reading"]
    return dict(glyph_id=glyph_id, reading=reading, glyph_revision=glyph["revision"], quorum_evidence_ids=ids,
                quorum_digest=lib.compute_quorum_digest(glyph_id, reading, ids),
                promotion_policy_version=lib.GLOBAL_PROMOTION_POLICY_VERSION, approval_source=lib.EXPLICIT_MANUAL_APPROVAL)


def correction_delivery(repo, action, payload):
    if action == "approval":
        return lib.approve_global_exact_glyph(repo, **payload)
    deliver = lib.deliver_global_direct_evidence if action == "direct" else lib.deliver_global_migration
    return deliver(repo, [payload])


def correction_worker(root, action, payload, barrier, queue):
    try:
        repo = lib.GlobalExactGlyphRepository.resolved(root, busy_timeout_ms=10000, write_retry_limit=2)
        barrier.wait(30)
        result = correction_delivery(repo, action, payload)
        queue.put(("ok", result))
    except Exception as exc:
        queue.put(("error", type(exc).__name__, str(exc)))


class MigrationReviewTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="migration-review-")
        self.addCleanup(temporary.cleanup)
        self.base = Path(temporary.name)

    def assert_source_blocked(self, project, message):
        source_before = tree_bytes(project.root)
        for existing in (False, True):
            with self.subTest(existing=existing):
                repo = lib.GlobalExactGlyphRepository.resolved(self.base / str(existing))
                if existing:
                    lib.deliver_global_direct_evidence(repo, [intent()])
                global_before = tree_bytes(repo.path.parent) if existing else {}
                with patch.object(lib.GlobalExactGlyphRepository, "initialize", side_effect=AssertionError), \
                     patch.object(lib, "_ingest_migration", side_effect=AssertionError), \
                     self.assertRaisesRegex(migration.MigrationValidationError, message):
                    migration.migrate_project(project.root, global_library_root=repo.path.parent, apply=True)
                self.assertEqual(tree_bytes(project.root), source_before)
                if existing:
                    self.assertEqual(tree_bytes(repo.path.parent), global_before)
                    self.assertEqual(sql_rows(repo, "migration_candidate"), [])
                else:
                    self.assertFalse(repo.path.parent.exists())

    def test_quarantined_retained_reading_missing_blocks_before_any_global_write(self):
        for cff in (False, True):
            with self.subTest(cff=cff):
                project = Project(self.base / f"missing-{cff}", cff=cff, reading="ㄇ")
                project.learning["verification_level"] = lib.QUARANTINED_CONFLICT
                project.save_learning()
                project.conflict(readings="ㄅ|ㄆ")
                self.assert_source_blocked(project, "retained reading missing")

    def test_quarantined_stale_sample_blocks_even_with_valid_conflict_samples(self):
        project = Project(self.base / "stale", count=2)
        project.learning["verification_level"] = lib.QUARANTINED_CONFLICT
        project.save_learning()
        project.conflict(ids=project.rows[0]["occurrence_id"])
        book = load_workbook(project.actual)
        sheet = book["實際注音"]
        headers = [cell.value for cell in sheet[1]]
        sheet.cell(3, headers.index("TTF字形SHA256") + 1).value = None
        book.save(project.actual)
        book.close()
        project.manifest["records"][1]["source_record"]["TTF字形SHA256"] = ""
        project.manifest["pdfs"][0]["actual_workbook_sha256"] = sha(project.actual.read_bytes())
        seal(project)
        self.assert_source_blocked(project, "different current identity")

    def test_valid_quarantined_learning_keeps_matching_reconfirmation_targets(self):
        for cff in (False, True):
            with self.subTest(cff=cff):
                project = Project(self.base / f"valid-{cff}", cff=cff, count=2)
                project.learning["verification_level"] = lib.QUARANTINED_CONFLICT
                project.save_learning()
                # Registry has one sample; learning retains an additional one.
                project.conflict(ids=project.rows[0]["occurrence_id"])
                before = tree_bytes(project.root)
                repo = lib.GlobalExactGlyphRepository.resolved(self.base / f"global-{cff}")
                result = migration.migrate_project(project.root, global_library_root=repo.path.parent, apply=True)
                self.assertEqual(result["mode"], "APPLIED")
                imports = {item["payload"]["reading"]: item["payload"] for item in result["intents"]}
                self.assertEqual(set(imports), {"ㄅ", "ㄆ"})
                self.assertEqual({item["legacy_evidence_sha256"] for item in imports.values()},
                                 {sha((project.evidence / review.GLYPH_CONFLICT_FILE).read_bytes())})
                targets = result["reconfirmation_targets"]
                for reading, expected in (("ㄅ", {row["occurrence_id"] for row in project.rows}),
                                          ("ㄆ", {project.rows[0]["occurrence_id"]})):
                    self.assertEqual({item["occurrence_id"] for item in targets
                                      if item["import_id"] == imports[reading]["import_id"]}, expected)
                self.assertEqual(len(targets), 3)  # The shared sample is not duplicated per import.
                self.assertTrue(all("reading" not in item for item in targets))
                self.assertEqual(tree_bytes(project.root), before)
                self.assertEqual(sql_rows(repo, "source_evidence"), [])
                self.assertEqual(sql_rows(repo, "promotion_approval"), [])
                snapshot = repo.load_snapshot()
                self.assertEqual(snapshot.trusted_identities, ())
                self.assertEqual(snapshot.quarantined_identities[0].conflicting_readings, ("ㄅ", "ㄆ"))

    def test_empty_apply_strict_validation_matrix(self):
        for donor in ("none", "noneligible"):
            project = Project(self.base / donor, record=composite_glyph() if donor == "noneligible" else None)
            if donor == "none":
                project.save_learning([])
            font = synthetic_sfnt([composite_glyph()], font_name_marker=b"test")
            for state in ("absent", "valid", "corrupt", "unknown", "invalid-row"):
                with self.subTest(donor=donor, state=state):
                    repo = lib.GlobalExactGlyphRepository.resolved(self.base / f"{donor}-{state}")
                    if state == "corrupt":
                        repo.path.parent.mkdir()
                        repo.path.write_bytes(b"corrupt sqlite")
                    elif state != "absent":
                        # Nonempty valid store makes no-op audit/generation assertions meaningful.
                        lib.deliver_global_direct_evidence(repo, [intent()])
                        if state in {"unknown", "invalid-row"}:
                            with closing(sqlite3.connect(repo.path)) as connection:
                                connection.execute("PRAGMA user_version=999" if state == "unknown" else
                                                   "UPDATE glyph_truth SET direct_source_count=999")
                                connection.commit()
                    project_before = tree_bytes(project.root)
                    global_before = store_evidence_bytes(repo) if repo.path.parent.exists() else {}
                    tables_before = ({name: sql_rows(repo, name) for name in lib._REQUIRED_COLUMNS}
                                     if state == "valid" else None)
                    args = ["--project-output", str(project.root), "--global-library-root", str(repo.path.parent)]
                    for apply in (False, True):
                        out, error = io.StringIO(), io.StringIO()
                        pdf = (patch.object(migration.fitz, "open", return_value=FakeDocument(font))
                               if donor == "noneligible" else nullcontext())
                        with pdf, redirect_stdout(out), redirect_stderr(error), \
                             patch.object(lib.GlobalExactGlyphRepository, "initialize", side_effect=AssertionError), \
                             patch.object(lib.GlobalExactGlyphRepository, "write_transaction", side_effect=AssertionError):
                            code = cli.main(args + (["--apply"] if apply else []))
                        if state in {"absent", "valid"}:
                            self.assertEqual(code, 0, error.getvalue())
                            result = json.loads(out.getvalue())
                            self.assertEqual(result["intents"], [])
                            self.assertEqual(len(result["skipped"]), int(donor == "noneligible"))
                            self.assertEqual(result["mode"], "APPLIED" if apply else "DRY_RUN")
                            if apply:
                                self.assertEqual(result["receipts"], [])
                        else:
                            self.assertEqual(code, 1)
                            self.assertEqual(json.loads(error.getvalue())["status"], "BLOCKED")
                            expected_error = {"corrupt": "GlobalLibraryCorruptError", "unknown": "GlobalLibrarySchemaError",
                                              "invalid-row": "GlobalLibraryValidationError"}[state]
                            self.assertEqual(json.loads(error.getvalue())["error_type"], expected_error)
                            self.assertEqual(out.getvalue(), "")
                        self.assertEqual(tree_bytes(project.root), project_before)
                        if state == "absent":
                            self.assertFalse(repo.path.parent.exists())
                        else:
                            self.assertEqual(store_evidence_bytes(repo), global_before)
                    if tables_before is not None:
                        self.assertEqual({name: sql_rows(repo, name) for name in lib._REQUIRED_COLUMNS}, tables_before)


class ReconfirmationBBoxTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="migration-bbox-")
        self.addCleanup(temporary.cleanup)
        self.base = Path(temporary.name)

    def change_source(self, project, changes, *, row_index=0, remove=()):
        """Reseal mutually agreeing source records without changing the PDF."""
        book = load_workbook(project.actual)
        sheet = book["實際注音"]
        headers = [cell.value for cell in sheet[1]]
        record = project.manifest["records"][row_index]
        for field, value in changes.items():
            sheet.cell(row_index + 2, headers.index(field) + 1).value = value
            project.rows[row_index][field] = value
            record["source_record"][field] = value
            if field in ("x0", "y0", "x1", "y1"):
                record[field] = value
            elif field in ("字元", "實體頁碼"):
                record[{"字元": "char", "實體頁碼": "physical_page"}[field]] = value
        for field in remove:
            sheet.delete_cols(headers.index(field) + 1)
            headers.remove(field)
            project.rows[row_index].pop(field, None)
            record.pop(field, None)
            record["source_record"].pop(field, None)
        book.save(project.actual)
        book.close()
        project.manifest["pdfs"][0]["actual_workbook_sha256"] = sha(project.actual.read_bytes())
        seal(project)

    def assert_blocked_before_global_write(self, project):
        source_before = tree_bytes(project.root)
        for existing in (False, True):
            with self.subTest(existing=existing):
                repo = lib.GlobalExactGlyphRepository.resolved(self.base / (project.root.name + str(existing)))
                if existing:
                    lib.deliver_global_direct_evidence(repo, [intent()])
                before = store_evidence_bytes(repo) if existing else {}
                tables = {name: sql_rows(repo, name) for name in lib._REQUIRED_COLUMNS} if existing else {}
                for apply in (False, True):
                    out, error = io.StringIO(), io.StringIO()
                    with redirect_stdout(out), redirect_stderr(error), \
                         patch.object(lib, "deliver_global_migration", side_effect=AssertionError), \
                         patch.object(lib.GlobalExactGlyphRepository, "initialize", side_effect=AssertionError), \
                         patch.object(lib.GlobalExactGlyphRepository, "write_transaction", side_effect=AssertionError):
                        code = cli.main(["--project-output", str(project.root), "--global-library-root", str(repo.path.parent)]
                                        + (["--apply"] if apply else []))
                    self.assertEqual(code, 1)
                    self.assertEqual(out.getvalue(), "")
                    report = json.loads(error.getvalue())
                    self.assertEqual(report["status"], "BLOCKED")
                    self.assertEqual(report["error_type"], "MigrationValidationError")
                    self.assertEqual(tree_bytes(project.root), source_before)
                    if existing:
                        self.assertEqual(store_evidence_bytes(repo), before)
                        self.assertEqual({name: sql_rows(repo, name) for name in lib._REQUIRED_COLUMNS}, tables)
                    else:
                        self.assertFalse(repo.path.parent.exists())

    def test_ttf_cff_targets_use_complete_pdf_bbox_including_tolerated_rounding(self):
        for cff in (False, True):
            for offset in (0, 0.5):
                with self.subTest(cff=cff, offset=offset):
                    project = Project(self.base / f"valid-{cff}-{offset}", cff=cff, count=2)
                    with fitz.open(project.pdf) as document:
                        boxes = [tuple(char[3]) for span in document[0].get_texttrace() for char in span["chars"]]
                    for index, box in enumerate(boxes):
                        self.change_source(project, {field: value + offset for field, value in
                            zip(("x0", "y0", "x1", "y1"), box)}, row_index=index)
                    before = tree_bytes(project.root)
                    repo = lib.GlobalExactGlyphRepository.resolved(self.base / f"valid-global-{cff}-{offset}")
                    for apply in (False, True):
                        result = migration.migrate_project(project.root, global_library_root=repo.path.parent, apply=apply)
                        targets = {item["occurrence_id"]: item for item in result["reconfirmation_targets"]}
                        self.assertEqual(len(targets), 2)
                        for row, box in zip(project.rows, boxes):
                            target = targets[row["occurrence_id"]]
                            self.assertEqual(target["bbox"], list(box))
                            self.assertEqual(target["identity"], project.exact)
                        self.assertEqual(tree_bytes(project.root), before)
                        if apply:
                            self.assertEqual(len(result["receipts"]), 1)

    def test_resealed_wrong_right_bottom_bbox_blocks_whole_project_before_write(self):
        for cff in (False, True):
            for index, changes in enumerate(({"x1": 9999}, {"y1": 9999}, {"x1": 9999, "y1": 9999})):
                with self.subTest(cff=cff, changes=changes):
                    project = Project(self.base / f"forged-{cff}-{index}", cff=cff, count=2)
                    pdf_before = project.pdf.read_bytes()
                    # First sample is valid: no partial import may precede rejection of the second.
                    self.change_source(project, changes, row_index=1)
                    self.assertEqual(project.pdf.read_bytes(), pdf_before)
                    self.assert_blocked_before_global_write(project)

    def test_missing_nonfinite_and_invalid_bbox_coordinates_block_before_write(self):
        for cff in (False, True):
            for field in ("x0", "y0", "x1", "y1"):
                for index, value in enumerate((None, "", "NaN", "Infinity", "-Infinity", True, "not-a-number")):
                    with self.subTest(cff=cff, field=field, value=value):
                        project = Project(self.base / f"invalid-{cff}-{field}-{index}", cff=cff)
                        self.change_source(project, {field: value})
                        self.assert_blocked_before_global_write(project)
            project = Project(self.base / f"missing-columns-{cff}", cff=cff)
            self.change_source(project, {}, remove=("x1", "y1"))
            self.assert_blocked_before_global_write(project)
            for index, changes in enumerate(({"x1": 60}, {"x1": 59}, {"y0": 90, "y1": 90},
                                               {"y0": 90, "y1": 89}, {"x0": -9999, "y0": -9999})):
                project = Project(self.base / f"extent-{cff}-{index}", cff=cff)
                self.change_source(project, changes)
                self.assert_blocked_before_global_write(project)

    def test_existing_locator_tolerance_is_not_widened_for_any_coordinate(self):
        for cff in (False, True):
            for field in ("x0", "y0", "x1", "y1"):
                with self.subTest(cff=cff, field=field):
                    project = Project(self.base / f"tolerance-{cff}-{field}", cff=cff)
                    self.change_source(project, {field: project.rows[0][field] + 0.801})
                    self.assert_blocked_before_global_write(project)

    def test_duplicate_pdf_text_trace_still_blocks_unique_locator(self):
        for cff in (False, True):
            with self.subTest(cff=cff):
                project = Project(self.base / f"duplicate-{cff}", cff=cff)
                with fitz.open(project.pdf) as document:
                    contents = document[0].get_contents()[0]
                    raw = document.xref_stream(contents)
                    document.update_stream(contents, raw + b"\n" + raw)
                    pdf = document.tobytes()
                project.pdf.write_bytes(pdf)
                digest = sha(pdf)
                project.manifest["pdfs"][0]["pdf_sha256"] = digest
                project.manifest["records"][0]["pdf_sha256"] = digest
                book = load_workbook(project.actual)
                meta = book["v5.2中繼資料"]
                for row in meta:
                    if row[0].value == "pdf_sha256":
                        row[1].value = digest
                book.save(project.actual)
                book.close()
                project.manifest["pdfs"][0]["actual_workbook_sha256"] = sha(project.actual.read_bytes())
                seal(project)
                with fitz.open(project.pdf) as document:
                    self.assertEqual(sum(len(span["chars"]) for span in document[0].get_texttrace()), 2)
                self.assert_blocked_before_global_write(project)

    def test_full_bbox_does_not_bypass_character_font_glyph_or_page_checks(self):
        for cff in (False, True):
            for index, changes in enumerate(({"字元": "另"}, {"font_xref": 1},
                                               {"glyph_id_字形索引": 0}, {"實體頁碼": 0})):
                with self.subTest(cff=cff, changes=changes):
                    project = Project(self.base / f"locator-{cff}-{index}", cff=cff)
                    self.change_source(project, changes)
                    self.assert_blocked_before_global_write(project)


class FrozenV0CompatibilityTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Execute the pinned historical implementation, including its real ID
        # algorithm, schema creator, transaction validation and strict reader.
        # Missing Git history is a failed verification, never a skipped test.
        source = subprocess.check_output(["git", "show", BASELINE + ":global_exact_glyph_library.py"], cwd=ROOT)
        cls.old = types.ModuleType("_phase5_pinned_baseline_library")
        sys.modules[cls.old.__name__] = cls.old
        cls.addClassCleanup(sys.modules.pop, cls.old.__name__)
        exec(compile(source, BASELINE + ":global_exact_glyph_library.py", "exec"), cls.old.__dict__)

    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="migration-v0-")
        self.addCleanup(temporary.cleanup)
        self.base = Path(temporary.name)
        self.exact = {"kind": lib.TTF_GLYF_SHA256, "style_group": "", "glyph_sha256": "a" * 64,
                      "identity_eligibility": lib.GLOBAL_ELIGIBLE_SIMPLE_GLYF_V1}

    def fixture(self, readings, *, quarantine=None, level="USER_VERIFIED_SINGLE", status="CANDIDATE", name="global",
                state="CANDIDATE"):
        old_repo = self.old.GlobalExactGlyphRepository.resolved(self.base / name)
        old_repo.initialize()
        with old_repo.write_transaction() as tx:
            tx.insert_candidate_identity(self.exact)
            glyph_id = self.old.compute_global_glyph_id(**self.exact)
            for index, reading in enumerate(readings):
                args = dict(glyph_id=glyph_id, reading=reading, legacy_project_id=f"old-project-{index}",
                            legacy_project_sha256="b" * 64, legacy_evidence_sha256="c" * 64,
                            old_verification_level=level)
                row = dict(import_id=self.old.compute_migration_import_id(**args), **args, status=status, imported_at=NOW)
                columns = self.old._REQUIRED_COLUMNS["migration_candidate"]
                tx._connection.execute("INSERT INTO migration_candidate (" + ",".join(columns) + ") VALUES (" +
                                       ",".join("?" for _ in columns) + ")", tuple(row[key] for key in columns))
            generation = tx.bump_generation()
            if quarantine is not None:
                tx._connection.execute("UPDATE glyph_truth SET state=?, revision=revision+1 WHERE glyph_id=?",
                                       (self.old.QUARANTINED_CONFLICT, glyph_id))
                tx._connection.execute("INSERT INTO glyph_conflict VALUES (?, ?, ?, ?, ?, ?, ?, NULL, NULL)",
                    (self.old.compute_glyph_conflict_id(glyph_id, quarantine, generation), glyph_id, self.old.OPEN,
                     self.old._canonical_json(sorted(quarantine)), NOW, NOW, generation))
        if state in {"PROMOTION_READY", "VERIFIED_GLOBAL"}:
            self.old.deliver_global_direct_evidence(old_repo, [intent(1, identity=self.exact), intent(2, identity=self.exact)])
            if state == "VERIFIED_GLOBAL":
                self.old.approve_global_exact_glyph(old_repo, **approval_request(old_repo, self.exact))
        self.assertEqual(old_repo.load_snapshot().store_status, self.old.VALID)
        return old_repo, lib.GlobalExactGlyphRepository.resolved(old_repo.path.parent)

    def test_frozen_schema_and_id_algorithm_remain_the_compatibility_boundary(self):
        # Adopted v1.1 section 5.6 overrides the general historical audit's
        # candidate rule; it does not relax the frozen ID or physical schema.
        self.assertEqual((self.old.GLOBAL_LIBRARY_SCHEMA_VERSION, self.old.GLOBAL_LIBRARY_SQLITE_USER_VERSION), ("1.0", 1))
        old, repo = self.fixture(["ㄅ", "ㄆ"])
        self.assertEqual(repo.load_snapshot().store_status, lib.VALID)
        for row in sql_rows(repo, "migration_candidate"):
            args = {key: row[key] for key in ("glyph_id", "reading", "legacy_project_id", "legacy_project_sha256",
                                            "legacy_evidence_sha256", "old_verification_level")}
            self.assertEqual(row["import_id"], self.old.compute_migration_import_id(**args))
            self.assertNotEqual(row["import_id"], lib.compute_migration_import_id(**args))

    def test_single_agreeing_insufficient_and_explicit_conflict_read_without_rewrite(self):
        cases = ((["ㄅ"], None, "USER_VERIFIED_SINGLE", "CANDIDATE"),
                 (["ㄅ", "ㄅ"], None, "VERIFIED_EXACT_GLYPH", "CANDIDATE"),
                 (["ㄅ"], None, "USER_VERIFIED_SINGLE", "INSUFFICIENT"),
                 (["ㄅ", "ㄆ"], ["ㄅ", "ㄆ"], "USER_VERIFIED_SINGLE", "CANDIDATE"),
                 (["ㄅ", "ㄆ"], ["ㄅ", "ㄆ"], "QUARANTINED_CONFLICT", "CONFLICT"))
        for index, (readings, quarantine, level, status) in enumerate(cases):
            with self.subTest(index=index):
                old, repo = self.fixture(readings, quarantine=quarantine, level=level, status=status, name=str(index))
                before = store_evidence_bytes(repo)
                snapshot = repo.load_snapshot()
                self.assertEqual(snapshot.store_status, lib.VALID)
                self.assertEqual(snapshot.trusted_identities, ())
                self.assertEqual(bool(snapshot.quarantined_identities), quarantine is not None)
                self.assertEqual(store_evidence_bytes(repo), before)
                for table in ("source_evidence", "promotion_approval", "processed_intent", "provenance_event"):
                    self.assertEqual(sql_rows(repo, table), [])
                with closing(sqlite3.connect(repo.path)) as connection:
                    self.assertEqual(connection.execute("PRAGMA user_version").fetchone()[0], 1)
                self.assertEqual(old.load_snapshot().store_status, self.old.VALID)

    def test_baseline_validator_omissions_stay_blocked_and_preserve_bytes(self):
        cases = ((["ㄅ", "ㄆ"], None, "QUARANTINED_CONFLICT", "CONFLICT", "requires quarantine"),
                 (["ㄅ", "ㄆ", "ㄇ"], ["ㄅ", "ㄆ"], "USER_VERIFIED_SINGLE", "CANDIDATE", "omits migration reading"))
        for index, (readings, quarantine, level, status, message) in enumerate(cases):
            with self.subTest(index=index):
                old, repo = self.fixture(readings, quarantine=quarantine, level=level, status=status, name=str(index))
                before = store_evidence_bytes(repo)
                self.assertEqual(old.load_snapshot().store_status, self.old.VALID)
                with self.assertRaisesRegex(lib.GlobalLibraryValidationError, message):
                    repo.load_snapshot()
                with self.assertRaises(lib.GlobalLibraryValidationError):
                    lib.deliver_global_direct_evidence(repo, [intent(identity=self.exact)])
                self.assertEqual(store_evidence_bytes(repo), before)

    def test_illegal_id_rejected_by_both_actual_validators(self):
        old, repo = self.fixture(["ㄅ"])
        with closing(sqlite3.connect(repo.path)) as connection:
            connection.execute("UPDATE migration_candidate SET import_id=?", ("f" * 64,))
            connection.commit()
        before = store_evidence_bytes(repo)
        for module, reader in ((self.old, old), (lib, repo)):
            with self.assertRaises(module.GlobalLibraryValidationError):
                reader.load_snapshot()
        self.assertEqual(store_evidence_bytes(repo), before)

    def test_current_imports_cannot_submit_v0_ids(self):
        old, repo = self.fixture(["ㄅ"])
        row = sql_rows(repo, "migration_candidate")[0]
        item = lib.make_migration_intent(self.exact, **{key: row[key] for key in (
            "reading", "legacy_project_id", "legacy_project_sha256", "legacy_evidence_sha256",
            "old_verification_level", "status")})
        self.assertNotEqual(item["payload"]["import_id"], row["import_id"])
        item["payload"]["import_id"] = row["import_id"]
        item["intent_id"] = "gmi_" + row["import_id"]
        item["payload_digest"] = lib._canonical_sha256(item["payload"])
        before = tree_bytes(repo.path.parent)
        with self.assertRaises(lib.GlobalLibraryValidationError):
            lib.deliver_global_migration(repo, [item])
        self.assertEqual(tree_bytes(repo.path.parent), before)

    def test_later_direct_and_migration_conflicts_never_rewrite_v0_rows(self):
        for delivery in ("direct", "migration"):
            with self.subTest(delivery=delivery):
                old, repo = self.fixture(["ㄅ"], name=delivery)
                before = sql_rows(repo, "migration_candidate")
                if delivery == "direct":
                    lib.deliver_global_direct_evidence(repo, [intent(identity=self.exact, reading="ㄆ")])
                else:
                    item = lib.make_migration_intent(self.exact, reading="ㄆ", legacy_project_id="new-project",
                        legacy_project_sha256="d" * 64, legacy_evidence_sha256="e" * 64,
                        old_verification_level="USER_VERIFIED_SINGLE", status=lib.MIGRATION_CANDIDATE)
                    lib.deliver_global_migration(repo, [item])
                    current = next(row for row in sql_rows(repo, "migration_candidate")
                                   if row["import_id"] == item["payload"]["import_id"])
                    self.assertEqual(current["status"], lib.MIGRATION_CONFLICT)
                self.assertEqual([row for row in sql_rows(repo, "migration_candidate")
                                  if row["import_id"] == before[0]["import_id"]], before)
                snapshot = repo.load_snapshot()
                self.assertEqual(snapshot.trusted_identities, ())
                self.assertEqual(snapshot.quarantined_identities[0].conflicting_readings, ("ㄅ", "ㄆ"))
                self.assertEqual(sql_rows(repo, "promotion_approval"), [])
                self.assertFalse(any(row["intent_id"] == "gmi_" + before[0]["import_id"]
                                     for row in sql_rows(repo, "processed_intent")))

    def test_matching_v0_audit_preserves_verified_digest_and_has_zero_quorum(self):
        old, repo = self.fixture(["ㄅ", "ㄅ"])
        before = sql_rows(repo, "migration_candidate")
        lib.deliver_global_direct_evidence(repo, [intent(1, identity=self.exact), intent(2, identity=self.exact)])
        approve(repo)
        exact = lib.canonical_global_exact_identity(**self.exact)
        snapshot = repo.load_snapshot()
        digest = lib.canonical_logical_subset_digest(snapshot, [exact]).sha256
        self.assertEqual(len(sql_rows(repo, "source_evidence")), 2)
        item = lib.make_migration_intent(self.exact, reading="ㄅ", legacy_project_id="new-project",
            legacy_project_sha256="d" * 64, legacy_evidence_sha256="e" * 64,
            old_verification_level="VERIFIED_EXACT_GLYPH", status=lib.MIGRATION_CANDIDATE)
        lib.deliver_global_migration(repo, [item])
        self.assertEqual(lib.canonical_logical_subset_digest(repo.load_snapshot(), [exact]).sha256, digest)
        self.assertEqual([row for row in sql_rows(repo, "migration_candidate")
                          if row["import_id"] in {r["import_id"] for r in before}], before)
        self.assertEqual(len(sql_rows(repo, "source_evidence")), 2)

    def fresh_request(self, repo, action):
        if action == "direct":
            return intent(3, identity=self.exact)
        if action == "approval":
            return approval_request(repo, self.exact)
        return lib.make_migration_intent(self.exact, reading="ㄅ", legacy_project_id="fresh-project",
            legacy_project_sha256="d" * 64, legacy_evidence_sha256="e" * 64,
            old_verification_level="USER_VERIFIED_SINGLE", status=lib.MIGRATION_CANDIDATE)

    def assert_effective_conflict(self, repo):
        snapshot = repo.load_snapshot()
        glyph_id = lib.compute_global_glyph_id(**self.exact)
        self.assertEqual(snapshot.store_status, lib.VALID)
        self.assertNotIn(glyph_id, {record.glyph_id for record in snapshot.trusted_identities})
        resolution = lib.resolve_exact_glyph_reuse(snapshot, self.exact,
            higher_priority_sources=[("PROJECT", "ㄅ")], lower_priority_sources=[("STATIC", "ㄅ")])
        self.assertTrue(resolution.conflict)
        self.assertEqual(resolution.reading, "")
        self.assertEqual(resolution.conflicting_readings, ("ㄅ", "ㄆ"))
        return snapshot

    def test_ordinary_v0_conflicts_are_readable_and_suppress_candidate_ready_and_verified(self):
        for state in (lib.CANDIDATE, lib.PROMOTION_READY, lib.VERIFIED_GLOBAL):
            # A single v0 contrary to direct/active reading also suppresses reuse.
            for readings in (("ㄅ", "ㄆ"), ("ㄆ",)) if state != lib.CANDIDATE else (("ㄅ", "ㄆ"),):
                with self.subTest(state=state, readings=readings):
                    old, repo = self.fixture(readings, state=state, name=state + str(len(readings)))
                    rows_before = {name: sql_rows(repo, name) for name in lib._REQUIRED_COLUMNS}
                    before = store_evidence_bytes(repo)
                    snapshot = self.assert_effective_conflict(repo)
                    diagnostic = snapshot.read_side_conflicts[0]
                    self.assertEqual(diagnostic.stored_state, state)
                    self.assertIn("durable quarantine/revocation not performed", diagnostic.diagnostic)
                    self.assertEqual(diagnostic.conflicting_readings, ("ㄅ", "ㄆ"))
                    self.assertEqual(snapshot.quarantined_identities, ())
                    self.assertEqual(lib.deliver_global_migration(repo, []), ())
                    self.assertEqual(store_evidence_bytes(repo), before)
                    self.assertEqual({name: sql_rows(repo, name) for name in lib._REQUIRED_COLUMNS}, rows_before)
                    self.assertEqual(old.load_snapshot().store_status, self.old.VALID)

    def test_read_side_conflict_preserves_unrelated_verified_result_and_subset(self):
        old, repo = self.fixture(["ㄅ", "ㄆ"], state=lib.VERIFIED_GLOBAL)
        unrelated = self.exact | {"glyph_sha256": "f" * 64}
        self.old.deliver_global_direct_evidence(old, [intent(8, identity=unrelated), intent(9, identity=unrelated)])
        self.old.approve_global_exact_glyph(old, **approval_request(old, unrelated))
        old_snapshot = old.load_snapshot()
        snapshot = self.assert_effective_conflict(repo)
        before = store_evidence_bytes(repo)
        for exact, same in ((unrelated, True), (self.exact, False)):
            # Convert only the test identity type; baseline uses its own actual digest implementation.
            baseline_digest = self.old.canonical_logical_subset_digest(old_snapshot, [exact]).sha256
            current_digest = lib.canonical_logical_subset_digest(snapshot, [exact]).sha256
            self.assertEqual(baseline_digest == current_digest, same)
            self.assertEqual(self.old.global_exact_glyph_evidence_hashes(old_snapshot, [exact]) ==
                             lib.global_exact_glyph_evidence_hashes(snapshot, [exact]), same)
        self.assertEqual(lib.resolve_exact_glyph_reuse(snapshot, unrelated).reading, "ㄅ")
        self.assertEqual(store_evidence_bytes(repo), before)

    def test_dry_run_and_empty_apply_report_read_side_conflict_without_writes(self):
        old, repo = self.fixture(["ㄅ", "ㄆ"], state=lib.VERIFIED_GLOBAL)
        project = Project(self.base / "project")
        project.save_learning([])
        source_before = tree_bytes(project.root)
        before = store_evidence_bytes(repo)
        for apply in (False, True):
            result = migration.migrate_project(project.root, global_library_root=repo.path.parent, apply=apply)
            self.assertEqual(result["mode"], "APPLIED" if apply else "DRY_RUN")
            self.assertEqual(result["global_store_status"], lib.VALID)
            self.assertEqual(result["global_read_side_conflicts"][0]["stored_state"], lib.VERIFIED_GLOBAL)
            if apply:
                self.assertEqual(result["receipts"], [])
            self.assertEqual(store_evidence_bytes(repo), before)
            self.assertEqual(tree_bytes(project.root), source_before)

    def test_fresh_delivery_and_approval_commit_quarantine_revoke_and_replay(self):
        for state in (lib.CANDIDATE, lib.PROMOTION_READY, lib.VERIFIED_GLOBAL):
            for action in (("direct", "migration") if state == lib.CANDIDATE else ("direct", "migration", "approval")):
                with self.subTest(state=state, action=action):
                    old, repo = self.fixture(["ㄅ", "ㄆ"], state=state, name=state + action)
                    before_rows = sql_rows(repo, "migration_candidate")
                    snapshot = repo.load_snapshot()
                    effective_digest = lib.global_exact_glyph_evidence_hashes(snapshot, [self.exact])
                    payload = self.fresh_request(repo, action)  # Matches one v0 reading, not the other.
                    result = correction_delivery(repo, action, payload)
                    after = self.assert_effective_conflict(repo)
                    self.assertEqual(after.read_side_conflicts, ())
                    self.assertEqual(after.generation, snapshot.generation + 1)
                    self.assertEqual(after.quarantined_identities[0].conflicting_readings, ("ㄅ", "ㄆ"))
                    # Persistence alone does not change the already effective conflict dependency.
                    self.assertEqual(lib.global_exact_glyph_evidence_hashes(after, [self.exact]), effective_digest)
                    glyph = sql_rows(repo, "glyph_truth")[0]
                    self.assertEqual((glyph["state"], glyph["active_reading"]), (lib.QUARANTINED_CONFLICT, None))
                    self.assertEqual(sql_rows(repo, "migration_candidate")[:len(before_rows)], before_rows)
                    self.assertEqual({row["status"] for row in sql_rows(repo, "promotion_approval")},
                                     {lib.APPROVED, lib.REVOKED} if state == lib.VERIFIED_GLOBAL else set())
                    self.assertEqual(len(sql_rows(repo, "source_evidence")),
                                     (0 if state == lib.CANDIDATE else 2) + int(action == "direct"))
                    self.assertFalse(any(row["intent_id"] == "gmi_" + item["import_id"]
                                         for row in sql_rows(repo, "processed_intent") for item in before_rows))
                    self.assertTrue(any(row["event_type"] == "CONFLICT_OPENED" for row in sql_rows(repo, "provenance_event")))
                    if action == "approval":
                        self.assertFalse(result["approval_granted"])
                        self.assertEqual(result["approval_id"], "")
                        self.assertEqual(result["state"], lib.QUARANTINED_CONFLICT)
                    tables = {name: sql_rows(repo, name) for name in lib._REQUIRED_COLUMNS}
                    replay = correction_delivery(repo, action, payload)
                    self.assertEqual(replay["receipt_digest"] if action == "approval" else replay[0].receipt_digest,
                                     result["receipt_digest"] if action == "approval" else result[0].receipt_digest)
                    self.assertTrue(replay["already_processed"] if action == "approval" else replay[0].already_processed)
                    self.assertEqual({name: sql_rows(repo, name) for name in lib._REQUIRED_COLUMNS}, tables)

    def test_conflict_transaction_failure_rolls_back_all_three_delivery_paths(self):
        for action in ("direct", "migration", "approval"):
            with self.subTest(action=action):
                old, repo = self.fixture(["ㄅ", "ㄆ"], state=lib.VERIFIED_GLOBAL, name=action)
                payload = self.fresh_request(repo, action)
                before = {name: sql_rows(repo, name) for name in lib._REQUIRED_COLUMNS}
                original = lib.GlobalWriteTransaction.record_processed_intent
                def fail_after_receipt(*args, **kwargs):
                    original(*args, **kwargs)
                    raise sqlite3.OperationalError("injected after conflict, revocation, generation and receipt")
                with patch.object(lib.GlobalWriteTransaction, "record_processed_intent", fail_after_receipt), \
                     self.assertRaises(lib.GlobalLibraryError):
                    correction_delivery(repo, action, payload)
                self.assertEqual({name: sql_rows(repo, name) for name in lib._REQUIRED_COLUMNS}, before)
                self.assertEqual(self.assert_effective_conflict(repo).read_side_conflicts[0].stored_state, lib.VERIFIED_GLOBAL)

    def test_concurrent_mutations_and_approval_replay_preserve_v0_counterevidence(self):
        context = multiprocessing.get_context("spawn")
        for actions in (("direct", "migration"), ("approval", "approval")):
            with self.subTest(actions=actions):
                old, repo = self.fixture(["ㄅ", "ㄆ"], state=lib.VERIFIED_GLOBAL, name="-".join(actions))
                before = sql_rows(repo, "migration_candidate")
                generation = repo.load_snapshot().generation
                barrier, queue = context.Barrier(2), context.Queue()
                processes = [context.Process(target=correction_worker,
                    args=(str(repo.path.parent), action, self.fresh_request(repo, action), barrier, queue)) for action in actions]
                try:
                    for process in processes:
                        process.start()
                    results = [queue.get(timeout=45) for _ in processes]
                    for process in processes:
                        process.join(45)
                        self.assertEqual(process.exitcode, 0)
                    self.assertTrue(all(result[0] == "ok" for result in results), results)
                    snapshot = self.assert_effective_conflict(repo)
                    self.assertEqual(snapshot.generation, generation + (1 if actions[0] == "approval" else 2))
                    self.assertEqual(snapshot.read_side_conflicts, ())
                    self.assertEqual(sql_rows(repo, "migration_candidate")[:len(before)], before)
                    self.assertEqual({row["status"] for row in sql_rows(repo, "promotion_approval")}, {lib.APPROVED, lib.REVOKED})
                    if actions[0] == "approval":
                        self.assertEqual(results[0][1]["receipt_digest"], results[1][1]["receipt_digest"])
                finally:
                    for process in processes:
                        if process.is_alive():
                            process.terminate()
                            process.join()
                    queue.close()
                    queue.join_thread()

    def test_read_only_compatibility_preserves_nonempty_wal_and_every_v0_column(self):
        old, repo = self.fixture(["ㄅ", "ㄆ"])
        with closing(sqlite3.connect(repo.path)) as holder:
            holder.execute("PRAGMA wal_autocheckpoint=0")
            holder.execute("BEGIN")
            holder.execute("SELECT generation FROM library_meta").fetchall()
            with old.write_transaction() as tx:
                tx.bump_generation()
            wal = repo.path.with_name(repo.path.name + "-wal")
            self.assertGreater(wal.stat().st_size, 0)
            rows = sql_rows(repo, "migration_candidate")
            before = store_evidence_bytes(repo)
            self.assert_effective_conflict(repo)
            lib.deliver_global_migration(repo, [])
            self.assertEqual(store_evidence_bytes(repo), before)
            self.assertEqual(sql_rows(repo, "migration_candidate"), rows)

    def test_occurrence_direct_override_remains_first_and_decoder_fallback_remains_independent(self):
        from test_global_exact_glyph_read_reuse_v580 import TTFWorkbookReadReuseIntegrationTests
        old, repo = self.fixture(["ㄅ", "ㄆ"], state=lib.VERIFIED_GLOBAL)
        snapshot = repo.load_snapshot()
        helper = TTFWorkbookReadReuseIntegrationTests()
        fallback, audit = helper._decode_fixture(self.base, snapshot, name="fallback")
        self.assertEqual(fallback["實際注音"], "ㄇ")
        self.assertEqual(audit[0]["已由獨立fallback解碼筆數"], 1)
        actual, audit = helper._decode_fixture(self.base, snapshot, occurrence_override="ㄈ", name="direct")
        self.assertEqual(actual["自動解碼原值"], "ㄇ")
        self.assertEqual(actual["實際注音"], "ㄈ")
        self.assertEqual(actual["解碼依據"], "使用者原頁人工覆核（出現位置限定）")
        self.assertEqual(audit[0]["已由獨立fallback解碼筆數"], 0)

    def test_cff_v0_conflict_uses_the_same_style_scoped_dependency_gate(self):
        self.exact = dict(kind=lib.CFF_GLYPH_SHA256, style_group="BIAOKAI_W5", glyph_sha256="a" * 64,
                          identity_eligibility=lib.GLOBAL_ELIGIBLE_COMPLETE_CFF_RECORDING_V1)
        old, repo = self.fixture(["ㄅ", "ㄆ"], state=lib.VERIFIED_GLOBAL)
        snapshot = self.assert_effective_conflict(repo)
        other_style = self.exact | {"style_group": "HEI_W5"}
        self.assertFalse(lib.resolve_exact_glyph_reuse(snapshot, other_style).conflict)
        self.assertEqual(lib.resolve_exact_glyph_reuse(snapshot, other_style).global_effective_state, lib.ABSENT)

    def test_ordinary_v0_exception_cannot_hide_direct_approval_schema_or_count_corruption(self):
        for bad in ("direct-candidate", "direct-verified", "approval", "counts", "schema"):
            with self.subTest(bad=bad):
                old, repo = self.fixture(["ㄅ", "ㄆ"], state=lib.VERIFIED_GLOBAL, name=bad)
                if bad.startswith("direct-"):
                    self.old.deliver_global_direct_evidence(old, [intent(3, reading="ㄆ", identity=self.exact)])
                with closing(sqlite3.connect(repo.path)) as connection:
                    if bad.startswith("direct-"):
                        connection.execute("DELETE FROM glyph_conflict")
                        connection.execute("UPDATE glyph_truth SET state=?, active_reading=?",
                            (lib.CANDIDATE, None) if bad == "direct-candidate" else (lib.VERIFIED_GLOBAL, "ㄅ"))
                    elif bad == "approval":
                        connection.execute("UPDATE promotion_approval SET quorum_digest=?", ("f" * 64,))
                    elif bad == "counts":
                        connection.execute("UPDATE glyph_truth SET direct_source_count=99")
                    else:
                        connection.execute("PRAGMA user_version=999")
                    connection.commit()
                before = store_evidence_bytes(repo)
                for module, reader in ((self.old, old), (lib, repo)):
                    with self.assertRaises(module.GlobalLibraryError):
                        reader.load_snapshot()
                self.assertEqual(store_evidence_bytes(repo), before)

    def test_stale_or_invalid_approval_request_cannot_mutate_compatible_store(self):
        old, repo = self.fixture(["ㄅ", "ㄆ"], state=lib.PROMOTION_READY)
        request = approval_request(repo, self.exact)
        before = {name: sql_rows(repo, name) for name in lib._REQUIRED_COLUMNS}
        for change in (dict(glyph_revision=request["glyph_revision"] + 1), dict(quorum_digest="f" * 64),
                       dict(quorum_evidence_ids=request["quorum_evidence_ids"][:1]), dict(approval_source="FUTURE")):
            with self.subTest(change=change), self.assertRaises(lib.GlobalLibraryValidationError):
                lib.approve_global_exact_glyph(repo, **(request | change))
            self.assertEqual({name: sql_rows(repo, name) for name in lib._REQUIRED_COLUMNS}, before)

    def test_new_import_and_approval_conflict_receipt_provenance_remain_required(self):
        for action, table in (("migration", "processed_intent"), ("migration", "provenance_event"),
                              ("approval", "processed_intent"), ("approval", "provenance_event")):
            with self.subTest(action=action, table=table):
                old, repo = self.fixture(["ㄅ", "ㄆ"], state=lib.PROMOTION_READY, name=action + table)
                payload = self.fresh_request(repo, action)
                correction_delivery(repo, action, payload)
                with closing(sqlite3.connect(repo.path)) as connection:
                    connection.execute("DELETE FROM " + table)
                    connection.commit()
                before = store_evidence_bytes(repo)
                with self.assertRaises(lib.GlobalLibraryValidationError):
                    repo.load_snapshot()
                self.assertEqual(store_evidence_bytes(repo), before)
