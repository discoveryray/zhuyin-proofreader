from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
import sys
sys.path.insert(0, str(ROOT))

from occurrence_ledger import (  # noqa: E402
    LEDGER_SCHEMA_VERSION,
    REVIEW_ID_SCHEMA_VERSION,
    SESSION_SCHEMA_VERSION,
    WORKBOOK_SCHEMA_VERSION,
    build_occurrence_ledger,
    prepare_occurrence_rows,
)
from standalone_proofread import (  # noqa: E402
    VERSION,
    _clear_actual_dependent_events,
    export_pending_for_gpt,
    import_gpt_decisions,
    json_load,
    json_save,
    normalize_db,
    seal_manifest,
)


def make_manifest(*, state="EXPECTED_UNRESOLVED", actual="ㄐㄩㄝˊ", expected_set=None, expected_evidence=""):
    source = {
        "pdf_sha256": "a" * 64, "pdf": "book.pdf", "pdf_name": "book.pdf",
        "實體頁碼": 1, "課本頁": 1, "字元": "角", "實際注音": actual,
        "解碼依據": "glyph evidence", "穩定注音鍵": "F#1", "font": "F", "font_xref": 1,
        "glyph_id_字形索引": 2, "注音元件ID": 1, "x0": 1, "y0": 2, "x1": 3, "y1": 4,
        "source_row_number": 1,
    }
    prepare_occurrence_rows("a" * 64, [source])
    classification = {
        "state": state,
        "actual": actual,
        "actual_evidence": "glyph evidence",
        "expected_set": expected_set or [],
        "expected_evidence": expected_evidence,
        "context_evidence": "角色/角@1",
        "comparison_result": "MATCH" if state == "PASS" else ("MISMATCH" if state == "DIFFERENCE_PENDING_CONFIRMATION" else ""),
        "source_view": "未建立預期音" if state == "EXPECTED_UNRESOLVED" else ("單音一致" if state == "PASS" else "差異候選"),
        "source_record": source,
    }
    ledger = build_occurrence_ledger([source], {source["occurrence_id"]: classification})
    return seal_manifest({
        "version": VERSION,
        "session_schema_version": SESSION_SCHEMA_VERSION,
        "workbook_schema_version": WORKBOOK_SCHEMA_VERSION,
        "ledger_schema_version": LEDGER_SCHEMA_VERSION,
        "review_id_schema_version": REVIEW_ID_SCHEMA_VERSION,
        "session_id": "session-safe-rebase",
        "records": ledger,
        "actual_source_ids": [source["occurrence_id"]],
        "regression_gate": {"ok": False, "required": 1, "executed": 1, "passed": 0, "failed": 1, "not_executed": 0, "duplicate_case_id": 0},
    })


def fill_expected_action(xlsx: Path, *, action="補建expected證據", expected="ㄐㄩㄝˊ"):
    wb = load_workbook(xlsx)
    ws = wb["待判定候選"]
    headers = {str(cell.value or ""): cell.column for cell in ws[1]}
    ws.cell(2, headers["action"], action)
    ws.cell(2, headers["proposed_expected_set"], expected)
    ws.cell(2, headers["proposed_expected_evidence"], "現版完整詞條")
    ws.cell(2, headers["proposed_context_evidence"], "角色/角@1")
    wb.save(xlsx)


class GptSafeRebaseTests(unittest.TestCase):
    def test_export_contains_expected_domain_snapshot(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            json_save(output / "校對工作階段.json", make_manifest())
            json_save(output / "人工判定資料庫.json", normalize_db({}))
            xlsx = export_pending_for_gpt(output)
            wb = load_workbook(xlsx, read_only=True, data_only=True)
            try:
                headers = [str(c.value or "") for c in wb["待判定候選"][1]]
                self.assertIn("expected_snapshot", headers)
            finally:
                wb.close()

    def test_expected_action_rebases_after_actual_only_change(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            old_manifest = make_manifest()
            json_save(output / "校對工作階段.json", old_manifest)
            json_save(output / "人工判定資料庫.json", normalize_db({}))
            xlsx = export_pending_for_gpt(output)
            fill_expected_action(xlsx)

            new_manifest = make_manifest(actual="ㄐㄧㄠˇ")
            new_manifest["session_id"] = old_manifest["session_id"]
            seal_manifest(new_manifest)
            json_save(output / "校對工作階段.json", new_manifest)
            fake_report = output / "fake.xlsx"
            with patch("standalone_proofread.generate_report", return_value=fake_report):
                imported, skipped, _ = import_gpt_decisions(output, xlsx)
            self.assertEqual(imported, 1)
            self.assertEqual(skipped, 0)
            db = json_load(output / "人工判定資料庫.json", {})
            event = next(iter(db["events"].values()))
            self.assertTrue(event.get("_safe_expected_rebase"))

    def test_expected_action_rebases_when_actual_change_makes_current_state_pass(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            old_manifest = make_manifest(
                state="DIFFERENCE_PENDING_CONFIRMATION",
                actual="ㄐㄧㄠˇ",
                expected_set=["ㄐㄩㄝˊ"],
                expected_evidence="舊完整詞條",
            )
            json_save(output / "校對工作階段.json", old_manifest)
            json_save(output / "人工判定資料庫.json", normalize_db({}))
            xlsx = export_pending_for_gpt(output)
            fill_expected_action(xlsx, action="解決expected證據", expected="ㄐㄩㄝˊ")

            # Only actual changed.  The same expected/context chain is still valid,
            # and the mechanical state is now PASS.
            new_manifest = make_manifest(
                state="PASS",
                actual="ㄐㄩㄝˊ",
                expected_set=["ㄐㄩㄝˊ"],
                expected_evidence="舊完整詞條",
            )
            new_manifest["session_id"] = old_manifest["session_id"]
            seal_manifest(new_manifest)
            json_save(output / "校對工作階段.json", new_manifest)
            fake_report = output / "fake.xlsx"
            with patch("standalone_proofread.generate_report", return_value=fake_report):
                imported, skipped, _ = import_gpt_decisions(output, xlsx)
            self.assertEqual((imported, skipped), (1, 0))

    def test_changed_expected_domain_still_rejects_stale_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            old_manifest = make_manifest()
            json_save(output / "校對工作階段.json", old_manifest)
            json_save(output / "人工判定資料庫.json", normalize_db({}))
            xlsx = export_pending_for_gpt(output)
            fill_expected_action(xlsx, expected="ㄐㄩㄝˊ")

            new_manifest = make_manifest(
                state="DIFFERENCE_PENDING_CONFIRMATION",
                actual="ㄐㄩㄝˊ",
                expected_set=["ㄐㄧㄠˇ"],
                expected_evidence="另一份現版詞條",
            )
            new_manifest["session_id"] = old_manifest["session_id"]
            seal_manifest(new_manifest)
            json_save(output / "校對工作階段.json", new_manifest)
            with self.assertRaisesRegex(ValueError, "expected 已由目前 resolver 建立為不同讀音"):
                import_gpt_decisions(output, xlsx)

    def test_untouched_stale_rows_do_not_reject_batch(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            old_manifest = make_manifest()
            json_save(output / "校對工作階段.json", old_manifest)
            json_save(output / "人工判定資料庫.json", normalize_db({}))
            xlsx = export_pending_for_gpt(output)

            new_manifest = make_manifest(actual="ㄐㄧㄠˇ")
            new_manifest["session_id"] = old_manifest["session_id"]
            seal_manifest(new_manifest)
            json_save(output / "校對工作階段.json", new_manifest)
            fake_report = output / "fake.xlsx"
            with patch("standalone_proofread.generate_report", return_value=fake_report):
                imported, skipped, _ = import_gpt_decisions(output, xlsx)
            self.assertEqual((imported, skipped), (0, 1))

    def test_actual_refresh_preserves_expected_events_but_revokes_confirmation(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            manifest = make_manifest()
            entry = manifest["records"][0]
            rid = entry["review_id"]
            oid = entry["occurrence_id"]

            db = normalize_db({})
            db["events"][rid] = {"action": "解決expected證據", "expected_set": "ㄐㄩㄝˊ"}
            json_save(output / "人工判定資料庫.json", db)
            removed = _clear_actual_dependent_events(output, [entry], {oid})
            self.assertEqual(removed, 0)
            kept = json_load(output / "人工判定資料庫.json", {})
            self.assertIn(rid, kept["events"])

            kept["events"][rid] = {"action": "確認現版差異"}
            json_save(output / "人工判定資料庫.json", kept)
            removed = _clear_actual_dependent_events(output, [entry], {oid})
            self.assertEqual(removed, 1)
            after = json_load(output / "人工判定資料庫.json", {})
            self.assertNotIn(rid, after["events"])

    def test_existing_expected_event_replays_on_pass_after_actual_refresh(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            manifest = make_manifest(
                state="PASS",
                actual="ㄐㄩㄝˊ",
                expected_set=["ㄐㄩㄝˊ"],
                expected_evidence="resolver evidence",
            )
            json_save(output / "校對工作階段.json", manifest)
            db = normalize_db({})
            rid = manifest["records"][0]["review_id"]
            db["events"][rid] = {
                "action": "解決expected證據",
                "expected_set": ["ㄐㄩㄝˊ"],
                "expected_evidence": "人工完整詞條",
                "context_evidence": "角色/角@1",
                "source": "人工 expected",
            }
            json_save(output / "人工判定資料庫.json", db)
            from standalone_proofread import materialize_ledger
            ledger = materialize_ledger(manifest, db)
            self.assertEqual(ledger[0]["state"], "PASS")
            self.assertEqual(ledger[0]["expected_evidence"], "人工完整詞條")

    def test_actual_refresh_downgrades_confirmation_to_expected_evidence(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            manifest = make_manifest(
                state="DIFFERENCE_PENDING_CONFIRMATION",
                actual="ㄐㄧㄠˇ",
                expected_set=["ㄐㄩㄝˊ"],
                expected_evidence="完整詞條",
            )
            entry = manifest["records"][0]
            rid = entry["review_id"]
            oid = entry["occurrence_id"]
            db = normalize_db({})
            db["events"][rid] = {
                "action": "確認現版差異",
                "expected_set": ["ㄐㄩㄝˊ"],
                "expected_evidence": "完整詞條",
                "context_evidence": "角色/角@1",
                "source": "人工 GUI 現版六閘門",
                "note": "已確認",
            }
            json_save(output / "人工判定資料庫.json", db)
            removed = _clear_actual_dependent_events(output, [entry], {oid})
            self.assertEqual(removed, 1)
            after = json_load(output / "人工判定資料庫.json", {})
            event = after["events"][rid]
            self.assertEqual(event["action"], "解決expected證據")
            self.assertEqual(event["expected_set"], ["ㄐㄩㄝˊ"])
            self.assertIn("撤銷", event["resolution_reason"])


if __name__ == "__main__":
    unittest.main()
