"""Tiny real sealed/PDF/XLSX transactions, no Tk or textbook fixtures."""
import copy
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

import pdf_portability as p
import standalone_proofread as sp
from review_save_service import ReviewSaveService, StaleReviewProjectError
from review_gui import review_lane, prepare_review_queue, expected_conflict_summary
from tests.test_pdf_portability import pdf, project, fixture_export_pending_for_gpt


@pytest.fixture
def batch(tmp_path, monkeypatch, request):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "user"))
    monkeypatch.setattr(sp, "USER_MEMORY", tmp_path / "user-memory.json")
    monkeypatch.setattr(sp, "REUSABLE_EXPECTED_RULES", tmp_path / "reusable-expected.json")
    left, right = tmp_path / "original-A.pdf", tmp_path / "different-B.pdf"
    pdf(left, title="A")
    pdf(right, title="B")
    source, target = tmp_path / "source", tmp_path / "target"
    boxes = [(10 + i * 30, 20, 25 + i * 30, 45) for i in range(6)]
    sm, sd = project(source, left, session="source-session", boxes=boxes)
    tm, td = project(target, right, session="target-session", boxes=boxes)
    assets = sp.validate_asset_manifest(Path(sp.__file__).parent)
    assert assets["ok"]
    fp = sp.compute_expected_asset_fingerprint(Path(sp.__file__).parent, assets,
        resolver_version=sp.EXPECTED_RESOLVER_VERSION, source_files=sp.EXPECTED_RESOLVER_SOURCE_FILES)
    for root, manifest in ((source, sm), (target, tm)):
        manifest["expected_asset_fingerprint"] = fp["fingerprint"]
        manifest["expected_asset_fingerprint_components"] = fp["components"]
        for entry in manifest["records"]:
            entry["source_record"]["所在行"] = "角色原文"
            entry["source_record"]["局部詞境"] = "角色"
        sp.seal_manifest(manifest)
        sp.json_save(root / "校對工作階段.json", manifest)
    for n, reading in ((1, "ㄐㄩㄝˊ"), (2, "ㄐㄩㄝˋ")):
        entry = tm["records"][n]
        entry.update(expected_set=[reading], expected_evidence="target independent resolver",
                     expected_status="RESOLVED")
        entry.update(sp.refresh_derived_state(entry, preserve_confirmed=False))
    for n, reading in ((3, "ㄐㄩㄝˊ"), (4, "ㄐㄩㄝˋ")):
        td["events"][tm["records"][n]["review_id"]] = {
            "action": "解決expected證據", "expected_set": [reading],
            "expected_evidence": "target independent human", "context_evidence": "target human context"}
    sp.seal_manifest(tm)
    if getattr(request, "param", None) == "no_conditions":
        for manifest in (sm, tm):
            for entry in manifest["records"]:
                entry["source_record"]["所在行"] = ""
                entry["source_record"]["局部詞境"] = ""
            sp.seal_manifest(manifest)
        for entry in sm["records"]:
            entry["context_evidence"] = ""
        sp.seal_manifest(sm)
        sp.json_save(source / "校對工作階段.json", sm)
    sp.json_save(target / "校對工作階段.json", tm)
    sp.json_save(target / "人工判定資料庫.json", td)
    exported = fixture_export_pending_for_gpt(source)
    filled = tmp_path / "proposal.xlsx"
    workbook = load_workbook(exported)
    sheet = workbook["待判定候選"]
    headers = [cell.value for cell in sheet[1]]
    for row in range(2, 7):
        for key, value in {"action": "補建expected證據", "proposed_expected_set": "ㄐㄩㄝˊ",
                           "proposed_expected_evidence": "source independent dictionary",
                           "proposed_context_evidence": "full source position and original context"}.items():
            sheet.cell(row, headers.index(key) + 1, value)
    sheet.cell(7, headers.index("action") + 1, "保留待人工")
    workbook.save(filled)
    workbook.close()
    return source, target, sm, tm, td, filled


def state(root):
    return {str(path.relative_to(root)): path.read_bytes() for path in root.rglob("*")
            if path.is_file() and path.name != ".global_exact_glyph_delivery.lock"}


def test_mixed_batch_preserves_rule_and_human_truth_and_counts(batch):
    source, target, sm, tm, before_db, filled = batch
    source_before = state(source)
    before_ledger = sp.materialize_ledger(tm, before_db)
    result = sp.import_gpt_decisions(target, filled)
    assert result.status == {"imported": 1, "consistent_supplements": 2, "duplicates": 0,
                             "skipped_unoperated": 0, "retained_manual": 1, "new_conflicts": 2,
                             "pending_conflicts": 2, "status": "PENDING_LOCAL_ADJUDICATION"}
    db = sp.json_load_strict(target / "人工判定資料庫.json")
    ledger = sp.materialize_ledger(tm, db)
    assert ledger[1:] == before_ledger[1:]
    assert all(sp.actual_confirmation_snapshot(a) == sp.actual_confirmation_snapshot(b)
               for a, b in zip(before_ledger, ledger))
    assert all(db["events"][rid] == event for rid, event in before_db["events"].items())
    receipt = sp.json_load_strict(target / p.CONFLICT_FILE)
    assert receipt["version"] == 2 and len(receipt["supplements"]) == len(receipt["conflicts"]) == 2
    assert [item["target_basis"]["kind"] for item in receipt["conflicts"]] == ["sealed_resolver_record", "review_event"]
    assert receipt["conflicts"][0]["target_basis"]["event"] is None
    assert "target_event" not in receipt["conflicts"][0]
    assert state(source) == source_before
    assert not (target / p.INCOMPLETE_FILE).exists()
    assert sp.json_load_strict(target / "pipeline_status.json")["completion_gate"]["hard_gates"]["expected_business_conflicts_zero"] is False
    pending = sp.json_load_strict(target / "待人工確認.json")["pending"]
    assert tm["records"][2]["review_id"] in {entry["review_id"] for entry in pending}


def test_adjudicate_retry_reload_retains_local_and_receipt_without_duplicates(batch):
    _, target, _, tm, _, filled = batch
    sp.import_gpt_decisions(target, filled)
    receipt_before = (target / p.CONFLICT_FILE).read_bytes()
    for n in (2, 4):
        db = sp.json_load_strict(target / "人工判定資料庫.json")
        view = p.expected_conflict_review_view(target, tm, db, sp.materialize_ledger(tm, db))
        row = view[n]
        assert review_lane(row) == "expected"
        text = expected_conflict_summary(row)
        assert "source independent dictionary" in text
        assert ("程式規則" if n == 2 else "人工判定") in text
        queue = prepare_review_queue(tm, view, [], 0, set(), set())
        assert row["review_id"] in {item["review_id"] for item in queue["records"]}
        event = sp.build_manual_expected_event(row, operation="ENTER_EXPECTED", expected_set="ㄐㄩㄝˋ",
                                              rationale="fresh local independent context")
        saved = ReviewSaveService(target).save_event(row["review_id"], event)
        assert saved.resolved_entry["expected_set"] == ["ㄐㄩㄝˋ"]
    chosen = sp.json_load_strict(target / "人工判定資料庫.json")
    assert p.validate_conflict_state(target, tm, chosen) == []
    result = sp.import_gpt_decisions(target, filled)
    assert result.status["duplicates"] == 5 and result.status["new_conflicts"] == 0
    assert result.status["consistent_supplements"] == 0 and result.status["pending_conflicts"] == 0
    assert sp.json_load_strict(target / "人工判定資料庫.json") == chosen
    assert (target / p.CONFLICT_FILE).read_bytes() == receipt_before
    reread = ReviewSaveService(target).load_resume_snapshot()
    assert reread.ledger == sp.materialize_ledger(tm, chosen)


def test_pending_retry_preserves_target_events_and_does_not_repeat(batch):
    _, target, _, tm, _, filled = batch
    sp.import_gpt_decisions(target, filled)
    before = (target / p.CONFLICT_FILE).read_bytes()
    db = sp.json_load_strict(target / "人工判定資料庫.json")
    result = sp.import_gpt_decisions(target, filled)
    assert result.status["duplicates"] == 5 and result.status["pending_conflicts"] == 2
    assert (target / p.CONFLICT_FILE).read_bytes() == before
    assert sp.json_load_strict(target / "人工判定資料庫.json") == db


def test_identical_bytes_renamed_and_relocated_are_only_duplicates(batch):
    _, target, _, _, _, filled = batch
    sp.import_gpt_decisions(target, filled)
    receipt_before = (target / p.CONFLICT_FILE).read_bytes()
    before_db = sp.json_load_strict(target / "人工判定資料庫.json")
    moved = filled.parent / "relocated-input" / "another-name.xlsx"
    moved.parent.mkdir()
    moved.write_bytes(filled.read_bytes())
    result = sp.import_gpt_decisions(target, moved)
    assert result.status["duplicates"] == 5
    assert result.status["new_conflicts"] == result.status["consistent_supplements"] == 0
    assert (target / p.CONFLICT_FILE).read_bytes() == receipt_before
    assert sp.json_load_strict(target / "人工判定資料庫.json") == before_db


def test_receipt_only_change_requires_display_reload_before_adjudication(batch):
    _, target, _, tm, _, filled = batch
    sp.import_gpt_decisions(target, filled)
    service = ReviewSaveService(target)
    snapshot = service.load_resume_snapshot()
    view = p.expected_conflict_review_view(target, tm, snapshot.db, snapshot.ledger)
    row = view[2]
    event = sp.build_manual_expected_event(row, operation="ENTER_EXPECTED", expected_set="ㄐㄩㄝˋ")
    workbook = load_workbook(filled)
    workbook.properties.title = "another independent saved file"
    second = filled.with_name("new-source.xlsx")
    workbook.save(second)
    workbook.close()
    sp.import_gpt_decisions(target, second)
    assert sp.json_load_strict(target / "人工判定資料庫.json") == snapshot.db
    before = state(target)
    with pytest.raises(StaleReviewProjectError, match="衝突來源已變更"):
        service.save_event(row["review_id"], event)
    assert state(target) == before
    service.load_resume_snapshot()
    saved = service.save_event(row["review_id"], event)
    assert row["review_id"] not in p.validate_conflict_state(target, tm, saved.db)


def test_receipt_changed_during_replay_stops_before_db_write(batch, monkeypatch):
    _, target, _, tm, _, filled = batch
    sp.import_gpt_decisions(target, filled)
    service = ReviewSaveService(target)
    snapshot = service.load_resume_snapshot()
    view = p.expected_conflict_review_view(target, tm, snapshot.db, snapshot.ledger)
    event = sp.build_manual_expected_event(view[2], operation="ENTER_EXPECTED", expected_set="ㄐㄩㄝˋ")
    original_capture = service._capture
    calls = 0
    def capture(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:
            path = target / p.CONFLICT_FILE
            receipt = sp.json_load_strict(path)
            receipt["integrity_sha256"] = "0" * 64
            sp.json_save(path, receipt)
        return original_capture(*args, **kwargs)
    monkeypatch.setattr(service, "_capture", capture)
    with pytest.raises(StaleReviewProjectError, match="驗證期間"):
        service.save_event(view[2]["review_id"], event)
    assert sp.json_load_strict(target / "人工判定資料庫.json") == snapshot.db


def test_actual_only_new_seal_preserves_expected_receipt_applicability(batch):
    _, target, _, tm, _, filled = batch
    sp.import_gpt_decisions(target, filled)
    db = sp.json_load_strict(target / "人工判定資料庫.json")
    original_seal = tm["manifest_integrity_sha256"]
    for entry in tm["records"]:
        entry["actual_evidence"] += " independent actual-only refinement"
    sp.seal_manifest(tm)
    assert tm["manifest_integrity_sha256"] != original_seal
    sp.json_save(target / "校對工作階段.json", tm)
    assert len(set(p.validate_conflict_state(target, tm, db))) == 2
    receipt = sp.json_load_strict(target / p.CONFLICT_FILE)
    assert all(item["target_basis"]["manifest_integrity_sha256"] == original_seal
               for item in [*receipt["conflicts"], *receipt["supplements"]])


def test_same_session_complete_proof_uses_business_classifier(batch):
    source, _, sm, _, _, filled = batch
    entry = sm["records"][0]
    entry.update(expected_set=["ㄐㄩㄝˊ"], expected_evidence="same-session independent rule",
                 expected_status="RESOLVED")
    entry.update(sp.refresh_derived_state(entry, preserve_confirmed=False))
    sp.seal_manifest(sm)
    sp.json_save(source / "校對工作階段.json", sm)
    result = sp.import_gpt_decisions(source, filled)
    assert result.status["consistent_supplements"] == 1
    assert result.status["imported"] == 4 and result.status["new_conflicts"] == 0
    assert sp.materialize_ledger(sm, sp.json_load_strict(source / "人工判定資料庫.json"))[0] == entry


def test_formal_pass_conflict_is_actionable_without_changing_truth(batch):
    _, target, _, tm, _, filled = batch
    workbook = load_workbook(filled)
    sheet = workbook["待判定候選"]
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(3, headers.index("proposed_expected_set") + 1, "ㄐㄩㄝˋ")
    workbook.save(filled)
    workbook.close()
    sp.import_gpt_decisions(target, filled)
    db = sp.json_load_strict(target / "人工判定資料庫.json")
    view = p.expected_conflict_review_view(target, tm, db, sp.materialize_ledger(tm, db))
    entry = view[1]
    assert entry["state"] == "PASS" and entry["expected_set"] == ["ㄐㄩㄝˊ"]
    assert review_lane(entry) == "expected"
    event = sp.build_manual_expected_event(entry, operation="ENTER_EXPECTED", expected_set="ㄐㄩㄝˊ")
    result = ReviewSaveService(target).save_event(entry["review_id"], event)
    assert entry["review_id"] not in p.validate_conflict_state(target, tm, result.db)


@pytest.mark.parametrize("batch", ["no_conditions"], indirect=True)
def test_missing_required_context_remains_fail_closed(batch):
    _, target, _, _, _, filled = batch
    before = state(target)
    with pytest.raises(ValueError, match="target|baseline|fingerprint|語境"):
        sp.import_gpt_decisions(target, filled)
    assert state(target) == before


@pytest.mark.parametrize("fault", ["identity", "context", "invalid_event", "fingerprint", "seal"])
def test_invalid_source_is_whole_batch_fail_closed(batch, fault):
    _, target, _, _, _, filled = batch
    workbook = load_workbook(filled)
    sheet = workbook["待判定候選"]
    headers = [cell.value for cell in sheet[1]]
    if fault in {"identity", "context", "invalid_event"}:
        key = {"identity": "occurrence_id", "context": "review_snapshot", "invalid_event": "proposed_expected_set"}[fault]
        sheet.cell(5, headers.index(key) + 1, "invalid")
    elif fault == "fingerprint":
        sheet = workbook["匯入中繼資料"]
        matching_rows = [row for row in sheet.iter_rows() if row[0].value == "expected_asset_fingerprint"]
        assert len(matching_rows) == 1
        matching_rows[0][1].value = "invalid"
    else:
        workbook[p.EXCEL_PROOF_SHEET].cell(2, 2, "invalid")
    workbook.save(filled)
    workbook.close()
    before = state(target)
    with pytest.raises(ValueError):
        sp.import_gpt_decisions(target, filled)
    assert state(target) == before


@pytest.mark.parametrize("phase", ["receipt", "db", "publication"])
def test_write_failure_restores_whole_original_state(batch, monkeypatch, phase):
    _, target, _, _, _, filled = batch
    before = state(target)
    if phase == "publication":
        def fail(*args, **kwargs):
            raise OSError("publication disk full")
        monkeypatch.setattr(p, "_publish_portable_outputs", fail)
    else:
        original = sp.json_save
        wanted = target / (p.CONFLICT_FILE if phase == "receipt" else "人工判定資料庫.json")
        def fail(path, *args, **kwargs):
            if Path(path) == wanted:
                raise OSError("transaction disk full")
            return original(path, *args, **kwargs)
        monkeypatch.setattr(sp, "json_save", fail)
    with pytest.raises(OSError, match="disk full"):
        sp.import_gpt_decisions(target, filled)
    assert state(target) == before


@pytest.mark.parametrize("change", ["integrity", "version", "basis", "identity"])
def test_receipt_corruption_blocks_reimport_without_writes(batch, change):
    _, target, _, _, _, filled = batch
    sp.import_gpt_decisions(target, filled)
    path = target / p.CONFLICT_FILE
    receipt = sp.json_load_strict(path)
    if change == "integrity":
        receipt["integrity_sha256"] = "0" * 64
    elif change == "version":
        receipt["version"] = 99
    else:
        if change == "basis":
            receipt["conflicts"][0]["target_basis"]["kind"] = "invented_human"
        else:
            receipt["conflicts"][0]["target_occurrence_id"] = "wrong"
        receipt = p._sealed_expected_receipt(receipt["conflicts"], receipt["supplements"])
    sp.json_save(path, receipt)
    before = state(target)
    with pytest.raises(ValueError):
        sp.import_gpt_decisions(target, filled)
    assert state(target) == before
