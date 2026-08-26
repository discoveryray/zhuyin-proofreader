from __future__ import annotations

import csv
import tempfile
import unittest
from pathlib import Path

import fitz
from openpyxl import load_workbook

from actual_review import (
    USER_GLYF_FILE,
    USER_CFF_FILE,
    build_actual_review_groups,
    build_actual_group_for_entry,
    apply_verified_actual_group,
    export_actual_review_package,
    import_actual_review_workbook,
    ensure_user_evidence_files,
    dynamic_actual_hashes,
    load_user_verified_cff,
)


SHA = "a" * 64


def make_pdf(path: Path):
    doc = fitz.open()
    page = doc.new_page(width=300, height=300)
    page.insert_text((60, 100), "ABC", fontsize=24)
    doc.save(path)
    doc.close()


def entry(oid: str, rid: str, pdf: Path, x: float, *, state="ACTUAL_DECODE_ERROR", source=None):
    src = {
        "TTF字形SHA256": SHA,
        "穩定注音鍵": f"key-{oid}",
    }
    if source:
        src.update(source)
    return {
        "occurrence_id": oid,
        "review_id": rid,
        "state": state,
        "pdf": str(pdf),
        "pdf_name": pdf.name,
        "physical_page": 1,
        "printed_page": "1",
        "char": "字",
        "stable_key": f"key-{oid}",
        "x0": x,
        "y0": 70,
        "x1": x + 20,
        "y1": 110,
        "actual": "",
        "actual_evidence": "decoder unresolved",
        "source_record": src,
    }


class ActualReviewV540Tests(unittest.TestCase):
    def test_exact_ttf_groups_and_promotes_only_with_two_examples(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            ensure_user_evidence_files(root)
            pdf = root / "book.pdf"
            make_pdf(pdf)
            a = entry("o1", "r1", pdf, 60)
            b = entry("o2", "r2", pdf, 120)
            group = build_actual_review_groups([a, b])[0]
            self.assertEqual(group["kind"], "TTF_GLYF_SHA256")
            self.assertEqual(group["occurrence_count"], 2)

            first = apply_verified_actual_group(root, group, "ㄒㄧ˙", checked_occurrence_ids=["o1"], source="test")
            self.assertEqual(first["reading"], "˙ㄒㄧ")
            self.assertEqual(first["learning_level"], "USER_VERIFIED_SINGLE")
            self.assertFalse(first["propagated_to_group"])

            second = apply_verified_actual_group(root, group, "˙ㄒㄧ", checked_occurrence_ids=["o1", "o2"], source="test")
            self.assertEqual(second["learning_level"], "VERIFIED_EXACT_GLYPH")
            self.assertTrue(second["propagated_to_group"])
            with (root / USER_GLYF_FILE).open(encoding="utf-8-sig", newline="") as fh:
                rows = list(csv.DictReader(fh))
            self.assertEqual(rows[0]["source_count"], "2")
            self.assertEqual(rows[0]["bopomofo"], "˙ㄒㄧ")

    def test_user_flagged_non_actual_state_can_build_group(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            pdf = root / "book.pdf"
            make_pdf(pdf)
            a = entry("o1", "r1", pdf, 60, state="DIFFERENCE_PENDING_CONFIRMATION")
            b = entry("o2", "r2", pdf, 120, state="PASS")
            group = build_actual_group_for_entry([a, b], a)
            self.assertEqual(group["occurrence_count"], 2)
            self.assertEqual(group["kind"], "TTF_GLYF_SHA256")

    def test_cff_learning_uses_full_glyph_sha_not_annotation_signature(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            pdf = root / "book.pdf"
            make_pdf(pdf)
            shared_annotation = {
                "TTF字形SHA256": "",
                "CFF樣式群組": "KAICHU_MD",
                "CFF符號簽名": "body1;body2",
                "CFF聲調簽名": "same-tone",
                "CFF輕聲簽名": "",
                "CFF完整注音簽名": "same-annotation-full-signature",
            }
            a = entry("o1", "r1", pdf, 60, source={**shared_annotation, "CFF整字字形SHA256": "b" * 64})
            b = entry("o2", "r2", pdf, 120, source={**shared_annotation, "CFF整字字形SHA256": "c" * 64})
            groups = build_actual_review_groups([a, b])
            self.assertEqual(len(groups), 2, "same annotation signature must not merge different complete CFF glyphs")
            self.assertTrue(all(g["kind"] == "CFF_GLYPH_SHA256" for g in groups))

    def test_cff_same_full_glyph_sha_can_group_for_double_visual_confirmation(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            pdf = root / "book.pdf"
            make_pdf(pdf)
            src = {
                "TTF字形SHA256": "",
                "CFF樣式群組": "KAICHU_MD",
                "CFF符號簽名": "body1;body2",
                "CFF聲調簽名": "tone",
                "CFF整字字形SHA256": "d" * 64,
            }
            a = entry("o1", "r1", pdf, 60, source=src)
            b = entry("o2", "r2", pdf, 120, source=src)
            groups = build_actual_review_groups([a, b])
            self.assertEqual(len(groups), 1)
            self.assertEqual(groups[0]["kind"], "CFF_GLYPH_SHA256")
            self.assertEqual(groups[0]["occurrence_count"], 2)

    def test_gpt_package_has_no_expected_columns_and_imports_transactionally(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            ensure_user_evidence_files(root)
            pdf = root / "book.pdf"
            make_pdf(pdf)
            a = entry("o1", "r1", pdf, 60)
            b = entry("o2", "r2", pdf, 120)
            ledger = [a, b]
            meta = {
                "version": "5.4.0",
                "session_id": "sess",
                "session_schema_version": "2.5.0",
                "workbook_schema_version": "2.5.0",
                "review_id_schema_version": "2.5.0",
            }
            zip_path = export_actual_review_package(root, ledger, **meta)
            self.assertTrue(zip_path.exists())
            xlsx = root / "actual待判定_GPT包" / "actual待判定_給GPT.xlsx"
            wb = load_workbook(xlsx)
            ws = wb["actual待判定"]
            headers = [c.value for c in ws[1]]
            self.assertFalse(any("expected" in str(h).lower() or "預期" in str(h) or "應標" in str(h) for h in headers))
            self.assertFalse(any(str(h).endswith("_char") for h in headers), "GPT actual sheet must not expose target Han character columns")
            self.assertIn("annotation", str(ws.cell(2, headers.index("sample_a_image") + 1).value))
            self.assertNotIn("context", str(ws.cell(2, headers.index("sample_a_image") + 1).value))
            ix = {h: i + 1 for i, h in enumerate(headers)}
            ws.cell(2, ix["decision"], "VERIFIED")
            ws.cell(2, ix["actual_reading"], "ㄒㄧ˙")
            ws.cell(2, ix["confidence"], "高")
            ws.cell(2, ix["sample_a_checked"], "Y")
            ws.cell(2, ix["sample_b_checked"], "Y")
            wb.save(xlsx)
            wb.close()

            groups = build_actual_review_groups(ledger)
            result = import_actual_review_workbook(root, xlsx, groups, expected_metadata=meta)
            self.assertEqual(result["imported_groups"], 1)
            with (root / USER_GLYF_FILE).open(encoding="utf-8-sig", newline="") as fh:
                rows = list(csv.DictReader(fh))
            self.assertEqual(rows[0]["verification_level"], "VERIFIED_EXACT_GLYPH")

    def test_cff_verified_truth_is_keyed_by_full_glyph_sha(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            ensure_user_evidence_files(root)
            p = root / USER_CFF_FILE
            p.write_text(
                "style_group,glyph_sha256,full_signature,bopomofo,verification_level,source_count,source_examples,updated_at,notes\n"
                + f"KAICHU_MD,{('e'*64)},shared-annotation,ㄋㄚˇ,VERIFIED_EXACT_GLYPH,2,o1|o2,2026-08-20,test\n",
                encoding="utf-8-sig",
            )
            loaded = load_user_verified_cff(p)
            self.assertIn(("KAICHU_MD", "e" * 64), loaded)
            self.assertEqual(loaded[("KAICHU_MD", "e" * 64)]["bopomofo"], "ㄋㄚˇ")
            self.assertNotIn(("KAICHU_MD", "shared-annotation"), loaded)

    def test_v540_migrates_only_the_known_bad_p122_override(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            override = root / "manual_actual_occurrence_overrides.csv"
            override.write_text(
                "pdf_contains,pdf_excludes,page,target_char,stable_key,x0,y0,actual_reading,source,note\n"
                "15_國小健體3下課本_單元5第2課_(學),,122,那,CFF:DFKaiChuIn-Md-BPMF-BF#1311,142.01,76.595,ㄋㄚˋ,使用者原頁回標＋GPT原頁視覺覆核（2026-08-20；出現位置限定）,bad\n"
                "other,,1,字,key,1,2,ㄗ,人工自訂,keep\n",
                encoding="utf-8-sig",
            )
            ensure_user_evidence_files(root)
            with override.open(encoding="utf-8-sig", newline="") as fh:
                rows = list(csv.DictReader(fh))
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["pdf_contains"], "other")

    def test_dynamic_hash_changes_when_user_truth_changes(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            ensure_user_evidence_files(root)
            first = dynamic_actual_hashes(root)
            p = root / USER_GLYF_FILE
            p.write_text(
                "glyph_sha256,bopomofo,verification_level,source_count,source_examples,updated_at,notes\n"
                + f"{SHA},ㄒㄧ,USER_VERIFIED_SINGLE,1,o1,2026-08-20,test\n",
                encoding="utf-8-sig",
            )
            second = dynamic_actual_hashes(root)
            self.assertNotEqual(first[USER_GLYF_FILE], second[USER_GLYF_FILE])


if __name__ == "__main__":
    unittest.main()
