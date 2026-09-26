"""Real pipeline orchestration with a synthetic, deterministic actual decoder.

The PDF rendering, candidate resolver, artifact seals, GUI preview, save
service, and portability adapter are real. Only the glyph decoder's output is
injected: this is synthetic acceptance, not real textbook acceptance.
"""

import base64
import copy
import hashlib
import json
import sys
import tkinter as tk
from threading import Barrier, Event, Thread, current_thread
from pathlib import Path
from unittest.mock import patch

import fitz
import pytest
from openpyxl import Workbook, load_workbook

import pdf_portability as portability
import review_gui
import standalone_proofread as sp
import occurrence_ledger as ol
import actual_review as ar
from review_display import occurrence_preview
from review_save_service import ReviewSaveService
from tests.test_project_repair_integration_v562 import ACTUAL_COLUMNS, EXCLUDED_COLUMNS
from tests.test_pdf_portability import pdf as small_pdf, project as small_project


def _pdfs(folder):
    first, second = folder / "download-a.pdf", folder / "download-b.pdf"
    with fitz.open() as document:
        page = document.new_page(width=400, height=220)
        page.insert_text((50, 90), "角角", fontname="china-t", fontsize=20)
        page.insert_text((80, 76), "ㄐㄩㄝˊ", fontname="china-t", fontsize=8)
        page.insert_text((114, 76), "ㄐㄩㄝˊ", fontname="china-t", fontsize=8)
        document.set_metadata({"title": "download A"})
        document.save(first, deflate=False, garbage=0)
        document.set_metadata({"title": "download B"})
        document.save(second, deflate=True, garbage=4)
    assert hashlib.sha256(first.read_bytes()).digest() != hashlib.sha256(second.read_bytes()).digest()
    assert portability._page_signatures(first) == portability._page_signatures(second)
    return first, second


def _synthetic_decode(pdf, output, *args, global_snapshot=None):
    """Emit actual rows from visible PDF char boxes, without expected input."""
    fingerprint = args[-2]
    digest = sp.sha256_file(pdf)
    with fitz.open(pdf) as document:
        chars = [char for block in document[0].get_text("rawdict")["blocks"]
                 for line in block.get("lines", []) for span in line.get("spans", [])
                 for char in span.get("chars", []) if char.get("c") == "角"]
    assert len(chars) == 2
    rows = []
    for index, char in enumerate(chars, 1):
        x0, y0, x1, y1 = char["bbox"]
        rows.append({
            "實體頁碼": 1, "課本頁": 1, "頁面標籤": "1", "字元": "角",
            "實際注音": "ㄐㄩㄝˊ", "解碼狀態": "已映射",
            "解碼依據": "synthetic independent PDF glyph fixture",
            "穩定注音鍵": f"fixture-{index}", "群組注音鍵": f"fixture-{index}",
            "font": "synthetic-font", "font_xref": 1, "注音元件ID": index,
            "glyph_id_字形索引": index, "x0": x0, "y0": y0, "x1": x1, "y1": y1,
            "source_row_number": index + 1,
            "ledger_schema_version": ol.LEDGER_SCHEMA_VERSION,
            "workbook_schema_version": ol.WORKBOOK_SCHEMA_VERSION,
        })
    ol.prepare_occurrence_rows(digest, rows)
    workbook = Workbook()
    ws = workbook.active
    ws.title = "實際注音"
    columns = [*ACTUAL_COLUMNS, "注音結構偵測來源", "CFF符號簽名", "CFF聲調"]
    ws.append(columns)
    for row in rows:
        ws.append([row.get(column, "") for column in columns])
    excluded = workbook.create_sheet("結構偵測排除")
    excluded.append(EXCLUDED_COLUMNS)
    metadata = workbook.create_sheet("v5.2中繼資料")
    metadata.append(["項目", "內容"])
    metadata.append(["workbook_schema_version", ol.WORKBOOK_SCHEMA_VERSION])
    metadata.append(["ledger_schema_version", ol.LEDGER_SCHEMA_VERSION])
    metadata.append(["pdf_sha256", digest])
    metadata.append(["actual_asset_fingerprint", fingerprint["fingerprint"]])
    metadata.append(["actual_asset_fingerprint_components", json.dumps(fingerprint["components"], ensure_ascii=False, sort_keys=True)])
    summary = workbook.create_sheet("摘要")
    summary.append(["項目", "內容"])
    summary.append(["偵測到注音字形筆數", len(rows)])
    workbook.save(output)
    workbook.close()


def _synthetic_unresolved_decode(pdf, output, *args, global_snapshot=None):
    _synthetic_decode(pdf, output, *args, global_snapshot=global_snapshot)
    override_rows = ar._read_csv(sp.project_actual_evidence_root(Path(output).parent.parent)
                                 / ar.OCCURRENCE_OVERRIDE_FILE, ar.OVERRIDE_HEADERS)
    workbook = load_workbook(output)
    sheet = workbook["實際注音"]
    headers = [cell.value for cell in sheet[1]]
    for row in range(2, sheet.max_row + 1):
        stable_key = str(sheet.cell(row, headers.index("穩定注音鍵") + 1).value or "")
        matched = next((item for item in override_rows if item["pdf_contains"] == Path(pdf).stem
                        and item["stable_key"] == stable_key), None)
        sheet.cell(row, headers.index("實際注音") + 1,
                   matched["actual_reading"] if matched else "")
        sheet.cell(row, headers.index("解碼狀態") + 1,
                   "已映射" if matched else "未辨識")
        sheet.cell(row, headers.index("解碼依據") + 1, "synthetic unresolved glyph")
    workbook.save(output)
    workbook.close()


def test_full_pipeline_different_sha_continues_and_returns(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = _pdfs(tmp_path)
    a, b = tmp_path / "project-a", tmp_path / "project-b"
    with patch.object(sp, "decode", side_effect=_synthetic_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], a, defer_excel_reports=True)
        a_manifest = sp.json_load_strict(a / "校對工作階段.json")
        assert len(a_manifest["records"]) == 2
        first_event = {"action": "確認非校對範圍", "exclusion_reason": "synthetic fixture",
                       "exclusion_evidence": "PDF page 1 visual position 1"}
        ReviewSaveService(a).save_event(a_manifest["records"][0]["review_id"], first_event)
        portability.prepare_portable_project(a)
        first.rename(tmp_path / "source-pdf-not-on-b.pdf")
        result = portability.continue_project(a, second, b)
        assert result["imported"] == 1
        b_manifest = sp.json_load_strict(b / "校對工作階段.json")
        assert len(b_manifest["records"]) == 2
        assert len(sp.json_load_strict(b / "待人工確認.json")["pending"]) == 1
        pending_workbook = load_workbook(b / "待判定候選_給GPT.xlsx", read_only=True)
        try:
            assert portability.EXCEL_PROOF_SHEET in pending_workbook.sheetnames
        finally:
            pending_workbook.close()
        b_status = sp.json_load_strict(b / "pipeline_status.json")
        assert b_status["user_report"] == str(b / "注音校對_最終報告.xlsx")
        workbook = load_workbook(b / "注音校對_最終報告.xlsx", read_only=True)
        try:
            assert any(row[0] == "人工判定" for row in workbook["人工與公司規則"].values)
        finally:
            workbook.close()
        pixels, box, _ = occurrence_preview(b_manifest["records"][1], output_dir=b)
        assert pixels.width > 0 and box is not None
        second_event = {"action": "確認非校對範圍", "exclusion_reason": "synthetic fixture",
                        "exclusion_evidence": "PDF page 1 visual position 2"}
        ReviewSaveService(b).save_event(b_manifest["records"][1]["review_id"], second_event)
        portability.prepare_portable_project(b)
        (tmp_path / "source-pdf-not-on-b.pdf").rename(first)
        second.rename(tmp_path / "target-pdf-not-on-a.pdf")
        back = portability.import_project_decisions(b, a)
        assert back["imported"] == 1 and back["duplicates"] == 1
        assert len(sp.json_load_strict(a / "人工判定資料庫.json")["events"]) == 2
        assert sp.json_load_strict(a / "待人工確認.json")["pending"] == []
        assert (a / "注音校對_最終報告.xlsx").is_file()
        assert sp.json_load_strict(a / "pipeline_status.json")["user_report"] == str(a / "注音校對_最終報告.xlsx")
        assert len(sp.materialize_ledger(a_manifest, sp.json_load_strict(a / "人工判定資料庫.json"))) == 2
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


def test_actual_excel_conflict_requires_fresh_local_visual_adjudication(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = _pdfs(tmp_path)
    a, b = tmp_path / "source", tmp_path / "target"
    with patch.object(sp, "decode", side_effect=_synthetic_unresolved_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], a, defer_excel_reports=True)
        sp.run_pipeline_pdfs([second], b, defer_excel_reports=True)
        sp.export_actual_pending_for_gpt(a)
        workbook = load_workbook(a / "actual待判定_GPT包" / "actual待判定_給GPT.xlsx")
        sheet = workbook["actual待判定"]
        headers = [cell.value for cell in sheet[1]]
        for key, value in {"decision": "VERIFIED", "actual_reading": "ㄐㄩㄝˊ",
                           "confidence": "高", "sample_a_checked": "Y"}.items():
            sheet.cell(2, headers.index(key) + 1, value)
        first_excel = tmp_path / "first-actual.xlsx"
        workbook.save(first_excel)
        sheet.cell(2, headers.index("actual_reading") + 1, "ㄐㄩㄝˇ")
        conflict_excel = tmp_path / "conflict-actual.xlsx"
        workbook.save(conflict_excel)
        sheet.cell(2, headers.index("actual_reading") + 1, "ㄐㄩㄝˋ")
        later_conflict_excel = tmp_path / "later-conflict-actual.xlsx"
        workbook.save(later_conflict_excel)
        workbook.close()
        assert sp.import_actual_gpt_decisions(b, first_excel)[0] == 1
        report_read = Event()
        release_report = Event()
        import_started = Event()
        import_done = Event()
        race_errors = []
        original_hold = portability.hold_actual_excel_completion

        def pause_after_report_gate(*args):
            result = original_hold(*args)
            if current_thread().name == "actual-conflict-report":
                report_read.set()
                assert release_report.wait(20)
            return result

        def report_worker():
            try:
                sp.regenerate_report(b)
            except Exception as exc:
                race_errors.append(("report", exc))

        def import_worker():
            import_started.set()
            try:
                sp.import_actual_gpt_decisions(b, conflict_excel)
            except ValueError as exc:
                if "不同判定" not in str(exc):
                    race_errors.append(("import", exc))
            else:
                race_errors.append(("import", AssertionError("conflicting import succeeded")))
            finally:
                import_done.set()

        with patch.object(portability, "hold_actual_excel_completion", side_effect=pause_after_report_gate):
            report_thread = Thread(target=report_worker, name="actual-conflict-report")
            report_thread.start()
            try:
                assert report_read.wait(20)
                import_thread = Thread(target=import_worker, name="actual-conflict-import")
                import_thread.start()
                assert import_started.wait(5)
                assert not import_done.wait(0.3), "衝突匯入在報告發布鎖外寫入"
            finally:
                release_report.set()
                report_thread.join(20)
                if "import_thread" in locals():
                    import_thread.join(20)
        assert not report_thread.is_alive() and not import_thread.is_alive()
        assert not race_errors
        gate = sp.json_load_strict(b / "pipeline_status.json")["completion_gate"]
        assert gate["hard_gates"]["actual_excel_conflicts_zero"] is False
        receipt_path = b / portability.ACTUAL_EXCEL_CONFLICT_FILE
        receipt_bytes = receipt_path.read_bytes()
        blocked_paths = [receipt_path, b / "人工判定資料庫.json",
                         b / "pipeline_status.json", b / "注音校對_最終報告.xlsx"]
        blocked_bytes = {path: path.read_bytes() if path.exists() else None for path in blocked_paths}
        for operation in ("--report-only", "--repair-project"):
            with patch.object(sys, "argv", ["standalone_proofread.py", "-o", str(b), operation]):
                with pytest.raises(ValueError, match="actual.*衝突"):
                    sp.main()
            assert {path: path.read_bytes() if path.exists() else None
                    for path in blocked_paths} == blocked_bytes
        with pytest.raises(ValueError, match="actual.*衝突"):
            sp.export_actual_pending_for_gpt(b)
        assert {path: path.read_bytes() if path.exists() else None
                for path in blocked_paths} == blocked_bytes
        from standalone_gui import App as ProjectApp
        report_gui = ProjectApp.__new__(ProjectApp)
        report_gui.output = type("SelectedProject", (), {"get": lambda self: str(b)})()
        with patch("standalone_gui.messagebox.showerror") as report_error, \
             patch("standalone_gui.os.startfile", create=True) as startfile, \
             patch("standalone_gui.subprocess.Popen") as open_other:
            report_gui.open_user_report()
            report_gui.open_technical_report()
            assert report_error.call_count == 2
            startfile.assert_not_called()
            open_other.assert_not_called()
        manifest = sp.json_load_strict(b / "校對工作階段.json")
        db = sp.json_load_strict(b / "人工判定資料庫.json")
        pending = portability.actual_excel_conflict_state(b, manifest, db)
        assert len(pending) == 1
        review_id = pending[0]
        ledger = sp.materialize_ledger(manifest, db)
        queue = review_gui.prepare_review_queue(manifest, ledger, [], 0, set(), set(),
                                                actual_conflict_review_ids=pending)
        assert review_id in [item["review_id"] for item in queue["records"]]
        window_root = tk.Tk()
        window_root.withdraw()
        try:
            window = tk.Toplevel(window_root)
            try:
                app = review_gui.ReviewApp(window, b)
                app.index = next(index for index, item in enumerate(app.records)
                                 if item["review_id"] == review_id)
                app.show()
                window.update()
                assert app.primary.cget("state") == "normal"
                assert "actual 衝突" in app.primary.cget("text")
                assert app.image.photo is not None
            finally:
                window.destroy()
        finally:
            window_root.destroy()
        watched = [b / name for name in ("校對工作階段.json", "人工判定資料庫.json",
                                          "注音校對_最終報告.xlsx", "pipeline_status.json",
                                          portability.INCOMPLETE_FILE)]
        before = {path: path.read_bytes() if path.exists() else None for path in watched}
        before_override = (sp.project_actual_evidence_root(b) / ar.OCCURRENCE_OVERRIDE_FILE).read_bytes()
        with patch.object(sp, "stage_manual_actual_group", side_effect=OSError("staging disk full")):
            with pytest.raises(OSError, match="staging disk full"):
                sp.stage_manual_actual_correction(b, review_id, "ㄐㄩㄝˊ")
        assert receipt_path.read_bytes() == receipt_bytes
        assert (sp.project_actual_evidence_root(b) / ar.OCCURRENCE_OVERRIDE_FILE).read_bytes() == before_override
        assert {path: path.read_bytes() if path.exists() else None for path in watched} == before
        tampered = sp.json_load_strict(receipt_path)
        tampered["conflicts"][0]["target_review_id"] = "foreign-review"
        sp.json_save(receipt_path, tampered)
        with pytest.raises(ValueError, match="完整性"):
            portability._load_project(b)
        receipt_path.write_bytes(receipt_bytes)
        for mutation in ("target_session", "source_occurrence", "source_review",
                         "source_decision", "source_pdf_sha", "duplicate"):
            altered = sp.json_load_strict(receipt_path)
            item = altered["conflicts"][0]
            source = item["source_record"]
            if mutation == "target_session":
                source["target_session_id"] = "foreign-session"
            elif mutation == "source_occurrence":
                source["source_occurrence_ids"] = ["foreign-occurrence"]
            elif mutation == "source_review":
                source["source_review_ids"] = ["foreign-review"]
            elif mutation == "source_decision":
                source["source_decision"]["actual_reading"] = "ㄐㄩㄝˋ"
            elif mutation == "source_pdf_sha":
                source["source_pdf_sha256"] = ["0" * 64]
            else:
                altered["conflicts"].append(copy.deepcopy(item))
            altered = portability._sealed_actual_excel_conflicts(altered["conflicts"])
            sp.json_save(receipt_path, altered)
            with pytest.raises(ValueError, match="衝突.*識別|衝突.*格式|衝突.*完整性"):
                portability._load_project(b)
            receipt_path.write_bytes(receipt_bytes)
        sp.stage_manual_actual_correction(b, review_id, "ㄐㄩㄝˊ", note="原頁重新核對")
        staged = ar.load_manual_actual_staging(sp.project_actual_evidence_root(b))["staged_groups"]
        assert len(staged) == 1 and "portable_actual_conflict:" in staged[0]["note"]
        old_report = (b / "注音校對_最終報告.xlsx").read_bytes()
        with patch.object(sp, "regenerate_report", side_effect=OSError("report disk full")):
            with pytest.raises(sp.ManualActualPostApplyError, match="report disk full"):
                sp.apply_staged_manual_actual_corrections(b)
        assert (b / "注音校對_最終報告.xlsx").read_bytes() == old_report
        assert sp.json_load_strict(b / "pipeline_status.json")["excel_report_deferred"] is True
        assert receipt_path.read_bytes() == receipt_bytes
        reopened_manifest, _ = portability._load_project(b)
        assert reopened_manifest["session_id"] == manifest["session_id"]
        assert portability.actual_excel_conflict_state(
            b, reopened_manifest, sp.json_load_strict(b / "人工判定資料庫.json")) == []
        assert "報告尚待更新" in __import__("standalone_gui").project_status_details(b)[0]
        with patch("standalone_gui.messagebox.showerror") as report_error, \
             patch("standalone_gui.os.startfile", create=True) as startfile, \
             patch("standalone_gui.subprocess.Popen") as open_other:
            report_gui.open_user_report()
            assert report_error.call_count == 1
            startfile.assert_not_called()
            open_other.assert_not_called()
        assert sp.regenerate_report(b).is_file()
        assert "actual 同位置判定衝突尚待" not in __import__("standalone_gui").project_status_details(b)[0]
        with patch("standalone_gui.messagebox.showerror") as report_error, \
             patch("standalone_gui.os.startfile", create=True) as startfile, \
             patch("standalone_gui.subprocess.Popen") as open_other:
            report_gui.open_user_report()
            report_error.assert_not_called()
            assert startfile.called or open_other.called
        actual_root = sp.project_actual_evidence_root(b)
        override_path = actual_root / ar.OCCURRENCE_OVERRIDE_FILE
        reviewed_entry = next(item for item in sp.json_load_strict(b / "校對工作階段.json")["records"]
                              if item["review_id"] == review_id)
        override_fields = ("pdf_contains", "pdf_excludes", "page", "target_char",
                           "stable_key", "x0", "y0")
        reviewed_key = ar._override_key_from_entry(reviewed_entry)

        def reviewed_override():
            return next(row for row in ar._read_csv(override_path, ar.OVERRIDE_HEADERS)
                        if tuple(row[field] for field in override_fields) == reviewed_key)

        adjudicated = reviewed_override()
        assert adjudicated["source"] == "人工 GUI actual 視覺確認"
        assert "portable_actual_conflict:" in adjudicated["note"]
        conflict_record = sp.json_load_strict(receipt_path)["conflicts"][0]
        assert hashlib.sha256(first_excel.read_bytes()).hexdigest() in conflict_record["target_record"]["note"]
        assert conflict_record["source_record"]["source_excel_sha256"] == hashlib.sha256(
            conflict_excel.read_bytes()).hexdigest()
        repeat_paths = [receipt_path, override_path, actual_root / ar.GLYPH_PROVENANCE_FILE,
                        b / "校對工作階段.json", b / "人工判定資料庫.json",
                        b / "pipeline_status.json", b / "注音校對_最終報告.xlsx"]
        before_repeat = {path: path.read_bytes() for path in repeat_paths}
        repeated = sp.import_actual_gpt_decisions(b, first_excel)
        assert repeated[0] == 0 and repeated.status["project_actual_commit"] == "NO_CHANGES"
        assert {path: path.read_bytes() for path in repeat_paths} == before_repeat
        assert portability.actual_excel_conflict_state(
            b, sp.json_load_strict(b / "校對工作階段.json"),
            sp.json_load_strict(b / "人工判定資料庫.json")) == []
        same_reading_workbook = load_workbook(first_excel)
        same_reading_sheet = same_reading_workbook["actual待判定"]
        same_reading_sheet.cell(2, headers.index("note") + 1, "另一份同讀音來源")
        new_same_reading = tmp_path / "new-same-reading.xlsx"
        same_reading_workbook.save(new_same_reading)
        same_reading_workbook.close()
        assert hashlib.sha256(new_same_reading.read_bytes()).hexdigest() != hashlib.sha256(
            first_excel.read_bytes()).hexdigest()
        sp.import_actual_gpt_decisions(b, new_same_reading)
        annotated = reviewed_override()
        assert annotated["source"] == "人工 GUI actual 視覺確認"
        assert "portable_actual_conflict:" in annotated["note"]
        assert hashlib.sha256(new_same_reading.read_bytes()).hexdigest() in annotated["note"]
        assert portability.actual_excel_conflict_state(
            b, sp.json_load_strict(b / "校對工作階段.json"),
            sp.json_load_strict(b / "人工判定資料庫.json")) == []
        before_new_repeat = {path: path.read_bytes() for path in repeat_paths}
        assert sp.import_actual_gpt_decisions(b, new_same_reading)[0] == 0
        assert {path: path.read_bytes() for path in repeat_paths} == before_new_repeat
        provenance = ar._read_csv(sp.project_actual_evidence_root(b) / ar.GLYPH_PROVENANCE_FILE,
                                  ar.GLYPH_PROVENANCE_HEADERS)
        assert any("portable_actual_conflict:" in row["note"]
                   and row["source"] == "人工 GUI actual 視覺確認" for row in provenance)
        with pytest.raises(ValueError, match="不同判定"):
            sp.import_actual_gpt_decisions(b, later_conflict_excel)
        later_receipt = sp.json_load_strict(receipt_path)
        assert len(later_receipt["conflicts"]) == 2
        assert portability.actual_excel_conflict_state(
            b, sp.json_load_strict(b / "校對工作階段.json"),
            sp.json_load_strict(b / "人工判定資料庫.json")) == [review_id]
        sp.stage_manual_actual_correction(b, review_id, "ㄐㄩㄝˊ", note="第二次原頁核對")
        sp.apply_staged_manual_actual_corrections(b)
        assert len(sp.json_load_strict(receipt_path)["conflicts"]) == 2
        assert portability.actual_excel_conflict_state(
            b, sp.json_load_strict(b / "校對工作階段.json"),
            sp.json_load_strict(b / "人工判定資料庫.json")) == []
        assert portability._load_project(b)[0]["session_id"] == manifest["session_id"]
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


def test_real_pipeline_actual_override_refresh_publishes_current_manifest(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = _pdfs(tmp_path)
    a, b = tmp_path / "source", tmp_path / "continued"
    with patch.object(sp, "decode", side_effect=_synthetic_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], a, defer_excel_reports=True)
        original = sp.json_load_strict(a / "校對工作階段.json")
        entry = original["records"][0]
        key = ar._override_key_from_entry(entry)
        fields = ("pdf_contains", "pdf_excludes", "page", "target_char", "stable_key", "x0", "y0")
        override = dict(zip(fields, key)) | {
            "actual_reading": "ㄐㄩㄝˊ", "source": "source direct visual", "note": "verified page 1",
        }
        ar._write_csv(sp.project_actual_evidence_root(a) / ar.OCCURRENCE_OVERRIDE_FILE,
                      ar.OVERRIDE_HEADERS, [override])
        sp.refresh_actual_project(a)
        refreshed = sp.json_load_strict(a / "校對工作階段.json")
        assert refreshed["session_id"] == original["session_id"]
        portability.prepare_portable_project(a)
        first.rename(tmp_path / "source-away.pdf")
        result = portability.continue_project(a, second, b)
        assert result["matched_actual_overrides"] == 1
        assert not (b / portability.INCOMPLETE_FILE).exists()
        b_manifest = sp.json_load_strict(b / "校對工作階段.json")
        sp.validate_manifest_integrity(b_manifest)
        assert b_manifest["pdfs"][0]["actual_workbook_sha256"] == sp.sha256_file(
            b / "01_實際注音" / Path(b_manifest["pdfs"][0]["actual_workbook"]).name)
        assert (b / "待人工確認.json").is_file()
        assert (b / "注音校對_最終報告.xlsx").is_file()
        target_overrides = ar._read_csv(sp.project_actual_evidence_root(b) / ar.OCCURRENCE_OVERRIDE_FILE,
                                        ar.OVERRIDE_HEADERS)
        assert len([row for row in target_overrides if row["source"] ==
                    "portable occurrence-local actual (no Global quorum)"]) == 1
        assert portability._load_project(b)[0]["session_id"] == b_manifest["session_id"]
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


def test_partial_portable_publication_requires_bound_repair_or_new_success(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = _pdfs(tmp_path)
    source, partial, completed, partial_proof = [tmp_path / name for name in
                                                  ("source", "partial", "completed", "partial-proof")]
    with patch.object(sp, "decode", side_effect=_synthetic_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], source, defer_excel_reports=True)
        source_manifest = sp.json_load_strict(source / "校對工作階段.json")
        ReviewSaveService(source).save_event(source_manifest["records"][0]["review_id"], {
            "action": "確認非校對範圍", "exclusion_reason": "source visual",
            "exclusion_evidence": "source PDF page 1",
        })
        portability.prepare_portable_project(source)
        with patch.object(portability, "_publish_portable_outputs", side_effect=OSError("disk full")):
            with pytest.raises(OSError, match="disk full"):
                portability.continue_project(source, second, partial)
        marker = partial / portability.INCOMPLETE_FILE
        assert marker.is_file()
        watched = [marker, partial / "校對工作階段.json", partial / "人工判定資料庫.json",
                   partial / "待人工確認.json", partial / "pipeline_status.json",
                   partial / "注音校對_最終報告.xlsx"]
        before = {path: path.read_bytes() if path.exists() else None for path in watched}
        with patch.object(sys, "argv", ["standalone_proofread.py", "-o", str(partial), "--report-only"]):
            with pytest.raises(ValueError, match="接續未完成"):
                sp.main()
        assert {path: path.read_bytes() if path.exists() else None for path in watched} == before
        # A bound, already-committed presentation is resumed only by the
        # explicit repair entrypoint. The original source remains untouched.
        with patch.object(sys, "argv", ["standalone_proofread.py", "-o", str(partial), "--repair-project"]):
            assert sp.main() == 0
        assert not marker.exists()
        assert (partial / "注音校對_最終報告.xlsx").is_file()

        source_db_before = (source / "人工判定資料庫.json").read_bytes()
        original_writer = portability.write_excel_content_proof
        def fail_final_excel(*args, **kwargs):
            if kwargs.get("allow_internal_incomplete"):
                raise OSError("final proof disk full")
            return original_writer(*args, **kwargs)
        with patch.object(portability, "write_excel_content_proof", side_effect=fail_final_excel):
            with pytest.raises(OSError, match="final proof disk full"):
                portability.continue_project(source, second, partial_proof)
        assert (partial_proof / portability.INCOMPLETE_FILE).exists()
        assert (source / "人工判定資料庫.json").read_bytes() == source_db_before
        with patch.object(sys, "argv", ["standalone_proofread.py", "-o", str(partial_proof), "--report-only"]):
            with pytest.raises(ValueError, match="接續未完成"):
                sp.main()
        with patch.object(sys, "argv", ["standalone_proofread.py", "-o", str(partial_proof), "--repair-project"]):
            assert sp.main() == 0
        assert not (partial_proof / portability.INCOMPLETE_FILE).exists()
        assert (partial_proof / "待判定候選_給GPT.xlsx").is_file()

        portability.continue_project(source, second, completed)
        assert not (completed / portability.INCOMPLETE_FILE).exists()
        with patch.object(sys, "argv", ["standalone_proofread.py", "-o", str(completed), "--report-only"]):
            assert sp.main() == 0
        assert (completed / "注音校對_最終報告.xlsx").is_file()
        assert sp.json_load_strict(completed / "pipeline_status.json")["user_report"] == str(
            completed / "注音校對_最終報告.xlsx")
        # The successfully published B project can still use its ordinary
        # same-session Excel import after a separate interrupted B was denied.
        exported = sp.export_pending_for_gpt(completed)
        workbook = load_workbook(exported)
        sheet = workbook["待判定候選"]
        headers = [cell.value for cell in sheet[1]]
        assert sheet.max_row == 2
        sheet.cell(2, headers.index("action") + 1, "確認非校對範圍")
        sheet.cell(2, headers.index("exclusion_reason") + 1, "completed B local visual")
        sheet.cell(2, headers.index("exclusion_evidence") + 1, "B PDF page 1")
        filled = tmp_path / "completed-same-session.xlsx"
        workbook.save(filled)
        workbook.close()
        assert sp.import_gpt_decisions(completed, filled)[0] == 1
        assert not (completed / portability.INCOMPLETE_FILE).exists()
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


@pytest.mark.parametrize("kind", ["expected", "actual"])
@pytest.mark.parametrize("entrypoint", ["direct", "auto"])
def test_incomplete_portable_project_rejects_same_session_excel_before_writes(
        tmp_path, monkeypatch, kind, entrypoint):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, _ = _pdfs(tmp_path)
    target = tmp_path / "target"
    decoder = _synthetic_decode if kind == "expected" else _synthetic_unresolved_decode
    with patch.object(sp, "decode", side_effect=decoder), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], target, defer_excel_reports=True)
        if kind == "expected":
            workbook_path = sp.export_pending_for_gpt(target)
            sheet_name = "待判定候選"
            values = {"action": "確認非校對範圍", "exclusion_reason": "local visual",
                      "exclusion_evidence": "PDF page 1"}
        else:
            sp.export_actual_pending_for_gpt(target)
            workbook_path = target / "actual待判定_GPT包" / "actual待判定_給GPT.xlsx"
            sheet_name = "actual待判定"
            values = {"decision": "VERIFIED", "actual_reading": "ㄐㄩㄝˊ",
                      "confidence": "高", "sample_a_checked": "Y"}
        workbook = load_workbook(workbook_path)
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        for key, value in values.items():
            sheet.cell(2, headers.index(key) + 1, value)
        filled = tmp_path / f"{kind}-filled.xlsx"
        workbook.save(filled)
        workbook.close()
        marker = target / portability.INCOMPLETE_FILE
        sp.json_save(marker, {"status": "PRESENTATION_PENDING", "source": "interrupted transfer"})
        watched = [marker, target / "人工判定資料庫.json", target / "pipeline_status.json",
                   target / "注音校對_最終報告.xlsx", target / "待人工確認.json",
                   sp.project_actual_evidence_root(target) / ar.OCCURRENCE_OVERRIDE_FILE]
        before = {path: path.read_bytes() if path.exists() else None for path in watched}
        with pytest.raises(ValueError, match="接續未完成"):
            if entrypoint == "direct":
                (sp.import_gpt_decisions if kind == "expected" else sp.import_actual_gpt_decisions)(target, filled)
            else:
                with patch.object(sys, "argv", ["standalone_proofread.py", "-o", str(target),
                                                "--import-gpt-auto", str(filled)]):
                    sp.main()
        assert {path: path.read_bytes() if path.exists() else None for path in watched} == before
        marker.unlink()
        if kind == "expected":
            assert sp.import_gpt_decisions(target, filled)[0] == 1
        else:
            assert sp.import_actual_gpt_decisions(target, filled)[0] == 1
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


@pytest.mark.parametrize("sha_mode", ["same", "different"])
def test_direct_expected_excel_import_maps_different_sha_session(tmp_path, monkeypatch, sha_mode):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, different = _pdfs(tmp_path)
    second = first if sha_mode == "same" else different
    a, b = tmp_path / "project-a", tmp_path / "project-b"
    with patch.object(sp, "decode", side_effect=_synthetic_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], a, defer_excel_reports=True)
        sp.run_pipeline_pdfs([second], b, defer_excel_reports=True)
        a_manifest = sp.json_load_strict(a / "校對工作階段.json")
        b_manifest = sp.json_load_strict(b / "校對工作階段.json")
        assert a_manifest["session_id"] != b_manifest["session_id"]
        assert (a_manifest["pdfs"][0]["pdf_sha256"] == b_manifest["pdfs"][0]["pdf_sha256"]) == (sha_mode == "same")
        workbook_path = sp.export_pending_for_gpt(a)
        workbook = load_workbook(workbook_path)
        assert workbook[portability.EXCEL_PROOF_SHEET].sheet_state == "veryHidden"
        sheet = workbook["待判定候選"]
        headers = [cell.value for cell in sheet[1]]
        sheet.cell(3, headers.index("action") + 1, "確認非校對範圍")
        sheet.cell(3, headers.index("exclusion_reason") + 1, "source visual")
        sheet.cell(3, headers.index("exclusion_evidence") + 1, "PDF page 1")
        renamed = tmp_path / "expected與差異判定_匯入.xlsx"
        workbook.save(renamed)
        workbook.close()
        assert sp.detect_gpt_workbook_kind(renamed) == "expected"
        count, _, report = sp.import_gpt_decisions(b, renamed)
        assert count == 1 and report.is_file()
        b_db = sp.json_load_strict(b / "人工判定資料庫.json")
        event = b_db["events"][b_manifest["records"][1]["review_id"]]
        assert event["portability_source"]["pdf_sha256"] == a_manifest["pdfs"][0]["pdf_sha256"]
        assert event["portability_source"]["target_pdf_sha256"] == b_manifest["pdfs"][0]["pdf_sha256"]
        assert sp.import_gpt_decisions(b, renamed)[0] == 0
        assert len(sp.json_load_strict(b / "待人工確認.json")["pending"]) == 1
        assert b_manifest["records"][0]["review_id"] not in b_db["events"]
        assert portability._load_project(b)[0]["session_id"] == b_manifest["session_id"]
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


@pytest.mark.parametrize("sha_mode", ["same", "different"])
def test_direct_actual_excel_import_is_local_only_across_different_sha(tmp_path, monkeypatch, sha_mode):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, different = _pdfs(tmp_path)
    second = first if sha_mode == "same" else different
    a, b = tmp_path / "source", tmp_path / "target"
    with patch.object(sp, "decode", side_effect=_synthetic_unresolved_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], a, defer_excel_reports=True)
        sp.run_pipeline_pdfs([second], b, defer_excel_reports=True)
        a_manifest = sp.json_load_strict(a / "校對工作階段.json")
        b_manifest = sp.json_load_strict(b / "校對工作階段.json")
        assert a_manifest["session_id"] != b_manifest["session_id"]
        assert (a_manifest["pdfs"][0]["pdf_sha256"] == b_manifest["pdfs"][0]["pdf_sha256"]) == (sha_mode == "same")
        exported = sp.export_actual_pending_for_gpt(a)
        assert exported is not None
        workbook_path = a / "actual待判定_GPT包" / "actual待判定_給GPT.xlsx"
        workbook = load_workbook(workbook_path)
        assert workbook[portability.EXCEL_PROOF_SHEET].sheet_state == "veryHidden"
        sheet = workbook["actual待判定"]
        headers = [cell.value for cell in sheet[1]]
        for key, value in {"decision": "VERIFIED", "actual_reading": "ㄐㄩㄝˊ",
                           "confidence": "高", "sample_a_checked": "Y", "note": "source page 1"}.items():
            sheet.cell(2, headers.index(key) + 1, value)
        renamed = tmp_path / "actual視覺判定_匯入.xlsx"
        workbook.save(renamed)
        workbook.close()
        assert sp.detect_gpt_workbook_kind(renamed) == "actual"
        result = sp.import_actual_gpt_decisions(b, renamed)
        assert result[0] == 1
        assert result.status["global_promotion_delivery"]["status"] == "NO_PENDING"
        assert not (b / portability.INCOMPLETE_FILE).exists()
        refreshed = sp.json_load_strict(b / "校對工作階段.json")
        assert refreshed["session_id"] == b_manifest["session_id"]
        assert portability._load_project(b)[0]["session_id"] == b_manifest["session_id"]
        assert sp.import_actual_gpt_decisions(b, renamed)[0] == 0
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


def test_old_filled_expected_excel_imported_in_a_can_gain_present_proof(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = _pdfs(tmp_path)
    a, b = tmp_path / "source", tmp_path / "target"
    with patch.object(sp, "decode", side_effect=_synthetic_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], a, defer_excel_reports=True)
        sp.run_pipeline_pdfs([second], b, defer_excel_reports=True)
        old = tmp_path / "已完成的舊版判定.xlsx"
        workbook = load_workbook(sp.export_pending_for_gpt(a))
        del workbook[portability.EXCEL_PROOF_SHEET]
        sheet = workbook["待判定候選"]
        headers = [cell.value for cell in sheet[1]]
        for key, value in {"action": "確認非校對範圍", "exclusion_reason": "原頁範圍",
                           "exclusion_evidence": "A 原頁檢查"}.items():
            sheet.cell(2, headers.index(key) + 1, value)
        workbook.save(old)
        workbook.close()
        assert sp.import_gpt_decisions(a, old)[0] == 1
        before = old.read_bytes()
        portable = portability.attach_filled_excel_proof(a, old)
        assert old.read_bytes() == before and portable.is_file()
        assert sp.import_gpt_decisions(b, portable)[0] == 1
        entry = sp.json_load_strict(b / "校對工作階段.json")["records"][0]
        event = sp.json_load_strict(b / "人工判定資料庫.json")["events"][entry["review_id"]]
        assert event["portability_source"]["source_excel_decision"]["action"] == "確認非校對範圍"
        assert sp.import_gpt_decisions(b, portable)[0] == 0
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


def test_old_filled_actual_excel_imported_in_a_can_gain_present_proof(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = _pdfs(tmp_path)
    a, b = tmp_path / "source", tmp_path / "target"
    with patch.object(sp, "decode", side_effect=_synthetic_unresolved_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], a, defer_excel_reports=True)
        sp.run_pipeline_pdfs([second], b, defer_excel_reports=True)
        sp.export_actual_pending_for_gpt(a)
        old = tmp_path / "已完成的舊版actual.xlsx"
        workbook = load_workbook(a / "actual待判定_GPT包" / "actual待判定_給GPT.xlsx")
        del workbook[portability.EXCEL_PROOF_SHEET]
        sheet = workbook["actual待判定"]
        headers = [cell.value for cell in sheet[1]]
        for key, value in {"decision": "VERIFIED", "actual_reading": "ㄐㄩㄝˊ",
                           "confidence": "高", "sample_a_checked": "Y", "note": "A 原頁"}.items():
            sheet.cell(2, headers.index(key) + 1, value)
        workbook.save(old)
        workbook.close()
        assert sp.import_actual_gpt_decisions(a, old)[0] == 1
        before = old.read_bytes()
        portable = portability.attach_filled_excel_proof(a, old)
        assert old.read_bytes() == before and portable.is_file()
        assert sp.import_actual_gpt_decisions(b, portable)[0] == 1
        assert sp.import_actual_gpt_decisions(b, portable)[0] == 0
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


def test_excel_expected_conflict_and_same_decision_keep_sources(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = _pdfs(tmp_path)
    a, b = tmp_path / "source", tmp_path / "target"
    with patch.object(sp, "decode", side_effect=_synthetic_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], a, defer_excel_reports=True)
        sp.run_pipeline_pdfs([second], b, defer_excel_reports=True)
        exported = sp.export_pending_for_gpt(a)
        first_xlsx, same_xlsx, conflict_xlsx = [tmp_path / name for name in
                                               ("first.xlsx", "same.xlsx", "conflict.xlsx")]
        workbook = load_workbook(exported)
        sheet = workbook["待判定候選"]
        headers = [cell.value for cell in sheet[1]]
        sheet.cell(2, headers.index("action") + 1, "確認非校對範圍")
        sheet.cell(2, headers.index("exclusion_reason") + 1, "same judgment")
        sheet.cell(2, headers.index("exclusion_evidence") + 1, "source page")
        workbook.save(first_xlsx)
        workbook.properties.title = "same decision from another saved workbook"
        workbook.save(same_xlsx)
        sheet.cell(2, headers.index("exclusion_reason") + 1, "different judgment")
        workbook.save(conflict_xlsx)
        workbook.close()
        assert sp.import_gpt_decisions(b, first_xlsx)[0] == 1
        assert sp.import_gpt_decisions(b, same_xlsx)[0] == 0
        # A second, byte-distinct Excel with the same decision is provenance,
        # not a second active verdict.
        target_id = sp.json_load_strict(b / "校對工作階段.json")["records"][0]["review_id"]
        assert len(sp.json_load_strict(b / "人工判定資料庫.json")["events"][target_id]
                   .get("portability_duplicate_sources", [])) == 1
        db_path = b / "人工判定資料庫.json"
        marker_path = b / portability.INCOMPLETE_FILE
        receipt_path = b / portability.CONFLICT_FILE
        before = {path: path.read_bytes() if path.exists() else None
                  for path in (db_path, marker_path, receipt_path)}
        original_save = sp.json_save
        def fail_target_db(path, *args, **kwargs):
            if Path(path) == db_path:
                raise OSError("disk full before DB commit")
            return original_save(path, *args, **kwargs)
        with patch.object(sp, "json_save", side_effect=fail_target_db):
            with pytest.raises(OSError, match="before DB commit"):
                sp.import_gpt_decisions(b, conflict_xlsx)
        assert {path: path.read_bytes() if path.exists() else None
                for path in (db_path, marker_path, receipt_path)} == before
        watched = [db_path, marker_path, receipt_path, b / "待人工確認.json",
                   b / "注音校對_最終報告.xlsx", b / "pipeline_status.json"]
        before_publish = {path: path.read_bytes() if path.exists() else None for path in watched}
        with patch.object(portability, "_publish_portable_outputs", side_effect=OSError("report disk full")):
            with pytest.raises(OSError, match="report disk full"):
                sp.import_gpt_decisions(b, conflict_xlsx)
        assert {path: path.read_bytes() if path.exists() else None for path in watched} == before_publish
        with pytest.raises(ValueError, match="不同判定"):
            sp.import_gpt_decisions(b, conflict_xlsx)
        receipt = sp.json_load_strict(b / portability.CONFLICT_FILE)
        assert len(receipt["conflicts"]) == 1
        assert receipt["conflicts"][0]["source_event"]["exclusion_reason"] == "different judgment"
        assert receipt["conflicts"][0]["target_event"]["exclusion_reason"] == "same judgment"
        assert target_id not in sp.json_load_strict(b / "人工判定資料庫.json")["events"]
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


def test_expected_excel_sequential_conflicts_keep_local_adjudication_and_history(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = _pdfs(tmp_path)
    source, target = tmp_path / "source", tmp_path / "target"
    with patch.object(sp, "decode", side_effect=_synthetic_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], source, defer_excel_reports=True)
        sp.run_pipeline_pdfs([second], target, defer_excel_reports=True)
        exported = sp.export_pending_for_gpt(source)
        workbook = load_workbook(exported)
        sheet = workbook["待判定候選"]
        headers = [cell.value for cell in sheet[1]]
        paths = []
        for label in ("first", "second", "third"):
            sheet.cell(2, headers.index("action") + 1, "確認非校對範圍")
            sheet.cell(2, headers.index("exclusion_reason") + 1, label)
            sheet.cell(2, headers.index("exclusion_evidence") + 1, "PDF page 1")
            path = tmp_path / f"{label}.xlsx"
            workbook.save(path)
            paths.append(path)
        workbook.close()
        manifest = sp.json_load_strict(target / "校對工作階段.json")
        conflicted_id, other_id = [record["review_id"] for record in manifest["records"]]
        ReviewSaveService(target).save_event(other_id, {
            "action": "確認非校對範圍", "exclusion_reason": "unrelated position",
            "exclusion_evidence": "PDF page 1"})
        assert sp.import_gpt_decisions(target, paths[0])[0] == 1
        with pytest.raises(ValueError, match="不同判定"):
            sp.import_gpt_decisions(target, paths[1])
        first_receipt = sp.json_load_strict(target / portability.CONFLICT_FILE)
        assert len(first_receipt["conflicts"]) == 1
        ReviewSaveService(target).save_event(conflicted_id, {
            "action": "確認非校對範圍", "exclusion_reason": "local adjudication",
            "exclusion_evidence": "B PDF direct visual"})
        adjudication = sp.json_load_strict(target / "人工判定資料庫.json")["events"][conflicted_id]
        assert adjudication["portability_conflict_resolution"]["original_conflicts"] == first_receipt["conflicts"]
        watched = [target / "人工判定資料庫.json", target / portability.CONFLICT_FILE,
                   target / portability.INCOMPLETE_FILE, target / "待人工確認.json",
                   target / "注音校對_最終報告.xlsx", target / "pipeline_status.json"]
        before_failed_write = {path: path.read_bytes() if path.exists() else None for path in watched}
        original_save = sp.json_save
        def fail_new_conflict_db(path, *args, **kwargs):
            if Path(path) == target / "人工判定資料庫.json":
                raise OSError("second conflict DB disk full")
            return original_save(path, *args, **kwargs)
        with patch.object(sp, "json_save", side_effect=fail_new_conflict_db):
            with pytest.raises(OSError, match="second conflict DB disk full"):
                sp.import_gpt_decisions(target, paths[2])
        assert {path: path.read_bytes() if path.exists() else None for path in watched} == before_failed_write
        with pytest.raises(ValueError, match="不同判定"):
            sp.import_gpt_decisions(target, paths[2])
        receipt = sp.json_load_strict(target / portability.CONFLICT_FILE)
        assert len(receipt["conflicts"]) == 2
        assert first_receipt["conflicts"][0] in receipt["conflicts"]
        assert receipt["conflicts"][1]["target_event"] == adjudication
        db = sp.json_load_strict(target / "人工判定資料庫.json")
        assert conflicted_id not in db["events"]
        assert db["events"][other_id]["exclusion_reason"] == "unrelated position"
        assert portability.validate_conflict_state(target, manifest, db) == [conflicted_id]
        before = (target / portability.CONFLICT_FILE).read_bytes()
        with pytest.raises(ValueError, match="衝突|不同判定"):
            sp.import_gpt_decisions(target, paths[2])
        assert (target / portability.CONFLICT_FILE).read_bytes() == before
        assert sp.materialize_ledger(manifest, db)[0]["active_review"] is True
        assert len(sp.json_load_strict(target / "待人工確認.json")["pending"]) == 1
        ReviewSaveService(target).save_event(conflicted_id, {
            "action": "確認非校對範圍", "exclusion_reason": "second local adjudication",
            "exclusion_evidence": "B PDF direct visual again"})
        resolved_db = sp.json_load_strict(target / "人工判定資料庫.json")
        resolution = resolved_db["events"][conflicted_id]["portability_conflict_resolution"]
        assert resolution["original_conflicts"] == receipt["conflicts"]
        assert portability.validate_conflict_state(target, manifest, resolved_db) == []
        assert portability._load_project(target)[0]["session_id"] == manifest["session_id"]
        assert sp.json_load_strict(target / portability.CONFLICT_FILE) == receipt
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


@pytest.mark.parametrize("kind", ["expected", "actual"])
def test_old_proofless_excel_only_maps_exact_same_sha_and_snapshot(tmp_path, monkeypatch, kind):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, different = _pdfs(tmp_path)
    a, b, c = [tmp_path / name for name in ("source", "exact-target", "different-target")]
    decoder = _synthetic_decode if kind == "expected" else _synthetic_unresolved_decode
    with patch.object(sp, "decode", side_effect=decoder), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], a, defer_excel_reports=True)
        sp.run_pipeline_pdfs([first], b, defer_excel_reports=True)
        sp.run_pipeline_pdfs([different], c, defer_excel_reports=True)
        if kind == "expected":
            exported = sp.export_pending_for_gpt(a)
            sheet_name, filled = "待判定候選", {
                "action": "確認非校對範圍", "exclusion_reason": "source",
                "exclusion_evidence": "page 1"}
        else:
            sp.export_actual_pending_for_gpt(a)
            exported = a / "actual待判定_GPT包" / "actual待判定_給GPT.xlsx"
            sheet_name, filled = "actual待判定", {
                "decision": "VERIFIED", "actual_reading": "ㄐㄩㄝˊ",
                "confidence": "高", "sample_a_checked": "Y"}
        workbook = load_workbook(exported)
        del workbook[portability.EXCEL_PROOF_SHEET]
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        for key, value in filled.items():
            sheet.cell(2, headers.index(key) + 1, value)
        old = tmp_path / f"old-{kind}.xlsx"
        workbook.save(old)
        workbook.close()
        importer = sp.import_gpt_decisions if kind == "expected" else sp.import_actual_gpt_decisions
        with pytest.raises(ValueError, match="沒有全頁證據|不同 SHA"):
            importer(c, old)
        assert importer(b, old)[0] == 1
        assert importer(b, old)[0] == 0
        assert portability._load_project(b)[0]["session_id"] != sp.json_load_strict(a / "校對工作階段.json")["session_id"]
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


@pytest.mark.parametrize("change", ["text", "zhuyin", "image"])
def test_excel_full_page_proof_rejects_same_filename_content_changes(tmp_path, monkeypatch, change):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, _ = _pdfs(tmp_path)
    altered_dir = tmp_path / "other-download"
    altered_dir.mkdir()
    altered = altered_dir / first.name
    with fitz.open(first) as document:
        page = document[0]
        if change == "text":
            page.insert_text((200, 110), "教材新文字", fontname="china-t", fontsize=14)
        elif change == "zhuyin":
            page.insert_text((200, 70), "ㄅㄆㄇ", fontname="china-t", fontsize=12)
        else:
            page.draw_rect(fitz.Rect(200, 130, 245, 170), color=(1, 0, 0), fill=(1, 0, 0))
        document.save(altered)
    assert sp.sha256_file(first) != sp.sha256_file(altered)
    a, b = tmp_path / "source", tmp_path / "target"
    with patch.object(sp, "decode", side_effect=_synthetic_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], a, defer_excel_reports=True)
        sp.run_pipeline_pdfs([altered], b, defer_excel_reports=True)
        workbook = load_workbook(sp.export_pending_for_gpt(a))
        sheet = workbook["待判定候選"]
        headers = [cell.value for cell in sheet[1]]
        for key, value in {"action": "確認非校對範圍", "exclusion_reason": "source",
                           "exclusion_evidence": "page 1"}.items():
            sheet.cell(2, headers.index(key) + 1, value)
        filled = tmp_path / "filled.xlsx"
        workbook.save(filled)
        workbook.close()
        db_path = b / "人工判定資料庫.json"
        before = db_path.read_bytes()
        with pytest.raises(ValueError, match="頁面內容|無法唯一對應"):
            sp.import_gpt_decisions(b, filled)
        assert db_path.read_bytes() == before
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


def test_actual_excel_conflict_and_precommit_write_failure_preserve_target(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = _pdfs(tmp_path)
    a, b, c = [tmp_path / name for name in ("source", "conflict-target", "write-fail-target")]
    with patch.object(sp, "decode", side_effect=_synthetic_unresolved_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], a, defer_excel_reports=True)
        sp.run_pipeline_pdfs([second], b, defer_excel_reports=True)
        sp.run_pipeline_pdfs([second], c, defer_excel_reports=True)
        sp.export_actual_pending_for_gpt(a)
        workbook = load_workbook(a / "actual待判定_GPT包" / "actual待判定_給GPT.xlsx")
        sheet = workbook["actual待判定"]
        headers = [cell.value for cell in sheet[1]]
        for key, value in {"decision": "VERIFIED", "actual_reading": "ㄐㄩㄝˊ",
                           "confidence": "高", "sample_a_checked": "Y"}.items():
            sheet.cell(2, headers.index(key) + 1, value)
        filled = tmp_path / "first-actual.xlsx"
        workbook.save(filled)
        sheet.cell(2, headers.index("actual_reading") + 1, "ㄐㄩㄝˇ")
        different = tmp_path / "different-actual.xlsx"
        workbook.save(different)
        sheet.cell(2, headers.index("actual_reading") + 1, "ㄐㄩㄝˋ")
        third = tmp_path / "third-actual.xlsx"
        workbook.save(third)
        workbook.close()
        root_c = sp.project_actual_evidence_root(c)
        original_c = (root_c / ar.OCCURRENCE_OVERRIDE_FILE).read_bytes()
        manifest_c = (c / "校對工作階段.json").read_bytes()
        with patch.object(ar, "_write_csv", side_effect=OSError("disk full")):
            with pytest.raises(OSError, match="disk full"):
                sp.import_actual_gpt_decisions(c, filled)
        assert (root_c / ar.OCCURRENCE_OVERRIDE_FILE).read_bytes() == original_c
        assert (c / "校對工作階段.json").read_bytes() == manifest_c
        assert not (c / portability.INCOMPLETE_FILE).exists()
        assert sp.import_actual_gpt_decisions(b, filled)[0] == 1
        root_b = sp.project_actual_evidence_root(b)
        original_b = (root_b / ar.OCCURRENCE_OVERRIDE_FILE).read_bytes()
        receipt_path = b / portability.ACTUAL_EXCEL_CONFLICT_FILE
        preserved = {path: path.read_bytes() for path in (
            b / "人工判定資料庫.json", b / "校對工作階段.json",
            b / "pipeline_status.json", root_b / ar.OCCURRENCE_OVERRIDE_FILE,
        )}
        original_json_save = sp.json_save

        def fail_receipt_write(path, *args, **kwargs):
            if Path(path) == receipt_path:
                raise OSError("receipt disk full")
            return original_json_save(path, *args, **kwargs)

        with patch.object(sp, "json_save", side_effect=fail_receipt_write):
            with pytest.raises(OSError, match="receipt disk full"):
                sp.import_actual_gpt_decisions(b, different)
        assert not receipt_path.exists()
        assert not (b / portability.INCOMPLETE_FILE).exists()
        for path, before in preserved.items():
            assert path.read_bytes() == before
        with pytest.raises(ValueError, match="不同判定"):
            sp.import_actual_gpt_decisions(b, different)
        with pytest.raises(ValueError, match="不同判定"):
            sp.import_actual_gpt_decisions(b, different)
        with pytest.raises(ValueError, match="不同判定"):
            sp.import_actual_gpt_decisions(b, third)
        assert (root_b / ar.OCCURRENCE_OVERRIDE_FILE).read_bytes() == original_b
        receipt = sp.json_load_strict(receipt_path)
        assert receipt["conflicts"][0]["source_reading"] == "ㄐㄩㄝˇ"
        assert receipt["conflicts"][0]["target_record"]["actual_reading"] == "ㄐㄩㄝˊ"
        assert [item["source_reading"] for item in receipt["conflicts"]] == ["ㄐㄩㄝˇ", "ㄐㄩㄝˋ"]
        with pytest.raises(ValueError, match="actual.*衝突|衝突.*actual"):
            portability._load_project(b)
        with pytest.raises(ValueError, match="actual.*衝突|衝突.*actual"):
            portability.prepare_portable_project(b)
        with pytest.raises(ValueError, match="actual.*衝突|衝突.*actual"):
            sp.regenerate_report(b)
        from standalone_gui import project_status_details
        assert "actual" in project_status_details(b)[0] and "衝突" in project_status_details(b)[0]
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


def test_actual_excel_prepared_and_post_ack_crash_windows_can_resume(tmp_path, monkeypatch):
    import global_glyph_promotion as promotion

    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = _pdfs(tmp_path)
    a, before_commit, after_refresh, after_reseal = [tmp_path / name for name in
                                                     ("source", "before-commit", "after-refresh", "after-reseal")]
    with patch.object(sp, "decode", side_effect=_synthetic_unresolved_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], a, defer_excel_reports=True)
        sp.run_pipeline_pdfs([second], before_commit, defer_excel_reports=True)
        sp.run_pipeline_pdfs([second], after_refresh, defer_excel_reports=True)
        sp.run_pipeline_pdfs([second], after_reseal, defer_excel_reports=True)
        sp.export_actual_pending_for_gpt(a)
        workbook = load_workbook(a / "actual待判定_GPT包" / "actual待判定_給GPT.xlsx")
        sheet = workbook["actual待判定"]
        headers = [cell.value for cell in sheet[1]]
        for key, value in {"decision": "VERIFIED", "actual_reading": "ㄐㄩㄝˊ",
                           "confidence": "高", "sample_a_checked": "Y"}.items():
            sheet.cell(2, headers.index(key) + 1, value)
        filled = tmp_path / "filled-actual.xlsx"
        workbook.save(filled)
        workbook.close()

        original_before = portability._actual_excel_file_snapshot(before_commit)
        with patch.object(ar, "_write_csv", side_effect=SystemExit("power loss before commit")):
            with pytest.raises(SystemExit, match="power loss"):
                sp.import_actual_gpt_decisions(before_commit, filled)
        assert (before_commit / portability.INCOMPLETE_FILE).exists()
        assert portability.resume_actual_excel_project(before_commit)["project_actual_commit"] == "ROLLED_BACK"
        assert not (before_commit / portability.INCOMPLETE_FILE).exists()
        assert portability._actual_excel_file_snapshot(before_commit) == original_before

        with patch.object(promotion, "acknowledge_project_refresh", side_effect=SystemExit("power loss after refresh")):
            with pytest.raises(SystemExit, match="power loss"):
                sp.import_actual_gpt_decisions(after_refresh, filled)
        marker = sp.json_load_strict(after_refresh / portability.INCOMPLETE_FILE)
        assert marker["status"] == "ACTUAL_EXCEL_RECOVERED"
        assert promotion.committed_project_recovery(sp.project_actual_evidence_root(after_refresh)) is not None
        resumed = portability.resume_actual_excel_project(after_refresh)
        assert resumed["project_actual_commit"] == "COMMITTED"
        assert not (after_refresh / portability.INCOMPLETE_FILE).exists()
        assert promotion.committed_project_recovery(sp.project_actual_evidence_root(after_refresh)) is None
        assert portability._load_project(after_refresh)[0]["session_id"] == sp.json_load_strict(
            after_refresh / "校對工作階段.json")["session_id"]
        assert sp.import_actual_gpt_decisions(after_refresh, filled)[0] == 0

        original_save = sp.json_save
        def interrupt_marker_transition(path, payload, **kwargs):
            if (Path(path) == after_reseal / portability.INCOMPLETE_FILE
                    and payload.get("status") == "ACTUAL_EXCEL_RECOVERED"):
                raise SystemExit("power loss after manifest reseal")
            return original_save(path, payload, **kwargs)
        with patch.object(sp, "json_save", side_effect=interrupt_marker_transition):
            with pytest.raises(SystemExit, match="power loss after manifest reseal"):
                sp.import_actual_gpt_decisions(after_reseal, filled)
        reseal_marker = sp.json_load_strict(after_reseal / portability.INCOMPLETE_FILE)
        assert reseal_marker["status"] == "ACTUAL_EXCEL_REFRESH_PENDING"
        assert portability._sha(after_reseal / "校對工作階段.json") != reseal_marker["pre_state"]["manifest"]
        sp.validate_output_artifact_hashes(sp.json_load_strict(after_reseal / "校對工作階段.json"))
        assert promotion.committed_project_recovery(sp.project_actual_evidence_root(after_reseal)) is not None
        assert portability.resume_actual_excel_project(after_reseal)["project_actual_commit"] == "COMMITTED"
        assert not (after_reseal / portability.INCOMPLETE_FILE).exists()
        assert sp.import_actual_gpt_decisions(after_reseal, filled)[0] == 0
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


@pytest.mark.parametrize("crash_after_commit", [False, True])
def test_parallel_actual_excel_imports_preserve_marker_and_recovery(tmp_path, monkeypatch, crash_after_commit):
    import global_glyph_promotion as promotion

    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = _pdfs(tmp_path)
    source, target = tmp_path / "source", tmp_path / "target"
    with patch.object(sp, "decode", side_effect=_synthetic_unresolved_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], source, defer_excel_reports=True)
        sp.run_pipeline_pdfs([second], target, defer_excel_reports=True)
        sp.export_actual_pending_for_gpt(source)
        workbook = load_workbook(source / "actual待判定_GPT包" / "actual待判定_給GPT.xlsx")
        sheet = workbook["actual待判定"]
        headers = [cell.value for cell in sheet[1]]
        for key, value in {"decision": "VERIFIED", "actual_reading": "ㄐㄩㄝˊ",
                           "confidence": "高", "sample_a_checked": "Y"}.items():
            sheet.cell(2, headers.index(key) + 1, value)
        filled = tmp_path / "filled-actual.xlsx"
        workbook.save(filled)
        workbook.close()
        root = sp.project_actual_evidence_root(target)
        original_overrides = ar._read_csv(root / ar.OCCURRENCE_OVERRIDE_FILE, ar.OVERRIDE_HEADERS)
        simultaneous_preflight = Barrier(2)
        first_marker = Event()
        second_resume = Event()
        first_done = Event()
        marker_owner = []
        results, errors = {}, {}
        original_proof = portability.load_excel_content_proof
        original_save = sp.json_save
        original_resume = portability.resume_actual_excel_project

        def preflight(*args, **kwargs):
            result = original_proof(*args, **kwargs)
            simultaneous_preflight.wait(20)
            return result

        def delayed_marker(path, value, **kwargs):
            if Path(path).name == portability.INCOMPLETE_FILE and value.get("status") == "ACTUAL_EXCEL_REFRESH_PENDING":
                if current_thread().name == "first":
                    result = original_save(path, value, **kwargs)
                    marker_owner.append(current_thread().name)
                    first_marker.set()
                    second_resume.wait(3)
                    return result
                first_marker.wait(3)
                result = original_save(path, value, **kwargs)
                marker_owner.append(current_thread().name)
                return result
            return original_save(path, value, **kwargs)

        def delayed_resume(path):
            if current_thread().name == "second":
                second_resume.set()
                first_done.wait(3)
            if crash_after_commit and marker_owner == [current_thread().name]:
                raise SystemExit("power loss after COMMITTED actual Excel")
            return original_resume(path)

        def run():
            try:
                results[current_thread().name] = sp.import_actual_gpt_decisions(target, filled)
            except BaseException as exc:
                errors[current_thread().name] = exc
            finally:
                if current_thread().name == "first":
                    first_done.set()

        with patch.object(portability, "load_excel_content_proof", side_effect=preflight), \
             patch.object(sp, "json_save", side_effect=delayed_marker), \
             patch.object(portability, "resume_actual_excel_project", side_effect=delayed_resume):
            threads = [Thread(target=run, name=name) for name in ("first", "second")]
            for thread in threads:
                thread.start()
            for thread in threads:
                thread.join(30)
        assert not any(thread.is_alive() for thread in threads)
        if crash_after_commit:
            assert not results
            assert len(errors) == 2
            assert any(isinstance(error, SystemExit) for error in errors.values())
            assert (target / portability.INCOMPLETE_FILE).exists()
            assert promotion.committed_project_recovery(root) is not None
            resumed = portability.resume_actual_excel_project(target)
            assert resumed["project_actual_commit"] == "COMMITTED"
            assert resumed["global_promotion_delivery"]["status"] == "NO_PENDING"
        else:
            assert len(results) == 1, {name: str(error) for name, error in errors.items()}
            assert len(errors) == 1
            result = next(iter(results.values()))
            assert result[0] == 1
            assert result.status["global_promotion_delivery"]["status"] == "NO_PENDING"
        assert not (target / portability.INCOMPLETE_FILE).exists()
        assert not (root / promotion.PROJECT_TRANSACTION_FILE).exists()
        assert promotion.committed_project_recovery(root) is None
        overrides = ar._read_csv(root / ar.OCCURRENCE_OVERRIDE_FILE, ar.OVERRIDE_HEADERS)
        assert len(overrides) == len(original_overrides) + 1
        imported = [row for row in overrides if row["pdf_contains"] == second.stem]
        assert len(imported) == 1 and "source_excel_sha256" in imported[0]["note"]
        assert sp.json_load_strict(target / "pipeline_status.json")["user_report"] == str(
            target / "注音校對_最終報告.xlsx")
        assert portability._load_project(target)[0]["session_id"] == sp.json_load_strict(
            target / "校對工作階段.json")["session_id"]
        if not crash_after_commit:
            assert sp.import_actual_gpt_decisions(target, filled)[0] == 0
            original_load = portability.load_excel_content_proof
            def incomplete_after_preflight(*args, **kwargs):
                result = original_load(*args, **kwargs)
                sp.json_save(target / portability.INCOMPLETE_FILE, {
                    "status": "ACTUAL_EXCEL_REFRESH_PENDING",
                    "source_excel_sha256": portability._sha(filled),
                    "pre_state": portability._actual_excel_file_snapshot(target),
                })
                return result
            with patch.object(portability, "load_excel_content_proof", side_effect=incomplete_after_preflight):
                with pytest.raises(ValueError, match="待專用恢復"):
                    sp.import_actual_gpt_decisions(target, filled)
            assert (target / portability.INCOMPLETE_FILE).exists()
            assert portability.resume_actual_excel_project(target)["project_actual_commit"] == "ROLLED_BACK"
        assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


def test_actual_excel_midrefresh_workbook_rewrite_can_resume_twice(tmp_path, monkeypatch):
    import global_glyph_promotion as promotion

    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = _pdfs(tmp_path)
    source, target = tmp_path / "source", tmp_path / "target"
    with patch.object(sp, "decode", side_effect=_synthetic_unresolved_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], source, defer_excel_reports=True)
        sp.run_pipeline_pdfs([second], target, defer_excel_reports=True)
        manifest = sp.json_load_strict(target / "校對工作階段.json")
        expected_id = manifest["records"][1]["review_id"]
        ReviewSaveService(target).save_event(expected_id, {
            "action": "補建expected證據", "expected_set": ["ㄐㄩㄝˊ"],
            "expected_evidence": "independent manual expected", "context_evidence": "PDF page 1",
        })
        expected_event = sp.json_load_strict(target / "人工判定資料庫.json")["events"][expected_id]
        sp.export_actual_pending_for_gpt(source)
        workbook = load_workbook(source / "actual待判定_GPT包" / "actual待判定_給GPT.xlsx")
        sheet = workbook["actual待判定"]
        headers = [cell.value for cell in sheet[1]]
        for key, value in {"decision": "VERIFIED", "actual_reading": "ㄐㄩㄝˊ",
                           "confidence": "高", "sample_a_checked": "Y"}.items():
            sheet.cell(2, headers.index(key) + 1, value)
        filled = tmp_path / "filled-actual.xlsx"
        workbook.save(filled)
        workbook.close()
        root = sp.project_actual_evidence_root(target)
        marker = target / portability.INCOMPLETE_FILE
        old_manifest = sp.json_load_strict(target / "校對工作階段.json")
        def rewrite_then_crash(*args, **kwargs):
            _synthetic_unresolved_decode(*args, **kwargs)
            raise RuntimeError("power loss after workbook rewrite")
        with patch.object(sp, "decode", side_effect=rewrite_then_crash):
            with pytest.raises(Exception, match="power loss after workbook rewrite"):
                sp.import_actual_gpt_decisions(target, filled)
        assert marker.exists() and promotion.committed_project_recovery(root) is not None
        with pytest.raises(ValueError, match="hash 不符"):
            sp.validate_output_artifact_hashes(old_manifest)
        intact_marker = sp.json_load_strict(marker)
        for item in intact_marker["sealed_workbooks"]["entries"]:
            portability._restore_exact_file(target / item["path"], base64.b64decode(item["bytes"]))
        sp.validate_output_artifact_hashes(old_manifest)
        sp.json_save(marker, {"status": "ACTUAL_EXCEL_REFRESH_PENDING",
                              "source_excel_sha256": "0" * 64})
        watched = [marker, root / promotion.PROJECT_TRANSACTION_FILE,
                   target / "人工判定資料庫.json", root / ar.OCCURRENCE_OVERRIDE_FILE,
                   target / "校對工作階段.json"]
        before_refusal = {path: path.read_bytes() for path in watched}
        with pytest.raises(ValueError, match="恢復材料不完整"):
            portability.resume_actual_excel_project(target)
        assert {path: path.read_bytes() for path in watched} == before_refusal
        sp.json_save(marker, intact_marker)
        damaged_marker = copy.deepcopy(intact_marker)
        damaged_marker["sealed_workbooks"]["entries"][0]["bytes"] = "AA=="
        sp.json_save(marker, damaged_marker)
        with pytest.raises(ValueError, match="備份內容不符"):
            portability.resume_actual_excel_project(target)
        assert promotion.committed_project_recovery(root) is not None
        sp.json_save(marker, intact_marker)
        with patch.object(sp, "decode", side_effect=rewrite_then_crash):
            with pytest.raises(Exception, match="power loss after workbook rewrite"):
                portability.resume_actual_excel_project(target)
        assert marker.exists() and promotion.committed_project_recovery(root) is not None
        resumed = portability.resume_actual_excel_project(target)
        assert resumed["project_actual_commit"] == "COMMITTED"
        assert resumed["global_promotion_delivery"]["status"] == "NO_PENDING"
        assert not marker.exists() and promotion.committed_project_recovery(root) is None
        current = sp.json_load_strict(target / "校對工作階段.json")
        sp.validate_output_artifact_hashes(current)
        assert current["session_id"] == old_manifest["session_id"]
        assert sp.json_load_strict(target / "人工判定資料庫.json")["events"][expected_id] == expected_event
        override_rows = ar._read_csv(root / ar.OCCURRENCE_OVERRIDE_FILE, ar.OVERRIDE_HEADERS)
        assert any(row["pdf_contains"] == second.stem and
                   "source_excel_sha256" in row["note"] for row in override_rows)
        assert (target / "pipeline_status.json").is_file()
        assert sp.import_actual_gpt_decisions(target, filled)[0] == 0
        assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


@pytest.mark.parametrize("saved_proof", [False, True])
def test_expected_export_without_local_pdf_preserves_same_session_work(saved_proof, tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = _pdfs(tmp_path)
    source, target = tmp_path / "source", tmp_path / "target"
    with patch.object(sp, "decode", side_effect=_synthetic_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], source, defer_excel_reports=True)
        sp.run_pipeline_pdfs([second], target, defer_excel_reports=True)
        if saved_proof:
            portability.prepare_portable_project(source)
        away = tmp_path / "source-pdf-away.pdf"
        first.rename(away)
        exported = sp.export_pending_for_gpt(source)
        workbook = load_workbook(exported)
        assert (portability.EXCEL_PROOF_SHEET in workbook.sheetnames) is saved_proof
        sheet = workbook["待判定候選"]
        headers = [cell.value for cell in sheet[1]]
        sheet.cell(2, headers.index("action") + 1, "確認非校對範圍")
        sheet.cell(2, headers.index("exclusion_reason") + 1, "source local visual")
        sheet.cell(2, headers.index("exclusion_evidence") + 1, "source PDF page 1")
        filled = tmp_path / "filled-expected.xlsx"
        workbook.save(filled)
        workbook.close()
        assert sp.import_gpt_decisions(source, filled)[0] == 1
        if saved_proof:
            assert sp.import_gpt_decisions(target, filled)[0] == 1
        else:
            before = (target / "人工判定資料庫.json").read_bytes()
            with pytest.raises(ValueError, match="證據|session|SHA|PDF"):
                sp.import_gpt_decisions(target, filled)
            assert (target / "人工判定資料庫.json").read_bytes() == before
        assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


def test_existing_target_event_cannot_bypass_six_gate_actual_evidence_match(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = tmp_path / "A.pdf", tmp_path / "B.pdf"
    small_pdf(first, title="A")
    small_pdf(second, title="B")
    a, b = tmp_path / "source", tmp_path / "target"
    a_manifest, _ = small_project(a, first, session="A")
    b_manifest, b_db = small_project(b, second, session="B")
    for root, manifest in ((a, a_manifest), (b, b_manifest)):
        record = manifest["records"][0]
        record.update(state="DIFFERENCE_PENDING_CONFIRMATION",
                      expected_set=["ㄐㄧㄠˇ"], expected_evidence="independent rule",
                      context_evidence="角@1")
        if root == b:
            record["actual_evidence"] = "different B glyph evidence"
        sp.seal_manifest(manifest)
        sp.json_save(root / "校對工作階段.json", manifest)
    b_id = b_manifest["records"][0]["review_id"]
    b_db["events"][b_id] = {"action": "確認非校對範圍",
                             "exclusion_reason": "B source scope", "exclusion_evidence": "B page"}
    sp.json_save(b / "人工判定資料庫.json", b_db)
    assert not any(row.get("review_event_replay_status") for row in sp.materialize_ledger(b_manifest, b_db))
    exported = sp.export_pending_for_gpt(a)
    workbook = load_workbook(exported)
    sheet = workbook["待判定候選"]
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(2, headers.index("action") + 1, "確認現版差異")
    for gate in sp.CONFIRMATION_GATES:
        sheet.cell(2, headers.index(gate) + 1, "Y")
    filled = tmp_path / "confirm.xlsx"
    workbook.save(filled)
    workbook.close()
    before = (b / "人工判定資料庫.json").read_bytes()
    with pytest.raises(ValueError, match="actual/expected 證據|當前 actual"):
        sp.import_gpt_decisions(b, filled)
    assert (b / "人工判定資料庫.json").read_bytes() == before
    assert not (b / portability.INCOMPLETE_FILE).exists()


def test_review_app_reopens_continued_project(tmp_path, monkeypatch):
    try:
        root = tk.Tk()
    except tk.TclError as exc:
        pytest.skip(f"Tk display unavailable: {exc}")
    root.withdraw()
    try:
        monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
        first, second = _pdfs(tmp_path)
        a, b = tmp_path / "project-a", tmp_path / "project-b"
        with patch.object(sp, "decode", side_effect=_synthetic_decode), \
             patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
             patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
            sp.run_pipeline_pdfs([first], a, defer_excel_reports=True)
            manifest = sp.json_load_strict(a / "校對工作階段.json")
            ReviewSaveService(a).save_event(
                manifest["records"][0]["review_id"],
                {"action": "確認非校對範圍", "exclusion_reason": "synthetic fixture",
                 "exclusion_evidence": "PDF page 1 visual position 1"},
            )
            portability.prepare_portable_project(a)
            first.rename(tmp_path / "pdf-unavailable-on-b.pdf")
            portability.continue_project(a, second, b)
            window = tk.Toplevel(root)
            try:
                app = review_gui.ReviewApp(window, b)
                window.update()
                assert len(app.manifest["records"]) == 2
                assert len(app.db["events"]) == 1
                assert app.image.photo is not None
                assert app.current()["source_row_number"] == 3
            finally:
                window.destroy()
    finally:
        root.destroy()


def test_real_pipeline_conflict_stays_pending_until_new_local_adjudication(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = _pdfs(tmp_path)
    third = tmp_path / "download-c.pdf"
    with fitz.open(first) as document:
        document.set_metadata({"title": "download C"})
        document.save(third, deflate=True, garbage=4)
    assert portability._page_signatures(first) == portability._page_signatures(third)
    a, b, merged = [tmp_path / name for name in ("project-a", "project-b", "merged")]
    with patch.object(sp, "decode", side_effect=_synthetic_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], a, defer_excel_reports=True)
        sp.run_pipeline_pdfs([second], b, defer_excel_reports=True)
        a_manifest = sp.json_load_strict(a / "校對工作階段.json")
        b_manifest = sp.json_load_strict(b / "校對工作階段.json")
        ReviewSaveService(a).save_event(
            a_manifest["records"][0]["review_id"],
            {"action": "確認非校對範圍", "exclusion_reason": "A local visual decision",
             "exclusion_evidence": "A PDF page 1"},
        )
        ReviewSaveService(b).save_event(
            b_manifest["records"][0]["review_id"],
            {"action": "確認非校對範圍", "exclusion_reason": "B different decision",
             "exclusion_evidence": "B PDF page 1"},
        )
        result = portability.merge_projects([a, b], third, merged)
        assert len(result["conflicts"]) == 1
        manifest = sp.json_load_strict(merged / "校對工作階段.json")
        db = sp.json_load_strict(merged / "人工判定資料庫.json")
        conflicted_id = manifest["records"][0]["review_id"]
        assert conflicted_id not in db["events"]
        assert sp.materialize_ledger(manifest, db)[0]["active_review"] is True
        assert len(sp.json_load_strict(merged / "待人工確認.json")["pending"]) == 2
        assert sp.json_load_strict(merged / "pipeline_status.json")["status"] != "PROOFREAD_COMPLETE"
        workbook = load_workbook(merged / "注音校對_最終報告.xlsx", read_only=True)
        try:
            assert not any(row[0] == "人工判定" for row in workbook["人工與公司規則"].values)
        finally:
            workbook.close()
        receipt = sp.json_load_strict(merged / portability.CONFLICT_FILE)
        ReviewSaveService(merged).save_event(
            conflicted_id,
            {"action": "確認非校對範圍", "exclusion_reason": "new local check",
             "exclusion_evidence": "C PDF page 1 direct visual"},
        )
        resolved = sp.json_load_strict(merged / "人工判定資料庫.json")
        assert resolved["events"][conflicted_id]["portability_conflict_resolution"]["original_conflicts"] == receipt["conflicts"]
        sp.save_pending_json(merged, manifest, resolved)
        sp.generate_report(merged, manifest, resolved)
        assert len(sp.json_load_strict(merged / "待人工確認.json")["pending"]) == 1
        assert sp.json_load_strict(merged / portability.CONFLICT_FILE) == receipt
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))
