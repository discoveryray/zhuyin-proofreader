"""Bounded 1.2 source-drawing contract; no expected or reusable actual truth.

The small generated PDF substitutes only the fixed audited PDF digest, allowing
the exact binding/real-byte guard to run without shipping教材. It does not claim
that synthetic font resources prove the two real clipped drawings; those remain
the independent, retained source audit and isolated formal-entry acceptance.
"""
import copy
import base64
import csv
import hashlib
import json
import zipfile
import zlib
from unittest.mock import patch

import fitz
import pytest
from openpyxl import load_workbook

import actual_render_geometry as geometry
import actual_review as ar
import pdf_portability as portable
import standalone_proofread as sp
from tests.test_actual_review_v540 import entry, make_pdf


META = dict(version="5.4.0", session_id="sess", session_schema_version="2.5.0",
            workbook_schema_version="2.5.0", review_id_schema_version="2.5.0")
# Independent recorded source values, deliberately not imported from allowlist.
FIXED = [
    ("occ_04baa04670870eaf8b222b2851434d15ead10a5fd6329919df546efd332d3fa2",
     "rev_9dd88223dd16012ccd76670eeb980f6ab5c4d3e35c0add6486d235899228bcd3",
     (271.002, -20.159, 285.219, -5.941), 256, 29856, 15100,
     "2d621e553dc7f804064dd5beb8f5e682edb213a51a47abc6ca72001a5d042e2e"),
    ("occ_c299a82eb43907d3df41b0bde08d6796d0b4eb5c502f6f1cb4a6780a29315331",
     "rev_c352f3a5ee8b2cbe2117642c74df52fbdc51872b17d16f33316da11e9eaee255",
     (110.926, -5.036, 125.143, 9.181), 276, 19586, 14310,
     "e6bd2a27ab4862895519a59322786f7d776bf8e200d29c9518b54764ba07be25"),
]


@pytest.fixture
def clipped(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "localappdata"))
    pdf = tmp_path / "drawing.pdf"
    doc = fitz.open()
    for i in range(4):
        page = doc.new_page(width=606.614, height=756.85)
        page.insert_text((60, 100), f"visible page {i+1}")
    doc.save(pdf)
    doc.close()
    sha = hashlib.sha256(pdf.read_bytes()).hexdigest()
    monkeypatch.setattr(geometry, "SOURCE_CLIPPED_PDF_SHA256", sha, raising=False)
    rows = []
    for oid, rid, bbox, xref, gid, component, component_sha in FIXED:
        bounds = dict(zip(("x0", "y0", "x1", "y1"), bbox))
        source = dict(occurrence_id=oid, review_id=rid, pdf_sha256=sha, font_xref=xref, **bounds)
        source.update({"實體頁碼": 4, "glyph_id_字形索引": gid, "注音元件ID": component,
                       "TTF字形SHA256": component_sha})
        rows.append(dict(occurrence_id=oid, review_id=rid, pdf=str(pdf), pdf_name=pdf.name,
                         pdf_sha256=sha, physical_page=4, printed_page="4", char="定位標籤",
                         font_xref=xref, glyph_id=gid, zhuyin_component_id=component,
                         state="ACTUAL_DECODE_ERROR", actual="", actual_evidence="unresolved",
                         source_record=source, **bounds))
    return pdf, rows


def export(root, rows):
    ar.export_actual_review_package(root, rows, **META)
    return root / "actual待判定_GPT包" / "actual待判定_給GPT.xlsx"


def edit_row(workbook, values, *, index=2):
    sheet = workbook["actual待判定"]
    headers = [cell.value for cell in sheet[1]]
    for key, value in values.items():
        sheet.cell(index, headers.index(key) + 1, value)


def test_bounded_export_keeps_all_ids_and_only_readable_images(clipped, tmp_path):
    pdf, blocked = clipped
    rows = list(blocked)
    for group in range(50):
        for sample in range(2 if group < 32 else 1):
            oid = "occ_" + hashlib.sha256(f"{group}-{sample}".encode()).hexdigest()
            source_sha = hashlib.sha256(f"group{group}".encode()).hexdigest()
            row = entry(oid, f"r{group}-{sample}", pdf, 60 + sample * 40,
                        state="PASS" if sample else "ACTUAL_DECODE_ERROR",
                        source={"TTF字形SHA256": source_sha})
            row["pdf_sha256"] = hashlib.sha256(pdf.read_bytes()).hexdigest()
            rows.append(row)
    original = copy.deepcopy(rows)
    groups = ar.build_actual_review_groups(rows)
    expected_ids = [member["occurrence_id"] for group in groups for member in group["members"][:2]]
    xlsx = export(tmp_path, rows)
    meta, worksheet = ar._load_sheet_rows(xlsx, "actual待判定")
    assert len(groups) == len(worksheet) == 52
    assert len(expected_ids) == 84
    assert meta["readable_sample_count"] == 82 and meta["blocked_sample_count"] == 2
    assert [row[k] for row in worksheet for k in ("sample_a_occurrence_id", "sample_b_occurrence_id") if row[k]] == expected_ids
    assert rows == original
    assert len(list(xlsx.parent.glob("images/*.png"))) == 164
    assert ar.load_actual_sample_profile(xlsx) == ar.validate_actual_sample_profile(xlsx, groups, tmp_path)
    for row in worksheet:
        if row["sample_a_occurrence_id"] in {item[0] for item in FIXED}:
            assert (row["decision"], row["actual_reading"], row["confidence"], row["sample_a_checked"], row["sample_a_image"]) == ("UNRESOLVED", None, None, "N", None)
    with zipfile.ZipFile(tmp_path / "actual待判定_GPT包.zip") as archive:
        assert len([name for name in archive.namelist() if name.endswith(".png")]) == 164
        assert "82 可判定" in archive.read("README.txt").decode("utf-8")


@pytest.mark.parametrize("part,key", [
    ("entry", key) for key in ("occurrence_id", "review_id", "pdf_sha256", "physical_page", "font_xref", "glyph_id", "zhuyin_component_id", "x0")
] + [("source", key) for key in ("occurrence_id", "review_id", "pdf_sha256", "實體頁碼", "font_xref", "glyph_id_字形索引", "注音元件ID", "TTF字形SHA256", "y1")])
def test_known_clipped_binding_missing_field_rejects(clipped, tmp_path, part, key):
    pdf, rows = clipped
    row = copy.deepcopy(rows[1])  # partial page overlap is still fully clipped
    assert ar.actual_sample_availability(row, tmp_path) == "SOURCE_DRAWING_CLIPPED"
    target = row if part == "entry" else row["source_record"]
    del target[key]
    with pytest.raises((ValueError, FileNotFoundError)):
        ar.actual_sample_availability(row, tmp_path)


def test_clipped_bytes_spoof_unknown_box_and_changed_normal_reject(clipped, tmp_path):
    pdf, rows = clipped
    pdf.write_bytes(pdf.read_bytes() + b"changed")
    with pytest.raises((ValueError, FileNotFoundError)):
        ar.actual_sample_availability(rows[0], tmp_path)
    unrelated = entry("unknown", "review", pdf, 60)
    unrelated.update(x0=110.926, y0=-5.036, x1=125.143, y1=9.181)
    with pytest.raises(ValueError, match="未知"):
        ar.actual_sample_availability(unrelated, tmp_path)
    normal = entry("normal", "review", pdf, 60)
    normal["pdf_sha256"] = "f" * 64
    with pytest.raises((ValueError, FileNotFoundError)):
        ar.actual_sample_availability(normal, tmp_path)


@pytest.mark.parametrize("change", ["verified", "reading", "Y", "flag", "downgrade", "profile_deleted", "chunk_count", "plan_deleted"])
def test_clipped_workbook_forgery_zero_transaction(clipped, tmp_path, change):
    _, rows = clipped
    xlsx = export(tmp_path, rows)
    workbook = load_workbook(xlsx)
    if change in ("verified", "reading", "Y", "flag"):
        edit_row(workbook, {"verified": {"decision": "VERIFIED"}, "reading": {"actual_reading": "ㄇㄥˊ"},
                            "Y": {"sample_a_checked": "Y"}, "flag": {"sample_a_availability": "READABLE"}}[change])
    elif change == "downgrade":
        for row in workbook["匯入中繼資料"]:
            if row[0].value == "actual_review_schema_version":
                row[1].value = "1.1"
    elif change == "profile_deleted":
        del workbook[ar.ACTUAL_SAMPLE_PROFILE_SHEET]
    elif change == "chunk_count":
        workbook[ar.ACTUAL_SAMPLE_PROFILE_SHEET].cell(3, 2, 999)
    else:
        workbook["actual待判定"].delete_rows(2)
    workbook.save(xlsx)
    workbook.close()
    with patch.object(ar, "apply_direct_visual_actual_batch") as transaction:
        with pytest.raises(ValueError):
            ar.import_actual_review_workbook(tmp_path / "actual", xlsx, ar.build_actual_review_groups(rows), expected_metadata=META)
    transaction.assert_not_called()
    assert not list((tmp_path / "localappdata").rglob("*.sqlite*"))


def test_self_sealed_readable_profile_cannot_override_current_source(clipped, tmp_path):
    _, rows = clipped
    xlsx = export(tmp_path, rows)
    workbook = load_workbook(xlsx)
    profile = ar.load_actual_sample_profile(xlsx)
    for group in profile["groups"]:
        group["samples"]["a"]["availability"] = "READABLE"
        group["samples"]["a"]["reason"] = ""
    del workbook[ar.ACTUAL_SAMPLE_PROFILE_SHEET]
    sheet = workbook["匯入中繼資料"]
    for i in range(sheet.max_row, 1, -1):
        if sheet.cell(i, 1).value in ar.ACTUAL_SAMPLE_PROFILE_META:
            sheet.delete_rows(i)
    ar._write_actual_sample_profile(workbook, profile)
    edit_row(workbook, dict(sample_a_availability="READABLE", source_availability_reason="",
                            decision="VERIFIED", actual_reading="ㄇㄥˊ", confidence="高", sample_a_checked="Y"))
    workbook.save(xlsx)
    workbook.close()
    with pytest.raises(ValueError, match="可信"):
        ar.validate_actual_sample_profile(xlsx, ar.build_actual_review_groups(rows), tmp_path)


def make_historical(workbook):
    del workbook[ar.ACTUAL_SAMPLE_PROFILE_SHEET]
    sheet = workbook["匯入中繼資料"]
    for i in range(sheet.max_row, 1, -1):
        key = sheet.cell(i, 1).value
        if key in ar.ACTUAL_SAMPLE_PROFILE_META:
            sheet.delete_rows(i)
        elif key == "actual_review_schema_version":
            sheet.cell(i, 2, "1.1")
    sheet = workbook["actual待判定"]
    for i in range(sheet.max_column, 0, -1):
        if sheet.cell(1, i).value in {"sample_a_availability", "sample_b_availability", "source_availability_reason"}:
            sheet.delete_cols(i)


def test_historical_complete_samples_adapter_and_clipped_downgrade(clipped, tmp_path):
    pdf, blocked = clipped
    normal = [entry("normal", "review", pdf, 60)]
    xlsx = export(tmp_path, normal)
    workbook = load_workbook(xlsx)
    make_historical(workbook)
    workbook.save(xlsx)
    workbook.close()
    assert ar.load_actual_sample_profile(xlsx) is None
    assert ar.validate_actual_sample_profile(xlsx, ar.build_actual_review_groups(normal), tmp_path)
    blocked_xlsx = export(tmp_path, blocked)
    workbook = load_workbook(blocked_xlsx)
    make_historical(workbook)
    workbook.save(blocked_xlsx)
    workbook.close()
    with pytest.raises(ValueError, match="historical_complete"):
        ar.validate_actual_sample_profile(blocked_xlsx, ar.build_actual_review_groups(blocked), tmp_path)


def test_normal_content_proof_deleted_rejects_before_same_session_recovery(tmp_path, monkeypatch):
    pdf = tmp_path / "normal.pdf"
    make_pdf(pdf)
    rows = [entry("normal", "review", pdf, 60)]
    xlsx = export(tmp_path, rows)
    manifest = dict(session_id="sess", pdfs=[{"pdf_sha256": hashlib.sha256(pdf.read_bytes()).hexdigest()}])
    (tmp_path / "校對工作階段.json").write_text(json.dumps(manifest), encoding="utf-8")
    with patch.object(sp, "validate_manifest_integrity"), patch.object(sp, "validate_output_artifact_hashes"), \
         patch.object(sp, "materialize_ledger", return_value=rows), \
         patch.object(sp, "recover_pending_project_actual_write") as recovery, \
         patch.object(sp, "load_or_initialize_db") as initialize:
        with pytest.raises(ValueError, match="proof deleted"):
            sp._import_same_session_actual_snapshot(tmp_path, xlsx, _original_xlsx=xlsx, _snapshot_sha=hashlib.sha256(xlsx.read_bytes()).hexdigest())
    recovery.assert_not_called()
    initialize.assert_not_called()


@pytest.mark.parametrize("already", [False, True])
def test_csv_clipped_rejects_before_recovery_initialize_and_shortcut(clipped, tmp_path, already):
    _, rows = clipped
    row = rows[0]
    if already:
        row["actual"] = "ㄔㄠˊ"
        row["actual_evidence"] = "local prior evidence"
    manifest = dict(session_id="sess")
    (tmp_path / "校對工作階段.json").write_text(json.dumps(manifest), encoding="utf-8")
    csv_path = tmp_path / "decisions.csv"
    decision = dict(occurrence_id=row["occurrence_id"], review_id=row["review_id"], pdf_name=row["pdf_name"],
                    physical_page=4, char=row["char"], exported_actual="", exported_actual_evidence_sha256="",
                    verified_actual="ㄔㄠˊ", visual_confirmation="Y", note="")
    with csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(decision))
        writer.writeheader()
        writer.writerow(decision)
    with patch.object(sp, "validate_manifest_integrity"), patch.object(sp, "validate_output_artifact_hashes"), \
         patch.object(sp, "_validate_gpt_bundle_metadata"), patch.object(sp, "materialize_ledger", return_value=rows), \
         patch.object(sp, "recover_pending_project_actual_write") as recovery, \
         patch.object(sp, "load_or_initialize_db") as initialize, patch.object(sp, "generate_report") as report:
        with pytest.raises(ValueError, match="SOURCE_DRAWING_CLIPPED"):
            sp.import_actual_occurrence_decisions(tmp_path, csv_path, package_meta={})
    recovery.assert_not_called()
    initialize.assert_not_called()
    report.assert_not_called()
    assert not list((tmp_path / "localappdata").rglob("*.sqlite*"))


def test_portable_legacy_shortcut_cannot_hide_deleted_12_profile(clipped, tmp_path):
    _, rows = clipped
    xlsx = export(tmp_path, rows)
    workbook = load_workbook(xlsx)
    del workbook[ar.ACTUAL_SAMPLE_PROFILE_SHEET]
    workbook.save(xlsx)
    workbook.close()
    manifest = dict(session_id="target", pdfs=[])
    with patch.object(portable, "_load_project", return_value=(manifest, {})), \
         patch.object(portable, "actual_excel_conflict_state", return_value=[]), \
         patch.object(portable, "_actual_transfer_state", return_value={"status": "NONE"}), \
         patch.object(portable, "_legacy_actual_already_imported", return_value=True) as shortcut:
        with pytest.raises(ValueError, match="profile"):
            portable.import_actual_excel(tmp_path, xlsx)
    shortcut.assert_not_called()


def seal_content_proof(workbook, manifest, pages):
    """Fixture's real-page signatures; ordinary sealed test source, no decoder."""
    info = manifest["pdfs"][0]
    payload = dict(version=portable.EXCEL_PROOF_VERSION, kind="actual", exported_at="fixture",
                   manifest=manifest, db={}, attestations={},
                   bound_rows=portable._excel_bound_rows(workbook, "actual"),
                   actual_profile_metadata=ar._actual_profile_metadata(workbook),
                   proof=dict(version=portable.PROOF_VERSION, session_id=manifest["session_id"],
                              manifest_integrity_sha256=manifest["manifest_integrity_sha256"],
                              pdfs=[dict(pdf_sha256=info["pdf_sha256"], pdf_name=info["pdf_name"], pages=pages)]))
    raw = portable._canonical(payload)
    packed = base64.b64encode(zlib.compress(raw)).decode("ascii")
    chunks = [packed[i:i+30000] for i in range(0, len(packed), 30000)]
    sheet = workbook.create_sheet(portable.EXCEL_PROOF_SHEET)
    sheet.append(["version", portable.EXCEL_PROOF_VERSION])
    sheet.append(["sha256", hashlib.sha256(raw).hexdigest()])
    sheet.append(["chunk_count", len(chunks)])
    for index, chunk in enumerate(chunks, 1):
        sheet.append([index, chunk])
    return payload


@pytest.mark.parametrize("source_present", [True, False])
def test_clipped_portable_original_bytes_required_for_different_sha(clipped, tmp_path, source_present):
    source_pdf, rows = clipped
    xlsx = export(tmp_path, rows)
    source_sha, pages = portable._pdf_content_snapshot(source_pdf)
    target_pdf = tmp_path / "target.pdf"
    target_pdf.write_bytes(source_pdf.read_bytes() + b"\n% byte-only difference\n")
    target_sha = hashlib.sha256(target_pdf.read_bytes()).hexdigest()
    source_manifest = dict(session_id="sess", manifest_integrity_sha256="source-seal",
                           pdfs=[dict(pdf=str(source_pdf), pdf_name=source_pdf.name, pdf_sha256=source_sha)], records=rows)
    target_rows = copy.deepcopy(rows)
    for row in target_rows:
        row.update(pdf=str(target_pdf), pdf_name=target_pdf.name, pdf_sha256=target_sha)
    target_manifest = dict(session_id="target", manifest_integrity_sha256="target-seal",
                           pdfs=[dict(pdf=str(target_pdf), pdf_name=target_pdf.name, pdf_sha256=target_sha)], records=target_rows)
    workbook = load_workbook(xlsx)
    seal_content_proof(workbook, source_manifest, pages)
    workbook.save(xlsx)
    workbook.close()
    # Real full-page raster/text mapping proves displayed equivalence. It cannot
    # stand in for original bytes for the finite source-drawing adapter.
    mapping, _ = portable._match_pdfs({"pdfs": [dict(pdf_sha256=source_sha, pdf_name=source_pdf.name, pages=pages)]}, source_manifest, target_manifest, tmp_path)
    assert mapping == {source_sha: target_sha}
    if not source_present:
        source_pdf.unlink()
    if source_present:
        assert portable._guard_actual_excel_sources(xlsx, rows, tmp_path, target_manifest, mapping)
    else:
        with pytest.raises(ValueError, match="SOURCE_DRAWING_CLIPPED_v1"):
            portable._guard_actual_excel_sources(xlsx, rows, tmp_path, target_manifest, mapping)


@pytest.mark.parametrize("forgery", ["decision", "metadata", "downgrade"])
def test_portable_clipped_and_profile_proof_forgery_zero_writes(clipped, tmp_path, forgery):
    pdf, rows = clipped
    sha, pages = portable._pdf_content_snapshot(pdf)
    manifest = dict(session_id="sess", manifest_integrity_sha256="fixture-seal",
                    pdfs=[dict(pdf=str(pdf), pdf_name=pdf.name, pdf_sha256=sha)], records=rows)
    xlsx = export(tmp_path, rows)
    workbook = load_workbook(xlsx)
    seal_content_proof(workbook, manifest, pages)
    if forgery == "decision":
        edit_row(workbook, dict(decision="VERIFIED", actual_reading="ㄇㄥˊ", confidence="高", sample_a_checked="Y"))
    elif forgery == "metadata":
        for row in workbook["匯入中繼資料"]:
            if row[0].value == "actual_sample_profile_sha256":
                row[1].value = "f" * 64
    else:
        make_historical(workbook)  # carried 1.2 proof cannot become historical
    workbook.save(xlsx)
    workbook.close()
    with patch.object(portable, "_load_project", return_value=(manifest, {})), \
         patch.object(portable, "actual_excel_conflict_state", return_value=[]), \
         patch.object(portable, "_actual_transfer_state", return_value={"status": "NONE"}), \
         patch.object(sp, "validate_manifest_integrity"), patch.object(sp, "materialize_ledger", return_value=rows), \
         patch.object(portable, "resume_actual_excel_project") as recovery, \
         patch.object(ar, "_write_csv") as writes, patch.object(sp, "json_save") as dbwrite:
        with pytest.raises(ValueError):
            portable.import_actual_excel(tmp_path, xlsx)
    writes.assert_not_called()
    dbwrite.assert_not_called()
    recovery.assert_not_called()
    assert not list((tmp_path / "localappdata").rglob("*.sqlite*"))


def test_local_prewrite_rechecks_real_source_before_initialization(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "localappdata"))
    pdf = tmp_path / "normal.pdf"
    make_pdf(pdf)
    row = entry("normal", "review", pdf, 60)
    row["pdf_sha256"] = hashlib.sha256(pdf.read_bytes()).hexdigest()
    xlsx = export(tmp_path, [row])
    workbook = load_workbook(xlsx)
    edit_row(workbook, dict(decision="VERIFIED", actual_reading="ㄐㄩㄝˊ", confidence="高", sample_a_checked="Y"))
    workbook.save(xlsx)
    workbook.close()
    calls = []
    def late_source_change():
        calls.append(True)
        if len(calls) == 2:  # after initial guard, under actual transaction lock
            pdf.write_bytes(pdf.read_bytes() + b"changed")
    with patch.object(ar, "apply_verified_actual_group") as apply, patch.object(sp, "initialize_project_actual_evidence") as initialize:
        with pytest.raises((ValueError, FileNotFoundError)):
            ar.import_actual_review_workbook(tmp_path / "actual", xlsx, ar.build_actual_review_groups([row]),
                                            expected_metadata=META, initialize_evidence=initialize,
                                            prewrite_guard=late_source_change)
    assert len(calls) == 2
    initialize.assert_not_called()
    apply.assert_not_called()
    assert not list((tmp_path / "localappdata").rglob("*.sqlite*"))


def test_csv_empty_rows_source_guard_precedes_no_row_report(clipped, tmp_path):
    pdf, rows = clipped
    csv_path = tmp_path / "empty.csv"
    csv_path.write_text("occurrence_id,review_id,pdf_name,physical_page,char,exported_actual,exported_actual_evidence_sha256,verified_actual,visual_confirmation,note\n", encoding="utf-8-sig")
    (tmp_path / "校對工作階段.json").write_text('{"session_id":"sess"}', encoding="utf-8")
    pdf.write_bytes(pdf.read_bytes() + b"changed")
    with patch.object(sp, "validate_manifest_integrity"), patch.object(sp, "validate_output_artifact_hashes"), \
         patch.object(sp, "_validate_gpt_bundle_metadata"), patch.object(sp, "materialize_ledger", return_value=rows), \
         patch.object(sp, "recover_pending_project_actual_write") as recovery, \
         patch.object(sp, "load_or_initialize_db") as initialize, patch.object(sp, "generate_report") as report:
        with pytest.raises((ValueError, FileNotFoundError)):
            sp.import_actual_occurrence_decisions(tmp_path, csv_path, package_meta={})
    recovery.assert_not_called()
    initialize.assert_not_called()
    report.assert_not_called()


def test_zip_clipped_guard_precedes_expected_dry_run(clipped, tmp_path):
    _, rows = clipped
    row = rows[0]
    (tmp_path / "校對工作階段.json").write_text('{"session_id":"sess"}', encoding="utf-8")
    decision = dict(occurrence_id=row["occurrence_id"], review_id=row["review_id"], pdf_name=row["pdf_name"],
                    physical_page=4, char=row["char"], exported_actual="", exported_actual_evidence_sha256="",
                    verified_actual="ㄔㄠˊ", visual_confirmation="Y", note="")
    import io
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=list(decision))
    writer.writeheader()
    writer.writerow(decision)
    bundle = tmp_path / "bundle.zip"
    with zipfile.ZipFile(bundle, "w") as archive:
        archive.writestr("actual_occurrence_decisions.csv", stream.getvalue().encode("utf-8-sig"))
        archive.writestr("expected.xlsx", b"fixture")
    with patch.object(sp, "validate_manifest_integrity"), patch.object(sp, "validate_output_artifact_hashes"), \
         patch.object(sp, "_validate_gpt_bundle_metadata"), patch.object(sp, "_load_gpt_bundle_metadata", return_value={"expected_workbook": "expected.xlsx"}), \
         patch.object(sp, "materialize_ledger", return_value=rows), patch.object(sp, "import_gpt_decisions") as expected, \
         patch.object(sp, "import_actual_occurrence_decisions") as actual:
        with pytest.raises(ValueError, match="SOURCE_DRAWING_CLIPPED"):
            sp.import_gpt_decision_bundle(tmp_path, bundle)
    expected.assert_not_called()
    actual.assert_not_called()
