"""Expected-only operation indexes retain complete identities and durable views."""
import copy
from pathlib import Path
from unittest.mock import patch
import pytest
import pdf_portability as pp
import standalone_proofread as sp
from tests.test_excel_text_storage import isolated_pipeline, exported, fill
from tests.test_pdf_portability_replay import edited_excel


def identity(row=2, **extra):
    return {'target_review_id': 'target', 'source_excel_sha256': 'a'*64,
            'source_excel_row': row, 'session_id': 'source', 'review_id': 'review',
            'occurrence_id': 'occurrence', 'full_source_payload': {'independent': '=source'}, **extra}


def event(**extra):
    return {'action': '解決expected證據', 'expected_set': ['ㄅ'],
            'expected_evidence': 'independent rule', 'context_evidence': 'source context', **extra}


def old_receipt_scan(records, incoming, answer):
    return any(item.get('source_identity') == incoming
               and pp._excel_expected_decision(item.get('source_event', {})) == pp._excel_expected_decision(answer)
               for item in records)


@pytest.mark.parametrize('contract', ['legacy', 'v2'])
def test_history_index_matches_original_full_replay_oracle(contract):
    source=identity();target=identity(3);duplicate=identity(4)
    source_event=event(portability_source=source)
    target_event=event(expected_evidence='retained target original',portability_source=target,
                       portability_duplicate_sources=[duplicate])
    record={'target_review_id':'target','source_identity':source,'source_event':source_event,
            'target_identity':target,'target_event':target_event}
    if contract=='v2':
        del record['target_event'];record.update(contract=pp.EXPECTED_BUSINESS_CONTRACT,
                                               target_basis={'event':target_event})
    records=[record];index=pp._ExpectedExcelHistory(records,[]);original=copy.deepcopy(records)
    cases=[(source,event(),target_event), (target,source_event,target_event),
           (duplicate,source_event,target_event),
           (identity(full_source_payload={'independent':'different identity'}),event(),target_event),
           (source,event(expected_evidence='changed answer'),event(expected_evidence='local adjudication')),
           (identity(target_review_id='other'),event(),target_event)]
    for incoming,source_answer,target_answer in cases:
        assert index.historical_replay(incoming,source_answer,target_answer) == pp._historical_import_replay(
            records,incoming,source_answer,target_answer,decision_payload=pp._excel_expected_decision)
    assert records==original
    records[0]['source_event']['expected_evidence']='caller mutation'
    assert index.historical_replay(source,event(),target_event)


def test_receipt_index_matches_original_conflict_and_supplement_scan():
    source=identity();answer=event();supplements=[{'source_identity':source,'source_event':answer}]
    conflicts=[{'target_review_id':'target','source_identity':identity(3),
                'source_event':event(note='retained conflict source')}]
    index=pp._ExpectedExcelHistory(conflicts,supplements)
    for incoming,proposed in [(source,answer),(identity(3),event(note='retained conflict source')),
            (identity(session_id='different full identity'),answer),(source,event(note='new answer')),
            (source,event(expected_evidence='local adjudication'))]:
        assert index.receipt_replay(incoming,proposed)==old_receipt_scan([*supplements,*conflicts],incoming,proposed)
    # The original validated receipt's mutable reading list may not mutate the
    # operation's fixed index when a caller later changes its own parsed view.
    supplements[0]['source_event']['expected_set'].append('ㄆ')
    assert index.receipt_replay(source,event())
    assert not index.receipt_replay(source,event(expected_set=['ㄅ','ㄆ']))


def prepare(isolated_pipeline, monkeypatch):
    import tests.test_excel_text_storage as fixture
    monkeypatch.setattr(fixture,'TEXTS',fixture.TEXTS[:2])
    root=isolated_pipeline('expected');path,sheet=exported(root,'expected')
    wb,_,_=fill(path,sheet,'expected');wb.save(path);wb.close()
    return root,path


def test_exact_reimport_keeps_db_bytes_and_uses_durable_report_only_publication(isolated_pipeline,monkeypatch):
    root,path=prepare(isolated_pipeline,monkeypatch)
    sp.import_gpt_decisions(root,path);before=(root/'人工判定資料庫.json').read_bytes()
    saved=[];save=sp.json_save;publish=pp._publish_portable_outputs;published=[]
    def capture(file,*args,**kwargs):
        saved.append(Path(file));return save(file,*args,**kwargs)
    def report(*args,**kwargs):
        assert (root/pp.INCOMPLETE_FILE).exists()
        published.append(True);return publish(*args,**kwargs)
    monkeypatch.setattr(sp,'json_save',capture);monkeypatch.setattr(pp,'_publish_portable_outputs',report)
    result=sp.import_gpt_decisions(root,path)
    assert result.status['duplicates']==2 and result[0]==0
    assert root/'人工判定資料庫.json' not in saved
    assert root/pp.INCOMPLETE_FILE in saved and published==[True]
    assert (root/'人工判定資料庫.json').read_bytes()==before and not (root/pp.INCOMPLETE_FILE).exists()


def test_new_matching_source_persists_receipt_without_rewriting_db(isolated_pipeline,monkeypatch,tmp_path):
    root,path=prepare(isolated_pipeline,monkeypatch);sp.import_gpt_decisions(root,path)
    before=(root/'人工判定資料庫.json').read_bytes()
    other=edited_excel(path,tmp_path/'another-source.xlsx',title='new exact source bytes')
    saved=[];save=sp.json_save
    def capture(file,*args,**kwargs):saved.append(Path(file));return save(file,*args,**kwargs)
    monkeypatch.setattr(sp,'json_save',capture)
    first=sp.import_gpt_decisions(root,other)
    assert first.status['consistent_supplements']==2
    assert root/pp.CONFLICT_FILE in saved and root/'人工判定資料庫.json' not in saved
    receipt=sp.json_load_strict(root/pp.CONFLICT_FILE)
    assert len(receipt['supplements'])==2
    assert all(item['source_identity']['source_excel_sha256']==pp._sha(other) for item in receipt['supplements'])
    saved.clear();second=sp.import_gpt_decisions(root,other)
    assert second.status['duplicates']==2 and root/pp.CONFLICT_FILE not in saved
    assert (root/'人工判定資料庫.json').read_bytes()==before


def test_duplicate_report_failure_preserves_original_data_and_views(isolated_pipeline,monkeypatch):
    root,path=prepare(isolated_pipeline,monkeypatch);sp.import_gpt_decisions(root,path)
    names=['人工判定資料庫.json',pp.CONFLICT_FILE,'待人工確認.json','注音校對_最終報告.xlsx','pipeline_status.json']
    before={name:(root/name).read_bytes() if (root/name).exists() else None for name in names}
    def failed(*args,**kwargs):
        (root/'待人工確認.json').write_bytes(b'partial unpublished view')
        raise OSError('required report-only publication failed')
    monkeypatch.setattr(pp,'_publish_portable_outputs',failed)
    with pytest.raises(OSError,match='report-only publication'):
        sp.import_gpt_decisions(root,path)
    assert {name:(root/name).read_bytes() if (root/name).exists() else None for name in names}==before
    assert not (root/pp.INCOMPLETE_FILE).exists()


def test_same_answer_new_source_appends_db_provenance_even_when_duplicates_nonzero(isolated_pipeline,monkeypatch,tmp_path):
    import tests.test_excel_text_storage as fixture
    from openpyxl import load_workbook
    monkeypatch.setattr(fixture,'TEXTS',fixture.TEXTS[:2]);root=isolated_pipeline('expected')
    path,_=exported(root,'expected');wb=load_workbook(path);ws=wb['待判定候選'];columns={c.value:c.column for c in ws[1]}
    for row in range(2,ws.max_row+1):
        ws.cell(row,columns['action'],'確認非校對範圍')
        ws.cell(row,columns['exclusion_reason'],'independent selected scope')
        ws.cell(row,columns['exclusion_evidence'],'PDF source position verified')
    wb.save(path);wb.close();sp.import_gpt_decisions(root,path)
    before=(root/'人工判定資料庫.json').read_bytes()
    other=edited_excel(path,tmp_path/'new-same-answer.xlsx',title='new source must remain auditable')
    saved=[];save=sp.json_save
    def capture(file,*args,**kwargs):saved.append(Path(file));return save(file,*args,**kwargs)
    monkeypatch.setattr(sp,'json_save',capture)
    result=sp.import_gpt_decisions(root,other)
    assert result.status['duplicates']==2 and root/'人工判定資料庫.json' in saved
    assert (root/'人工判定資料庫.json').read_bytes()!=before
    db=sp.json_load_strict(root/'人工判定資料庫.json')
    assert all(len(item['portability_duplicate_sources'])==1 for item in db['events'].values())
    assert all(item['portability_duplicate_sources'][0]['source_excel_sha256']==pp._sha(other) for item in db['events'].values())
