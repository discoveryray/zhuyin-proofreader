"""Real portable actual import survives a source-only teardown failure.

The decoder supplies explicit isolated readings, as in the existing portable
pipeline fixture. PDF/native source proof, Excel seals, actual transaction,
expected save, refresh and duplicate detection all execute production logic.
"""
from __future__ import annotations

from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

import fitz
import pytest
from openpyxl import load_workbook

import actual_review as actual
import actual_source_visual_safety as safety
import global_glyph_promotion as promotion
import pdf_portability as portable
import standalone_proofread as pipeline
from review_save_service import ReviewSaveService
from tests.test_pdf_portability_integration import _pdfs, _synthetic_unresolved_decode


@pytest.fixture
def real_portable_actual_import(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = _pdfs(tmp_path)
    source, target = tmp_path / "source", tmp_path / "target"
    with patch.object(pipeline, "decode", side_effect=_synthetic_unresolved_decode), \
         patch.object(pipeline, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        pipeline.run_pipeline_pdfs([first], source, defer_excel_reports=True)
        pipeline.run_pipeline_pdfs([second], target, defer_excel_reports=True)
        manifest = pipeline.json_load_strict(target / "校對工作階段.json")
        first_entry = manifest["records"][0]
        source_manifest = pipeline.json_load_strict(source / "校對工作階段.json")
        source_entry = source_manifest["records"][0]
        mapped_fields = ("physical_page", "char", "font_xref", "glyph_id", "x0", "y0", "x1", "y1")
        assert [source_entry[key] for key in mapped_fields] == [first_entry[key] for key in mapped_fields]
        original_overrides = actual._read_csv(pipeline.project_actual_evidence_root(target)
                                             / actual.OCCURRENCE_OVERRIDE_FILE, actual.OVERRIDE_HEADERS)
        # A durable expected event is independent of the incoming manual actual.
        ReviewSaveService(target).save_event(first_entry["review_id"], {
            "action": "補建expected證據", "expected_set": ["ㄐㄧㄠˇ"],
            "expected_evidence": "isolated independent expected rule",
            "context_evidence": "角色（隔離規則）",
        })
        before_db = pipeline.json_load_strict(target / "人工判定資料庫.json")
        assert before_db["events"][first_entry["review_id"]]["expected_set"] == ["ㄐㄧㄠˇ"]
        with safety.source_visual_operation():
            for entry in manifest["records"]:
                result = actual.actual_visual_source_safety(entry, target)
                assert result["location_readable"] is True
        pipeline.export_actual_pending_for_gpt(source)
        exported = source / "actual待判定_GPT包" / "actual待判定_給GPT.xlsx"
        profile = actual.load_actual_sample_profile(exported)
        samples = [group["samples"]["a"] for group in profile["groups"]
                   if group["samples"]["a"]["source"]["occurrence_id"] == source_entry["occurrence_id"]]
        assert len(samples) == 1 and samples[0]["availability"] == "READABLE"
        assert samples[0]["source"]["pdf_sha256"] == source_entry["pdf_sha256"]
        assert all(samples[0]["source"][field] == source_entry[field]
                   for field in ("physical_page", "font_xref", "glyph_id", "x0", "y0", "x1", "y1"))
        workbook = load_workbook(exported)
        try:
            assert workbook[portable.EXCEL_PROOF_SHEET].sheet_state == "veryHidden"
            sheet = workbook["actual待判定"]
            headers = [cell.value for cell in sheet[1]]
            answer_rows = [index for index in range(2, sheet.max_row + 1)
                           if sheet.cell(index, headers.index("sample_a_occurrence_id") + 1).value
                           == source_entry["occurrence_id"]]
            assert len(answer_rows) == 1
            for name, value in {"decision": "VERIFIED", "actual_reading": "ㄐㄩㄝˊ",
                                "confidence": "高", "sample_a_checked": "Y",
                                "note": "isolated PDF physical page 1"}.items():
                sheet.cell(answer_rows[0], headers.index(name) + 1, value)
            filled = tmp_path / "actual-import.xlsx"
            workbook.save(filled)
        finally:
            workbook.close()
        yield target, filled, first_entry, before_db, original_overrides
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


def _assert_durable_import(target, original_entry, before_db, original_overrides, result):
    assert isinstance(result, pipeline.ActualImportResult)
    assert result[0] == 1
    status = pipeline.json_load_strict(target / "pipeline_status.json")
    assert result[2] == Path(status["user_report"])
    # The genuine pipeline fixture intentionally defers publication. A durable
    # actual commit/refresh does not imply an Excel file was already rendered.
    assert status["excel_report_deferred"] is True
    assert not result[2].exists()
    assert result.status["project_actual_commit"] == "COMMITTED"
    assert result.status["global_promotion_delivery"]["status"] == "NO_PENDING"
    manifest, db = portable._load_project(target)
    assert db["events"] == before_db["events"]
    rows = pipeline.materialize_ledger(manifest, db)
    entry = next(row for row in rows if row["occurrence_id"] == original_entry["occurrence_id"])
    assert entry["actual"] == "ㄐㄩㄝˊ"
    assert entry["expected_set"] == ["ㄐㄧㄠˇ"]
    root = pipeline.project_actual_evidence_root(target)
    overrides = actual._read_csv(root / actual.OCCURRENCE_OVERRIDE_FILE, actual.OVERRIDE_HEADERS)
    local = [row for row in overrides if actual._override_applies_to_pdf(row, Path(original_entry["pdf"]))]
    assert len(local) == 1 and local[0]["actual_reading"] == "ㄐㄩㄝˊ"
    assert tuple(local[0][field] for field in ("pdf_contains", "pdf_excludes", "page", "target_char",
                                             "stable_key", "x0", "y0")) == actual._override_key_from_entry(original_entry)
    key = lambda row: tuple(row[field] for field in actual.OVERRIDE_HEADERS)
    assert sorted([row for row in overrides if row not in local], key=key) == sorted(original_overrides, key=key)
    assert not (target / portable.INCOMPLETE_FILE).exists()
    assert not (root / promotion.PROJECT_TRANSACTION_FILE).exists()
    return root / actual.OCCURRENCE_OVERRIDE_FILE, manifest, db


def test_actual_excel_import_normal_keeps_actual_and_expected_sources(real_portable_actual_import):
    target, workbook, first_entry, before_db, original_overrides = real_portable_actual_import
    result = pipeline.import_actual_gpt_decisions(target, workbook)
    _assert_durable_import(target, first_entry, before_db, original_overrides, result)


@contextmanager
def _owned_source_close_fault(monkeypatch):
    real_page, real_close = safety.SourceVisualOperation._page, fitz.Document.close
    owned_documents, owners, attempted = {}, {}, []

    def observe_owned_page(operation, *args, **kwargs):
        result = real_page(operation, *args, **kwargs)
        owners[id(operation)] = operation
        for _raw, document in operation._documents.values():
            owned_documents[id(document)] = document
        return result

    def close_owned_source(document):
        real_close(document)
        if owned_documents.get(id(document)) is document:
            # No ordinary source snapshot, pipeline renderer, or transaction
            # writer is altered. The real owned native document already closed.
            attempted.append(id(document))
            raise RuntimeError("isolated owned-source close failure")

    with monkeypatch.context() as injection:
        injection.setattr(safety.SourceVisualOperation, "_page", observe_owned_page)
        injection.setattr(fitz.Document, "close", close_owned_source)
        yield owned_documents, owners, attempted


def _assert_owned_cleanup(error, tracking):
    owned_documents, owners, attempted = tracking
    assert error.commit_state == "COMMITTED_REFRESH_PENDING"
    assert isinstance(error.cause, safety.SourceVisualCleanupError)
    assert "isolated owned-source close failure" in str(error.cause)
    assert owned_documents and set(attempted) == set(owned_documents)
    assert all(document.is_closed for document in owned_documents.values())
    assert owners and all(owner._closed and not owner._documents and not owner._pages
                          and not owner._sources and not owner._raws for owner in owners.values())
    assert safety._CURRENT_OPERATION.get() is None


def test_actual_excel_postcommit_source_cleanup_fault_retains_result_without_reapply(
        real_portable_actual_import, monkeypatch):
    target, workbook, first_entry, before_db, original_overrides = real_portable_actual_import
    with _owned_source_close_fault(monkeypatch) as tracking:
        with pytest.raises(safety.SourceVisualPostCommitCleanupError) as caught:
            pipeline.import_actual_gpt_decisions(target, workbook)
    error = caught.value
    _assert_owned_cleanup(error, tracking)
    override, manifest, db = _assert_durable_import(
        target, first_entry, before_db, original_overrides, error.completed_result)
    committed_bytes = override.read_bytes()
    # The public entry recognizes the already committed workbook. Observe the
    # actual writer rather than mocking duplicate detection or transaction PASS.
    with patch.object(actual, "_write_csv", wraps=actual._write_csv) as writes:
        duplicate = pipeline.import_actual_gpt_decisions(target, workbook)
    assert duplicate[0] == 0
    assert duplicate.status["project_actual_commit"] == "NO_CHANGES"
    assert not [call for call in writes.call_args_list if Path(call.args[0]) == override]
    assert override.read_bytes() == committed_bytes
    reopened, reopened_db = portable._load_project(target)
    assert reopened == manifest and reopened_db == db


def test_project_actual_import_postcommit_source_cleanup_fault_retains_result_without_reapply(
        real_portable_actual_import, monkeypatch):
    incoming, workbook, first_entry, before_db, original_overrides = real_portable_actual_import
    imported = pipeline.import_actual_gpt_decisions(incoming, workbook)
    incoming_override, _manifest, _db = _assert_durable_import(
        incoming, first_entry, before_db, original_overrides, imported)
    incoming_bytes = incoming_override.read_bytes()
    destination = incoming.parent / "source"
    # This workbook belongs to destination's own sealed session/PDF/sample A.
    # Its handwritten reading is revalidated and committed by the public local
    # path. Project transfer requires this independently checked local actual;
    # no incoming CSV or checked-source record is copied to satisfy the guard.
    locally_checked = pipeline.import_actual_gpt_decisions(destination, workbook)
    assert locally_checked[0] == 1
    assert locally_checked.status["project_actual_commit"] == "COMMITTED"
    original_manifest, original_db = portable._load_project(destination)
    assert original_manifest["records"][0]["actual"] == "ㄐㄩㄝˊ"
    assert not original_db["events"]
    root = pipeline.project_actual_evidence_root(destination)
    override = root / actual.OCCURRENCE_OVERRIDE_FILE
    original_bytes = override.read_bytes()
    portable.prepare_portable_project(incoming)
    with _owned_source_close_fault(monkeypatch) as tracking:
        with pytest.raises(safety.SourceVisualPostCommitCleanupError) as caught:
            portable.import_project_decisions(incoming, destination)
    error = caught.value
    _assert_owned_cleanup(error, tracking)
    result = error.completed_result
    assert result["imported"] == 1 and result["matched_actual_overrides"] == 1
    assert result["actual_staging"] == "NONE" and not result["conflicts"]
    # Mapping a current actual override validates both real source locations;
    # this project-transfer contract never turns that match into a new vote or
    # directly applies it to the destination. The durable mutation is expected.
    manifest, db = portable._load_project(destination)
    entry = manifest["records"][0]
    assert entry["occurrence_id"] == original_manifest["records"][0]["occurrence_id"]
    assert db["events"][entry["review_id"]]["expected_set"] == ["ㄐㄧㄠˇ"]
    assert override.read_bytes() == original_bytes
    assert incoming_override.read_bytes() == incoming_bytes
    assert not actual.load_manual_actual_staging(root)["staged_groups"]
    assert not (root / promotion.PROJECT_TRANSACTION_FILE).exists()
    assert not (destination / portable.INCOMPLETE_FILE).exists()
    with patch.object(actual, "_write_csv", wraps=actual._write_csv) as writes:
        duplicate = portable.import_project_decisions(incoming, destination)
    assert duplicate["imported"] == 0 and duplicate["duplicates"] == 1
    assert duplicate["matched_actual_overrides"] == 1
    assert not [call for call in writes.call_args_list if Path(call.args[0]) == override]
    assert override.read_bytes() == original_bytes
    assert portable._load_project(destination) == (manifest, db)
