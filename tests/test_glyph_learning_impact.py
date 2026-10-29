"""Finite registered usage and real isolated PDF/workbook dependency checks."""
import copy
import hashlib
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

import actual_review as ar
import glyph_learning_impact as impact
import global_exact_glyph_library as library
import standalone_proofread as sp
from review_save_service import ReviewSaveService
from tests.actual_source_test_support import register_fixture_glyph, write_fixture_pdf, bind_source_from_pdf
from tests.test_global_exact_glyph_read_reuse_v580 import _identity, _record, _snapshot, _absent_snapshot
from tests.test_excel_text_storage import isolated_pipeline


def fixture_project(root, *, source_snapshot=None, glyphs=("impact-a", "impact-b")):
    root.mkdir()
    infos, identities = [], []
    source_snapshot = source_snapshot or _absent_snapshot()
    for index, seed in enumerate(glyphs):
        sha = register_fixture_glyph(seed)
        identity = _identity(sha)
        identities.append(identity)
        pdf = root.parent / f"{root.name}-{index}.pdf"
        write_fixture_pdf(pdf, [{"glyph_sha256": sha}])
        entry = bind_source_from_pdf({"x0": 10, "y0": 20}, pdf)
        path = root / "01_實際注音" / f"actual-{index}.xlsx"
        path.parent.mkdir(exist_ok=True)
        book = Workbook()
        sheet = book.active
        sheet.title = "實際注音"
        sheet.append(["字形架構", "font_xref", "注音元件ID", "TTF字形SHA256", "CFF樣式群組", "CFF整字字形SHA256"])
        sheet.append(["TrueType複合元件", entry["font_xref"], entry["zhuyin_component_id"], sha, "", ""])
        metadata = book.create_sheet("v5.2中繼資料")
        metadata.append(["pdf_sha256", entry["pdf_sha256"]])
        metadata.append(["actual_asset_fingerprint_components", json.dumps({
            "global_exact_glyph_evidence_hashes": library.global_exact_glyph_evidence_hashes(source_snapshot, (identity,))})])
        book.save(path)
        book.close()
        infos.append({"pdf": str(pdf), "pdf_name": pdf.name, "pdf_sha256": entry["pdf_sha256"],
                      "actual_workbook": str(path.relative_to(root)),
                      "actual_workbook_sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
    manifest = sp.seal_manifest({"session_id": "session-" + root.name, "pdfs": infos, "records": []})
    (root / "校對工作階段.json").write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    return manifest, identities


def repository(snapshot):
    return SimpleNamespace(load_snapshot=lambda: snapshot)


def test_conflict_invalidates_only_dependent_pdf_and_recheck_is_read_only(tmp_path):
    identity = _identity(register_fixture_glyph("impact-a"))
    verified = _snapshot(trusted=(_record(identity),))
    root = tmp_path / "project"
    manifest, identities = fixture_project(root, source_snapshot=verified)
    before = {str(path): path.read_bytes() for path in tmp_path.rglob("*") if path.is_file()}
    assert impact.recheck_project_dependencies(root, repository=repository(verified))["status"] == "CURRENT"
    conflict = _snapshot(quarantined=(_record(identity, conflict=("ㄅ", "ㄆ")),))
    result = impact.recheck_project_dependencies(root, repository=repository(conflict))
    assert result["status"] == "STALE"
    assert [row["status"] for row in result["pdfs"]] == ["STALE", "CURRENT"]
    assert result["pdfs"][0]["dependencies"][0]["glyph_sha256"] == identities[0].glyph_sha256
    assert before == {str(path): path.read_bytes() for path in tmp_path.rglob("*") if path.is_file()}


def test_registered_impacts_are_finite_and_report_source_changes(tmp_path, monkeypatch):
    root, unregistered = tmp_path / "registered", tmp_path / "not-registered"
    manifest, identities = fixture_project(root)
    fixture_project(unregistered)
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "appdata"))
    assert impact.recheck_project_dependencies(root)["status"] == "CURRENT"
    assert not (tmp_path / "appdata").exists()
    registry = tmp_path / "appdata" / "usage.json"
    impact.register_usage(root, repository=repository(_absent_snapshot()), registry_path=registry)
    glyph_id = library.compute_global_glyph_id(identities[0].kind, identities[0].style_group,
        identities[0].glyph_sha256, identity_eligibility=identities[0].identity_eligibility)
    monkeypatch.setattr(Path, "rglob", lambda *_: (_ for _ in ()).throw(AssertionError("no scan")))
    found = impact.known_impacts([glyph_id], registry_path=registry)
    assert found["scope"] == "REGISTERED_PROJECTS_ONLY"
    assert [item["output_dir"] for item in found["projects"]] == [str(root)]
    assert found["projects"][0]["status"] == "REGISTERED"
    source = Path(manifest["pdfs"][0]["pdf"])
    source.write_bytes(source.read_bytes() + b"changed source")
    assert impact.known_impacts(glyph_id, registry_path=registry)["projects"][0]["status"] == "SOURCE_CHANGED"
    source.unlink()
    assert impact.known_impacts(glyph_id, registry_path=registry)["projects"][0]["status"] == "UNAVAILABLE"


def test_unknown_dependency_contract_and_registry_corruption_fail_closed(tmp_path):
    root = tmp_path / "project"
    manifest, _ = fixture_project(root)
    path = root / manifest["pdfs"][0]["actual_workbook"]
    book = load_workbook(path)
    book["v5.2中繼資料"]["B2"] = "{}"
    book.save(path)
    book.close()
    manifest["pdfs"][0]["actual_workbook_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
    manifest = sp.seal_manifest(manifest)
    (root / "校對工作階段.json").write_text(json.dumps(manifest), encoding="utf-8")
    assert impact.recheck_project_dependencies(root, repository=repository(_absent_snapshot()))["status"] == "UNPROVABLE"
    registry = tmp_path / "usage.json"
    with pytest.raises(ValueError, match="無法核實"):
        impact.register_usage(root, repository=repository(_absent_snapshot()), registry_path=registry)
    assert not registry.exists()
    registry.write_text('{"contract":"future","projects":[]}', encoding="utf-8")
    with pytest.raises(ValueError, match="contract"):
        impact.known_impacts([], registry_path=registry)


def test_recheck_rejects_source_and_manifest_drift_without_writes(tmp_path):
    root = tmp_path / "project"
    manifest, _ = fixture_project(root)
    old = copy.deepcopy(manifest)
    manifest["session_id"] += "-changed"
    (root / "校對工作階段.json").write_text(json.dumps(sp.seal_manifest(manifest)), encoding="utf-8")
    with pytest.raises(ValueError, match="工作階段已變更"):
        impact.recheck_project_dependencies(root, manifest=old, repository=repository(_absent_snapshot()))
    source = Path(manifest["pdfs"][0]["pdf"])
    source.write_bytes(b"invalid source")
    with pytest.raises(ValueError, match="SOURCE_INVALID|SHA"):
        impact.recheck_project_dependencies(root, repository=repository(_absent_snapshot()))


def test_unprovable_reopen_does_not_refresh(tmp_path, monkeypatch):
    monkeypatch.setattr(impact, "recheck_project_dependencies", lambda *a, **k: {"status": "UNPROVABLE"})
    monkeypatch.setattr(sp, "refresh_actual_project", lambda *a, **k: pytest.fail("must not write"))
    with pytest.raises(ValueError, match="無法核實"):
        impact.refresh_stale_project_dependencies(tmp_path)


def test_global_reopen_uses_decoder_refresh_and_preserves_manual_expected(isolated_pipeline, monkeypatch):
    # The existing fixture substitutes only decoded rows. This exercises the
    # real source admission, staging, publication, and manual expected storage.
    import tests.test_excel_text_storage as fixture_module
    monkeypatch.setattr(fixture_module, "TEXTS", ("角",))
    root = isolated_pipeline("actual")
    sp.enable_surface_actual_project(root, confirmation=True)
    service = ReviewSaveService(root)
    snapshot = service.load_resume_snapshot()
    row = snapshot.ledger[0]
    result = service.save_event(row["review_id"], sp.build_manual_expected_event(row,
        operation="ENTER_EXPECTED", expected_set="ㄐㄧㄠˇ"),
        expected_manifest=snapshot.manifest, expected_db=snapshot.db)
    saved_expected = copy.deepcopy(result.db["events"][row["review_id"]])
    checks = iter(({"status": "STALE"}, {"status": "CURRENT"}))
    monkeypatch.setattr(impact, "recheck_project_dependencies", lambda *a, **k: next(checks))
    monkeypatch.setattr(sp, "_surface_projection_refresh", lambda *a, **k: pytest.fail("old seals must not be projected"))
    assert impact.refresh_stale_project_dependencies(root)["status"] == "CURRENT"
    refreshed = ReviewSaveService(root).load_resume_snapshot()
    assert refreshed.db["events"][row["review_id"]] == saved_expected
    assert refreshed.ledger[0]["expected_set"] == ["ㄐㄧㄠˇ"]
