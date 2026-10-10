"""Isolated end-to-end contracts for same-content, different-byte PDF transfer."""

import hashlib
import json
import threading
from pathlib import Path
from unittest.mock import patch

import fitz
import pytest
import actual_review as ar
from openpyxl import Workbook, load_workbook

import pdf_portability as portable
import standalone_proofread as sp
from review_display import occurrence_preview
from review_save_service import ReviewSaveService
from occurrence_ledger import (
    LEDGER_SCHEMA_VERSION, REVIEW_ID_SCHEMA_VERSION, SESSION_SCHEMA_VERSION,
    WORKBOOK_SCHEMA_VERSION, build_occurrence_ledger, prepare_occurrence_rows,
)


@pytest.fixture(autouse=True)
def synthetic_project_report_renderer(monkeypatch):
    # These small sealed-project fixtures use intentionally synthetic candidate
    # workbooks. The separate real-pipeline integration test exercises actual
    # pending/status/Excel publication after transfer.
    monkeypatch.setattr(portable, "_publish_portable_outputs", lambda *_args: None)


def pdf(path: Path, *, title: str, text: str = "角角ㄐㄩㄝˊ", picture=(0, 0, 1), tone="rising"):
    document = fitz.open()
    page = document.new_page(width=300, height=200)
    page.insert_text((20, 40), text, fontsize=16)
    page.draw_rect(fitz.Rect(210, 100, 245, 135), color=picture, fill=picture)
    # A vector tone mark stands in for the embedded textbook glyph program;
    # the default PDF font cannot encode Bopomofo on every CI host.
    page.draw_line((70, 100), (85, 88 if tone == "rising" else 112), color=(0, 0, 0), width=2)
    document.set_metadata({"title": title})
    document.save(path)
    document.close()


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def project(root: Path, source_pdf: Path, *, session: str, event_positions=(), boxes=None):
    root.mkdir(exist_ok=True)
    actual = root / "01_實際注音"
    candidate = root / "02_候選報告"
    actual.mkdir()
    candidate.mkdir()
    actual_file = actual / "actual.xlsx"
    candidate_file = candidate / "candidate.xlsx"
    evidence_root = sp.project_actual_evidence_root(root)
    ar.ensure_user_evidence_files(evidence_root)
    dynamic = ar.dynamic_actual_hashes(evidence_root, pdf_path=source_pdf, dependencies=None, read_only=True)
    workbook = Workbook()
    workbook.active.title = "實際注音"
    workbook.active.append(["occurrence_id"])
    metadata = workbook.create_sheet("v5.2中繼資料")
    metadata.append(["項目", "內容"])
    metadata.append(["pdf_sha256", sha(source_pdf)])
    metadata.append(["actual_asset_fingerprint_components", json.dumps({"dynamic_actual_evidence_hashes": dynamic})])
    workbook.save(actual_file)
    workbook.close()
    candidate_file.write_bytes(b"sealed candidate test artifact")
    digest = sha(source_pdf)
    rows = []
    boxes = boxes or ((20, 20, 40, 45), (90, 20, 110, 45))
    for index, (x0, y0, x1, y1) in enumerate(boxes, 1):
        rows.append({
            "ledger_schema_version": LEDGER_SCHEMA_VERSION,
            "workbook_schema_version": WORKBOOK_SCHEMA_VERSION,
            "pdf_sha256": digest, "pdf": str(source_pdf), "pdf_name": source_pdf.name,
            "實體頁碼": 1, "課本頁": 1, "字元": "角", "實際注音": "ㄐㄩㄝˊ",
            "解碼依據": "glyph evidence", "穩定注音鍵": f"F#{index}", "font": "F", "font_xref": 1,
            "TTF字形SHA256": "a" * 64,
            "glyph_id_字形索引": index, "注音元件ID": index,
            "x0": x0, "y0": y0, "x1": x1, "y1": y1, "source_row_number": index,
        })
    prepare_occurrence_rows(digest, rows)
    ledger = build_occurrence_ledger(rows, {
        row["occurrence_id"]: {
            "state": "EXPECTED_UNRESOLVED", "actual": "ㄐㄩㄝˊ", "actual_evidence": "glyph evidence",
            "context_evidence": f"角色/角@{index}", "source_view": "未建立預期音", "source_record": row,
        }
        for index, row in enumerate(rows, 1)
    })
    manifest = sp.seal_manifest({
        "version": sp.VERSION,
        "session_schema_version": SESSION_SCHEMA_VERSION,
        "workbook_schema_version": WORKBOOK_SCHEMA_VERSION,
        "ledger_schema_version": LEDGER_SCHEMA_VERSION,
        "review_id_schema_version": REVIEW_ID_SCHEMA_VERSION,
        "session_id": session,
        "pdfs": [{"pdf": str(source_pdf), "pdf_name": source_pdf.name,
                  "pdf_sha256": digest, "actual_workbook": str(actual_file),
                  "actual_workbook_sha256": sha(actual_file), "candidate_workbook": str(candidate_file),
                  "candidate_workbook_sha256": sha(candidate_file)}],
        "records": ledger,
        "actual_source_ids": [row["occurrence_id"] for row in rows],
        "regression_gate": {"ok": False, "required": 1, "executed": 1,
                            "passed": 0, "failed": 1, "not_executed": 0, "duplicate_case_id": 0},
    })
    sp.json_save(root / "校對工作階段.json", manifest)
    events = {}
    for index in event_positions:
        events[rows[index]["review_id"]] = {
            "action": "補建expected證據", "expected_set": ["ㄐㄩㄝˊ"],
            "expected_evidence": "現版手冊", "context_evidence": f"角色/角@{index + 1}",
        }
    db = sp.normalize_db({})
    db["events"] = events
    sp.json_save(root / "人工判定資料庫.json", db)
    return manifest, db


def fixture_export_pending_for_gpt(source: Path) -> Path:
    """Fixture-only writer-entry shim; all original proof producer checks run.

    Existing root-first decorator incorrectly receives Workbook at this one
    export entry. This shim ends before any tested import/transaction call.
    Production export validation is not claimed by these import regressions.
    """
    from sealed_workbook_paths import trusted_root, binding_scope
    original = portable.write_excel_content_proof
    def fixture_writer(workbook, output_dir, *args, **kwargs):
        trusted_root(output_dir)
        with binding_scope():
            return original.__wrapped__(workbook, output_dir, *args, **kwargs)
    with patch.object(portable, "write_excel_content_proof", fixture_writer):
        return sp.export_pending_for_gpt(source)


def _filled_expected_excel(source: Path, target: Path, *, action: str = "解決expected證據") -> Path:
    exported = fixture_export_pending_for_gpt(source)
    workbook = load_workbook(exported)
    sheet = workbook["待判定候選"]
    headers = [cell.value for cell in sheet[1]]
    for key, value in {"action": action, "proposed_expected_set": "ㄐㄩㄝˊ",
                       "proposed_expected_evidence": "A independent rule",
                       "proposed_context_evidence": "角色/角@1"}.items():
        sheet.cell(2, headers.index(key) + 1, value)
    workbook.save(target)
    workbook.close()
    return target


@pytest.mark.parametrize("action,reading,actual,state", [
    ("解決expected證據", "ㄐㄩㄝˇ", "ㄐㄩㄝˊ", "DIFFERENCE_PENDING_CONFIRMATION"),
    ("解決expected證據", "ㄐㄩㄝˊ", "ㄐㄩㄝˇ", "DIFFERENCE_PENDING_CONFIRMATION"),
    ("解決expected證據", "ㄐㄩㄝˊ", "ㄐㄩㄝˊ", "PASS"),
    ("補建expected證據", "ㄐㄩㄝˊ", "ㄐㄩㄝˇ", "DIFFERENCE_PENDING_CONFIRMATION"),
])
def test_cross_session_expected_excel_preserves_formal_target_expected(
        tmp_path, monkeypatch, action, reading, actual, state):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    a_pdf, b_pdf = tmp_path / "A.pdf", tmp_path / "B.pdf"
    pdf(a_pdf, title="A")
    pdf(b_pdf, title="B")
    a, b = tmp_path / "A", tmp_path / "B"
    project(a, a_pdf, session="A")
    manifest, _ = project(b, b_pdf, session="B")
    entry = manifest["records"][0]
    entry.update(state=state, actual=actual, expected_status="RESOLVED",
                 expected_set=[reading], expected_evidence="B independent rule",
                 context_evidence="角色/角@1")
    sp.seal_manifest(manifest)
    sp.json_save(b / "校對工作階段.json", manifest)
    filled = _filled_expected_excel(a, tmp_path / "filled.xlsx", action=action)
    watched = [b / name for name in ("人工判定資料庫.json", portable.INCOMPLETE_FILE,
                                    portable.CONFLICT_FILE, "待人工確認.json",
                                    "注音校對_最終報告.xlsx", "pipeline_status.json")]
    before = {path: path.read_bytes() if path.exists() else None for path in watched}
    before_row = sp.materialize_ledger(manifest, sp.json_load_strict(b / "人工判定資料庫.json"))[0]
    result = sp.import_gpt_decisions(b, filled)
    same = reading == "ㄐㄩㄝˊ"
    assert result.status["consistent_supplements"] == int(same)
    assert result.status["new_conflicts"] == int(not same)
    assert sp.json_load_strict(b / "人工判定資料庫.json") == sp.normalize_db({})
    assert not (b / portable.INCOMPLETE_FILE).exists()
    after_row = sp.materialize_ledger(manifest, sp.json_load_strict(b / "人工判定資料庫.json"))[0]
    assert (after_row["state"], after_row["expected_set"], after_row["expected_evidence"]) == (
        before_row["state"], before_row["expected_set"], before_row["expected_evidence"])


def test_expected_excel_rechecks_incomplete_marker_after_preflight(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    a_pdf, b_pdf = tmp_path / "A.pdf", tmp_path / "B.pdf"
    pdf(a_pdf, title="A")
    pdf(b_pdf, title="B")
    a, b = tmp_path / "A", tmp_path / "B"
    project(a, a_pdf, session="A")
    project(b, b_pdf, session="B")
    filled = _filled_expected_excel(a, tmp_path / "filled.xlsx", action="補建expected證據")
    marker = b / portable.INCOMPLETE_FILE
    watched = [b / name for name in ("人工判定資料庫.json", portable.CONFLICT_FILE,
                                    "待人工確認.json", "注音校對_最終報告.xlsx",
                                    "pipeline_status.json")]
    before = {path: path.read_bytes() if path.exists() else None for path in watched}
    real = portable.load_excel_content_proof
    injected = {}
    def after_preflight(*args, **kwargs):
        result = real(*args, **kwargs)
        sp.json_save(marker, {"status": "ACTUAL_EXCEL_REFRESH_PENDING", "source": "concurrent actual"})
        injected["marker"] = marker.read_bytes()
        return result
    with patch.object(portable, "load_excel_content_proof", side_effect=after_preflight):
        with pytest.raises(ValueError, match="未完成|恢復|標記"):
            sp.import_gpt_decisions(b, filled)
    assert marker.read_bytes() == injected["marker"]
    assert {path: path.read_bytes() if path.exists() else None for path in watched} == before


@pytest.mark.parametrize("publish_fails", [False, True])
def test_expected_excel_does_not_clear_another_publication_marker(tmp_path, monkeypatch, publish_fails):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    a_pdf, b_pdf = tmp_path / "A.pdf", tmp_path / "B.pdf"
    pdf(a_pdf, title="A")
    pdf(b_pdf, title="B")
    a, b = tmp_path / "A", tmp_path / "B"
    project(a, a_pdf, session="A")
    project(b, b_pdf, session="B")
    filled = _filled_expected_excel(a, tmp_path / "filled.xlsx", action="補建expected證據")
    marker = b / portable.INCOMPLETE_FILE
    foreign = {"status": "ACTUAL_EXCEL_REFRESH_PENDING", "source": "concurrent actual"}
    def replace_during_publish(*_args):
        assert sp.json_load_strict(marker)["status"] == "PRESENTATION_PENDING"
        sp.json_save(marker, foreign)
        if publish_fails:
            raise OSError("report disk full after foreign marker")
    with patch.object(portable, "_publish_portable_outputs", side_effect=replace_during_publish):
        with pytest.raises((ValueError, OSError), match="標記|其他|未完成|foreign marker"):
            sp.import_gpt_decisions(b, filled)
    assert sp.json_load_strict(marker) == foreign
    assert sp.json_load_strict(b / "人工判定資料庫.json")["events"]


def test_expected_excel_precommit_rollback_preserves_foreign_marker_and_receipt(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    a_pdf, b_pdf = tmp_path / "A.pdf", tmp_path / "B.pdf"
    pdf(a_pdf, title="A")
    pdf(b_pdf, title="B")
    a, b = tmp_path / "A", tmp_path / "B"
    project(a, a_pdf, session="A")
    project(b, b_pdf, session="B")
    first = _filled_expected_excel(a, tmp_path / "first.xlsx")
    assert sp.import_gpt_decisions(b, first)[0] == 1
    workbook = load_workbook(first)
    sheet = workbook["待判定候選"]
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(2, headers.index("proposed_expected_evidence") + 1, "different independent rule")
    second = tmp_path / "second.xlsx"
    workbook.save(second)
    workbook.close()
    db_path = b / "人工判定資料庫.json"
    receipt = b / portable.CONFLICT_FILE
    marker = b / portable.INCOMPLETE_FILE
    before_db = db_path.read_bytes()
    foreign_marker = {"status": "ACTUAL_EXCEL_REFRESH_PENDING", "source": "other transaction"}
    foreign_receipt = {"version": 1, "conflicts": [], "source": "other transaction"}
    original_save = sp.json_save
    injected = {}
    def fail_after_handoff(path, value, **kwargs):
        if Path(path) == db_path:
            original_save(marker, foreign_marker)
            original_save(receipt, foreign_receipt)
            injected["marker"] = marker.read_bytes()
            injected["receipt"] = receipt.read_bytes()
            raise OSError("disk full after marker handoff")
        return original_save(path, value, **kwargs)
    with patch.object(sp, "json_save", side_effect=fail_after_handoff):
        with pytest.raises(ValueError, match="標記|其他交易"):
            sp.import_gpt_decisions(b, second)
    assert db_path.read_bytes() == before_db
    assert marker.read_bytes() == injected["marker"]
    assert receipt.read_bytes() == injected["receipt"]


@pytest.mark.parametrize("difference", ["same", "reading", "evidence", "expected",
                                         "expected_missing", "expected_evidence", "context",
                                         "source_snapshot", "missing_snapshot"])
def test_project_six_gate_transfer_requires_target_actual_snapshot(tmp_path, monkeypatch, difference):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = tmp_path / "source.pdf", tmp_path / "target.pdf"
    pdf(first, title="A")
    pdf(second, title="B")
    source, target = tmp_path / "source", tmp_path / "target"
    source_manifest, source_db = project(source, first, session="A")
    target_manifest, target_db = project(target, second, session="B")
    expected = "ㄐㄧㄠˇ"
    for manifest, reading in ((source_manifest, "ㄐㄩㄝˊ"),
                              (target_manifest, "ㄐㄩㄝˇ" if difference == "reading" else "ㄐㄩㄝˊ")):
        entry = manifest["records"][0]
        entry.update(state="DIFFERENCE_PENDING_CONFIRMATION", actual=reading,
                     actual_evidence="independent glyph " + reading,
                     expected_set=[expected], expected_evidence="independent expected rule",
                     context_evidence="same printed context")
        if manifest is target_manifest:
            if difference == "evidence":
                entry["actual_evidence"] = "other independent glyph evidence"
            elif difference == "expected":
                entry["expected_set"] = ["ㄐㄧㄠˊ"]
            elif difference == "expected_missing":
                entry["expected_set"] = []
                entry["state"] = "EXPECTED_UNRESOLVED"
            elif difference == "expected_evidence":
                entry["expected_evidence"] = "different expected rule"
            elif difference == "context":
                entry["context_evidence"] = "different printed context"
        sp.seal_manifest(manifest)
    sp.json_save(source / "校對工作階段.json", source_manifest)
    sp.json_save(target / "校對工作階段.json", target_manifest)
    source_row = source_manifest["records"][0]
    source_db["events"][source_row["review_id"]] = {
        "action": "確認現版差異", "expected_set": [expected],
        "expected_evidence": "independent expected rule",
        "context_evidence": "same printed context",
        "confirmation_actual_snapshot": sp.actual_confirmation_snapshot(source_row),
        **{gate: True for gate in sp.CONFIRMATION_GATES},
    }
    if difference == "source_snapshot":
        source_db["events"][source_row["review_id"]]["confirmation_actual_snapshot"]["actual_evidence"] = "stale"
    elif difference == "missing_snapshot":
        source_db["events"][source_row["review_id"]].pop("confirmation_actual_snapshot")
    sp.json_save(source / "人工判定資料庫.json", source_db)
    db_path = target / "人工判定資料庫.json"
    before = db_path.read_bytes()
    assert sp.materialize_ledger(source_manifest, source_db)[0]["state"] == "TEXTBOOK_ERROR_CONFIRMED"
    if difference != "same":
        with pytest.raises(ValueError, match="actual|確認|證據"):
            portable.import_project_decisions(source, target)
        assert db_path.read_bytes() == before
        assert not (target / portable.INCOMPLETE_FILE).exists()
        assert sp.materialize_ledger(target_manifest, target_db)[0]["state"] != "TEXTBOOK_ERROR_CONFIRMED"
    else:
        assert portable.import_project_decisions(source, target)["imported"] == 1
        imported = sp.json_load_strict(db_path)
        assert sp.materialize_ledger(target_manifest, imported)[0]["state"] == "TEXTBOOK_ERROR_CONFIRMED"
        assert imported["events"][target_manifest["records"][0]["review_id"]][
            "portability_source"]["pdf_sha256"] == sha(first)


def test_project_import_rechecks_conflict_receipt_after_preflight(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = tmp_path / "source.pdf", tmp_path / "target.pdf"
    pdf(first, title="A")
    pdf(second, title="B")
    source, target = tmp_path / "source", tmp_path / "target"
    source_manifest, _ = project(source, first, session="A", event_positions=(0,))
    target_manifest, _ = project(target, second, session="B", event_positions=(1,))
    source_row, target_row = source_manifest["records"][0], target_manifest["records"][0]
    receipt_path = target / portable.CONFLICT_FILE
    conflict = {
        "source_review_id": source_row["review_id"],
        "target_review_id": target_row["review_id"],
        "source_pdf_sha256": source_row["pdf_sha256"],
        "target_pdf_sha256": target_row["pdf_sha256"],
        "source_event": {"action": "補建expected證據"},
        "target_event": {"action": "確認非校對範圍"},
    }
    db_path = target / "人工判定資料庫.json"
    before_db = db_path.read_bytes()
    pending_path = target / "待人工確認.json"
    status_path = target / "pipeline_status.json"
    report_path = target / "注音校對_最終報告.xlsx"
    sp.json_save(pending_path, {"pending": [target_row["review_id"]]})
    sp.json_save(status_path, {"status": "PENDING_LOCAL_ADJUDICATION"})
    report_path.write_bytes(b"previous report")
    prior_views = {path: path.read_bytes() for path in (pending_path, status_path, report_path)}
    original_mapper = portable._mapped_occurrence_overrides
    def completed_concurrent_conflict(*args):
        sp.json_save(receipt_path, {"version": 1, "conflicts": [conflict]})
        return original_mapper(*args)
    with patch.object(portable, "_mapped_occurrence_overrides", side_effect=completed_concurrent_conflict):
        with pytest.raises(ValueError, match="衝突|未裁決"):
            portable.import_project_decisions(source, target)
    assert db_path.read_bytes() == before_db
    assert {path: path.read_bytes() for path in prior_views} == prior_views
    assert target_manifest["records"][1]["review_id"] in sp.json_load_strict(db_path)["events"]
    assert sp.json_load_strict(receipt_path)["conflicts"] == [conflict]
    assert portable.validate_conflict_state(target, target_manifest, sp.json_load_strict(db_path)) == [
        target_row["review_id"]]
    assert not (target / portable.INCOMPLETE_FILE).exists()


def reseal_actual_dynamic(root: Path, source_pdf: Path):
    """Test producer: bind a changed project-local CSV to its sealed workbook."""
    evidence_root = sp.project_actual_evidence_root(root)
    dynamic = ar.dynamic_actual_hashes(evidence_root, pdf_path=source_pdf, dependencies=None, read_only=True)
    manifest = sp.json_load_strict(root / "校對工作階段.json")
    actual_file = root / "01_實際注音" / "actual.xlsx"
    workbook = load_workbook(actual_file)
    metadata = workbook["v5.2中繼資料"]
    for row in metadata.iter_rows(min_row=2, max_col=2):
        if row[0].value == "actual_asset_fingerprint_components":
            row[1].value = json.dumps({"dynamic_actual_evidence_hashes": dynamic})
    workbook.save(actual_file)
    workbook.close()
    manifest["pdfs"][0]["actual_workbook_sha256"] = sha(actual_file)
    sp.seal_manifest(manifest)
    sp.json_save(root / "校對工作階段.json", manifest)


def add_override(root: Path, entry, reading="ㄐㄩㄝˊ"):
    key = ar._override_key_from_entry(entry)
    fields = ("pdf_contains", "pdf_excludes", "page", "target_char", "stable_key", "x0", "y0")
    row = dict(zip(fields, key)) | {
        "actual_reading": reading, "source": "original direct visual",
        "note": "source page checked",
    }
    ar._write_csv(sp.project_actual_evidence_root(root) / ar.OCCURRENCE_OVERRIDE_FILE,
                  ar.OVERRIDE_HEADERS, [row])
    return row


def pending_actual_excel_conflict(target: Path, target_manifest, source_manifest):
    """Model a sealed, internally valid unresolved actual-import receipt."""
    target_entry, source_entry = target_manifest["records"][0], source_manifest["records"][0]
    previous = ar._read_csv(sp.project_actual_evidence_root(target) / ar.OCCURRENCE_OVERRIDE_FILE,
                            ar.OVERRIDE_HEADERS)[0]
    group_id, snapshot = "verified-source-group", "verified-source-snapshot"
    source_reading = "ㄐㄩㄝˇ"
    provenance = {
        "source_session_id": source_manifest["session_id"],
        "source_manifest_integrity_sha256": source_manifest["manifest_integrity_sha256"],
        "source_pdf_sha256": [source_entry["pdf_sha256"]],
        "source_group_id": group_id, "source_group_snapshot": snapshot,
        "source_occurrence_ids": [source_entry["occurrence_id"]],
        "source_review_ids": [source_entry["review_id"]],
        "source_roster": [{"occurrence_id": source_entry["occurrence_id"],
                           "review_id": source_entry["review_id"],
                           "pdf_sha256": source_entry["pdf_sha256"]}],
        "target_session_id": target_manifest["session_id"],
        "target_occurrence_ids": [target_entry["occurrence_id"]],
        "source_excel_sha256": "b" * 64, "source_excel_row": 2,
        "source_decision": {"decision": "VERIFIED", "actual_reading": source_reading,
                            "group_id": group_id, "group_snapshot": snapshot,
                            "sample_a_occurrence_id": source_entry["occurrence_id"]},
    }
    receipt = portable._sealed_actual_excel_conflicts([{
        "target_occurrence_id": target_entry["occurrence_id"],
        "target_review_id": target_entry["review_id"],
        "target_pdf_sha256": target_entry["pdf_sha256"],
        "target_record": previous, "source_record": provenance,
        "source_reading": source_reading,
    }])
    path = target / portable.ACTUAL_EXCEL_CONFLICT_FILE
    sp.json_save(path, receipt)
    assert portable.actual_excel_conflict_state(
        target, target_manifest, sp.json_load_strict(target / "人工判定資料庫.json")) == [target_entry["review_id"]]
    return path.read_bytes()


@pytest.mark.parametrize("malformed", [False, True])
def test_expected_excel_rechecks_actual_conflict_after_preflight(tmp_path, monkeypatch, malformed):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    a_pdf, b_pdf = tmp_path / "A.pdf", tmp_path / "B.pdf"
    pdf(a_pdf, title="A")
    pdf(b_pdf, title="B")
    a, b = tmp_path / "A", tmp_path / "B"
    source_manifest, _ = project(a, a_pdf, session="A")
    target_manifest, _ = project(b, b_pdf, session="B")
    add_override(b, target_manifest["records"][0])
    reseal_actual_dynamic(b, b_pdf)
    target_manifest = sp.json_load_strict(b / "校對工作階段.json")
    filled = _filled_expected_excel(a, tmp_path / "filled.xlsx", action="補建expected證據")
    watched = [b / name for name in ("人工判定資料庫.json", "校對工作階段.json",
                                    portable.INCOMPLETE_FILE, portable.CONFLICT_FILE,
                                    "待人工確認.json", "注音校對_最終報告.xlsx", "pipeline_status.json")]
    before = {path: path.read_bytes() if path.exists() else None for path in watched}
    real = portable.load_excel_content_proof
    injected = {}

    def after_preflight(*args, **kwargs):
        result = real(*args, **kwargs)
        injected["receipt"] = pending_actual_excel_conflict(b, target_manifest, source_manifest)
        if malformed:
            receipt_path = b / portable.ACTUAL_EXCEL_CONFLICT_FILE
            receipt = sp.json_load_strict(receipt_path)
            receipt["integrity_sha256"] = "0" * 64
            sp.json_save(receipt_path, receipt)
            injected["receipt"] = receipt_path.read_bytes()
        return result

    with patch.object(portable, "load_excel_content_proof", side_effect=after_preflight):
        with pytest.raises(ValueError, match="actual.*衝突|衝突.*actual|完整性"):
            sp.import_gpt_decisions(b, filled)
    assert (b / portable.ACTUAL_EXCEL_CONFLICT_FILE).read_bytes() == injected["receipt"]
    assert {path: path.read_bytes() if path.exists() else None for path in watched} == before
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


@pytest.mark.parametrize("dry_run", [False, True])
def test_expected_excel_rejects_target_pdf_changed_after_mapping(tmp_path, monkeypatch, dry_run):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    a_pdf, b_pdf, changed_pdf = (tmp_path / name for name in ("A.pdf", "B.pdf", "changed.pdf"))
    pdf(a_pdf, title="A")
    pdf(b_pdf, title="B")
    pdf(changed_pdf, title="changed", text="different textbook content")
    source, target = tmp_path / "source", tmp_path / "target"
    project(source, a_pdf, session="A")
    project(target, b_pdf, session="B")
    filled = _filled_expected_excel(source, tmp_path / "filled.xlsx", action="補建expected證據")
    watched = [target / name for name in (
        "人工判定資料庫.json", portable.INCOMPLETE_FILE, portable.CONFLICT_FILE,
        "待人工確認.json", "pipeline_status.json", "注音校對_最終報告.xlsx",
    )]
    watched.append(sp.project_actual_evidence_root(target) / ar.OCCURRENCE_OVERRIDE_FILE)
    before = {path: path.read_bytes() if path.exists() else None for path in watched}
    original = portable._verify_excel_import_snapshot
    calls = 0

    def replace_after_mapping(*args):
        nonlocal calls
        result = original(*args)
        calls += 1
        if calls == (1 if dry_run else 2):
            b_pdf.write_bytes(changed_pdf.read_bytes())
        return result

    with patch.object(portable, "_verify_excel_import_snapshot", side_effect=replace_after_mapping):
        with pytest.raises(ValueError, match="目標 PDF|PDF SHA|頁面內容"):
            sp.import_gpt_decisions(target, filled, dry_run=dry_run)
    assert calls == (1 if dry_run else 2)
    assert {path: path.read_bytes() if path.exists() else None for path in watched} == before
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


def test_project_import_rejects_target_pdf_changed_after_mapping(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    a_pdf, b_pdf, changed_pdf = (tmp_path / name for name in ("A.pdf", "B.pdf", "changed.pdf"))
    pdf(a_pdf, title="A")
    pdf(b_pdf, title="B")
    pdf(changed_pdf, title="changed", text="different textbook content")
    source, target = tmp_path / "source", tmp_path / "target"
    project(source, a_pdf, session="A", event_positions=(0,))
    project(target, b_pdf, session="B")
    portable.prepare_portable_project(source)
    watched = [target / name for name in (
        "人工判定資料庫.json", portable.INCOMPLETE_FILE, portable.CONFLICT_FILE,
        "待人工確認.json", "pipeline_status.json", "注音校對_最終報告.xlsx",
    )]
    watched.append(sp.project_actual_evidence_root(target) / ar.OCCURRENCE_OVERRIDE_FILE)
    before = {path: path.read_bytes() if path.exists() else None for path in watched}
    original = portable._mapped_occurrence_overrides
    changed = False

    def replace_after_mapping(*args):
        nonlocal changed
        result = original(*args)
        changed = True
        b_pdf.write_bytes(changed_pdf.read_bytes())
        return result

    with patch.object(portable, "_mapped_occurrence_overrides", side_effect=replace_after_mapping):
        with pytest.raises(ValueError, match="目標 PDF|PDF SHA|頁面內容"):
            portable.import_project_decisions(source, target)
    assert changed
    assert {path: path.read_bytes() if path.exists() else None for path in watched} == before
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


def test_incomplete_actual_excel_marker_cannot_be_cleared_without_owner_plan(tmp_path):
    source_pdf = tmp_path / "textbook.pdf"
    pdf(source_pdf, title="textbook")
    target = tmp_path / "target"
    project(target, source_pdf, session="B")
    marker = target / portable.INCOMPLETE_FILE
    sp.json_save(marker, {"status": "ACTUAL_EXCEL_REFRESH_PENDING",
                          "source_excel_sha256": "0" * 64,
                          "pre_state": portable._actual_excel_file_snapshot(target)})
    before = marker.read_bytes()
    db_before = (target / "人工判定資料庫.json").read_bytes()
    with pytest.raises(ValueError, match="標記|恢復材料|不完整"):
        portable.resume_actual_excel_project(target)
    assert marker.read_bytes() == before
    assert (target / "人工判定資料庫.json").read_bytes() == db_before


@pytest.mark.parametrize("mutation", ["missing_owner", "foreign_owner", "future_version", "changed_plan"])
def test_actual_excel_marker_owner_and_plan_must_match_before_recovery(tmp_path, mutation):
    from global_glyph_promotion import post_commit_recovery_plan_from_results

    source_pdf = tmp_path / "textbook.pdf"
    pdf(source_pdf, title="textbook")
    target = tmp_path / "target"
    manifest, _ = project(target, source_pdf, session="B")
    plan = post_commit_recovery_plan_from_results([])
    complete = portable._sealed_actual_excel_marker({
        "status": "ACTUAL_EXCEL_REFRESH_PENDING", "source_excel_sha256": "a" * 64,
        "import_token": "b" * 32, "pre_state": portable._actual_excel_file_snapshot(target),
        "sealed_workbooks": portable._sealed_workbook_snapshot(target, manifest),
        "recovery_plan": plan,
        "recovery_plan_sha256": hashlib.sha256(portable._canonical(plan)).hexdigest(),
    })
    if mutation == "missing_owner":
        complete.pop("import_token")
    elif mutation == "foreign_owner":
        complete["import_token"] = "c" * 32
    elif mutation == "future_version":
        complete["version"] += 1
    else:
        complete["recovery_plan_sha256"] = "0" * 64
    marker = target / portable.INCOMPLETE_FILE
    sp.json_save(marker, complete)
    watched = [marker, target / "人工判定資料庫.json", target / "校對工作階段.json"]
    before = {path: path.read_bytes() for path in watched}
    with pytest.raises(ValueError, match="標記|恢復材料"):
        portable.resume_actual_excel_project(target)
    assert {path: path.read_bytes() for path in watched} == before


@pytest.mark.parametrize("malformed", [False, True])
def test_project_import_rechecks_actual_conflict_after_preflight(tmp_path, monkeypatch, malformed):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    a_pdf, b_pdf = tmp_path / "A.pdf", tmp_path / "B.pdf"
    pdf(a_pdf, title="A")
    pdf(b_pdf, title="B")
    source, target = tmp_path / "source", tmp_path / "target"
    source_manifest, _ = project(source, a_pdf, session="A", event_positions=(0,))
    target_manifest, _ = project(target, b_pdf, session="B")
    add_override(target, target_manifest["records"][0])
    reseal_actual_dynamic(target, b_pdf)
    target_manifest = sp.json_load_strict(target / "校對工作階段.json")
    portable.prepare_portable_project(source)
    watched = [target / name for name in ("人工判定資料庫.json", "校對工作階段.json",
                                         portable.INCOMPLETE_FILE, portable.CONFLICT_FILE,
                                         "待人工確認.json", "注音校對_最終報告.xlsx", "pipeline_status.json")]
    before = {path: path.read_bytes() if path.exists() else None for path in watched}
    real = portable._mapped_occurrence_overrides
    injected = {}

    def after_preflight(*args, **kwargs):
        result = real(*args, **kwargs)
        injected["receipt"] = pending_actual_excel_conflict(target, target_manifest, source_manifest)
        if malformed:
            receipt_path = target / portable.ACTUAL_EXCEL_CONFLICT_FILE
            receipt = sp.json_load_strict(receipt_path)
            receipt["integrity_sha256"] = "0" * 64
            sp.json_save(receipt_path, receipt)
            injected["receipt"] = receipt_path.read_bytes()
        return result

    with patch.object(portable, "_mapped_occurrence_overrides", side_effect=after_preflight):
        with pytest.raises(ValueError, match="actual.*衝突|衝突.*actual|完整性"):
            portable.import_project_decisions(source, target)
    assert (target / portable.ACTUAL_EXCEL_CONFLICT_FILE).read_bytes() == injected["receipt"]
    assert {path: path.read_bytes() if path.exists() else None for path in watched} == before
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


def test_different_sha_same_pages_a_to_b_to_a_with_reopen(tmp_path):
    a_pdf, b_pdf = tmp_path / "A.pdf", tmp_path / "B.pdf"
    pdf(a_pdf, title="download A")
    pdf(b_pdf, title="download B")
    assert sha(a_pdf) != sha(b_pdf)
    a, b = tmp_path / "A_project", tmp_path / "B_project"
    a_manifest, _ = project(a, a_pdf, session="A", event_positions=(0,))
    b_manifest, _ = project(b, b_pdf, session="B")
    portable.prepare_portable_project(a)
    a_pdf.rename(tmp_path / "A-away.pdf")
    result = portable.import_project_decisions(a, b)
    assert result["imported"] == 1 and not result["conflicts"]
    assert result["pdf_sha_map"] == {a_manifest["pdfs"][0]["pdf_sha256"]: b_manifest["pdfs"][0]["pdf_sha256"]}
    b_db = sp.json_load_strict(b / "人工判定資料庫.json")
    target_event = b_db["events"][b_manifest["records"][0]["review_id"]]
    assert target_event["portability_source"]["review_id"] == a_manifest["records"][0]["review_id"]
    assert target_event["portability_source"]["pdf_sha256"] != target_event["portability_source"]["target_pdf_sha256"]
    assert sp.materialize_ledger(b_manifest, b_db)[0]["expected_set"] == ["ㄐㄩㄝˊ"]
    assert portable.import_project_decisions(a, b)["duplicates"] == 1
    b_db["events"][b_manifest["records"][1]["review_id"]] = {
        "action": "補建expected證據", "expected_set": ["ㄐㄩㄝˊ"],
        "expected_evidence": "現版手冊", "context_evidence": "角色/角@2",
    }
    sp.json_save(b / "人工判定資料庫.json", b_db)
    portable.prepare_portable_project(b)
    (tmp_path / "A-away.pdf").rename(a_pdf)
    back = portable.import_project_decisions(b, a)
    assert back["duplicates"] == 1 and back["imported"] == 1
    reopened = sp.json_load_strict(a / "人工判定資料庫.json")
    assert len(reopened["events"]) == 2
    assert len(sp.materialize_ledger(a_manifest, reopened)) == 2


@pytest.mark.parametrize("action,difference", [
    ("解決expected證據", "reading"), ("補建expected證據", "reading"),
    ("解決expected證據", "evidence"), ("補建expected證據", "evidence"),
    ("解決expected證據", "context"), ("補建expected證據", "context"),
])
def test_project_transfer_preserves_formal_target_expected(tmp_path, action, difference):
    a_pdf, b_pdf = tmp_path / "A.pdf", tmp_path / "B.pdf"
    pdf(a_pdf, title="A")
    pdf(b_pdf, title="B")
    assert sha(a_pdf) != sha(b_pdf)
    source, target = tmp_path / "source", tmp_path / "target"
    source_manifest, source_db = project(source, a_pdf, session="A", event_positions=(0,))
    target_manifest, target_db = project(target, b_pdf, session="B")
    source_event = source_db["events"][source_manifest["records"][0]["review_id"]]
    source_event["action"] = action
    sp.json_save(source / "人工判定資料庫.json", source_db)
    entry = target_manifest["records"][0]
    entry.update(state="DIFFERENCE_PENDING_CONFIRMATION" if difference == "reading" else "PASS",
                 expected_status="RESOLVED",
                 expected_set=["ㄐㄩㄝˇ"] if difference == "reading" else ["ㄐㄩㄝˊ"],
                 expected_evidence="B independent rule" if difference == "evidence" else "現版手冊",
                 context_evidence="B independent context" if difference == "context" else "角色/角@1")
    sp.seal_manifest(target_manifest)
    sp.json_save(target / "校對工作階段.json", target_manifest)
    watched = [target / name for name in ("校對工作階段.json", "人工判定資料庫.json",
                                         portable.INCOMPLETE_FILE, portable.CONFLICT_FILE,
                                         "待人工確認.json", "注音校對_最終報告.xlsx", "pipeline_status.json")]
    before = {path: path.read_bytes() if path.exists() else None for path in watched}
    before_entry = sp.materialize_ledger(target_manifest, target_db)[0]
    with pytest.raises(ValueError, match="正式 expected|獨立 expected"):
        portable.import_project_decisions(source, target)
    assert {path: path.read_bytes() if path.exists() else None for path in watched} == before
    after_entry = sp.materialize_ledger(target_manifest, sp.json_load_strict(target / "人工判定資料庫.json"))[0]
    assert (after_entry["state"], after_entry["expected_set"], after_entry["expected_evidence"],
            after_entry["context_evidence"]) == (
                before_entry["state"], before_entry["expected_set"], before_entry["expected_evidence"],
                before_entry["context_evidence"])


@pytest.mark.parametrize("action,difference", [
    ("補建expected證據", "line"), ("解決expected證據", "line"),
    ("補建expected證據", "local_context"), ("解決expected證據", "local_context"),
    ("補建expected證據", "context_evidence"), ("解決expected證據", "context_evidence"),
])
def test_project_expected_transfer_rechecks_independent_target_context(tmp_path, action, difference):
    a_pdf, b_pdf = tmp_path / "A.pdf", tmp_path / "B.pdf"
    pdf(a_pdf, title="A")
    pdf(b_pdf, title="B")
    assert sha(a_pdf) != sha(b_pdf)
    source, target = tmp_path / "source", tmp_path / "target"
    source_manifest, source_db = project(source, a_pdf, session="A", event_positions=(0,))
    source_db["events"][source_manifest["records"][0]["review_id"]]["action"] = action
    sp.json_save(source / "人工判定資料庫.json", source_db)
    target_manifest, target_db = project(target, b_pdf, session="B")
    entry = target_manifest["records"][0]
    if difference == "line":
        entry["source_record"]["所在行"] = "B independent line"
    elif difference == "local_context":
        entry["source_record"]["局部詞境"] = "B independent lexical context"
    else:
        entry["context_evidence"] = "B independent context evidence"
    sp.seal_manifest(target_manifest)
    sp.json_save(target / "校對工作階段.json", target_manifest)
    watched = [target / name for name in ("校對工作階段.json", "人工判定資料庫.json",
                                         portable.INCOMPLETE_FILE, portable.CONFLICT_FILE,
                                         "待人工確認.json", "注音校對_最終報告.xlsx", "pipeline_status.json")]
    before = {path: path.read_bytes() if path.exists() else None for path in watched}
    source_row = sp.materialize_ledger(source_manifest, source_db)[0]
    target_row = sp.materialize_ledger(target_manifest, target_db)[0]
    assert source_row["context_evidence"] != target_row["context_evidence"] or difference != "context_evidence"
    with pytest.raises(ValueError, match="詞境|context"):
        portable.import_project_decisions(source, target)
    assert {path: path.read_bytes() if path.exists() else None for path in watched} == before
    reopened = sp.materialize_ledger(target_manifest, sp.json_load_strict(target / "人工判定資料庫.json"))[0]
    assert reopened["state"] == target_row["state"] == "EXPECTED_UNRESOLVED"
    assert reopened["context_evidence"] == target_row["context_evidence"]


@pytest.mark.parametrize("action", ["補建expected證據", "解決expected證據"])
def test_project_expected_transfer_accepts_more_specific_event_context(tmp_path, action):
    a_pdf, b_pdf = tmp_path / "A.pdf", tmp_path / "B.pdf"
    pdf(a_pdf, title="A")
    pdf(b_pdf, title="B")
    assert sha(a_pdf) != sha(b_pdf)
    source, target = tmp_path / "source", tmp_path / "target"
    source_manifest, source_db = project(source, a_pdf, session="A", event_positions=(0,))
    target_manifest, target_db = project(target, b_pdf, session="B")
    source_row, target_row = source_manifest["records"][0], target_manifest["records"][0]
    assert source_row["context_evidence"] == target_row["context_evidence"]
    assert source_row["source_record"].get("所在行") == target_row["source_record"].get("所在行")
    assert source_row["source_record"].get("局部詞境") == target_row["source_record"].get("局部詞境")
    event = source_db["events"][source_row["review_id"]]
    event["action"] = action
    event["context_evidence"] = "角色／角；同頁句子已由人工核對"
    sp.json_save(source / "人工判定資料庫.json", source_db)
    assert sp.materialize_ledger(source_manifest, source_db)[0]["context_evidence"] == event["context_evidence"]
    result = portable.import_project_decisions(source, target)
    assert result["imported"] == 1 and result["conflicts"] == []
    imported = sp.json_load_strict(target / "人工判定資料庫.json")["events"][target_row["review_id"]]
    assert imported["context_evidence"] == event["context_evidence"]
    assert imported["portability_source"]["review_id"] == source_row["review_id"]
    assert sp.materialize_ledger(target_manifest, sp.json_load_strict(target / "人工判定資料庫.json"))[0]["context_evidence"] == event["context_evidence"]


def test_project_expected_binding_transfers_only_after_proven_mapping_and_unchanged_baseline(tmp_path):
    a_pdf, b_pdf = tmp_path / "A.pdf", tmp_path / "B.pdf"
    pdf(a_pdf, title="A")
    pdf(b_pdf, title="B")
    source, target = tmp_path / "source", tmp_path / "target"
    source_manifest, source_db = project(source, a_pdf, session="A", event_positions=(0,))
    target_manifest, _ = project(target, b_pdf, session="B")
    for directory, manifest in ((source, source_manifest), (target, target_manifest)):
        # Explicit synthetic expected contract and common printed text context.
        manifest["expected_asset_fingerprint"] = "b" * 64
        manifest["expected_asset_fingerprint_components"] = {"expected_resolver_semantics_epoch": "1"}
        for row in manifest["records"]:
            row["source_record"].update({"所在行": "角色", "局部詞境": "角色"})
        sp.json_save(directory / "校對工作階段.json", sp.seal_manifest(manifest))
    source_row, target_row = source_manifest["records"][0], target_manifest["records"][0]
    event = source_db["events"][source_row["review_id"]]
    bound = sp.bind_expected_resolution(source_row, event, source_manifest)
    bound["undo_previous_event"] = copy_binding = sp.bind_expected_resolution(source_row, event, source_manifest)
    source_db["events"][source_row["review_id"]] = bound
    sp.json_save(source / "人工判定資料庫.json", source_db)
    result = portable.import_project_decisions(source, target)
    assert result["imported"] == 1 and result["conflicts"] == []
    db = sp.json_load_strict(target / "人工判定資料庫.json")
    transferred = db["events"][target_row["review_id"]]
    assert transferred["expected_resolution_binding"]["session_id"] == "B"
    assert transferred["undo_previous_event"]["expected_resolution_binding"]["session_id"] == "B"
    assert copy_binding["expected_resolution_binding"]["session_id"] == "A"
    assert transferred["portability_source"]["source_expected_resolution_binding"] == bound["expected_resolution_binding"]
    assert sp.materialize_ledger(target_manifest, db)[0]["state"] == "PASS"
    for update in ({"expected_evidence": "same-reading new independent source"},
                   {"context_evidence": "new independent context"}):
        with pytest.raises(ValueError, match="baseline"):
            portable._transfer_event(bound, source_row, {**target_row, **update},
                source_manifest=source_manifest, target_manifest=target_manifest,
                source_baseline=source_row, target_baseline={**target_row, **update})


@pytest.mark.parametrize("same_decision", [True, False])
def test_project_transfer_keeps_prior_human_expected_duplicate_or_conflict(tmp_path, same_decision):
    a_pdf, b_pdf = tmp_path / "A.pdf", tmp_path / "B.pdf"
    pdf(a_pdf, title="A")
    pdf(b_pdf, title="B")
    source, target = tmp_path / "source", tmp_path / "target"
    source_manifest, source_db = project(source, a_pdf, session="A", event_positions=(0,))
    source_event = source_db["events"][source_manifest["records"][0]["review_id"]]
    source_event["action"] = "解決expected證據"
    sp.json_save(source / "人工判定資料庫.json", source_db)
    target_manifest, target_db = project(target, b_pdf, session="B")
    entry = target_manifest["records"][0]
    entry.update(state="DIFFERENCE_PENDING_CONFIRMATION", expected_status="RESOLVED",
                 expected_set=["ㄐㄩㄝˇ"], expected_evidence="B independent rule",
                 context_evidence="角色/角@1")
    sp.seal_manifest(target_manifest)
    sp.json_save(target / "校對工作階段.json", target_manifest)
    target_id = entry["review_id"]
    target_db["events"][target_id] = {
        "action": "解決expected證據", "expected_set": ["ㄐㄩㄝˊ" if same_decision else "ㄐㄩㄝˋ"],
        "expected_evidence": "現版手冊", "context_evidence": "角色/角@1",
    }
    sp.json_save(target / "人工判定資料庫.json", target_db)
    result = portable.import_project_decisions(source, target)
    db = sp.json_load_strict(target / "人工判定資料庫.json")
    if same_decision:
        assert result["duplicates"] == 1 and not result["conflicts"]
        assert db["events"][target_id]["expected_set"] == ["ㄐㄩㄝˊ"]
        assert len(db["events"][target_id]["portability_duplicate_sources"]) == 1
    else:
        assert len(result["conflicts"]) == 1
        assert target_id not in db["events"]
        receipt = sp.json_load_strict(target / portable.CONFLICT_FILE)
        assert receipt["conflicts"][0]["target_event"]["expected_set"] == ["ㄐㄩㄝˋ"]
        assert receipt["conflicts"][0]["source_event"]["expected_set"] == ["ㄐㄩㄝˊ"]


@pytest.mark.parametrize("change", ["text", "picture", "zhuyin"])
def test_same_filename_different_content_rejected(tmp_path, change):
    left, right = tmp_path / "left", tmp_path / "right"
    left.mkdir(); right.mkdir()
    first, second = left / "book.pdf", right / "book.pdf"
    pdf(first, title="A")
    kwargs = {"title": "B"}
    if change == "text":
        kwargs["text"] = "不一樣ㄐㄩㄝˊ"
    elif change == "picture":
        kwargs["picture"] = (1, 0, 0)
    else:
        kwargs["tone"] = "falling"
    pdf(second, **kwargs)
    source, target = tmp_path / "source", tmp_path / "target"
    project(source, first, session="source", event_positions=(0,))
    project(target, second, session="target")
    before = (target / "人工判定資料庫.json").read_bytes()
    with pytest.raises(ValueError, match="頁面內容"):
        portable.import_project_decisions(source, target)
    assert (target / "人工判定資料庫.json").read_bytes() == before


def test_conflict_preserves_both_and_write_failure_keeps_target(tmp_path):
    a_pdf, b_pdf = tmp_path / "A.pdf", tmp_path / "B.pdf"
    pdf(a_pdf, title="A")
    pdf(b_pdf, title="B")
    source, target = tmp_path / "source", tmp_path / "target"
    project(source, a_pdf, session="source", event_positions=(0,))
    manifest, db = project(target, b_pdf, session="target")
    rid = manifest["records"][0]["review_id"]
    db["events"][rid] = {"action": "保留待人工"}
    sp.json_save(target / "人工判定資料庫.json", db)
    before = (target / "人工判定資料庫.json").read_bytes()
    result = portable.import_project_decisions(source, target)
    assert result["conflicts"] == [{"source_review_id": sp.json_load_strict(source / "校對工作階段.json")["records"][0]["review_id"], "target_review_id": rid}]
    assert (target / "人工判定資料庫.json").read_bytes() != before
    assert rid not in sp.json_load_strict(target / "人工判定資料庫.json")["events"]
    receipt = sp.json_load_strict(target / portable.CONFLICT_FILE)
    assert receipt["conflicts"][0]["target_event"]["action"] == "保留待人工"
    assert receipt["conflicts"][0]["source_event"]["action"] == "補建expected證據"
    with pytest.raises(ValueError, match="未裁決"):
        portable.prepare_portable_project(target)
    resolved = ReviewSaveService(target).save_event(
        rid, {"action": "確認非校對範圍", "exclusion_reason": "new local check",
              "exclusion_evidence": "PDF page 1"})
    assert resolved is not None
    decision = sp.json_load_strict(target / "人工判定資料庫.json")["events"][rid]
    assert decision["portability_conflict_resolution"]["original_conflicts"] == receipt["conflicts"]
    ReviewSaveService(target).save_event(
        rid, {"action": "確認非校對範圍", "exclusion_reason": "revised local check",
              "exclusion_evidence": "PDF page 1 inspected again"})
    revised = sp.json_load_strict(target / "人工判定資料庫.json")
    assert revised["events"][rid]["portability_conflict_resolution"] == decision["portability_conflict_resolution"]
    assert portable.validate_conflict_state(target, manifest, revised) == []
    assert sp.json_load_strict(target / portable.CONFLICT_FILE) == receipt
    another = tmp_path / "another-target"
    project(another, b_pdf, session="another")
    before = (another / "人工判定資料庫.json").read_bytes()
    with patch("standalone_proofread.json_save", side_effect=OSError("disk full")):
        with pytest.raises(OSError, match="disk full"):
            portable.import_project_decisions(source, another)
    assert (another / "人工判定資料庫.json").read_bytes() == before


def test_repeated_character_requires_its_own_position(tmp_path):
    a_pdf, b_pdf = tmp_path / "A.pdf", tmp_path / "B.pdf"
    pdf(a_pdf, title="A")
    pdf(b_pdf, title="B")
    source, target = tmp_path / "source", tmp_path / "target"
    project(source, a_pdf, session="source", event_positions=(1,))
    target_manifest, _ = project(target, b_pdf, session="target")
    portable.import_project_decisions(source, target)
    events = sp.json_load_strict(target / "人工判定資料庫.json")["events"]
    assert list(events) == [target_manifest["records"][1]["review_id"]]


def test_concurrent_save_is_retained_during_explicit_import(tmp_path):
    first, second = tmp_path / "first.pdf", tmp_path / "second.pdf"
    pdf(first, title="A")
    pdf(second, title="B")
    source, target = tmp_path / "source", tmp_path / "target"
    project(source, first, session="source", event_positions=(0,))
    target_manifest, _ = project(target, second, session="target")
    entered, proceed = threading.Event(), threading.Event()
    real_map = portable._map_reviews
    result, errors = [], []

    def paused_map(*args):
        mapped = real_map(*args)
        entered.set()
        assert proceed.wait(10)
        return mapped

    def importer():
        try:
            result.append(portable.import_project_decisions(source, target))
        except Exception as exc:
            errors.append(exc)

    with patch.object(portable, "_map_reviews", side_effect=paused_map):
        thread = threading.Thread(target=importer)
        thread.start()
        try:
            assert entered.wait(10)
            ReviewSaveService(target).save_event(
                target_manifest["records"][1]["review_id"],
                {"action": "確認非校對範圍", "exclusion_reason": "different visual position",
                 "exclusion_evidence": "PDF page 1, position 2"},
            )
        finally:
            proceed.set()
            thread.join(10)
    assert not thread.is_alive() and not errors
    assert result[0]["imported"] == 1
    assert len(sp.json_load_strict(target / "人工判定資料庫.json")["events"]) == 2


def test_different_pdf_object_storage_and_identity_mapping(tmp_path):
    first, second = tmp_path / "uncompressed.pdf", tmp_path / "compressed.pdf"
    document = fitz.open()
    page = document.new_page(width=300, height=200)
    page.insert_text((20, 40), "same complete page", fontsize=16)
    page.draw_rect(fitz.Rect(210, 100, 245, 135), color=(0, 0, 1), fill=(0, 0, 1))
    document.save(first, deflate=False, garbage=0)
    document.save(second, deflate=True, garbage=4)
    document.close()
    assert sha(first) != sha(second)
    assert portable._page_signatures(first) == portable._page_signatures(second)
    source, target = tmp_path / "source", tmp_path / "target"
    source_manifest, _ = project(source, first, session="source", event_positions=(0,))
    target_manifest, _ = project(target, second, session="target")
    # The PDF's font/xref/glyph identity may differ even though the page and
    # occurrence position are equal. Original identities remain distinct.
    assert source_manifest["records"][0]["review_id"] != target_manifest["records"][0]["review_id"]
    result = portable.import_project_decisions(source, target)
    assert result["imported"] == 1


def test_rotation_storage_on_unreviewed_equal_page_does_not_reject(tmp_path):
    first, second = tmp_path / "A.pdf", tmp_path / "B.pdf"
    document = fitz.open()
    page = document.new_page(width=300, height=200)
    page.insert_text((20, 40), "identical lesson", fontsize=16)
    annex = document.new_page(width=200, height=200)
    annex.draw_rect(fitz.Rect(90, 90, 110, 110), fill=(0, 0, 0))
    document.save(first)
    annex.set_rotation(180)
    document.save(second)
    document.close()
    source_pages, target_pages = portable._page_signatures(first), portable._page_signatures(second)
    assert sha(first) != sha(second)
    assert source_pages[1]["rotation"] == 0 and target_pages[1]["rotation"] == 180
    assert source_pages[1]["rgb_sha256"] == target_pages[1]["rgb_sha256"]
    source, target = tmp_path / "source", tmp_path / "target"
    project(source, first, session="A", event_positions=(0,))
    project(target, second, session="B")
    assert portable.import_project_decisions(source, target)["imported"] == 1


def test_rotated_reviewed_page_maps_same_char_by_display_position(tmp_path):
    first, second = tmp_path / "source.pdf", tmp_path / "rotated.pdf"
    source_boxes = ((20, 20, 40, 45), (90, 20, 110, 45))
    target_boxes = ((260, 155, 280, 180), (190, 155, 210, 180))
    for path, boxes, rotation in ((first, source_boxes, 0), (second, target_boxes, 180)):
        document = fitz.open()
        page = document.new_page(width=300, height=200)
        for box in boxes:
            page.draw_rect(fitz.Rect(box), color=(0, 0, 0), fill=(0, 0, 0))
        page.set_rotation(rotation)
        document.save(path)
        document.close()
    assert sha(first) != sha(second)
    signatures_a, signatures_b = portable._page_signatures(first), portable._page_signatures(second)
    assert signatures_a[0]["rgb_sha256"] == signatures_b[0]["rgb_sha256"]
    source, target = tmp_path / "source-project", tmp_path / "target-project"
    source_manifest, _ = project(source, first, session="A", event_positions=(0,), boxes=source_boxes)
    target_manifest, _ = project(target, second, session="B", boxes=target_boxes)
    assert portable.import_project_decisions(source, target)["imported"] == 1
    assert list(sp.json_load_strict(target / "人工判定資料庫.json")["events"]) == [
        target_manifest["records"][0]["review_id"]
    ]
    assert source_manifest["records"][0]["occurrence_id"] != target_manifest["records"][0]["occurrence_id"]


def test_missing_source_pdf_requires_prior_proof_and_moved_artifact_is_checked(tmp_path):
    source_pdf, target_pdf = tmp_path / "source.pdf", tmp_path / "target.pdf"
    pdf(source_pdf, title="A")
    pdf(target_pdf, title="B")
    source, target = tmp_path / "source", tmp_path / "target"
    project(source, source_pdf, session="source", event_positions=(0,))
    project(target, target_pdf, session="target")
    before = (target / "人工判定資料庫.json").read_bytes()
    source_pdf.rename(tmp_path / "absent.pdf")
    with pytest.raises(FileNotFoundError, match="原 PDF"):
        portable.import_project_decisions(source, target)
    assert (target / "人工判定資料庫.json").read_bytes() == before
    (tmp_path / "absent.pdf").rename(source_pdf)
    portable.prepare_portable_project(source)
    source_pdf.rename(tmp_path / "absent.pdf")
    # An intact workbook at the copied project-relative path is accepted even
    # though the sealed manifest retains the first computer's absolute path.
    assert portable.import_project_decisions(source, target)["imported"] == 1
    (source / "01_實際注音" / "actual.xlsx").write_bytes(b"tampered")
    with pytest.raises(ValueError, match="來源工作簿"):
        portable.import_project_decisions(source, target)


def test_pending_actual_transaction_rejects_transfer_without_touching_target(tmp_path):
    source_pdf, target_pdf = tmp_path / "source.pdf", tmp_path / "target.pdf"
    pdf(source_pdf, title="A")
    pdf(target_pdf, title="B")
    source, target = tmp_path / "source", tmp_path / "target"
    project(source, source_pdf, session="source", event_positions=(0,))
    project(target, target_pdf, session="target")
    root = sp.project_actual_evidence_root(source)
    root.mkdir(parents=True, exist_ok=True)
    (root / "global_exact_glyph_project_transaction.json").write_text("pending", encoding="utf-8")
    before = (target / "人工判定資料庫.json").read_bytes()
    with pytest.raises(ValueError, match="交易尚待原專案恢復"):
        portable.import_project_decisions(source, target)
    assert (target / "人工判定資料庫.json").read_bytes() == before


def test_multi_hop_provenance_retains_first_source(tmp_path):
    files = [tmp_path / f"{letter}.pdf" for letter in "ABC"]
    for letter, path in zip("ABC", files):
        pdf(path, title=letter)
    dirs = [tmp_path / letter for letter in "ABC"]
    manifests = [project(folder, path, session=letter, event_positions=(0,) if letter == "A" else ())[0]
                 for letter, folder, path in zip("ABC", dirs, files)]
    portable.import_project_decisions(dirs[0], dirs[1])
    portable.import_project_decisions(dirs[1], dirs[2])
    event = sp.json_load_strict(dirs[2] / "人工判定資料庫.json")["events"][manifests[2]["records"][0]["review_id"]]
    assert event["portability_source"]["session_id"] == "B"
    assert event["portability_source"]["prior"]["session_id"] == "A"
    assert event["portability_source"]["prior"]["review_id"] == manifests[0]["records"][0]["review_id"]


def test_imported_project_can_show_local_page_save_and_reopen(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    source_pdf, local_pdf = tmp_path / "download-a.pdf", tmp_path / "download-b.pdf"
    pdf(source_pdf, title="computer A")
    pdf(local_pdf, title="computer B")
    source, target = tmp_path / "source", tmp_path / "target"
    project(source, source_pdf, session="A", event_positions=(0,))
    target_manifest, _ = project(target, local_pdf, session="B")
    portable.prepare_portable_project(source)
    source_pdf.rename(tmp_path / "not-on-B.pdf")
    portable.import_project_decisions(source, target)
    row = target_manifest["records"][1]
    pixels, rectangle, _ = occurrence_preview(row, output_dir=target)
    assert pixels.width > 0 and pixels.height > 0 and rectangle is not None
    event = {"action": "補建expected證據", "expected_set": ["ㄐㄩㄝˊ"],
             "expected_evidence": "現版手冊", "context_evidence": "角色/角@2"}
    result = ReviewSaveService(target).save_event(row["review_id"], event)
    assert result.resolved_entry["expected_set"] == ["ㄐㄩㄝˊ"]
    reopened = sp.json_load_strict(target / "人工判定資料庫.json")
    assert len(reopened["events"]) == 2
    assert sp.materialize_ledger(target_manifest, reopened)[1]["expected_set"] == ["ㄐㄩㄝˊ"]


def test_source_actual_staging_is_preserved_as_recheck_not_new_global_evidence(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    source_pdf, target_pdf = tmp_path / "source.pdf", tmp_path / "target.pdf"
    pdf(source_pdf, title="A")
    pdf(target_pdf, title="B")
    source, target = tmp_path / "source", tmp_path / "target"
    manifest, _ = project(source, source_pdf, session="source", event_positions=(0,))
    project(target, target_pdf, session="target")
    group = ar.build_actual_group_for_entry(manifest["records"], manifest["records"][0])
    root = sp.project_actual_evidence_root(source)
    ar.stage_manual_actual_group(
        root, group, "ㄐㄩㄝˊ", checked_occurrence_ids=[manifest["records"][0]["occurrence_id"]],
        source="manual visual actual confirmation",
    )
    source_staging = ar.load_manual_actual_staging(root)
    result = portable.import_project_decisions(source, target)
    assert result["actual_staging"] == "RECHECK_LOCAL_PDF_REQUIRED"
    receipt = sp.json_load_strict(target / portable.PENDING_ACTUAL_FILE)
    assert receipt["status"] == "RECHECK_LOCAL_PDF_REQUIRED"
    assert len(receipt["sources"]) == 1
    assert receipt["sources"][0]["staging"] == source_staging
    assert receipt["sources"][0]["source_session_id"] == "source"
    assert ar.load_manual_actual_staging(sp.project_actual_evidence_root(target))["staged_groups"] == []
    from standalone_gui import project_status_details
    assert "重新核對" in project_status_details(target)[0]
    assert not (tmp_path / "isolated-localappdata").exists()


def test_staged_actual_receipt_survives_a_to_b_to_c(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    files = [tmp_path / f"{letter}.pdf" for letter in "ABC"]
    for letter, path in zip("ABC", files):
        pdf(path, title=letter)
    a, b, c = [tmp_path / name for name in ("project-a", "project-b", "project-c")]
    a_manifest, _ = project(a, files[0], session="A", event_positions=(0,))
    project(b, files[1], session="B")
    group = ar.build_actual_group_for_entry(a_manifest["records"], a_manifest["records"][0])
    a_root = sp.project_actual_evidence_root(a)
    ar.stage_manual_actual_group(a_root, group, "ㄐㄩㄝˊ",
                                 checked_occurrence_ids=[a_manifest["records"][0]["occurrence_id"]],
                                 source="manual visual actual confirmation")
    original = ar.load_manual_actual_staging(a_root)
    portable.import_project_decisions(a, b)
    portable.prepare_portable_project(b)

    def build_target(paths, output_dir, **kwargs):
        project(output_dir, files[2], session="C")

    with patch.object(sp, "run_pipeline_pdfs", side_effect=build_target):
        result = portable.continue_project(b, files[2], c)
    assert result["sources"][0]["actual_staging"] == "RECHECK_LOCAL_PDF_REQUIRED"
    receipt = sp.json_load_strict(c / portable.PENDING_ACTUAL_FILE)
    assert len(receipt["sources"]) == 1
    assert receipt["sources"][0]["source_session_id"] == "A"
    assert receipt["sources"][0]["staging"] == original
    assert receipt["sources"][0]["forwarded_via"][0]["source_session_id"] == "B"
    assert ar.load_manual_actual_staging(sp.project_actual_evidence_root(c))["staged_groups"] == []
    assert not (tmp_path / "isolated-localappdata").exists()


def test_tampered_forwarded_actual_receipt_is_rejected_before_target_creation(tmp_path):
    files = [tmp_path / f"{letter}.pdf" for letter in "ABC"]
    for letter, path in zip("ABC", files):
        pdf(path, title=letter)
    a, b, c = [tmp_path / name for name in ("project-a", "project-b", "project-c")]
    a_manifest, _ = project(a, files[0], session="A", event_positions=(0,))
    project(b, files[1], session="B")
    group = ar.build_actual_group_for_entry(a_manifest["records"], a_manifest["records"][0])
    ar.stage_manual_actual_group(
        sp.project_actual_evidence_root(a), group, "ㄐㄩㄝˊ",
        checked_occurrence_ids=[a_manifest["records"][0]["occurrence_id"]],
        source="manual visual actual confirmation",
    )
    portable.import_project_decisions(a, b)
    portable.prepare_portable_project(b)
    receipt_path = b / portable.PENDING_ACTUAL_FILE
    receipt = sp.json_load_strict(receipt_path)
    first_key = next(iter(receipt["sources"][0]["mapped_occurrences"]))
    receipt["sources"][0]["mapped_occurrences"][first_key] = "wrong-target-occurrence"
    sp.json_save(receipt_path, receipt)
    with pytest.raises(ValueError, match="暫存與目前工作階段不符"):
        portable.continue_project(b, files[2], c)
    assert not c.exists()


def test_continue_operation_uses_local_pdf_and_only_clears_success_marker(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, local = tmp_path / "first.pdf", tmp_path / "local.pdf"
    pdf(first, title="A")
    pdf(local, title="B")
    source, target = tmp_path / "source", tmp_path / "continued"
    project(source, first, session="source", event_positions=(0,))
    portable.prepare_portable_project(source)
    first.rename(tmp_path / "first-away.pdf")

    def build_target(paths, output_dir, **kwargs):
        assert paths == [local.resolve()]
        project(output_dir, local, session="continued")

    with patch.object(sp, "run_pipeline_pdfs", side_effect=build_target):
        result = portable.continue_project(source, local, target)
    assert result["imported"] == 1
    assert not (target / portable.INCOMPLETE_FILE).exists()
    assert sp.json_load_strict(target / "校對工作階段.json")["pdfs"][0]["pdf_sha256"] == sha(local)
    assert len(sp.json_load_strict(target / "人工判定資料庫.json")["events"]) == 1


def test_failed_continue_keeps_source_and_marks_target_incomplete(tmp_path):
    first, local = tmp_path / "first.pdf", tmp_path / "local.pdf"
    pdf(first, title="A")
    pdf(local, title="B")
    source, target = tmp_path / "source", tmp_path / "continued"
    project(source, first, session="source", event_positions=(0,))
    portable.prepare_portable_project(source)
    source_db = (source / "人工判定資料庫.json").read_bytes()
    with patch.object(sp, "run_pipeline_pdfs", side_effect=OSError("write failed")):
        with pytest.raises(OSError, match="write failed"):
            portable.continue_project(source, local, target)
    assert (source / "人工判定資料庫.json").read_bytes() == source_db
    assert (target / portable.INCOMPLETE_FILE).is_file()
    from standalone_gui import project_status_details
    assert "未完成" in project_status_details(target)[0]
    with pytest.raises(ValueError, match="接續未完成"):
        ReviewSaveService(target).save_event("any", {"action": "保留待人工"})


def test_sealed_but_incomplete_target_cannot_be_reexported(tmp_path):
    first, local = tmp_path / "source.pdf", tmp_path / "local.pdf"
    pdf(first, title="A")
    pdf(local, title="B")
    source, partial, another = tmp_path / "source", tmp_path / "partial", tmp_path / "another"
    project(source, first, session="A", event_positions=(0,))
    before = (source / "人工判定資料庫.json").read_bytes()

    def build_target(paths, output_dir, **kwargs):
        project(output_dir, local, session="B")

    with patch.object(sp, "run_pipeline_pdfs", side_effect=build_target), \
         patch.object(portable, "_mapped_occurrence_overrides", side_effect=ValueError("mapping stopped")):
        with pytest.raises(ValueError, match="mapping stopped"):
            portable.continue_project(source, local, partial)
    assert (partial / "校對工作階段.json").exists()
    assert (partial / portable.INCOMPLETE_FILE).exists()
    with pytest.raises(ValueError, match="接續未完成"):
        portable.prepare_portable_project(partial)
    with pytest.raises(ValueError, match="接續未完成"):
        portable.continue_project(partial, first, another)
    with pytest.raises(ValueError, match="接續未完成"):
        portable.import_project_decisions(partial, source)
    assert not another.exists()
    assert (source / "人工判定資料庫.json").read_bytes() == before


def test_report_publication_failure_retains_data_and_blocks_partial_target(tmp_path):
    first, local = tmp_path / "first.pdf", tmp_path / "local.pdf"
    pdf(first, title="A")
    pdf(local, title="B")
    source, target = tmp_path / "source", tmp_path / "target"
    project(source, first, session="A", event_positions=(0,))
    project(target, local, session="B")
    before = (source / "人工判定資料庫.json").read_bytes()
    with patch.object(portable, "_publish_portable_outputs", side_effect=OSError("report disk full")):
        with pytest.raises(OSError, match="report disk full"):
            portable.import_project_decisions(source, target)
    assert (source / "人工判定資料庫.json").read_bytes() == before
    assert len(sp.json_load_strict(target / "人工判定資料庫.json")["events"]) == 1
    assert (target / portable.INCOMPLETE_FILE).exists()
    from standalone_gui import project_status_details
    assert "未完成" in project_status_details(target)[0]
    with pytest.raises(ValueError, match="接續未完成"):
        portable.prepare_portable_project(target)
    with pytest.raises(ValueError, match="接續未完成"):
        ReviewSaveService(target).save_event("any", {"action": "保留待人工"})


def test_tampered_source_actual_override_cannot_be_imported(tmp_path):
    first, second = tmp_path / "first.pdf", tmp_path / "second.pdf"
    pdf(first, title="A")
    pdf(second, title="B")
    source, target = tmp_path / "source", tmp_path / "target"
    source_manifest, _ = project(source, first, session="source", event_positions=(0,))
    project(target, second, session="target")
    add_override(source, source_manifest["records"][0])
    before = (target / "人工判定資料庫.json").read_bytes()
    with pytest.raises(ValueError, match="fingerprint 不符"):
        portable.import_project_decisions(source, target)
    assert (target / "人工判定資料庫.json").read_bytes() == before


@pytest.mark.parametrize("restore_source_before_import", [False, True])
def test_source_actual_override_change_after_fingerprint_check_rejects_public_import(
        tmp_path, monkeypatch, restore_source_before_import):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = tmp_path / "first.pdf", tmp_path / "second.pdf"
    pdf(first, title="A")
    pdf(second, title="B")
    assert sha(first) != sha(second)
    source, target = tmp_path / "source", tmp_path / "target"
    source_manifest, _ = project(source, first, session="A", event_positions=(0,))
    target_manifest, _ = project(target, second, session="B")
    add_override(source, source_manifest["records"][0], reading="ㄐㄩㄝˊ")
    reseal_actual_dynamic(source, first)
    add_override(target, target_manifest["records"][0], reading="ㄐㄩㄝˇ")
    reseal_actual_dynamic(target, second)
    watched = [target / name for name in (
        "人工判定資料庫.json", portable.CONFLICT_FILE, portable.INCOMPLETE_FILE,
        "pipeline_status.json", "注音校對_最終報告.xlsx",
    )]
    before = {path: path.read_bytes() if path.exists() else None for path in watched}
    original_hashes = ar.dynamic_actual_hashes
    original_read_csv = ar._read_csv
    source_root = sp.project_actual_evidence_root(source).resolve()
    changed = False
    restored = False

    def change_source_after_hash(root, *args, **kwargs):
        nonlocal changed
        result = original_hashes(root, *args, **kwargs)
        if Path(root).resolve() == source_root and not changed:
            changed = True
            add_override(source, source_manifest["records"][0], reading="ㄐㄩㄝˇ")
        return result

    def read_then_optionally_restore(path, headers):
        nonlocal restored
        rows = original_read_csv(path, headers)
        if (restore_source_before_import and changed and not restored
                and Path(path).resolve() == source_root / ar.OCCURRENCE_OVERRIDE_FILE):
            # A path rehash would see the sealed bytes again, while the caller
            # has already consumed the unsealed rows returned above.
            restored = True
            add_override(source, source_manifest["records"][0], reading="ㄐㄩㄝˊ")
        return rows

    with patch.object(ar, "dynamic_actual_hashes", side_effect=change_source_after_hash), \
         patch.object(ar, "_read_csv", side_effect=read_then_optionally_restore):
        with pytest.raises(ValueError, match="fingerprint 不符"):
            portable.import_project_decisions(source, target)
    assert changed
    assert restored == restore_source_before_import
    assert {path: path.read_bytes() if path.exists() else None for path in watched} == before
    assert ar._read_csv(source_root / ar.OCCURRENCE_OVERRIDE_FILE, ar.OVERRIDE_HEADERS)[0][
        "actual_reading"] == ("ㄐㄩㄝˊ" if restore_source_before_import else "ㄐㄩㄝˇ")
    assert not (tmp_path / "isolated-localappdata").exists()


def test_unrelated_pdf_actual_override_does_not_invalidate_sealed_source(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = tmp_path / "first.pdf", tmp_path / "second.pdf"
    pdf(first, title="A")
    pdf(second, title="B")
    source, target = tmp_path / "source", tmp_path / "target"
    source_manifest, _ = project(source, first, session="A", event_positions=(0,))
    target_manifest, _ = project(target, second, session="B")
    sealed_row = add_override(source, source_manifest["records"][0])
    reseal_actual_dynamic(source, first)
    add_override(target, target_manifest["records"][0])
    reseal_actual_dynamic(target, second)
    other_row = dict(sealed_row, pdf_contains="unrelated-pdf", actual_reading="ㄐㄩㄝˇ")
    ar._write_csv(sp.project_actual_evidence_root(source) / ar.OCCURRENCE_OVERRIDE_FILE,
                  ar.OVERRIDE_HEADERS, [sealed_row, other_row])

    result = portable.import_project_decisions(source, target)
    assert result["imported"] == 1
    assert result["matched_actual_overrides"] == 1
    assert sp.json_load_strict(target / "人工判定資料庫.json")["events"]
    assert not (tmp_path / "isolated-localappdata").exists()


def test_saved_actual_override_maps_without_global_promotion(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, local = tmp_path / "first.pdf", tmp_path / "local.pdf"
    pdf(first, title="A")
    pdf(local, title="B")
    source, continued = tmp_path / "source", tmp_path / "continued"
    source_manifest, _ = project(source, first, session="A", event_positions=(0,))
    add_override(source, source_manifest["records"][0])
    reseal_actual_dynamic(source, first)
    portable.prepare_portable_project(source)
    first.rename(tmp_path / "first-away.pdf")

    def build_target(paths, output_dir, **kwargs):
        project(output_dir, local, session="B")

    def refresh_target(output_dir, **kwargs):
        reseal_actual_dynamic(output_dir, local)

    with patch.object(sp, "run_pipeline_pdfs", side_effect=build_target), \
         patch.object(sp, "refresh_actual_project", side_effect=refresh_target):
        result = portable.continue_project(source, local, continued)
    assert result["matched_actual_overrides"] == 1
    target_rows = ar._read_csv(sp.project_actual_evidence_root(continued) / ar.OCCURRENCE_OVERRIDE_FILE,
                               ar.OVERRIDE_HEADERS)
    assert len(target_rows) == 1
    assert target_rows[0]["pdf_contains"] == local.stem
    assert target_rows[0]["actual_reading"] == "ㄐㄩㄝˊ"
    assert source_manifest["records"][0]["occurrence_id"] in target_rows[0]["note"]
    portable.prepare_portable_project(continued)
    (tmp_path / "first-away.pdf").rename(first)
    back = portable.import_project_decisions(continued, source)
    assert back["matched_actual_overrides"] == 1
    assert back["duplicates"] == 1
    assert ar._read_csv(sp.project_actual_evidence_root(source) / ar.OCCURRENCE_OVERRIDE_FILE,
                        ar.OVERRIDE_HEADERS)[0]["actual_reading"] == "ㄐㄩㄝˊ"
    assert not (tmp_path / "isolated-localappdata").exists()


def test_existing_target_rejects_unmapped_actual_without_changing_either_project(tmp_path):
    first, second = tmp_path / "first.pdf", tmp_path / "second.pdf"
    pdf(first, title="A")
    pdf(second, title="B")
    source, target = tmp_path / "source", tmp_path / "target"
    source_manifest, _ = project(source, first, session="A", event_positions=(0,))
    project(target, second, session="B")
    add_override(source, source_manifest["records"][0])
    reseal_actual_dynamic(source, first)
    before_source = (source / "人工判定資料庫.json").read_bytes()
    before_target = (target / "人工判定資料庫.json").read_bytes()
    with pytest.raises(ValueError, match="缺少 1 筆"):
        portable.import_project_decisions(source, target)
    assert (source / "人工判定資料庫.json").read_bytes() == before_source
    assert (target / "人工判定資料庫.json").read_bytes() == before_target


def test_two_independent_projects_merge_in_new_project_and_preserve_sources(tmp_path):
    files = [tmp_path / f"{letter}.pdf" for letter in "ABC"]
    for letter, path in zip("ABC", files):
        pdf(path, title=letter)
    first, second, merged = [tmp_path / name for name in ("first", "second", "merged")]
    project(first, files[0], session="A", event_positions=(0,))
    project(second, files[1], session="B", event_positions=(1,))
    original_a = (first / "人工判定資料庫.json").read_bytes()
    original_b = (second / "人工判定資料庫.json").read_bytes()

    def build_target(paths, output_dir, **kwargs):
        project(output_dir, files[2], session="C")

    with patch.object(sp, "run_pipeline_pdfs", side_effect=build_target):
        result = portable.merge_projects([first, second], files[2], merged)
    assert result["imported"] == 2 and result["conflicts"] == []
    assert len(sp.json_load_strict(merged / "人工判定資料庫.json")["events"]) == 2
    assert (first / "人工判定資料庫.json").read_bytes() == original_a
    assert (second / "人工判定資料庫.json").read_bytes() == original_b
    assert not (merged / portable.INCOMPLETE_FILE).exists()


def test_same_decision_from_two_sources_keeps_both_exact_identities(tmp_path):
    files = [tmp_path / f"{letter}.pdf" for letter in "ABC"]
    for letter, path in zip("ABC", files):
        pdf(path, title=letter)
    first, second, merged = [tmp_path / name for name in ("first", "second", "merged")]
    first_manifest, _ = project(first, files[0], session="A", event_positions=(0,))
    second_manifest, _ = project(second, files[1], session="B", event_positions=(0,))

    def build_target(paths, output_dir, **kwargs):
        project(output_dir, files[2], session="C")

    with patch.object(sp, "run_pipeline_pdfs", side_effect=build_target):
        result = portable.merge_projects([first, second], files[2], merged)
    assert result["sources"][1]["duplicates"] == 1
    target_db = sp.json_load_strict(merged / "人工判定資料庫.json")
    event = next(iter(target_db["events"].values()))
    primary = event["portability_source"]
    duplicate = event["portability_duplicate_sources"]
    assert (primary["session_id"], primary["pdf_sha256"], primary["occurrence_id"], primary["review_id"]) == (
        "A", sha(files[0]), first_manifest["records"][0]["occurrence_id"],
        first_manifest["records"][0]["review_id"],
    )
    assert len(duplicate) == 1
    assert (duplicate[0]["session_id"], duplicate[0]["pdf_sha256"],
            duplicate[0]["occurrence_id"], duplicate[0]["review_id"]) == (
        "B", sha(files[1]), second_manifest["records"][0]["occurrence_id"],
        second_manifest["records"][0]["review_id"],
    )
    assert portable.import_project_decisions(second, merged)["duplicates"] == 1
    assert sp.json_load_strict(merged / "人工判定資料庫.json") == target_db


def test_two_sources_keep_both_actual_staging_originals_without_promotion(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    files = [tmp_path / f"{letter}.pdf" for letter in "ABC"]
    for letter, path in zip("ABC", files):
        pdf(path, title=letter)
    first, second, merged = [tmp_path / name for name in ("first", "second", "merged")]
    first_manifest, _ = project(first, files[0], session="A", event_positions=(0,))
    second_manifest, _ = project(second, files[1], session="B", event_positions=(1,))
    originals = []
    for folder, manifest, index in ((first, first_manifest, 0), (second, second_manifest, 1)):
        group = ar.build_actual_group_for_entry(manifest["records"], manifest["records"][index])
        evidence_root = sp.project_actual_evidence_root(folder)
        ar.stage_manual_actual_group(
            evidence_root, group, "ㄐㄩㄝˊ",
            checked_occurrence_ids=[manifest["records"][index]["occurrence_id"]],
            source="manual visual actual confirmation",
        )
        originals.append(ar.load_manual_actual_staging(evidence_root))

    def build_target(paths, output_dir, **kwargs):
        project(output_dir, files[2], session="C")

    with patch.object(sp, "run_pipeline_pdfs", side_effect=build_target):
        result = portable.merge_projects([first, second], files[2], merged)
    assert result["imported"] == 2
    receipt = sp.json_load_strict(merged / portable.PENDING_ACTUAL_FILE)
    assert [source["source_session_id"] for source in receipt["sources"]] == ["A", "B"]
    assert [source["staging"] for source in receipt["sources"]] == originals
    assert portable.import_project_decisions(first, merged)["duplicates"] == 1
    assert sp.json_load_strict(merged / portable.PENDING_ACTUAL_FILE) == receipt
    assert ar.load_manual_actual_staging(sp.project_actual_evidence_root(merged))["staged_groups"] == []
    assert not (tmp_path / "isolated-localappdata").exists()
    assert not (merged / portable.INCOMPLETE_FILE).exists()


def test_b_new_actual_override_returns_to_a_in_new_project(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first_pdf, b_pdf = tmp_path / "first.pdf", tmp_path / "download-b.pdf"
    pdf(first_pdf, title="A")
    pdf(b_pdf, title="B")
    first, b, back = [tmp_path / name for name in ("first-project", "b-project", "back-on-a")]
    project(first, first_pdf, session="A", event_positions=(0,))
    portable.prepare_portable_project(first)
    first_pdf.rename(tmp_path / "missing-on-b.pdf")

    def build_target(paths, output_dir, **kwargs):
        project(output_dir, paths[0], session="B" if output_dir == b else "A-return")

    def refresh_target(output_dir, **kwargs):
        reseal_actual_dynamic(output_dir, b_pdf if output_dir == b else first_pdf)

    with patch.object(sp, "run_pipeline_pdfs", side_effect=build_target), \
         patch.object(sp, "refresh_actual_project", side_effect=refresh_target):
        portable.continue_project(first, b_pdf, b)
        b_manifest = sp.json_load_strict(b / "校對工作階段.json")
        add_override(b, b_manifest["records"][1])
        reseal_actual_dynamic(b, b_pdf)
        portable.prepare_portable_project(b)
        (tmp_path / "missing-on-b.pdf").rename(first_pdf)
        result = portable.merge_projects([first, b], first_pdf, back)
    assert result["matched_actual_overrides"] == 1
    assert result["sources"][1]["duplicates"] == 1
    returned = ar._read_csv(sp.project_actual_evidence_root(back) / ar.OCCURRENCE_OVERRIDE_FILE,
                            ar.OVERRIDE_HEADERS)
    assert len(returned) == 1
    assert returned[0]["actual_reading"] == "ㄐㄩㄝˊ"
    assert b_manifest["records"][1]["occurrence_id"] in returned[0]["note"]
    assert len(sp.json_load_strict(back / "人工判定資料庫.json")["events"]) == 1
    assert not (tmp_path / "isolated-localappdata").exists()


def test_two_project_conflicting_decisions_keep_both_sources(tmp_path):
    files = [tmp_path / f"{letter}.pdf" for letter in "ABC"]
    for letter, path in zip("ABC", files):
        pdf(path, title=letter)
    first, second, merged = [tmp_path / name for name in ("first", "second", "merged")]
    project(first, files[0], session="A", event_positions=(0,))
    second_manifest, db = project(second, files[1], session="B", event_positions=(0,))
    db["events"][second_manifest["records"][0]["review_id"]]["expected_evidence"] = "另一來源"
    sp.json_save(second / "人工判定資料庫.json", db)
    before_a = (first / "人工判定資料庫.json").read_bytes()
    before_b = (second / "人工判定資料庫.json").read_bytes()

    def build_target(paths, output_dir, **kwargs):
        project(output_dir, files[2], session="C")

    with patch.object(sp, "run_pipeline_pdfs", side_effect=build_target):
        result = portable.merge_projects([first, second], files[2], merged)
    assert len(result["conflicts"]) == 1
    receipt = sp.json_load_strict(merged / portable.CONFLICT_FILE)
    assert receipt["conflicts"][0]["source_event"]["expected_evidence"] == "另一來源"
    assert receipt["conflicts"][0]["target_event"]["expected_evidence"] == "現版手冊"
    conflicted_id = receipt["conflicts"][0]["target_review_id"]
    merged_manifest = sp.json_load_strict(merged / "校對工作階段.json")
    merged_db = sp.json_load_strict(merged / "人工判定資料庫.json")
    assert conflicted_id not in merged_db["events"]
    assert sp.materialize_ledger(merged_manifest, merged_db)[0]["active_review"] is True
    assert (first / "人工判定資料庫.json").read_bytes() == before_a
    assert (second / "人工判定資料庫.json").read_bytes() == before_b
    from standalone_gui import project_status_details
    assert "衝突" in project_status_details(merged)[0]
    fourth, transferred = tmp_path / "D.pdf", tmp_path / "transferred"
    pdf(fourth, title="D")
    with pytest.raises(ValueError, match="未裁決"):
        portable.continue_project(merged, fourth, transferred)
    assert not transferred.exists()
    ReviewSaveService(merged).save_event(
        conflicted_id, {"action": "確認非校對範圍", "exclusion_reason": "fresh local adjudication",
                        "exclusion_evidence": "PDF page 1 visual check"})
    portable.prepare_portable_project(merged)

    def build_fourth(paths, output_dir, **kwargs):
        project(output_dir, fourth, session="D")

    with patch.object(sp, "run_pipeline_pdfs", side_effect=build_fourth):
        portable.continue_project(merged, fourth, transferred)
    later_event = next(iter(sp.json_load_strict(transferred / "人工判定資料庫.json")["events"].values()))
    assert later_event["portability_conflict_resolution"]["original_conflicts"] == receipt["conflicts"]


def test_manual_expected_event_rebinds_exact_target_and_keeps_origin(tmp_path):
    first, second = tmp_path / "source.pdf", tmp_path / "local.pdf"
    pdf(first, title="source")
    pdf(second, title="local")
    assert sha(first) != sha(second)
    source, target = tmp_path / "source", tmp_path / "target"
    source_manifest, _ = project(source, first, session="A")
    target_manifest, _ = project(target, second, session="B")
    source_row = source_manifest["records"][0]
    event = sp.build_manual_expected_event(source_row, operation="ENTER_EXPECTED",
                                            expected_set=["ㄐㄩㄝˊ"], rationale="教材頁面核對")
    source_db = sp.json_load_strict(source / "人工判定資料庫.json")
    source_db["events"][source_row["review_id"]] = event
    sp.json_save(source / "人工判定資料庫.json", source_db)
    portable.import_project_decisions(source, target)
    target_row = target_manifest["records"][0]
    target_event = sp.json_load_strict(target / "人工判定資料庫.json")["events"][target_row["review_id"]]
    assert target_event["manual_expected_decision"]["target"]["pdf_sha256"] == sha(second)
    assert target_event["portability_source"]["source_manual_expected_decision"] == event["manual_expected_decision"]
    assert target_event["portability_source"]["review_id"] == source_row["review_id"]
    assert sp.materialize_ledger(target_manifest, sp.json_load_strict(target / "人工判定資料庫.json"))[0]["expected_set"] == ["ㄐㄩㄝˊ"]
    ReviewSaveService(target).save_event(target_row["review_id"], target_event)
    assert portable._load_project(target)[0]["session_id"] == "B"

    tampered, rejected = tmp_path / "tampered", tmp_path / "rejected"
    project(tampered, first, session="T")
    project(rejected, second, session="R")
    tampered_db = sp.json_load_strict(tampered / "人工判定資料庫.json")
    tampered_event = dict(event)
    tampered_event["manual_expected_decision"] = json.loads(json.dumps(event["manual_expected_decision"]))
    tampered_event["manual_expected_decision"]["target"]["pdf_sha256"] = "0" * 64
    tampered_id = sp.json_load_strict(tampered / "校對工作階段.json")["records"][0]["review_id"]
    tampered_db["events"][tampered_id] = tampered_event
    sp.json_save(tampered / "人工判定資料庫.json", tampered_db)
    before = (rejected / "人工判定資料庫.json").read_bytes()
    with pytest.raises(Exception, match="review event replay 失敗"):
        portable.import_project_decisions(tampered, rejected)
    assert (rejected / "人工判定資料庫.json").read_bytes() == before


@pytest.mark.parametrize("root", ["[]", "null", "false", "0"])
@pytest.mark.parametrize("corrupt_source", [True, False])
def test_falsy_nonobject_database_root_is_rejected_without_mutation(tmp_path, root, corrupt_source):
    first, second = tmp_path / "a.pdf", tmp_path / "b.pdf"
    pdf(first, title="A")
    pdf(second, title="B")
    source, target = tmp_path / "source", tmp_path / "target"
    project(source, first, session="A", event_positions=(0,))
    project(target, second, session="B")
    bad = (source if corrupt_source else target) / "人工判定資料庫.json"
    bad.write_text(root, encoding="utf-8")
    before = (target / "人工判定資料庫.json").read_bytes()
    with pytest.raises(ValueError, match="根節點"):
        portable.import_project_decisions(source, target)
    assert (target / "人工判定資料庫.json").read_bytes() == before


def test_target_db_becomes_nonobject_before_locked_candidate_read(tmp_path):
    first, second = tmp_path / "a.pdf", tmp_path / "b.pdf"
    pdf(first, title="A")
    pdf(second, title="B")
    source, target = tmp_path / "source", tmp_path / "target"
    project(source, first, session="A", event_positions=(0,))
    project(target, second, session="B")
    original_mapper = portable._mapped_occurrence_overrides

    def change_target_db(*args):
        (target / "人工判定資料庫.json").write_text("[]", encoding="utf-8")
        return original_mapper(*args)

    with patch.object(portable, "_mapped_occurrence_overrides", side_effect=change_target_db):
        with pytest.raises(ValueError, match="根節點"):
            portable.import_project_decisions(source, target)
    assert (target / "人工判定資料庫.json").read_text(encoding="utf-8") == "[]"


def test_missing_source_uses_only_explicit_identical_sha_local_pdf(tmp_path):
    other_download = tmp_path / "other-download"
    other_download.mkdir()
    original, local, changed = tmp_path / "original.pdf", tmp_path / "local.pdf", other_download / "original.pdf"
    pdf(original, title="A")
    local.write_bytes(original.read_bytes())
    pdf(changed, title="B")
    source, target, refused = [tmp_path / name for name in ("source", "target", "refused")]
    project(source, original, session="A", event_positions=(0,))
    original.rename(tmp_path / "source-away.pdf")

    def build_target(paths, output_dir, **kwargs):
        project(output_dir, paths[0], session="B")

    with patch.object(sp, "run_pipeline_pdfs", side_effect=build_target):
        with pytest.raises(FileNotFoundError):
            portable.continue_project(source, changed, refused)
        result = portable.continue_project(source, local, target)
    assert result["imported"] == 1
    assert not refused.exists()
    assert portable._load_project(target)[0]["pdfs"][0]["pdf_sha256"] == sha(local)


def test_noop_and_clear_do_not_adjudicate_conflict(tmp_path):
    first, second, third = [tmp_path / f"{name}.pdf" for name in "ABC"]
    for name, path in zip("ABC", (first, second, third)):
        pdf(path, title=name)
    a, b, merged = [tmp_path / name for name in ("a", "b", "merged")]
    project(a, first, session="A", event_positions=(0,))
    manifest, db = project(b, second, session="B", event_positions=(0,))
    db["events"][manifest["records"][0]["review_id"]]["expected_evidence"] = "other"
    sp.json_save(b / "人工判定資料庫.json", db)

    def build(paths, output_dir, **kwargs):
        project(output_dir, paths[0], session="M")

    with patch.object(sp, "run_pipeline_pdfs", side_effect=build):
        portable.merge_projects([a, b], third, merged)
    merged_manifest = sp.json_load_strict(merged / "校對工作階段.json")
    review_id = merged_manifest["records"][0]["review_id"]
    receipt = sp.json_load_strict(merged / portable.CONFLICT_FILE)
    saver = ReviewSaveService(merged)
    saver.save_event(review_id, {"action": "保留待人工"})
    after_noop = sp.json_load_strict(merged / "人工判定資料庫.json")
    assert "portability_conflict_resolution" not in after_noop["events"][review_id]
    assert portable.validate_conflict_state(merged, merged_manifest, after_noop) == [review_id]
    saver.save_event(review_id, None)
    assert portable.validate_conflict_state(merged, merged_manifest, sp.json_load_strict(merged / "人工判定資料庫.json")) == [review_id]
    before_imported = (merged / "人工判定資料庫.json").read_bytes()
    with pytest.raises(ValueError, match="本地重新核對"):
        saver.save_event(review_id, {"action": "確認非校對範圍", "exclusion_reason": "copied",
                                     "exclusion_evidence": "source page",
                                     "portability_source": {"session_id": "A"}})
    assert (merged / "人工判定資料庫.json").read_bytes() == before_imported
    with pytest.raises(ValueError, match="未裁決"):
        portable.prepare_portable_project(merged)
    saver.save_event(review_id, {"action": "確認非校對範圍", "exclusion_reason": "local",
                                 "exclusion_evidence": "local PDF page 1"})
    resolved = sp.json_load_strict(merged / "人工判定資料庫.json")
    assert portable.validate_conflict_state(merged, merged_manifest, resolved) == []
    assert resolved["events"][review_id]["portability_conflict_resolution"]["original_conflicts"] == receipt["conflicts"]


def test_duplicate_carries_source_conflict_adjudication_across_hops(tmp_path):
    files = [tmp_path / f"{letter}.pdf" for letter in "ABCDE"]
    for letter, path in zip("ABCDE", files):
        pdf(path, title=letter)
    a, b, c, d, e = [tmp_path / letter for letter in "abcde"]
    project(a, files[0], session="A", event_positions=(0,))
    b_manifest, b_db = project(b, files[1], session="B", event_positions=(0,))
    b_db["events"][b_manifest["records"][0]["review_id"]]["expected_evidence"] = "other"
    sp.json_save(b / "人工判定資料庫.json", b_db)

    def build(paths, output_dir, **kwargs):
        project(output_dir, paths[0], session="C")

    with patch.object(sp, "run_pipeline_pdfs", side_effect=build):
        portable.merge_projects([a, b], files[2], c)
    c_id = sp.json_load_strict(c / "校對工作階段.json")["records"][0]["review_id"]
    ReviewSaveService(c).save_event(c_id, {"action": "確認非校對範圍", "exclusion_reason": "adjudicated",
                                          "exclusion_evidence": "C page 1"})
    c_event = sp.json_load_strict(c / "人工判定資料庫.json")["events"][c_id]
    d_manifest, d_db = project(d, files[3], session="D")
    d_id = d_manifest["records"][0]["review_id"]
    d_db["events"][d_id] = {key: value for key, value in c_event.items()
                             if key != "portability_conflict_resolution"}
    sp.json_save(d / "人工判定資料庫.json", d_db)
    portable.prepare_portable_project(c)
    assert portable.import_project_decisions(c, d)["duplicates"] == 1
    duplicate = sp.json_load_strict(d / "人工判定資料庫.json")["events"][d_id]
    origin = duplicate["portability_duplicate_sources"][0]
    assert origin["source_conflict_resolution"] == c_event["portability_conflict_resolution"]
    project(e, files[4], session="E")
    portable.prepare_portable_project(d)
    portable.import_project_decisions(d, e)
    e_event = next(iter(sp.json_load_strict(e / "人工判定資料庫.json")["events"].values()))
    assert e_event["portability_source"]["prior_duplicate_sources"][0]["source_conflict_resolution"] == c_event["portability_conflict_resolution"]
