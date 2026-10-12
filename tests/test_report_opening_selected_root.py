"""Real report handlers and sealed-copy validation; no Tk or PDF import.

The finite synthetic producer writes actual XLSX files, a sealed manifest,
normalized decision DB and a sealed conflict receipt with local CSV provenance.
Those bytes are test evidence, never a real user's visual attestation.
"""
import hashlib
import shutil
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
import pytest
import actual_review as ar
import pdf_portability as portable
import sealed_workbook_paths as locations
import standalone_gui as gui
import standalone_proofread as sp
from occurrence_ledger import build_occurrence_ledger, prepare_occurrence_rows

REPORTS = {"open_user_report": "注音校對_最終報告.xlsx",
           "open_technical_report": "注音校對_技術稽核.xlsx"}


@pytest.fixture
def report_copy(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-global"))
    origin, selected = tmp_path / "original project", tmp_path / "selected copy"
    origin.mkdir()
    source = origin / "source.pdf"
    source.write_bytes(b"synthetic sealed PDF identity; not imported or rendered")
    source_sha = hashlib.sha256(source.read_bytes()).hexdigest()
    row = {"pdf_sha256": source_sha, "pdf": str(source), "pdf_name": source.name,
           "實體頁碼": 1, "課本頁": 1, "字元": "一", "實際注音": "ㄧ",
           "解碼依據": "independent synthetic glyph", "穩定注音鍵": "fixture#1",
           "font": "fixture", "font_xref": 1, "TTF字形SHA256": "a" * 64,
           "glyph_id_字形索引": 1, "注音元件ID": 1,
           "x0": 20, "y0": 20, "x1": 40, "y1": 40, "source_row_number": 1}
    prepare_occurrence_rows(source_sha, [row])
    records = build_occurrence_ledger([row], {row["occurrence_id"]: {
        "state": "EXPECTED_UNRESOLVED", "actual": "ㄧ", "actual_evidence": "synthetic glyph",
        "context_evidence": "fixture context", "source_view": "未建立預期音", "source_record": row}})
    info = {"pdf": str(source), "pdf_name": source.name, "pdf_sha256": source_sha}
    for kind, role in locations.ROLES.items():
        path = origin / role / (kind + ".xlsx");path.parent.mkdir()
        workbook = Workbook();workbook.active.append([kind, "independent fixture"])
        workbook.save(path);workbook.close()
        info[kind + "_workbook"] = str(path)
        info[kind + "_workbook_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
    manifest = sp.seal_manifest({"session_id": "report-selected-root-fixture", "pdfs": [info],
                                 "records": records, "actual_source_ids": [row["occurrence_id"]]})
    db = sp.normalize_db({})
    entry = records[0]
    key = ar._override_key_from_entry(entry)
    fields = ("pdf_contains", "pdf_excludes", "page", "target_char", "stable_key", "x0", "y0")
    target = dict(zip(fields, key)) | {"actual_reading": "ㄧ", "source": "original fixture", "note": ""}
    source_excel = origin / "synthetic-source-decisions.xlsx"
    workbook = Workbook();workbook.active.append(["decision", "actual_reading"])
    workbook.active.append(["VERIFIED", "ㄧˊ"]);workbook.save(source_excel);workbook.close()
    source_manifest = sp.seal_manifest({"session_id": "independent-source-session", "records": []})
    provenance = {"source_session_id": source_manifest["session_id"],
        "source_manifest_integrity_sha256": source_manifest["manifest_integrity_sha256"],
        "source_pdf_sha256": ["d" * 64], "source_group_id": "source-group",
        "source_group_snapshot": "source-snapshot", "source_occurrence_ids": ["source-occ"],
        "source_review_ids": ["source-review"],
        "source_roster": [{"occurrence_id": "source-occ", "review_id": "source-review", "pdf_sha256": "d" * 64}],
        "target_session_id": manifest["session_id"], "target_occurrence_ids": [entry["occurrence_id"]],
        "source_excel_sha256": hashlib.sha256(source_excel.read_bytes()).hexdigest(), "source_excel_row": 2,
        "source_decision": {"decision": "VERIFIED", "actual_reading": "ㄧˊ", "group_id": "source-group",
                            "group_snapshot": "source-snapshot", "sample_a_occurrence_id": "source-occ"}}
    conflict = {"target_occurrence_id": entry["occurrence_id"], "target_review_id": entry["review_id"],
                "target_pdf_sha256": source_sha, "target_record": target,
                "source_record": provenance, "source_reading": "ㄧˊ"}
    evidence_root = sp.project_actual_evidence_root(origin)
    ar.ensure_user_evidence_files(evidence_root)
    note = portable._actual_excel_conflict_token(conflict)
    ar._write_csv(evidence_root / ar.OCCURRENCE_OVERRIDE_FILE, ar.OVERRIDE_HEADERS,
                  [target | {"source": "人工 GUI actual 視覺確認", "note": note}])
    ar._write_csv(evidence_root / ar.GLYPH_PROVENANCE_FILE, ar.GLYPH_PROVENANCE_HEADERS,
                  [{"event_type": "OCCURRENCE_OVERRIDE", "source": "人工 GUI actual 視覺確認",
                    "note": note, "occurrence_ids": entry["occurrence_id"], "bopomofo": "ㄧ"}])
    sp.json_save(origin / "校對工作階段.json", manifest)
    sp.json_save(origin / "人工判定資料庫.json", db)
    sp.json_save(origin / portable.ACTUAL_EXCEL_CONFLICT_FILE, portable._sealed_actual_excel_conflicts([conflict]))
    sp.json_save(origin / "pipeline_status.json", {"excel_report_deferred": False})
    for name in REPORTS.values():
        workbook = Workbook();workbook.active.append(["report", name]);workbook.save(origin / name);workbook.close()
    shutil.copytree(origin, selected)
    assert portable.actual_excel_conflict_state(selected, manifest, db) == []
    return origin, selected, manifest, entry


def invoke(selected, method):
    app = gui.App.__new__(gui.App)
    app.output = type("SelectedRoot", (), {"get": lambda self: str(selected)})()
    with patch.object(gui.messagebox, "showerror") as errors, patch.object(gui.messagebox, "showinfo") as infos, \
            patch.object(gui.os, "startfile", create=True) as windows_open, \
            patch.object(gui.subprocess, "Popen") as other_open, \
            patch.object(locations, "_digest", wraps=locations._digest) as digest:
        getattr(app, method)()
        return errors.call_args_list, infos.call_args_list, windows_open.call_args_list, other_open.call_args_list, digest.call_args_list


@pytest.mark.parametrize("method", list(REPORTS), ids=["user", "technical"])
@pytest.mark.parametrize("old_state", ["intact", "absent", "different"])
def test_both_reports_read_only_selected_copy(report_copy, method, old_state):
    origin, selected, manifest, _ = report_copy
    if old_state == "absent":shutil.rmtree(origin)
    if old_state == "different":
        for info in manifest["pdfs"]:
            for kind in locations.ROLES:Path(info[kind + "_workbook"]).write_bytes(b"old root different bytes")
    errors, infos, opened, other, hashes = invoke(selected, method)
    assert not errors and not infos
    assert opened == [((str(selected / REPORTS[method]),), {})] if gui.os.name == "nt" else not opened
    if gui.os.name != "nt":assert other == [((["xdg-open", str(selected / REPORTS[method])],), {})]
    else:assert not other
    assert len(hashes) >= 2
    assert all(Path(call.args[0]).is_relative_to(selected) for call in hashes)


@pytest.mark.parametrize("kind,mutation,method", [
    ("actual", "missing", "open_user_report"), ("actual", "wrong-sha", "open_technical_report"),
    ("candidate", "missing", "open_technical_report"), ("candidate", "wrong-sha", "open_user_report")],
    ids=["actual-missing", "actual-sha", "candidate-missing", "candidate-sha"])
def test_invalid_selected_workbook_never_falls_back_to_intact_origin(report_copy, kind, mutation, method):
    origin, selected, manifest, _ = report_copy
    descriptor = locations.descriptor(manifest["pdfs"][0], kind)
    assert (origin / descriptor).is_file()
    path = selected / descriptor
    if mutation == "missing":path.unlink()
    else:path.write_bytes(b"selected bytes do not match seal")
    errors, _, opened, other, hashes = invoke(selected, method)
    assert errors and not opened and not other
    assert all(Path(call.args[0]).is_relative_to(selected) for call in hashes)


@pytest.mark.parametrize("method", list(REPORTS), ids=["user", "technical"])
@pytest.mark.parametrize("state", ["pending", "deferred"])
def test_real_unresolved_conflict_and_deferred_report_remain_blocked(report_copy, method, state):
    _, selected, manifest, entry = report_copy
    if state == "pending":
        root = sp.project_actual_evidence_root(selected)
        ar._write_csv(root / ar.GLYPH_PROVENANCE_FILE, ar.GLYPH_PROVENANCE_HEADERS, [])
        assert portable.actual_excel_conflict_state(selected, manifest, sp.json_load_strict(selected / "人工判定資料庫.json")) == [entry["review_id"]]
    else:sp.json_save(selected / "pipeline_status.json", {"excel_report_deferred": True})
    errors, _, opened, other, _ = invoke(selected, method)
    assert errors[0].args[0] == ("報告尚未完成" if state == "pending" else "報告尚待更新")
    assert not opened and not other


@pytest.mark.parametrize("guard,method", [("manifest", "open_user_report"), ("receipt", "open_technical_report"),
                                        ("missing-report", "open_user_report"), ("incomplete", "open_technical_report")],
                         ids=["manifest", "receipt", "missing-report", "incomplete"])
def test_seal_receipt_missing_report_and_incomplete_guards_are_preserved(report_copy, guard, method):
    _, selected, _, _ = report_copy
    if guard == "manifest":
        manifest = sp.json_load_strict(selected / "校對工作階段.json");manifest["session_id"] = "tampered"
        sp.json_save(selected / "校對工作階段.json", manifest)
    elif guard == "receipt":
        path = selected / portable.ACTUAL_EXCEL_CONFLICT_FILE
        receipt = sp.json_load_strict(path);receipt["conflicts"][0]["source_reading"] = "ㄧˋ";sp.json_save(path, receipt)
    elif guard == "missing-report":(selected / REPORTS[method]).unlink()
    else:(selected / "跨電腦接續未完成.json").write_text("{}")
    errors, infos, opened, other, _ = invoke(selected, method)
    assert infos if guard == "missing-report" else errors
    assert not opened and not other
