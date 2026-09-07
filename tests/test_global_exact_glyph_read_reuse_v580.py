from __future__ import annotations

import hashlib
import inspect
import json
import struct
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from openpyxl import Workbook, load_workbook

import actual_review
import cff_zhuyin_decoder as cff_decoder
import check_pronunciation_candidates as candidates
import cross_version_compat
import export_zhuyin_readings as decoder
import global_exact_glyph_library as library
import runtime_source_validation
import standalone_proofread
from exact_glyph_identity import (
    CFF_GLYPH_SHA256,
    GLOBAL_ELIGIBLE_COMPLETE_CFF_RECORDING_V1,
    GLOBAL_ELIGIBLE_SIMPLE_GLYF_V1,
    NON_GLOBAL_ELIGIBLE,
    TTF_GLYF_SHA256,
)


ROOT = Path(__file__).resolve().parents[1]
SHA_A = "a" * 64
SHA_B = "b" * 64
SHA_C = "c" * 64
NOW = "2026-09-03T00:00:00Z"


def _identity(
    sha256: str = SHA_A,
    *,
    kind: str = TTF_GLYF_SHA256,
    style_group: str = "",
) -> library.GlobalExactGlyphIdentity:
    eligibility = (
        GLOBAL_ELIGIBLE_SIMPLE_GLYF_V1
        if kind == TTF_GLYF_SHA256
        else GLOBAL_ELIGIBLE_COMPLETE_CFF_RECORDING_V1
    )
    return library.canonical_global_exact_identity(
        kind,
        style_group,
        sha256,
        eligibility,
    )


def _record(
    identity: library.GlobalExactGlyphIdentity,
    *,
    reading: str = "ㄅ",
    conflict: tuple[str, ...] = (),
    revision: int = 1,
    source_count: int = 2,
    updated_at: str = NOW,
) -> library.GlobalGlyphSnapshotRecord:
    quarantined = bool(conflict)
    return library.GlobalGlyphSnapshotRecord(
        glyph_id=library.compute_global_glyph_id(
            identity.kind,
            identity.style_group,
            identity.glyph_sha256,
            identity_eligibility=identity.identity_eligibility,
        ),
        identity=identity,
        state=(library.QUARANTINED_CONFLICT if quarantined else library.VERIFIED_GLOBAL),
        active_reading=("" if quarantined else reading),
        conflicting_readings=tuple(sorted(conflict)),
        direct_source_count=source_count,
        independent_source_count=source_count,
        revision=revision,
        updated_at=updated_at,
    )


def _snapshot(
    *,
    trusted: tuple[library.GlobalGlyphSnapshotRecord, ...] = (),
    quarantined: tuple[library.GlobalGlyphSnapshotRecord, ...] = (),
    status: str = library.VALID,
    generation: int | None = 1,
    loaded_at: str = NOW,
    provenance_count: int = 0,
    database_path: str = "C:/fixture/library.sqlite3",
) -> library.GlobalExactGlyphSnapshot:
    return library.GlobalExactGlyphSnapshot(
        store_status=status,
        schema_version=(library.GLOBAL_LIBRARY_SCHEMA_VERSION if status == library.VALID else ""),
        identity_contract_version=library.GLOBAL_IDENTITY_CONTRACT_VERSION,
        promotion_policy_version=library.GLOBAL_PROMOTION_POLICY_VERSION,
        generation=(generation if status == library.VALID else None),
        trusted_identities=tuple(sorted(trusted, key=lambda item: item.identity.tuple)),
        quarantined_identities=tuple(sorted(quarantined, key=lambda item: item.identity.tuple)),
        database_path=database_path,
        loaded_at=loaded_at,
        provenance_event_count=provenance_count,
    )


def _absent_snapshot() -> library.GlobalExactGlyphSnapshot:
    return _snapshot(status=library.ABSENT, generation=None)


def _fingerprint_payload(global_component: dict[str, str]) -> dict:
    components = {
        "fingerprint_schema_version": "3.0.0",
        "reuse_policy": "evidence_assets_with_global_exact_v1",
        "actual_decoder_semantics_epoch": "1",
        "pdf_sha256": SHA_C,
        "actual_asset_hashes": {"actual": SHA_A},
        "dynamic_actual_evidence_hashes": {"dynamic": SHA_B},
        "global_exact_glyph_evidence_hashes": global_component,
    }
    reusable = cross_version_compat.fingerprint_contract_components(
        components,
        chain="actual",
    )
    digest = hashlib.sha256(
        json.dumps(
            reusable,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()
    return {"fingerprint": digest, "components": components}


class GlobalFingerprintScopingTests(unittest.TestCase):
    def test_per_pdf_subset_changes_only_for_requested_identity(self):
        identity_x = _identity(SHA_A)
        identity_y = _identity(SHA_B)
        absent = _absent_snapshot()
        x_b = _snapshot(trusted=(_record(identity_x, reading="ㄅ"),))
        x_p = _snapshot(trusted=(_record(identity_x, reading="ㄆ"),))
        x_conflict = _snapshot(
            quarantined=(_record(identity_x, conflict=("ㄅ", "ㄆ")),)
        )
        y_b = _snapshot(trusted=(_record(identity_y, reading="ㄅ"),))

        a_absent = library.global_exact_glyph_evidence_hashes(absent, (identity_x,))
        b_absent = library.global_exact_glyph_evidence_hashes(absent, (identity_y,))
        self.assertNotEqual(
            a_absent,
            library.global_exact_glyph_evidence_hashes(x_b, (identity_x,)),
        )
        self.assertNotEqual(
            library.global_exact_glyph_evidence_hashes(x_b, (identity_x,)),
            library.global_exact_glyph_evidence_hashes(x_p, (identity_x,)),
        )
        self.assertNotEqual(
            library.global_exact_glyph_evidence_hashes(x_b, (identity_x,)),
            library.global_exact_glyph_evidence_hashes(x_conflict, (identity_x,)),
        )
        self.assertEqual(
            b_absent,
            library.global_exact_glyph_evidence_hashes(x_b, (identity_y,)),
        )
        self.assertEqual(
            a_absent,
            library.global_exact_glyph_evidence_hashes(y_b, (identity_x,)),
        )

    def test_generation_timestamps_path_provenance_revision_and_counts_are_audit_only(self):
        identity = _identity()
        first_record = _record(identity, reading="ㄅ")
        first = _snapshot(trusted=(first_record,))
        second = _snapshot(
            trusted=(replace(
                first_record,
                revision=91,
                direct_source_count=23,
                independent_source_count=23,
                updated_at="2026-09-03T12:34:56Z",
            ),),
            generation=999,
            loaded_at="2026-09-03T12:35:00Z",
            provenance_count=10000,
            database_path="D:/unrelated/location.sqlite3",
        )
        self.assertEqual(
            library.global_exact_glyph_evidence_hashes(first, (identity,)),
            library.global_exact_glyph_evidence_hashes(second, (identity,)),
        )

    def test_verified_global_store_disappearance_invalidates_dependent_subset(self):
        identity = _identity()
        present = _snapshot(trusted=(_record(identity, reading="ㄅ"),))
        self.assertNotEqual(
            library.global_exact_glyph_evidence_hashes(present, (identity,)),
            library.global_exact_glyph_evidence_hashes(_absent_snapshot(), (identity,)),
        )

    def test_empty_sealed_roster_is_distinct_from_nonreusable_provisional_roster(self):
        snapshot = _absent_snapshot()
        sealed = library.global_exact_glyph_evidence_hashes(snapshot, ())
        provisional = library.global_exact_glyph_evidence_hashes(snapshot, None)
        self.assertEqual(sealed["scope_mode"], library.GLOBAL_EXACT_PER_PDF_SCOPE_MODE)
        self.assertEqual(provisional["scope_mode"], library.GLOBAL_EXACT_PROVISIONAL_SCOPE_MODE)
        self.assertNotEqual(sealed, provisional)

        sealed_payload = _fingerprint_payload(sealed)
        provisional_payload = _fingerprint_payload(provisional)
        self.assertTrue(cross_version_compat.fingerprint_compatible(
            sealed_payload["fingerprint"],
            sealed_payload["components"],
            sealed_payload,
            chain="actual",
        ))
        self.assertFalse(cross_version_compat.fingerprint_compatible(
            provisional_payload["fingerprint"],
            provisional_payload["components"],
            provisional_payload,
            chain="actual",
        ))

    def test_global_component_is_required_and_v29_has_no_v30_adapter(self):
        component = library.global_exact_glyph_evidence_hashes(_absent_snapshot(), ())
        current = _fingerprint_payload(component)
        missing = dict(current["components"])
        missing.pop("global_exact_glyph_evidence_hashes")
        self.assertFalse(cross_version_compat.fingerprint_compatible(
            current["fingerprint"], missing, current, chain="actual"
        ))
        legacy = dict(current["components"])
        legacy["fingerprint_schema_version"] = "2.9.0"
        legacy["reuse_policy"] = "evidence_assets_v1"
        legacy.pop("global_exact_glyph_evidence_hashes")
        self.assertFalse(cross_version_compat.fingerprint_compatible(
            "legacy", legacy, current, chain="actual"
        ))
        self.assertNotIn(
            ("2.9.0", "3.0.0"),
            cross_version_compat._FINGERPRINT_SCHEMA_ADAPTERS["actual"],
        )

    def test_unknown_identity_or_promotion_contract_fails_closed(self):
        component = library.global_exact_glyph_evidence_hashes(_absent_snapshot(), ())
        for field in ("identity_contract_version", "promotion_policy_version"):
            with self.subTest(field=field):
                malformed = {**component, field: "999.0"}
                with self.assertRaises(library.GlobalLibraryValidationError):
                    library.canonical_global_exact_glyph_evidence_hashes(malformed)


class ExactReusePrecedenceTests(unittest.TestCase):
    def test_global_only_and_project_global_agreement_reuse(self):
        identity = _identity()
        snapshot = _snapshot(trusted=(_record(identity, reading="ㄅ"),))
        global_only = library.resolve_exact_glyph_reuse(snapshot, identity)
        self.assertEqual(global_only.reading, "ㄅ")
        self.assertEqual(global_only.sources, ("GLOBAL_VERIFIED_EXACT",))

        agreement = library.resolve_exact_glyph_reuse(
            snapshot,
            identity,
            higher_priority_sources=(("PROJECT_VERIFIED_EXACT", "ㄅ"),),
            lower_priority_sources=(("STATIC_VERIFIED_EXACT", "ㄅ"),),
        )
        self.assertEqual(agreement.reading, "ㄅ")
        self.assertEqual(
            agreement.sources,
            (
                "PROJECT_VERIFIED_EXACT",
                "GLOBAL_VERIFIED_EXACT",
                "STATIC_VERIFIED_EXACT",
            ),
        )

    def test_project_global_or_global_static_disagreement_suppresses_all_exact_donors(self):
        identity = _identity()
        snapshot = _snapshot(trusted=(_record(identity, reading="ㄅ"),))
        for kwargs in (
            {"higher_priority_sources": (("PROJECT_VERIFIED_EXACT", "ㄆ"),)},
            {"lower_priority_sources": (("STATIC_VERIFIED_EXACT", "ㄆ"),)},
        ):
            with self.subTest(kwargs=kwargs):
                result = library.resolve_exact_glyph_reuse(snapshot, identity, **kwargs)
                self.assertTrue(result.conflict)
                self.assertEqual(result.reading, "")
                self.assertEqual(result.sources, ())
                self.assertEqual(result.conflicting_readings, ("ㄅ", "ㄆ"))

    def test_project_and_global_quarantine_are_union_conflict_gate(self):
        identity = _identity()
        global_conflict = _snapshot(
            quarantined=(_record(identity, conflict=("ㄅ", "ㄆ")),)
        )
        result = library.resolve_exact_glyph_reuse(global_conflict, identity)
        self.assertTrue(result.conflict)
        self.assertEqual(result.reading, "")
        self.assertEqual(result.global_effective_state, library.QUARANTINED_CONFLICT)

        project_conflict = library.resolve_exact_glyph_reuse(
            _absent_snapshot(),
            identity,
            higher_priority_sources=(("PROJECT_VERIFIED_EXACT", "ㄅ"),),
            exact_identity_quarantined=True,
            quarantine_readings=("ㄅ", "ㄆ"),
        )
        self.assertTrue(project_conflict.conflict)
        self.assertEqual(project_conflict.reading, "")

    def test_candidate_and_promotion_ready_are_not_snapshot_reading_donors(self):
        # Phase 2 snapshots intentionally omit CANDIDATE and PROMOTION_READY;
        # their requested logical state is therefore ABSENT for read reuse.
        identity = _identity()
        result = library.resolve_exact_glyph_reuse(_snapshot(), identity)
        self.assertEqual(result.global_effective_state, library.ABSENT)
        self.assertFalse(result.conflict)
        self.assertEqual(result.reading, "")

    def test_ttf_structural_noneligibility_never_calls_global_resolver(self):
        with patch.object(decoder, "resolve_exact_glyph_reuse") as lookup:
            result = decoder.resolve_ttf_exact_reuse(
                _snapshot(trusted=(_record(_identity(), reading="ㄅ"),)),
                {
                    "eligibility": NON_GLOBAL_ELIGIBLE,
                    "glyph_sha256": SHA_A,
                    "reason": "COMPOSITE_GLYF",
                },
            )
        self.assertIsNone(result)
        lookup.assert_not_called()

    def test_ttf_global_label_and_exact_disagreement(self):
        identity = _identity()
        snapshot = _snapshot(trusted=(_record(identity, reading="ㄅ"),))
        structural = {
            "eligibility": GLOBAL_ELIGIBLE_SIMPLE_GLYF_V1,
            "glyph_sha256": SHA_A,
        }
        global_only = decoder.resolve_ttf_exact_reuse(snapshot, structural)
        self.assertEqual(global_only.reading, "ㄅ")
        self.assertIn("GLOBAL_VERIFIED_EXACT", global_only.sources)
        self.assertEqual(decoder.GLOBAL_TTF_SOURCE_LABEL, "Global exact TTF glyf SHA-256")

        conflict = decoder.resolve_ttf_exact_reuse(
            snapshot,
            structural,
            static_record={"bopomofo": "ㄆ"},
        )
        self.assertTrue(conflict.conflict)
        self.assertEqual(conflict.reading, "")


def _simple_glyph(x: int = 0) -> bytes:
    header = struct.pack(">hhhhh", 1, 0, 0, x, 0)
    return header + struct.pack(">HH", 0, 0) + (bytes((0x33, x)) if x else bytes((0x31, 0)))


def _composite_glyph(component_gid: int = 0) -> bytes:
    return struct.pack(">hhhhhHHbb", -1, 0, 0, 0, 0, 0x0002, component_gid, 0, 0)


def _synthetic_sfnt(glyph_records: list[bytes]) -> bytes:
    head = bytearray(54)
    struct.pack_into(">h", head, 50, 1)
    maxp = struct.pack(">IH", 0x00010000, len(glyph_records))
    offsets = [0]
    glyf = bytearray()
    for record in glyph_records:
        glyf.extend(record)
        offsets.append(len(glyf))
    loca = struct.pack(f">{len(offsets)}I", *offsets)
    tables = [(b"head", bytes(head)), (b"maxp", maxp), (b"loca", loca), (b"glyf", bytes(glyf))]
    cursor = 12 + 16 * len(tables)
    directory = bytearray()
    body = bytearray()
    for tag, data in tables:
        padding = (-cursor) % 4
        body.extend(b"\0" * padding)
        cursor += padding
        directory.extend(struct.pack(">4sIII", tag, 0, cursor, len(data)))
        body.extend(data)
        cursor += len(data)
    return struct.pack(">IHHHH", 0x00010000, len(tables), 0, 0, 0) + bytes(directory) + bytes(body)


DEPENDENCY_HEADERS = [
    "實際注音",
    "字形架構",
    "font_xref",
    "注音元件ID",
    "TTF字形SHA256",
    "CFF樣式群組",
    "CFF整字字形SHA256",
]


def _write_dependency_workbook(path: Path, rows: list[dict], headers=DEPENDENCY_HEADERS) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "實際注音"
    worksheet.append(list(headers))
    for row in rows:
        worksheet.append([row.get(name, "") for name in headers])
    workbook.save(path)
    workbook.close()


class _FakePdfDocument:
    def __init__(self, fonts: dict[int, bytes]):
        self.fonts = fonts
        self.closed = False

    def extract_font(self, xref: int):
        return (f"font-{xref}", "ttf", "TrueType", self.fonts[xref])

    def close(self):
        self.closed = True


class WorkbookDependencyRevalidationTests(unittest.TestCase):
    def test_ttf_candidates_are_reparsed_and_include_unresolved_eligible_and_noneligible_rows(self):
        simple = _simple_glyph(17)
        composite = _composite_glyph()
        zero = struct.pack(">hhhhh", 0, 0, 0, 0, 0)
        malformed = b"\0" * 9
        raw_records = [simple, composite, zero, malformed]
        fonts = {index + 1: _synthetic_sfnt([raw]) for index, raw in enumerate(raw_records)}
        rows = [
            {
                "實際注音": "",  # unresolved identities still participate
                "字形架構": "TrueType複合元件",
                "font_xref": index + 1,
                "注音元件ID": 0,
                "TTF字形SHA256": hashlib.sha256(raw).hexdigest(),
            }
            for index, raw in enumerate(raw_records)
        ]
        with tempfile.TemporaryDirectory(dir=ROOT) as directory:
            path = Path(directory) / "actual.xlsx"
            _write_dependency_workbook(path, rows)
            document = _FakePdfDocument(fonts)
            with patch.object(actual_review.fitz, "open", return_value=document):
                dependencies = actual_review.actual_workbook_global_exact_dependencies(
                    path,
                    Path(directory) / "book.pdf",
                )
        self.assertTrue(document.closed)
        self.assertEqual(len(dependencies), 4)
        by_sha = {item.glyph_sha256: item for item in dependencies}
        self.assertEqual(
            by_sha[hashlib.sha256(simple).hexdigest()].identity_eligibility,
            GLOBAL_ELIGIBLE_SIMPLE_GLYF_V1,
        )
        for raw in (composite, zero, malformed):
            self.assertEqual(
                by_sha[hashlib.sha256(raw).hexdigest()].identity_eligibility,
                NON_GLOBAL_ELIGIBLE,
            )

    def test_workbook_pdf_ttf_sha_mismatch_forces_decode(self):
        simple = _simple_glyph(11)
        with tempfile.TemporaryDirectory(dir=ROOT) as directory:
            path = Path(directory) / "actual.xlsx"
            _write_dependency_workbook(path, [{
                "字形架構": "TrueType複合元件",
                "font_xref": 1,
                "注音元件ID": 0,
                "TTF字形SHA256": SHA_A,
            }])
            with patch.object(
                actual_review.fitz,
                "open",
                return_value=_FakePdfDocument({1: _synthetic_sfnt([simple])}),
            ):
                dependencies = actual_review.actual_workbook_global_exact_dependencies(
                    path,
                    Path(directory) / "book.pdf",
                )
        self.assertIsNone(dependencies)

    def test_cff_dependency_is_style_plus_complete_glyph_sha(self):
        with tempfile.TemporaryDirectory(dir=ROOT) as directory:
            path = Path(directory) / "actual.xlsx"
            _write_dependency_workbook(path, [{
                "字形架構": "CFF整字注音",
                "CFF樣式群組": "BIAOKAI_W5",
                "CFF整字字形SHA256": SHA_A,
            }])
            with patch.object(
                actual_review.fitz,
                "open",
                return_value=_FakePdfDocument({}),
            ):
                dependencies = actual_review.actual_workbook_global_exact_dependencies(
                    path,
                    Path(directory) / "book.pdf",
                )
        self.assertEqual(dependencies, (_identity(
            SHA_A,
            kind=CFF_GLYPH_SHA256,
            style_group="BIAOKAI_W5",
        ),))

    def test_missing_exact_columns_force_decode_and_corrupt_workbook_is_integrity_error(self):
        with tempfile.TemporaryDirectory(dir=ROOT) as directory:
            base = Path(directory)
            legacy = base / "legacy.xlsx"
            _write_dependency_workbook(legacy, [], headers=["實際注音", "font_xref"])
            with patch.object(
                actual_review.fitz,
                "open",
                return_value=_FakePdfDocument({}),
            ):
                self.assertIsNone(actual_review.actual_workbook_global_exact_dependencies(
                    legacy,
                    base / "book.pdf",
                ))

            incomplete = base / "incomplete.xlsx"
            _write_dependency_workbook(incomplete, [{
                "字形架構": "TrueType複合元件",
                "font_xref": 1,
                "注音元件ID": 0,
                "TTF字形SHA256": "",
            }])
            with patch.object(
                actual_review.fitz,
                "open",
                return_value=_FakePdfDocument({}),
            ):
                self.assertIsNone(actual_review.actual_workbook_global_exact_dependencies(
                    incomplete,
                    base / "book.pdf",
                ))

            corrupt = base / "corrupt.xlsx"
            corrupt.write_bytes(b"not an xlsx")
            with self.assertRaisesRegex(ValueError, "workbook 損壞"):
                actual_review.actual_workbook_global_exact_dependencies(
                    corrupt,
                    base / "book.pdf",
                )


def _cff_recording():
    return [
        ("moveTo", ((1100, 0),)),
        ("lineTo", ((1210, 0),)),
        ("lineTo", ((1210, 200),)),
        ("lineTo", ((1100, 200),)),
        ("closePath", ()),
    ]


def _cff_inspector(
    snapshot: library.GlobalExactGlyphSnapshot,
    *,
    style_group: str = "BIAOKAI_W5",
    project_reading: str = "",
    symbol_reading: str = "",
    quarantined: bool = False,
):
    recording = _cff_recording()
    glyph_sha = cff_decoder.cff_glyph_outline_sha256(recording)
    contour = cff_decoder._split_contours(recording)[0]
    signature = cff_decoder._signature([contour])
    inspector = object.__new__(cff_decoder.CFFZhuyinInspector)
    inspector.font_name = "fixture"
    inspector.style_group = style_group
    inspector.symbol_map = (
        {(style_group, signature): {"symbol": symbol_reading}}
        if symbol_reading
        else {}
    )
    inspector.verified_glyph_fingerprints = (
        {(style_group, glyph_sha): {"bopomofo": project_reading}}
        if project_reading
        else {}
    )
    inspector.global_snapshot = snapshot
    inspector.quarantined_glyph_identities = (
        {(style_group, glyph_sha)} if quarantined else set()
    )
    inspector._recording = lambda _glyph_id: recording
    return inspector, glyph_sha


class CFFGlobalExactReuseTests(unittest.TestCase):
    def test_global_exact_rescues_normal_symbol_table_unresolved(self):
        probe, glyph_sha = _cff_inspector(_absent_snapshot())
        identity = _identity(
            glyph_sha,
            kind=CFF_GLYPH_SHA256,
            style_group="BIAOKAI_W5",
        )
        inspector, _ = _cff_inspector(
            _snapshot(trusted=(_record(identity, reading="ㄅ"),))
        )
        result = inspector.decode(1)
        self.assertEqual(result["reading"], "ㄅ")
        self.assertEqual(result["exact_source_label"], decoder.GLOBAL_CFF_SOURCE_LABEL)
        self.assertEqual(result["exact_reuse_sources"], ("GLOBAL_VERIFIED_EXACT",))

    def test_same_complete_sha_different_style_does_not_reuse(self):
        _probe, glyph_sha = _cff_inspector(_absent_snapshot())
        other_style = _identity(
            glyph_sha,
            kind=CFF_GLYPH_SHA256,
            style_group="YUAN_W7",
        )
        inspector, _ = _cff_inspector(
            _snapshot(trusted=(_record(other_style, reading="ㄅ"),)),
            style_group="BIAOKAI_W5",
        )
        result = inspector.decode(1)
        self.assertEqual(result["reading"], "")
        self.assertEqual(result["global_effective_state"], library.ABSENT)

    def test_project_global_agreement_label_and_disagreement_fallback(self):
        _probe, glyph_sha = _cff_inspector(_absent_snapshot())
        identity = _identity(
            glyph_sha,
            kind=CFF_GLYPH_SHA256,
            style_group="BIAOKAI_W5",
        )
        snapshot = _snapshot(trusted=(_record(identity, reading="ㄅ"),))
        agreement, _ = _cff_inspector(snapshot, project_reading="ㄅ")
        result = agreement.decode(1)
        self.assertEqual(result["reading"], "ㄅ")
        self.assertEqual(result["exact_source_label"], decoder.PROJECT_GLOBAL_AGREEMENT_LABEL)

        conflict_with_fallback, _ = _cff_inspector(
            snapshot,
            project_reading="ㄆ",
            symbol_reading="ㄇ",
        )
        result = conflict_with_fallback.decode(1)
        self.assertTrue(result["exact_conflict"])
        self.assertEqual(result["exact_reading"], "")
        self.assertEqual(result["reading"], "ㄇ")
        self.assertEqual(result["method"], "CFF整字右側注音輪廓直接解碼")

        conflict_without_fallback, _ = _cff_inspector(snapshot, project_reading="ㄆ")
        result = conflict_without_fallback.decode(1)
        self.assertTrue(result["exact_conflict"])
        self.assertEqual(result["reading"], "")

    def test_global_quarantine_is_not_donor_but_independent_decoder_survives(self):
        _probe, glyph_sha = _cff_inspector(_absent_snapshot())
        identity = _identity(
            glyph_sha,
            kind=CFF_GLYPH_SHA256,
            style_group="BIAOKAI_W5",
        )
        snapshot = _snapshot(
            quarantined=(_record(identity, conflict=("ㄅ", "ㄆ")),)
        )
        inspector, _ = _cff_inspector(snapshot, symbol_reading="ㄇ")
        result = inspector.decode(1)
        self.assertTrue(result["exact_conflict"])
        self.assertEqual(result["reading"], "ㄇ")
        self.assertEqual(result["global_effective_state"], library.QUARANTINED_CONFLICT)


class _DecodePage:
    rect = SimpleNamespace(width=600, height=800)

    def get_label(self):
        return "1"

    def get_fonts(self, full=True):
        return [(10, "ttf", "TrueType", "FixtureFont", "F1", "Identity-H", 0)]

    def get_texttrace(self):
        return [{
            "font": "FixtureFont",
            "size": 12,
            "opacity": 1,
            "seqno": 1,
            "chars": ((ord("字"), 5, (10, 10), (10, 10, 20, 20)),),
        }]

    def get_text(self, kind):
        if kind == "words":
            return []
        return {}


class _DecodeDocument:
    page_count = 1

    def __init__(self):
        self.page = _DecodePage()

    def __iter__(self):
        return iter((self.page,))

    def extract_font(self, xref):
        if xref != 10:
            raise KeyError(xref)
        return ("FixtureFont", "ttf", "TrueType", b"fixture-font-bytes")

    def close(self):
        pass


class _DecodeTrueTypeInspector:
    def __init__(self, _data):
        pass

    def components(self, glyph_id):
        if int(glyph_id) != 5:
            return []
        return [
            {"gid": 1, "x": 0, "y": 0, "flags": 0},
            {"gid": 7, "x": 1024, "y": 0, "flags": 0},
        ]

    def glyph_sha256(self, glyph_id):
        if int(glyph_id) != 7:
            raise IndexError(glyph_id)
        return SHA_A

    def global_exact_identity(self, glyph_id):
        if int(glyph_id) != 7:
            return {"eligibility": NON_GLOBAL_ELIGIBLE, "glyph_sha256": SHA_B}
        return {
            "eligibility": GLOBAL_ELIGIBLE_SIMPLE_GLYF_V1,
            "glyph_sha256": SHA_A,
            "reason": "VISIBLE_SIMPLE_GLYF",
        }

    def simple_contours(self, _glyph_id):
        return []


def _read_actual_and_audit(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        actual_sheet = workbook["實際注音"]
        headers = [cell.value for cell in next(actual_sheet.iter_rows(min_row=1, max_row=1))]
        values = [cell.value for cell in next(actual_sheet.iter_rows(min_row=2, max_row=2))]
        actual = dict(zip(headers, values))
        audit_sheet = workbook["Global exact reuse 稽核"]
        audit_headers = [cell.value for cell in next(audit_sheet.iter_rows(min_row=1, max_row=1))]
        audits = [
            dict(zip(audit_headers, (cell.value for cell in row)))
            for row in audit_sheet.iter_rows(min_row=2)
        ]
        return actual, audits
    finally:
        workbook.close()


class TTFWorkbookReadReuseIntegrationTests(unittest.TestCase):
    def _decode_fixture(
        self,
        root: Path,
        snapshot: library.GlobalExactGlyphSnapshot,
        *,
        map_reading: str = "ㄇ",
        occurrence_override: str = "",
        name: str,
    ):
        pdf = root / f"{name}.pdf"
        pdf.write_bytes(b"synthetic PDF fixture")
        output = root / f"{name}.xlsx"
        map_path = root / f"{name}-map.csv"
        map_path.write_text(
            "font,zhuyin_component_id,bopomofo,verification,notes\n"
            f"FixtureFont,7,{map_reading},fixture exact map,actual-only fixture\n",
            encoding="utf-8",
        )
        groups_path = root / f"{name}-missing-groups.csv"
        dynamic_root = root / f"{name}-dynamic"
        dynamic_root.mkdir()
        if occurrence_override:
            (dynamic_root / actual_review.OCCURRENCE_OVERRIDE_FILE).write_text(
                "pdf_contains,pdf_excludes,page,target_char,stable_key,x0,y0,actual_reading,source,note\n"
                f"{name},,1,字,,10,10,{occurrence_override},direct visual fixture,highest authority\n",
                encoding="utf-8",
            )
        with (
            patch.object(decoder.fitz, "open", return_value=_DecodeDocument()),
            patch.object(decoder, "TrueTypeGlyphInspector", _DecodeTrueTypeInspector),
        ):
            decoder.decode(
                pdf,
                output,
                map_path,
                groups_path,
                actual_asset_fingerprint={"fingerprint": "fixture", "components": {}},
                dynamic_evidence_root=dynamic_root,
                global_snapshot=snapshot,
            )
        return _read_actual_and_audit(output)

    def test_absent_preserves_existing_decoder_global_preempts_it_and_conflict_falls_back_with_audit(self):
        identity = _identity(SHA_A)
        with tempfile.TemporaryDirectory(dir=ROOT) as directory:
            root = Path(directory)
            absent_actual, absent_audit = self._decode_fixture(
                root,
                _absent_snapshot(),
                name="absent",
            )
            self.assertEqual(absent_actual["實際注音"], "ㄇ")
            self.assertEqual(absent_actual["解碼依據"], "精確字型對照")
            self.assertEqual(absent_audit, [])

            verified_actual, verified_audit = self._decode_fixture(
                root,
                _snapshot(trusted=(_record(identity, reading="ㄅ"),)),
                name="verified",
            )
            self.assertEqual(verified_actual["實際注音"], "ㄅ")
            self.assertEqual(verified_actual["解碼依據"], decoder.GLOBAL_TTF_SOURCE_LABEL)
            self.assertEqual(verified_actual["對照來源font"], decoder.GLOBAL_TTF_SOURCE_LABEL)
            self.assertEqual(verified_audit[0]["處置"], "GLOBAL_EXACT_REUSED")

            conflict_actual, conflict_audit = self._decode_fixture(
                root,
                _snapshot(quarantined=(_record(identity, conflict=("ㄅ", "ㄆ")),)),
                name="conflict",
            )
            self.assertEqual(conflict_actual["實際注音"], "ㄇ")
            self.assertEqual(conflict_actual["解碼依據"], "精確字型對照")
            self.assertEqual(
                conflict_audit[0]["處置"],
                decoder.GLOBAL_EXACT_RUNTIME_CONFLICT_LABEL,
            )
            self.assertEqual(conflict_audit[0]["已由獨立fallback解碼筆數"], 1)

    def test_occurrence_override_remains_highest_authority_during_global_conflict(self):
        identity = _identity(SHA_A)
        with tempfile.TemporaryDirectory(dir=ROOT) as directory:
            actual, audit = self._decode_fixture(
                Path(directory),
                _snapshot(quarantined=(_record(identity, conflict=("ㄅ", "ㄆ")),)),
                occurrence_override="ㄈ",
                name="override",
            )
        self.assertEqual(actual["自動解碼原值"], "ㄇ")
        self.assertEqual(actual["實際注音"], "ㄈ")
        self.assertEqual(actual["解碼依據"], "使用者原頁人工覆核（出現位置限定）")
        self.assertEqual(audit[0]["處置"], decoder.GLOBAL_EXACT_RUNTIME_CONFLICT_LABEL)
        self.assertEqual(audit[0]["已由獨立fallback解碼筆數"], 0)


class Phase3ArchitectureTests(unittest.TestCase):
    def test_actual_fingerprint_and_source_roster_contract(self):
        self.assertEqual(runtime_source_validation.ACTUAL_FINGERPRINT_SCHEMA_VERSION, "3.0.0")
        self.assertIn(
            "global_exact_glyph_evidence_hashes",
            cross_version_compat.FINGERPRINT_COMPATIBILITY_REQUIRED_KEYS["actual"],
        )
        self.assertEqual(
            standalone_proofread.ACTUAL_DECODER_SOURCE_FILES,
            decoder.ACTUAL_DECODER_SOURCE_FILES,
        )
        self.assertIn("global_exact_glyph_library.py", decoder.ACTUAL_DECODER_SOURCE_FILES)
        self.assertNotIn(
            "global_exact_glyph_library.py",
            candidates.EXPECTED_RESOLVER_SOURCE_FILES,
        )

    def test_read_paths_do_not_connect_global_mutation_promotion_outbox_or_migration(self):
        production_read_paths = "\n".join(
            inspect.getsource(function)
            for function in (
                standalone_proofread.run_pipeline_pdfs,
                standalone_proofread.regenerate_report,
                decoder.decode,
                candidates.analyze,
                actual_review.actual_workbook_global_exact_dependencies,
                cff_decoder.CFFZhuyinInspector.decode,
            )
        )
        for forbidden in (
            ".initialize(",
            ".write_transaction(",
            ".execute_idempotent_write(",
            ".compare_and_swap_glyph_revision(",
            "project_outbox",
            "promotion_approval",
            "migration_candidate",
            "INSERT INTO",
            "DELETE FROM",
        ):
            with self.subTest(forbidden=forbidden):
                self.assertNotIn(forbidden, production_read_paths)
        self.assertNotIn("library.sqlite3", (ROOT / "runtime_asset_manifest.json").read_text(encoding="utf-8"))

    def test_epochs_and_expected_fingerprint_contract_remain_unchanged(self):
        self.assertEqual(cross_version_compat.ACTUAL_DECODER_SEMANTICS_EPOCH, "1")
        self.assertEqual(cross_version_compat.EXPECTED_RESOLVER_SEMANTICS_EPOCH, "1")
        self.assertEqual(runtime_source_validation.EXPECTED_FINGERPRINT_SCHEMA_VERSION, "2.7.0")
        self.assertEqual(
            cross_version_compat.FINGERPRINT_COMPATIBILITY_REQUIRED_KEYS["expected"],
            (
                "fingerprint_schema_version",
                "reuse_policy",
                "expected_resolver_semantics_epoch",
                "expected_asset_hashes",
            ),
        )

    def test_absent_read_creates_neither_directory_nor_database(self):
        with tempfile.TemporaryDirectory(dir=ROOT) as directory:
            global_root = Path(directory) / "not-created"
            snapshot = library.GlobalExactGlyphRepository.resolved(global_root).load_snapshot()
            self.assertEqual(snapshot.store_status, library.ABSENT)
            self.assertFalse(global_root.exists())

    def test_present_invalid_store_blocks_before_project_or_decode_mutation(self):
        with tempfile.TemporaryDirectory(dir=ROOT) as directory:
            root = Path(directory)
            output = root / "output"
            pdf = root / "book.pdf"
            pdf.write_bytes(b"fixture")
            repository = SimpleNamespace()
            repository.load_snapshot = MagicMock(
                side_effect=library.GlobalLibraryCorruptError("fixture corruption")
            )
            with (
                patch.object(standalone_proofread, "validate_asset_manifest", return_value={"ok": True, "errors": []}),
                patch.object(standalone_proofread.GlobalExactGlyphRepository, "resolved", return_value=repository),
                patch.object(standalone_proofread, "write_pipeline_blocked", return_value=output / "pipeline_blocked.json") as blocked,
                patch.object(standalone_proofread, "initialize_project_actual_evidence") as initialize_project,
                patch.object(standalone_proofread, "decode") as decode_call,
            ):
                with self.assertRaises(runtime_source_validation.SourceValidationError):
                    standalone_proofread.run_pipeline_pdfs([pdf], output)
            repository.load_snapshot.assert_called_once_with()
            blocked.assert_called_once()
            self.assertEqual(blocked.call_args.args[2], "GLOBAL_EXACT_SOURCE_INVALID")
            initialize_project.assert_not_called()
            decode_call.assert_not_called()


if __name__ == "__main__":
    unittest.main()
