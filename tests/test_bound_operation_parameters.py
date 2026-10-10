"""Explicit project-parameter binding and real tiny export; no writer shim/Tk."""
import inspect
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

import pdf_portability as p
import sealed_workbook_paths as locations
import standalone_proofread as sp
from tests.test_pdf_portability import pdf, project


@pytest.fixture
def portable_project(tmp_path, monkeypatch):
    monkeypatch.setattr(sp, "USER_MEMORY", tmp_path / "memory.json")
    monkeypatch.setattr(sp, "REUSABLE_EXPECTED_RULES", tmp_path / "rules.json")
    source_pdf = tmp_path / "tiny.pdf"
    pdf(source_pdf, title="tiny bound-operation export fixture")
    root = tmp_path / "selected project"
    manifest, db = project(root, source_pdf, session="tiny-real-export")
    # Both sealed role files are real XLSX inputs, not opaque byte placeholders.
    candidate = Path(manifest["pdfs"][0]["candidate_workbook"])
    workbook = Workbook()
    workbook.active.title = "候選報告"
    workbook.active.append(["occurrence_id", "actual", "expected"])
    for entry in manifest["records"]:
        workbook.active.append([entry["occurrence_id"], entry["actual"], ""])
    workbook.save(candidate)
    workbook.close()
    manifest["pdfs"][0]["candidate_workbook_sha256"] = sp.sha256_file(candidate)
    sp.seal_manifest(manifest)
    sp.json_save(root / "校對工作階段.json", manifest)
    return root, manifest, db


@pytest.mark.parametrize("mode", ["positional", "keywords"])
def test_real_export_with_content_proof_is_readable(portable_project, mode):
    root, manifest, _ = portable_project
    before = {Path(item[key]): Path(item[key]).read_bytes() for item in manifest["pdfs"]
              for key in ("pdf", "actual_workbook", "candidate_workbook")}
    exported = sp.export_pending_for_gpt(root) if mode == "positional" else sp.export_pending_for_gpt(output_dir=root)
    with_export = p.load_excel_content_proof(exported, root, manifest, kind="expected",
                                           workbook_session=manifest["session_id"])
    assert with_export[0] == manifest
    assert len(with_export[3]) == len(manifest["records"])
    assert with_export[6]["kind"] == "expected"
    workbook = load_workbook(exported, read_only=True)
    try:
        assert workbook[p.EXCEL_PROOF_SHEET].sheet_state == "veryHidden"
        assert workbook["待判定候選"].max_row == len(manifest["records"]) + 1
    finally:
        workbook.close()
    assert all(path.read_bytes() == raw for path, raw in before.items())


@pytest.mark.parametrize("mode", ["positional", "keywords"])
def test_workbook_first_real_writer_binds_output_dir(portable_project, mode):
    root, manifest, db = portable_project
    exported = sp.export_pending_for_gpt(root)
    workbook = load_workbook(exported)
    workbook.remove(workbook[p.EXCEL_PROOF_SHEET])
    if mode == "positional":
        p.write_excel_content_proof(workbook, root, manifest, db, kind="expected")
    else:
        p.write_excel_content_proof(workbook=workbook, output_dir=root, manifest=manifest, db=db, kind="expected")
    path = root / (mode + ".xlsx")
    workbook.save(path)
    workbook.close()
    decoded = p.load_excel_content_proof(path, root, manifest, kind="expected", workbook_session=manifest["session_id"])
    assert decoded[0] == manifest and decoded[1] == db
    assert len(decoded[3]) == len(manifest["records"])


def test_bare_root_first_default_and_wrapping_contract(tmp_path):
    @locations.bound_operation
    def operation(output_dir=tmp_path, *, value=3):
        assert locations._bindings.get() is not None
        return output_dir, value
    assert operation(value=4) == (tmp_path, 4)
    assert operation(tmp_path, value=5) == (tmp_path, 5)
    assert operation(output_dir=tmp_path) == (tmp_path, 3)
    assert operation.__name__ == "operation"
    assert inspect.signature(operation) == inspect.signature(operation.__wrapped__)
    with locations.binding_scope():
        assert operation.__wrapped__(value=3) == (tmp_path, 3)
    assert locations._bindings.get() is None


def test_explicit_second_keyword_only_and_positional_only_roots(tmp_path):
    @locations.bound_operation(project_parameter="target_dir")
    def second(payload, target_dir):
        return payload, target_dir
    @locations.bound_operation(project_parameter="selected_root")
    def keyword_only(payload, *, selected_root):
        return selected_root
    @locations.bound_operation(project_parameter="output_dir")
    def positional_only(output_dir, /):
        return output_dir
    value = object()
    assert second(value, tmp_path) == (value, tmp_path)
    assert second(payload=value, target_dir=tmp_path) == (value, tmp_path)
    assert keyword_only(value, selected_root=tmp_path) == tmp_path
    assert positional_only(tmp_path) == tmp_path


@pytest.mark.parametrize("fault", ["missing", "duplicate", "unexpected"])
def test_signature_bind_rejects_wrong_arguments_before_entry(tmp_path, fault):
    entered = []
    @locations.bound_operation(project_parameter="target_dir")
    def operation(payload, target_dir):
        entered.append(target_dir)
    with pytest.raises(TypeError):
        if fault == "missing":
            operation(object())
        elif fault == "duplicate":
            operation(object(), tmp_path, target_dir=tmp_path)
        else:
            operation(payload=object(), target_dir=tmp_path, unknown=True)
    assert entered == [] and locations._bindings.get() is None


@pytest.mark.parametrize("name", ["absent", "args", "kwargs"])
def test_invalid_parameter_configuration_is_rejected_when_decorated(name):
    with pytest.raises(ValueError, match="project parameter"):
        @locations.bound_operation(project_parameter=name)
        def operation(output_dir, *args, **kwargs):
            pytest.fail("invalid decorator admitted entry")
    with pytest.raises(ValueError, match="project parameter"):
        @locations.bound_operation
        def unknown_bare_root(target_dir):
            pytest.fail("bare form guessed target_dir")


def test_invalid_root_is_rejected_for_keyword_writer_without_skipping_trust():
    workbook = Workbook()
    try:
        with pytest.raises(TypeError):
            p.write_excel_content_proof(workbook=workbook, output_dir=object(), manifest={}, db={}, kind="expected")
        assert p.EXCEL_PROOF_SHEET not in workbook.sheetnames
        assert locations._bindings.get() is None
    finally:
        workbook.close()


def test_nested_scope_and_exception_cleanup(tmp_path):
    seen = []
    @locations.bound_operation(project_parameter="target_dir")
    def inner(value, target_dir):
        seen.append(locations._bindings.get())
        raise RuntimeError("entry failed")
    @locations.bound_operation(project_parameter="output_dir")
    def outer(output_dir):
        seen.append(locations._bindings.get())
        with pytest.raises(RuntimeError, match="entry failed"):
            inner(object(), target_dir=output_dir)
        assert locations._bindings.get() is seen[0]
    outer(tmp_path)
    assert seen[0] is seen[1] and locations._bindings.get() is None
    with locations.binding_scope():
        caller_scope = locations._bindings.get()
        with pytest.raises(RuntimeError):
            inner(object(), tmp_path)
        assert locations._bindings.get() is caller_scope
    assert locations._bindings.get() is None
