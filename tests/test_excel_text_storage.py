"""Isolated XLSX round trips, not manual Excel or real textbook acceptance."""
import copy
import io
import json
from pathlib import Path
from unittest.mock import patch

import fitz
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from fontTools.fontBuilder import FontBuilder
from fontTools.pens.ttGlyphPen import TTGlyphPen

import actual_review as ar
import check_pronunciation_candidates as candidates
import export_zhuyin_readings as decoder
import occurrence_ledger as ol
import pdf_portability as portable
import standalone_proofread as sp
from excel_text import append_text_row, write_text_value
from tests.test_global_exact_glyph_read_reuse_v580 import (
    _DecodeDocument, _DecodeTrueTypeInspector, _absent_snapshot,
)
from export_pdf_text_diagnostics import TrueTypeGlyphInspector, build_font_lookup, normalize_basefont, resolve_font
from exact_glyph_identity import classify_ttf_glyf_record, NON_GLOBAL_ELIGIBLE
from actual_source_visual_safety import source_visual_safety, source_visual_operation
from tests.test_project_repair_integration_v562 import ACTUAL_COLUMNS, EXCLUDED_COLUMNS
from tests.test_import_transaction import make_manifest
from tests.test_actual_review_v540 import entry, make_pdf


# Literal, independent oracles; production has no symbol/prefix list.
TEXTS = (
    "=15,15個一換成1個十和5", "=6,表示6個一。", "=1+2", "+1", "-1", "@1", "=",
    "001", "1/2", "1E3", "#N/A", "＝15", "×", "÷", "−", "≤", "≥",
    "中文×3÷2−1≤5≥0", " 前後空白 ", "第一行\n=第二行", "'原本的單引號", "\t空白\r\n末行 ",
)
ROOT = Path(__file__).resolve().parents[1]


def assert_text(path, sheet_name, field, expected, *, identity_field=None):
    for data_only in (False, True):
        wb = load_workbook(path, data_only=data_only)
        try:
            sheet = wb[sheet_name]
            columns = {cell.value: cell.column for cell in sheet[1]}
            cells = [sheet.cell(row, columns[field]) for row in range(2, sheet.max_row + 1)]
            if identity_field is None:
                assert [cell.value for cell in cells] == list(expected)
            else:
                identities = [sheet.cell(row, columns[identity_field]).value
                              for row in range(2, sheet.max_row + 1)]
                assert len(identities) == len(set(identities)) == len(expected)
                assert dict(zip(identities, (cell.value for cell in cells))) == expected
            assert all(cell.data_type == "s" for cell in cells)
        finally:
            wb.close()


def project_bytes(root):
    # Byte-level atomicity: include all project artifacts and dynamic evidence.
    return {str(p.relative_to(root)): p.read_bytes() for p in root.rglob("*")
            if p.is_file() and p.suffix != ".lock"}


def literal_source_pdf(path, count):
    """Real distinct project-only programs, created before source/ID sealing."""
    with fitz.open() as document:
        page = document.new_page(width=300, height=max(200, 80 + count * 32))
        for index in range(count):
            names = [".notdef", "base", "whole"]
            pen = TTGlyphPen(None)
            pen.moveTo((0, 0)); pen.lineTo((200 + index * 8, 0))
            pen.lineTo((200 + index * 8, 350 + index * 6))
            pen.lineTo((0, 350 + index * 6)); pen.closePath()
            glyphs = {".notdef": TTGlyphPen(None).glyph(), "base": pen.glyph()}
            composite = TTGlyphPen(glyphs)
            composite.addComponent("base", (1, 0, 0, 1, 0, 0))
            glyphs["whole"] = composite.glyph()
            builder = FontBuilder(1000, isTTF=True)
            builder.setupGlyphOrder(names); builder.setupCharacterMap({ord("角"): "whole"})
            builder.setupGlyf(glyphs)
            builder.setupHorizontalMetrics({name: (600, 0) for name in names})
            builder.setupHorizontalHeader(ascent=800, descent=-200)
            builder.setupNameTable({"familyName": f"LiteralSource{index}", "styleName": "Regular"})
            builder.setupOS2(sTypoAscender=800, sTypoDescender=-200, usWinAscent=800, usWinDescent=200)
            builder.setupPost()
            stream = io.BytesIO(); builder.save(stream)
            name = f"Literal{index}"
            page.insert_font(fontname=name, fontbuffer=stream.getvalue())
            page.insert_text((50, 56 + index * 32), "角", fontname=name, fontsize=20)
        document.save(path)


def literal_native_sources(path):
    records = []
    with fitz.open(path) as document:
        page = document[0]
        lookup = build_font_lookup(page.get_fonts(full=True))
        for span in page.get_texttrace():
            xref, *_ = resolve_font(lookup, normalize_basefont(span["font"]))
            assert type(xref) is int
            inspector = TrueTypeGlyphInspector(document.extract_font(xref)[3])
            for char in span["chars"]:
                assert char[0] == ord("角")
                identity = classify_ttf_glyf_record(inspector.glyph_bytes(char[1]))
                assert identity["reason"] == "COMPOSITE_GLYF"
                assert identity["eligibility"] == NON_GLOBAL_ELIGIBLE
                records.append({"字元": chr(char[0]), "font": span["font"], "font_xref": xref,
                                "glyph_id_字形索引": char[1], "注音元件ID": char[1],
                                "TTF字形SHA256": inspector.glyph_sha256(char[1]),
                                **dict(zip(("x0", "y0", "x1", "y1"), char[3]))})
    assert len({row["TTF字形SHA256"] for row in records}) == len(records)
    boxes = [fitz.Rect(*(row[key] for key in ("x0", "y0", "x1", "y1"))) for row in records]
    assert all(not first.intersects(second) for index, first in enumerate(boxes) for second in boxes[index+1:])
    return records


@pytest.fixture
def isolated_pipeline(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "localappdata"))
    source_snapshots = {}

    def controlled_decode(pdf, output, *args, global_snapshot=None):
        # Only glyph acquisition is synthetic. The real candidate analyzer,
        # seals, both exporters/importers, actual transaction and refresh run.
        sealed_sha = sp.sha256_file(pdf)
        snapshot = source_snapshots[sealed_sha]
        texts, native_sources, roster = (snapshot[key] for key in ("texts", "native_sources", "roster"))
        unresolved = snapshot["unresolved"]
        assert literal_native_sources(pdf) == native_sources
        wb = Workbook(); sheet = wb.active; sheet.title = "實際注音"
        columns = [*ACTUAL_COLUMNS, "注音結構偵測來源", "CFF符號簽名", "CFF聲調"]
        append_text_row(sheet, columns)
        overrides = ar._read_csv(args[13] / ar.OCCURRENCE_OVERRIDE_FILE, ar.OVERRIDE_HEADERS)
        records = []
        for base, text in zip(roster, texts):
            stable = base["穩定注音鍵"]
            matched = next((item for item in overrides if item["stable_key"] == stable), None)
            rec = {**base,
                   "實際注音": matched["actual_reading"] if matched else ("" if unresolved else "ㄐㄩㄝˇ"),
                   "解碼狀態": "已映射" if matched or not unresolved else "未辨識",
                   "解碼依據": text}
            records.append(rec)
        for rec in records:
            append_text_row(sheet, [rec.get(key, "") for key in columns])
        excluded = wb.create_sheet("結構偵測排除"); append_text_row(excluded, EXCLUDED_COLUMNS)
        metadata = wb.create_sheet("v5.2中繼資料")
        fingerprint = args[-2]
        # This fixture replaces fresh decoding, so it must emit the producer's
        # epoch receipt before the real dependency-sealing guard may accept it.
        for pair in [("項目", "內容"), ("workbook_schema_version", ol.WORKBOOK_SCHEMA_VERSION),
                     ("ledger_schema_version", ol.LEDGER_SCHEMA_VERSION), ("pdf_sha256", sealed_sha),
                     ("actual_decoder_semantics_epoch", fingerprint["components"]["actual_decoder_semantics_epoch"]),
                     ("actual_asset_fingerprint", fingerprint["fingerprint"]),
                     ("actual_asset_fingerprint_components", json.dumps(fingerprint["components"], ensure_ascii=False, sort_keys=True))]:
            append_text_row(metadata, pair)
        summary = wb.create_sheet("摘要")
        append_text_row(summary, ["項目", "內容"])
        append_text_row(summary, ["偵測到注音字形筆數", len(texts)])
        wb.save(output)
        wb.close()

    def context(rec, _index):
        snapshot = source_snapshots[rec["pdf_sha256"]]
        index = int(str(rec["穩定注音鍵"]).rsplit("-", 1)[1])
        assert rec["occurrence_id"] == snapshot["roster"][index]["occurrence_id"]
        text = snapshot["texts"][index]
        return text, text, "", "", "", None, text

    monkeypatch.setattr(sp, "decode", controlled_decode)
    monkeypatch.setattr(sp, "actual_workbook_global_exact_dependencies", lambda *_: ())
    monkeypatch.setattr(candidates, "actual_workbook_global_exact_dependencies", lambda *_: ())
    monkeypatch.setattr(candidates, "locate_line_context", context)

    def create(lane):
        # Callers can select their literal roster after fixture setup, before create.
        texts = tuple(TEXTS)
        source_dir = tmp_path / f"literal-source-{len(source_snapshots)}"
        source_dir.mkdir()
        pdf = source_dir / "literal-sources.pdf"
        literal_source_pdf(pdf, len(texts))
        sealed_sha = sp.sha256_file(pdf)
        native_sources = literal_native_sources(pdf)
        assert len(native_sources) == len(texts)
        roster = [{**native, "實體頁碼": 1, "課本頁": text, "頁面標籤": text,
                   "穩定注音鍵": f"text-fixture-{index}", "群組注音鍵": f"text-fixture-{index}",
                   "source_row_number": index + 1, "ledger_schema_version": ol.LEDGER_SCHEMA_VERSION,
                   "workbook_schema_version": ol.WORKBOOK_SCHEMA_VERSION}
                  for index, (native, text) in enumerate(zip(native_sources, texts))]
        ol.prepare_occurrence_rows(sealed_sha, roster)
        with source_visual_operation():
            for rec in roster:
                proof = source_visual_safety({
                    "pdf": str(pdf), "pdf_name": pdf.name, "pdf_sha256": sealed_sha, "physical_page": 1,
                    "char": rec["字元"], "font_xref": rec["font_xref"],
                    "glyph_id": rec["glyph_id_字形索引"], "zhuyin_component_id": rec["注音元件ID"],
                    "occurrence_id": rec["occurrence_id"], "review_id": rec["review_id"],
                    "source_record": rec, **{key: rec[key] for key in ("x0", "y0", "x1", "y1")}})
                assert proof["location_readable"] is True and proof["donor_eligible"] is True

        # Refresh is bound to this immutable source/roster, not later TEXTS edits.
        assert sealed_sha not in source_snapshots
        source_snapshots[sealed_sha] = {"texts": texts, "native_sources": native_sources,
                                        "roster": roster, "unresolved": lane == "actual"}
        root = tmp_path / lane
        sp.run_pipeline_pdfs([pdf], root, defer_excel_reports=True)
        return root
    yield create
    assert not list((tmp_path / "localappdata").rglob("*.sqlite*"))


def exported(root, lane):
    if lane == "expected":
        return sp.export_pending_for_gpt(root), "待判定候選"
    assert sp.export_actual_pending_for_gpt(root) is not None
    return root / "actual待判定_GPT包" / "actual待判定_給GPT.xlsx", "actual待判定"


def fill(path, sheet_name, lane):
    wb = load_workbook(path)
    sheet = wb[sheet_name]
    cols = {cell.value: cell.column for cell in sheet[1]}
    answers = ({"action": "解決expected證據", "proposed_expected_set": "ㄐㄩㄝˊ",
                "proposed_expected_evidence": "獨立的正式應標證據",
                "proposed_context_evidence": "獨立的本列語境"} if lane == "expected" else
               {"decision": "VERIFIED", "actual_reading": "ㄐㄩㄝˊ", "confidence": "高",
                "sample_a_checked": "Y", "sample_b_checked": "Y"})
    for row in range(2, sheet.max_row + 1):
        for key, value in answers.items():
            sheet.cell(row, cols[key], value)
    return wb, sheet, cols


@pytest.mark.parametrize("lane", ["expected", "actual"])
def test_formal_export_save_reopen_fill_import_preserves_literal_text(isolated_pipeline, lane):
    root = isolated_pipeline(lane)
    manifest = sp.json_load_strict(root / "校對工作階段.json")
    before = copy.deepcopy(manifest["records"])
    candidate = Path(manifest["pdfs"][0]["candidate_workbook"])
    assert_text(candidate, "Occurrence Ledger", "printed_page", TEXTS)
    for data_only in (False, True):
        wb = load_workbook(candidate, data_only=data_only)
        try:
            ws = wb["Occurrence Ledger"]
            cols = {cell.value: cell.column for cell in ws[1]}
            sources = [json.loads(ws.cell(row, cols["source_record_json"]).value)
                       for row in range(2, ws.max_row + 1)]
            assert [source["所在行"] for source in sources] == list(TEXTS)
        finally:
            wb.close()
    path, sheet_name = exported(root, lane)
    fields = ("課本頁", "所在行", "局部詞境", "actual_evidence") if lane == "expected" else ("sample_a_page", "actual_evidence")
    # Actual groups have their own deterministic order. Bind literal oracles to
    # original occurrence IDs rather than imposing expected-lane row ordering.
    text_oracle = {entry["occurrence_id"]: text for entry, text in zip(before, TEXTS)}
    kwargs = {} if lane == "expected" else {"identity_field": "sample_a_occurrence_id"}
    for field in fields:
        assert_text(path, sheet_name, field, TEXTS if lane == "expected" else text_oracle, **kwargs)
    wb = load_workbook(path)
    try:
        assert portable.EXCEL_PROOF_SHEET in wb.sheetnames
        ws = wb[sheet_name]
        assert len(ws.data_validations.dataValidation) > 0
        cols = {cell.value: cell.column for cell in ws[1]}
        assert ws.cell(2, cols["編號"]).value == 1 and ws.cell(2, cols["編號"]).data_type == "n"
        assert ws.cell(2, cols["action" if lane == "expected" else "decision"]).value is None
        if lane == "expected":
            assert ws.cell(2, cols["實體頁碼"]).value == 1
    finally:
        wb.close()
    wb, _, _ = fill(path, sheet_name, lane)
    wb.save(path)
    wb.close()
    for field in fields:
        assert_text(path, sheet_name, field, TEXTS if lane == "expected" else text_oracle, **kwargs)
    importer = sp.import_gpt_decisions if lane == "expected" else sp.import_actual_gpt_decisions
    assert importer(root, path)[0] == len(TEXTS)
    db = sp.json_load_strict(root / "人工判定資料庫.json")
    latest = sp.json_load_strict(root / "校對工作階段.json")
    ledger = sp.materialize_ledger(latest, db)
    if lane == "expected":
        assert [entry["actual"] for entry in ledger] == [entry["actual"] for entry in before]
        assert [entry["actual_evidence"] for entry in ledger] == [entry["actual_evidence"] for entry in before]
    else:
        assert [entry["expected_set"] for entry in ledger] == [entry["expected_set"] for entry in before]
        assert [entry["expected_evidence"] for entry in ledger] == [entry["expected_evidence"] for entry in before]
        assert [entry["actual"] for entry in ledger] == ["ㄐㄩㄝˊ"] * len(TEXTS)


@pytest.mark.parametrize("lane", ["expected", "actual"])
@pytest.mark.parametrize("tamper", ["source", "identity", "seal"])
def test_protected_tamper_rejects_entire_batch_without_project_writes(isolated_pipeline, lane, tamper):
    root = isolated_pipeline(lane)
    path, sheet_name = exported(root, lane)
    wb, sheet, cols = fill(path, sheet_name, lane)
    if tamper == "seal":
        wb[portable.EXCEL_PROOF_SHEET].cell(2, 2, "0" * 64)
    else:
        field = ("所在行" if lane == "expected" else "sample_a_page") if tamper == "source" else ("review_id" if lane == "expected" else "sample_a_occurrence_id")
        write_text_value(sheet, sheet.max_row, cols[field], "modified protected source")
    wb.save(path)
    wb.close()
    before = project_bytes(root)
    importer = sp.import_gpt_decisions if lane == "expected" else sp.import_actual_gpt_decisions
    with pytest.raises(ValueError):
        importer(root, path)
    assert project_bytes(root) == before


def test_real_first_actual_writer_preserves_each_label_and_evidence(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "localappdata"))
    source = tmp_path / "source.pdf"
    with fitz.open() as real:
        real.new_page()
        real.save(source)
    document = _DecodeDocument()
    pages = []
    for text in TEXTS:
        page = copy.copy(document.page)
        page.literal = text
        page.get_label = lambda value=text: value
        pages.append(page)
    document.page_count = len(pages)
    # A local acquisition adapter; production decode and worksheet writes run.
    class Document:
        page_count = len(pages)
        def __iter__(self): return iter(pages)
        def extract_font(self, xref): return document.extract_font(xref)
        def close(self): pass
    output = tmp_path / "actual.xlsx"
    with patch.object(decoder.fitz, "open", return_value=Document()), \
         patch.object(decoder, "TrueTypeGlyphInspector", _DecodeTrueTypeInspector), \
         patch.object(decoder, "infer_printed_page_number", side_effect=lambda page, _: page.literal):
        decoder.decode(source, output, tmp_path / "absent-map.csv", tmp_path / "absent-groups.csv",
                       dynamic_evidence_root=tmp_path / "dynamic", global_snapshot=_absent_snapshot())
    assert_text(output, "實際注音", "課本頁", TEXTS)
    assert_text(output, "實際注音", "頁面標籤", TEXTS)
    wb = load_workbook(output, data_only=True)
    try:
        assert all(type(row[0]) is int for row in wb["實際注音"].iter_rows(min_row=2, values_only=True))
    finally:
        wb.close()


@pytest.mark.parametrize("bad,reason", [("x" * 32768, "32767"), ("😀" * 16384, "UTF-16"),
                                       ("bad\x00text", "U+0000"), ("\ud800", "U+D800"), ("\uffff", "U+FFFF")],
                         ids=["too-long-ascii", "too-long-utf16", "nul", "surrogate", "noncharacter"])
def test_unrepresentable_text_names_sheet_field_and_fails_before_mutation(bad, reason):
    wb = Workbook()
    ws = wb.active
    ws.title = "教材"
    ws.append(["原文"])
    with pytest.raises(ValueError, match=reason.replace("+", r"\+")) as exc:
        append_text_row(ws, [bad])
    assert "教材/原文 (A2)" in str(exc.value)
    assert ws.max_row == 1
    write_text_value(ws, 2, 1, "original")
    with pytest.raises(ValueError):
        write_text_value(ws, 2, 1, bad)
    assert ws.cell(2, 1).value == "original"


def test_native_scalars_required_formula_and_validation_are_preserved(tmp_path):
    wb = Workbook()
    ws = wb.active
    append_text_row(ws, ["text", "integer", "float", "boolean", "empty", "formula"])
    append_text_row(ws, ["#N/A", 15, 1.25, True, None, None])
    ws.cell(2, 6, "=B2+C2")  # An actual formula writer is deliberately separate.
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("D2")
    path = tmp_path / "native.xlsx"
    wb.save(path)
    loaded = load_workbook(path)
    try:
        assert [cell.value for cell in loaded.active[2]] == ["#N/A", 15, 1.25, True, None, "=B2+C2"]
        assert [cell.data_type for cell in loaded.active[2]][:4] == ["s", "n", "n", "b"]
        assert loaded.active.cell(2, 6).data_type == "f"
        assert len(loaded.active.data_validations.dataValidation) == 1
    finally:
        loaded.close()


def test_exact_capacity_boundary_round_trip(tmp_path):
    wb = Workbook()
    ws = wb.active
    append_text_row(ws, ["原文"])
    append_text_row(ws, ["a" * 32767])
    append_text_row(ws, ["😀" * 16383 + "a"])
    path = tmp_path / "boundary.xlsx"
    wb.save(path)
    assert_text(path, ws.title, "原文", ["a" * 32767, "😀" * 16383 + "a"])


@pytest.mark.parametrize("lane", ["expected", "actual"])
def test_formal_export_limit_refusal_preserves_previous_output(tmp_path, monkeypatch, lane):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "localappdata"))
    root = tmp_path / "project"
    root.mkdir()
    if lane == "expected":
        # A new, isolated source fixture: sealed once with its declared input,
        # never an importer adapter that reseals a modified carried workbook.
        source = make_manifest()
        source["records"][0]["source_record"]["所在行"] = "x" * 32768
        sp.json_save(root / "校對工作階段.json", sp.seal_manifest(source))
        sp.json_save(root / "人工判定資料庫.json", sp.normalize_db({}))
        path = root / "待判定候選_給GPT.xlsx"
        path.write_bytes(b"previous publication")
        before = path.read_bytes()
        with pytest.raises(ValueError, match="待判定候選/所在行.*32767"):
            sp.export_pending_for_gpt(root)
        assert path.read_bytes() == before
    else:
        pdf = tmp_path / "source.pdf"
        make_pdf(pdf)
        row = entry("occ_" + "a" * 64, "rev_" + "b" * 64, pdf, 60)
        meta = dict(version=sp.VERSION, session_id="isolated",
                    session_schema_version=ol.SESSION_SCHEMA_VERSION,
                    workbook_schema_version=ol.WORKBOOK_SCHEMA_VERSION,
                    review_id_schema_version=ol.REVIEW_ID_SCHEMA_VERSION)
        ar.export_actual_review_package(root, [row], **meta)
        before = project_bytes(root)
        row["printed_page"] = "x" * 32768
        with pytest.raises(ValueError, match="actual待判定/sample_a_page.*32767"):
            ar.export_actual_review_package(root, [row], **meta)
        assert project_bytes(root) == before
