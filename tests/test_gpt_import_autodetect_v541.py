from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
import sys

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from standalone_proofread import detect_gpt_workbook_kind, plan_gpt_auto_imports  # noqa: E402
from standalone_gui import project_status_text  # noqa: E402


class GptImportAutodetectV541Tests(unittest.TestCase):
    def _book(self, names):
        td = tempfile.TemporaryDirectory()
        path = Path(td.name) / "x.xlsx"
        wb = Workbook()
        wb.remove(wb.active)
        for name in names:
            wb.create_sheet(name)
        wb.save(path)
        return td, path

    def test_expected_workbook_detected_from_sheet_structure(self):
        td, path = self._book(["使用說明", "匯入中繼資料", "待判定候選"])
        try:
            self.assertEqual(detect_gpt_workbook_kind(path), "expected")
        finally:
            td.cleanup()

    def test_actual_workbook_detected_from_sheet_structure(self):
        td, path = self._book(["使用說明", "匯入中繼資料", "actual待判定"])
        try:
            self.assertEqual(detect_gpt_workbook_kind(path), "actual")
        finally:
            td.cleanup()

    def test_mixed_expected_actual_workbook_is_fail_closed(self):
        td, path = self._book(["匯入中繼資料", "待判定候選", "actual待判定"])
        try:
            with self.assertRaises(ValueError):
                detect_gpt_workbook_kind(path)
        finally:
            td.cleanup()

    def test_unknown_workbook_is_rejected(self):
        td, path = self._book(["Sheet1"])
        try:
            with self.assertRaises(ValueError):
                detect_gpt_workbook_kind(path)
        finally:
            td.cleanup()

    def test_multi_select_plan_orders_actual_before_expected(self):
        td = tempfile.TemporaryDirectory()
        try:
            folder = Path(td.name)
            expected = folder / "expected.xlsx"
            actual = folder / "actual.xlsx"
            wb = Workbook(); wb.remove(wb.active)
            for name in ["使用說明", "匯入中繼資料", "待判定候選"]:
                wb.create_sheet(name)
            wb.save(expected)
            wb = Workbook(); wb.remove(wb.active)
            for name in ["使用說明", "匯入中繼資料", "actual待判定"]:
                wb.create_sheet(name)
            wb.save(actual)
            planned = plan_gpt_auto_imports([expected, actual])
            self.assertEqual([kind for _path, kind in planned], ["actual", "expected"])
            self.assertEqual(planned[0][0], actual)
        finally:
            td.cleanup()

    def test_project_status_reads_new_pending_count_after_status_rewrite(self):
        import json
        with tempfile.TemporaryDirectory() as td:
            folder = Path(td)
            status = folder / "pipeline_status.json"
            status.write_text(json.dumps({
                "status": "PROCESSING_FINISHED",
                "status_text": "處理結束，校對尚未完成。",
                "completion_gate": {"state_counts": {"PASS": 1, "EXPECTED_AMBIGUOUS": 305}},
            }, ensure_ascii=False), encoding="utf-8")
            self.assertIn("待處理 305 筆", project_status_text(folder))
            status.write_text(json.dumps({
                "status": "PROCESSING_FINISHED",
                "status_text": "處理結束，校對尚未完成。",
                "completion_gate": {"state_counts": {"PASS": 109, "EXPECTED_AMBIGUOUS": 197}},
            }, ensure_ascii=False), encoding="utf-8")
            self.assertIn("待處理 197 筆", project_status_text(folder))


if __name__ == "__main__":
    unittest.main()
