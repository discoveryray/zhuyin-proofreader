"""Portability publication and candidate-text safety regressions."""

import hashlib
import copy
import os
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

import fitz
import pytest
from openpyxl import load_workbook

import pdf_portability as portable
import standalone_proofread as sp
from tests.test_pdf_portability import pdf, project
from tests.test_pdf_portability_integration import _pdfs, _synthetic_decode
from review_save_service import ReviewSaveService


def read(path):
    return path.read_bytes() if path.exists() else None


def real_project(root):
    root.mkdir()
    first, _ = _pdfs(root)
    output = root / "project"
    with patch.object(sp, "decode", side_effect=_synthetic_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], output)
    return output


def expected_workbook(root, output):
    workbook = load_workbook(sp.export_pending_for_gpt(output))
    sheet = workbook["待判定候選"]
    headers = [cell.value for cell in sheet[1]]
    for key, value in {"action": "確認非校對範圍", "exclusion_reason": "local visual",
                       "exclusion_evidence": "PDF page 1"}.items():
        sheet.cell(2, headers.index(key) + 1, value)
    filled = root / "filled.xlsx"
    workbook.save(filled)
    workbook.close()
    return filled


def test_same_session_expected_rechecks_foreign_marker_before_writes(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    output = real_project(tmp_path / "a")
    filled = expected_workbook(tmp_path, output)
    marker = output / portable.INCOMPLETE_FILE
    foreign = {"status": "PRESENTATION_PENDING", "owner": "other transaction"}
    watched = [output / name for name in ("人工判定資料庫.json", "待人工確認.json",
                                          "pipeline_status.json", "注音校對_最終報告.xlsx")]
    before = {path: read(path) for path in watched}
    original = sp.materialize_ledger
    injected = False

    def inject(*args, **kwargs):
        nonlocal injected
        if not injected:
            injected = True
            sp.json_save(marker, foreign)
        return original(*args, **kwargs)

    with patch.object(sp, "materialize_ledger", side_effect=inject):
        with pytest.raises(ValueError, match="未完成|標記"):
            sp.import_gpt_decisions(output, filled)
    assert injected and sp.json_load_strict(marker) == foreign
    assert {path: read(path) for path in watched} == before
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


def test_report_only_rechecks_foreign_marker_inside_publication_lock(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    output = real_project(tmp_path / "b")
    marker = output / portable.INCOMPLETE_FILE
    foreign = {"status": "PRESENTATION_PENDING", "owner": "other transaction"}
    watched = [output / name for name in ("校對工作階段.json", "人工判定資料庫.json",
                                          "待人工確認.json", "pipeline_status.json",
                                          "注音校對_最終報告.xlsx")]
    before = {path: read(path) for path in watched}
    original_lock = sp.project_delivery_lock
    injected = False

    @contextmanager
    def inject(root):
        nonlocal injected
        if not injected:
            injected = True
            sp.json_save(marker, foreign)
        with original_lock(root):
            yield

    with patch.object(sp, "project_delivery_lock", side_effect=inject), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        with pytest.raises(ValueError, match="未完成|標記"):
            sp.regenerate_report(output)
    assert injected and sp.json_load_strict(marker) == foreign
    assert read(output / "注音校對_最終報告.xlsx") == before[output / "注音校對_最終報告.xlsx"]
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


def simple_projects(root):
    first, second = _pdfs(root)
    source, target = root / "source", root / "target"
    with patch.object(sp, "decode", side_effect=_synthetic_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], source)
        sp.run_pipeline_pdfs([second], target)
    manifest = sp.json_load_strict(source / "校對工作階段.json")
    ReviewSaveService(source).save_event(manifest["records"][0]["review_id"], {
        "action": "確認非校對範圍", "exclusion_reason": "isolated fixture",
        "exclusion_evidence": "PDF page 1 visual position 1"})
    return source, target


def test_project_import_never_clears_foreign_marker_after_publish(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    source, target = simple_projects(tmp_path)
    marker = target / portable.INCOMPLETE_FILE
    foreign = {"status": "PRESENTATION_PENDING", "owner": "other transaction"}

    def inject(*args, **kwargs):
        sp.json_save(marker, foreign)

    with patch.object(portable, "_publish_portable_outputs", side_effect=inject):
        with pytest.raises(ValueError, match="標記|其他交易|變動"):
            portable.import_project_decisions(source, target)
    assert sp.json_load_strict(marker) == foreign
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


def test_committed_project_import_presentation_can_resume_safely(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    source, target = simple_projects(tmp_path)
    marker = target / portable.INCOMPLETE_FILE
    source_before = read(source / "人工判定資料庫.json")
    with patch.object(portable, "_publish_portable_outputs", side_effect=OSError("publication disk full")):
        with pytest.raises(OSError, match="publication disk full"):
            portable.import_project_decisions(source, target)
    assert marker.exists() and len(sp.json_load_strict(target / "人工判定資料庫.json")["events"]) == 1
    assert read(source / "人工判定資料庫.json") == source_before
    assert sp.repair_project_state(target) == target / "注音校對_最終報告.xlsx"
    assert not marker.exists()
    assert len(sp.json_load_strict(target / "人工判定資料庫.json")["events"]) == 1
    assert not list((tmp_path / "isolated-localappdata").rglob("*.sqlite*"))


def test_presentation_second_failure_then_repair_and_unknown_marker_refused(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    source, target = simple_projects(tmp_path)
    marker = target / portable.INCOMPLETE_FILE
    with patch.object(portable, "_publish_portable_outputs", side_effect=OSError("first disk fault")):
        with pytest.raises(OSError, match="first disk fault"):
            portable.import_project_decisions(source, target)
    original = read(marker)
    with patch.object(portable, "_publish_portable_outputs", side_effect=OSError("second disk fault")):
        with pytest.raises(OSError, match="second disk fault"):
            sp.repair_project_state(target)
    assert read(marker) == original
    with pytest.raises(ValueError, match="未完成"):
        sp.regenerate_report(target)
    assert sp.repair_project_state(target) == target / "注音校對_最終報告.xlsx"
    assert not marker.exists()
    sp.json_save(marker, {"status": "PRESENTATION_PENDING", "owner": "unknown legacy"})
    with pytest.raises(ValueError, match="專用發布恢復計畫"):
        sp.repair_project_state(target)
    assert marker.exists()


def test_presentation_plan_tamper_keeps_committed_db_and_marker(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    source, target = simple_projects(tmp_path)
    marker = target / portable.INCOMPLETE_FILE
    with patch.object(portable, "_publish_portable_outputs", side_effect=OSError("disk fault")):
        with pytest.raises(OSError, match="disk fault"):
            portable.import_project_decisions(source, target)
    payload = sp.json_load_strict(marker)
    payload["db_after_digest"] = "0" * 64
    sp.json_save(marker, payload)
    before = {path: read(path) for path in
              (target / "人工判定資料庫.json", marker, target / "pipeline_status.json")}
    with pytest.raises(ValueError, match="完整性"):
        sp.repair_project_state(target)
    assert {path: read(path) for path in before} == before


def test_presentation_precommit_failure_restores_receipts_and_clears_only_own_marker(
        tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    source, target = simple_projects(tmp_path)
    marker = target / portable.INCOMPLETE_FILE
    watched = [target / name for name in ("人工判定資料庫.json", portable.CONFLICT_FILE,
                                          portable.PENDING_ACTUAL_FILE, "pipeline_status.json",
                                          "注音校對_最終報告.xlsx")]
    before = {path: read(path) for path in watched}
    db_path = target / "人工判定資料庫.json"
    original_save = sp.json_save

    def fail_db(path, *args, **kwargs):
        if Path(path) == db_path:
            raise OSError("decision DB disk fault")
        return original_save(path, *args, **kwargs)

    with patch.object(sp, "json_save", side_effect=fail_db):
        with pytest.raises(OSError, match="decision DB disk fault"):
            portable.import_project_decisions(source, target)
    assert marker.exists()
    with pytest.raises(ValueError, match="已還原提交前"):
        sp.repair_project_state(target)
    assert not marker.exists()
    assert {path: read(path) for path in watched} == before


@pytest.mark.parametrize("lane", ["project", "same_session_expected"])
def test_foreign_marker_inserted_at_creation_is_not_overwritten(tmp_path, monkeypatch, lane):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    if lane == "project":
        source, target = simple_projects(tmp_path)
        operation = lambda: portable.import_project_decisions(source, target)
    else:
        target = real_project(tmp_path / "same-session")
        filled = expected_workbook(tmp_path, target)
        operation = lambda: sp.import_gpt_decisions(target, filled)
    marker = target / portable.INCOMPLETE_FILE
    foreign = {"status": "PRESENTATION_PENDING", "owner": "foreign transaction"}
    watched = [target / name for name in ("人工判定資料庫.json", portable.CONFLICT_FILE,
                                          portable.PENDING_ACTUAL_FILE, "pipeline_status.json")]
    before = {path: read(path) for path in watched}
    original_save = sp.json_save
    inserted = False

    def insert_at_marker(path, *args, **kwargs):
        nonlocal inserted
        if Path(path) == marker and not inserted:
            inserted = True
            original_save(marker, foreign, expected_sha256="")
        return original_save(path, *args, **kwargs)

    with patch.object(sp, "json_save", side_effect=insert_at_marker):
        with pytest.raises(ValueError, match="其他操作變更"):
            operation()
    assert inserted and sp.json_load_strict(marker) == foreign
    assert {path: read(path) for path in watched} == before


def test_same_session_expected_postcommit_report_failure_uses_repair(tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    output = real_project(tmp_path / "same-session")
    filled = expected_workbook(tmp_path, output)
    marker = output / portable.INCOMPLETE_FILE
    before = read(output / "人工判定資料庫.json")
    with patch.object(portable, "_publish_portable_outputs", side_effect=OSError("report disk fault")):
        with pytest.raises(OSError, match="report disk fault"):
            sp.import_gpt_decisions(output, filled)
    assert marker.exists()
    assert read(output / "人工判定資料庫.json") != before
    with pytest.raises(ValueError, match="未完成"):
        sp.regenerate_report(output)
    assert sp.repair_project_state(output) == output / "注音校對_最終報告.xlsx"
    assert not marker.exists()
    assert len(sp.json_load_strict(output / "人工判定資料庫.json")["events"]) == 1


def test_same_render_different_candidate_text_is_not_equivalent(tmp_path):
    first, second = tmp_path / "first.pdf", tmp_path / "second.pdf"
    for path, hidden in ((first, "sourcehiddentext"), (second, "differenthiddentext")):
        with fitz.open() as document:
            page = document.new_page(width=300, height=200)
            page.insert_text((20, 40), "visible lesson", fontsize=14)
            page.insert_text((20, 85), hidden, fontsize=11, render_mode=3)
            document.save(path)
    assert hashlib.sha256(first.read_bytes()).digest() != hashlib.sha256(second.read_bytes()).digest()
    from check_pronunciation_candidates import build_pdf_line_index
    assert build_pdf_line_index(first) != build_pdf_line_index(second)
    assert not portable._pages_equivalent(portable._page_signatures(first),
                                          portable._page_signatures(second))


def test_legacy_visual_proof_requires_exact_source_to_map_different_sha(tmp_path):
    source_pdf, same_bytes, different_bytes = (tmp_path / name for name in
                                               ("source.pdf", "copy.pdf", "metadata.pdf"))
    pdf(source_pdf, title="A")
    same_bytes.write_bytes(source_pdf.read_bytes())
    pdf(different_bytes, title="B")
    source = tmp_path / "source-project"
    manifest, _ = project(source, source_pdf, session="source")
    proof_path = portable.prepare_portable_project(source)
    proof = sp.json_load_strict(proof_path)
    old = copy.deepcopy(proof)
    for page in old["pdfs"][0]["pages"]:
        page.pop("candidate_text_sha256")
    old["integrity_sha256"] = hashlib.sha256(portable._canonical({
        key: value for key, value in old.items() if key != "integrity_sha256"})).hexdigest()
    sp.json_save(proof_path, old)
    source_pdf.unlink()
    target_same = tmp_path / "target-same"
    same_manifest, _ = project(target_same, same_bytes, session="target-same")
    source_proof = portable._load_proof(source, manifest)
    assert portable._match_pdfs(source_proof, manifest, same_manifest, target_same)[0] == {
        manifest["pdfs"][0]["pdf_sha256"]: same_manifest["pdfs"][0]["pdf_sha256"]}
    target_changed = tmp_path / "target-changed"
    changed_manifest, _ = project(target_changed, different_bytes, session="target-changed")
    with pytest.raises(ValueError, match="缺少校對文字層"):
        portable._match_pdfs(source_proof, manifest, changed_manifest, target_changed)
    source_pdf.write_bytes(same_bytes.read_bytes())
    upgraded = portable._load_proof(source, manifest)
    assert portable._match_pdfs(upgraded, manifest, changed_manifest, target_changed)[0] == {
        manifest["pdfs"][0]["pdf_sha256"]: changed_manifest["pdfs"][0]["pdf_sha256"]}


def test_pdf_sha_and_page_proof_share_the_same_bytes_snapshot(tmp_path):
    source, replacement = tmp_path / "source.pdf", tmp_path / "replacement.pdf"
    pdf(source, title="original", text="original lesson")
    pdf(replacement, title="replacement", text="other lesson")
    sealed_sha = hashlib.sha256(source.read_bytes()).hexdigest()
    original_pages = portable._page_signatures(source)
    replacement_pages = portable._page_signatures(replacement)
    manifest = {"session_id": "A", "manifest_integrity_sha256": "sealed-A",
                "pdfs": [{"pdf": str(source), "pdf_name": source.name,
                          "pdf_sha256": sealed_sha}]}
    original_source_pdf = portable._source_pdf

    def swap_source(*args):
        selected = original_source_pdf(*args)
        os.replace(replacement, source)
        return selected

    with patch.object(portable, "_source_pdf", side_effect=swap_source):
        with pytest.raises((ValueError, FileNotFoundError), match="SHA|變動|內容"):
            portable._proof_payload(tmp_path, manifest)
    assert original_pages != replacement_pages

    good, target, swap = (tmp_path / name for name in ("good.pdf", "target.pdf", "swap.pdf"))
    pdf(good, title="good", text="same lesson")
    pdf(target, title="target", text="different lesson")
    swap.write_bytes(good.read_bytes())
    source_sha = hashlib.sha256(good.read_bytes()).hexdigest()
    target_sha = hashlib.sha256(target.read_bytes()).hexdigest()
    proof = {"pdfs": [{"pdf_sha256": source_sha, "pdf_name": good.name,
                       "pages": portable._page_signatures(good)}]}
    target_manifest = {"pdfs": [{"pdf": str(target), "pdf_name": target.name,
                                 "pdf_sha256": target_sha}]}

    def swap_target(*args):
        selected = original_source_pdf(*args)
        os.replace(swap, target)
        return selected

    with patch.object(portable, "_source_pdf", side_effect=swap_target):
        with pytest.raises((ValueError, FileNotFoundError), match="SHA|變動|內容"):
            portable._match_pdfs(proof, manifest, target_manifest, tmp_path)

    missing_source = tmp_path / "missing.pdf"
    local = tmp_path / "local.pdf"
    pdf(local, title="old local", text="unrelated lesson")
    stale_pages = portable._page_signatures(local)
    local.write_bytes(good.read_bytes())
    missing_manifest = {"pdfs": [{"pdf": str(missing_source), "pdf_name": missing_source.name,
                                  "pdf_sha256": source_sha}]}
    with pytest.raises((ValueError, FileNotFoundError), match="SHA|變動|頁面|內容"):
        portable._proof_from_explicit_local(tmp_path, missing_manifest, local, stale_pages)


@pytest.mark.parametrize("partial_presentation", [False, True])
def test_cross_session_expected_postcommit_interrupt_can_repair(
        tmp_path, monkeypatch, partial_presentation):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = _pdfs(tmp_path)
    source, target = tmp_path / "source", tmp_path / "target"
    with patch.object(sp, "decode", side_effect=_synthetic_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], source)
        sp.run_pipeline_pdfs([second], target)
    filled = expected_workbook(tmp_path, source)
    source_db_before = read(source / "人工判定資料庫.json")
    target_manifest = sp.json_load_strict(target / "校對工作階段.json")

    def interrupt(output_dir, manifest, db):
        if partial_presentation:
            sp.save_pending_json(output_dir, manifest, db)
        raise KeyboardInterrupt("abrupt stop after commit")

    with patch.object(portable, "_publish_portable_outputs", side_effect=interrupt):
        with pytest.raises(KeyboardInterrupt, match="abrupt stop"):
            sp.import_gpt_decisions(target, filled)
    marker = target / portable.INCOMPLETE_FILE
    assert marker.exists()
    assert len(sp.json_load_strict(target / "人工判定資料庫.json")["events"]) == 1
    assert read(source / "人工判定資料庫.json") == source_db_before
    with pytest.raises(ValueError, match="未完成"):
        sp.regenerate_report(target)
    assert sp.repair_project_state(target) == target / "注音校對_最終報告.xlsx"
    assert not marker.exists()
    assert sp.json_load_strict(target / "校對工作階段.json")["session_id"] == target_manifest["session_id"]
    assert sp.json_load_strict(target / "pipeline_status.json")["user_report"] == str(
        target / "注音校對_最終報告.xlsx")


def test_failed_rawdict_extraction_cannot_seal_or_match_pdf_text(tmp_path, monkeypatch):
    first, second = tmp_path / "first.pdf", tmp_path / "second.pdf"
    for path, hidden in ((first, "hidden source"), (second, "hidden target")):
        with fitz.open() as document:
            page = document.new_page(width=300, height=200)
            page.insert_text((20, 40), "visible lesson")
            page.insert_text((20, 80), hidden, render_mode=3)
            document.save(path)
    assert portable._page_signatures(first) != portable._page_signatures(second)
    source, target = tmp_path / "source", tmp_path / "target"
    source_manifest, _ = project(source, first, session="source")
    target_manifest, _ = project(target, second, session="target")
    watched = [source / portable.PROOF_FILE, source / "人工判定資料庫.json",
               target / "人工判定資料庫.json", target / portable.INCOMPLETE_FILE]
    before = {path: read(path) for path in watched}
    real_get_text = fitz.Page.get_text

    def failed_rawdict(page, option="text", *args, **kwargs):
        if option == "rawdict":
            raise RuntimeError("RAWDICT unavailable")
        return real_get_text(page, option, *args, **kwargs)

    with patch.object(fitz.Page, "get_text", failed_rawdict):
        with pytest.raises((ValueError, RuntimeError), match="RAWDICT|文字|抽取"):
            portable.prepare_portable_project(source)
        from check_pronunciation_candidates import build_pdf_line_index
        assert build_pdf_line_index(first) == {1: []}
    assert {path: read(path) for path in watched} == before
    proof = portable._proof_payload(source, source_manifest)
    with patch.object(fitz.Page, "get_text", failed_rawdict):
        with pytest.raises((ValueError, RuntimeError), match="RAWDICT|文字|抽取"):
            portable._match_pdfs(proof, source_manifest, target_manifest, target)
    assert {path: read(path) for path in watched} == before


def test_stale_same_session_excel_cannot_overwrite_unresolved_project_conflict(
        tmp_path, monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "isolated-localappdata"))
    first, second = _pdfs(tmp_path)
    source, target = tmp_path / "source", tmp_path / "target"
    with patch.object(sp, "decode", side_effect=_synthetic_decode), \
         patch.object(sp, "actual_workbook_global_exact_dependencies", return_value=()), \
         patch("check_pronunciation_candidates.actual_workbook_global_exact_dependencies", return_value=()):
        sp.run_pipeline_pdfs([first], source)
        sp.run_pipeline_pdfs([second], target)
    target_manifest = sp.json_load_strict(target / "校對工作階段.json")
    target_ids = [item["review_id"] for item in target_manifest["records"]]
    old_excel = expected_workbook(tmp_path, target)
    source_id = sp.json_load_strict(source / "校對工作階段.json")["records"][0]["review_id"]
    ReviewSaveService(source).save_event(source_id, {
        "action": "確認非校對範圍", "exclusion_reason": "source decision",
        "exclusion_evidence": "A PDF page 1"})
    ReviewSaveService(target).save_event(target_ids[0], {
        "action": "確認非校對範圍", "exclusion_reason": "target decision",
        "exclusion_evidence": "B PDF page 1"})
    assert portable.import_project_decisions(source, target)["conflicts"]
    db = sp.json_load_strict(target / "人工判定資料庫.json")
    assert portable.validate_conflict_state(target, target_manifest, db) == [target_ids[0]]
    watched = [target / name for name in ("人工判定資料庫.json", portable.CONFLICT_FILE,
                                          portable.INCOMPLETE_FILE, "待人工確認.json",
                                          "pipeline_status.json", "注音校對_最終報告.xlsx")]
    before = {path: read(path) for path in watched}
    with pytest.raises(ValueError, match="衝突|裁決"):
        sp.import_gpt_decisions(target, old_excel)
    assert {path: read(path) for path in watched} == before
    assert portable.validate_conflict_state(
        target, target_manifest, sp.json_load_strict(target / "人工判定資料庫.json")) == [target_ids[0]]
    ReviewSaveService(target).save_event(target_ids[0], {
        "action": "確認非校對範圍", "exclusion_reason": "fresh local visual",
        "exclusion_evidence": "B PDF page 1 rechecked"})
    assert portable.validate_conflict_state(
        target, target_manifest, sp.json_load_strict(target / "人工判定資料庫.json")) == []
    filled = expected_workbook(tmp_path, target)
    assert sp.import_gpt_decisions(target, filled)[0] == 1
    assert not (target / portable.INCOMPLETE_FILE).exists()
    assert len(sp.json_load_strict(target / "人工判定資料庫.json")["events"]) == 2
