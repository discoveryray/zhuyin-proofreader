from __future__ import annotations

import csv
import hashlib
import json
import os
import shutil
import zipfile
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

import fitz
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from occurrence_ledger import canonical_bopomofo
from cff_zhuyin_decoder import cff_full_annotation_signature
from cross_version_compat import review_id_schema_compatible, schema_compatible
from export_pdf_text_diagnostics import TrueTypeGlyphInspector
from exact_glyph_identity import (
    ACTUAL_GROUP_IDENTITY_KINDS,
    CFF_GLYPH_SHA256,
    GLOBAL_ELIGIBLE_COMPLETE_CFF_RECORDING_V1,
    GLOBAL_ELIGIBLE_SIMPLE_GLYF_V1,
    NON_GLOBAL_ELIGIBLE,
    TTF_GLYF_SHA256,
    actual_group_id,
    actual_group_identity_payload,
    actual_group_match_key,
    canonical_actual_group_identity,
    legacy_v1_actual_group_id,
)
from global_exact_glyph_library import (
    GlobalExactGlyphIdentity,
    canonical_global_exact_identity,
)

ACTUAL_REVIEW_SCHEMA_VERSION = "1.1"
MANUAL_ACTUAL_STAGING_SCHEMA_VERSION = "1.1"
LEGACY_MANUAL_ACTUAL_STAGING_SCHEMA_VERSION = "1.0"
MANUAL_ACTUAL_STAGING_FILE = "manual_actual_staging.json"
ACTUAL_PENDING_STATES = frozenset({"ACTUAL_DECODE_ERROR", "ACTUAL_UNRESOLVED"})
USER_GLYF_FILE = "user_verified_glyf_fingerprints.csv"
USER_CFF_FILE = "user_verified_cff_glyph_fingerprints.csv"
OCCURRENCE_OVERRIDE_FILE = "manual_actual_occurrence_overrides.csv"
GLYPH_CONFLICT_FILE = "glyph_truth_conflicts.csv"
GLYPH_PROVENANCE_FILE = "glyph_truth_provenance.csv"
STATIC_TTF_FINGERPRINT_FILE = "ttf_verified_glyf_fingerprints.csv"

USER_GLYF_HEADERS = [
    "glyph_sha256", "bopomofo", "verification_level", "source_count",
    "source_examples", "updated_at", "notes",
]
USER_CFF_HEADERS = [
    "style_group", "glyph_sha256", "full_signature", "bopomofo", "verification_level",
    "source_count", "source_examples", "updated_at", "notes",
]
OVERRIDE_HEADERS = [
    "pdf_contains", "pdf_excludes", "page", "target_char", "stable_key",
    "x0", "y0", "actual_reading", "source", "note",
]
GLYPH_CONFLICT_HEADERS = [
    "kind", "style_group", "glyph_sha256", "status", "readings",
    "source_occurrence_ids", "updated_at", "notes",
]
GLYPH_PROVENANCE_HEADERS = [
    "event_id", "kind", "style_group", "glyph_sha256", "bopomofo",
    "event_type", "occurrence_ids", "source", "created_at", "note",
]

_AUTHORITATIVE_DYNAMIC_ACTUAL_FILES = (
    OCCURRENCE_OVERRIDE_FILE,
    USER_GLYF_FILE,
    USER_CFF_FILE,
    GLYPH_CONFLICT_FILE,
    GLYPH_PROVENANCE_FILE,
)

_MANUAL_ACTUAL_STAGING_ROOT_FIELDS = frozenset({"schema_version", "staged_groups"})
_MANUAL_ACTUAL_STAGING_GROUP_FIELDS = frozenset({
    "group_id", "group_snapshot", "kind", "exact_key", "style_group",
    "reading", "checked_occurrence_ids", "member_occurrence_ids", "note",
    "staged_at", "updated_at", "source",
})
_MANUAL_ACTUAL_STAGING_GROUP_KINDS = ACTUAL_GROUP_IDENTITY_KINDS


def _text(value: Any) -> str:
    return str(value or "").strip()


def _canon(value: Any) -> str:
    return canonical_bopomofo(value)


def _hash_payload(payload: Mapping[str, Any]) -> str:
    raw = json.dumps(dict(payload), ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def _style_sheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    if ws.max_row:
        fill = PatternFill("solid", fgColor="1F4E78")
        font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        max_len = 0
        for row in range(1, min(ws.max_row, 120) + 1):
            value = ws.cell(row, col).value
            max_len = max(max_len, len(str(value or "")))
        ws.column_dimensions[get_column_letter(col)].width = min(max(10, max_len + 2), 48)


def _read_csv(path: Path, headers: Sequence[str]) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        actual_headers = [str(x or "") for x in (reader.fieldnames or [])]
        if actual_headers != list(headers):
            raise ValueError(f"{path.name} 欄位格式不符，禁止自動覆寫")
        return [{h: _text(row.get(h)) for h in headers} for row in reader]


def _write_csv(path: Path, headers: Sequence[str], rows: Iterable[Mapping[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with tmp.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(headers))
        writer.writeheader()
        for row in rows:
            writer.writerow({h: row.get(h, "") for h in headers})
    tmp.replace(path)


def manual_actual_staging_path(root: Path) -> Path:
    """Return the project-local pending-intent store path.

    The caller supplies ``project_actual_evidence_root(output_dir)``.  This file
    deliberately lives beside dynamic actual evidence, but it is not itself
    authoritative evidence and is not part of ``dynamic_actual_hashes``.
    """
    return Path(root) / MANUAL_ACTUAL_STAGING_FILE


def _empty_manual_actual_staging() -> dict[str, Any]:
    return {
        "schema_version": MANUAL_ACTUAL_STAGING_SCHEMA_VERSION,
        "staged_groups": [],
    }


def _reject_duplicate_json_keys(pairs: Sequence[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f"manual actual staging JSON key 重複：{key}")
        result[key] = value
    return result


def _staging_text(value: Any, field: str, *, allow_empty: bool = False) -> str:
    if not isinstance(value, str):
        raise ValueError(f"manual actual staging {field} 必須是字串")
    normalized = value.strip()
    if value != normalized:
        raise ValueError(f"manual actual staging {field} 不得含首尾空白")
    if not allow_empty and not normalized:
        raise ValueError(f"manual actual staging {field} 不得空白")
    return normalized


def _staging_id_list(value: Any, field: str) -> list[str]:
    if not isinstance(value, list):
        raise ValueError(f"manual actual staging {field} 必須是 list")
    result = [_staging_text(item, f"{field}[{index}]") for index, item in enumerate(value)]
    if len(result) != len(set(result)):
        raise ValueError(f"manual actual staging {field} 含重複 occurrence_id")
    return result


def _staging_timestamp(value: Any, field: str) -> tuple[str, datetime]:
    text = _staging_text(value, field)
    if "T" not in text:
        raise ValueError(f"manual actual staging {field} 必須是 ISO 8601 datetime")
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError as exc:
        raise ValueError(f"manual actual staging {field} 不是合法 ISO 8601 datetime") from exc
    return text, parsed


def _validate_manual_actual_staged_group(
    raw: Any,
    index: int,
    *,
    schema_version: str = MANUAL_ACTUAL_STAGING_SCHEMA_VERSION,
) -> dict[str, Any]:
    label = f"staged_groups[{index}]"
    if not isinstance(raw, dict):
        raise ValueError(f"manual actual staging {label} 必須是 object")
    missing = sorted(_MANUAL_ACTUAL_STAGING_GROUP_FIELDS - set(raw))
    unexpected = sorted(set(raw) - _MANUAL_ACTUAL_STAGING_GROUP_FIELDS)
    if missing:
        raise ValueError(f"manual actual staging {label} 缺少欄位：{missing}")
    if unexpected:
        raise ValueError(f"manual actual staging {label} 含未知欄位：{unexpected}")

    group_id = _staging_text(raw.get("group_id"), f"{label}.group_id")
    group_snapshot = _staging_text(raw.get("group_snapshot"), f"{label}.group_snapshot").lower()
    if len(group_snapshot) != 64 or any(ch not in "0123456789abcdef" for ch in group_snapshot):
        raise ValueError(f"manual actual staging {label}.group_snapshot 必須是 SHA-256")

    kind = _staging_text(raw.get("kind"), f"{label}.kind")
    if kind not in _MANUAL_ACTUAL_STAGING_GROUP_KINDS:
        raise ValueError(f"manual actual staging {label}.kind 不支援：{kind}")
    exact_key = _staging_text(raw.get("exact_key"), f"{label}.exact_key")
    style_group = _staging_text(raw.get("style_group"), f"{label}.style_group", allow_empty=True)
    try:
        identity = canonical_actual_group_identity(kind, exact_key, style_group)
    except ValueError as exc:
        raise ValueError(f"manual actual staging {label} exact identity 無效：{exc}") from exc
    kind = identity["kind"]
    exact_key = identity["exact_key"]
    style_group = identity["style_group"]

    if schema_version == LEGACY_MANUAL_ACTUAL_STAGING_SCHEMA_VERSION:
        expected_group_id = legacy_v1_actual_group_id(identity)
    else:
        expected_group_id = actual_group_id(identity)
    if group_id != expected_group_id:
        raise ValueError(f"manual actual staging {label}.group_id 與 exact group identity 不一致")

    reading = _staging_text(raw.get("reading"), f"{label}.reading")
    canonical = _canon(reading)
    if not canonical or reading != canonical:
        raise ValueError(f"manual actual staging {label}.reading 必須是 canonical Bopomofo")

    checked_ids = _staging_id_list(raw.get("checked_occurrence_ids"), f"{label}.checked_occurrence_ids")
    member_ids = _staging_id_list(raw.get("member_occurrence_ids"), f"{label}.member_occurrence_ids")
    if not member_ids:
        raise ValueError(f"manual actual staging {label}.member_occurrence_ids 不得為空")
    if not checked_ids:
        raise ValueError(f"manual actual staging {label}.checked_occurrence_ids 不得為空")
    unknown_checked = sorted(set(checked_ids) - set(member_ids))
    if unknown_checked:
        raise ValueError(
            f"manual actual staging {label}.checked_occurrence_ids 不屬於 group members：{unknown_checked}"
        )

    note = _staging_text(raw.get("note"), f"{label}.note", allow_empty=True)
    source = _staging_text(raw.get("source"), f"{label}.source")
    staged_at, staged_dt = _staging_timestamp(raw.get("staged_at"), f"{label}.staged_at")
    updated_at, updated_dt = _staging_timestamp(raw.get("updated_at"), f"{label}.updated_at")
    if (staged_dt.tzinfo is None) != (updated_dt.tzinfo is None):
        raise ValueError(f"manual actual staging {label} timestamps timezone contract 不一致")
    if updated_dt < staged_dt:
        raise ValueError(f"manual actual staging {label}.updated_at 不得早於 staged_at")

    return {
        "group_id": group_id,
        "group_snapshot": group_snapshot,
        "kind": kind,
        "exact_key": exact_key,
        "style_group": style_group,
        "reading": reading,
        "checked_occurrence_ids": checked_ids,
        "member_occurrence_ids": member_ids,
        "note": note,
        "staged_at": staged_at,
        "updated_at": updated_at,
        "source": source,
    }


def _validate_manual_actual_staging_document(raw: Any) -> dict[str, Any]:
    if not isinstance(raw, dict):
        raise ValueError("manual actual staging root 必須是 object")
    missing = sorted(_MANUAL_ACTUAL_STAGING_ROOT_FIELDS - set(raw))
    unexpected = sorted(set(raw) - _MANUAL_ACTUAL_STAGING_ROOT_FIELDS)
    if missing:
        raise ValueError(f"manual actual staging root 缺少欄位：{missing}")
    if unexpected:
        raise ValueError(f"manual actual staging root 含未知欄位：{unexpected}")
    schema_version = raw.get("schema_version")
    if not isinstance(schema_version, str):
        raise ValueError(
            "manual actual staging schema 不相容："
            f"required={MANUAL_ACTUAL_STAGING_SCHEMA_VERSION!r}, observed={schema_version!r}"
        )
    if schema_version not in {
        MANUAL_ACTUAL_STAGING_SCHEMA_VERSION,
        LEGACY_MANUAL_ACTUAL_STAGING_SCHEMA_VERSION,
    }:
        raise ValueError(
            "manual actual staging schema 不相容："
            f"required={MANUAL_ACTUAL_STAGING_SCHEMA_VERSION!r}, observed={raw.get('schema_version')!r}"
        )
    groups_raw = raw.get("staged_groups")
    if not isinstance(groups_raw, list):
        raise ValueError("manual actual staging staged_groups 必須是 list")

    groups: list[dict[str, Any]] = []
    seen_group_ids: set[str] = set()
    seen_snapshots: set[str] = set()
    seen_occurrence_ids: set[str] = set()
    for index, item in enumerate(groups_raw):
        group = _validate_manual_actual_staged_group(
            item,
            index,
            schema_version=str(schema_version),
        )
        group_id = group["group_id"]
        if group_id in seen_group_ids:
            raise ValueError(f"manual actual staging group_id 重複：{group_id}")
        if group["group_snapshot"] in seen_snapshots:
            raise ValueError(f"manual actual staging group_snapshot 重複：{group['group_snapshot']}")
        overlap = seen_occurrence_ids & set(group["member_occurrence_ids"])
        if overlap:
            raise ValueError(f"manual actual staging occurrence identity 同時屬於多組：{sorted(overlap)}")
        seen_group_ids.add(group_id)
        seen_snapshots.add(group["group_snapshot"])
        seen_occurrence_ids.update(group["member_occurrence_ids"])
        groups.append(group)
    if (
        schema_version == LEGACY_MANUAL_ACTUAL_STAGING_SCHEMA_VERSION
        and any(group["kind"] == CFF_GLYPH_SHA256 for group in groups)
    ):
        raise ValueError(
            "manual actual staging v1.0 含 pending CFF decision；"
            "舊 group_id 未綁定 style_group，禁止靜默重新解釋，請重新進行 actual review"
        )
    groups.sort(key=lambda group: group["group_id"])
    return {
        "schema_version": MANUAL_ACTUAL_STAGING_SCHEMA_VERSION,
        "staged_groups": groups,
    }


def load_manual_actual_staging(root: Path) -> dict[str, Any]:
    """Load and fully validate pending manual actual intent.

    Absence represents an empty queue.  The only legacy adapters are the
    explicitly validated v1.0 empty and non-CFF forms; a pending v1.0 CFF
    decision cannot prove the new style-bound identity and fails closed without
    rewriting the file.  Malformed JSON, duplicate keys, unknown schema/fields,
    or contradictory identities also fail closed.
    """
    path = manual_actual_staging_path(root)
    if not path.exists():
        return _empty_manual_actual_staging()
    try:
        raw = json.loads(
            path.read_text(encoding="utf-8-sig"),
            object_pairs_hook=_reject_duplicate_json_keys,
        )
    except Exception as exc:
        if isinstance(exc, ValueError) and str(exc).startswith("manual actual staging"):
            raise
        raise ValueError(f"manual actual staging JSON 損壞：{path}") from exc
    return _validate_manual_actual_staging_document(raw)


def _write_manual_actual_staging(root: Path, document: Mapping[str, Any]) -> Path:
    validated = _validate_manual_actual_staging_document(dict(document))
    path = manual_actual_staging_path(root)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    try:
        with tmp.open("w", encoding="utf-8", newline="\n") as handle:
            json.dump(validated, handle, ensure_ascii=False, indent=2, sort_keys=True)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        tmp.replace(path)
    except Exception:
        try:
            tmp.unlink()
        except FileNotFoundError:
            pass
        raise
    return path


def stage_manual_actual_group(
    root: Path,
    group: Mapping[str, Any],
    reading: str,
    *,
    checked_occurrence_ids: Sequence[str],
    source: str,
    note: str = "",
) -> dict[str, Any]:
    """Durably upsert one uncommitted manual actual decision.

    This function extracts only actual group identity and audit fields.  It does
    not initialize/mutate authoritative evidence, call promotion logic, clear
    review events, refresh a PDF, decode, or regenerate reports.
    """
    if not isinstance(group, Mapping):
        raise ValueError("manual actual staging group 必須是 mapping")
    group_id = _staging_text(group.get("group_id"), "input.group_id")
    group_snapshot = _staging_text(group.get("group_snapshot"), "input.group_snapshot").lower()
    kind = _staging_text(group.get("kind"), "input.kind")
    exact_key = _staging_text(group.get("exact_key"), "input.exact_key")
    style_group = _staging_text(group.get("style_group", ""), "input.style_group", allow_empty=True)
    try:
        identity = canonical_actual_group_identity(kind, exact_key, style_group)
    except ValueError as exc:
        raise ValueError(f"manual actual staging input exact identity 無效：{exc}") from exc
    kind = identity["kind"]
    exact_key = identity["exact_key"]
    style_group = identity["style_group"]
    if group_id != actual_group_id(identity):
        raise ValueError("manual actual staging input.group_id 與 canonical exact identity 不一致")

    members = group.get("members")
    if not isinstance(members, list) or not members:
        raise ValueError("manual actual staging group.members 必須是非空 list")
    member_ids: list[str] = []
    for index, member in enumerate(members):
        if not isinstance(member, Mapping):
            raise ValueError(f"manual actual staging group.members[{index}] 必須是 mapping")
        member_ids.append(_staging_text(member.get("occurrence_id"), f"input.members[{index}].occurrence_id"))
        try:
            member_key = actual_group_match_key(actual_group_identity(member))
        except ValueError as exc:
            raise ValueError(
                f"manual actual staging group.members[{index}] exact identity 無效：{exc}"
            ) from exc
        if member_key != actual_group_match_key(identity):
            raise ValueError(
                f"manual actual staging group.members[{index}] 不屬於 canonical exact group"
            )
    if len(member_ids) != len(set(member_ids)):
        raise ValueError("manual actual staging group.members 含重複 occurrence_id")
    occurrence_count = group.get("occurrence_count")
    if occurrence_count is not None:
        if isinstance(occurrence_count, bool):
            raise ValueError("manual actual staging group.occurrence_count 無效")
        try:
            observed_count = int(occurrence_count)
        except (TypeError, ValueError) as exc:
            raise ValueError("manual actual staging group.occurrence_count 無效") from exc
        if observed_count != len(member_ids):
            raise ValueError("manual actual staging group.occurrence_count 與 members 不一致")

    if isinstance(checked_occurrence_ids, (str, bytes)) or not isinstance(checked_occurrence_ids, Sequence):
        raise ValueError("manual actual staging checked_occurrence_ids 必須是 sequence")
    checked_input = [
        _staging_text(item, f"input.checked_occurrence_ids[{index}]")
        for index, item in enumerate(checked_occurrence_ids)
    ]
    if len(checked_input) != len(set(checked_input)):
        raise ValueError("manual actual staging checked_occurrence_ids 含重複 occurrence_id")
    checked_set = set(checked_input)
    if not checked_set:
        raise ValueError("manual actual staging checked_occurrence_ids 不得為空")
    unknown_checked = sorted(checked_set - set(member_ids))
    if unknown_checked:
        raise ValueError(f"manual actual staging checked occurrence 不屬於 group members：{unknown_checked}")
    checked_ids = [occurrence_id for occurrence_id in member_ids if occurrence_id in checked_set]

    canonical_reading = _canon(reading)
    if not canonical_reading:
        raise ValueError("manual actual staging reading 不是合法單一注音")
    source_text = _staging_text(source, "input.source")
    note_text = _staging_text(note, "input.note", allow_empty=True)

    document = load_manual_actual_staging(root)
    existing = next(
        (item for item in document["staged_groups"] if item["group_id"] == group_id),
        None,
    )
    immutable_identity = {
        "group_snapshot": group_snapshot,
        "kind": kind,
        "exact_key": exact_key,
        "style_group": style_group,
        "member_occurrence_ids": member_ids,
    }
    if existing:
        changed = [key for key, value in immutable_identity.items() if existing.get(key) != value]
        if changed:
            raise ValueError(
                f"manual actual staging stale/contradictory group identity：{group_id} changed={changed}；請重新確認"
            )

    now_dt = datetime.now().astimezone()
    if existing:
        previous_updated = datetime.fromisoformat(existing["updated_at"].replace("Z", "+00:00"))
        if now_dt <= previous_updated:
            now_dt = previous_updated + timedelta(microseconds=1)
    now = now_dt.isoformat(timespec="microseconds")
    staged = {
        "group_id": group_id,
        **immutable_identity,
        "reading": canonical_reading,
        "checked_occurrence_ids": checked_ids,
        "note": note_text,
        "staged_at": existing["staged_at"] if existing else now,
        "updated_at": now,
        "source": source_text,
    }
    groups = [item for item in document["staged_groups"] if item["group_id"] != group_id]
    groups.append(staged)
    _write_manual_actual_staging(root, {
        "schema_version": MANUAL_ACTUAL_STAGING_SCHEMA_VERSION,
        "staged_groups": groups,
    })
    return dict(staged)


def remove_staged_manual_actual_group(
    root: Path,
    group_id: str,
    *,
    group_snapshot: str,
    updated_at: str,
) -> bool:
    """Atomically acknowledge one exact staged decision.

    ``updated_at`` is part of the compare-and-remove guard so a later upsert of
    the same group cannot be accidentally cleared by a stale batch-apply worker.
    """
    wanted_id = _staging_text(group_id, "remove.group_id")
    wanted_snapshot = _staging_text(group_snapshot, "remove.group_snapshot").lower()
    wanted_updated_at = _staging_text(updated_at, "remove.updated_at")
    document = load_manual_actual_staging(root)
    existing = next(
        (item for item in document["staged_groups"] if item["group_id"] == wanted_id),
        None,
    )
    if existing is None:
        return False
    if existing["group_snapshot"] != wanted_snapshot or existing["updated_at"] != wanted_updated_at:
        raise ValueError(f"manual actual staging decision 已更新：{wanted_id}；拒絕移除較新的 staged intent")
    remaining = [item for item in document["staged_groups"] if item["group_id"] != wanted_id]
    _write_manual_actual_staging(root, {
        "schema_version": MANUAL_ACTUAL_STAGING_SCHEMA_VERSION,
        "staged_groups": remaining,
    })
    return True


def _manual_actual_staging_ack_tokens(
    staged_decisions: Sequence[Mapping[str, Any]],
) -> list[dict[str, str]]:
    if isinstance(staged_decisions, (str, bytes)) or not isinstance(staged_decisions, Sequence):
        raise ValueError("manual actual staging batch acknowledgement 必須是 sequence")
    tokens: list[dict[str, str]] = []
    seen_group_ids: set[str] = set()
    for index, decision in enumerate(staged_decisions):
        if not isinstance(decision, Mapping):
            raise ValueError(f"manual actual staging acknowledgement[{index}] 必須是 mapping")
        group_id = _staging_text(decision.get("group_id"), f"acknowledgement[{index}].group_id")
        if group_id in seen_group_ids:
            raise ValueError(f"manual actual staging acknowledgement group_id 重複：{group_id}")
        seen_group_ids.add(group_id)
        tokens.append({
            "group_id": group_id,
            "group_snapshot": _staging_text(
                decision.get("group_snapshot"),
                f"acknowledgement[{index}].group_snapshot",
            ),
            "updated_at": _staging_text(
                decision.get("updated_at"),
                f"acknowledgement[{index}].updated_at",
            ),
        })
    tokens.sort(key=lambda token: token["group_id"])
    return tokens


def _validate_manual_actual_staging_ack_tokens(
    document: Mapping[str, Any],
    tokens: Sequence[Mapping[str, str]],
) -> None:
    current_by_id = {
        item["group_id"]: item
        for item in document.get("staged_groups", [])
    }
    for token in tokens:
        group_id = token["group_id"]
        current = current_by_id.get(group_id)
        if current is None:
            raise ValueError(
                f"manual actual staging decision 已不存在：{group_id}；拒絕 acknowledge stale batch"
            )
        if (
            current["group_snapshot"] != token["group_snapshot"]
            or current["updated_at"] != token["updated_at"]
        ):
            raise ValueError(
                f"manual actual staging decision 已更新：{group_id}；拒絕移除較新的 staged intent"
            )


def remove_staged_manual_actual_groups(
    root: Path,
    staged_decisions: Sequence[Mapping[str, Any]],
) -> dict[str, Any]:
    """Atomically acknowledge an exact frozen batch of staged decisions.

    All selected ``group_id + group_snapshot + updated_at`` tokens are checked
    against one freshly loaded document before a single atomic rewrite.  New
    unrelated decisions are retained from that current document; a missing or
    updated selected decision rejects the entire acknowledgement without write.
    """
    tokens = _manual_actual_staging_ack_tokens(staged_decisions)
    document = load_manual_actual_staging(root)
    _validate_manual_actual_staging_ack_tokens(document, tokens)
    if not tokens:
        return {
            "acknowledged_group_count": 0,
            "acknowledged_group_ids": [],
            "staging_remaining_count": len(document["staged_groups"]),
        }
    selected_ids = {token["group_id"] for token in tokens}
    remaining = [
        item for item in document["staged_groups"]
        if item["group_id"] not in selected_ids
    ]
    _write_manual_actual_staging(root, {
        "schema_version": MANUAL_ACTUAL_STAGING_SCHEMA_VERSION,
        "staged_groups": remaining,
    })
    return {
        "acknowledged_group_count": len(tokens),
        "acknowledged_group_ids": sorted(selected_ids),
        "staging_remaining_count": len(remaining),
    }



KNOWN_BAD_LEGACY_OVERRIDE = {
    "pdf_contains": "15_國小健體3下課本_單元5第2課_(學)",
    "page": "122",
    "target_char": "那",
    "stable_key": "CFF:DFKaiChuIn-Md-BPMF-BF#1311",
    "x0": "142.01",
    "y0": "76.595",
    "actual_reading": "ㄋㄚˋ",
}


def _migrate_v540_known_bad_override(root: Path) -> bool:
    """Remove one v5.3.1 patch that was later disproved by direct PDF glyph review.

    The migration is intentionally exact and source-scoped.  It does not infer
    any replacement actual.  Once removed, the normal CFF glyph decoder
    independently reconstructs the printed reading from the PDF.
    """
    path = Path(root) / OCCURRENCE_OVERRIDE_FILE
    if not path.exists():
        return False
    rows = _read_csv(path, OVERRIDE_HEADERS)
    kept = []
    removed = False
    for row in rows:
        exact = all(_text(row.get(k)) == v for k, v in KNOWN_BAD_LEGACY_OVERRIDE.items() if k != "actual_reading")
        source = _text(row.get("source"))
        reading = _canon(row.get("actual_reading"))
        generated_by_bad_patch = "2026-08-20" in source and ("GPT" in source or "原頁回標" in source)
        if exact and reading == KNOWN_BAD_LEGACY_OVERRIDE["actual_reading"] and generated_by_bad_patch:
            removed = True
            continue
        kept.append(row)
    if removed:
        _write_csv(path, OVERRIDE_HEADERS, kept)
    return removed

def ensure_user_evidence_files(root: Path) -> None:
    root = Path(root)
    _migrate_v540_known_bad_override(root)
    for name, headers in (
        (USER_GLYF_FILE, USER_GLYF_HEADERS),
        (USER_CFF_FILE, USER_CFF_HEADERS),
        (GLYPH_CONFLICT_FILE, GLYPH_CONFLICT_HEADERS),
        (GLYPH_PROVENANCE_FILE, GLYPH_PROVENANCE_HEADERS),
    ):
        path = root / name
        if not path.exists():
            _write_csv(path, headers, [])


def _snapshot_dynamic_actual_evidence(root: Path) -> dict[Path, bytes | None]:
    root = Path(root)
    return {
        root / name: (root / name).read_bytes() if (root / name).exists() else None
        for name in _AUTHORITATIVE_DYNAMIC_ACTUAL_FILES
    }


def _restore_dynamic_actual_evidence(backups: Mapping[Path, bytes | None]) -> None:
    for path, data in backups.items():
        if data is None:
            try:
                path.unlink()
            except FileNotFoundError:
                pass
        else:
            path.write_bytes(data)


def validate_dynamic_actual_evidence(root: Path) -> dict[str, Any]:
    root = Path(root)
    ensure_user_evidence_files(root)
    result = {"ok": True, "errors": [], "hashes": {}}
    specs = [
        (OCCURRENCE_OVERRIDE_FILE, OVERRIDE_HEADERS),
        (USER_GLYF_FILE, USER_GLYF_HEADERS),
        (USER_CFF_FILE, USER_CFF_HEADERS),
        (GLYPH_CONFLICT_FILE, GLYPH_CONFLICT_HEADERS),
    ]
    for name, headers in specs:
        path = root / name
        try:
            if name == OCCURRENCE_OVERRIDE_FILE and not path.exists():
                # A fresh installation may legitimately have no occurrence
                # overrides yet.  Absence is a stable empty dynamic evidence
                # state and is still fingerprinted distinctly from a real file.
                result["hashes"][name] = "ABSENT"
                continue
            rows = _read_csv(path, headers)
            for idx, row in enumerate(rows, 2):
                reading_col = "actual_reading" if name == OCCURRENCE_OVERRIDE_FILE else "bopomofo"
                reading = _canon(row.get(reading_col))
                if row.get(reading_col) and not reading:
                    raise ValueError(f"row {idx} {reading_col} 不是合法單一注音")
                if name == USER_GLYF_FILE and row.get("verification_level") == "VERIFIED_EXACT_GLYPH":
                    if not row.get("glyph_sha256") or len(row.get("glyph_sha256", "")) != 64:
                        raise ValueError(f"row {idx} glyph_sha256 無效")
                if name == USER_CFF_FILE and row.get("verification_level") == "VERIFIED_EXACT_GLYPH":
                    sha = _text(row.get("glyph_sha256")).lower()
                    if not row.get("style_group") or len(sha) != 64 or any(ch not in "0123456789abcdef" for ch in sha):
                        raise ValueError(f"row {idx} CFF exact full-glyph SHA-256 key 不完整")
                if name == GLYPH_CONFLICT_FILE:
                    sha = _text(row.get("glyph_sha256")).lower()
                    kind = _text(row.get("kind"))
                    if kind not in {"TTF_GLYF_SHA256", "CFF_GLYPH_SHA256"}:
                        raise ValueError(f"row {idx} kind 無效")
                    if len(sha) != 64 or any(ch not in "0123456789abcdef" for ch in sha):
                        raise ValueError(f"row {idx} glyph_sha256 無效")
                    if _text(row.get("status")) != "GLYPH_TRUTH_CONFLICT":
                        raise ValueError(f"row {idx} status 必須為 GLYPH_TRUTH_CONFLICT")
                    readings = [_canon(x) for x in _text(row.get("readings")).split("|") if _text(x)]
                    if len(set(filter(None, readings))) < 2:
                        raise ValueError(f"row {idx} conflict 至少需要兩個互斥讀音")
            result["hashes"][name] = _sha256_file(path)
        except Exception as exc:
            result["ok"] = False
            result["errors"].append(f"{name}：{exc}")
    return result


def _subset_sha256(rows: Sequence[Mapping[str, Any]], headers: Sequence[str]) -> str:
    payload = [
        {h: _text(row.get(h)) for h in headers}
        for row in rows
    ]
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def _override_applies_to_pdf(row: Mapping[str, Any], pdf_path: Path) -> bool:
    stem = Path(pdf_path).stem
    scope = _text(row.get("pdf_contains"))
    if scope and scope not in stem:
        return False
    excludes = _text(row.get("pdf_excludes"))
    if excludes:
        parts = [x.strip() for x in excludes.replace(";", "|").replace(",", "|").split("|") if x.strip()]
        if any(part in stem for part in parts):
            return False
    return True


def actual_workbook_dynamic_dependencies(path: Path) -> dict[str, Any] | None:
    """Read exact-glyph dependencies already observed in one actual workbook.

    This is intentionally a lightweight cache-inspection step, not a decoder.
    It lets a new user-verified glyph invalidate only PDFs that actually contain
    that exact glyph.  Older workbooks without the v5.4 evidence columns return
    None, which safely falls back to hashing the complete dynamic truth files.
    """
    path = Path(path)
    if not path.exists():
        return None
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            if "實際注音" not in wb.sheetnames:
                return None
            ws = wb["實際注音"]
            iterator = ws.iter_rows(values_only=True)
            headers = [str(v or "") for v in next(iterator)]
            idx = {name: i for i, name in enumerate(headers)}
            ttf_col = idx.get("TTF字形SHA256")
            cff_col = idx.get("CFF整字字形SHA256")
            style_col = idx.get("CFF樣式群組")
            if ttf_col is None and cff_col is None:
                return None
            ttf: set[str] = set()
            cff: set[tuple[str, str]] = set()
            for row in iterator:
                if ttf_col is not None and ttf_col < len(row):
                    sha = _text(row[ttf_col]).lower()
                    if len(sha) == 64 and all(ch in "0123456789abcdef" for ch in sha):
                        ttf.add(sha)
                if cff_col is not None and cff_col < len(row):
                    sha = _text(row[cff_col]).lower()
                    style = _text(row[style_col]) if style_col is not None and style_col < len(row) else ""
                    if style and len(sha) == 64 and all(ch in "0123456789abcdef" for ch in sha):
                        cff.add((style, sha))
            return {
                "ttf_glyph_sha256": sorted(ttf),
                "cff_glyph_keys": [list(item) for item in sorted(cff)],
            }
        finally:
            wb.close()
    except Exception:
        return None


def actual_workbook_global_exact_dependencies(
    path: Path,
    pdf_path: Path,
) -> tuple[GlobalExactGlyphIdentity, ...] | None:
    """Revalidate one workbook's global-exact roster against current PDF bytes.

    Existing workbook SHA columns are candidates, never TTF eligibility proof.
    Every TTF candidate is re-parsed from its current embedded font and passed
    through ``TrueTypeGlyphInspector.global_exact_identity``.  ``None`` means
    that the dependency contract is missing/unprovable and therefore forces a
    decode; a structurally unreadable XLSX raises so artifact corruption cannot
    be mistaken for an empty roster.
    """
    path = Path(path)
    pdf_path = Path(pdf_path)
    if not path.exists():
        return None

    required_columns = {
        "字形架構",
        "font_xref",
        "注音元件ID",
        "TTF字形SHA256",
        "CFF樣式群組",
        "CFF整字字形SHA256",
    }
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"actual workbook 損壞，無法驗證 global exact dependency：{path}") from exc

    document = None
    try:
        if "實際注音" not in workbook.sheetnames:
            return None
        worksheet = workbook["實際注音"]
        iterator = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = next(iterator)
        except StopIteration:
            return None
        headers = [str(value or "") for value in raw_headers]
        if len(headers) != len(set(headers)):
            raise ValueError(f"actual workbook 實際注音欄位重複：{path}")
        if not required_columns.issubset(headers):
            return None
        index = {name: offset for offset, name in enumerate(headers)}

        try:
            document = fitz.open(pdf_path)
        except Exception:
            return None
        inspectors: dict[int, TrueTypeGlyphInspector | None] = {}
        dependencies: dict[tuple[str, str, str], GlobalExactGlyphIdentity] = {}

        for row in iterator:
            def cell(name: str) -> Any:
                offset = index[name]
                return row[offset] if offset < len(row) else None

            raw_ttf_sha = str(cell("TTF字形SHA256") or "").strip()
            raw_cff_sha = str(cell("CFF整字字形SHA256") or "").strip()
            raw_architecture = str(cell("字形架構") or "")
            architecture = raw_architecture.strip()
            if raw_architecture != architecture:
                return None
            if architecture == "TrueType複合元件":
                if not raw_ttf_sha or raw_cff_sha:
                    return None
            elif architecture == "CFF整字注音":
                if raw_ttf_sha or not raw_cff_sha:
                    return None
            else:
                # Every row in the actual sheet is a detected occurrence.  An
                # unknown/blank architecture cannot prove that its exact
                # dependency roster is complete.
                return None
            if raw_ttf_sha and raw_cff_sha:
                return None

            if raw_ttf_sha:
                if (
                    raw_ttf_sha != raw_ttf_sha.lower()
                    or len(raw_ttf_sha) != 64
                    or any(ch not in "0123456789abcdef" for ch in raw_ttf_sha)
                ):
                    return None
                try:
                    xref_value = cell("font_xref")
                    glyph_value = cell("注音元件ID")
                    if isinstance(xref_value, bool) or isinstance(glyph_value, bool):
                        return None
                    if isinstance(xref_value, float) and not xref_value.is_integer():
                        return None
                    if isinstance(glyph_value, float) and not glyph_value.is_integer():
                        return None
                    xref = int(xref_value)
                    glyph_id = int(glyph_value)
                except (TypeError, ValueError, OverflowError):
                    return None
                if xref < 0 or glyph_id < 0:
                    return None
                if xref not in inspectors:
                    try:
                        _name, extension, _font_type, data = document.extract_font(xref)
                        inspectors[xref] = (
                            TrueTypeGlyphInspector(data)
                            if str(extension or "").lower() == "ttf"
                            else None
                        )
                    except Exception:
                        inspectors[xref] = None
                inspector = inspectors[xref]
                if inspector is None:
                    return None
                result = inspector.global_exact_identity(glyph_id)
                current_sha = str(result.get("glyph_sha256") or "")
                if current_sha != raw_ttf_sha:
                    return None
                eligibility = str(result.get("eligibility") or "")
                if eligibility == GLOBAL_ELIGIBLE_SIMPLE_GLYF_V1:
                    identity = canonical_global_exact_identity(
                        TTF_GLYF_SHA256,
                        "",
                        current_sha,
                        eligibility,
                    )
                elif eligibility == NON_GLOBAL_ELIGIBLE:
                    identity = GlobalExactGlyphIdentity(
                        TTF_GLYF_SHA256,
                        "",
                        current_sha,
                        NON_GLOBAL_ELIGIBLE,
                    )
                else:
                    return None
                existing = dependencies.get(identity.tuple)
                if existing is not None and existing.identity_eligibility != identity.identity_eligibility:
                    return None
                dependencies[identity.tuple] = identity

            if raw_cff_sha:
                if (
                    raw_cff_sha != raw_cff_sha.lower()
                    or len(raw_cff_sha) != 64
                    or any(ch not in "0123456789abcdef" for ch in raw_cff_sha)
                ):
                    return None
                raw_style_value = cell("CFF樣式群組")
                raw_style = str(raw_style_value or "")
                style = raw_style.strip()
                # Unknown CFF families have no canonical global identity.  They
                # remain available to the existing direct/zero-map decoders.
                if not style:
                    continue
                if raw_style != style:
                    return None
                try:
                    identity = canonical_global_exact_identity(
                        CFF_GLYPH_SHA256,
                        style,
                        raw_cff_sha,
                        GLOBAL_ELIGIBLE_COMPLETE_CFF_RECORDING_V1,
                    )
                except Exception:
                    return None
                dependencies[identity.tuple] = identity

        return tuple(dependencies[key] for key in sorted(dependencies))
    finally:
        if document is not None:
            document.close()
        workbook.close()


def dynamic_actual_hashes(
    root: Path,
    *,
    pdf_path: Path | None = None,
    dependencies: Mapping[str, Any] | None = None,
) -> dict[str, str]:
    """Hash dynamic actual evidence, optionally scoped to one PDF's dependencies.

    Global validation is always performed first.  With a PDF and dependency
    roster, occurrence overrides are filtered by PDF scope and learned exact
    glyph truths are filtered by hashes actually present in that PDF.  This
    preserves fail-closed fingerprints while avoiding whole-project cache
    invalidation after an unrelated correction.
    """
    report = validate_dynamic_actual_evidence(root)
    if not report.get("ok"):
        raise ValueError("動態 actual 證據檔驗證失敗：" + "；".join(report.get("errors") or []))
    if pdf_path is None:
        return dict(sorted((report.get("hashes") or {}).items()))

    root = Path(root)
    pdf_path = Path(pdf_path)
    override_rows = _read_csv(root / OCCURRENCE_OVERRIDE_FILE, OVERRIDE_HEADERS) if (root / OCCURRENCE_OVERRIDE_FILE).exists() else []
    override_rows = [row for row in override_rows if _override_applies_to_pdf(row, pdf_path)]

    glyf_rows = _read_csv(root / USER_GLYF_FILE, USER_GLYF_HEADERS)
    cff_rows = _read_csv(root / USER_CFF_FILE, USER_CFF_HEADERS)
    conflict_rows = _read_csv(root / GLYPH_CONFLICT_FILE, GLYPH_CONFLICT_HEADERS)
    mode = "global_fallback"
    if dependencies is not None:
        mode = "per_pdf_exact_dependency_v1"
        ttf_keys = {_text(x).lower() for x in (dependencies.get("ttf_glyph_sha256") or []) if _text(x)}
        cff_keys = {
            (_text(item[0]), _text(item[1]).lower())
            for item in (dependencies.get("cff_glyph_keys") or [])
            if isinstance(item, (list, tuple)) and len(item) >= 2
        }
        glyf_rows = [row for row in glyf_rows if _text(row.get("glyph_sha256")).lower() in ttf_keys]
        cff_rows = [
            row for row in cff_rows
            if (_text(row.get("style_group")), _text(row.get("glyph_sha256")).lower()) in cff_keys
        ]
        conflict_rows = [
            row for row in conflict_rows
            if (
                (_text(row.get("kind")) == "TTF_GLYF_SHA256" and _text(row.get("glyph_sha256")).lower() in ttf_keys)
                or (
                    _text(row.get("kind")) == "CFF_GLYPH_SHA256"
                    and (_text(row.get("style_group")), _text(row.get("glyph_sha256")).lower()) in cff_keys
                )
            )
        ]

    return {
        "scope_mode": hashlib.sha256(mode.encode("utf-8")).hexdigest(),
        OCCURRENCE_OVERRIDE_FILE: _subset_sha256(override_rows, OVERRIDE_HEADERS),
        USER_GLYF_FILE: _subset_sha256(glyf_rows, USER_GLYF_HEADERS),
        USER_CFF_FILE: _subset_sha256(cff_rows, USER_CFF_HEADERS),
        GLYPH_CONFLICT_FILE: _subset_sha256(conflict_rows, GLYPH_CONFLICT_HEADERS),
    }


def _sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _source(entry: Mapping[str, Any]) -> Mapping[str, Any]:
    return entry.get("source_record") or entry.get("row") or {}


def _cff_full_signature(source: Mapping[str, Any]) -> str:
    existing = _text(source.get("CFF完整注音簽名"))
    if existing:
        return existing
    style = _text(source.get("CFF樣式群組"))
    body = _text(source.get("CFF符號簽名"))
    tone = _text(source.get("CFF聲調簽名"))
    neutral = _text(source.get("CFF輕聲簽名"))
    if not style or not body:
        return ""
    return cff_full_annotation_signature(style, body.split(";"), tone, neutral)


def actual_group_identity(entry: Mapping[str, Any]) -> dict[str, str]:
    source = _source(entry)
    sha = _text(source.get("TTF字形SHA256")).lower()
    if len(sha) == 64 and all(ch in "0123456789abcdef" for ch in sha):
        return {
            **canonical_actual_group_identity(TTF_GLYF_SHA256, sha),
            "full_signature": "",
        }
    cff_glyph_sha = _text(source.get("CFF整字字形SHA256")).lower()
    cff_style = _text(source.get("CFF樣式群組"))
    if (
        cff_style
        and len(cff_glyph_sha) == 64
        and all(ch in "0123456789abcdef" for ch in cff_glyph_sha)
    ):
        return {
            **canonical_actual_group_identity(CFF_GLYPH_SHA256, cff_glyph_sha, cff_style),
            "full_signature": _cff_full_signature(source),
        }
    # Older workbooks may carry only an annotation signature.  That signature
    # is audit evidence, not a reusable learning key: different complete CFF
    # glyphs can share annotation components.  Keep them occurrence-specific
    # until a current decoder supplies the complete-glyph SHA-256.
    stable = _text(entry.get("stable_key") or source.get("穩定注音鍵"))
    if stable:
        return {
            **canonical_actual_group_identity(
                "SESSION_STABLE_KEY",
                _text(entry.get("occurrence_id") or stable),
            ),
            "full_signature": _cff_full_signature(source),
        }
    return {
        **canonical_actual_group_identity("OCCURRENCE_ONLY", _text(entry.get("occurrence_id"))),
        "full_signature": "",
    }


def build_actual_review_groups(entries: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    """Build GPT/manual actual groups from pending seeds, expanded to all exact peers.

    A verified exact glyph may already occur in non-pending rows of the current
    project.  Expanding each pending seed to every same exact identity ensures
    promotion writes occurrence-scoped overrides for all currently affected
    PDFs, so per-PDF fingerprints can safely invalidate only true dependants.
    """
    all_grouped: dict[tuple[str, str, str], list[dict[str, Any]]] = {}
    pending_keys: set[tuple[str, str, str]] = set()
    for raw in entries:
        entry = dict(raw)
        ident = actual_group_identity(entry)
        key = actual_group_match_key(ident)
        all_grouped.setdefault(key, []).append(entry)
        if _text(raw.get("state")) in ACTUAL_PENDING_STATES:
            pending_keys.add(key)
    result = []
    for kind, style_group, exact_key in sorted(pending_keys):
        members = all_grouped.get((kind, style_group, exact_key), [])
        if not members:
            continue
        members.sort(key=lambda e: (_text(e.get("pdf_name")), int(e.get("physical_page") or 0), float(e.get("y0") or 0), float(e.get("x0") or 0)))
        ident = actual_group_identity(members[0])
        identity_payload = actual_group_identity_payload(ident)
        group_id = actual_group_id(ident)
        snapshot = _hash_payload({
            "group_id": group_id,
            **identity_payload,
            "members": [
                {
                    "occurrence_id": e.get("occurrence_id"),
                    "review_id": e.get("review_id"),
                    "state": e.get("state"),
                    "actual": e.get("actual"),
                    "actual_evidence": e.get("actual_evidence"),
                }
                for e in members
            ],
        })
        result.append({
            "group_id": group_id,
            "group_snapshot": snapshot,
            "kind": kind,
            "exact_key": exact_key,
            "style_group": style_group,
            "full_signature": ident.get("full_signature", ""),
            "members": members,
            "occurrence_count": len(members),
        })
    result.sort(key=lambda g: (_text(g["members"][0].get("pdf_name")), int(g["members"][0].get("physical_page") or 0), _text(g["group_id"])))
    return result



def build_actual_group_for_entry(entries: Sequence[Mapping[str, Any]], target: Mapping[str, Any]) -> dict[str, Any]:
    """Build one review group for a user-flagged actual misread.

    Unlike the automatic pending queue, this may start from any ledger state.
    Exact TTF/CFF identities are grouped so two independent page examples can
    promote reusable truth; otherwise the correction remains occurrence-only.
    """
    ident = actual_group_identity(target)
    kind, exact_key = ident["kind"], ident["exact_key"]
    group_key = actual_group_match_key(ident)
    members = [dict(e) for e in entries if actual_group_match_key(actual_group_identity(e)) == group_key]
    if not members:
        members = [dict(target)]
    members.sort(key=lambda e: (_text(e.get("pdf_name")), int(e.get("physical_page") or 0), float(e.get("y0") or 0), float(e.get("x0") or 0)))
    identity_payload = actual_group_identity_payload(ident)
    group_id = actual_group_id(ident)
    snapshot = _hash_payload({
        "group_id": group_id,
        **identity_payload,
        "members": [
            {"occurrence_id": e.get("occurrence_id"), "review_id": e.get("review_id"), "state": e.get("state"), "actual": e.get("actual")}
            for e in members
        ],
    })
    return {
        "group_id": group_id, "group_snapshot": snapshot, "kind": kind, "exact_key": exact_key,
        "style_group": ident.get("style_group", ""), "full_signature": ident.get("full_signature", ""),
        "members": members, "occurrence_count": len(members),
    }

def _resolve_pdf_path(entry: Mapping[str, Any], output_dir: Path | None = None) -> Path:
    stored = Path(_text(entry.get("pdf")))
    if stored.exists():
        return stored
    name = _text(entry.get("pdf_name"))
    if output_dir and name:
        candidates = []
        for base in [Path(output_dir).parent, Path(output_dir)]:
            direct = base / name
            if direct.exists():
                candidates.append(direct)
            try:
                candidates.extend(p for p in base.rglob(name) if p.is_file())
            except Exception:
                pass
        unique = []
        seen = set()
        for p in candidates:
            rp = str(p.resolve())
            if rp not in seen:
                seen.add(rp); unique.append(p)
        if len(unique) == 1:
            return unique[0]
    raise FileNotFoundError(f"找不到現版 PDF：{name or stored}")


def render_occurrence_png(entry: Mapping[str, Any], output_path: Path, *, context: bool = False, output_dir: Path | None = None) -> Path:
    pdf = _resolve_pdf_path(entry, output_dir)
    page_number = int(entry.get("physical_page") or 0) - 1
    if page_number < 0:
        raise ValueError("缺少實體頁碼")
    x0, y0, x1, y1 = [float(entry.get(k) or 0) for k in ("x0", "y0", "x1", "y1")]
    if x1 <= x0 or y1 <= y0:
        raise ValueError("缺少可用 bbox")
    doc = fitz.open(pdf)
    try:
        page = doc[page_number]
        if context:
            px, py, scale = 95, 60, 3.0
        else:
            px, py, scale = 16, 14, 7.0
        clip = fitz.Rect(max(0, x0-px), max(0, y0-py), min(page.rect.width, x1+px), min(page.rect.height, y1+py))
        pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), clip=clip, alpha=False)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        pix.save(str(output_path))
    finally:
        doc.close()
    return output_path


def render_annotation_only_png(entry: Mapping[str, Any], output_path: Path, *, output_dir: Path | None = None) -> Path:
    """Render only the right-side Bopomofo region of one annotated glyph.

    The GPT actual-review path deliberately minimizes Chinese semantic context.
    The crop is visual evidence only; if a font layout falls outside this
    right-side convention the package still includes the whole-glyph tight crop
    as a second visual reference, never expected/dictionary data.
    """
    pdf = _resolve_pdf_path(entry, output_dir)
    page_number = int(entry.get("physical_page") or 0) - 1
    if page_number < 0:
        raise ValueError("缺少實體頁碼")
    x0, y0, x1, y1 = [float(entry.get(k) or 0) for k in ("x0", "y0", "x1", "y1")]
    if x1 <= x0 or y1 <= y0:
        raise ValueError("缺少可用 bbox")
    width = x1 - x0
    # In the supported textbook fonts, Bopomofo occupies the right ~45% of the
    # annotated glyph advance.  Keep a small overlap so the first symbol is not
    # clipped, but omit most of the Han character to reduce semantic priming.
    ax0 = x0 + width * 0.50
    pad_x, pad_y, scale = max(1.5, width * 0.04), 3.0, 10.0
    doc = fitz.open(pdf)
    try:
        page = doc[page_number]
        clip = fitz.Rect(max(0, ax0-pad_x), max(0, y0-pad_y), min(page.rect.width, x1+pad_x), min(page.rect.height, y1+pad_y))
        pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), clip=clip, alpha=False)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        pix.save(str(output_path))
    finally:
        doc.close()
    return output_path


def export_actual_review_package(
    output_dir: Path,
    ledger: Sequence[Mapping[str, Any]],
    *,
    version: str,
    session_id: str,
    session_schema_version: str,
    workbook_schema_version: str,
    review_id_schema_version: str,
) -> Path:
    output_dir = Path(output_dir)
    groups = build_actual_review_groups(ledger)
    package_dir = output_dir / "actual待判定_GPT包"
    images_dir = package_dir / "images"
    if package_dir.exists():
        shutil.rmtree(package_dir)
    images_dir.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    guide = wb.active
    guide.title = "使用說明"
    instructions = [
        ("用途", "只判定 PDF 原頁實際可見注音 actual；不得使用 expected、字典、中文字讀音常識反推。"),
        ("判定方式", "先看 annotation-only 圖直接讀注音；同組有2筆以上時，VERIFIED 必須同時查看樣本B。兩張不一致就填 UNRESOLVED。whole-glyph 圖僅用於確認裁切位置，不得用中文字讀音反推。"),
        ("聲調", "輕聲點、二三四聲都必須直接從圖片確認。聲調前後位置可等價輸入，程式會 canonicalize。"),
        ("匯入安全", "匯入採交易驗證；group/session/snapshot 不一致、非法注音、未勾樣本等任何錯誤會整批拒絕。"),
        ("expected 隔離", "本包不輸出 expected 欄位，也不允許匯入修改 expected。"),
        ("待判定 unique groups", len(groups)),
        ("待判定 occurrences", sum(g["occurrence_count"] for g in groups)),
    ]
    guide.append(["項目", "說明"])
    for row in instructions:
        guide.append(row)

    meta = wb.create_sheet("匯入中繼資料")
    meta.append(["項目", "內容"])
    for row in [
        ("version", version),
        ("session_id", session_id),
        ("session_schema_version", session_schema_version),
        ("workbook_schema_version", workbook_schema_version),
        ("review_id_schema_version", review_id_schema_version),
        ("actual_review_schema_version", ACTUAL_REVIEW_SCHEMA_VERSION),
        ("exported_group_count", len(groups)),
    ]:
        meta.append(row)

    ws = wb.create_sheet("actual待判定")
    headers = [
        "編號", "group_id", "group_snapshot", "group_kind", "exact_key", "occurrence_count",
        "sample_a_image", "sample_b_image", "sample_a_occurrence_id", "sample_b_occurrence_id",
        "sample_a_pdf", "sample_a_page", "sample_b_pdf", "sample_b_page",
        "current_actual", "actual_evidence", "decision", "actual_reading", "confidence",
        "sample_a_checked", "sample_b_checked", "note",
    ]
    ws.append(headers)

    for index, group in enumerate(groups, 1):
        members = group["members"]
        a = members[0]
        b = members[1] if len(members) > 1 else None
        image_names = []
        for label, entry in (("a", a), ("b", b)):
            if entry is None:
                image_names.append("")
                continue
            annotation = images_dir / f"{group['group_id']}_{label}_annotation.png"
            tight = images_dir / f"{group['group_id']}_{label}_whole_glyph.png"
            render_annotation_only_png(entry, annotation, output_dir=output_dir)
            render_occurrence_png(entry, tight, context=False, output_dir=output_dir)
            image_names.append(f"images/{annotation.name} | images/{tight.name}")
        ws.append([
            index, group["group_id"], group["group_snapshot"], group["kind"], group["exact_key"], group["occurrence_count"],
            image_names[0], image_names[1], a.get("occurrence_id", ""), b.get("occurrence_id", "") if b else "",
            a.get("pdf_name", ""), a.get("printed_page", ""),
            b.get("pdf_name", "") if b else "", b.get("printed_page", "") if b else "",
            a.get("actual", "") or "", a.get("actual_evidence", "") or "", "", "", "", "", "", "",
        ])

    decision_col = headers.index("decision") + 1
    confidence_col = headers.index("confidence") + 1
    a_checked_col = headers.index("sample_a_checked") + 1
    b_checked_col = headers.index("sample_b_checked") + 1
    for col, values in [
        (decision_col, '"VERIFIED,UNRESOLVED"'),
        (confidence_col, '"高,中,低"'),
        (a_checked_col, '"Y,N"'),
        (b_checked_col, '"Y,N"'),
    ]:
        dv = DataValidation(type="list", formula1=values, allow_blank=True)
        ws.add_data_validation(dv)
        if ws.max_row >= 2:
            letter = get_column_letter(col)
            dv.add(f"{letter}2:{letter}{ws.max_row}")

    meta.sheet_state = "hidden"
    for col_name in ("group_snapshot", "exact_key", "sample_a_occurrence_id", "sample_b_occurrence_id"):
        ws.column_dimensions[get_column_letter(headers.index(col_name) + 1)].hidden = True
    for sheet in wb.worksheets:
        _style_sheet(sheet)
    xlsx = package_dir / "actual待判定_給GPT.xlsx"
    wb.save(xlsx)

    zip_path = output_dir / "actual待判定_GPT包.zip"
    if zip_path.exists():
        zip_path.unlink()
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for path in package_dir.rglob("*"):
            if path.is_file():
                zf.write(path, path.relative_to(package_dir))
    return zip_path


def _load_sheet_rows(path: Path, sheet: str) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "匯入中繼資料" not in wb.sheetnames or sheet not in wb.sheetnames:
            raise ValueError("actual GPT 報表缺少必要工作表")
        meta_ws = wb["匯入中繼資料"]
        meta = {}
        for row in meta_ws.iter_rows(values_only=True):
            if row and row[0] not in (None, "", "項目"):
                meta[_text(row[0])] = row[1] if len(row) > 1 else ""
        ws = wb[sheet]
        iterator = ws.iter_rows(values_only=True)
        try:
            headers = [_text(x) for x in next(iterator)]
        except StopIteration:
            headers = []
        if len(headers) != len(set(headers)) or any(not h for h in headers):
            raise ValueError("actual GPT 報表 header 無效或重複")
        rows = [{headers[i]: row[i] if i < len(row) else None for i in range(len(headers))} for row in iterator]
        return meta, rows
    finally:
        wb.close()


def load_glyph_truth_quarantine(root: Path) -> dict[str, set[Any]]:
    """Return exact glyph identities that are forbidden from global reuse.

    Quarantine is behavioral evidence: the exact SHA has contradictory direct
    visual readings.  It blocks reusable SHA bridging but never invents an
    occurrence reading.
    """
    root = Path(root)
    ensure_user_evidence_files(root)
    rows = _read_csv(root / GLYPH_CONFLICT_FILE, GLYPH_CONFLICT_HEADERS)
    ttf: set[str] = set()
    cff: set[tuple[str, str]] = set()
    for row in rows:
        if _text(row.get("status")) != "GLYPH_TRUTH_CONFLICT":
            continue
        sha = _text(row.get("glyph_sha256")).lower()
        if _text(row.get("kind")) == "TTF_GLYF_SHA256":
            ttf.add(sha)
        elif _text(row.get("kind")) == "CFF_GLYPH_SHA256":
            cff.add((_text(row.get("style_group")), sha))
    return {"ttf": ttf, "cff": cff}


def _learning_location(root: Path, group: Mapping[str, Any]):
    kind = _text(group.get("kind"))
    exact_key = _text(group.get("exact_key")).lower()
    if kind == "TTF_GLYF_SHA256":
        return Path(root) / USER_GLYF_FILE, USER_GLYF_HEADERS, ("glyph_sha256",), {"glyph_sha256": exact_key}
    if kind == "CFF_GLYPH_SHA256":
        return (
            Path(root) / USER_CFF_FILE,
            USER_CFF_HEADERS,
            ("style_group", "glyph_sha256"),
            {
                "style_group": _text(group.get("style_group")),
                "glyph_sha256": exact_key,
                "full_signature": _text(group.get("full_signature")),
            },
        )
    return None


def _current_learning_record(root: Path, group: Mapping[str, Any]) -> dict[str, str] | None:
    info = _learning_location(root, group)
    if not info:
        return None
    path, headers, key_fields, base = info
    key = tuple(base.get(k, "") for k in key_fields)
    for row in _read_csv(path, headers):
        if tuple(row.get(k, "") for k in key_fields) == key:
            return row
    return None


def _static_ttf_reading(group: Mapping[str, Any]) -> str:
    if _text(group.get("kind")) != "TTF_GLYF_SHA256":
        return ""
    sha = _text(group.get("exact_key")).lower()
    path = Path(__file__).resolve().parent / STATIC_TTF_FINGERPRINT_FILE
    if not path.exists():
        return ""
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh):
            if _text(row.get("glyph_sha256")).lower() == sha:
                return _canon(row.get("bopomofo"))
    return ""


def _append_provenance_event(
    root: Path,
    group: Mapping[str, Any],
    reading: str,
    event_type: str,
    occurrence_ids: Sequence[str],
    *,
    source: str,
    note: str = "",
) -> None:
    path = Path(root) / GLYPH_PROVENANCE_FILE
    rows = _read_csv(path, GLYPH_PROVENANCE_HEADERS)
    ids = sorted({_text(x) for x in occurrence_ids if _text(x)})
    payload = {
        "kind": _text(group.get("kind")),
        "style_group": _text(group.get("style_group")),
        "glyph_sha256": _text(group.get("exact_key")).lower(),
        "bopomofo": _canon(reading),
        "event_type": _text(event_type),
        "occurrence_ids": "|".join(ids),
        "source": _text(source),
        "note": _text(note),
    }
    event_id = "gpe_" + _hash_payload(payload)[:24]
    if any(_text(row.get("event_id")) == event_id for row in rows):
        return
    rows.append({
        "event_id": event_id,
        **payload,
        "created_at": datetime.now().isoformat(timespec="seconds"),
    })
    rows.sort(key=lambda r: (_text(r.get("created_at")), _text(r.get("event_id"))))
    _write_csv(path, GLYPH_PROVENANCE_HEADERS, rows)


def _upsert_conflict(
    root: Path,
    group: Mapping[str, Any],
    readings: Sequence[str],
    occurrence_ids: Sequence[str],
    *,
    note: str = "",
) -> None:
    path = Path(root) / GLYPH_CONFLICT_FILE
    rows = _read_csv(path, GLYPH_CONFLICT_HEADERS)
    kind = _text(group.get("kind"))
    style = _text(group.get("style_group"))
    sha = _text(group.get("exact_key")).lower()
    key = (kind, style, sha)
    existing = next((row for row in rows if (_text(row.get("kind")), _text(row.get("style_group")), _text(row.get("glyph_sha256")).lower()) == key), None)
    all_readings = {_canon(x) for x in readings if _canon(x)}
    all_ids = {_text(x) for x in occurrence_ids if _text(x)}
    if existing:
        all_readings.update(_canon(x) for x in _text(existing.get("readings")).split("|") if _canon(x))
        all_ids.update(_text(x) for x in _text(existing.get("source_occurrence_ids")).split("|") if _text(x))
    if len(all_readings) < 2:
        raise ValueError("GLYPH_TRUTH_CONFLICT 必須保留至少兩個互斥讀音")
    new_row = {
        "kind": kind,
        "style_group": style,
        "glyph_sha256": sha,
        "status": "GLYPH_TRUTH_CONFLICT",
        "readings": "|".join(sorted(all_readings)),
        "source_occurrence_ids": "|".join(sorted(all_ids)),
        "updated_at": datetime.now().isoformat(timespec="seconds"),
        "notes": _text(note),
    }
    if existing:
        comparable_old = {k: _text(existing.get(k)) for k in GLYPH_CONFLICT_HEADERS if k != "updated_at"}
        comparable_new = {k: _text(new_row.get(k)) for k in GLYPH_CONFLICT_HEADERS if k != "updated_at"}
        if comparable_old == comparable_new:
            return
        existing.update(new_row)
    else:
        rows.append(new_row)
    rows.sort(key=lambda r: (_text(r.get("kind")), _text(r.get("style_group")), _text(r.get("glyph_sha256"))))
    _write_csv(path, GLYPH_CONFLICT_HEADERS, rows)


def _demote_learning_record_to_conflict(root: Path, group: Mapping[str, Any], conflict_note: str) -> None:
    info = _learning_location(root, group)
    if not info:
        return
    path, headers, key_fields, base = info
    rows = _read_csv(path, headers)
    key = tuple(base.get(k, "") for k in key_fields)
    changed = False
    for row in rows:
        if tuple(row.get(k, "") for k in key_fields) != key:
            continue
        if _text(row.get("verification_level")) != "QUARANTINED_CONFLICT":
            row["verification_level"] = "QUARANTINED_CONFLICT"
            row["updated_at"] = datetime.now().isoformat(timespec="seconds")
            row["notes"] = (_text(row.get("notes")) + "；" + _text(conflict_note)).strip("；")
            changed = True
    if changed:
        _write_csv(path, headers, rows)


def _override_key_from_entry(entry: Mapping[str, Any]) -> tuple[str, ...]:
    return (
        Path(_text(entry.get("pdf_name"))).stem,
        "",
        _text(entry.get("printed_page") or entry.get("physical_page")),
        _text(entry.get("char")),
        _text(entry.get("stable_key") or _source(entry).get("穩定注音鍵")),
        _text(entry.get("x0")),
        _text(entry.get("y0")),
    )


def _remove_unverified_group_overrides(
    root: Path,
    members: Sequence[Mapping[str, Any]],
    preserve_occurrence_ids: set[str],
) -> list[str]:
    """Remove propagated local overrides after an exact-glyph conflict.

    Only occurrence IDs that were themselves direct visual examples are kept.
    This prevents an earlier two-sample global promotion from leaving every
    peer occurrence pinned to the contaminated reading.
    """
    path = Path(root) / OCCURRENCE_OVERRIDE_FILE
    if not path.exists():
        return []
    rows = _read_csv(path, OVERRIDE_HEADERS)
    member_by_key = {_override_key_from_entry(entry): _text(entry.get("occurrence_id")) for entry in members}
    kept = []
    removed_ids = []
    for row in rows:
        key = tuple(row.get(h, "") for h in ("pdf_contains", "pdf_excludes", "page", "target_char", "stable_key", "x0", "y0"))
        oid = member_by_key.get(key)
        if oid and oid not in preserve_occurrence_ids:
            removed_ids.append(oid)
            continue
        kept.append(row)
    if len(kept) != len(rows):
        _write_csv(path, OVERRIDE_HEADERS, kept)
    return sorted(set(removed_ids))


def _append_or_update_overrides(root: Path, entries: Sequence[Mapping[str, Any]], reading: str, *, source: str, note: str) -> None:
    path = Path(root) / OCCURRENCE_OVERRIDE_FILE
    rows = _read_csv(path, OVERRIDE_HEADERS)
    by_key = {
        tuple(row[h] for h in ("pdf_contains", "pdf_excludes", "page", "target_char", "stable_key", "x0", "y0")): row
        for row in rows
    }
    for entry in entries:
        pdf_name = Path(_text(entry.get("pdf_name"))).stem
        key_values = {
            "pdf_contains": pdf_name,
            "pdf_excludes": "",
            "page": _text(entry.get("printed_page") or entry.get("physical_page")),
            "target_char": _text(entry.get("char")),
            "stable_key": _text(entry.get("stable_key") or _source(entry).get("穩定注音鍵")),
            "x0": _text(entry.get("x0")),
            "y0": _text(entry.get("y0")),
        }
        key = tuple(key_values[h] for h in ("pdf_contains", "pdf_excludes", "page", "target_char", "stable_key", "x0", "y0"))
        new_row = {**key_values, "actual_reading": reading, "source": source, "note": note}
        old = by_key.get(key)
        if old and _canon(old.get("actual_reading")) not in {"", reading}:
            # This function is called only after explicit visual verification.
            # Replacing a stale occurrence override is therefore allowed, but
            # the superseded value remains in the audit note instead of being
            # silently forgotten.
            old_reading = _canon(old.get("actual_reading")) or _text(old.get("actual_reading"))
            audit = f"覆寫舊 occurrence actual={old_reading}；舊來源={_text(old.get('source'))}"
            new_row["note"] = (note + "；" + audit).strip("；")
        by_key[key] = new_row
    ordered = sorted(by_key.values(), key=lambda r: (r["pdf_contains"], r["page"], r["target_char"], r["stable_key"], r["y0"], r["x0"]))
    _write_csv(path, OVERRIDE_HEADERS, ordered)


def _update_learning_file(root: Path, group: Mapping[str, Any], reading: str, verified_entries: Sequence[Mapping[str, Any]], note: str) -> dict[str, Any]:
    kind = _text(group.get("kind"))
    exact_key = _text(group.get("exact_key")).lower()
    info = _learning_location(root, group)
    if not info or not exact_key:
        return {"learning_level": "OCCURRENCE_ONLY", "conflict": False, "old_verified_ids": set(), "known_readings": set()}
    path, headers, key_fields, base = info
    rows = _read_csv(path, headers)
    key = tuple(base.get(k, "") for k in key_fields)
    existing = next((row for row in rows if tuple(row.get(k, "") for k in key_fields) == key), None)
    examples = {_text(x.get("occurrence_id")) for x in verified_entries if _text(x.get("occurrence_id"))}
    old_verified_ids = set()
    known_readings = set()
    if existing:
        old_reading = _canon(existing.get("bopomofo"))
        if old_reading:
            known_readings.add(old_reading)
        old_verified_ids.update(filter(None, (_text(x) for x in _text(existing.get("source_examples")).split("|"))))
    static_reading = _static_ttf_reading(group)
    if static_reading:
        known_readings.add(static_reading)
    quarantine = load_glyph_truth_quarantine(root)
    already_quarantined = (
        (kind == "TTF_GLYF_SHA256" and exact_key in quarantine["ttf"])
        or (kind == "CFF_GLYPH_SHA256" and (_text(group.get("style_group")), exact_key) in quarantine["cff"])
    )
    conflicting = {r for r in known_readings if r and r != reading}
    if conflicting or already_quarantined:
        all_readings = set(known_readings) | {reading}
        # If the row was already quarantined, the registry itself may contain
        # the second reading even when the demoted learning row carries only one.
        conflict_rows = _read_csv(Path(root) / GLYPH_CONFLICT_FILE, GLYPH_CONFLICT_HEADERS)
        for crow in conflict_rows:
            same = _text(crow.get("kind")) == kind and _text(crow.get("glyph_sha256")).lower() == exact_key
            if kind == "CFF_GLYPH_SHA256":
                same = same and _text(crow.get("style_group")) == _text(group.get("style_group"))
            if same:
                all_readings.update(_canon(x) for x in _text(crow.get("readings")).split("|") if _canon(x))
                old_verified_ids.update(filter(None, (_text(x) for x in _text(crow.get("source_occurrence_ids")).split("|"))))
        conflict_note = (
            f"exact glyph 出現互斥視覺讀音：{','.join(sorted(all_readings))}；"
            "已停止全域 reusable truth，保留 occurrence-local 證據"
        )
        _upsert_conflict(root, group, sorted(all_readings), sorted(old_verified_ids | examples), note=(note + "；" + conflict_note).strip("；"))
        _demote_learning_record_to_conflict(root, group, conflict_note)
        _append_provenance_event(root, group, reading, "GLYPH_TRUTH_CONFLICT", sorted(examples), source="visual actual review", note=note or conflict_note)
        return {
            "learning_level": "GLYPH_TRUTH_CONFLICT",
            "conflict": True,
            "old_verified_ids": old_verified_ids,
            "known_readings": all_readings,
        }

    examples.update(old_verified_ids)
    source_count = len(examples)
    level = "VERIFIED_EXACT_GLYPH" if source_count >= 2 else "USER_VERIFIED_SINGLE"
    new_row = {
        **base,
        "bopomofo": reading,
        "verification_level": level,
        "source_count": str(source_count),
        "source_examples": "|".join(sorted(examples)),
        "updated_at": datetime.now().isoformat(timespec="seconds"),
        "notes": note,
    }
    if existing:
        existing.update(new_row)
    else:
        rows.append(new_row)
    rows.sort(key=lambda r: tuple(r.get(k, "") for k in key_fields))
    _write_csv(path, headers, rows)
    _append_provenance_event(root, group, reading, "PROMOTED" if level == "VERIFIED_EXACT_GLYPH" else "SINGLE_VISUAL_EVIDENCE", sorted(examples), source="visual actual review", note=note)
    return {"learning_level": level, "conflict": False, "old_verified_ids": old_verified_ids, "known_readings": {reading}}


def apply_verified_actual_group(
    root: Path,
    group: Mapping[str, Any],
    reading: str,
    *,
    checked_occurrence_ids: Sequence[str] | None = None,
    source: str,
    note: str = "",
) -> dict[str, Any]:
    """Apply direct visual actual evidence without coupling it to global learning.

    v5.5.1 rule: occurrence-local truth is authoritative for the checked printed
    positions even when reusable exact-glyph promotion conflicts.  A promotion
    conflict quarantines the SHA and reopens unverified peers instead of rolling
    back the local correction.
    """
    root = Path(root)
    ensure_user_evidence_files(root)
    reading = _canon(reading)
    if not reading:
        raise ValueError("實際注音不是合法單一注音")
    members = list(group.get("members") or [])
    checked = set(checked_occurrence_ids or [])
    if not checked:
        checked = {_text(members[0].get("occurrence_id"))} if members else set()
    verified_entries = [m for m in members if _text(m.get("occurrence_id")) in checked]
    if not verified_entries:
        raise ValueError("沒有任何已核對 occurrence")

    learning = _update_learning_file(root, group, reading, verified_entries, note)
    conflict = bool(learning.get("conflict"))
    verified_ids = {_text(e.get("occurrence_id")) for e in verified_entries if _text(e.get("occurrence_id"))}
    if conflict:
        preserve_ids = set(learning.get("old_verified_ids") or set()) | verified_ids
        removed_ids = _remove_unverified_group_overrides(root, members, preserve_ids)
        target_entries = verified_entries
        propagate_all = False
    else:
        removed_ids = []
        propagate_all = len(members) <= 1 or len(verified_entries) >= 2
        target_entries = members if propagate_all else verified_entries

    # Local visual correction is committed regardless of reusable promotion.
    _append_or_update_overrides(root, target_entries, reading, source=source, note=note)
    _append_provenance_event(
        root, group, reading, "OCCURRENCE_OVERRIDE",
        [_text(e.get("occurrence_id")) for e in target_entries], source=source, note=note,
    )
    affected = sorted({_text(e.get("occurrence_id")) for e in members if _text(e.get("occurrence_id"))} if conflict else {_text(e.get("occurrence_id")) for e in target_entries if _text(e.get("occurrence_id"))})
    return {
        "reading": reading,
        "target_occurrence_ids": [_text(e.get("occurrence_id")) for e in target_entries],
        "affected_occurrence_ids": affected,
        "verified_occurrence_ids": sorted(verified_ids),
        "learning_level": learning.get("learning_level", "OCCURRENCE_ONLY"),
        "propagated_to_group": propagate_all,
        "glyph_truth_conflict": conflict,
        "quarantined": conflict,
        "quarantine_key": _text(group.get("exact_key")) if conflict else "",
        "reopened_occurrence_ids": removed_ids,
        "known_conflicting_readings": sorted(learning.get("known_readings") or []),
    }


def _index_live_manual_actual_groups(
    live_groups: Sequence[Mapping[str, Any]],
) -> dict[str, Mapping[str, Any]]:
    if isinstance(live_groups, (str, bytes)) or not isinstance(live_groups, Sequence):
        raise ValueError("manual actual batch live_groups 必須是 sequence")
    by_id: dict[str, Mapping[str, Any]] = {}
    for index, group in enumerate(live_groups):
        if not isinstance(group, Mapping):
            raise ValueError(f"manual actual batch live_groups[{index}] 必須是 mapping")
        group_id = _staging_text(group.get("group_id"), f"live_groups[{index}].group_id")
        if group_id in by_id:
            raise ValueError(f"manual actual batch live group_id 重複：{group_id}")
        by_id[group_id] = group
    return by_id


def _revalidate_staged_manual_actual_group(
    staged: Mapping[str, Any],
    live_group: Mapping[str, Any],
) -> dict[str, Any]:
    group_id = staged["group_id"]
    comparisons = {
        "group_id": _staging_text(live_group.get("group_id"), f"live.{group_id}.group_id"),
        "group_snapshot": _staging_text(
            live_group.get("group_snapshot"), f"live.{group_id}.group_snapshot"
        ),
        "kind": _staging_text(live_group.get("kind"), f"live.{group_id}.kind"),
        "exact_key": _staging_text(live_group.get("exact_key"), f"live.{group_id}.exact_key"),
        "style_group": _staging_text(
            live_group.get("style_group", ""),
            f"live.{group_id}.style_group",
            allow_empty=True,
        ),
    }
    for field, observed in comparisons.items():
        if observed != staged[field]:
            raise ValueError(
                f"manual actual batch live group identity 已變更：{group_id} field={field}；整批拒絕"
            )
    try:
        staged_identity = canonical_actual_group_identity(
            staged.get("kind"), staged.get("exact_key"), staged.get("style_group", "")
        )
        live_identity = canonical_actual_group_identity(
            comparisons["kind"], comparisons["exact_key"], comparisons["style_group"]
        )
    except ValueError as exc:
        raise ValueError(
            f"manual actual batch live group exact identity 無效：{group_id}；整批拒絕"
        ) from exc
    if actual_group_id(live_identity) != comparisons["group_id"]:
        raise ValueError(
            f"manual actual batch live group_id 未綁定 canonical exact identity：{group_id}；整批拒絕"
        )
    if actual_group_match_key(staged_identity) != actual_group_match_key(live_identity):
        raise ValueError(
            f"manual actual batch live group canonical exact identity 已變更：{group_id}；整批拒絕"
        )

    members = live_group.get("members")
    if not isinstance(members, list) or not members:
        raise ValueError(f"manual actual batch live group members 無效：{group_id}")
    member_ids: list[str] = []
    for index, member in enumerate(members):
        if not isinstance(member, Mapping):
            raise ValueError(f"manual actual batch live group member 無效：{group_id}[{index}]")
        member_ids.append(
            _staging_text(
                member.get("occurrence_id"),
                f"live.{group_id}.members[{index}].occurrence_id",
            )
        )
        try:
            member_key = actual_group_match_key(actual_group_identity(member))
        except ValueError as exc:
            raise ValueError(
                f"manual actual batch live group member exact identity 無效：{group_id}[{index}]；整批拒絕"
            ) from exc
        if member_key != actual_group_match_key(live_identity):
            raise ValueError(
                f"manual actual batch live group member 不屬於 canonical exact identity："
                f"{group_id}[{index}]；整批拒絕"
            )
    if len(member_ids) != len(set(member_ids)):
        raise ValueError(f"manual actual batch live group members 含重複 occurrence_id：{group_id}")
    invalid_checked = sorted(set(staged["checked_occurrence_ids"]) - set(member_ids))
    if invalid_checked:
        raise ValueError(
            f"manual actual batch checked occurrence 已失效：{group_id} ids={invalid_checked}；整批拒絕"
        )
    if member_ids != staged["member_occurrence_ids"]:
        raise ValueError(f"manual actual batch live group member roster 已變更：{group_id}；整批拒絕")

    occurrence_count = live_group.get("occurrence_count")
    if isinstance(occurrence_count, bool):
        raise ValueError(f"manual actual batch live group occurrence_count 無效：{group_id}")
    try:
        observed_count = int(occurrence_count)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"manual actual batch live group occurrence_count 無效：{group_id}") from exc
    if observed_count != len(member_ids):
        raise ValueError(f"manual actual batch live group occurrence_count 與 members 不一致：{group_id}")

    return dict(live_group)


def _empty_manual_actual_batch_result() -> dict[str, Any]:
    return {
        "applied_group_count": 0,
        "group_results": [],
        "affected_occurrence_ids": [],
        "reopened_occurrence_ids": [],
        "quarantined_group_ids": [],
        "affected_pdf_names": [],
        "staging_remaining_count": 0,
    }


def apply_staged_manual_actual_batch(
    root: Path,
    live_groups: Sequence[Mapping[str, Any]],
) -> dict[str, Any]:
    """Apply one frozen staged-manual-actual queue as a transaction.

    Pending staging is revalidated against authoritative live review groups
    before any dynamic actual evidence mutation.  Every selected decision is
    applied through ``apply_verified_actual_group`` in sorted ``group_id``
    order.  The five authoritative files are restored byte-for-byte on any
    late exception or stale final acknowledgement; staging is acknowledged in
    one compare-and-remove rewrite only after all group applications succeed.
    """
    root = Path(root)
    staging = load_manual_actual_staging(root)
    frozen = sorted(staging["staged_groups"], key=lambda item: item["group_id"])
    if not frozen:
        return _empty_manual_actual_batch_result()

    live_by_id = _index_live_manual_actual_groups(live_groups)
    selected: list[tuple[dict[str, Any], dict[str, Any]]] = []
    occurrence_pdf_names: dict[str, str] = {}
    for staged in frozen:
        live = live_by_id.get(staged["group_id"])
        if live is None:
            raise ValueError(
                f"manual actual batch 找不到 live group：{staged['group_id']}；整批拒絕"
            )
        validated_live = _revalidate_staged_manual_actual_group(staged, live)
        selected.append((staged, validated_live))
        for member in validated_live["members"]:
            occurrence_id = _text(member.get("occurrence_id"))
            pdf_name = _text(member.get("pdf_name"))
            if occurrence_id and pdf_name:
                occurrence_pdf_names[occurrence_id] = pdf_name

    tokens = _manual_actual_staging_ack_tokens(frozen)
    current_staging = load_manual_actual_staging(root)
    _validate_manual_actual_staging_ack_tokens(current_staging, tokens)

    backups = _snapshot_dynamic_actual_evidence(root)
    group_results: list[dict[str, Any]] = []
    affected_occurrence_ids: set[str] = set()
    reopened_occurrence_ids: set[str] = set()
    quarantined_group_ids: set[str] = set()
    try:
        for staged, live_group in selected:
            applied = apply_verified_actual_group(
                root,
                live_group,
                staged["reading"],
                checked_occurrence_ids=staged["checked_occurrence_ids"],
                source=staged["source"],
                note=staged["note"],
            )
            if not isinstance(applied, Mapping):
                raise ValueError(
                    f"manual actual batch apply result 無效：{staged['group_id']}"
                )
            item = {
                **dict(applied),
                "group_id": staged["group_id"],
                "group_snapshot": staged["group_snapshot"],
            }
            group_results.append(item)
            affected_occurrence_ids.update(
                _text(value) for value in (applied.get("affected_occurrence_ids") or []) if _text(value)
            )
            reopened_occurrence_ids.update(
                _text(value) for value in (applied.get("reopened_occurrence_ids") or []) if _text(value)
            )
            if applied.get("quarantined") or applied.get("glyph_truth_conflict"):
                quarantined_group_ids.add(staged["group_id"])

        acknowledgement = remove_staged_manual_actual_groups(root, frozen)
    except Exception:
        _restore_dynamic_actual_evidence(backups)
        raise

    affected_ids = sorted(affected_occurrence_ids)
    return {
        "applied_group_count": len(group_results),
        "group_results": group_results,
        "affected_occurrence_ids": affected_ids,
        "reopened_occurrence_ids": sorted(reopened_occurrence_ids),
        "quarantined_group_ids": sorted(quarantined_group_ids),
        "affected_pdf_names": sorted({
            occurrence_pdf_names[occurrence_id]
            for occurrence_id in affected_ids
            if occurrence_id in occurrence_pdf_names
        }),
        "staging_remaining_count": acknowledgement["staging_remaining_count"],
    }


def import_actual_review_workbook(
    root: Path,
    xlsx: Path,
    groups: Sequence[Mapping[str, Any]],
    *,
    expected_metadata: Mapping[str, Any],
) -> dict[str, Any]:
    meta, rows = _load_sheet_rows(Path(xlsx), "actual待判定")
    required_meta = dict(expected_metadata)
    required_meta["actual_review_schema_version"] = ACTUAL_REVIEW_SCHEMA_VERSION
    mismatched = {}
    for key, required in required_meta.items():
        observed = meta.get(key)
        if key == "version":
            # Application release is audit-only from v5.6 onward.
            continue
        if key in {"session_schema_version", "workbook_schema_version"}:
            if not schema_compatible(observed, required):
                mismatched[key] = (required, observed)
            continue
        if key == "review_id_schema_version":
            if not review_id_schema_compatible(observed, required):
                mismatched[key] = (required, observed)
            continue
        if str(observed) != str(required):
            mismatched[key] = (required, observed)
    if mismatched:
        raise ValueError(f"actual GPT workbook session/schema 不相容：{mismatched}")
    required_columns = {
        "group_id", "group_snapshot", "decision", "actual_reading", "confidence",
        "sample_a_checked", "sample_b_checked", "note",
    }
    if rows:
        missing = sorted(required_columns - set(rows[0]))
        if missing:
            raise ValueError(f"actual GPT workbook 缺少欄位：{missing}")
    by_id = {g["group_id"]: g for g in groups}
    seen = set()
    staged = []
    errors = []
    for row_no, row in enumerate(rows, 2):
        gid = _text(row.get("group_id"))
        decision = _text(row.get("decision")).upper()
        if not gid and not decision:
            continue
        if gid in seen:
            errors.append(f"row {row_no}: group_id 重複 {gid}")
            continue
        seen.add(gid)
        group = by_id.get(gid)
        if not group:
            errors.append(f"row {row_no}: unknown group_id {gid}")
            continue
        if _text(row.get("group_snapshot")) != _text(group.get("group_snapshot")):
            errors.append(f"row {row_no}: stale/modified group snapshot")
            continue
        if decision not in {"", "VERIFIED", "UNRESOLVED"}:
            errors.append(f"row {row_no}: decision 必須是 VERIFIED/UNRESOLVED")
            continue
        if decision != "VERIFIED":
            continue
        reading = _canon(row.get("actual_reading"))
        if not reading:
            errors.append(f"row {row_no}: VERIFIED 缺少合法 actual_reading")
            continue
        a_ok = _text(row.get("sample_a_checked")).upper() == "Y"
        b_ok = _text(row.get("sample_b_checked")).upper() == "Y"
        if not a_ok:
            errors.append(f"row {row_no}: VERIFIED 必須 sample_a_checked=Y")
        if int(group.get("occurrence_count") or 0) >= 2 and not b_ok:
            errors.append(f"row {row_no}: 同組多 occurrence 時 VERIFIED 必須 sample_b_checked=Y")
        conf = _text(row.get("confidence"))
        if conf not in {"高", "中"}:
            errors.append(f"row {row_no}: VERIFIED confidence 只能是高/中")
        checked_ids = [_text(group["members"][0].get("occurrence_id"))]
        if len(group["members"]) > 1:
            checked_ids.append(_text(group["members"][1].get("occurrence_id")))
        staged.append((group, reading, checked_ids, _text(row.get("note"))))
    if errors:
        raise ValueError("actual GPT 匯入整批拒絕；未寫入任何一列：\n" + "\n".join(errors[:100]))

    # v5.5.1: reusable glyph-truth conflicts are no longer batch-fatal.
    # Occurrence-local visual corrections must still commit; the promotion layer
    # will quarantine contradictory exact SHA evidence during apply.
    ensure_user_evidence_files(root)

    # Transactional write: if any late occurrence-level validation fails,
    # restore all dynamic actual evidence files byte-for-byte.
    backups = _snapshot_dynamic_actual_evidence(root)
    results = []
    try:
        for group, reading, ids, note in staged:
            results.append(apply_verified_actual_group(
                root, group, reading, checked_occurrence_ids=ids,
                source="GPT actual 視覺證據匯入", note=note,
            ))
    except Exception:
        _restore_dynamic_actual_evidence(backups)
        raise
    return {"imported_groups": len(results), "results": results}


def load_user_verified_glyf(path: Path) -> dict[str, dict[str, str]]:
    rows = _read_csv(Path(path), USER_GLYF_HEADERS)
    out = {}
    for row in rows:
        if row.get("verification_level") != "VERIFIED_EXACT_GLYPH":
            continue
        sha = _text(row.get("glyph_sha256")).lower()
        bop = _canon(row.get("bopomofo"))
        if sha and bop:
            out[sha] = {"bopomofo": bop, "verification": row.get("verification_level", ""), "notes": row.get("notes", ""), "source_font": "USER-VERIFIED-GLYF", "source_gid": -1}
    return out


def load_user_verified_cff(path: Path) -> dict[tuple[str, str], dict[str, str]]:
    rows = _read_csv(Path(path), USER_CFF_HEADERS)
    out = {}
    for row in rows:
        if row.get("verification_level") != "VERIFIED_EXACT_GLYPH":
            continue
        style = _text(row.get("style_group"))
        sha = _text(row.get("glyph_sha256")).lower()
        bop = _canon(row.get("bopomofo"))
        if style and len(sha) == 64 and bop:
            out[(style, sha)] = {
                "bopomofo": bop,
                "verification": row.get("verification_level", ""),
                "notes": row.get("notes", ""),
                "full_signature": row.get("full_signature", ""),
            }
    return out
