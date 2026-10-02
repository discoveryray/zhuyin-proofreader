"""Explicit, fail-closed transfer of a sealed review project to another PDF.

This module never changes a source session, its workbooks, or its identities.
The transfer receipt is local project evidence, not Global glyph evidence.
"""

from __future__ import annotations

import copy
from contextlib import contextmanager, ExitStack
import base64
import binascii
import functools
import hashlib
import json
import os
import tempfile
import uuid
import zlib
from datetime import datetime
from pathlib import Path
from typing import Any, Mapping, Sequence

import fitz


PROOF_FILE = "PDF內容對應證據.json"
PROOF_VERSION = 2
PENDING_ACTUAL_FILE = "來源actual待重新核對.json"
INCOMPLETE_FILE = "跨電腦接續未完成.json"
ACTUAL_EXCEL_CONFLICT_FILE = "跨Excel_actual衝突.json"
ACTUAL_EXCEL_MARKER_VERSION = 1
CONFLICT_FILE = "跨專案判定衝突.json"
EXCEL_PROOF_SHEET = "跨電腦內容證據"
EXCEL_PROOF_VERSION = 1
PRESENTATION_PLAN_VERSION = 1
PRESENTATION_RECEIPTS = (CONFLICT_FILE, PENDING_ACTUAL_FILE)


def _sha(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


@contextmanager
def _workbook_snapshot(path: Path, *, expected_sha256: str | None = None):
    """All workbook consumers read the same bytes whose digest was verified."""
    original = Path(path).resolve()
    data = original.read_bytes()
    digest = hashlib.sha256(data).hexdigest()
    if expected_sha256 is not None and digest != expected_sha256:
        raise ValueError("來源工作簿 SHA 不符；未使用變動的 workbook")
    with tempfile.TemporaryDirectory(prefix="portable-excel-snapshot-") as folder:
        snapshot = Path(folder) / original.name
        snapshot.write_bytes(data)
        yield snapshot, digest


def _stable_excel_import(importer):
    """Read user Excel bytes once before any workbook metadata or decision read."""
    @functools.wraps(importer)
    def wrapped(target_dir, xlsx, *args, **kwargs):
        original = Path(xlsx).resolve()
        with _workbook_snapshot(original) as (snapshot, digest):
            return importer(target_dir, snapshot, *args,
                            _original_xlsx=original, _snapshot_sha=digest, **kwargs)
    return wrapped


def _verify_excel_import_snapshot(snapshot: Path, original: Path, digest: str) -> None:
    if _sha(snapshot) != digest or _sha(original) != digest:
        raise ValueError("Excel 來源檔於匯入期間變動；未寫入目標判定")


def _canonical(value: Any) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


def _pdf_content_snapshot(path: Path, *, expected_sha256: str | None = None):
    """Read PDF bytes once, then derive SHA, candidate text and raster from them.

    The candidate resolver uses RAWDICT text even when a PDF marks it invisible.
    Comparing its normalized characters in displayed coordinates prevents a
    same-raster PDF from silently changing proofreading context. Font names and
    PDF object identifiers are storage details, so they are not part of proof.
    """
    from check_pronunciation_candidates import build_pdf_line_index

    data = Path(path).read_bytes()
    digest = hashlib.sha256(data).hexdigest()
    if expected_sha256 is not None and digest != expected_sha256:
        raise ValueError(f"PDF SHA 在建立內容證據前已變動：{path}")
    # The production candidate index accepts a path. Feed it a private copy
    # of the already-read bytes; raster rendering uses the same in-memory bytes.
    with tempfile.TemporaryDirectory(prefix="portable-pdf-snapshot-") as folder:
        snapshot_path = Path(folder) / "content.pdf"
        snapshot_path.write_bytes(data)
        line_index = build_pdf_line_index(snapshot_path, strict_extraction=True)
        if _sha(snapshot_path) != digest:
            raise ValueError("PDF 校對文字快照於讀取期間變動")
    signatures = []
    with fitz.open(stream=data, filetype="pdf") as document:
        if not document.page_count:
            raise ValueError("PDF 沒有頁面")
        for page in document:
            if page.rect.is_empty:
                raise ValueError("PDF 頁面尺寸無效")
            pix = page.get_pixmap(matrix=fitz.Matrix(3, 3), colorspace=fitz.csRGB, alpha=False, annots=True)
            display = page.rotation_matrix
            candidate_blocks = []
            for block in line_index.get(page.number + 1, []):
                chars = []
                for char in block.get("chars", []):
                    box = fitz.Rect(char["bbox"]) * display
                    chars.append((char["c"], *(round(value, 2) for value in box)))
                box = fitz.Rect(block["bbox"]) * display
                candidate_blocks.append({
                    "text": block["text"],
                    "bbox": [round(value, 2) for value in box],
                    "chars": chars,
                })
            candidate_blocks.sort(key=lambda item: (item["bbox"], item["text"], item["chars"]))
            signatures.append({
                "width": round(page.rect.width, 4),
                "height": round(page.rect.height, 4),
                "rotation": page.rotation,
                "display_matrix": [round(value, 6) for value in page.rotation_matrix],
                "pixels_width": pix.width,
                "pixels_height": pix.height,
                "rgb_sha256": hashlib.sha256(pix.samples).hexdigest(),
                "candidate_text_sha256": hashlib.sha256(_canonical(candidate_blocks)).hexdigest(),
            })
    return digest, signatures


def _page_signatures(path: Path) -> list[dict[str, Any]]:
    return _pdf_content_snapshot(path)[1]


def _source_pdf(manifest: Mapping[str, Any], info: Mapping[str, Any], output_dir: Path) -> Path:
    """Only exact source bytes may create a portable proof."""
    wanted = str(info.get("pdf_sha256") or "")
    name = str(info.get("pdf_name") or "")
    candidates = [Path(str(info.get("pdf") or "")), output_dir.parent / name]
    matching = [path.resolve() for path in candidates if path.is_file() and _sha(path) == wanted]
    if not matching:
        raise FileNotFoundError(f"無法建立跨電腦證據：原 PDF 不存在或 SHA 不符：{name}")
    return matching[0]


def _verify_available_session_pdfs(output_dir: Path, manifest) -> None:
    """Expected can work offline, but a known local replacement is not offline."""
    import standalone_proofread as sp
    for info in manifest.get("pdfs", []):
        try:
            sp._resolve_session_pdfs(output_dir, {"pdfs": [info]})
        except FileNotFoundError:
            name = str(info.get("pdf_name") or Path(str(info.get("pdf") or "")).name)
            stored = Path(str(info.get("pdf") or ""))
            if stored.is_file() or (name and any(path.is_file() for path in output_dir.parent.rglob(name))):
                raise ValueError(f"目前 PDF 與工作階段 SHA 不符：{name}；未寫入判定")


@contextmanager
def _project_transfer_locks(*directories):
    """Acquire source and target roots in the same order for every transfer."""
    import standalone_proofread as sp
    roots = {sp.project_actual_evidence_root(Path(path).resolve()) for path in directories}
    with ExitStack() as stack:
        for root in sorted(roots, key=lambda path: str(path).casefold()):
            stack.enter_context(sp.project_delivery_lock(root))
        yield


def _artifact(output_dir: Path, info: Mapping[str, Any], kind: str) -> Path:
    stored = Path(str(info.get(f"{kind}_workbook") or ""))
    subdir = "01_實際注音" if kind == "actual" else "02_候選報告"
    local = output_dir / subdir / stored.name
    wanted = str(info.get(f"{kind}_workbook_sha256") or "")
    candidates = {str(path.resolve()): path for path in (stored, local) if path.is_file()}
    matches = [path for path in candidates.values() if wanted and _sha(path) == wanted]
    if not matches:
        raise ValueError(f"DATA_INTEGRITY_ERROR：來源工作簿遺失或 SHA 不符：{stored.name}")
    return matches[0]


def _load_project(output_dir: Path, *, allow_incomplete: bool = False,
                  allow_unresolved_conflict: bool = False,
                  allow_actual_excel_conflict: bool = False):
    import standalone_proofread as sp

    output_dir = Path(output_dir).resolve()
    if not allow_incomplete and (output_dir / INCOMPLETE_FILE).exists():
        raise ValueError(f"跨電腦接續未完成，不能當作完整來源或匯入目標：{output_dir}")
    manifest = sp.json_load_strict(output_dir / "校對工作階段.json")
    sp.validate_manifest_integrity(manifest)
    for info in manifest.get("pdfs", []):
        _artifact(output_dir, info, "actual")
        _artifact(output_dir, info, "candidate")
    raw_db = sp.json_load_strict(output_dir / "人工判定資料庫.json")
    if not isinstance(raw_db, dict):
        raise ValueError("DATA_INTEGRITY_ERROR：人工判定資料庫根節點必須是物件")
    db = sp.normalize_db(raw_db)
    unresolved = validate_conflict_state(output_dir, manifest, db)
    if unresolved and not allow_unresolved_conflict:
        raise ValueError(f"來源專案有未裁決的同位置判定衝突，不能再轉送：{unresolved}")
    actual_conflicts = actual_excel_conflict_state(output_dir, manifest, db)
    if actual_conflicts and not allow_actual_excel_conflict:
        raise ValueError(f"Excel actual 衝突尚未經本地原頁重核，不能當完整專案轉送：{actual_conflicts}")
    ledger = sp.materialize_ledger(manifest, db)
    stale = [row["review_id"] for row in ledger if row.get("review_event_replay_status")]
    if stale:
        raise ValueError(f"來源/目標工作階段已有失效判定，不能當有效成果轉用：{stale[:10]}")
    return manifest, db


def _actual_excel_conflict_token(conflict: Mapping[str, Any]) -> str:
    return "portable_actual_conflict:" + hashlib.sha256(_canonical(conflict)).hexdigest()


def _sealed_actual_excel_conflicts(conflicts: list[dict[str, Any]]) -> dict[str, Any]:
    payload = {"version": 1, "conflicts": conflicts}
    return {**payload, "integrity_sha256": hashlib.sha256(_canonical(payload)).hexdigest()}


def _validated_actual_excel_receipt(path: Path, *, with_digest: bool = False
                                    ) -> dict[str, Any] | tuple[dict[str, Any], str]:
    import standalone_proofread as sp

    receipt, digest = sp.json_load_snapshot(path)
    if (not isinstance(receipt, dict)
            or set(receipt) != {"version", "conflicts", "integrity_sha256"}
            or receipt.get("version") != 1
            or not isinstance(receipt.get("conflicts"), list) or not receipt["conflicts"]):
        raise ValueError("Excel actual 衝突紀錄格式無法驗證")
    expected = hashlib.sha256(_canonical({"version": 1, "conflicts": receipt["conflicts"]})).hexdigest()
    if receipt["integrity_sha256"] != expected:
        raise ValueError("Excel actual 衝突紀錄完整性無法驗證")
    return (receipt, digest) if with_digest else receipt


def actual_excel_conflict_state(output_dir: Path, manifest: Mapping[str, Any],
                                db: Mapping[str, Any]) -> list[str]:
    """Keep imported actual disagreements pending until new local visual proof exists."""
    import actual_review as ar
    import standalone_proofread as sp
    from global_glyph_promotion import committed_project_recovery

    receipt_path = Path(output_dir) / ACTUAL_EXCEL_CONFLICT_FILE
    if not receipt_path.exists():
        return []
    receipt = _validated_actual_excel_receipt(receipt_path)
    entries = {entry["review_id"]: entry for entry in manifest.get("records", [])}
    ledger = {entry["review_id"]: entry for entry in sp.materialize_ledger(manifest, db)}
    root = sp.project_actual_evidence_root(output_dir)
    override_rows = ar._read_csv(root / ar.OCCURRENCE_OVERRIDE_FILE, ar.OVERRIDE_HEADERS)
    overrides = {tuple(row[field] for field in (
        "pdf_contains", "pdf_excludes", "page", "target_char", "stable_key", "x0", "y0"
    )): row for row in override_rows}
    provenance = ar._read_csv(root / ar.GLYPH_PROVENANCE_FILE, ar.GLYPH_PROVENANCE_HEADERS)
    refresh_pending = committed_project_recovery(root) is not None
    pending = []
    seen = set()
    for conflict in receipt["conflicts"]:
        if not isinstance(conflict, dict) or set(conflict) != {
            "target_occurrence_id", "target_review_id", "target_pdf_sha256",
            "target_record", "source_record", "source_reading",
        }:
            raise ValueError("Excel actual 衝突來源格式無法驗證")
        review_id = conflict.get("target_review_id")
        entry = entries.get(review_id)
        source = conflict.get("source_record")
        target = conflict.get("target_record")
        reading = ar._canon(conflict.get("source_reading"))
        decision = source.get("source_decision") if isinstance(source, dict) else None
        source_occurrences = source.get("source_occurrence_ids") if isinstance(source, dict) else None
        source_reviews = source.get("source_review_ids") if isinstance(source, dict) else None
        source_pdfs = source.get("source_pdf_sha256") if isinstance(source, dict) else None
        source_roster = source.get("source_roster") if isinstance(source, dict) else None
        identity = (conflict.get("target_occurrence_id"),
                    source.get("source_excel_sha256") if isinstance(source, dict) else None,
                    source.get("source_excel_row") if isinstance(source, dict) else None,
                    reading)
        if (entry is None or conflict.get("target_occurrence_id") != entry.get("occurrence_id")
                or conflict.get("target_pdf_sha256") != entry.get("pdf_sha256")
                or not isinstance(source, dict) or not isinstance(target, dict)
                or not reading or reading != conflict.get("source_reading")
                or not isinstance(decision, dict) or decision.get("decision") != "VERIFIED"
                or decision.get("actual_reading") != reading
                or decision.get("group_id") != source.get("source_group_id")
                or decision.get("group_snapshot") != source.get("source_group_snapshot")
                or decision.get("sample_a_occurrence_id") not in (source_occurrences or [])
                or not isinstance(source_occurrences, list) or not source_occurrences
                or not isinstance(source_reviews, list) or len(source_reviews) != len(source_occurrences)
                or not all(isinstance(item, str) and item for item in source_reviews)
                or not isinstance(source_roster, list) or len(source_roster) != len(source_occurrences)
                or [(item.get("occurrence_id"), item.get("review_id"))
                    for item in source_roster if isinstance(item, dict)]
                != list(zip(source_occurrences, source_reviews))
                or not isinstance(source_pdfs, list) or not source_pdfs
                or not all(isinstance(item, str) and len(item) == 64
                           and all(char in "0123456789abcdef" for char in item) for item in source_pdfs)
                or any(item.get("pdf_sha256") not in source_pdfs for item in source_roster)
                or not isinstance(source.get("source_manifest_integrity_sha256"), str)
                or len(source["source_manifest_integrity_sha256"]) != 64
                or not source.get("source_session_id")
                or source.get("target_session_id") != manifest.get("session_id")
                or not isinstance(source.get("source_excel_sha256"), str)
                or len(source["source_excel_sha256"]) != 64
                or not all(char in "0123456789abcdef" for char in source["source_excel_sha256"])
                or type(source.get("source_excel_row")) is not int or source["source_excel_row"] < 2
                or entry["occurrence_id"] not in (source.get("target_occurrence_ids") or [])
                or not ar._canon(target.get("actual_reading"))
                or ar._canon(target.get("actual_reading")) == reading
                or tuple(target.get(field) for field in (
                    "pdf_contains", "pdf_excludes", "page", "target_char", "stable_key", "x0", "y0"
                )) != ar._override_key_from_entry(entry)):
            raise ValueError("Excel actual 衝突來源或目標位置識別無法驗證")
        if identity in seen:
            raise ValueError("Excel actual 衝突來源識別重複，無法驗證")
        seen.add(identity)
        token = _actual_excel_conflict_token(conflict)
        current = overrides.get(ar._override_key_from_entry(entry))
        current_reading = ar._canon((current or {}).get("actual_reading"))
        if (current is None or current.get("source") != "人工 GUI actual 視覺確認"
                or token not in current.get("note", "")
                or current_reading != ar._canon(ledger[review_id].get("actual"))
                or refresh_pending):
            pending.append(review_id)
            continue
        matched = any(
            row.get("event_type") == "OCCURRENCE_OVERRIDE"
            and row.get("source") == "人工 GUI actual 視覺確認"
            and token in row.get("note", "")
            and entry["occurrence_id"] in str(row.get("occurrence_ids") or "").split("|")
            and ar._canon(row.get("bopomofo")) == current_reading
            for row in provenance
        )
        if not matched:
            pending.append(review_id)
    return sorted(set(pending))


def actual_excel_conflict_notes(output_dir: Path, manifest: Mapping[str, Any],
                                db: Mapping[str, Any], checked_ids: Sequence[str]) -> str:
    """Bind a fresh local visual action to every pending conflict it checks."""
    import standalone_proofread as sp

    receipt_path = Path(output_dir) / ACTUAL_EXCEL_CONFLICT_FILE
    if not receipt_path.exists():
        return ""
    pending = set(actual_excel_conflict_state(output_dir, manifest, db))
    if not pending:
        return ""
    receipt = _validated_actual_excel_receipt(receipt_path)
    checked = set(checked_ids)
    tokens = sorted({_actual_excel_conflict_token(item) for item in receipt["conflicts"]
                     if item["target_review_id"] in pending
                     and item["target_occurrence_id"] in checked})
    return "；".join(tokens)


def hold_actual_excel_completion(output_dir: Path, manifest: Mapping[str, Any],
                                 db: Mapping[str, Any], gate: dict[str, Any]) -> dict[str, Any]:
    """Keep an independently computed ledger from declaring conflict completion."""
    pending = actual_excel_conflict_state(output_dir, manifest, db)
    if not pending:
        return gate
    from occurrence_ledger import PROCESSING_FINISHED, PROOFREAD_COMPLETE

    held = copy.deepcopy(gate)
    held["hard_gates"]["actual_excel_conflicts_zero"] = False
    held["failed_gates"].append("actual_excel_conflicts_zero")
    held["complete"] = False
    if held["status"] == PROOFREAD_COMPLETE:
        held["status"] = PROCESSING_FINISHED
    held["actual_excel_conflict_review_ids"] = pending
    return held


def _actual_excel_conflict_status(status: Mapping[str, Any], conflicts: list[dict[str, Any]]) -> dict[str, Any]:
    from occurrence_ledger import PROCESSING_FINISHED, PROOFREAD_COMPLETE

    updated = copy.deepcopy(dict(status))
    if updated.get("status") == PROOFREAD_COMPLETE:
        updated["status"] = PROCESSING_FINISHED
    updated["status_text"] = "Excel actual 同位置判定衝突尚待原頁重新核對。"
    ids = sorted({item["target_review_id"] for item in conflicts})
    updated["actual_excel_conflict_review_ids"] = ids
    gate = updated.get("completion_gate")
    if isinstance(gate, dict):
        gate["hard_gates"] = dict(gate.get("hard_gates") or {})
        gate["hard_gates"]["actual_excel_conflicts_zero"] = False
        gate["failed_gates"] = list(dict.fromkeys([*(gate.get("failed_gates") or []),
                                                     "actual_excel_conflicts_zero"]))
        gate["complete"] = False
        if gate.get("status") == PROOFREAD_COMPLETE:
            gate["status"] = PROCESSING_FINISHED
        gate["actual_excel_conflict_review_ids"] = ids
    return updated


def validate_conflict_state(output_dir: Path, manifest: Mapping[str, Any], db: Mapping[str, Any]):
    """A conflicted row is inactive until a fresh local event cites its receipt."""
    import standalone_proofread as sp

    path = Path(output_dir) / CONFLICT_FILE
    if not path.exists():
        return []
    receipt = sp.json_load_strict(path)
    if (not isinstance(receipt, dict) or receipt.get("version") != 1
            or not isinstance(receipt.get("conflicts"), list) or not receipt["conflicts"]):
        raise ValueError("跨專案判定衝突紀錄格式無法驗證")
    valid_ids = {entry["review_id"] for entry in manifest.get("records", [])}
    unresolved = []
    def locally_resolves(event, conflict):
        resolution = event.get("portability_conflict_resolution") if isinstance(event, dict) else None
        return (isinstance(resolution, dict)
                and event.get("action") != "保留待人工"
                and event.get("portability_source") is None
                and hashlib.sha256(_canonical(conflict)).hexdigest()
                in resolution.get("conflict_sha256", [])
                and conflict in resolution.get("original_conflicts", []))

    for index, conflict in enumerate(receipt["conflicts"]):
        if (not isinstance(conflict, dict)
                or conflict.get("target_review_id") not in valid_ids
                or not isinstance(conflict.get("source_event"), dict)
                or not isinstance(conflict.get("target_event"), dict)
                or not conflict.get("source_pdf_sha256")
                or not conflict.get("target_pdf_sha256")):
            raise ValueError("跨專案判定衝突來源或位置識別不完整")
        review_id = conflict["target_review_id"]
        # A later conflict removes the active local adjudication. Its target
        # snapshot is the durable, exact record of that adjudication and the
        # earlier receipt it resolved; neither source verdict is discarded.
        historical_resolution = any(
            later.get("target_review_id") == review_id
            and locally_resolves(later.get("target_event"), conflict)
            for later in receipt["conflicts"][index + 1:]
            if isinstance(later, dict))
        if historical_resolution:
            continue
        event = db["events"].get(review_id)
        if event is None or event.get("action") == "保留待人工":
            unresolved.append(review_id)
        else:
            if not locally_resolves(event, conflict):
                raise ValueError(f"衝突位置有未經本地裁決的有效事件：{review_id}")
    return unresolved


def conflict_resolution_evidence(output_dir: Path, manifest: Mapping[str, Any],
                                 db: Mapping[str, Any], review_id: str):
    unresolved = validate_conflict_state(output_dir, manifest, db)
    if review_id not in unresolved:
        return None
    import standalone_proofread as sp

    path = Path(output_dir) / CONFLICT_FILE
    receipt = sp.json_load_strict(path)
    selected = [copy.deepcopy(item) for item in receipt["conflicts"]
                if item["target_review_id"] == review_id]
    return {
        "conflict_sha256": [hashlib.sha256(_canonical(item)).hexdigest() for item in selected],
        "original_conflicts": selected,
    }


def _actual_transfer_state(output_dir: Path, manifest: Mapping[str, Any]):
    """Preserve unresolved actual intent without creating new human evidence."""
    import actual_review as ar
    import standalone_proofread as sp
    from global_glyph_promotion import PROJECT_TRANSACTION_FILE, load_promotion_outbox

    root = sp.project_actual_evidence_root(output_dir)
    transaction = root / PROJECT_TRANSACTION_FILE
    if transaction.exists():
        # PREPARED must roll back; COMMITTED must complete its source-side
        # postconditions. Neither journal can be replayed under a new session.
        raise ValueError(
            f"來源 actual 交易尚待原專案恢復：{transaction}；先在來源完成 recovery，"
            "再準備/匯入跨電腦成果。未搬移交易或增加 Global quorum"
        )
    # An outbox can be pending even after project evidence committed.
    # Do not transplant it into another project or deliver it as fresh proof.
    if any(item["status"] == "PENDING" for item in load_promotion_outbox(root)["items"]):
        raise ValueError("來源 Global promotion outbox 尚有未交付意圖；先在來源完成恢復")
    staging = ar.load_manual_actual_staging(root)
    source_ids = {str(entry.get("occurrence_id") or "") for entry in manifest.get("records", [])}
    for group in staging["staged_groups"]:
        if not set(group["member_occurrence_ids"]) <= source_ids:
            raise ValueError("來源 actual 暫存含工作階段不存在的 occurrence；未轉接")
    prior = []
    receipt_path = output_dir / PENDING_ACTUAL_FILE
    if receipt_path.exists():
        receipt = sp.json_load_strict(receipt_path)
        if (not isinstance(receipt, dict)
                or receipt.get("status") != "RECHECK_LOCAL_PDF_REQUIRED"
                or not isinstance(receipt.get("sources"), list)):
            raise ValueError("前次來源 actual 暫存紀錄格式無法驗證；未轉接")
        for item in receipt["sources"]:
            if not isinstance(item, dict) or not isinstance(item.get("mapped_occurrences"), dict):
                raise ValueError("前次來源 actual 暫存位置映射缺失；未轉接")
            original_staging = ar._validate_manual_actual_staging_document(item.get("staging"))
            if (not item.get("source_session_id") or not item.get("source_manifest_integrity_sha256")
                    or not item.get("source_pdf_sha256")
                    or not original_staging["staged_groups"]):
                raise ValueError("前次來源 actual 暫存原始身分不完整；未轉接")
            original_mapping = item["mapped_occurrences"]
            staged_ids = {occurrence_id for group in original_staging["staged_groups"]
                          for occurrence_id in group["member_occurrence_ids"]}
            if not staged_ids <= set(original_mapping):
                raise ValueError("前次來源 actual 暫存 occurrence 原位置未對應；未轉接")
            hops = item.get("forwarded_via", [])
            if not isinstance(hops, list) or any(not isinstance(hop, dict) for hop in hops):
                raise ValueError("前次來源 actual 暫存轉送鏈無效；未轉接")
            previous_session = item.get("target_session_id")
            previous_ids = set(original_mapping.values())
            if not previous_session or len(previous_ids) != len(original_mapping):
                raise ValueError("前次來源 actual 暫存首跳位置不唯一；未轉接")
            for hop in hops:
                hop_mapping = hop.get("mapped_occurrences")
                if (hop.get("source_session_id") != previous_session
                        or not hop.get("target_session_id")
                        or not isinstance(hop_mapping, dict)
                        or set(hop_mapping) != previous_ids
                        or len(set(hop_mapping.values())) != len(hop_mapping)):
                    raise ValueError("前次來源 actual 暫存轉送鏈位置無法逐跳核對；未轉接")
                previous_session = hop["target_session_id"]
                previous_ids = set(hop_mapping.values())
            if previous_session != manifest["session_id"] or previous_ids != source_ids:
                raise ValueError("前次來源 actual 暫存與目前工作階段不符；未轉接")
            prior.append(item)
    return {
        "source_session_id": manifest["session_id"],
        "source_manifest_integrity_sha256": manifest["manifest_integrity_sha256"],
        "staging": staging,
        "prior_receipts": prior,
        "status": "RECHECK_LOCAL_PDF_REQUIRED" if staging["staged_groups"] or prior else "NONE",
    }


def _proof_payload(output_dir: Path, manifest: Mapping[str, Any]) -> dict[str, Any]:
    pdfs = []
    for info in manifest.get("pdfs", []):
        path = _source_pdf(manifest, info, output_dir)
        observed_sha, pages = _pdf_content_snapshot(path, expected_sha256=str(info["pdf_sha256"]))
        pdfs.append({
            "pdf_sha256": observed_sha,
            "pdf_name": str(info["pdf_name"]),
            "pages": pages,
        })
    if not pdfs or len({item["pdf_sha256"] for item in pdfs}) != len(pdfs):
        raise ValueError("來源工作階段 PDF 清單空白或 SHA 不唯一")
    return {
        "version": PROOF_VERSION,
        "session_id": str(manifest.get("session_id") or ""),
        "manifest_integrity_sha256": str(manifest["manifest_integrity_sha256"]),
        "pdfs": pdfs,
    }


def _excel_bound_rows(workbook, kind: str):
    """Bind every source row while leaving only the documented answer cells editable."""
    import standalone_proofread as sp

    sheet_name = "待判定候選" if kind == "expected" else "actual待判定"
    editable = ({"action", "proposed_expected_set", "proposed_expected_evidence",
                 "proposed_context_evidence", "exclusion_reason", "exclusion_evidence",
                 "note", *sp.CONFIRMATION_GATES}
                if kind == "expected" else
                {"decision", "actual_reading", "confidence", "sample_a_checked",
                 "sample_b_checked", "note"})
    sheet = workbook[sheet_name]
    rows = list(sheet.values)
    if not rows:
        raise ValueError("Excel 來源判定工作表空白")
    headers = [str(value or "") for value in rows[0]]
    if not all(headers) or len(headers) != len(set(headers)):
        raise ValueError("Excel 來源判定欄位無效")
    return [{header: ("" if value is None else value)
             for header, value in zip(headers, row) if header not in editable}
            for row in rows[1:]]


def write_excel_content_proof(workbook, output_dir: Path, manifest, db, *, kind: str,
                              snapshot_db=None, attestations=None,
                              allow_internal_incomplete: bool = False) -> None:
    """Embed a sealed source snapshot and all-page proof in an invisible sheet."""
    import standalone_proofread as sp

    if kind not in {"expected", "actual"} or EXCEL_PROOF_SHEET in workbook.sheetnames:
        raise ValueError("Excel 跨電腦內容證據種類或工作表無效")
    if not manifest.get("pdfs"):
        # Historical synthetic/session-only exports remain locally importable;
        # they cannot claim portable content evidence without any PDF bytes.
        return
    source_manifest, source_db = _load_project(
        output_dir, allow_incomplete=allow_internal_incomplete,
        allow_unresolved_conflict=allow_internal_incomplete)
    if source_manifest != manifest or source_db != sp.normalize_db(db):
        raise ValueError("Excel 匯出期間來源專案判定或封印已改變")
    if snapshot_db is None:
        snapshot_db = source_db
    else:
        snapshot_db = sp.normalize_db(snapshot_db)
        if any(row.get("review_event_replay_status") for row in
               sp.materialize_ledger(source_manifest, snapshot_db)):
            raise ValueError("Excel 補證來源快照判定失效")
    try:
        page_proof = _proof_payload(Path(output_dir), source_manifest)
    except FileNotFoundError:
        if kind != "expected":
            raise
        saved = Path(output_dir) / PROOF_FILE
        if not saved.exists():
            # The ordinary same-session workbook remains usable. It carries
            # no cross-session claim without the exact PDF or sealed proof.
            return
        page_proof = _load_proof(Path(output_dir), source_manifest)
    payload = {"version": EXCEL_PROOF_VERSION, "kind": kind,
               "exported_at": datetime.now().astimezone().isoformat(timespec="microseconds"),
               "manifest": source_manifest, "db": snapshot_db,
               "proof": page_proof,
               "bound_rows": _excel_bound_rows(workbook, kind),
               "attestations": attestations or {}}
    raw = _canonical(payload)
    encoded = base64.b64encode(zlib.compress(raw, level=9)).decode("ascii")
    chunks = [encoded[index:index + 30000] for index in range(0, len(encoded), 30000)]
    sheet = workbook.create_sheet(EXCEL_PROOF_SHEET)
    sheet.append(["version", EXCEL_PROOF_VERSION])
    sheet.append(["sha256", hashlib.sha256(raw).hexdigest()])
    sheet.append(["chunk_count", len(chunks)])
    for index, chunk in enumerate(chunks, 1):
        sheet.append([index, chunk])
    sheet.sheet_state = "veryHidden"


def load_excel_content_proof(xlsx: Path, target_dir: Path, target_manifest, *,
                             kind: str, workbook_session: str):
    """Validate the carried source and map every occurrence by displayed position."""
    import standalone_proofread as sp
    from openpyxl import load_workbook

    workbook = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        if EXCEL_PROOF_SHEET not in workbook.sheetnames:
            return _legacy_exact_excel_mapping(xlsx, target_dir, target_manifest,
                                               kind=kind, workbook_session=workbook_session)
        rows = list(workbook[EXCEL_PROOF_SHEET].values)
        bound_rows = _excel_bound_rows(workbook, kind)
    finally:
        workbook.close()
    if (len(rows) < 4 or rows[0][:2] != ("version", EXCEL_PROOF_VERSION)
            or rows[1][0] != "sha256" or rows[2][0] != "chunk_count"
            or type(rows[2][1]) is not int or rows[2][1] < 1
            or len(rows) != rows[2][1] + 3
            or any(row[0] != index or not isinstance(row[1], str)
                   for index, row in enumerate(rows[3:], 1))):
        raise ValueError("Excel 跨電腦內容證據格式損壞")
    try:
        packed = base64.b64decode("".join(row[1] for row in rows[3:]), validate=True)
        decompressor = zlib.decompressobj()
        raw = decompressor.decompress(packed, 128 * 1024 * 1024 + 1)
        if (len(raw) > 128 * 1024 * 1024 or not decompressor.eof
                or decompressor.unconsumed_tail or decompressor.unused_data):
            raise ValueError("Excel 跨電腦內容證據超過大小上限或壓縮串流不完整")
        payload = json.loads(raw)
    except (ValueError, zlib.error, TypeError) as exc:
        raise ValueError("Excel 跨電腦內容證據無法解碼") from exc
    if (len(raw) > 128 * 1024 * 1024 or hashlib.sha256(raw).hexdigest() != rows[1][1]
            or not isinstance(payload, dict) or payload.get("version") != EXCEL_PROOF_VERSION
            or payload.get("kind") != kind or _canonical(payload) != raw):
        raise ValueError("Excel 跨電腦內容證據完整性或種類不符")
    source_manifest, source_db, proof = (payload.get(key) for key in ("manifest", "db", "proof"))
    if payload.get("bound_rows") != bound_rows:
        raise ValueError("Excel 來源判定列的原始識別、快照或位置已被修改")
    if not isinstance(source_manifest, dict) or not isinstance(source_db, dict) or not isinstance(proof, dict):
        raise ValueError("Excel 來源專案證據格式不完整")
    sp.validate_manifest_integrity(source_manifest)
    if str(source_manifest.get("session_id")) != workbook_session:
        raise ValueError("Excel 來源 session 識別與內容證據不符")
    source_db = sp.normalize_db(source_db)
    source_ledger = sp.materialize_ledger(source_manifest, source_db)
    if any(row.get("review_event_replay_status") for row in source_ledger):
        raise ValueError("Excel 來源專案含失效判定")
    if (proof.get("version") != PROOF_VERSION
            or proof.get("session_id") != source_manifest["session_id"]
            or proof.get("manifest_integrity_sha256") != source_manifest["manifest_integrity_sha256"]
            or [(item.get("pdf_sha256"), item.get("pdf_name")) for item in proof.get("pdfs", [])]
            != [(item.get("pdf_sha256"), item.get("pdf_name")) for item in source_manifest.get("pdfs", [])]):
        raise ValueError("Excel 來源全頁內容證據與封印不符")
    pdf_map, geometry = _match_pdfs(proof, source_manifest, target_manifest, Path(target_dir))
    reviews = _map_reviews(source_manifest, target_manifest, pdf_map, geometry)
    return (source_manifest, source_db, source_ledger, reviews, pdf_map,
            payload["exported_at"], payload, geometry)


def _legacy_exact_excel_mapping(xlsx: Path, target_dir: Path, target_manifest, *,
                                kind: str, workbook_session: str):
    """Narrow adapter for proofless files whose original IDs and snapshots match."""
    import actual_review as ar
    import standalone_proofread as sp

    if not workbook_session or workbook_session == target_manifest["session_id"]:
        raise ValueError("舊 Excel 缺少可驗證的跨 session 來源識別")
    geometry = {}
    for info in target_manifest.get("pdfs", []):
        path = _source_pdf(target_manifest, info, Path(target_dir))
        _, pages = _pdf_content_snapshot(path, expected_sha256=info["pdf_sha256"])
        geometry[info["pdf_sha256"]] = (pages, pages)
    source_db = sp.normalize_db({})
    source_ledger = sp.materialize_ledger(target_manifest, source_db)
    by_review = {entry["review_id"]: entry for entry in source_ledger}
    if kind == "expected":
        rows = sp.workbook_rows(xlsx, "待判定候選", ["occurrence_id", "review_id",
                          "review_snapshot", "action", "expected_target_snapshot"])
        sp.assert_unique_ids(rows, "occurrence_id")
        sp.assert_unique_ids(rows, "review_id")
        checked = 0
        for number, row in enumerate(rows, 2):
            if not str(row.get("action") or "").strip():
                continue
            entry = by_review.get(str(row.get("review_id") or ""))
            if (entry is None or str(row.get("occurrence_id") or "") != entry["occurrence_id"]
                    or str(row.get("review_snapshot") or "") != sp._review_snapshot(entry)
                    or not sp._expected_target_unchanged(row, entry)):
                raise ValueError(f"舊 Excel row {number} 沒有全頁證據，且 exact ID/快照不符；不同 SHA 或檔名需由 A 補證")
            checked += 1
    else:
        _, rows = ar._load_sheet_rows(xlsx, "actual待判定")
        groups = {group["group_id"]: group for group in ar.build_actual_review_groups(source_ledger)}
        checked = 0
        for number, row in enumerate(rows, 2):
            if str(row.get("decision") or "").strip().upper() != "VERIFIED":
                continue
            group = groups.get(str(row.get("group_id") or ""))
            if (group is None or str(row.get("group_snapshot") or "") != group["group_snapshot"]
                    or str(row.get("sample_a_occurrence_id") or "") != group["members"][0]["occurrence_id"]
                    or str(row.get("sample_b_occurrence_id") or "") !=
                       (group["members"][1]["occurrence_id"] if len(group["members"]) > 1 else "")):
                raise ValueError(f"舊 actual row {number} 沒有全頁證據，且 exact group/快照不符；不同 SHA 需由 A 補證")
            checked += 1
    if not checked:
        raise ValueError("舊 Excel 沒有可核對的已填判定；無法跨 session 匯入")
    # This is an observed exact-ID relationship, not a recovered source seal.
    source_context = {
        "session_id": workbook_session,
        "manifest_integrity_sha256": "",
        "expected_asset_fingerprint": target_manifest.get("expected_asset_fingerprint"),
        "pdfs": copy.deepcopy(target_manifest["pdfs"]),
        "records": copy.deepcopy(target_manifest["records"]),
    }
    reviews = {entry["review_id"]: (entry, entry) for entry in source_ledger}
    pdf_map = {item["pdf_sha256"]: item["pdf_sha256"] for item in target_manifest["pdfs"]}
    payload = {"attestations": {"basis": "LEGACY_EXACT_IDS_NO_SOURCE_SEAL"}}
    return (source_context, source_db, source_ledger, reviews, pdf_map,
            datetime.now().astimezone().isoformat(timespec="microseconds"), payload, geometry)


@_stable_excel_import
def attach_filled_excel_proof(source_dir: Path, xlsx: Path, *,
                              _original_xlsx: Path, _snapshot_sha: str) -> Path:
    """Attest an old filled workbook against the *current* sealed A project.

    This does not assert that the workbook caused an earlier A decision.  It
    verifies every filled answer against A's current evidence, then carries
    that present-time attestation and the original visible answers together.
    """
    import actual_review as ar
    import standalone_proofread as sp
    from openpyxl import load_workbook

    source_dir, xlsx = Path(source_dir).resolve(), Path(xlsx).resolve()
    manifest, current_db = _load_project(source_dir)
    workbook = load_workbook(xlsx)
    temporary = None
    try:
        if EXCEL_PROOF_SHEET in workbook.sheetnames:
            raise ValueError("Excel 已有跨電腦內容證據；不得重新封裝既有來源")
        kinds = [kind for kind, sheet in (("expected", "待判定候選"),
                                          ("actual", "actual待判定")) if sheet in workbook.sheetnames]
        if len(kinds) != 1:
            raise ValueError("舊 Excel 判定工作表種類無法唯一辨識")
        kind = kinds[0]
        metadata = sp.workbook_metadata(xlsx, "匯入中繼資料") if kind == "expected" else ar._load_sheet_rows(xlsx, "actual待判定")[0]
        if str(metadata.get("session_id") or "") != str(manifest["session_id"]):
            raise ValueError("舊 Excel session 與 A 現行封印不符，不能補證")
        if kind == "expected":
            rows = sp.workbook_rows(xlsx, "待判定候選", [
                "occurrence_id", "review_id", "review_snapshot", "action",
                "proposed_expected_set", "proposed_expected_evidence", "proposed_context_evidence",
                "exclusion_reason", "exclusion_evidence", *sp.CONFIRMATION_GATES, "note"])
            if str(metadata.get("expected_asset_fingerprint") or "") != str(manifest.get("expected_asset_fingerprint") or ""):
                raise ValueError("舊 Excel expected fingerprint 與 A 封印不符")
            snapshot_db = copy.deepcopy(current_db)
            attestations = {"basis": "CURRENT_SEALED_PROJECT_AT_ATTACHMENT", "expected_events": {}}
            for number, row in enumerate(rows, 2):
                action = str(row.get("action") or "").strip()
                if not action or action == "保留待人工":
                    continue
                review_id = str(row.get("review_id") or "")
                saved = current_db["events"].get(review_id)
                if saved is not None:
                    compared = {
                        "action": action,
                        "expected_set": str(row.get("proposed_expected_set") or "").strip(),
                        "expected_evidence": str(row.get("proposed_expected_evidence") or "").strip(),
                        "context_evidence": str(row.get("proposed_context_evidence") or "").strip(),
                        "exclusion_reason": str(row.get("exclusion_reason") or "").strip(),
                        "exclusion_evidence": str(row.get("exclusion_evidence") or "").strip(),
                        "note": str(row.get("note") or "").strip(),
                    }
                    compared.update({gate: str(row.get(gate) or "").strip().upper() == "Y"
                                     for gate in sp.CONFIRMATION_GATES})
                    if any(saved.get(key) != value for key, value in compared.items()):
                        raise ValueError(f"舊 Excel row {number} 判定與 A 目前保存事件不同，不能補證")
                    attestations["expected_events"][review_id] = copy.deepcopy(saved)
                    snapshot_db["events"].pop(review_id)
            ledger = sp.materialize_ledger(manifest, snapshot_db)
            by_id = {entry["review_id"]: entry for entry in ledger}
            for number, row in enumerate(rows, 2):
                action = str(row.get("action") or "").strip()
                if not action or action == "保留待人工":
                    continue
                entry = by_id.get(str(row.get("review_id") or ""))
                if (entry is None or str(row.get("occurrence_id") or "") != entry["occurrence_id"]
                        or str(row.get("review_snapshot") or "") != sp._review_snapshot(entry)
                        or not sp._expected_target_unchanged(row, entry)):
                    raise ValueError(f"舊 Excel row {number} 原始識別/快照/位置無法從 A 封印重建")
        else:
            metadata, rows = ar._load_sheet_rows(xlsx, "actual待判定")
            if str(metadata.get("actual_review_schema_version")) != ar.ACTUAL_REVIEW_SCHEMA_VERSION:
                raise ValueError("舊 Excel actual schema 與 A 不相容")
            snapshot_db = current_db
            attestations = {"basis": "CURRENT_SEALED_PROJECT_AT_ATTACHMENT", "actual_groups": {}}
            ledger = sp.materialize_ledger(manifest, current_db)
            by_occ = {entry["occurrence_id"]: entry for entry in ledger}
            current_pending = {group["group_id"]: group for group in ar.build_actual_review_groups(ledger)}
            override_path = sp.project_actual_evidence_root(source_dir) / ar.OCCURRENCE_OVERRIDE_FILE
            overrides = ar._read_csv(override_path, ar.OVERRIDE_HEADERS)
            fields = ("pdf_contains", "pdf_excludes", "page", "target_char", "stable_key", "x0", "y0")
            override_by_key = {tuple(item[field] for field in fields): item for item in overrides}
            for number, row in enumerate(rows, 2):
                if str(row.get("decision") or "").strip().upper() != "VERIFIED":
                    continue
                sample = by_occ.get(str(row.get("sample_a_occurrence_id") or ""))
                if sample is None:
                    raise ValueError(f"舊 actual row {number} 樣本 occurrence 已無法在 A 找到")
                group = current_pending.get(str(row.get("group_id") or ""))
                if group is None:
                    group = ar.build_actual_group_for_entry(ledger, sample)
                members = group["members"]
                reading = ar._canon(row.get("actual_reading"))
                if (str(row.get("group_id") or "") != group["group_id"]
                        or int(row.get("occurrence_count") or 0) != len(members)
                        or str(row.get("group_kind") or "") != group["kind"]
                        or str(row.get("exact_key") or "") != group["exact_key"]
                        or str(row.get("sample_a_occurrence_id") or "") != members[0]["occurrence_id"]
                        or str(row.get("sample_b_occurrence_id") or "") !=
                           (members[1]["occurrence_id"] if len(members) > 1 else "")
                        or not reading or str(row.get("actual_reading") or "").strip() != reading
                        or str(row.get("sample_a_checked") or "").upper() != "Y"
                        or (len(members) > 1 and str(row.get("sample_b_checked") or "").upper() != "Y")):
                    raise ValueError(f"舊 actual row {number} 群組/樣本/判定無法與 A 現況核對")
                if str(row.get("group_snapshot") or "") != group["group_snapshot"]:
                    for member in members:
                        local = override_by_key.get(ar._override_key_from_entry(member))
                        if local is None or ar._canon(local["actual_reading"]) != reading or member.get("actual") != reading:
                            raise ValueError(f"舊 actual row {number} A 目前 occurrence-local 判定不足以補證")
                    attestations["actual_groups"][str(number)] = {
                        "old_group_snapshot": str(row.get("group_snapshot") or ""),
                        "current_group_snapshot": group["group_snapshot"],
                        "group_id": group["group_id"], "reading": reading,
                        "member_occurrence_ids": [item["occurrence_id"] for item in members],
                        "override_rows": [copy.deepcopy(override_by_key[ar._override_key_from_entry(item)])
                                          for item in members],
                    }
        write_excel_content_proof(workbook, source_dir, manifest, current_db, kind=kind,
                                  snapshot_db=snapshot_db, attestations=attestations)
        _verify_excel_import_snapshot(xlsx, _original_xlsx, _snapshot_sha)
        descriptor, temporary = tempfile.mkstemp(prefix=f".{_original_xlsx.stem}.proof-", suffix=".xlsx", dir=_original_xlsx.parent)
        os.close(descriptor)
        workbook.save(temporary)
        destination = _original_xlsx.with_name(f"{_original_xlsx.stem}_跨電腦_{uuid.uuid4().hex[:8]}.xlsx")
        os.replace(temporary, destination)
        temporary = None
        return destination
    finally:
        workbook.close()
        if temporary is not None:
            Path(temporary).unlink(missing_ok=True)


def prepare_portable_project(output_dir: Path) -> Path:
    """Capture source-page proof while original bytes are still available."""
    import standalone_proofread as sp

    output_dir = Path(output_dir).resolve()
    manifest, _ = _load_project(output_dir)
    _actual_transfer_state(output_dir, manifest)
    payload = _proof_payload(output_dir, manifest)
    proof = {**payload, "integrity_sha256": hashlib.sha256(_canonical(payload)).hexdigest()}
    path = output_dir / PROOF_FILE
    sp.json_save(path, proof)
    return path


def _load_proof(output_dir: Path, manifest: Mapping[str, Any]):
    import standalone_proofread as sp

    path = output_dir / PROOF_FILE
    if path.exists():
        proof = sp.json_load_strict(path)
        payload = {key: value for key, value in proof.items() if key != "integrity_sha256"}
        if (proof.get("version") != PROOF_VERSION
                or proof.get("integrity_sha256") != hashlib.sha256(_canonical(payload)).hexdigest()
                or proof.get("manifest_integrity_sha256") != manifest.get("manifest_integrity_sha256")
                or proof.get("session_id") != manifest.get("session_id")):
            raise ValueError("跨電腦 PDF 內容證據損壞或不屬於來源工作階段")
        expected = [(info.get("pdf_sha256"), info.get("pdf_name")) for info in manifest.get("pdfs", [])]
        observed = [(info.get("pdf_sha256"), info.get("pdf_name")) for info in proof.get("pdfs", [])]
        if observed != expected:
            raise ValueError("跨電腦 PDF 內容證據與來源 PDF 清單不符")
        if any("candidate_text_sha256" not in page
               for item in proof["pdfs"] for page in item.get("pages", [])):
            # A legacy visual-only proof cannot establish the candidate text
            # of different PDF bytes. Rebuild from the exact sealed source only.
            upgraded = copy.deepcopy(proof)
            for saved, info in zip(upgraded["pdfs"], manifest["pdfs"]):
                try:
                    original = _source_pdf(manifest, info, output_dir)
                except FileNotFoundError:
                    # _match_pdfs may still prove byte-for-byte identity with
                    # the target; different bytes require this missing layer.
                    return proof
                _, fresh = _pdf_content_snapshot(original, expected_sha256=str(info["pdf_sha256"]))
                if not _pages_visual_equivalent(saved["pages"], fresh):
                    raise ValueError("舊 PDF 內容證據與來源原頁不符；未補證")
                saved["pages"] = fresh
            return upgraded
        return proof
    # Source bytes may still be present without a previously exported proof.
    return _proof_payload(output_dir, manifest)


def _proof_from_explicit_local(output_dir: Path, manifest: Mapping[str, Any],
                               local_pdf: Path, local_pages):
    infos = manifest.get("pdfs", [])
    observed_sha, observed_pages = _pdf_content_snapshot(local_pdf)
    if observed_pages != local_pages:
        raise ValueError("明確指定的本地 PDF 在頁面內容取證後已變動")
    try:
        return _load_proof(output_dir, manifest)
    except FileNotFoundError:
        if len(infos) != 1 or observed_sha != infos[0].get("pdf_sha256"):
            raise
        # The selected file has the sealed source bytes; no path search or
        # inferred equivalence is needed to obtain its complete page proof.
        return {"pdfs": [{"pdf_sha256": infos[0]["pdf_sha256"],
                           "pdf_name": infos[0]["pdf_name"], "pages": local_pages}]}


def _pages_visual_equivalent(source_pages, target_pages):
    if len(source_pages) != len(target_pages):
        return False
    for source, target in zip(source_pages, target_pages):
        # Rotation and its matrix describe PDF storage. Reviewed positions are
        # separately transformed into the displayed page coordinates.
        fields = ("width", "height", "pixels_width", "pixels_height", "rgb_sha256")
        if any(source.get(field) != target.get(field) for field in fields):
            return False
    return True


def _pages_equivalent(source_pages, target_pages):
    return (_pages_visual_equivalent(source_pages, target_pages)
            and all(source.get("candidate_text_sha256")
                    and source.get("candidate_text_sha256") == target.get("candidate_text_sha256")
                    for source, target in zip(source_pages, target_pages)))


def _match_pdfs(proof: Mapping[str, Any], source_manifest: Mapping[str, Any],
                target_manifest: Mapping[str, Any], target_dir: Path):
    source_pdfs = list(proof["pdfs"])
    target_pdfs = list(target_manifest.get("pdfs", []))
    if len(source_pdfs) != len(target_pdfs):
        raise ValueError("PDF 數量不同；不支援部分教材成果移轉")
    rendered = []
    for target in target_pdfs:
        path = _source_pdf(target_manifest, target, target_dir)
        _, pages = _pdf_content_snapshot(path, expected_sha256=str(target["pdf_sha256"]))
        rendered.append((target, pages))
    mapping = {}
    geometry = {}
    for source in source_pdfs:
        missing_text = any("candidate_text_sha256" not in page for page in source["pages"])
        candidates = []
        for target, pages in rendered:
            exact_bytes = source["pdf_sha256"] == target["pdf_sha256"]
            source_pages = pages if missing_text and exact_bytes else source["pages"]
            if (not missing_text or exact_bytes) and _pages_visual_equivalent(source["pages"], pages) \
                    and _pages_equivalent(source_pages, pages):
                candidates.append((target, pages, source_pages))
        if len(candidates) != 1:
            if missing_text and not any(source["pdf_sha256"] == target["pdf_sha256"]
                                        for target, _ in rendered):
                raise ValueError("舊 PDF 內容證據缺少校對文字層；不同 SHA 須由 A 原 PDF 精確 SHA 補證")
            raise ValueError(
                f"PDF 頁面內容/圖片/文字/注音字形或頁面位置無法唯一對應：{source['pdf_name']}；"
                "請確認是同版完整教材"
            )
        target, target_pages, source_pages = candidates[0]
        rendered.remove((target, target_pages))
        mapping[source["pdf_sha256"]] = target["pdf_sha256"]
        geometry[source["pdf_sha256"]] = (source_pages, target_pages)
    return mapping, geometry


def _verify_mapped_target_pdfs(target_dir: Path, target_manifest, pdf_map, geometry) -> None:
    """Bind current B bytes to B's already verified full-page snapshot at each gate."""
    by_sha = {str(info["pdf_sha256"]): info for info in target_manifest.get("pdfs", [])}
    if len(by_sha) != len(target_manifest.get("pdfs", [])) or len(pdf_map) != len(by_sha):
        raise ValueError("目標 PDF 內容對應清單已變動；未匯入判定")
    for source_sha, target_sha in pdf_map.items():
        info = by_sha.get(target_sha)
        expected = geometry.get(source_sha) if geometry is not None else None
        if info is None or (geometry is not None and expected is None) \
                or (geometry is None and source_sha != target_sha):
            raise ValueError("目標 PDF 內容對應已變動；未匯入判定")
        try:
            path = _source_pdf(target_manifest, info, target_dir)
            if _sha(path) != target_sha:
                raise ValueError("已驗證的目標 PDF bytes 已變動")
        except (FileNotFoundError, ValueError) as exc:
            raise ValueError(f"目標 PDF 在內容對應後已變動；未匯入判定：{info['pdf_name']}") from exc


def _anchor(entry: Mapping[str, Any], pages):
    """Position is mandatory; text alone cannot select an occurrence."""
    coordinates = []
    for field in ("x0", "y0", "x1", "y1"):
        value = entry.get(field)
        if value is None or isinstance(value, bool):
            raise ValueError(f"occurrence 位置缺少 {field}")
        coordinates.append(round(float(value), 2))
    page = int(entry.get("physical_page") or 0)
    if page < 1 or coordinates[2] <= coordinates[0] or coordinates[3] <= coordinates[1]:
        raise ValueError("occurrence 頁碼或 bbox 無效")
    if page > len(pages) or len(pages[page - 1].get("display_matrix", [])) != 6:
        raise ValueError("occurrence 顯示頁面座標轉換資料缺失")
    displayed = fitz.Rect(coordinates) * fitz.Matrix(*pages[page - 1]["display_matrix"])
    if displayed.is_empty:
        raise ValueError("occurrence 顯示位置無效")
    return (page, *(round(value, 2) for value in displayed), str(entry.get("char") or ""))


def _map_reviews(source_manifest: Mapping[str, Any], target_manifest: Mapping[str, Any], pdf_map, geometry):
    source_by_target = {target_sha: source_sha for source_sha, target_sha in pdf_map.items()}
    target_by_anchor = {}
    for entry in target_manifest.get("records", []):
        target_sha = str(entry.get("pdf_sha256") or "")
        source_sha = source_by_target.get(target_sha)
        if source_sha is None:
            raise ValueError("目標 occurrence PDF 無來源內容對應")
        key = (target_sha, _anchor(entry, geometry[source_sha][1]))
        if key in target_by_anchor:
            raise ValueError(f"目標位置不唯一：{key}")
        target_by_anchor[key] = entry
    mapped = {}
    used = set()
    for entry in source_manifest.get("records", []):
        source_sha = str(entry.get("pdf_sha256") or "")
        if source_sha not in pdf_map:
            raise ValueError("來源 occurrence PDF 無目標內容對應")
        key = (pdf_map[source_sha], _anchor(entry, geometry[source_sha][0]))
        target = target_by_anchor.get(key)
        if target is None:
            raise ValueError(f"來源 occurrence 無法對應目標位置：{entry.get('review_id')}")
        if target["review_id"] in used:
            raise ValueError("多個來源 occurrence 對應同一目標位置")
        used.add(target["review_id"])
        mapped[entry["review_id"]] = (entry, target)
    if len(mapped) != len(target_by_anchor):
        raise ValueError("來源與目標 occurrence 數量/位置不同；不支援部分移轉")
    return mapped


def _event_payload(event: Mapping[str, Any]):
    return {key: value for key, value in event.items()
            if key not in {"portability_source", "portability_duplicate_sources",
                           "portability_conflict_resolution"}}


def _excel_expected_decision(event: Mapping[str, Any]):
    import standalone_proofread as sp

    fields = ("action", "expected_set", "expected_evidence", "context_evidence",
              "exclusion_reason", "exclusion_evidence", "note", *sp.CONFIRMATION_GATES)
    return {key: event.get(key) for key in fields}


def _historical_import_replay(conflicts, identity, source_event, target_event, *, decision_payload):
    """Recognize only exact inputs retained in already validated conflict history.

    Source-side project events retain their original coordinates; target-side
    events retain the transferred payload and any equal-decision provenance.
    Identity alone is insufficient because a project's event can later change.
    """
    for conflict in conflicts:
        if conflict.get("target_review_id") != identity["target_review_id"]:
            continue
        for side, incoming in (("source", source_event), ("target", target_event)):
            saved = conflict.get(f"{side}_event")
            if not isinstance(saved, dict):
                continue
            known = [conflict.get(f"{side}_identity"), saved.get("portability_source"),
                     *saved.get("portability_duplicate_sources", [])]
            if identity in known and decision_payload(saved) == decision_payload(incoming):
                return True
    return False


def _restore_exact_file(path: Path, original: bytes | None) -> None:
    if original is None:
        path.unlink(missing_ok=True)
        return
    descriptor, temporary = tempfile.mkstemp(prefix=f".{path.name}.restore-", dir=path.parent)
    try:
        with os.fdopen(descriptor, "wb") as stream:
            stream.write(original)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        Path(temporary).unlink(missing_ok=True)


def _document_digest(value: Any) -> str:
    return hashlib.sha256(_canonical(value)).hexdigest()


def _read_presentation_receipts(target_dir: Path):
    """Keep strict parsed documents and rollback bytes from the same read."""
    import standalone_proofread as sp
    snapshots = {}
    for name in PRESENTATION_RECEIPTS:
        path = target_dir / name
        try:
            raw = path.read_bytes()
        except FileNotFoundError:
            raw = None
        snapshots[name] = (raw, sp.json_parse_strict(raw, path) if raw is not None else None)
    return snapshots


def _receipt_snapshot_shas(snapshots):
    return {name: hashlib.sha256(raw).hexdigest() if raw is not None else ""
            for name, (raw, _) in snapshots.items()}


def _check_receipt_shas(target_dir: Path, expected):
    for name, digest in expected.items():
        path = target_dir / name
        if (_sha(path) if path.exists() else "") != digest:
            raise ValueError(f"來源回執於判定期間變動；保留新紀錄，未覆寫：{name}")


def _presentation_receipt_plan(target_dir: Path, planned: Mapping[str, Any], *, snapshots=None):
    """Freeze receipt backup, candidate and digest from one retained snapshot."""
    snapshots = _read_presentation_receipts(target_dir) if snapshots is None else snapshots
    receipts = {}
    for name, (old, current) in snapshots.items():
        desired = planned.get(name, current)
        receipts[name] = {
            "before_bytes": base64.b64encode(old).decode("ascii") if old is not None else None,
            "before_digest": _document_digest(current) if current is not None else None,
            "after_digest": _document_digest(desired) if desired is not None else None,
        }
    _check_receipt_shas(target_dir, _receipt_snapshot_shas(snapshots))
    return receipts


def _new_presentation_marker(target_dir: Path, manifest, before_db_sha: str,
                             candidate_db, planned_receipts: Mapping[str, Any], *,
                             phase: str, export_pending: bool = False, _receipt_snapshots=None):
    if phase not in {"PROJECT_IMPORT", "NEW_PROJECT_FINAL", "SAME_SESSION_EXPECTED",
                     "CROSS_SESSION_EXPECTED"}:
        raise ValueError("未知專案發布階段")
    import standalone_proofread as sp
    _actual_transfer_state(target_dir, manifest)
    if _sha(target_dir / "人工判定資料庫.json") != before_db_sha:
        raise ValueError("判定資料庫於建立發布計畫前變動；未寫入")
    live_manifest, manifest_sha = sp.json_load_snapshot(target_dir / "校對工作階段.json")
    if live_manifest != manifest:
        raise ValueError("工作階段已變動；未建立不一致的發布計畫")
    payload = {
        "status": "PRESENTATION_PENDING", "plan_version": PRESENTATION_PLAN_VERSION,
        "owner_token": uuid.uuid4().hex, "phase": phase,
        "manifest_file_sha256": manifest_sha,
        "manifest_integrity_sha256": manifest["manifest_integrity_sha256"],
        "db_before_sha256": before_db_sha,
        "db_after_digest": _document_digest(candidate_db),
        "receipts": _presentation_receipt_plan(target_dir, planned_receipts, snapshots=_receipt_snapshots),
        "export_pending": bool(export_pending),
    }
    return {**payload, "plan_integrity_sha256": _document_digest(payload)}


def _validated_presentation_marker(value):
    import standalone_proofread as sp
    if not isinstance(value, dict) or set(value) != {
        "status", "plan_version", "owner_token", "phase", "manifest_file_sha256",
        "manifest_integrity_sha256", "db_before_sha256", "db_after_digest",
        "receipts", "export_pending", "plan_integrity_sha256",
    }:
        raise ValueError("未完成專案缺少專用發布恢復計畫；不可猜測清除標記")
    payload = {key: item for key, item in value.items() if key != "plan_integrity_sha256"}
    if (value["status"] != "PRESENTATION_PENDING"
            or value["plan_version"] != PRESENTATION_PLAN_VERSION
            or value["phase"] not in {"PROJECT_IMPORT", "NEW_PROJECT_FINAL", "SAME_SESSION_EXPECTED",
                                       "CROSS_SESSION_EXPECTED"}
            or not isinstance(value["owner_token"], str) or len(value["owner_token"]) != 32
            or any(char not in "0123456789abcdef" for char in value["owner_token"])
            or type(value["export_pending"]) is not bool
            or not isinstance(value["receipts"], dict)
            or set(value["receipts"]) != set(PRESENTATION_RECEIPTS)
            or any(not isinstance(value[name], str) or len(value[name]) != 64
                   for name in ("manifest_file_sha256", "manifest_integrity_sha256",
                                "db_before_sha256", "db_after_digest", "plan_integrity_sha256"))
            or value["plan_integrity_sha256"] != _document_digest(payload)):
        raise ValueError("未完成專案發布計畫格式或完整性無法驗證")
    for item in value["receipts"].values():
        if (not isinstance(item, dict)
                or set(item) != {"before_bytes", "before_digest", "after_digest"}
                or any(digest is not None and (not isinstance(digest, str) or len(digest) != 64)
                       for digest in (item["before_digest"], item["after_digest"]))):
            raise ValueError("未完成專案來源回執恢復材料無法驗證")
        encoded = item["before_bytes"]
        if encoded is None:
            if item["before_digest"] is not None:
                raise ValueError("未完成專案回執原狀不一致")
        else:
            try:
                original = base64.b64decode(encoded, validate=True)
                if _document_digest(sp.json_parse_strict(original, Path("presentation receipt backup"))) != item["before_digest"]:
                    raise ValueError("未完成專案回執原文與摘要不符")
            except (ValueError, TypeError) as exc:
                raise ValueError("未完成專案回執原文無法驗證") from exc
    return value


def _presentation_current_receipts(target_dir: Path):
    import standalone_proofread as sp

    return {name: (_document_digest(sp.json_load_strict(target_dir / name))
                   if (target_dir / name).exists() else None)
            for name in PRESENTATION_RECEIPTS}


def resume_portable_project(output_dir: Path, *, _expected_marker_sha: str | None = None):
    """Redo only a fully bound committed presentation, or undo precommit receipts."""
    import standalone_proofread as sp

    output_dir = Path(output_dir).resolve()
    root = sp.project_actual_evidence_root(output_dir)
    with sp.project_delivery_lock(root):
        marker_path = output_dir / INCOMPLETE_FILE
        if _expected_marker_sha is not None:
            _assert_actual_marker_owner(marker_path, _expected_marker_sha)
        raw_marker, marker_sha = sp.json_load_snapshot(marker_path)
        marker = _validated_presentation_marker(raw_marker)
        _assert_actual_marker_owner(marker_path, _expected_marker_sha or marker_sha)
        manifest_path = output_dir / "校對工作階段.json"
        if _sha(manifest_path) != marker["manifest_file_sha256"]:
            raise ValueError("未完成專案封印檔已變動；未恢復或清標記")
        manifest = sp.json_load_strict(manifest_path)
        sp.validate_manifest_integrity(manifest)
        sp.validate_output_artifact_hashes(manifest)
        if manifest["manifest_integrity_sha256"] != marker["manifest_integrity_sha256"]:
            raise ValueError("未完成專案封印與發布計畫不符")
        db_path = output_dir / "人工判定資料庫.json"
        raw_db = sp.json_load_strict(db_path)
        if not isinstance(raw_db, dict):
            raise ValueError("未完成專案判定資料庫格式錯誤")
        db = sp.normalize_db(raw_db)
        db_after = _document_digest(db) == marker["db_after_digest"]
        current_receipts = _presentation_current_receipts(output_dir)
        receipt_after = all(current_receipts[name] == marker["receipts"][name]["after_digest"]
                            for name in PRESENTATION_RECEIPTS)
        if db_after and receipt_after:
            _load_project(output_dir, allow_incomplete=True, allow_unresolved_conflict=True)
            _actual_transfer_state(output_dir, manifest)
            _assert_actual_marker_owner(marker_path, marker_sha)
            published_report = _publish_portable_outputs(output_dir, manifest, db)
            if marker["export_pending"]:
                if any(sp.needs_expected_review(row) for row in sp.materialize_ledger(manifest, db)):
                    sp.export_pending_for_gpt(output_dir, _allow_internal_portable=True)
            if (sp.json_load_strict(marker_path) != marker
                    or _sha(manifest_path) != marker["manifest_file_sha256"]
                    or _document_digest(sp.normalize_db(sp.json_load_strict(db_path))) != marker["db_after_digest"]
                    or _presentation_current_receipts(output_dir) != current_receipts):
                raise ValueError("發布期間來源資料或未完成標記已由其他交易變動；保留標記")
            _assert_actual_marker_owner(marker_path, marker_sha)
            marker_path.unlink()
            return {"status": "PRESENTATION_RECOVERED", "report": str(published_report)}
        if (_sha(db_path) == marker["db_before_sha256"]
                and all(current_receipts[name] in {
                    marker["receipts"][name]["before_digest"],
                    marker["receipts"][name]["after_digest"]
                } for name in PRESENTATION_RECEIPTS)):
            if sp.json_load_strict(marker_path) != marker:
                raise ValueError("未完成標記已由其他交易變動；未回滾")
            for name in PRESENTATION_RECEIPTS:
                _assert_actual_marker_owner(marker_path, marker_sha)
                encoded = marker["receipts"][name]["before_bytes"]
                _restore_exact_file(output_dir / name,
                                    base64.b64decode(encoded) if encoded is not None else None)
            if sp.json_load_strict(marker_path) != marker:
                raise ValueError("回滾期間未完成標記已變動；保留現場")
            _assert_actual_marker_owner(marker_path, marker_sha)
            marker_path.unlink()
            return {"status": "PRECOMMIT_ROLLED_BACK"}
        raise ValueError("未完成專案 DB／來源回執與專用恢復計畫不符；保留資料及標記")


def _require_transferable_expected_lane(action: str, target, prior_event, label: str) -> None:
    """An imported expected event cannot become B's independent resolver truth."""
    import standalone_proofread as sp

    if (action in {"補建expected證據", "解決expected證據"}
            and not prior_event and sp.infer_expected_status(target) == "RESOLVED"):
        raise ValueError(f"{label}目標已有正式 expected；保留雙方獨立證據，須在目標明確裁決")


def _require_mapped_textbook_context(source, target, label: str) -> None:
    """An occurrence anchor cannot substitute for B's independently built text context."""
    source_record = source.get("source_record") or {}
    target_record = target.get("source_record") or {}
    if any(str(source_record.get(field) or "") != str(target_record.get(field) or "")
           for field in ("所在行", "局部詞境")):
        raise ValueError(f"{label}教材詞境不同，不能沿用判定")


@_stable_excel_import
def import_expected_excel(target_dir: Path, xlsx: Path, *, dry_run: bool = False,
                          _original_xlsx: Path, _snapshot_sha: str):
    """Import filled expected rows after full-page and occurrence-local mapping."""
    import standalone_proofread as sp

    target_dir, xlsx = Path(target_dir).resolve(), Path(xlsx).resolve()
    target_manifest, target_db = _load_project(target_dir)
    metadata = sp.workbook_metadata(xlsx, "匯入中繼資料")
    source_manifest, _, source_ledger, reviews, pdf_map, exported_at, _, geometry = load_excel_content_proof(
        xlsx, target_dir, target_manifest, kind="expected",
        workbook_session=str(metadata.get("session_id") or ""))
    if str(metadata.get("expected_asset_fingerprint") or "") != str(source_manifest.get("expected_asset_fingerprint") or ""):
        raise ValueError("Excel expected fingerprint 與來源封印不符")
    allowed = {str(target_manifest.get("expected_asset_fingerprint") or ""),
               *map(str, target_manifest.get("compatible_expected_asset_fingerprints") or [])}
    if str(metadata.get("expected_asset_fingerprint") or "") not in allowed:
        raise ValueError("Excel 與目標 expected 資產 fingerprint 不相容")
    required = ["occurrence_id", "review_id", "review_snapshot", "state", "action",
                "proposed_expected_set", "proposed_expected_evidence", "proposed_context_evidence",
                "exclusion_reason", "exclusion_evidence", *sp.CONFIRMATION_GATES, "note"]
    rows = sp.workbook_rows(xlsx, "待判定候選", required)
    _verify_excel_import_snapshot(xlsx, _original_xlsx, _snapshot_sha)
    sp.assert_unique_ids(rows, "occurrence_id")
    sp.assert_unique_ids(rows, "review_id")
    source_by_id = {entry["review_id"]: entry for entry in source_ledger}
    target_ledger = sp.materialize_ledger(target_manifest, target_db)
    target_by_id = {entry["review_id"]: entry for entry in target_ledger}
    target_base_by_id = {entry["review_id"]: entry for entry in
                         sp.materialize_ledger(target_manifest, sp.normalize_db({}))}
    actions = []
    skipped = 0
    for number, row in enumerate(rows, 2):
        action = str(row.get("action") or "").strip()
        if not action or action == "保留待人工":
            skipped += 1
            continue
        source = source_by_id.get(str(row.get("review_id") or ""))
        if source is None or str(row.get("occurrence_id") or "") != source["occurrence_id"]:
            raise ValueError(f"Excel row {number} 來源 occurrence/review 識別無法驗證")
        if (str(row.get("review_snapshot") or "") != sp._review_snapshot(source)
                or not sp._expected_target_unchanged(row, source)):
            raise ValueError(f"Excel row {number} 來源 snapshot/文字位置已變更")
        pair = reviews.get(source["review_id"])
        if pair is None:
            raise ValueError(f"Excel row {number} 來源位置無法映射")
        target = target_by_id[pair[1]["review_id"]]
        prior_event = target_db["events"].get(target["review_id"]) or {}
        target_for_validation = target_base_by_id[target["review_id"]] if prior_event else target
        prior_sources = [prior_event.get("portability_source"),
                         *prior_event.get("portability_duplicate_sources", [])]
        if any(isinstance(item, dict) and item.get("source_excel_sha256") == _snapshot_sha
               and item.get("review_id") == source["review_id"]
               and item.get("source_excel_row") == number for item in prior_sources):
            skipped += 1
            continue
        _require_mapped_textbook_context(source, target, f"Excel row {number} ")
        expected_only = action in {"補建expected證據", "解決expected證據"}
        if action not in sp.DECISIONS:
            raise ValueError(f"Excel row {number} action 無效")
        if not expected_only:
            semantic = ("state", "actual", "actual_evidence", "expected_set",
                        "expected_evidence", "context_evidence")
            if any(source.get(key) != target_for_validation.get(key) for key in semantic):
                raise ValueError(f"Excel row {number} 目標 actual/expected 證據或狀態不同；須重新核對")
            if action == "確認現版差異" and any(
                    source.get(key) != target.get(key) for key in
                    ("actual", "actual_evidence", "expected_set", "expected_evidence", "context_evidence")):
                raise ValueError(f"Excel row {number} 目標當前 actual/expected 證據不同；六個確認 gate 不可沿用")
        _require_transferable_expected_lane(action, target, prior_event, f"Excel row {number} ")
        event = {
            "action": action, "expected_set": str(row.get("proposed_expected_set") or "").strip(),
            "expected_evidence": str(row.get("proposed_expected_evidence") or "").strip(),
            "context_evidence": str(row.get("proposed_context_evidence") or "").strip(),
            "exclusion_reason": str(row.get("exclusion_reason") or "").strip(),
            "exclusion_evidence": str(row.get("exclusion_evidence") or "").strip(),
            "note": str(row.get("note") or "").strip(),
            "source": "跨電腦 Excel expected 判定轉接", "updated_at": exported_at,
            "_safe_expected_rebase": expected_only and sp._review_snapshot(source) != sp._review_snapshot(target),
        }
        for gate in sp.CONFIRMATION_GATES:
            value = str(row.get(gate) or "").strip().upper()
            if value not in {"", "Y", "N"}:
                raise ValueError(f"Excel row {number} gate {gate} 必須是 Y/N")
            event[gate] = value == "Y"
        for entry, label in ((source, "來源"), (target_for_validation, "目標")):
            replay = sp._apply_review_event(entry, event)
            if replay.get("review_event_replay_status"):
                raise ValueError(f"Excel row {number} {label}判定失效：{replay['review_event_replay_status']}")
        if expected_only and target_manifest.get("expected_asset_fingerprint"):
            # Full-page/text mapping and both original replay checks above are
            # still mandatory. Bind the validated transaction to B's own baseline.
            event = sp.bind_expected_resolution(target_base_by_id[target["review_id"]], event, target_manifest)
        identity = _source_identity(source_manifest, source, source["review_id"],
                                    target_manifest, target, event)
        identity["source_excel_sha256"] = _snapshot_sha
        identity["source_excel_row"] = number
        identity["source_excel_decision"] = copy.deepcopy(row)
        actions.append((target, event, identity))
    if dry_run:
        with sp.project_delivery_lock(sp.project_actual_evidence_root(target_dir)):
            _verify_mapped_target_pdfs(target_dir, target_manifest, pdf_map, geometry)
        return len(actions), skipped, target_dir / "注音校對_最終報告.xlsx"

    written_candidate_sha = None
    with sp.project_delivery_lock(sp.project_actual_evidence_root(target_dir)):
        _verify_excel_import_snapshot(xlsx, _original_xlsx, _snapshot_sha)
        marker_path = target_dir / INCOMPLETE_FILE
        if marker_path.exists():
            raise ValueError("Excel 匯入期間目標出現未完成標記；未寫入任何判定，先完成專用恢復")
        live_manifest = sp.json_load_strict(target_dir / "校對工作階段.json")
        sp.validate_manifest_integrity(live_manifest)
        if live_manifest["manifest_integrity_sha256"] != target_manifest["manifest_integrity_sha256"]:
            raise ValueError("Excel 匯入期間目標封印已變動")
        _verify_mapped_target_pdfs(target_dir, live_manifest, pdf_map, geometry)
        db_path = target_dir / "人工判定資料庫.json"
        raw_db, before = sp.json_load_snapshot(db_path)
        if not isinstance(raw_db, dict):
            raise ValueError("目標判定資料庫根節點不是物件")
        live_db = sp.normalize_db(raw_db)
        if live_db != target_db:
            raise ValueError("Excel 匯入期間目標判定資料庫已變動")
        pending_actual = actual_excel_conflict_state(target_dir, live_manifest, live_db)
        if pending_actual:
            raise ValueError(f"Excel 匯入期間目標出現未裁決 actual 衝突；未寫入判定：{pending_actual}")
        receipt_snapshots = _read_presentation_receipts(target_dir)
        receipt_shas = _receipt_snapshot_shas(receipt_snapshots)
        receipt_path = target_dir / CONFLICT_FILE
        prior_conflicts = []
        if receipt_snapshots[CONFLICT_FILE][0] is not None:
            # The receipt and target DB are one logical state. Validate both
            # under the same project lock before extending conflict history.
            if validate_conflict_state(target_dir, live_manifest, target_db):
                raise ValueError("Excel 匯入期間目標出現未裁決衝突，不能轉入新判定")
            prior_conflicts = copy.deepcopy(receipt_snapshots[CONFLICT_FILE][1]["conflicts"])
        candidate = copy.deepcopy(target_db)
        imported, duplicates, conflicts = 0, 0, []
        for target, event, identity in actions:
            review_id = target["review_id"]
            if _historical_import_replay(prior_conflicts, identity, event, event,
                                         decision_payload=_excel_expected_decision):
                duplicates += 1
                continue
            current = candidate["events"].get(review_id)
            if current is not None:
                known = [current.get("portability_source"),
                         *current.get("portability_duplicate_sources", [])]
                if identity in known:
                    duplicates += 1
                    continue
                if _excel_expected_decision(current) == _excel_expected_decision(event):
                    current.setdefault("portability_duplicate_sources", []).append(identity)
                    duplicates += 1
                    continue
                conflicts.append({"target_review_id": review_id,
                                  "source_review_id": identity["review_id"],
                                  "source_session_id": source_manifest["session_id"],
                                  "source_pdf_sha256": identity["pdf_sha256"],
                                  "source_occurrence_id": identity["occurrence_id"],
                                  "target_session_id": target_manifest["session_id"],
                                  "target_pdf_sha256": target["pdf_sha256"],
                                  "target_occurrence_id": target["occurrence_id"],
                                  "source_identity": identity,
                                  "target_identity": copy.deepcopy(current.get("portability_source")),
                                  "source_event": {**event, "portability_source": identity},
                                  "target_event": copy.deepcopy(current)})
                candidate["events"].pop(review_id)
                continue
            candidate["events"][review_id] = {**event, "portability_source": identity}
            imported += 1
        ledger = sp.materialize_ledger(target_manifest, candidate)
        if any(row.get("review_event_replay_status") for row in ledger):
            raise ValueError("Excel 匯入後判定無法在目標安全重播")
        if imported or duplicates or conflicts:
            receipt_path = target_dir / CONFLICT_FILE
            presentation_paths = [target_dir / name for name in
                                  ("待人工確認.json", "注音校對_最終報告.xlsx", "pipeline_status.json")]
            original_files = {path: path.read_bytes() if path.exists() else None
                              for path in [db_path, receipt_path, marker_path, *presentation_paths]}
            marker_before = marker_path.read_bytes() if marker_path.exists() else None
            receipt_before = receipt_snapshots[CONFLICT_FILE][0]
            original_files[receipt_path] = receipt_before
            planned_conflicts = ({"version": 1, "conflicts": [*prior_conflicts, *conflicts]}
                                 if conflicts else None)
            owned_marker = _new_presentation_marker(
                target_dir, live_manifest, before, candidate,
                {CONFLICT_FILE: planned_conflicts} if planned_conflicts is not None else {},
                phase="CROSS_SESSION_EXPECTED", _receipt_snapshots=receipt_snapshots)
            def validate_write():
                _actual_transfer_state(target_dir, live_manifest)
                _check_receipt_shas(target_dir, receipt_shas)
                _verify_mapped_target_pdfs(target_dir, live_manifest, pdf_map, geometry)
            validate_write()
            db_committed = False
            try:
                marker_sha = sp.json_save(marker_path, owned_marker, expected_sha256="")
                _assert_actual_marker_owner(marker_path, marker_sha)
                validate_write()
                if conflicts:
                    receipt_shas[CONFLICT_FILE] = sp.json_save(
                        receipt_path, planned_conflicts, expected_sha256=receipt_shas[CONFLICT_FILE])
                if sp.json_load_strict(marker_path) != owned_marker:
                    raise ValueError("Excel expected 匯入標記已由其他交易變動；未寫入判定")
                validate_write()
                sp.json_save(db_path, candidate, expected_sha256=before)
                db_committed = True
                written_candidate_sha = _sha(db_path)
            except Exception as exc:
                if not db_committed and _sha(db_path) == before:
                    if not marker_path.exists() or sp.json_load_strict(marker_path) != owned_marker:
                        raise ValueError("Excel expected 寫入失敗且未完成標記已由其他交易變動；保留衝突紀錄與現場待查核") from exc
                    _check_receipt_shas(target_dir, receipt_shas)
                    _restore_exact_file(receipt_path, receipt_before)
                    _restore_exact_file(marker_path, marker_before)
                raise
    if imported or duplicates or conflicts:
        try:
            resume_portable_project(target_dir, _expected_marker_sha=marker_sha)
        except Exception:
            with sp.project_delivery_lock(sp.project_actual_evidence_root(target_dir)):
                if (_sha(target_dir / "人工判定資料庫.json") != written_candidate_sha
                        or not (target_dir / INCOMPLETE_FILE).exists()
                        or sp.json_load_strict(target_dir / INCOMPLETE_FILE) != owned_marker):
                    raise ValueError("Excel expected 發布失敗且目標資料或未完成標記已由其他交易變動；保留現場待查核")
                _check_receipt_shas(target_dir, receipt_shas)
                # Restore the original user views, conflict receipt and DB;
                # leave the fail-closed marker in place until every restore succeeds.
                for path in [*presentation_paths, target_dir / CONFLICT_FILE,
                             target_dir / "人工判定資料庫.json"]:
                    _restore_exact_file(path, original_files[path])
                _restore_exact_file(target_dir / INCOMPLETE_FILE,
                                    original_files[target_dir / INCOMPLETE_FILE])
            raise
    if conflicts:
        raise ValueError(f"Excel expected 同位置有 {len(conflicts)} 筆不同判定；兩份來源已保留於 {CONFLICT_FILE}，目標位置已標待人工裁決（不是寫入回滾）")
    return imported, skipped + duplicates, target_dir / "注音校對_最終報告.xlsx"


@_stable_excel_import
def import_actual_excel(target_dir: Path, xlsx: Path, *,
                        _original_xlsx: Path, _snapshot_sha: str):
    """Carry source visual decisions as local occurrence evidence, never a vote."""
    import actual_review as ar
    import standalone_proofread as sp
    from global_glyph_promotion import (
        direct_visual_project_transaction, load_promotion_outbox,
        post_commit_recovery_plan_from_results, PROJECT_TRANSACTION_FILE,
    )

    target_dir, xlsx = Path(target_dir).resolve(), Path(xlsx).resolve()
    target_manifest, target_db = _load_project(target_dir, allow_actual_excel_conflict=True)
    existing_actual_conflicts = actual_excel_conflict_state(target_dir, target_manifest, target_db)
    target_actual_state = _actual_transfer_state(target_dir, target_manifest)
    if target_actual_state["status"] != "NONE":
        raise ValueError("目標尚有未核對的 actual 暫存/來源轉送；先完成本地恢復")
    metadata, rows = ar._load_sheet_rows(xlsx, "actual待判定")
    _verify_excel_import_snapshot(xlsx, _original_xlsx, _snapshot_sha)
    from openpyxl import load_workbook
    workbook = load_workbook(xlsx, read_only=True)
    try:
        legacy = EXCEL_PROOF_SHEET not in workbook.sheetnames
    finally:
        workbook.close()
    if legacy:
        with sp.project_delivery_lock(sp.project_actual_evidence_root(target_dir)):
            if _legacy_actual_already_imported(target_dir, target_manifest, target_db, xlsx, rows):
                _verify_excel_import_snapshot(xlsx, _original_xlsx, _snapshot_sha)
                live_manifest, live_db = _load_project(target_dir)
                if live_manifest != target_manifest or live_db != target_db:
                    raise ValueError("目標封印或判定於舊 actual Excel 核對期間變動")
                for info in live_manifest["pdfs"]:
                    _source_pdf(live_manifest, info, target_dir)
                return sp.ActualImportResult(0, 0, target_dir / "注音校對_最終報告.xlsx",
                                             {"project_actual_commit": "NO_CHANGES",
                                              "global_promotion_delivery": "NOT_ATTEMPTED"})
    source_manifest, _, source_ledger, reviews, pdf_map, _, proof_payload, geometry = load_excel_content_proof(
        xlsx, target_dir, target_manifest, kind="actual",
        workbook_session=str(metadata.get("session_id") or ""))
    if (str(metadata.get("actual_review_schema_version")) != ar.ACTUAL_REVIEW_SCHEMA_VERSION
            or str(metadata.get("session_id")) != str(source_manifest["session_id"])):
        raise ValueError("Excel actual group/session 來源契約不符")
    groups = {group["group_id"]: group for group in ar.build_actual_review_groups(source_ledger)}
    source_by_occ = {entry["occurrence_id"]: entry for entry in source_ledger}
    attached_groups = (proof_payload.get("attestations") or {}).get("actual_groups") or {}
    target_ledger = sp.materialize_ledger(target_manifest, target_db)
    by_target_occ = {entry["occurrence_id"]: entry for entry in target_ledger}
    workbook_sha = _snapshot_sha
    decisions = []
    seen_groups = set()
    for number, row in enumerate(rows, 2):
        group_id = str(row.get("group_id") or "").strip()
        decision = str(row.get("decision") or "").strip().upper()
        if not group_id and not decision:
            continue
        if group_id in seen_groups:
            raise ValueError(f"Excel actual row {number} group_id 重複")
        seen_groups.add(group_id)
        group = groups.get(group_id)
        attached = attached_groups.get(str(number))
        if attached is not None:
            sample = source_by_occ.get(str(row.get("sample_a_occurrence_id") or ""))
            if sample is None or not isinstance(attached, dict):
                raise ValueError(f"Excel actual row {number} 補證樣本無效")
            group = ar.build_actual_group_for_entry(source_ledger, sample)
            members_for_check = group["members"]
            supplied_overrides = attached.get("override_rows")
            if (attached.get("group_id") != group_id
                    or attached.get("old_group_snapshot") != str(row.get("group_snapshot") or "")
                    or attached.get("current_group_snapshot") != group["group_snapshot"]
                    or attached.get("member_occurrence_ids") !=
                       [entry["occurrence_id"] for entry in members_for_check]
                    or attached.get("reading") != ar._canon(row.get("actual_reading"))
                    or not isinstance(supplied_overrides, list)
                    or len(supplied_overrides) != len(members_for_check)):
                raise ValueError(f"Excel actual row {number} 補證群組/判定與來源不符")
            for member, override in zip(members_for_check, supplied_overrides):
                fields = ("pdf_contains", "pdf_excludes", "page", "target_char", "stable_key", "x0", "y0")
                if (not isinstance(override, dict)
                        or tuple(override.get(field) for field in fields) != ar._override_key_from_entry(member)
                        or ar._canon(override.get("actual_reading")) != attached["reading"]
                        or member.get("actual") != attached["reading"]):
                    raise ValueError(f"Excel actual row {number} 補證 occurrence-local 判定無效")
        elif group is None or str(row.get("group_snapshot") or "") != group["group_snapshot"]:
            raise ValueError(f"Excel actual row {number} 來源 group/snapshot 失效")
        members = group["members"]
        if (str(row.get("group_kind") or "") != group["kind"]
                or str(row.get("exact_key") or "") != group["exact_key"]
                or int(row.get("occurrence_count") or 0) != len(members)
                or str(row.get("sample_a_pdf") or "") != str(members[0].get("pdf_name") or "")
                or str(row.get("sample_a_page") or "") != str(members[0].get("printed_page") or "")):
            raise ValueError(f"Excel actual row {number} 來源 group 內容/位置欄位已變更")
        if (str(row.get("sample_a_occurrence_id") or "") != members[0]["occurrence_id"]
                or str(row.get("sample_b_occurrence_id") or "")
                != (members[1]["occurrence_id"] if len(members) > 1 else "")):
            raise ValueError(f"Excel actual row {number} 樣本 occurrence 與來源 group 不符")
        if decision not in {"", "VERIFIED", "UNRESOLVED"}:
            raise ValueError(f"Excel actual row {number} decision 無效")
        if decision != "VERIFIED":
            continue
        reading = ar._canon(row.get("actual_reading"))
        if not reading or str(row.get("actual_reading") or "").strip() != reading:
            raise ValueError(f"Excel actual row {number} 注音不是 canonical")
        a_checked = str(row.get("sample_a_checked") or "").upper() == "Y"
        b_checked = str(row.get("sample_b_checked") or "").upper() == "Y"
        if (not a_checked or (len(members) > 1 and not b_checked)
                or str(row.get("confidence") or "") not in {"高", "中"}):
            raise ValueError(f"Excel actual row {number} 視覺核對樣本/信心不足")
        # The source's direct visual decision would cover its exact peers only
        # after both source samples were checked. Every target position is still
        # independently mapped from the source's complete occurrence roster.
        selected = members if len(members) <= 1 or b_checked else members[:1]
        target_entries = []
        for member in selected:
            pair = reviews.get(member["review_id"])
            if pair is None:
                raise ValueError(f"Excel actual row {number} 來源 occurrence 無位置映射")
            target_entries.append(by_target_occ[pair[1]["occurrence_id"]])
        provenance = {
            "source_session_id": source_manifest["session_id"],
            "source_manifest_integrity_sha256": source_manifest["manifest_integrity_sha256"],
            "source_pdf_sha256": [item["pdf_sha256"] for item in source_manifest["pdfs"]],
            "source_group_id": group_id, "source_group_snapshot": group["group_snapshot"],
            "source_occurrence_ids": [item["occurrence_id"] for item in selected],
            "source_review_ids": [item["review_id"] for item in selected],
            "source_roster": [
                {"occurrence_id": item["occurrence_id"], "review_id": item["review_id"],
                 "pdf_sha256": item["pdf_sha256"]} for item in selected
            ],
            "target_session_id": target_manifest["session_id"],
            "target_occurrence_ids": [item["occurrence_id"] for item in target_entries],
            "source_excel_sha256": workbook_sha, "source_excel_row": number,
            "source_decision": copy.deepcopy(row),
        }
        decisions.append((group_id, reading, target_entries, provenance))

    root = sp.project_actual_evidence_root(target_dir)
    outbox = load_promotion_outbox(root)
    if any(item["status"] == "PENDING" for item in outbox["items"]):
        raise ValueError("目標尚有待配送的原有 Global outbox；先完成來源交易恢復")
    override_path = root / ar.OCCURRENCE_OVERRIDE_FILE
    original_rows = ar._read_csv(override_path, ar.OVERRIDE_HEADERS)
    fields = ("pdf_contains", "pdf_excludes", "page", "target_char", "stable_key", "x0", "y0")
    by_key = {tuple(row[field] for field in fields): row for row in original_rows}
    conflict_receipt_path = target_dir / ACTUAL_EXCEL_CONFLICT_FILE
    conflict_receipt_sha = _sha(conflict_receipt_path) if conflict_receipt_path.exists() else None
    prior_conflicts = (_validated_actual_excel_receipt(conflict_receipt_path)["conflicts"]
                       if conflict_receipt_sha is not None else [])
    changed = False
    conflict_records = []
    recovery_results = []
    for group_id, reading, entries, provenance in decisions:
        target_ids = []
        for entry in entries:
            key = ar._override_key_from_entry(entry)
            old = by_key.get(key)
            if old is not None and ar._canon(old["actual_reading"]) not in {"", reading}:
                conflict_records.append({"target_occurrence_id": entry["occurrence_id"],
                                         "target_review_id": entry["review_id"],
                                         "target_pdf_sha256": entry["pdf_sha256"],
                                         "target_record": copy.deepcopy(old),
                                         "source_record": copy.deepcopy(provenance),
                                         "source_reading": reading})
                continue
            audit = json.dumps(provenance, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
            if old is not None and audit in old.get("note", ""):
                continue
            # A fresh local visual decision replaces the occurrence row, but
            # the sealed conflict receipt still retains the previously imported
            # workbook's original row. Replaying that workbook is not a new
            # decision and must not replace the local adjudication or its token.
            if old is not None and any(
                item["target_occurrence_id"] == entry["occurrence_id"]
                and ar._canon(item["target_record"]["actual_reading"]) == reading
                and audit in item["target_record"].get("note", "")
                for item in prior_conflicts
            ):
                continue
            new_row = dict(zip(fields, key)) | {
                "actual_reading": reading,
                "source": (old["source"] if old is not None
                           and ar._canon(old["actual_reading"]) == reading else
                           "portable Excel occurrence-local actual (no Global quorum)"),
                "note": ((old.get("note", "") + "；") if old else "") + audit,
            }
            by_key[key] = new_row
            target_ids.append(entry["occurrence_id"])
            changed = True
        if target_ids:
            recovery_results.append({"group_id": group_id, "reading": reading,
                                     "verified_occurrence_ids": sorted(target_ids),
                                     "affected_occurrence_ids": sorted(target_ids)})
    def validate_live_target():
        _verify_excel_import_snapshot(xlsx, _original_xlsx, _snapshot_sha)
        if (target_dir / INCOMPLETE_FILE).exists() or (root / PROJECT_TRANSACTION_FILE).exists():
            raise ValueError("目標 actual Excel 交易待專用恢復；未寫入或回報完成")
        live_manifest, live_db = _load_project(target_dir, allow_actual_excel_conflict=True)
        if live_manifest != target_manifest or live_db != target_db:
            raise ValueError("目標封印或人工判定於 actual Excel 匯入期間變動；未寫入")
        _verify_mapped_target_pdfs(target_dir, live_manifest, pdf_map, geometry)
        if (_sha(conflict_receipt_path) if conflict_receipt_path.exists() else None) != conflict_receipt_sha:
            raise ValueError("Excel actual 衝突紀錄於匯入期間變動；未寫入")

    _verify_excel_import_snapshot(xlsx, _original_xlsx, _snapshot_sha)
    if conflict_records:
        receipt = target_dir / ACTUAL_EXCEL_CONFLICT_FILE
        with sp.project_delivery_lock(root):
            validate_live_target()
            if ar._read_csv(override_path, ar.OVERRIDE_HEADERS) != original_rows:
                raise ValueError("目標 actual 證據於衝突登記前已變動；未覆寫任何紀錄")
            try:
                previous, previous_sha = _validated_actual_excel_receipt(receipt, with_digest=True)
            except FileNotFoundError:
                previous, previous_sha = {"conflicts": []}, ""
            all_conflicts = copy.deepcopy(previous["conflicts"])
            for conflict in conflict_records:
                source = conflict["source_record"]
                key = (conflict["target_occurrence_id"], source["source_excel_sha256"],
                       source["source_excel_row"], conflict["source_reading"])
                matching = [item for item in all_conflicts if isinstance(item, dict)
                            and isinstance(item.get("source_record"), dict)
                            and (item.get("target_occurrence_id"),
                                 item["source_record"].get("source_excel_sha256"),
                                 item["source_record"].get("source_excel_row"),
                                 item.get("source_reading")) == key]
                if matching:
                    if len(matching) != 1 or matching[0] != conflict:
                        raise ValueError("既有 Excel actual 衝突來源識別有矛盾；未覆寫")
                else:
                    all_conflicts.append(conflict)
            if all_conflicts != previous["conflicts"]:
                status_path = target_dir / "pipeline_status.json"
                original_status = status_path.read_bytes() if status_path.exists() else None
                next_status = None
                if original_status is not None:
                    live_status = sp.json_parse_strict(original_status, status_path)
                    if not isinstance(live_status, dict):
                        raise ValueError("既有 pipeline status 無法驗證；未登記 actual 衝突")
                    next_status = _actual_excel_conflict_status(live_status, all_conflicts)
                # The receipt is the durable conflict gate. Publish it before
                # status so a rejected receipt CAS cannot require status rollback.
                sp.json_save(receipt, _sealed_actual_excel_conflicts(all_conflicts),
                             expected_sha256=previous_sha)
                if next_status is not None:
                    sp.json_save(status_path, next_status,
                                 expected_sha256=hashlib.sha256(original_status).hexdigest())
        raise ValueError("Excel actual 同位置已有不同判定；既有與本次來源均保留於跨Excel_actual衝突.json，目標未覆寫")
    if existing_actual_conflicts:
        raise ValueError("Excel actual 衝突尚未在原頁重新核對；未匯入其他判定或回報完成")
    marker = target_dir / INCOMPLETE_FILE
    if not changed:
        with sp.project_delivery_lock(root):
            validate_live_target()
            if (_sha(conflict_receipt_path) if conflict_receipt_path.exists() else None) != conflict_receipt_sha:
                raise ValueError("Excel actual 衝突紀錄於匯入期間變動；請重新核對")
            if actual_excel_conflict_state(target_dir, target_manifest, target_db):
                raise ValueError("Excel actual 衝突於匯入期間仍待本地核對；不能回報完成")
            if marker.exists() or (root / PROJECT_TRANSACTION_FILE).exists():
                raise ValueError("目標 actual Excel 交易待專用恢復；不能回報匯入完成")
            if ar._read_csv(override_path, ar.OVERRIDE_HEADERS) != original_rows:
                raise ValueError("目標 actual occurrence 證據於匯入期間變動；請重新匯入核對")
        return sp.ActualImportResult(0, 0, target_dir / "注音校對_最終報告.xlsx",
                                     {"project_actual_commit": "NO_CHANGES", "global_promotion_delivery": "NOT_ATTEMPTED"})
    # Preflight may take time, so bind the live target, marker and durable
    # transaction under one lock. A second import must never own our marker.
    with sp.project_delivery_lock(root):
        validate_live_target()
        if (_sha(conflict_receipt_path) if conflict_receipt_path.exists() else None) != conflict_receipt_sha:
            raise ValueError("Excel actual 衝突紀錄於匯入期間變動；未寫入")
        if actual_excel_conflict_state(target_dir, target_manifest, target_db):
            raise ValueError("Excel actual 衝突於匯入期間仍待本地核對；未寫入")
        if marker.exists() or (root / PROJECT_TRANSACTION_FILE).exists():
            raise ValueError("目標 actual Excel 交易待專用恢復；未開始新匯入")
        live_manifest = sp.json_load_strict(target_dir / "校對工作階段.json")
        sp.validate_manifest_integrity(live_manifest)
        sp.validate_output_artifact_hashes(live_manifest)
        if live_manifest["manifest_integrity_sha256"] != target_manifest["manifest_integrity_sha256"]:
            raise ValueError("目標封印於 actual Excel 匯入期間變動；未寫入")
        live_db = sp.normalize_db(sp.json_load_strict(target_dir / "人工判定資料庫.json"))
        if live_db != target_db or _actual_transfer_state(target_dir, live_manifest)["status"] != "NONE":
            raise ValueError("目標人工判定或 actual 暫存於匯入期間變動；未寫入")
        if ar._read_csv(override_path, ar.OVERRIDE_HEADERS) != original_rows:
            raise ValueError("目標 actual occurrence 證據於匯入期間變動；未覆寫")
        if any(item["status"] == "PENDING" for item in load_promotion_outbox(root)["items"]):
            raise ValueError("目標 Global outbox 於匯入期間變動；未寫入 actual")
        recovery_plan = post_commit_recovery_plan_from_results(recovery_results)
        owned_marker = _sealed_actual_excel_marker({
            "status": "ACTUAL_EXCEL_REFRESH_PENDING",
            "source_excel_sha256": workbook_sha,
            "import_token": uuid.uuid4().hex,
            "pre_state": _actual_excel_file_snapshot(target_dir),
            "sealed_workbooks": _sealed_workbook_snapshot(target_dir, live_manifest),
            "recovery_plan": recovery_plan,
            "recovery_plan_sha256": hashlib.sha256(_canonical(recovery_plan)).hexdigest(),
        })
        marker_sha = sp.json_save(marker, owned_marker, expected_sha256="")
        try:
            with direct_visual_project_transaction(root) as bind_recovery_plan:
                _assert_actual_marker_owner(marker, marker_sha)
                _verify_mapped_target_pdfs(target_dir, live_manifest, pdf_map, geometry)
                ar._write_csv(override_path, ar.OVERRIDE_HEADERS,
                              sorted(by_key.values(), key=lambda row: tuple(row[field] for field in fields)))
                _assert_actual_marker_owner(marker, marker_sha)
                _verify_mapped_target_pdfs(target_dir, live_manifest, pdf_map, geometry)
                bind_recovery_plan(recovery_plan)
        except Exception:
            # A PREPARED/COMMITTED journal owns recovery. Only remove our own
            # marker after a fully rolled-back precommit failure.
            if not (root / PROJECT_TRANSACTION_FILE).exists() and marker.exists():
                _assert_actual_marker_owner(marker, marker_sha)
                marker.unlink()
            raise
        recovered = resume_actual_excel_project(target_dir)
        if recovered.get("project_actual_commit") != "COMMITTED":
            raise ValueError("Excel actual 交易未取得 COMMITTED 恢復紀錄")
        return sp.ActualImportResult(len(decisions), recovered["cleared_actual_dependent_event_count"],
                                     Path(recovered["refresh_report"]),
                                     {"project_actual_commit": "COMMITTED",
                                      "global_promotion_delivery": recovered["global_promotion_delivery"],
                                      "project_refresh": recovered["project_refresh"]})


def _legacy_actual_already_imported(target_dir, manifest, db, xlsx, rows):
    """Recognize the exact old workbook already recorded on every local member."""
    import actual_review as ar
    import standalone_proofread as sp

    sha = _sha(xlsx)
    ledger = sp.materialize_ledger(manifest, db)
    by_occ = {entry["occurrence_id"]: entry for entry in ledger}
    override_path = sp.project_actual_evidence_root(target_dir) / ar.OCCURRENCE_OVERRIDE_FILE
    fields = ("pdf_contains", "pdf_excludes", "page", "target_char", "stable_key", "x0", "y0")
    overrides = {tuple(item[field] for field in fields): item
                 for item in ar._read_csv(override_path, ar.OVERRIDE_HEADERS)}
    checked = 0
    for row in rows:
        if str(row.get("decision") or "").strip().upper() != "VERIFIED":
            continue
        sample = by_occ.get(str(row.get("sample_a_occurrence_id") or ""))
        if sample is None:
            return False
        group = ar.build_actual_group_for_entry(ledger, sample)
        members = group["members"]
        reading = ar._canon(row.get("actual_reading"))
        if (not reading or group["group_id"] != str(row.get("group_id") or "")
                or len(members) != int(row.get("occurrence_count") or 0)
                or members[0]["occurrence_id"] != str(row.get("sample_a_occurrence_id") or "")
                or (members[1]["occurrence_id"] if len(members) > 1 else "") !=
                   str(row.get("sample_b_occurrence_id") or "")):
            return False
        for member in members:
            override = overrides.get(ar._override_key_from_entry(member))
            if (override is None or ar._canon(override["actual_reading"]) != reading
                    or member.get("actual") != reading):
                return False
            note = override.get("note", "")
            decoder = json.JSONDecoder()
            audits = []
            for index, char in enumerate(note):
                if char == "{":
                    try:
                        item, _ = decoder.raw_decode(note[index:])
                    except (ValueError, TypeError):
                        continue
                    audits.append(item)
            if not any(isinstance(audit, dict) and audit.get("source_excel_sha256") == sha
                       and audit.get("source_group_id") == group["group_id"] for audit in audits):
                return False
        checked += 1
    return checked > 0


def _actual_excel_file_snapshot(output_dir: Path):
    import actual_review as ar
    import standalone_proofread as sp

    output_dir = Path(output_dir)
    root = sp.project_actual_evidence_root(output_dir)
    paths = {
        "manifest": output_dir / "校對工作階段.json",
        "db": output_dir / "人工判定資料庫.json",
        "pending": output_dir / "待人工確認.json",
        "report": output_dir / "注音校對_最終報告.xlsx",
        "status": output_dir / "pipeline_status.json",
        "override": root / ar.OCCURRENCE_OVERRIDE_FILE,
    }
    return {name: _sha(path) if path.exists() else None for name, path in paths.items()}


def _sealed_actual_excel_marker(fields: Mapping[str, Any]) -> dict[str, Any]:
    marker = {"version": ACTUAL_EXCEL_MARKER_VERSION, **fields}
    marker["integrity_sha256"] = hashlib.sha256(_canonical(marker)).hexdigest()
    return marker


def _valid_sha256(value: Any) -> bool:
    return isinstance(value, str) and len(value) == 64 and all(ch in "0123456789abcdef" for ch in value)


def _validate_actual_excel_marker(output_dir: Path, marker: Any) -> None:
    """Reject unknown owners/plans before recovery can touch a journal or marker."""
    import standalone_proofread as sp
    from global_exact_glyph_library import GlobalLibraryValidationError
    from global_glyph_promotion import validate_post_commit_recovery_plan

    base = {"version", "status", "source_excel_sha256", "import_token", "integrity_sha256"}
    pending = {"pre_state", "sealed_workbooks", "recovery_plan", "recovery_plan_sha256"}
    recovered = {"recovery_token", "post_state"}
    if not isinstance(marker, dict) or marker.get("version") != ACTUAL_EXCEL_MARKER_VERSION:
        raise ValueError("actual Excel 恢復標記版本或恢復材料不完整；保留原標記")
    expected = (base | pending if marker.get("status") == "ACTUAL_EXCEL_REFRESH_PENDING"
                else base | recovered if marker.get("status") == "ACTUAL_EXCEL_RECOVERED" else set())
    if (set(marker) != expected or not _valid_sha256(marker.get("source_excel_sha256"))
            or not isinstance(marker.get("import_token"), str)
            or len(marker["import_token"]) != 32
            or any(ch not in "0123456789abcdef" for ch in marker["import_token"])
            or not _valid_sha256(marker.get("integrity_sha256"))
            or marker["integrity_sha256"] != hashlib.sha256(_canonical({
                key: value for key, value in marker.items() if key != "integrity_sha256"
            })).hexdigest()):
        raise ValueError("actual Excel 恢復標記所有權、完整性或恢復材料不完整；保留原標記")
    state = marker.get("pre_state") if pending <= expected else marker.get("post_state")
    required_state = {"manifest", "db", "pending", "report", "status", "override"}
    if (not isinstance(state, dict) or set(state) != required_state
            or not _valid_sha256(state.get("manifest")) or not _valid_sha256(state.get("db"))
            or any(value is not None and not _valid_sha256(value) for value in state.values())):
        raise ValueError("actual Excel 恢復標記檔案狀態不完整；保留原標記")
    if marker["status"] == "ACTUAL_EXCEL_RECOVERED":
        if not _valid_sha256(marker.get("recovery_token")):
            raise ValueError("actual Excel 已恢復標記缺少交易所有權；保留原標記")
        return
    plan = marker["recovery_plan"]
    try:
        if (validate_post_commit_recovery_plan(plan) != plan
                or not _valid_sha256(marker.get("recovery_plan_sha256"))
                or marker["recovery_plan_sha256"] != hashlib.sha256(_canonical(plan)).hexdigest()):
            raise ValueError("actual Excel 恢復計畫與標記不符")
        snapshot = marker["sealed_workbooks"]
        if not isinstance(snapshot, dict) or set(snapshot) != {
                "manifest_integrity_sha256", "manifest_bytes", "entries"}:
            raise ValueError("actual Excel 封印工作簿恢復材料不完整")
        old_raw = base64.b64decode(snapshot["manifest_bytes"], validate=True)
        old_manifest = json.loads(old_raw)
        if not isinstance(old_manifest, dict):
            raise ValueError("actual Excel 舊封印 manifest 格式不符")
        if (hashlib.sha256(old_raw).hexdigest() != state["manifest"]
                or snapshot["manifest_integrity_sha256"] != old_manifest.get("manifest_integrity_sha256")):
            raise ValueError("actual Excel 舊封印與恢復標記不符")
        sp.validate_manifest_integrity(old_manifest)
        expected_entries = []
        for info in old_manifest.get("pdfs", []):
            for kind in ("actual", "candidate"):
                path = Path(str(info.get(f"{kind}_workbook") or "")).resolve()
                relative = path.relative_to(Path(output_dir).resolve())
                expected_entries.append((str(relative), str(info.get(f"{kind}_workbook_sha256") or "")))
        entries = snapshot["entries"]
        if not isinstance(entries, list) or len(entries) != len(expected_entries):
            raise ValueError("actual Excel 舊封印備份數量不符")
        for entry, (relative, wanted) in zip(entries, expected_entries):
            if (not isinstance(entry, dict) or set(entry) != {"path", "sha256", "bytes"}
                    or entry["path"] != relative or entry["sha256"] != wanted
                    or not _valid_sha256(wanted)
                    or hashlib.sha256(base64.b64decode(entry["bytes"], validate=True)).hexdigest() != wanted):
                raise ValueError("actual Excel 舊封印備份內容不符")
    except (KeyError, TypeError, ValueError, binascii.Error, GlobalLibraryValidationError) as exc:
        raise ValueError(f"actual Excel 恢復材料不完整或不符；保留原標記：{exc}") from exc


def _sealed_workbook_snapshot(output_dir: Path, manifest: Mapping[str, Any]):
    """Keep exact pre-commit workbook bytes for a dedicated COMMITTED redo."""
    output_dir = Path(output_dir).resolve()
    manifest_bytes = (output_dir / "校對工作階段.json").read_bytes()
    entries = []
    for info in manifest.get("pdfs", []):
        for kind in ("actual", "candidate"):
            path = Path(str(info.get(f"{kind}_workbook") or "")).resolve()
            try:
                relative = path.relative_to(output_dir)
            except ValueError as exc:
                raise ValueError("actual Excel 封印工作簿不在專案目錄內") from exc
            expected_sha = str(info.get(f"{kind}_workbook_sha256") or "")
            raw = path.read_bytes()
            if hashlib.sha256(raw).hexdigest() != expected_sha:
                raise ValueError("actual Excel 舊封印工作簿備份 hash 不符")
            entries.append({"path": str(relative), "sha256": expected_sha,
                            "bytes": base64.b64encode(raw).decode("ascii")})
    return {"manifest_integrity_sha256": manifest["manifest_integrity_sha256"],
            "manifest_bytes": base64.b64encode(manifest_bytes).decode("ascii"),
            "entries": entries}


def _restore_sealed_workbooks_for_actual_resume(output_dir: Path, marker: Mapping[str, Any], pending):
    """Bind COMMITTED source/journal, then restore only old sealed derived workbooks."""
    import actual_review as ar
    import standalone_proofread as sp

    output_dir = Path(output_dir).resolve()
    manifest_path = output_dir / "校對工作階段.json"
    snapshot = marker.get("sealed_workbooks")
    if (not isinstance(snapshot, dict)
            or not isinstance(marker.get("pre_state"), dict)
            or not isinstance(snapshot.get("entries"), list)
            or marker.get("recovery_plan_sha256") != hashlib.sha256(_canonical(pending[1])).hexdigest()):
        raise ValueError("actual Excel COMMITTED 舊封印恢復材料不完整或不符；未略過完整性檢查")
    try:
        old_raw = base64.b64decode(snapshot.get("manifest_bytes"), validate=True)
        old_manifest = json.loads(old_raw)
    except (ValueError, TypeError) as exc:
        raise ValueError("actual Excel COMMITTED 舊封印無法解碼") from exc
    if (hashlib.sha256(old_raw).hexdigest() != marker["pre_state"].get("manifest")
            or not isinstance(old_manifest, dict)
            or snapshot.get("manifest_integrity_sha256") != old_manifest.get("manifest_integrity_sha256")):
        raise ValueError("actual Excel COMMITTED 舊封印與標記不符")
    sp.validate_manifest_integrity(old_manifest)
    expected = []
    for info in old_manifest.get("pdfs", []):
        for kind in ("actual", "candidate"):
            path = Path(str(info.get(f"{kind}_workbook") or "")).resolve()
            try:
                relative = path.relative_to(output_dir)
            except ValueError as exc:
                raise ValueError("actual Excel 封印工作簿不在專案目錄內") from exc
            expected.append((path, str(relative), str(info.get(f"{kind}_workbook_sha256") or "")))
    if len(snapshot["entries"]) != len(expected):
        raise ValueError("actual Excel 舊封印工作簿備份數量不符")
    validated = []
    for item, (path, relative, wanted) in zip(snapshot["entries"], expected):
        if (not isinstance(item, dict) or item.get("path") != relative
                or item.get("sha256") != wanted or not wanted):
            raise ValueError("actual Excel 舊封印工作簿備份識別不符")
        try:
            raw = base64.b64decode(item.get("bytes"), validate=True)
        except (ValueError, TypeError) as exc:
            raise ValueError("actual Excel 舊封印工作簿備份無法解碼") from exc
        if hashlib.sha256(raw).hexdigest() != wanted:
            raise ValueError("actual Excel 舊封印工作簿備份內容不符")
        validated.append((path, raw, wanted))
    records = {entry["occurrence_id"]: entry for entry in old_manifest.get("records", [])}
    root = sp.project_actual_evidence_root(output_dir)
    fields = ("pdf_contains", "pdf_excludes", "page", "target_char", "stable_key", "x0", "y0")
    overrides = {tuple(row[field] for field in fields): row for row in
                 ar._read_csv(root / ar.OCCURRENCE_OVERRIDE_FILE, ar.OVERRIDE_HEADERS)}
    conditions = pending[1]["checked_postconditions"]
    if not conditions:
        raise ValueError("actual Excel COMMITTED journal 沒有來源判定位置")
    for condition in conditions:
        entry = records.get(condition["occurrence_id"])
        override = overrides.get(ar._override_key_from_entry(entry)) if entry else None
        if override is None or ar._canon(override["actual_reading"]) != condition["reading"]:
            raise ValueError("actual Excel COMMITTED journal 與 occurrence-local 判定不符")
        decoder = json.JSONDecoder()
        audits = []
        for index, char in enumerate(override.get("note", "")):
            if char == "{":
                try:
                    audit, _ = decoder.raw_decode(override["note"][index:])
                except (ValueError, TypeError):
                    continue
                audits.append(audit)
        if not any(isinstance(audit, dict)
                   and audit.get("source_excel_sha256") == marker["source_excel_sha256"]
                   and audit.get("source_group_id") == condition["group_id"]
                   for audit in audits):
            raise ValueError("actual Excel COMMITTED 來源 SHA/group 與 journal 不符")
    manifest = sp.json_load_strict(manifest_path)
    sp.validate_manifest_integrity(manifest)
    if _sha(manifest_path) != marker["pre_state"]["manifest"]:
        if (manifest.get("session_id") != old_manifest.get("session_id")
                or [(info.get("pdf_sha256"), info.get("pdf_name")) for info in manifest.get("pdfs", [])]
                != [(info.get("pdf_sha256"), info.get("pdf_name")) for info in old_manifest.get("pdfs", [])]
                or [(row.get("occurrence_id"), row.get("review_id")) for row in manifest.get("records", [])]
                != [(row.get("occurrence_id"), row.get("review_id")) for row in old_manifest.get("records", [])]):
            raise ValueError("actual Excel COMMITTED 已重封專案的來源識別不符")
        sp.validate_output_artifact_hashes(manifest)
        return
    for path, raw, wanted in validated:
        if not path.is_file() or _sha(path) != wanted:
            _restore_exact_file(path, raw)
    sp.validate_output_artifact_hashes(manifest)


def resume_actual_excel_project(output_dir: Path):
    """Resolve only the durable actual-Excel marker, never a publication marker."""
    import standalone_proofread as sp

    output_dir = Path(output_dir).resolve()
    with sp.project_delivery_lock(sp.project_actual_evidence_root(output_dir)):
        return _resume_actual_excel_project_locked(output_dir)


def _assert_actual_marker_owner(path: Path, expected_sha: str) -> None:
    if not path.is_file() or _sha(path) != expected_sha:
        raise ValueError("actual Excel 未完成標記已由其他交易變動；保留標記與 journal 待查核")


def _acknowledge_owned_actual_refresh(root: Path, token, marker_path: Path, marker_sha: str):
    from global_glyph_promotion import acknowledge_project_refresh, PROJECT_TRANSACTION_FILE

    journal = root / PROJECT_TRANSACTION_FILE
    raw = journal.read_bytes()
    _assert_actual_marker_owner(marker_path, marker_sha)
    acknowledge_project_refresh(root, token)
    try:
        _assert_actual_marker_owner(marker_path, marker_sha)
    except ValueError:
        # A foreign marker published during acknowledgement cannot consume our
        # recovery evidence. Restore only an absent journal; never another owner.
        try:
            with journal.open("xb") as stream:
                stream.write(raw)
                stream.flush()
                os.fsync(stream.fileno())
        except FileExistsError:
            pass
        raise


def _resume_actual_excel_project_locked(output_dir: Path):
    import standalone_proofread as sp
    from global_glyph_promotion import (
        acknowledge_project_refresh, committed_project_recovery,
        recover_pending_project_actual_write,
    )

    output_dir = Path(output_dir).resolve()
    marker_path = output_dir / INCOMPLETE_FILE
    marker_sha = _sha(marker_path)
    marker = sp.json_load_strict(marker_path)
    _assert_actual_marker_owner(marker_path, marker_sha)
    _validate_actual_excel_marker(output_dir, marker)
    root = sp.project_actual_evidence_root(output_dir)
    _assert_actual_marker_owner(marker_path, marker_sha)
    recover_pending_project_actual_write(root)
    _assert_actual_marker_owner(marker_path, marker_sha)
    pending = committed_project_recovery(root)
    if marker["status"] == "ACTUAL_EXCEL_RECOVERED":
        if marker.get("post_state") != _actual_excel_file_snapshot(output_dir):
            raise ValueError("actual Excel 已恢復標記與現存封印/報告不符；未清除")
        if pending is not None:
            if pending[0] != marker.get("recovery_token"):
                raise ValueError("actual Excel COMMITTED journal 與已恢復標記不符")
            _acknowledge_owned_actual_refresh(root, pending[0], marker_path, marker_sha)
        _assert_actual_marker_owner(marker_path, marker_sha)
        marker_path.unlink()
        return {"project_actual_commit": "COMMITTED", "project_refresh": "SUCCESS",
                "cleared_actual_dependent_event_count": 0,
                "refresh_report": str(output_dir / "注音校對_最終報告.xlsx"),
                "global_promotion_delivery": {"status": "NO_PENDING"}}
    if pending is None:
        if marker.get("pre_state") != _actual_excel_file_snapshot(output_dir):
            raise ValueError("actual Excel 未提交交易的原資料狀態已變動；保留標記待人工查核")
        _assert_actual_marker_owner(marker_path, marker_sha)
        marker_path.unlink()
        return {"project_actual_commit": "ROLLED_BACK", "project_refresh": "NOT_STARTED"}
    _restore_sealed_workbooks_for_actual_resume(output_dir, marker, pending)
    recovered = sp.recover_committed_actual_project(output_dir, acknowledge=False)
    if recovered is None or recovered.get("project_refresh") != "SUCCESS":
        raise ValueError("actual Excel COMMITTED 交易未完成安全恢復")
    _assert_actual_marker_owner(marker_path, marker_sha)
    marker_sha = sp.json_save(marker_path, _sealed_actual_excel_marker({
        "status": "ACTUAL_EXCEL_RECOVERED",
        "source_excel_sha256": marker["source_excel_sha256"],
        "import_token": marker["import_token"],
        "recovery_token": pending[0],
        "post_state": _actual_excel_file_snapshot(output_dir),
    }), expected_sha256=marker_sha)
    _acknowledge_owned_actual_refresh(root, pending[0], marker_path, marker_sha)
    _assert_actual_marker_owner(marker_path, marker_sha)
    marker_path.unlink()
    return recovered


def _source_identity(source_manifest, source, source_id, target_manifest, target, event):
    return {
        "session_id": source_manifest["session_id"],
        "pdf_sha256": source["pdf_sha256"],
        "occurrence_id": source["occurrence_id"],
        "review_id": source_id,
        "target_session_id": target_manifest["session_id"],
        "target_pdf_sha256": target["pdf_sha256"],
        "target_occurrence_id": target["occurrence_id"],
        "target_review_id": target["review_id"],
        "prior": copy.deepcopy(event.get("portability_source")),
        "prior_duplicate_sources": copy.deepcopy(event.get("portability_duplicate_sources", [])),
        "source_conflict_resolution": copy.deepcopy(event.get("portability_conflict_resolution")),
        "source_manual_expected_decision": copy.deepcopy(event.get("manual_expected_decision")),
        **({"source_expected_resolution_binding": copy.deepcopy(event["expected_resolution_binding"])}
           if "expected_resolution_binding" in event else {}),
    }


def _transfer_event(event, source, target, *, source_manifest=None, target_manifest=None,
                    source_baseline=None, target_baseline=None):
    """Rebind a validated human expected target to the proven local occurrence."""
    from occurrence_ledger import manual_expected_target, valid_manual_expected_decision

    transferred = copy.deepcopy(event)
    if "expected_resolution_binding" in event:
        import standalone_proofread as sp
        if source_baseline is None or target_baseline is None:
            raise ValueError("跨專案 expected binding 缺少已驗證雙方 baseline")
        resolution = sp._validate_expected_resolution_binding(source_baseline, event, source_manifest)
        if (resolution is None or target_manifest is None
                or source_manifest.get("expected_asset_fingerprint") != target_manifest.get("expected_asset_fingerprint")
                or sp.confirmation_expected_snapshot(source_baseline) != sp.confirmation_expected_snapshot(target_baseline)):
            raise ValueError("跨專案 expected binding baseline／fingerprint 不同；未轉接")
        rebound = sp.bind_expected_resolution(target_baseline, resolution, target_manifest,
            legacy_evidence=event["expected_resolution_binding"]["legacy_evidence"])
        transferred["expected_resolution_binding"] = rebound["expected_resolution_binding"]
    previous = event.get("undo_previous_event")
    if isinstance(previous, dict) and "expected_resolution_binding" in previous:
        transferred["undo_previous_event"] = _transfer_event(previous, source, target,
            source_manifest=source_manifest, target_manifest=target_manifest,
            source_baseline=source_baseline, target_baseline=target_baseline)
    if "manual_expected_decision" in event:
        if not valid_manual_expected_decision({**source, **event}):
            raise ValueError("來源人工 expected 判定位置或語境無法驗證")
        transferred["manual_expected_decision"]["target"] = manual_expected_target(target)
        if not valid_manual_expected_decision({**target, **transferred}):
            raise ValueError("目標人工 expected 判定位置或語境無法驗證")
    return transferred


def _validate_project_confirmation_transfer(event, source, target):
    """Bind a transferred six-gate verdict to the target's independent evidence."""
    import standalone_proofread as sp

    if event.get("action") != "確認現版差異":
        return
    source_actual = sp.actual_confirmation_snapshot(source)
    recorded = event.get("confirmation_actual_snapshot")
    if (not isinstance(recorded, dict) or set(recorded) != set(source_actual)
            or recorded != source_actual):
        raise ValueError("來源六閘門缺少有效 actual 確認 snapshot；不能跨專案沿用")
    if source_actual != sp.actual_confirmation_snapshot(target):
        raise ValueError("目標 actual 讀音或證據不同；來源六閘門確認不可直接沿用")
    event_expected = list(sp.normalize_expected_set(event.get("expected_set")))
    source_expected = list(sp.normalize_expected_set(source.get("expected_set")))
    target_expected = list(sp.normalize_expected_set(target.get("expected_set")))
    if (not event_expected or event_expected != source_expected or source_expected != target_expected
            or any(not str(event.get(key) or "")
                   or str(event.get(key) or "") != str(source.get(key) or "")
                   or str(source.get(key) or "") != str(target.get(key) or "")
                   for key in ("expected_evidence", "context_evidence"))
            or any(event.get(gate) is not True for gate in sp.CONFIRMATION_GATES)):
        raise ValueError("目標 expected、詞境或六閘門證據不同；來源確認不可直接沿用")


def _mapped_occurrence_overrides(source_dir: Path, source_manifest, target_manifest, reviews):
    """Carry occurrence-local actual only; never mint Global glyph evidence."""
    import actual_review as ar
    import standalone_proofread as sp

    source_root = sp.project_actual_evidence_root(source_dir)
    source_names = [Path(str(info["pdf_name"])) for info in source_manifest.get("pdfs", [])]
    # Hold the source delivery lock across verification and consumption. The
    # per-PDF hash below also binds the exact in-memory rows being transferred
    # to the sealed workbook, even if an external writer bypasses that lock.
    with sp.project_delivery_lock(source_root):
        sealed_by_pdf = []
        for info, name in zip(source_manifest.get("pdfs", []), source_names):
            workbook = _artifact(source_dir, info, "actual")
            with _workbook_snapshot(workbook, expected_sha256=info["actual_workbook_sha256"]) as (workbook, _):
                metadata = sp.workbook_metadata(workbook)
                if metadata.get("pdf_sha256") != info.get("pdf_sha256"):
                    raise ValueError("來源 actual workbook PDF SHA 不符")
                try:
                    components = json.loads(str(metadata["actual_asset_fingerprint_components"]))
                    sealed_dynamic = components["dynamic_actual_evidence_hashes"]
                except (KeyError, TypeError, ValueError) as exc:
                    raise ValueError("來源 actual workbook 缺少可核對的動態 evidence fingerprint") from exc
                dependency_roster = ar.actual_workbook_dynamic_dependencies(workbook)
            current_dynamic = ar.dynamic_actual_hashes(
                source_root, pdf_path=name, dependencies=dependency_roster, read_only=True,
            )
            if sealed_dynamic != current_dynamic:
                raise ValueError(f"來源動態 actual 證據與封印 workbook fingerprint 不符：{info['pdf_name']}")
            sealed_by_pdf.append((name, sealed_dynamic[ar.OCCURRENCE_OVERRIDE_FILE]))
        rows = ar._read_csv(source_root / ar.OCCURRENCE_OVERRIDE_FILE, ar.OVERRIDE_HEADERS)
        for name, sealed_hash in sealed_by_pdf:
            consumed = [row for row in rows if ar._override_applies_to_pdf(row, name)]
            if ar._subset_sha256(consumed, ar.OVERRIDE_HEADERS) != sealed_hash:
                raise ValueError(f"來源動態 actual 證據與封印 workbook fingerprint 不符：{name}")
    rows = [row for row in rows if any(ar._override_applies_to_pdf(row, name) for name in source_names)]
    if not rows:
        return []
    by_source_key = {}
    for source, target in reviews.values():
        key = ar._override_key_from_entry(source)[2:]
        if key in by_source_key:
            raise ValueError("來源 actual occurrence override 位置不唯一")
        by_source_key[key] = (source, target)
    mapped = []
    seen_target = set()
    for row in rows:
        key = tuple(row.get(field, "") for field in ("page", "target_char", "stable_key", "x0", "y0"))
        pair = by_source_key.get(key)
        if pair is None:
            raise ValueError(f"來源 actual occurrence override 無法唯一對應：{key}")
        source, target = pair
        reading = ar._canon(row.get("actual_reading"))
        if not reading or reading != row["actual_reading"]:
            raise ValueError("來源 actual occurrence override 讀音不是 canonical")
        target_key = ar._override_key_from_entry(target)
        if target_key in seen_target:
            raise ValueError("多個來源 actual override 對應同一目標位置")
        seen_target.add(target_key)
        mapped.append(dict(zip(
            ("pdf_contains", "pdf_excludes", "page", "target_char", "stable_key", "x0", "y0"),
            target_key,
        )) | {
            "actual_reading": reading,
            "source": "portable occurrence-local actual (no Global quorum)",
            "note": (
                f"source_session={source_manifest['session_id']};source_pdf_sha256={source['pdf_sha256']};"
                f"source_occurrence_id={source['occurrence_id']};source={row['source']};note={row['note']}"
            ),
        })
    return mapped


def _publish_portable_outputs(target_dir: Path, manifest, db):
    """Rebuild every derived user view from the committed target DB."""
    import standalone_proofread as sp

    sp.save_pending_json(target_dir, manifest, db)
    return sp.generate_report(target_dir, manifest, db)


def import_project_decisions(source_dir: Path, target_dir: Path, *,
                             _allow_incomplete_target: bool = False,
                             _selected_local_pdf: Path | None = None,
                             _target_marker_sha: str | None = None) -> dict[str, Any]:
    """Import only proven mapped review events; conflicts remain untouched."""
    import standalone_proofread as sp

    source_dir, target_dir = Path(source_dir).resolve(), Path(target_dir).resolve()
    if source_dir == target_dir:
        raise ValueError("來源與目標專案必須不同")
    if _allow_incomplete_target:
        if not _target_marker_sha:
            raise ValueError("內部接續缺少明確 owner；未匯入")
        _assert_actual_marker_owner(target_dir / INCOMPLETE_FILE, _target_marker_sha)
    source_manifest, source_db = _load_project(source_dir)
    target_manifest, _ = _load_project(target_dir, allow_incomplete=_allow_incomplete_target)
    source_current = {entry["review_id"]: entry for entry in
                      sp.materialize_ledger(source_manifest, source_db)}
    source_actual = _actual_transfer_state(source_dir, source_manifest)
    _actual_transfer_state(target_dir, target_manifest)
    if _selected_local_pdf is None:
        proof = _load_proof(source_dir, source_manifest)
    else:
        local_pdf = Path(_selected_local_pdf).resolve()
        target_pdfs = target_manifest.get("pdfs", [])
        if (len(target_pdfs) != 1 or not local_pdf.is_file()
                or _sha(local_pdf) != target_pdfs[0].get("pdf_sha256")):
            raise ValueError("明確指定的本地 PDF 與目標封印 SHA 不符")
        _, local_pages = _pdf_content_snapshot(local_pdf, expected_sha256=target_pdfs[0]["pdf_sha256"])
        proof = _proof_from_explicit_local(source_dir, source_manifest,
                                            local_pdf, local_pages)
    pdf_map, geometry = _match_pdfs(proof, source_manifest, target_manifest, target_dir)
    reviews = _map_reviews(source_manifest, target_manifest, pdf_map, geometry)
    mapped_overrides = _mapped_occurrence_overrides(source_dir, source_manifest, target_manifest, reviews)
    if mapped_overrides:
        import actual_review as ar
        target_actual_root = sp.project_actual_evidence_root(target_dir)
        target_overrides = ar._read_csv(target_actual_root / ar.OCCURRENCE_OVERRIDE_FILE, ar.OVERRIDE_HEADERS)
        fields = ("pdf_contains", "pdf_excludes", "page", "target_char", "stable_key", "x0", "y0")
        existing = {tuple(row[field] for field in fields): row for row in target_overrides}
        absent = [row for row in mapped_overrides if tuple(row[field] for field in fields) not in existing]
        different = [row for row in mapped_overrides if (
            (old := existing.get(tuple(row[field] for field in fields))) is not None
            and old["actual_reading"] != row["actual_reading"]
        )]
        if absent or different:
            raise ValueError(
                f"來源有已保存 occurrence-local actual 判定，目標缺少 {len(absent)} 筆、"
                f"衝突 {len(different)} 筆；需先在隔離的新接續專案安全重建 actual，未匯入任何判定"
            )
    # The same project lock protects the read, replay check, candidate and
    # compare-and-swap write. A concurrent ReviewSaveService save is therefore
    # included in the candidate rather than lost by a stale earlier read.
    with _project_transfer_locks(source_dir, target_dir):
        live_manifest = sp.json_load_strict(target_dir / "校對工作階段.json")
        sp.validate_manifest_integrity(live_manifest)
        if live_manifest["manifest_integrity_sha256"] != target_manifest["manifest_integrity_sha256"]:
            raise ValueError("目標工作階段於內容對應後已變動；未匯入")
        _verify_mapped_target_pdfs(target_dir, live_manifest, pdf_map, geometry)
        path = target_dir / "人工判定資料庫.json"
        raw_db, before = sp.json_load_snapshot(path)
        if not isinstance(raw_db, dict):
            raise ValueError("DATA_INTEGRITY_ERROR：目標人工判定資料庫根節點必須是物件")
        target_db = sp.normalize_db(raw_db)
        receipt_snapshots = _read_presentation_receipts(target_dir)
        receipt_shas = _receipt_snapshot_shas(receipt_snapshots)
        if (target_dir / INCOMPLETE_FILE).exists() and not _allow_incomplete_target:
            raise ValueError("目標跨電腦接續未完成；匯入期間不得寫入判定")
        pending_actual = actual_excel_conflict_state(target_dir, live_manifest, target_db)
        if pending_actual:
            raise ValueError(f"目標匯入期間出現未裁決 actual 衝突；未匯入判定：{pending_actual}")
        unresolved = validate_conflict_state(target_dir, live_manifest, target_db)
        if unresolved:
            raise ValueError(f"目標匯入期間出現未裁決衝突；未匯入判定：{unresolved}")
        prior_conflicts = (receipt_snapshots[CONFLICT_FILE][1]["conflicts"]
                           if receipt_snapshots[CONFLICT_FILE][0] is not None else [])
        target_current = {entry["review_id"]: entry for entry in
                          sp.materialize_ledger(live_manifest, target_db)}
        source_baseline = {entry["review_id"]: entry for entry in
                           sp.prepare_review_ledger(source_manifest, sp.normalize_db({}))[0]}
        target_baseline = {entry["review_id"]: entry for entry in
                           sp.prepare_review_ledger(live_manifest, sp.normalize_db({}))[0]}
        candidate = copy.deepcopy(target_db)
        imported, duplicates, provenance_updates, conflicts, conflict_records = 0, 0, 0, [], []
        for source_id, event in source_db["events"].items():
            if source_id not in reviews:
                raise ValueError(f"來源 event 無法對應 occurrence：{source_id}")
            source, target = reviews[source_id]
            target_id = target["review_id"]
            current = candidate["events"].get(target_id)
            identity = _source_identity(source_manifest, source, source_id, target_manifest, target, event)
            _require_transferable_expected_lane(event.get("action"), target_current[target_id],
                                                current, f"來源判定 {source_id} ")
            if event.get("action") in {"補建expected證據", "解決expected證據"}:
                source_entry, target_entry = source_current[source_id], target_current[target_id]
                _require_mapped_textbook_context(source_entry, target_entry, f"來源判定 {source_id} ")
                target_context = str(target_entry.get("context_evidence") or "").strip()
                if (current is None and target_context
                        and str(source.get("context_evidence") or "").strip() != target_context):
                    raise ValueError(f"來源判定 {source_id} 目標獨立 context_evidence 詞境不同；未覆寫")
            _validate_project_confirmation_transfer(
                event, source_current[source_id], target_current[target_id])
            transferred = _transfer_event(event, source, target,
                source_manifest=source_manifest, target_manifest=live_manifest,
                source_baseline=source_baseline[source_id], target_baseline=target_baseline[target_id])
            if _historical_import_replay(prior_conflicts, identity, event, transferred,
                                         decision_payload=_event_payload):
                duplicates += 1
                continue
            if current is not None:
                if _event_payload(current) == _event_payload(transferred):
                    duplicates += 1
                    known = [current.get("portability_source"),
                             *current.get("portability_duplicate_sources", [])]
                    if identity not in known:
                        current.setdefault("portability_duplicate_sources", []).append(identity)
                        provenance_updates += 1
                else:
                    conflicts.append({"source_review_id": source_id, "target_review_id": target_id})
                    conflict_records.append({
                        "source_review_id": source_id,
                        "target_review_id": target_id,
                        "source_session_id": source_manifest["session_id"],
                        "source_pdf_sha256": source["pdf_sha256"],
                        "source_occurrence_id": source["occurrence_id"],
                        "target_session_id": target_manifest["session_id"],
                        "target_pdf_sha256": target["pdf_sha256"],
                        "target_occurrence_id": target["occurrence_id"],
                        "source_identity": identity,
                        "target_identity": copy.deepcopy(current.get("portability_source")),
                        "source_event": copy.deepcopy(event),
                        "target_event": copy.deepcopy(current),
                    })
                    candidate["events"].pop(target_id)
                continue
            transferred["portability_source"] = identity
            candidate["events"][target_id] = transferred
            imported += 1
        # Existing replay validators are the final authority. A changed actual
        # or expected lane can invalidate a conclusion, never guessed valid.
        ledger = sp.materialize_ledger(target_manifest, candidate)
        invalid = [entry["review_id"] for entry in ledger if entry.get("review_event_replay_status")]
        if invalid:
            raise ValueError(f"目標 actual/expected 證據已改變，判定不能直接沿用：{invalid[:10]}")
        mutations = bool(imported or provenance_updates or conflict_records
                         or source_actual["status"] == "RECHECK_LOCAL_PDF_REQUIRED")
        planned_receipts = {}
        if conflict_records:
            conflict_path = target_dir / CONFLICT_FILE
            existing_conflicts = (receipt_snapshots[CONFLICT_FILE][1]
                                  if receipt_snapshots[CONFLICT_FILE][0] is not None
                                  else {"version": 1, "conflicts": []})
            if (not isinstance(existing_conflicts, dict)
                    or existing_conflicts.get("version") != 1
                    or not isinstance(existing_conflicts.get("conflicts"), list)):
                raise ValueError("目標既有衝突紀錄格式無法驗證；未覆寫")
            planned_receipts[CONFLICT_FILE] = {
                "version": 1,
                "conflicts": [*existing_conflicts["conflicts"], *conflict_records],
            }
        if source_actual["status"] == "RECHECK_LOCAL_PDF_REQUIRED":
            # Every original group remains a read-only receipt across hops.
            # No checked ID becomes a fresh local direct visual or Global vote.
            hop_mapping = {source["occurrence_id"]: target["occurrence_id"]
                           for source, target in reviews.values()}
            incoming = []
            if source_actual["staging"]["staged_groups"]:
                incoming.append({
                    "source_session_id": source_actual["source_session_id"],
                    "source_manifest_integrity_sha256": source_actual["source_manifest_integrity_sha256"],
                    "staging": source_actual["staging"],
                    "status": "RECHECK_LOCAL_PDF_REQUIRED",
                    "source_pdf_sha256": [info["pdf_sha256"] for info in source_manifest["pdfs"]],
                    "target_session_id": target_manifest["session_id"],
                    "mapped_occurrences": hop_mapping,
                })
            for prior in source_actual["prior_receipts"]:
                forwarded = copy.deepcopy(prior)
                forwarded.setdefault("forwarded_via", []).append({
                    "source_session_id": source_manifest["session_id"],
                    "target_session_id": target_manifest["session_id"],
                    "mapped_occurrences": hop_mapping,
                })
                incoming.append(forwarded)
            receipt_path = target_dir / PENDING_ACTUAL_FILE
            existing = (receipt_snapshots[PENDING_ACTUAL_FILE][1]
                        if receipt_snapshots[PENDING_ACTUAL_FILE][0] is not None
                        else {"status": "RECHECK_LOCAL_PDF_REQUIRED", "sources": []})
            if (existing.get("status") != "RECHECK_LOCAL_PDF_REQUIRED"
                    or not isinstance(existing.get("sources"), list)):
                raise ValueError("目標來源 actual 暫存紀錄格式無法驗證；未覆寫")
            sources = list(existing["sources"])
            for source_receipt in incoming:
                same_origin = [item for item in sources if (
                    item.get("source_session_id"), item.get("source_manifest_integrity_sha256")
                ) == (
                    source_receipt["source_session_id"], source_receipt["source_manifest_integrity_sha256"]
                )]
                if any(item.get("staging") != source_receipt["staging"] for item in same_origin):
                    raise ValueError("同一來源 actual 暫存原文衝突；未覆寫")
                if source_receipt not in sources:
                    sources.append(source_receipt)
            if sources != existing["sources"]:
                planned_receipts[PENDING_ACTUAL_FILE] = {
                    "status": "RECHECK_LOCAL_PDF_REQUIRED",
                    "sources": sources,
                }
        def validate_source():
            if _load_project(source_dir) != (source_manifest, source_db):
                raise ValueError("來源工作階段或判定於匯入期間變動；未轉送")
            if _actual_transfer_state(source_dir, source_manifest) != source_actual:
                raise ValueError("來源 actual 狀態於匯入期間變動；未轉送")
            if _mapped_occurrence_overrides(source_dir, source_manifest, target_manifest, reviews) != mapped_overrides:
                raise ValueError("來源 actual occurrence 證據於匯入期間變動；未轉送")
        validate_source()
        _actual_transfer_state(target_dir, live_manifest)
        _check_receipt_shas(target_dir, receipt_shas)
        marker_path = target_dir / INCOMPLETE_FILE
        marker_sha = _target_marker_sha
        if _allow_incomplete_target:
            _assert_actual_marker_owner(marker_path, marker_sha)
        elif marker_path.exists():
            raise ValueError("目標出現未完成標記；未寫入")
        if mutations and not _allow_incomplete_target:
            owned_marker = _new_presentation_marker(
                target_dir, live_manifest, before, candidate, planned_receipts,
                phase="PROJECT_IMPORT", _receipt_snapshots=receipt_snapshots)
            _actual_transfer_state(target_dir, live_manifest)
            _check_receipt_shas(target_dir, receipt_shas)
            marker_sha = sp.json_save(marker_path, owned_marker, expected_sha256="")
        def validate_owned_write():
            validate_source()
            _actual_transfer_state(target_dir, live_manifest)
            _check_receipt_shas(target_dir, receipt_shas)
            if marker_sha:
                _assert_actual_marker_owner(marker_path, marker_sha)
            if sp.json_load_strict(target_dir / "校對工作階段.json") != live_manifest:
                raise ValueError("目標工作階段於寫入前變動")
            _verify_mapped_target_pdfs(target_dir, live_manifest, pdf_map, geometry)
        for name, document in planned_receipts.items():
            validate_owned_write()
            receipt_path = target_dir / name
            receipt_shas[name] = sp.json_save(receipt_path, document,
                                             expected_sha256=receipt_shas[name])
        if imported or provenance_updates or conflict_records:
            validate_owned_write()
            sp.json_save(path, candidate, expected_sha256=before)
        if marker_sha:
            _assert_actual_marker_owner(marker_path, marker_sha)
    if not _allow_incomplete_target and mutations:
        resume_portable_project(target_dir, _expected_marker_sha=marker_sha)
    return {"imported": imported, "duplicates": duplicates, "conflicts": conflicts,
            "provenance_updates": provenance_updates,
            "matched_actual_overrides": len(mapped_overrides),
            "actual_staging": source_actual["status"],
            "source_session_id": source_manifest["session_id"],
            "target_session_id": target_manifest["session_id"], "pdf_sha_map": pdf_map}


def continue_project(source_dir: Path, local_pdf: Path, target_dir: Path):
    """Create a new sealed project from local bytes, then import proven events."""
    return merge_projects([source_dir], local_pdf, target_dir)


def _preflight_merge_sources(source_dirs, local_pdf):
    if not 1 <= len(source_dirs) <= 2 or len(set(source_dirs)) != len(source_dirs):
        raise ValueError("跨專案合併需要一至兩個不同來源專案")
    if not local_pdf.is_file() or local_pdf.suffix.lower() != ".pdf":
        raise FileNotFoundError(f"找不到當地 PDF：{local_pdf}")
    _, local_pages = _pdf_content_snapshot(local_pdf)
    for source_dir in source_dirs:
        source_manifest, _ = _load_project(source_dir)
        _actual_transfer_state(source_dir, source_manifest)
        proof = _proof_from_explicit_local(source_dir, source_manifest, local_pdf, local_pages)
        if len(proof["pdfs"]) != 1:
            raise ValueError("此入口目前僅支援單一 PDF 專案；多 PDF 請先建立目標專案再匯入")
        if not _pages_equivalent(proof["pdfs"][0]["pages"], local_pages):
            raise ValueError(f"當地 PDF 頁面內容/位置與來源不同：{source_dir}；未建立目標專案")
    return local_pages


def merge_projects(source_dirs, local_pdf: Path, target_dir: Path):
    source_dirs = [Path(path).resolve() for path in source_dirs]
    local_pdf = Path(local_pdf).resolve()
    target_dir = Path(target_dir).resolve()
    if target_dir.exists() and any(target_dir.iterdir()):
        raise ValueError("接續目標資料夾必須為空，避免覆寫既有成果")
    # Invalid source/preflight must not create even a lock-only target.
    _preflight_merge_sources(source_dirs, local_pdf)
    with _project_transfer_locks(*source_dirs, target_dir):
        return _merge_projects_locked(source_dirs, local_pdf, target_dir)


def _merge_projects_locked(source_dirs, local_pdf: Path, target_dir: Path):
    """Build a new project from local bytes and merge one or two source projects.

    The existing projects are read-only. The first source's decision wins only
    when the second source has no decision at that position; differing decisions
    are retained in an explicit conflict receipt for human resolution.
    """
    import standalone_proofread as sp

    source_dirs = [Path(path).resolve() for path in source_dirs]
    local_pdf, target_dir = Path(local_pdf).resolve(), Path(target_dir).resolve()
    if not 1 <= len(source_dirs) <= 2 or len(set(source_dirs)) != len(source_dirs):
        raise ValueError("跨專案合併需要一至兩個不同來源專案")
    # Acquiring the lock creates only its own directory chain and lock file.
    # Recheck under that lock so a waiting concurrent builder cannot overwrite
    # a project completed since the outer empty-directory preflight.
    actual_root = sp.project_actual_evidence_root(target_dir)
    allowed = {actual_root / ".global_exact_glyph_delivery.lock", actual_root}
    allowed.update(parent for parent in actual_root.parents if target_dir in parent.parents)
    if any(path not in allowed for path in target_dir.rglob("*")):
        raise ValueError("接續目標資料夾必須為空，避免覆寫既有成果")
    local_pages = _preflight_merge_sources(source_dirs, local_pdf)
    target_dir.mkdir(parents=True, exist_ok=True)
    marker = target_dir / INCOMPLETE_FILE
    marker_sha = sp.json_save(marker, {
        "status": "INCOMPLETE",
        "source_projects": [str(path) for path in source_dirs],
        "local_pdf": str(local_pdf),
        "reason": "只有建立目標封印、匯入映射與核對均成功後才移除此標記",
    }, expected_sha256="")
    try:
        _assert_actual_marker_owner(marker, marker_sha)
        sp.run_pipeline_pdfs([local_pdf], target_dir,
                             _prewrite_guard=lambda: _assert_actual_marker_owner(marker, marker_sha))
        _assert_actual_marker_owner(marker, marker_sha)
        target_manifest, _ = _load_project(target_dir, allow_incomplete=True)
        mapped_by_key = {}
        for source_dir in source_dirs:
            source_manifest, _ = _load_project(source_dir)
            proof = _proof_from_explicit_local(source_dir, source_manifest, local_pdf, local_pages)
            pdf_map, geometry = _match_pdfs(proof, source_manifest, target_manifest, target_dir)
            reviews = _map_reviews(source_manifest, target_manifest, pdf_map, geometry)
            for row in _mapped_occurrence_overrides(source_dir, source_manifest, target_manifest, reviews):
                key = tuple(row[field] for field in (
                    "pdf_contains", "pdf_excludes", "page", "target_char", "stable_key", "x0", "y0"
                ))
                old = mapped_by_key.get(key)
                if old is not None and old["actual_reading"] != row["actual_reading"]:
                    raise ValueError(f"兩來源同位置 actual 判定衝突；原專案均保留，未套用：{key}")
                if old is None:
                    mapped_by_key[key] = row
                else:
                    old["note"] += "；同讀音另一來源：" + row["note"]
        mapped_overrides = list(mapped_by_key.values())
        if mapped_overrides:
            import actual_review as ar
            actual_root = sp.project_actual_evidence_root(target_dir)
            existing_rows = ar._read_csv(actual_root / ar.OCCURRENCE_OVERRIDE_FILE, ar.OVERRIDE_HEADERS)
            fields = ("pdf_contains", "pdf_excludes", "page", "target_char", "stable_key", "x0", "y0")
            existing_keys = {tuple(row[field] for field in fields) for row in existing_rows}
            if any(tuple(row[field] for field in fields) in existing_keys for row in mapped_overrides):
                raise ValueError("目標已有同位置 actual occurrence override；未覆寫")
            _assert_actual_marker_owner(marker, marker_sha)
            ar._write_csv(actual_root / ar.OCCURRENCE_OVERRIDE_FILE, ar.OVERRIDE_HEADERS,
                          [*existing_rows, *mapped_overrides])
            sp.refresh_actual_project(
                target_dir, _prewrite_guard=lambda: _assert_actual_marker_owner(marker, marker_sha))
            refreshed_manifest, _ = _load_project(target_dir, allow_incomplete=True)
            if (refreshed_manifest.get("session_id") != target_manifest.get("session_id")
                    or {(row["review_id"], row["occurrence_id"]) for row in refreshed_manifest["records"]}
                    != {(row["review_id"], row["occurrence_id"]) for row in target_manifest["records"]}):
                raise ValueError("actual 重建後目標工作階段或 occurrence 清單改變；未匯入")
            target_manifest = refreshed_manifest
        results = []
        conflicts = []
        for source_dir in source_dirs:
            result = import_project_decisions(source_dir, target_dir,
                                              _allow_incomplete_target=True,
                                              _selected_local_pdf=local_pdf,
                                              _target_marker_sha=marker_sha)
            results.append(result)
            conflicts.extend(result["conflicts"])
        with sp.project_delivery_lock(sp.project_actual_evidence_root(target_dir)):
            _assert_actual_marker_owner(marker, marker_sha)
            final_manifest, final_db = _load_project(target_dir, allow_incomplete=True,
                                                      allow_unresolved_conflict=True)
            if final_manifest.get("session_id") != target_manifest.get("session_id"):
                raise ValueError("匯入期間目標工作階段改變；未發布")
            final_db_path = target_dir / "人工判定資料庫.json"
            raw_final_db, final_db_sha = sp.json_load_snapshot(final_db_path)
            if sp.normalize_db(raw_final_db) != final_db:
                raise ValueError("匯入期間目標判定改變；未發布")
            owned_marker = _new_presentation_marker(
                target_dir, final_manifest, final_db_sha, final_db, {},
                phase="NEW_PROJECT_FINAL", export_pending=True)
            _assert_actual_marker_owner(marker, marker_sha)
            marker_sha = sp.json_save(marker, owned_marker, expected_sha256=marker_sha)
            _assert_actual_marker_owner(marker, marker_sha)
        resume_portable_project(target_dir, _expected_marker_sha=marker_sha)
        return {"sources": results, "conflicts": conflicts,
                "imported": sum(item["imported"] for item in results),
                "duplicates": sum(item["duplicates"] for item in results),
                "matched_actual_overrides": len(mapped_overrides)}
    except Exception:
        # The source remains unchanged. Partial target files are retained with
        # an explicit marker for forensic inspection, never shown as success.
        raise
