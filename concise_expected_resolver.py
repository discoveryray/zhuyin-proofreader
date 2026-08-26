from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from unicodedata import normalize
import re

from openpyxl import load_workbook

from pronunciation_rule_engine import RuleDecision

BOPOMOFO = set("ㄅㄆㄇㄈㄉㄊㄋㄌㄍㄎㄏㄐㄑㄒㄓㄔㄕㄖㄗㄘㄙㄧㄨㄩㄚㄛㄜㄝㄞㄟㄠㄡㄢㄣㄤㄥㄦ")
TONES = set("ˊˇˋ˙")

# These characters require a context/position/sandhi rule and must never be
# promoted merely because one simplified single-character reading happens to
# be unique in the concise-dictionary export.
EXCLUDED_SINGLE = set("一不的得著了個子頭兒們麼嗎呢啊呀哇啦吧喲哦呵嘍囉和地")

SOURCE_NAME = "教育部《國語辭典簡編本》資料 dict_concised_2014_20260626.xlsx"


def _is_variation_selector(ch: str) -> bool:
    cp = ord(ch)
    return 0xFE00 <= cp <= 0xFE0F or 0xE0100 <= cp <= 0xE01EF


def _norm_bopomofo(value) -> str:
    s = str(value or "").strip()
    if not s:
        return ""
    s = s.replace("‧", "˙").replace("・", "˙").replace("一", "ㄧ")
    s = re.sub(r"[（(][^）)]*[）)]", "", s).strip()
    s = s.replace("\u0307", "˙").replace(" ", "").replace("　", "")
    if s and all(ch in BOPOMOFO or ch in TONES for ch in s):
        return s
    return ""


def _split_pron(value) -> list[str]:
    if value in (None, ""):
        return []
    # The MOE export separates syllables with ASCII/ideographic whitespace.
    out: list[str] = []
    for part in re.split(r"\s+", str(value).strip()):
        p = _norm_bopomofo(part)
        if p:
            out.append(p)
    return out


def _is_cjk(ch: str) -> bool:
    return ("\u3400" <= ch <= "\u9fff") or ("\uf900" <= ch <= "\ufaff")


def _han_chars(text: str) -> list[str]:
    return [ch for ch in str(text or "") if _is_cjk(ch)]


def _normalize_word(value) -> str:
    raw = normalize("NFKC", str(value or ""))
    return "".join(ch for ch in raw if not ch.isspace() and not _is_variation_selector(ch))


def _normalize_context(text: str, target_pos: int | None) -> tuple[str, int | None]:
    """NFKC/space/IVS cleanup while preserving the target character offset."""
    if text is None:
        return "", None
    raw = str(text)
    out: list[str] = []
    new_pos: int | None = None
    for i, ch in enumerate(raw):
        nch = normalize("NFKC", ch)
        if _is_variation_selector(ch) or nch.isspace():
            continue
        if target_pos is not None and i == target_pos:
            new_pos = len(out)
        out.extend(nch)
    return "".join(out), new_pos


@dataclass(frozen=True)
class ConciseEntry:
    word: str
    primary: tuple[str, ...]
    pron_text: str
    poly_order: str
    variant_type: str
    variant_text: str


@dataclass
class ConciseDictionary:
    path: Path
    word_entries: dict[str, list[ConciseEntry]]
    char_primary: dict[str, set[str]]
    max_word_len: int


_CACHE: dict[str, tuple[int, int, ConciseDictionary]] = {}


def load_concise_dictionary(path: Path | None) -> ConciseDictionary | None:
    if path is None:
        return None
    path = Path(path)
    if not path.exists():
        return None
    st = path.stat()
    key = str(path.resolve())
    cached = _CACHE.get(key)
    if cached and cached[0] == st.st_mtime_ns and cached[1] == st.st_size:
        return cached[2]

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    required = {"字詞名", "注音一式"}
    if not required.issubset(idx):
        wb.close()
        raise ValueError(f"簡編本資料缺少必要欄位：{sorted(required - set(idx))}")

    words: dict[str, list[ConciseEntry]] = defaultdict(list)
    char_primary: dict[str, set[str]] = defaultdict(set)
    max_word_len = 1

    def get(row, name):
        i = idx.get(name)
        return row[i] if i is not None else None

    for row in ws.iter_rows(min_row=2, values_only=True):
        word = _normalize_word(get(row, "字詞名"))
        if not word:
            continue
        primary = tuple(_split_pron(get(row, "注音一式")))
        ent = ConciseEntry(
            word=word,
            primary=primary,
            pron_text=str(get(row, "注音一式") or "").strip(),
            poly_order=str(get(row, "多音排序") or "").strip(),
            variant_type=str(get(row, "變體類型 1:變 2:又音 3:語音 4:讀音") or "").strip(),
            variant_text=str(get(row, "變體注音") or "").strip(),
        )
        words[word].append(ent)
        hchars = _han_chars(word)
        max_word_len = max(max_word_len, len(word))
        if len(hchars) == 1 and len(primary) == 1:
            char_primary[hchars[0]].add(primary[0])
    wb.close()

    obj = ConciseDictionary(path=path, word_entries=dict(words), char_primary=dict(char_primary), max_word_len=max_word_len)
    _CACHE[key] = (st.st_mtime_ns, st.st_size, obj)
    return obj


def _reading_at(word: str, ent: ConciseEntry, target_index_in_word: int, target_char: str) -> str | None:
    chars = _han_chars(word)
    if len(chars) != len(ent.primary):
        return None
    if not (0 <= target_index_in_word < len(chars)):
        return None
    if chars[target_index_in_word] != target_char:
        return None
    return ent.primary[target_index_in_word]


def _external_overlap_competitor(
    concise: "ConciseDictionary",
    text: str,
    start: int,
    end: int,
    target_pos: int,
) -> str:
    """Return a competing multi-character word that crosses the candidate boundary.

    A concise-dictionary substring is unsafe when one of its *non-target*
    characters is simultaneously part of another dictionary word extending
    outside the candidate span.  This is direct lexical-boundary evidence, not
    a hand-maintained phrase blacklist.  Examples include 小人|和他 versus the
    false substring 人和, 背|書包 versus 背書, and 轉|成功 versus 轉成.

    Ambiguous overlaps are rejected (fail closed); this function never selects
    a pronunciation.
    """
    radius = min(14, max(2, concise.max_word_len))
    for overlap_pos in range(start, end):
        if overlap_pos == target_pos or not _is_cjk(text[overlap_pos]):
            continue
        for other_start in range(max(0, overlap_pos - radius + 1), overlap_pos + 1):
            for other_end in range(overlap_pos + 1, min(len(text), overlap_pos + radius) + 1):
                if other_start >= start and other_end <= end:
                    continue
                other = text[other_start:other_end]
                if other not in concise.word_entries:
                    continue
                if len(_han_chars(other)) < 2:
                    continue
                return other
    return ""


def _known_cross_boundary_false_hit(target_char: str, word: str, text: str, target_pos: int) -> bool:
    """Block previously demonstrated dictionary-substring boundary failures.

    These guards are independent of the observed PDF Zhuyin and therefore only
    reduce false lexical promotions; they never manufacture an expected sound.
    """
    if target_char == "中" and word == "中的":
        return True

    if target_char == "比":
        tail4 = text[target_pos: target_pos + 4]
        prev2 = text[max(0, target_pos - 2): target_pos]
        after = text[target_pos + len(word): target_pos + len(word) + 1]
        comparative_preds = set("高低大小快慢多少長短遠近重輕粗細寬窄強弱")
        comparative_left = {
            "臀部", "頭部", "肩部", "背部", "胸部", "腹部",
            "腰部", "腿部", "手臂", "身高", "腳掌", "膝蓋",
        }
        if word == "比肩" and (tail4.startswith("比肩膀") or (prev2 in comparative_left and after in comparative_preds)):
            return True
        if word == "比鄰" and tail4.startswith("比鄰居"):
            return True

    if target_char == "傳" and word == "左傳":
        next_char = text[target_pos + 1: target_pos + 2]
        if next_char in {"球", "給", "遞", "送", "出", "到", "回", "向", "接"}:
            return True

    if target_char == "背" and word == "背書":
        if text[target_pos: target_pos + 3].startswith("背書包"):
            return True

    if target_char == "轉" and word == "轉成":
        if text[target_pos: target_pos + 3].startswith("轉成功"):
            return True

    # 「人和」是固定詞，但教材常見「家人和他／家人和同學」是
    # 「家人」+ 連接詞「和」的跨詞界鄰接。只在左側明確形成常見
    # 「X人」詞時阻擋；真正「天時地利人和」不會命中此條件。
    if target_char == "和" and word == "人和" and target_pos >= 2:
        left_word = text[target_pos - 2: target_pos]
        if left_word in {
            "家人", "他人", "別人", "女人", "男人", "老人", "主人", "工人",
            "客人", "病人", "大人", "友人", "國人", "本人", "私人", "行人",
            "成人", "新人", "親人", "敵人", "商人", "軍人", "藝人", "名人",
        }:
            return True

    return False


def resolve_concise_word(
    concise: ConciseDictionary | None,
    target_char: str,
    line_text: str,
    target_pos: int | None,
    blocked_phrases: set[str] | None = None,
) -> tuple[RuleDecision | None, list[RuleDecision]]:
    """Resolve a target from the longest exact MOE concise-dictionary word.

    If equally specific complete words disagree at the target position, abstain
    and return the competing decisions as a conflict.  Only primary 「注音一式」
    is used here; dictionary variants do not silently widen the expected value.
    """
    if concise is None or target_pos is None or not line_text:
        return None, []
    text, pos = _normalize_context(line_text, target_pos)
    target_char = _normalize_word(target_char)
    if pos is None or len(target_char) != 1 or not (0 <= pos < len(text)) or text[pos] != target_char:
        return None, []

    blocked_phrases = {_normalize_word(x) for x in (blocked_phrases or set()) if _normalize_word(x)}
    found: list[tuple[int, int, str, str, ConciseEntry, int]] = []
    # Keep the tested v5.1.4 handoff search radius; it covers ordinary textbook
    # lexical entries while avoiding accidental sentence-sized dictionary keys.
    radius = min(14, max(2, concise.max_word_len))
    for start in range(max(0, pos - radius), pos + 1):
        for end in range(pos + 1, min(len(text), pos + radius + 1) + 1):
            word = text[start:end]
            if word in blocked_phrases or word not in concise.word_entries:
                continue
            if _known_cross_boundary_false_hit(target_char, word, text, pos):
                continue
            if _external_overlap_competitor(concise, text, start, end, pos):
                continue
            hchars = _han_chars(word)
            if len(hchars) < 2:
                continue
            cidx = sum(1 for ch in word[: pos - start] if _is_cjk(ch))
            for ent in concise.word_entries[word]:
                reading = _reading_at(word, ent, cidx, target_char)
                if reading:
                    found.append((len(hchars), len(word), word, reading, ent, cidx))

    if not found:
        return None, []
    best = max((x[0], x[1]) for x in found)
    top = [x for x in found if (x[0], x[1]) == best]

    decisions: list[RuleDecision] = []
    seen = set()
    for _, _, word, reading, ent, _ in top:
        key = (word, reading)
        if key in seen:
            continue
        seen.add(key)
        decisions.append(RuleDecision(
            expected_reading=reading,
            rule_id=f"MOE-CONCISE-WORD-{word}-{target_char}",
            rule_mode="concise_word",
            matched_phrase=word,
            source=SOURCE_NAME,
            source_url="",
            note=f"教育部《國語辭典簡編本》完整詞條「{word}」注音：{ent.pron_text}；取本座標字音 {reading}。",
            priority=140,
            evidence_level="moe_complete_word",
        ))
    readings = {d.expected_reading for d in decisions}
    if len(readings) != 1:
        return None, decisions
    return decisions[0], decisions


def resolve_concise_single(
    concise: ConciseDictionary | None,
    target_char: str,
) -> RuleDecision | None:
    if concise is None:
        return None
    target_char = _normalize_word(target_char)
    if len(target_char) != 1 or target_char in EXCLUDED_SINGLE:
        return None
    vals = sorted(concise.char_primary.get(target_char, set()))
    if len(vals) != 1:
        return None
    reading = vals[0]
    return RuleDecision(
        expected_reading=reading,
        rule_id=f"MOE-CONCISE-SINGLE-{target_char}",
        rule_mode="concise_single",
        matched_phrase=target_char,
        source=SOURCE_NAME,
        source_url="",
        note=f"教育部《國語辭典簡編本》單字「{target_char}」的主音唯一為 {reading}。",
        priority=125,
        evidence_level="moe_single_unique",
    )
