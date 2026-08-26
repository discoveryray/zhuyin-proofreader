from __future__ import annotations

import argparse
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from cff_unseen_family_bootstrap import (
    bootstrap_unseen_families,
    cff_unseen_family_key,
    cff_variant,
    decode_with_bootstrap,
    load_crossfamily_consensus,
)
from export_zhuyin_readings import (
    VERSION as DECODER_VERSION,
    decode,
    DEFAULT_MAP, DEFAULT_GROUPS, DEFAULT_CFF_MAP, DEFAULT_CFF_CONSENSUS,
    DEFAULT_XREF_OVERRIDES, DEFAULT_TRANSFORMS, DEFAULT_FINGERPRINTS,
    DEFAULT_OUTLINE_SIGNATURES, DEFAULT_MAPPING_CORRECTIONS,
    DEFAULT_SYMBOL_TEMPLATES, DEFAULT_ACTUAL_OVERRIDES, DEFAULT_STRUCTURAL_EXCLUSIONS,
)

PROGRAM = "CFF 未見家族跨檔批次零對照自舉器"
VERSION = "5.2.0"


def _headers(ws):
    return {str(c.value or ""): i for i, c in enumerate(ws[1], start=1)}


def _summary_map(ws):
    out = {}
    for r in range(1, ws.max_row + 1):
        k = ws.cell(r, 1).value
        if k is not None:
            out[str(k)] = r
    return out


def _set_summary(ws, label, value):
    rows = _summary_map(ws)
    r = rows.get(label)
    if r is None:
        r = ws.max_row + 1
        ws.cell(r, 1).value = label
    ws.cell(r, 2).value = value


def _collect_workbook(path: Path):
    wb = load_workbook(path)
    if "實際注音" not in wb.sheetnames:
        raise ValueError(f"工作簿缺少『實際注音』頁籤：{path}")
    ws = wb["實際注音"]
    h = _headers(ws)
    required = ["注音結構偵測來源", "font", "glyph_id_字形索引", "穩定注音鍵", "CFF符號簽名", "CFF聲調", "實際注音"]
    missing = [x for x in required if x not in h]
    if missing:
        raise ValueError(f"工作簿欄位不足 {missing}：{path}")
    items = []
    for row_no in range(2, ws.max_row + 1):
        if str(ws.cell(row_no, h["注音結構偵測來源"]).value or "") != "unseen-CFF-structure":
            continue
        font = str(ws.cell(row_no, h["font"]).value or "")
        sig_text = str(ws.cell(row_no, h["CFF符號簽名"]).value or "")
        sigs = tuple(x for x in sig_text.split(";") if x)
        try:
            gid = int(ws.cell(row_no, h["glyph_id_字形索引"]).value)
        except Exception:
            continue
        item = {
            "workbook": path,
            "row": row_no,
            "font": font,
            "glyph_id": gid,
            "stable_key": str(ws.cell(row_no, h["穩定注音鍵"]).value or ""),
            "variant": cff_variant(font),
            "family_key": cff_unseen_family_key(font),
            "signatures": sigs,
            "tone": str(ws.cell(row_no, h["CFF聲調"]).value or ""),
            "old_reading": str(ws.cell(row_no, h["實際注音"]).value or ""),
            "manual_override": str(ws.cell(row_no, h.get("人工實際注音覆寫", 0)).value or "") if h.get("人工實際注音覆寫") else "",
        }
        items.append(item)
    return wb, h, items


def _build_session_maps(items, variant_ref, gid_ref):
    family_maps, audit_rows, family_stats = bootstrap_unseen_families(
        items, variant_ref, gid_ref, min_source_styles=2
    )

    # v4.6: session-wide exact-signature propagation.  This is NOT shape
    # similarity: the 40-char normalized CFF signature must be byte-identical.
    # A signature propagates only if every accepted target-family label agrees.
    label_sets = defaultdict(set)
    label_families = defaultdict(set)
    for family, mp in family_maps.items():
        for sig, label in mp.items():
            label_sets[sig].add(label)
            label_families[sig].add(family)
    global_exact = {sig: next(iter(labels)) for sig, labels in label_sets.items() if len(labels) == 1}
    global_conflicts = {sig: sorted(labels) for sig, labels in label_sets.items() if len(labels) > 1}

    occurrence_sigs = defaultdict(set)
    for it in items:
        occurrence_sigs[it["family_key"]].update(it["signatures"])

    propagated = []
    final_maps = {}
    for family, base in family_maps.items():
        mp = dict(base)
        for sig in occurrence_sigs.get(family, set()):
            if sig in mp or sig not in global_exact:
                continue
            mp[sig] = global_exact[sig]
            propagated.append({
                "family_key": family,
                "signature": sig,
                "symbol": global_exact[sig],
                "source_family_count": len(label_families[sig]),
                "source_families": "|".join(sorted(label_families[sig])),
            })
        final_maps[family] = mp
    return final_maps, audit_rows, family_stats, propagated, global_conflicts


def _sync_key_sheets(wb, ws, h, unseen_keys):
    # Recompute each unseen stable key from the final occurrence rows in this
    # workbook.  A key is promoted only if every occurrence is decoded and all
    # readings agree; otherwise it remains abstained.
    key_values = defaultdict(list)
    key_fonts = {}
    key_gids = {}
    for row_no in range(2, ws.max_row + 1):
        if str(ws.cell(row_no, h["注音結構偵測來源"]).value or "") != "unseen-CFF-structure":
            continue
        key = str(ws.cell(row_no, h["穩定注音鍵"]).value or "")
        key_values[key].append(str(ws.cell(row_no, h["實際注音"]).value or ""))
        key_fonts[key] = str(ws.cell(row_no, h["font"]).value or "")
        key_gids[key] = ws.cell(row_no, h["glyph_id_字形索引"]).value

    if "注音鍵彙整" in wb.sheetnames:
        wk = wb["注音鍵彙整"]
        hk = _headers(wk)
        if "穩定注音鍵" in hk and "實際注音" in hk:
            for r in range(2, wk.max_row + 1):
                key = str(wk.cell(r, hk["穩定注音鍵"]).value or "")
                if key not in unseen_keys:
                    continue
                vals = key_values.get(key, [])
                value = vals[0] if vals and all(vals) and len(set(vals)) == 1 else ""
                wk.cell(r, hk["實際注音"]).value = value
                if "解碼依據" in hk:
                    wk.cell(r, hk["解碼依據"]).value = "CFF未見家族跨檔批次零對照自舉（批次暫存exact簽名）" if value else "未知CFF家族：批次安全門檻未通過"
                if "驗證方式" in hk:
                    wk.cell(r, hk["驗證方式"]).value = "跨PDF批次：CID雙重共識＋目標exact簽名＋無衝突exact跨家族傳播" if value else ""
                if "備註" in hk:
                    wk.cell(r, hk["備註"]).value = "target family mapping=0；暫存簽名不寫回持久mapping；未命中或衝突即棄權。"

    if "未解碼注音鍵" in wb.sheetnames:
        wm = wb["未解碼注音鍵"]
        hm = _headers(wm)
        if "穩定注音鍵" in hm:
            to_delete = []
            present = set()
            for r in range(2, wm.max_row + 1):
                key = str(wm.cell(r, hm["穩定注音鍵"]).value or "")
                present.add(key)
                if key in unseen_keys:
                    vals = key_values.get(key, [])
                    if vals and all(vals) and len(set(vals)) == 1:
                        to_delete.append(r)
            for r in reversed(to_delete):
                wm.delete_rows(r, 1)
            # Add newly abstained keys only if the earlier single-file pass did
            # not already list them (possible if a pooled conflict retracts a
            # local acceptance).
            present = {str(wm.cell(r, hm["穩定注音鍵"]).value or "") for r in range(2, wm.max_row + 1)}
            for key in sorted(unseen_keys):
                vals = key_values.get(key, [])
                if vals and all(vals) and len(set(vals)) == 1:
                    continue
                if key in present:
                    continue
                row = [""] * wm.max_column
                def put(name, val):
                    if name in hm:
                        row[hm[name]-1] = val
                put("穩定注音鍵", key)
                put("font", key_fonts.get(key, ""))
                put("注音元件ID", key_gids.get(key, ""))
                put("出現次數", len(vals))
                put("下一步", "批次exact證據仍不足或有衝突；保持棄權，不新增target mapping。")
                put("字形架構", "CFF整字注音")
                wm.append(row)


def _style_workbook(wb):
    # Keep existing workbook style; only style newly created audit sheets.
    if "CFF批次自舉稽核" not in wb.sheetnames:
        return
    ws = wb["CFF批次自舉稽核"]
    fill = PatternFill("solid", fgColor="1F4E78")
    for c in ws[1]:
        c.fill = fill
        c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    widths = [18, 36, 10, 22, 14, 16, 18, 42, 42]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)


def _patch_workbook(path, wb, h, items, final_maps, audit_rows, propagated, global_conflicts, global_stats):
    ws = wb["實際注音"]
    old_auto = sum(1 for it in items if it["old_reading"] and not it["manual_override"])
    additions = retractions = final_auto = 0
    unseen_keys = {it["stable_key"] for it in items}
    propagated_sigs = {(x["family_key"], x["signature"]): x for x in propagated}

    for it in items:
        if it["manual_override"]:
            # Human occurrence truth is a separate evidence layer.  Batch
            # bootstrap is not allowed to replace it.
            continue
        mp = final_maps.get(it["family_key"], {})
        new_reading = decode_with_bootstrap(it, mp)
        old = str(ws.cell(it["row"], h["實際注音"]).value or "")
        if new_reading:
            final_auto += 1
        if not old and new_reading:
            additions += 1
        if old and not new_reading:
            retractions += 1
        ws.cell(it["row"], h["實際注音"]).value = new_reading
        ws.cell(it["row"], h["解碼狀態"]).value = "已解碼" if new_reading else "待建立對照"
        ws.cell(it["row"], h["解碼依據"]).value = (
            "CFF未見家族跨檔批次零對照自舉（批次暫存exact簽名）" if new_reading
            else "未知CFF家族：跨檔批次安全門檻未通過"
        )
        if new_reading:
            auto_group = f"AUTOBOOT-SESSION:{it['family_key']}"
            ws.cell(it["row"], h["字型相容群組"]).value = f"CFF:{auto_group}"
            ws.cell(it["row"], h["群組注音鍵"]).value = f"CFF:{auto_group}#{it['glyph_id']}"
            ws.cell(it["row"], h["判定方式"]).value = "跨PDF CID雙重共識→目標exact簽名暫存→無衝突exact傳播→exact解碼"
            ws.cell(it["row"], h["CFF樣式群組"]).value = auto_group
        # Show which individual signatures are actually proven.  Missing exact
        # labels remain '?' so an abstention is visibly auditable.
        if "CFF符號序列" in h:
            ws.cell(it["row"], h["CFF符號序列"]).value = "".join(mp.get(s, "?") for s in it["signatures"])

    _sync_key_sheets(wb, ws, h, unseen_keys)

    # Replace/add the batch audit sheet.
    if "CFF批次自舉稽核" in wb.sheetnames:
        del wb["CFF批次自舉稽核"]
    wa = wb.create_sheet("CFF批次自舉稽核")
    wa.append(["階段", "未見家族鍵", "CFF符號簽名", "暫存符號", "處置", "目標glyph支持", "來源家族支持", "來源家族", "安全說明"])
    for a in audit_rows:
        wa.append([
            "CID雙重共識種子", a.get("family_key", ""), a.get("signature", ""), a.get("symbol", ""),
            a.get("status", ""), a.get("target_gid_count", 0), a.get("max_source_support", 0),
            a.get("source_styles", ""),
            "variant+CID 與 CID-only 兩路必須同時存在且讀音一致；target簽名無衝突才可接受。",
        ])
    for p in propagated:
        wa.append([
            "跨家族exact簽名傳播", p["family_key"], p["signature"], p["symbol"], "接受：session exact傳播",
            "", p["source_family_count"], p["source_families"],
            "只傳播完全相同的正規化CFF簽名；所有來源標籤必須一致；不使用模糊距離。",
        ])
    for sig, labels in sorted(global_conflicts.items()):
        wa.append([
            "跨家族exact簽名傳播", "", sig, "", "棄權：跨家族同簽名標籤衝突", "", "", "|".join(labels),
            "發現互斥標籤後全域停用此簽名，不猜測。",
        ])
    _style_workbook(wb)

    # Recalculate high-level summary directly from final rows.
    all_rows = list(range(2, ws.max_row + 1))
    total = len(all_rows)
    mapped = sum(1 for r in all_rows if str(ws.cell(r, h["實際注音"]).value or ""))
    cff_rows = [r for r in all_rows if str(ws.cell(r, h.get("字形架構", 0)).value or "") == "CFF整字注音"] if h.get("字形架構") else []
    cff_mapped = sum(1 for r in cff_rows if str(ws.cell(r, h["實際注音"]).value or ""))
    unseen_total = len(items)
    unseen_manual = sum(1 for it in items if it["manual_override"])
    unseen_auto_final = sum(1 for it in items if not it["manual_override"] and decode_with_bootstrap(it, final_maps.get(it["family_key"], {})))

    if "摘要" in wb.sheetnames:
        ss = wb["摘要"]
        _set_summary(ss, "工具", f"PDF 實際注音解碼器 v{VERSION}")
        _set_summary(ss, "已解碼筆數", mapped)
        _set_summary(ss, "其中：CFF輪廓解碼總筆數", cff_mapped)
        _set_summary(ss, "其中：未見CFF家族零對照自舉", unseen_auto_final)
        _set_summary(ss, "未見CFF家族結構偵測筆數", unseen_total)
        _set_summary(ss, "未見CFF家族自舉後棄權筆數", max(0, unseen_total - unseen_auto_final - unseen_manual))
        _set_summary(ss, "待建立對照筆數", total - mapped)
        _set_summary(ss, "目前解碼覆蓋率", (mapped / total) if total else 0)
        _set_summary(ss, "批次PDF數", global_stats["pdf_count"])
        _set_summary(ss, "批次未見CFF結構筆數", global_stats["unseen_total"])
        _set_summary(ss, "批次單檔原自舉已解碼筆數", global_stats["old_auto"])
        _set_summary(ss, "批次跨檔自舉最終解碼筆數", global_stats["final_auto"])
        _set_summary(ss, "批次跨檔自舉新增筆數", global_stats["additions"])
        _set_summary(ss, "批次重新評估撤回筆數", global_stats["retractions"])
        _set_summary(ss, "批次暫存exact簽名數", global_stats["signature_count"])
        _set_summary(ss, "批次跨家族exact簽名傳播數", global_stats["propagated_count"])
        _set_summary(ss, "批次衝突簽名數", global_stats["conflict_count"])
        _set_summary(ss, "批次安全原則", "target family mapping=0；批次暫存表只存在輸出/記憶體；CID雙重共識、exact簽名、聲調安全閘門；衝突/不足即棄權；不啟用模糊幾何。")

    wb.save(path)
    return {
        "file": path.name, "unseen_total": unseen_total, "old_auto": old_auto,
        "final_auto": unseen_auto_final, "additions": additions, "retractions": retractions,
        "manual": unseen_manual,
    }


def _make_batch_report(path: Path, file_stats, family_stats, audit_rows, propagated, global_conflicts, global_stats):
    wb = Workbook()
    ws = wb.active
    ws.title = "摘要"
    ws.append(["指標", "結果"])
    metrics = [
        ("工具", f"{PROGRAM} v{VERSION}"),
        ("底層單檔解碼器", DECODER_VERSION),
        ("批次PDF數", global_stats["pdf_count"]),
        ("未見CFF結構筆數", global_stats["unseen_total"]),
        ("單檔自舉原已解碼筆數", global_stats["old_auto"]),
        ("跨檔批次最終自動解碼筆數", global_stats["final_auto"]),
        ("跨檔批次新增筆數", global_stats["additions"]),
        ("重新評估撤回筆數", global_stats["retractions"]),
        ("跨檔批次覆蓋率", (global_stats["final_auto"] / global_stats["unseen_total"]) if global_stats["unseen_total"] else 0),
        ("批次暫存exact簽名數", global_stats["signature_count"]),
        ("跨家族exact簽名傳播數", global_stats["propagated_count"]),
        ("衝突簽名數", global_stats["conflict_count"]),
        ("安全策略", "已知家族CID雙重共識→目標家族exact簽名暫存→跨PDF合併證據→無衝突exact簽名傳播；任一缺漏/衝突/聲調不安全即棄權。"),
        ("禁止事項", "不建立target family持久mapping；不以中文字/詞義/句境反填；不啟用fuzzy/nearest-neighbour。"),
        ("建立時間", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for x in metrics: ws.append(x)

    wf = wb.create_sheet("檔案統計")
    wf.append(["檔案", "未見CFF筆數", "單檔原自舉", "批次最終自舉", "新增", "撤回", "人工覆寫"])
    for r in file_stats:
        wf.append([r["file"], r["unseen_total"], r["old_auto"], r["final_auto"], r["additions"], r["retractions"], r["manual"]])

    wfa = wb.create_sheet("家族統計")
    wfa.append(["未見家族鍵", "出現筆數", "CID安全種子筆數", "候選簽名數", "直接接受簽名數", "直接衝突簽名數"])
    for family, st in sorted(family_stats.items()):
        wfa.append([family, st.get("occurrences",0), st.get("seeded_occurrences",0), st.get("candidate_signatures",0), st.get("accepted_signatures",0), st.get("conflict_signatures",0)])

    wa = wb.create_sheet("簽名稽核")
    wa.append(["階段", "未見家族鍵", "CFF符號簽名", "暫存符號", "處置", "目標glyph支持", "來源支持", "來源家族"])
    for a in audit_rows:
        wa.append(["CID雙重共識種子", a.get("family_key",""), a.get("signature",""), a.get("symbol",""), a.get("status",""), a.get("target_gid_count",0), a.get("max_source_support",0), a.get("source_styles","")])
    for p in propagated:
        wa.append(["跨家族exact簽名傳播", p["family_key"], p["signature"], p["symbol"], "接受：session exact傳播", "", p["source_family_count"], p["source_families"]])
    for sig, labels in sorted(global_conflicts.items()):
        wa.append(["跨家族exact簽名傳播", "", sig, "", "棄權：同簽名標籤衝突", "", "", "|".join(labels)])

    for sh in wb.worksheets:
        sh.sheet_view.showGridLines = False
        sh.freeze_panes = "A2"
        for c in sh[1]:
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.font = Font(color="FFFFFF", bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        for col in range(1, sh.max_column+1):
            max_len = 0
            for r in range(1, min(sh.max_row, 200)+1):
                v = sh.cell(r,col).value
                if v is not None:
                    max_len = max(max_len, min(60, len(str(v))))
            sh.column_dimensions[get_column_letter(col)].width = max(10, min(48, max_len + 2))
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 90
    # Percentage formatting for report summary.
    for r in range(2, ws.max_row+1):
        if ws.cell(r,1).value == "跨檔批次覆蓋率":
            ws.cell(r,2).number_format = "0.00%"
    wb.save(path)


def _resolve_pdfs(args):
    out = [Path(x) for x in (args.pdfs or [])]
    if args.input_dir:
        out.extend(sorted(Path(args.input_dir).glob("*.pdf")))
    # Preserve order, remove duplicates.
    seen=set(); uniq=[]
    for p in out:
        rp=p.resolve()
        if rp not in seen:
            seen.add(rp); uniq.append(p)
    return uniq


def main():
    ap = argparse.ArgumentParser(description=f"{PROGRAM} v{VERSION}")
    ap.add_argument("pdfs", nargs="*", help="要以同一批次共同自舉的 PDF")
    ap.add_argument("--input-dir", help="批次讀取此資料夾內全部 PDF")
    ap.add_argument("-o", "--output-dir", required=True, help="單檔工作簿與批次總報告輸出資料夾")
    ap.add_argument("--existing-workbooks", nargs="*", help="跳過PDF單檔解碼，直接對既有解碼工作簿做批次自舉")
    ap.add_argument("--write-only-workbooks", nargs="*", help="仍讀取全部 existing workbooks 建立批次證據，但只改寫指定工作簿；供單 PDF 增量重解碼使用")
    ap.add_argument("--cff-consensus", default=str(DEFAULT_CFF_CONSENSUS))
    args = ap.parse_args()

    outdir = Path(args.output_dir)
    outdir.mkdir(parents=True, exist_ok=True)
    workbook_paths=[]
    pdfs=[]
    if args.existing_workbooks:
        workbook_paths=[Path(x) for x in args.existing_workbooks]
    else:
        pdfs=_resolve_pdfs(args)
        if not pdfs:
            print("未指定 PDF。", file=sys.stderr); return 2
        for i,pdf in enumerate(pdfs,1):
            if not pdf.exists():
                print(f"找不到 PDF：{pdf}", file=sys.stderr); return 2
            out=outdir/f"{pdf.stem}_實際注音解碼.xlsx"
            print(f"[{i}/{len(pdfs)}] 單檔結構解碼：{pdf.name}", flush=True)
            decode(pdf, out, DEFAULT_MAP, DEFAULT_GROUPS, DEFAULT_CFF_MAP,
                   DEFAULT_XREF_OVERRIDES, DEFAULT_TRANSFORMS, DEFAULT_FINGERPRINTS,
                   DEFAULT_OUTLINE_SIGNATURES, DEFAULT_MAPPING_CORRECTIONS,
                   DEFAULT_SYMBOL_TEMPLATES, DEFAULT_ACTUAL_OVERRIDES,
                   DEFAULT_STRUCTURAL_EXCLUSIONS, Path(args.cff_consensus))
            workbook_paths.append(out)

    loaded={}; all_items=[]
    for p in workbook_paths:
        wb,h,items=_collect_workbook(p)
        loaded[p]=(wb,h,items)
        all_items.extend(items)
    variant_ref,gid_ref=load_crossfamily_consensus(Path(args.cff_consensus))
    if not variant_ref or not gid_ref:
        print("CFF 跨家族 CID 共識表為空，無法安全自舉。", file=sys.stderr); return 3

    final_maps,audit_rows,family_stats,propagated,global_conflicts=_build_session_maps(all_items,variant_ref,gid_ref)

    old_auto=sum(1 for it in all_items if it["old_reading"] and not it["manual_override"])
    final_auto=sum(1 for it in all_items if not it["manual_override"] and decode_with_bootstrap(it,final_maps.get(it["family_key"],{})))
    additions=sum(1 for it in all_items if not it["manual_override"] and not it["old_reading"] and decode_with_bootstrap(it,final_maps.get(it["family_key"],{})))
    retractions=sum(1 for it in all_items if not it["manual_override"] and it["old_reading"] and not decode_with_bootstrap(it,final_maps.get(it["family_key"],{})))
    global_stats={
        "pdf_count": len(workbook_paths), "unseen_total": len(all_items), "old_auto": old_auto,
        "final_auto": final_auto, "additions": additions, "retractions": retractions,
        "signature_count": sum(len(x) for x in final_maps.values()),
        "propagated_count": len(propagated),
        "conflict_count": sum(1 for a in audit_rows if str(a.get("status","")).startswith("棄權：同一簽名")) + len(global_conflicts),
    }

    write_only = None
    if args.write_only_workbooks:
        write_only = {str(Path(x).resolve()) for x in args.write_only_workbooks}
    file_stats=[]
    for p,(wb,h,items) in loaded.items():
        if write_only is not None and str(Path(p).resolve()) not in write_only:
            # The workbook still participates in cross-file evidence, but an
            # unrelated actual correction must not rewrite its ZIP container.
            try:
                wb.close()
            except Exception:
                pass
            continue
        st=_patch_workbook(p,wb,h,items,final_maps,audit_rows,propagated,global_conflicts,global_stats)
        file_stats.append(st)

    report=outdir/"CFF未見家族跨檔批次零對照自舉_總報告.xlsx"
    _make_batch_report(report,file_stats,family_stats,audit_rows,propagated,global_conflicts,global_stats)
    cov=(final_auto/len(all_items)) if all_items else 0.0
    print(
        f"CFF 批次處理結束（不等同全冊校對完成）。未見CFF：{final_auto}/{len(all_items)}（{cov:.2%}）；"
        f"單檔原自舉 {old_auto}；批次新增 {additions}；撤回 {retractions}；"
        f"衝突簽名 {global_stats['conflict_count']}。\n總報告：{report}",
        flush=True,
    )
    return 0


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    raise SystemExit(main())
