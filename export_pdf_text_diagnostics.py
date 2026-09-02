from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
import re
import sys
import struct
import hashlib

import fitz  # PyMuPDF
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from exact_glyph_identity import classify_ttf_glyf_record

PROGRAM = "PDF 文字層診斷匯出工具"
VERSION = "0.4.1-poin-safeexcel"
ZH_FONT_RE = re.compile(r"(zhuyin|zhuin|zuinn|poin|chuin|bpmf)", re.IGNORECASE)


def unique_output_path(pdf_path: Path) -> Path:
    base = pdf_path.with_name(f"{pdf_path.stem}_文字層診斷.xlsx")
    if not base.exists():
        return base
    i = 2
    while True:
        candidate = pdf_path.with_name(f"{pdf_path.stem}_文字層診斷_{i}.xlsx")
        if not candidate.exists():
            return candidate
        i += 1


def safe_round(value, digits=3):
    try:
        return round(float(value), digits)
    except Exception:
        return None


def safe_excel_text(value):
    if value is None:
        return ""
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))

def safe_char(codepoint):
    try:
        return safe_excel_text(chr(int(codepoint)))
    except Exception:
        return ""


def style_sheet(ws, freeze="A2"):
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 0
        for cell in ws[letter][: min(ws.max_row, 500)]:
            val = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(val))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 42)


def get_page_fonts(page: fitz.Page):
    try:
        return page.get_fonts(full=True)
    except TypeError:
        return page.get_fonts()
    except Exception:
        return []


def normalize_basefont(basefont):
    s = str(basefont or "")
    if "+" in s:
        s = s.split("+", 1)[1]
    return s


def build_font_lookup(font_records):
    lookup = defaultdict(list)
    for rec in font_records:
        vals = list(rec) + [None] * 7
        xref, ext, ftype, basefont, name, encoding, referencer = vals[:7]
        short = normalize_basefont(basefont)
        lookup[short].append({
            "xref": xref,
            "basefont": basefont,
            "resource": name,
            "encoding": encoding,
        })
    return lookup


def resolve_font(font_lookup, font_name):
    matches = font_lookup.get(str(font_name or ""), [])
    if len(matches) == 1:
        m = matches[0]
        return m["xref"], m["basefont"], m["resource"], m["encoding"]
    if len(matches) > 1:
        return ",".join(str(m["xref"]) for m in matches), "; ".join(str(m["basefont"]) for m in matches), "; ".join(str(m["resource"]) for m in matches), "; ".join(str(m["encoding"]) for m in matches)
    return None, None, None, None


class TrueTypeGlyphInspector:
    """Minimal TrueType 'glyf' composite reader; uses only the standard library."""

    ARG_1_AND_2_ARE_WORDS = 0x0001
    ARGS_ARE_XY_VALUES = 0x0002
    WE_HAVE_A_SCALE = 0x0008
    MORE_COMPONENTS = 0x0020
    WE_HAVE_AN_X_AND_Y_SCALE = 0x0040
    WE_HAVE_A_TWO_BY_TWO = 0x0080

    def __init__(self, data: bytes):
        self.data = data
        self.tables = self._read_tables()
        if "head" not in self.tables or "loca" not in self.tables or "glyf" not in self.tables or "maxp" not in self.tables:
            raise ValueError("不是可解析的 TrueType glyf 字型")
        head_off, _ = self.tables["head"]
        maxp_off, _ = self.tables["maxp"]
        self.index_to_loc_format = struct.unpack_from(">h", data, head_off + 50)[0]
        self.num_glyphs = struct.unpack_from(">H", data, maxp_off + 4)[0]
        self.glyf_off, self.glyf_length = self.tables["glyf"]
        self.loca = self._read_loca()

    def _read_tables(self):
        if len(self.data) < 12:
            raise ValueError("字型資料過短")
        num_tables = struct.unpack_from(">H", self.data, 4)[0]
        out = {}
        pos = 12
        for _ in range(num_tables):
            tag_b, _checksum, off, length = struct.unpack_from(">4sIII", self.data, pos)
            pos += 16
            tag = tag_b.decode("latin-1")
            out[tag] = (off, length)
        return out

    def _read_loca(self):
        off, _ = self.tables["loca"]
        count = self.num_glyphs + 1
        if self.index_to_loc_format == 0:
            vals = struct.unpack_from(f">{count}H", self.data, off)
            return [v * 2 for v in vals]
        vals = struct.unpack_from(f">{count}I", self.data, off)
        return list(vals)

    def glyph_bytes(self, glyph_id: int) -> bytes:
        """Return the raw ``glyf`` record for one glyph.

        This remains the existing project-local byte-level fingerprint source.
        It does not establish global eligibility: equal composite records may
        reference the same GIDs whose component outlines differ between font
        programs.  An empty/invalid glyph returns ``b""``.
        """
        try:
            gid = int(glyph_id)
        except Exception:
            return b""
        if gid < 0 or gid >= self.num_glyphs:
            return b""
        start = self.glyf_off + self.loca[gid]
        end = self.glyf_off + self.loca[gid + 1]
        if end <= start:
            return b""
        return bytes(self.data[start:end])

    def glyph_sha256(self, glyph_id: int) -> str:
        raw = self.glyph_bytes(glyph_id)
        return hashlib.sha256(raw).hexdigest() if raw else ""

    def global_exact_identity(self, glyph_id: int):
        """Return the fail-closed simple-glyf identity admission result.

        ``glyph_bytes()`` and ``glyph_sha256()`` intentionally retain their
        existing project-local raw-record behavior.  This separate boundary
        proves the current record range and complete simple-glyph structure
        before declaring it eligible for any future cross-project reuse.
        """

        def invalid(reason):
            result = classify_ttf_glyf_record(b"", record_range_valid=False)
            result["reason"] = reason
            return result
        try:
            gid = int(glyph_id)
        except Exception:
            return invalid("INVALID_GLYPH_ID")
        if gid < 0 or gid >= self.num_glyphs or gid + 1 >= len(self.loca):
            return invalid("GLYPH_ID_OUT_OF_RANGE")

        head_off, head_length = self.tables["head"]
        maxp_off, maxp_length = self.tables["maxp"]
        loca_off, loca_length = self.tables["loca"]
        loca_entry_size = 2 if self.index_to_loc_format == 0 else 4
        required_loca_length = (self.num_glyphs + 1) * loca_entry_size
        table_ranges_valid = (
            self.index_to_loc_format in {0, 1}
            and head_length >= 54
            and maxp_length >= 6
            and loca_length >= required_loca_length
            and head_off >= 0
            and maxp_off >= 0
            and loca_off >= 0
            and self.glyf_off >= 0
            and self.glyf_length >= 0
            and head_off + head_length <= len(self.data)
            and maxp_off + maxp_length <= len(self.data)
            and loca_off + loca_length <= len(self.data)
            and self.glyf_off + self.glyf_length <= len(self.data)
        )
        relative_start = self.loca[gid]
        relative_end = self.loca[gid + 1]
        record_range_valid = (
            table_ranges_valid
            and 0 <= relative_start < relative_end <= self.glyf_length
        )
        if not record_range_valid:
            return invalid("INVALID_GLYF_RECORD_RANGE")

        start = self.glyf_off + relative_start
        end = self.glyf_off + relative_end
        raw = bytes(self.data[start:end])
        if len(raw) != relative_end - relative_start:
            return invalid("TRUNCATED_GLYF_RECORD_RANGE")
        return classify_ttf_glyf_record(raw)

    def simple_contours(self, glyph_id: int):
        """Decode a simple TrueType glyph into point contours.

        Each returned point is ``(x, y, on_curve)``.  Composite glyphs return
        an empty list because the zhuyin pronunciation component itself is
        expected to be a simple outline in the supported textbook fonts.
        This parser uses only the TrueType ``glyf`` specification and does not
        depend on a renderer or on Chinese-character semantics.
        """
        try:
            gid = int(glyph_id)
        except Exception:
            return []
        if gid < 0 or gid >= self.num_glyphs:
            return []
        start = self.glyf_off + self.loca[gid]
        end = self.glyf_off + self.loca[gid + 1]
        if end <= start or end - start < 10:
            return []
        number_of_contours = struct.unpack_from(">h", self.data, start)[0]
        # Annotation components are small simple glyphs.  Fail closed on
        # pathological values so a malformed or unexpected glyph cannot turn
        # outline self-audit into an unbounded parser loop.
        if number_of_contours <= 0 or number_of_contours > 256:
            return []

        pos = start + 10
        try:
            if pos + 2 * number_of_contours > end:
                return []
            end_points = list(struct.unpack_from(f">{number_of_contours}H", self.data, pos))
            pos += 2 * number_of_contours
            point_count = end_points[-1] + 1
            if point_count <= 0 or point_count > 4096:
                return []
            if pos + 2 > end:
                return []
            instruction_length = struct.unpack_from(">H", self.data, pos)[0]
            pos += 2
            if instruction_length < 0 or pos + instruction_length > end:
                return []
            pos += instruction_length

            flags = []
            while len(flags) < point_count:
                if pos >= end:
                    return []
                flag = self.data[pos]
                pos += 1
                flags.append(flag)
                if flag & 0x08:  # REPEAT_FLAG
                    if pos >= end:
                        return []
                    repeat = self.data[pos]
                    pos += 1
                    if len(flags) + repeat > point_count:
                        return []
                    flags.extend([flag] * repeat)
            if len(flags) != point_count:
                return []

            xs = []
            x = 0
            for flag in flags:
                if flag & 0x02:  # X_SHORT_VECTOR
                    dx = self.data[pos]
                    pos += 1
                    dx = dx if flag & 0x10 else -dx
                elif flag & 0x10:  # X_IS_SAME_OR_POSITIVE_X_SHORT_VECTOR
                    dx = 0
                else:
                    dx = struct.unpack_from(">h", self.data, pos)[0]
                    pos += 2
                x += dx
                xs.append(x)

            ys = []
            y = 0
            for flag in flags:
                if flag & 0x04:  # Y_SHORT_VECTOR
                    dy = self.data[pos]
                    pos += 1
                    dy = dy if flag & 0x20 else -dy
                elif flag & 0x20:  # Y_IS_SAME_OR_POSITIVE_Y_SHORT_VECTOR
                    dy = 0
                else:
                    dy = struct.unpack_from(">h", self.data, pos)[0]
                    pos += 2
                y += dy
                ys.append(y)
        except (IndexError, struct.error):
            return []

        points = [(xs[i], ys[i], bool(flags[i] & 0x01)) for i in range(point_count)]
        contours = []
        first = 0
        for last in end_points:
            contours.append(points[first : last + 1])
            first = last + 1
        return contours

    def components(self, glyph_id: int):
        try:
            gid = int(glyph_id)
        except Exception:
            return []
        if gid < 0 or gid >= self.num_glyphs:
            return []
        start = self.glyf_off + self.loca[gid]
        end = self.glyf_off + self.loca[gid + 1]
        if end <= start or end - start < 10:
            return []
        number_of_contours = struct.unpack_from(">h", self.data, start)[0]
        if number_of_contours >= 0:
            return []
        pos = start + 10
        result = []
        while pos + 4 <= end:
            flags, comp_gid = struct.unpack_from(">HH", self.data, pos)
            pos += 4
            x = y = None
            if flags & self.ARG_1_AND_2_ARE_WORDS:
                if pos + 4 > end:
                    break
                a1, a2 = struct.unpack_from(">hh", self.data, pos)
                pos += 4
            else:
                if pos + 2 > end:
                    break
                a1, a2 = struct.unpack_from(">bb", self.data, pos)
                pos += 2
            if flags & self.ARGS_ARE_XY_VALUES:
                x, y = a1, a2
            if flags & self.WE_HAVE_A_SCALE:
                pos += 2
            elif flags & self.WE_HAVE_AN_X_AND_Y_SCALE:
                pos += 4
            elif flags & self.WE_HAVE_A_TWO_BY_TWO:
                pos += 8
            result.append({"gid": comp_gid, "x": x, "y": y, "flags": flags})
            if not (flags & self.MORE_COMPONENTS):
                break
        return result


def guess_main_and_zhuyin_component(font_name, components):
    if not components:
        return None, None, "無複合元件"
    font_text = str(font_name or "")
    # Combined annotated fonts in this corpus place the pronunciation component around x=+1024.
    shifted = [c for c in components if isinstance(c.get("x"), int) and abs(c["x"]) >= 500]
    unshifted = [c for c in components if c not in shifted]
    if len(shifted) == 1:
        zh = shifted[0]["gid"]
        main = unshifted[0]["gid"] if unshifted else None
        return main, zh, "依元件水平位移判定"
    # ZhuYinNREG in the sample PDF is an annotation-only font: one component, no base Han glyph.
    if len(components) == 1 and re.search(r"zhuyin", font_text, re.IGNORECASE):
        return None, components[0]["gid"], "單獨注音字型"
    return None, None, "待判定"


def extract(pdf_path: Path, output_path: Path) -> dict:
    doc = fitz.open(pdf_path)
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "摘要"
    ws_chars = wb.create_sheet("字元資料")
    ws_trace = wb.create_sheet("字形索引")
    ws_fonts = wb.create_sheet("頁面字型")
    ws_dist = wb.create_sheet("字體分布")
    ws_zh = wb.create_sheet("注音字型候選")
    ws_components = wb.create_sheet("複合字形元件")
    ws_zhcomp = wb.create_sheet("注音元件彙整")

    ws_chars.append([
        "實體頁碼", "頁面標籤", "區塊序號", "行序號", "span序號", "字元序號",
        "字元", "Unicode", "x0", "y0", "x1", "y1", "origin_x", "origin_y",
        "font", "size", "flags", "color", "alpha", "ascender", "descender",
        "line_dir_x", "line_dir_y", "wmode"
    ])
    ws_trace.append([
        "實體頁碼", "頁面標籤", "trace_span序號", "trace_seqno", "trace字元序號",
        "字元", "Unicode", "glyph_id_字形索引", "origin_x", "origin_y",
        "x0", "y0", "x1", "y1", "font", "size", "font_xref", "basefont",
        "資源名稱", "encoding", "flags", "wmode", "type", "layer", "dir_x", "dir_y",
        "複合元件數", "主字元件ID_推定", "注音元件ID_推定", "元件判定方式"
    ])
    ws_fonts.append([
        "實體頁碼", "頁面標籤", "font_xref", "副檔名", "字型類型", "basefont",
        "資源名稱", "encoding", "referencer"
    ])
    ws_dist.append(["font", "size", "字元數"])
    ws_zh.append([
        "font", "font_xref", "glyph_id_字形索引", "字元", "Unicode", "出現次數", "出現頁碼"
    ])
    ws_components.append([
        "實體頁碼", "頁面標籤", "字元", "Unicode", "font", "font_xref", "glyph_id_字形索引",
        "複合元件數", "component1_gid", "component1_x", "component1_y",
        "component2_gid", "component2_x", "component2_y",
        "component3_gid", "component3_x", "component3_y",
        "主字元件ID_推定", "注音元件ID_推定", "元件判定方式"
    ])
    ws_zhcomp.append([
        "注音元件ID_推定", "font", "font_xref", "字元例", "來源字形索引例", "出現次數", "出現頁碼", "判定方式"
    ])

    distribution = Counter()
    zh_counter = Counter()
    zh_pages = defaultdict(set)
    zh_component_counter = Counter()
    zh_component_pages = defaultdict(set)
    zh_component_chars = defaultdict(list)
    zh_component_glyphs = defaultdict(list)
    zh_component_methods = defaultdict(set)
    font_inspector_cache = {}
    char_count = 0
    trace_char_count = 0
    font_record_count = 0

    def inspect_components(font_xref_value, glyph_id_value):
        if not isinstance(font_xref_value, int):
            return []
        if font_xref_value not in font_inspector_cache:
            try:
                _name, _ext, _type, font_bytes = doc.extract_font(font_xref_value)
                font_inspector_cache[font_xref_value] = TrueTypeGlyphInspector(font_bytes)
            except Exception:
                font_inspector_cache[font_xref_value] = None
        inspector = font_inspector_cache.get(font_xref_value)
        if inspector is None:
            return []
        try:
            return inspector.components(glyph_id_value)
        except Exception:
            return []

    for page_no, page in enumerate(doc, start=1):
        label = page.get_label() or str(page_no)
        font_records = get_page_fonts(page)
        font_lookup = build_font_lookup(font_records)

        raw = page.get_text("rawdict") or {}
        for block_i, block in enumerate(raw.get("blocks", [])):
            if block.get("type", 0) != 0:
                continue
            for line_i, line in enumerate(block.get("lines", [])):
                direction = line.get("dir") or (None, None)
                wmode = line.get("wmode")
                for span_i, span in enumerate(line.get("spans", [])):
                    font = span.get("font")
                    size = safe_round(span.get("size"))
                    flags = span.get("flags")
                    color = span.get("color")
                    alpha = span.get("alpha")
                    ascender = safe_round(span.get("ascender"))
                    descender = safe_round(span.get("descender"))
                    for char_i, ch in enumerate(span.get("chars", [])):
                        c = safe_excel_text(ch.get("c", ""))
                        bbox = ch.get("bbox") or (None, None, None, None)
                        origin = ch.get("origin") or (None, None)
                        ws_chars.append([
                            page_no, label, block_i, line_i, span_i, char_i,
                            c, f"U+{ord(c):04X}" if len(c) == 1 else "",
                            safe_round(bbox[0]), safe_round(bbox[1]), safe_round(bbox[2]), safe_round(bbox[3]),
                            safe_round(origin[0]), safe_round(origin[1]),
                            font, size, flags, color, alpha, ascender, descender,
                            safe_round(direction[0]), safe_round(direction[1]), wmode,
                        ])
                        distribution[(str(font or ""), size)] += 1
                        char_count += 1

        try:
            traces = page.get_texttrace() or []
        except Exception:
            traces = []
        for trace_i, span in enumerate(traces):
            font = span.get("font")
            size = safe_round(span.get("size"))
            font_xref, basefont, resource, encoding = resolve_font(font_lookup, font)
            direction = span.get("dir") or (None, None)
            for char_i, ch in enumerate(span.get("chars", ())):
                # PyMuPDF texttrace char tuple: (unicode, glyph_id, origin, bbox)
                if len(ch) < 4:
                    continue
                ucs, glyph_id, origin, bbox = ch[:4]
                c = safe_char(ucs)
                components = inspect_components(font_xref, glyph_id)
                main_component, zh_component, component_method = guess_main_and_zhuyin_component(font, components)
                ws_trace.append([
                    page_no, label, trace_i, span.get("seqno"), char_i,
                    c, f"U+{int(ucs):04X}" if isinstance(ucs, int) else "", glyph_id,
                    safe_round(origin[0]), safe_round(origin[1]),
                    safe_round(bbox[0]), safe_round(bbox[1]), safe_round(bbox[2]), safe_round(bbox[3]),
                    font, size, font_xref, basefont, resource, encoding,
                    span.get("flags"), span.get("wmode"), span.get("type"), span.get("layer"),
                    safe_round(direction[0]), safe_round(direction[1]),
                    len(components), main_component, zh_component, component_method,
                ])
                trace_char_count += 1
                if components:
                    flat = []
                    for comp in components[:3]:
                        flat.extend([comp.get("gid"), comp.get("x"), comp.get("y")])
                    while len(flat) < 9:
                        flat.append(None)
                    ws_components.append([
                        page_no, label, c, f"U+{int(ucs):04X}" if isinstance(ucs, int) else "",
                        font, font_xref, glyph_id, len(components), *flat,
                        main_component, zh_component, component_method
                    ])
                if font and ZH_FONT_RE.search(str(font)):
                    key = (str(font), str(font_xref or ""), glyph_id, c, f"U+{int(ucs):04X}" if isinstance(ucs, int) else "")
                    zh_counter[key] += 1
                    zh_pages[key].add(page_no)
                    if zh_component is not None:
                        ckey = (zh_component, str(font), str(font_xref or ""))
                        zh_component_counter[ckey] += 1
                        zh_component_pages[ckey].add(page_no)
                        if c and c not in zh_component_chars[ckey] and len(zh_component_chars[ckey]) < 12:
                            zh_component_chars[ckey].append(c)
                        if glyph_id not in zh_component_glyphs[ckey] and len(zh_component_glyphs[ckey]) < 12:
                            zh_component_glyphs[ckey].append(glyph_id)
                        zh_component_methods[ckey].add(component_method)

        for font_rec in font_records:
            vals = list(font_rec) + [None] * 7
            xref, ext, ftype, basefont, name, encoding, referencer = vals[:7]
            ws_fonts.append([page_no, label, xref, ext, ftype, basefont, name, encoding, referencer])
            font_record_count += 1

    for (font, size), count in sorted(distribution.items(), key=lambda x: (-x[1], x[0][0], x[0][1] or 0)):
        ws_dist.append([font, size, count])

    for key, count in sorted(zh_counter.items(), key=lambda x: (x[0][0], str(x[0][1]), int(x[0][2]) if isinstance(x[0][2], int) else -1, x[0][3])):
        font, font_xref, glyph_id, c, unicode_text = key
        pages = ",".join(str(x) for x in sorted(zh_pages[key]))
        ws_zh.append([font, font_xref, glyph_id, c, unicode_text, count, pages])

    for key, count in sorted(zh_component_counter.items(), key=lambda x: (int(x[0][0]), x[0][1], x[0][2])):
        zh_component, font, font_xref = key
        ws_zhcomp.append([
            zh_component, font, font_xref,
            "、".join(zh_component_chars[key]),
            ",".join(str(x) for x in zh_component_glyphs[key]),
            count, ",".join(str(x) for x in sorted(zh_component_pages[key])),
            "；".join(sorted(zh_component_methods[key])),
        ])

    fitz_version = getattr(fitz, "VersionBind", None) or getattr(fitz, "__version__", "")
    summary_rows = [
        ("工具", f"{PROGRAM} v{VERSION}"),
        ("來源 PDF", str(pdf_path)),
        ("匯出時間", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("PDF 頁數", doc.page_count),
        ("rawdict 字元資料筆數", char_count),
        ("texttrace 字形索引筆數", trace_char_count),
        ("頁面字型記錄數", font_record_count),
        ("字體/尺寸組合數", len(distribution)),
        ("注音字型候選組合數", len(zh_counter)),
        ("推定注音元件組合數", len(zh_component_counter)),
        ("PyMuPDF 版本", str(fitz_version)),
        ("v0.4 新增", "直接解析內嵌 TrueType 複合字形，匯出 component glyph ID 與位移，並推定主字元件/注音元件；不需額外安裝 fontTools。"),
        ("用途", "只做診斷匯出，不修改原 PDF，也不執行注音正誤判斷。"),
    ]
    for k, v in summary_rows:
        ws_summary.append([k, v])
    ws_summary["A1"].font = Font(bold=True)
    ws_summary.column_dimensions["A"].width = 26
    ws_summary.column_dimensions["B"].width = 90
    ws_summary.sheet_view.showGridLines = False

    style_sheet(ws_chars)
    style_sheet(ws_trace)
    style_sheet(ws_fonts)
    style_sheet(ws_dist)
    style_sheet(ws_zh)
    style_sheet(ws_components)
    style_sheet(ws_zhcomp)

    wb.save(output_path)
    doc.close()
    return {
        "pages": summary_rows[3][1],
        "chars": char_count,
        "trace_chars": trace_char_count,
        "fonts": font_record_count,
        "output": output_path,
    }


def choose_pdf_gui() -> Path | None:
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        filename = filedialog.askopenfilename(title="選擇要診斷的 PDF", filetypes=[("PDF", "*.pdf")])
        root.destroy()
        return Path(filename) if filename else None
    except Exception:
        return None


def notify_gui(title: str, message: str, error: bool = False):
    try:
        import tkinter as tk
        from tkinter import messagebox
        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        if error:
            messagebox.showerror(title, message)
        else:
            messagebox.showinfo(title, message)
        root.destroy()
    except Exception:
        pass


def main() -> int:
    ap = argparse.ArgumentParser(description=f"{PROGRAM} v{VERSION}")
    ap.add_argument("pdf", nargs="?", help="來源 PDF；省略時會跳出檔案選擇視窗")
    ap.add_argument("-o", "--output", help="輸出 xlsx 路徑；省略時自動放在 PDF 同一資料夾")
    args = ap.parse_args()

    pdf_path = Path(args.pdf) if args.pdf else choose_pdf_gui()
    if not pdf_path:
        return 0
    if not pdf_path.exists() or pdf_path.suffix.lower() != ".pdf":
        msg = f"找不到有效 PDF：{pdf_path}"
        notify_gui(PROGRAM, msg, error=True)
        print(msg, file=sys.stderr)
        return 2

    output_path = Path(args.output) if args.output else unique_output_path(pdf_path)
    try:
        result = extract(pdf_path, output_path)
    except Exception as exc:
        msg = f"匯出失敗：{exc}"
        notify_gui(PROGRAM, msg, error=True)
        print(msg, file=sys.stderr)
        return 1

    msg = (
        f"PDF 文字診斷匯出處理結束（不等同全冊校對完成）。\n\n頁數：{result['pages']}\n"
        f"rawdict 字元資料：{result['chars']} 筆\n"
        f"texttrace 字形索引：{result['trace_chars']} 筆\n"
        f"字型記錄：{result['fonts']} 筆\n\n"
        f"輸出：\n{result['output']}"
    )
    print(msg)
    notify_gui(PROGRAM, msg)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
